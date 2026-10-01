"""Pre-scoring four-arm process/assurance fixtures for Phase 6."""
from __future__ import annotations

import json
import os
import shutil
import subprocess
from pathlib import Path

from read_engine_phase6 import benchmark as b5

HERE = Path(__file__).resolve().parent

CASES = {
    "module_identity": """import json, sys, os, __main__
print(json.dumps({"name":__name__,"file":__file__,"package":__package__,"spec":str(__spec__),"loader":type(__loader__).__name__,"argv":sys.argv,"path0":sys.path[0],"cwd":os.getcwd(),"main_same":__main__.__dict__ is globals()}, sort_keys=True))
""",
    "imports_and_cache": """import json, sys, __main__, local_helper
import local_helper as again
print(json.dumps({"local":local_helper.VALUE,"same":again is local_helper,"main_same":__main__.__dict__ is globals(),"extra_parent_modules":sorted(x for x in ("librecalc_agent.config","librecalc_agent._frozen.capture") if x in sys.modules)}, sort_keys=True))
""",
    "paths_environment": """import json, os, sys
print(json.dumps({"argv":sys.argv,"path":sys.path,"cwd":os.getcwd(),"pythonpath":os.environ.get("PYTHONPATH"),"context":bool(os.environ.get("READ_ENGINE_PHASE3_CONTEXT")),"fd1":os.fstat(1).st_mode}, sort_keys=True))
""",
    "output_flush": 'import sys\nsys.stdout.write("out")\nsys.stderr.write("err")\n',
    "normal_return": 'print("returned")\n',
    "system_exit_zero": 'import sys\nprint("zero")\nsys.exit(0)\n',
    "system_exit_nonzero": 'import sys\nprint("nonzero")\nsys.exit(7)\n',
    "uncaught_exception": 'raise ValueError("fixture boom")\n',
    "finally_and_file_flush": """try:
    with open("effect.txt", "w") as f:
        f.write("closed")
        raise RuntimeError("caught")
except RuntimeError:
    pass
finally:
    with open("finally.txt", "w") as f: f.write("finally")
""",
    "atexit_workbook_write": """import atexit, openpyxl
def write():
    wb=openpyxl.load_workbook("input.xlsx")
    wb.active["A1"]="atexit-write"
    wb.save("input.xlsx")
atexit.register(write)
""",
    "atexit_main_import": """import atexit
MARKER="script-main"
def check():
    import __main__
    print("main-marker="+str(getattr(__main__, "MARKER", None)))
atexit.register(check)
""",
    "atexit_output": """import atexit, sys
atexit.register(lambda: (sys.stdout.write("late-out\\n"),sys.stderr.write("late-err\\n")))
""",
    "body_main_import": """import __main__
print(__main__.__dict__ is globals())
""",
    "subprocess": """import subprocess, sys
p=subprocess.run([sys.executable,"-c","import os;print('child:'+str(os.environ.get('FIXTURE_KEY')))"],capture_output=True,text=True)
print(p.stdout.strip(),p.returncode)
""",
    "signal_catch": """import signal
seen=[]
signal.signal(signal.SIGUSR1, lambda *_: seen.append("caught"))
signal.raise_signal(signal.SIGUSR1)
print(seen)
""",
    "sigint_keyboard_interrupt": """import signal
signal.raise_signal(signal.SIGINT)
""",
    "inherited_extra_fd": """import os
print(os.read(int(os.environ["TEST_FD"]), 1))
""",
    "normal_exit_143": """import sys
sys.exit(143)
""",
    "abrupt_os_exit": """import openpyxl, os
wb=openpyxl.load_workbook("input.xlsx")
wb.active["A1"]="abrupt-write"
wb.save("input.xlsx")
os._exit(7)
""",
    "fatal_signal": """import openpyxl, os, signal
wb=openpyxl.load_workbook("input.xlsx")
wb.active["A1"]="signal-write"
wb.save("input.xlsx")
os.kill(os.getpid(),signal.SIGTERM)
""",
}


def _run(case: str, source: str, arm: str) -> dict:
    base = HERE / "runs/process_semantics" / case / arm.lower()
    work = base / "work"
    if work.exists():
        shutil.rmtree(work)
    work.mkdir(parents=True)
    (work / "workload.py").write_text(source)
    (work / "local_helper.py").write_text("VALUE = 17\n")
    from openpyxl import Workbook, load_workbook
    wb = Workbook()
    wb.active["A1"] = "before"
    wb.save(work / "input.xlsx")
    env = {k: v for k, v in os.environ.items() if k not in
           {"PYTHONPATH", "LIBRECALC_CONFIG", "LIBRECALC_RUN_CONTEXT", "READ_ENGINE_PHASE3_CONTEXT", "READ_ENGINE_PHASE6_CONTEXT"}
           and not k.startswith("CANDIDATE_A_")}
    env["XDG_CACHE_HOME"] = str(base / "xdg")
    env["PYTHONDONTWRITEBYTECODE"] = "1"
    env["FIXTURE_KEY"] = "same"
    venv = str(b5.VENV)
    inherited_fd = None
    if case == "inherited_extra_fd":
        inherited_fd = os.open(work / "probe_fd.txt", os.O_RDWR | os.O_CREAT | os.O_TRUNC, 0o600)
        os.write(inherited_fd, b"Z")
        os.lseek(inherited_fd, 0, os.SEEK_SET)
        env["TEST_FD"] = str(inherited_fd)
    if arm == "PY":
        argv = [venv, str(work / "workload.py")]
    elif arm == "H2":
        argv = [str(HERE / "observer"), venv, str(work / "workload.py"),
                str(work), str(base / "persistent-cache"), str(base / "runs"), str(b5.ROOT)]
    else:
        harness = b5.ROOT / ("read_engine_phase4/harness.py" if arm == "H0" else "read_engine_phase5/harness.py")
        argv = [venv, str(harness), "run", "--workdir", str(work),
                "--cache-root", str(base / "persistent-cache"),
                "--run-root", str(base / "runs"), str(work / "workload.py")]
    try:
        p = subprocess.run(argv, cwd=work, env=env, capture_output=True, timeout=30,
                           pass_fds=(inherited_fd,) if inherited_fd is not None else ())
        code, out, err = p.returncode, p.stdout, p.stderr
    except subprocess.TimeoutExpired as exc:
        code, out, err = None, exc.stdout or b"", exc.stderr or b""
    if inherited_fd is not None:
        os.close(inherited_fd)
    summary = capture = receipt = None
    if arm != "PY":
        last = base / "runs/last_run.json"
        if last.exists():
            rd = Path(json.loads(last.read_text())["run_dir"])
            if arm == "H2" and (rd / "observer_receipt.json").exists():
                receipt = json.loads((rd / "observer_receipt.json").read_text())
                summary = receipt
            elif (rd / "summary.json").exists():
                summary = json.loads((rd / "summary.json").read_text())
            if (rd / "capture.json").exists():
                capture = json.loads((rd / "capture.json").read_text())
    cell = load_workbook(work / "input.xlsx").active["A1"].value
    return {"arm": arm, "exit_code": code, "stdout": out.decode(errors="replace"),
            "stderr": err.decode(errors="replace"), "summary_exists": summary is not None,
            "capture_records": len(capture) if capture is not None else None,
            "observer_receipt": receipt,
            "workbook_A1": cell,
            "effect_txt": (work / "effect.txt").read_text() if (work / "effect.txt").exists() else None,
            "finally_txt": (work / "finally.txt").read_text() if (work / "finally.txt").exists() else None}


def main() -> None:
    b5.verify()
    output = HERE / "process_semantics.jsonl"
    if output.exists():
        raise RuntimeError("Process-semantics ledger already exists")
    with output.open("w") as f:
        for case, source in CASES.items():
            results = {arm: _run(case, source, arm) for arm in b5.ARMS}
            py, h0, h1, h2 = (results[a] for a in b5.ARMS)
            observation = {"exit_equal": py["exit_code"] == h2["exit_code"],
                           "stdout_equal": py["stdout"] == h2["stdout"],
                           "stderr_equal": py["stderr"] == h2["stderr"],
                           "file_effect_equal": py["workbook_A1"] == h2["workbook_A1"],
                           "h0_capture": h0["capture_records"], "h1_capture": h1["capture_records"],
                           "h2_capture": h2["capture_records"], "h2_receipt": h2["summary_exists"]}
            if case in {"abrupt_os_exit", "fatal_signal", "atexit_workbook_write"}:
                expected = {"abrupt_os_exit": "abrupt-write", "fatal_signal": "signal-write",
                            "atexit_workbook_write": "atexit-write"}[case]
                exact = (h2["exit_code"] == py["exit_code"] and h2["workbook_A1"] == expected
                         and h2["capture_records"] == 1 and h2["summary_exists"])
                classification = "EXACT" if exact else "PRODUCT_CONTRACT_VIOLATION"
            elif case == "module_identity":
                def comparable(result):
                    obj = json.loads(result["stdout"])
                    return {k: obj[k] for k in ("name", "package", "spec", "loader", "main_same")}
                classification = "EXACT" if (comparable(py) == comparable(h2)
                                             and h2["exit_code"] == py["exit_code"]) else "PRODUCT_CONTRACT_VIOLATION"
            elif case in {"paths_environment", "imports_and_cache"}:
                # Interposition necessarily adds runtime modules/context. Require
                # normal execution and record all exact values for review.
                classification = ("OBSERVABLE_DIFFERENCE_BUT_ACCEPTABLE"
                                  if h2["exit_code"] == py["exit_code"] else "PRODUCT_CONTRACT_VIOLATION")
            elif case == "uncaught_exception":
                def normalized_error(result):
                    return result["stderr"].replace(str(HERE / "runs/process_semantics" / case / result["arm"].lower() / "work"), "<WORK>")
                classification = ("EXACT" if h2["exit_code"] == py["exit_code"]
                                  and normalized_error(h2) == normalized_error(py)
                                  else "PRODUCT_CONTRACT_VIOLATION")
            elif case == "sigint_keyboard_interrupt":
                classification = ("EXACT" if h2["exit_code"] == py["exit_code"]
                                  and h2["summary_exists"] else "PRODUCT_CONTRACT_VIOLATION")
            elif py["exit_code"] == h2["exit_code"] and py["stdout"] == h2["stdout"] and py["stderr"] == h2["stderr"]:
                classification = "EXACT"
            else:
                classification = "PRODUCT_CONTRACT_VIOLATION"
            row = {"case": case, "classification": classification,
                   "observations": observation, "results": results}
            f.write(json.dumps(row, sort_keys=True) + "\n")
            f.flush()
            print(case, classification, flush=True)


if __name__ == "__main__":
    main()
