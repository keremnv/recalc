#!/usr/bin/env python3
"""Phase13 Q4 (Gate-A): precedent-closure + color-preservation + volatility audit.

For same-formula roots: diff output-vs-gold VALUES at precedent cells
(UPSTREAM_CONTENT vs CACHE_SUSPECT vs DYNAMIC_REFS).
For color roots: input-vs-output font compare on value-untouched cells
(TOOLING_DESTRUCTION vs MISSED_FIX).
For UNSCORABLE rows: volatile/external-function scan of assessed formulas.
Checkpointed, task-grouped. Usage: q4_closure.py [budget_runs]
"""
import json, os, gc, sys
sys.path.insert(0, 'phase13')
sys.path.insert(0, 'benchmark-data/SpreadsheetBench-2/evaluation')
from analyze_roots import refs_of
from evaluation import compare_font_color
import openpyxl

DATA = 'benchmark-data/SpreadsheetBench-2/data'
DS = {c: {d['id']: d for d in json.load(open(f'{DATA}/{c}/dataset.json'))}
      for c in ('Template', 'Debugging', 'Financial_Model')}
scores = {x['id']: x for x in json.load(open('phase13/working/popB_scores.json'))}
led = {}
for line in open('phase12r/BASELINE_CORRECTION_LEDGER.jsonl'):
    r = json.loads(line)
    if r.get('stratum') == 'P1':
        led[r['id']] = r.get('classification')
CKPT = 'phase13/working/popB_q4.json'
VOLATILE = ('TODAY(', 'NOW(', 'RAND(', 'RANDBETWEEN(', 'OFFSET(', 'INDIRECT(', 'CELL(', 'INFO(')

def main(budget):
    det = json.load(open('phase13/working/popB_fullerr.json'))
    roots = json.load(open('phase13/working/popB_roots.json'))
    cause = json.load(open('phase13/working/popB_rootcause.json'))
    out = json.load(open(CKPT)) if os.path.exists(CKPT) else {}
    bytask = {}
    for id in det:
        if id not in out:
            bytask.setdefault(det[id]['task'], []).append(id)
    done = 0
    for task in sorted(bytask):
        if done >= budget:
            break
        cat, tid = task.split(':')[0], task.split(':')[1]
        d0 = DS[cat][tid]
        try:
            wg_f = openpyxl.load_workbook(f"{DATA}/{cat}/{d0['golden_response_path']}", data_only=False)
            wg_v = openpyxl.load_workbook(f"{DATA}/{cat}/{d0['golden_response_path']}", data_only=True)
            wi_f = openpyxl.load_workbook(f"{DATA}/{cat}/{d0['spreadsheet_path']}", data_only=False)
        except Exception as e:
            for id in bytask[task]:
                out[id] = {'fatal': f'gold-load: {e}'[:80]}
            continue
        for id in bytask[task]:
            if done >= budget:
                break
            d = det[id]
            if d.get('fatal'):
                out[id] = {'fatal': d['fatal']}
                continue
            # skip runs needing no closure: no same-formula roots, no color
            # roots, recalc-fair (volatility scan only matters when UNSCORABLE)
            _c = cause.get(id, {})
            _has_color = any(m['cls'] == 'COLOR_ONLY' for m in d.get('mm', []))
            if (_c.get('formula_same', 1) == 0 and not _has_color
                    and led.get(id) == 'RECALC_NO_MATERIAL_EFFECT'):
                out[id] = {'task': task, 'upstream': 0, 'cache_suspect': 0,
                           'dynamic': 0, 'sampled_roots': 0, 'up_edited': 0,
                           'up_untouched': 0, 'color_missed': 0, 'color_tooling': 0,
                           'volatile_cells': [], 'n_volatile': 0, 'skipped': True}
                done += 1
                continue
            try:
                wo_f = openpyxl.load_workbook(scores[id]['V0'], data_only=False)
                wo_v = openpyxl.load_workbook(scores[id]['V0'], data_only=True)
            except Exception as e:
                out[id] = {'fatal': f'v0-load: {e}'[:80]}
                continue
            rs = [(s, c) for s, c, *_ in roots[id]['roots']]
            rset = set(rs)
            mm_by_cell = {(m['sheet'], m['cell']): m for m in d['mm']}
            upstream, cache_sus, dynamic = 0, 0, 0
            sampled_n = 0
            upstream_untouched, upstream_edited = 0, 0
            col_missed, col_tooling = 0, 0
            vol_cells = []
            for (sh, cell) in rs:
                m = mm_by_cell.get((sh, cell))
                if not m:
                    continue
                if m['cls'] == 'COLOR_ONLY':
                    # value-untouched + official font-color changed in/out?
                    # -> tooling destruction; else agent-side (missed color fix)
                    try:
                        iv = wi_f[sh][cell].value
                        ov = wo_f[sh][cell].value
                        same_color = compare_font_color(wi_f[sh][cell].font,
                                                        wo_f[sh][cell].font)
                        if ov == iv and not same_color:
                            col_tooling += 1
                        else:
                            col_missed += 1
                    except Exception:
                        col_missed += 1
                    continue
                if m['cls'] in ('MISSING_OUTPUT_NONE',) or m.get('mode') == 'FORMULA':
                    continue
                # same-formula check via stored out_f/ans_f? use live compare
                try:
                    of = wo_f[sh][cell].value
                    gf = wg_f[sh][cell].value
                except Exception:
                    continue
                if hasattr(of, 'text'):
                    of = of.text
                if hasattr(gf, 'text'):
                    gf = gf.text
                def norm(f):
                    if not isinstance(f, str) or not f.startswith('='):
                        return None
                    f = f.replace('$', '').upper()
                    return f
                if norm(of) is None or norm(of) != norm(gf):
                    continue  # differs/no-formula: handled elsewhere
                # volatility scan (UNSCORABLE rows)
                if led.get(id) != 'RECALC_NO_MATERIAL_EFFECT' and isinstance(of, str):
                    up = of.upper()
                    if any(v in up for v in VOLATILE) or '[' in of:
                        vol_cells.append(f'{sh}!{cell}')
                # precedent closure (bounded: full check <=2000 refs,
                # else corners + deterministic stride sample, flagged)
                prefs = refs_of(of if isinstance(of, str) else '', sh)
                sampled = False
                if len(prefs) > 2000:
                    prefs = sorted(prefs)
                    stride = max(1, len(prefs) // 1500)
                    prefs = set(prefs[::stride]) | set(prefs[:4]) | set(prefs[-4:])
                    sampled = True
                if not prefs and isinstance(of, str) and of.startswith('='):
                    # has formula but no parseable refs: constants-only or dynamic
                    import re as _re
                    if _re.search(r'[A-Z]+\d+', of.upper()):
                        dynamic += 1  # refs exist but unparsed (named/dynamic)
                        continue
                    else:
                        cache_sus += 1  # constants-only formula, same text, diff value
                        continue
                if not prefs:
                    continue
                diff_found, unassessed = False, False
                # refs_of casefolds sheets; resolve against real names
                smap = {s.casefold(): s for s in wo_v.sheetnames}
                gmap = {s.casefold(): s for s in wg_v.sheetnames}
                for (psh, pcell) in prefs:
                    try:
                        pv_o = wo_v[smap[psh]][pcell].value
                        pv_g = wg_v[gmap[psh]][pcell].value
                    except Exception:
                        unassessed = True
                        continue
                    if pv_o != pv_g:
                        # tolerance for floats
                        try:
                            if (isinstance(pv_o, (int, float)) and isinstance(pv_g, (int, float))
                                    and abs(pv_o - pv_g) <= 0.01 * max(abs(pv_o), abs(pv_g), 1e-9)):
                                continue
                        except Exception:
                            pass
                        diff_found = True
                        # edited or untouched by agent?
                        try:
                            _om = {s.casefold(): s for s in wo_f.sheetnames}
                            _im = {s.casefold(): s for s in wi_f.sheetnames}
                            if wo_f[_om[psh]][pcell].value != wi_f[_im[psh]][pcell].value:
                                upstream_edited += 1
                            else:
                                upstream_untouched += 1
                        except Exception:
                            upstream_untouched += 1
                        break
                if diff_found:
                    upstream += 1
                elif unassessed:
                    dynamic += 1
                else:
                    cache_sus += 1
                if sampled:
                    sampled_n += 1
            out[id] = {'task': task, 'upstream': upstream, 'cache_suspect': cache_sus,
                       'dynamic': dynamic, 'sampled_roots': sampled_n,
                       'up_edited': upstream_edited,
                       'up_untouched': upstream_untouched, 'color_missed': col_missed,
                       'color_tooling': col_tooling, 'volatile_cells': vol_cells[:10],
                       'n_volatile': len(vol_cells)}
            done += 1
            del wo_f, wo_v
        del wg_f, wg_v, wi_f
        gc.collect()
        json.dump(out, open(CKPT + '.tmp', 'w'))
        os.replace(CKPT + '.tmp', CKPT)
        print(f'{task}: ckpt {len(out)}', flush=True)
    pend = sum(1 for id in det if id not in out)
    print('pending:', pend)
    if pend == 0:
        vals = [v for v in out.values() if 'upstream' in v]
        print('same-formula roots: upstream=%d cache_suspect=%d dynamic=%d' % (
            sum(v['upstream'] for v in vals), sum(v['cache_suspect'] for v in vals),
            sum(v['dynamic'] for v in vals)))
        print('upstream edited/untouched:', sum(v['up_edited'] for v in vals),
              sum(v['up_untouched'] for v in vals))
        print('color missed/tooling:', sum(v['color_missed'] for v in vals),
              sum(v['color_tooling'] for v in vals))
        print('rows with volatile assessed cells:', sum(1 for v in vals if v['n_volatile']))

if __name__ == '__main__':
    main(int(sys.argv[1]) if len(sys.argv) > 1 else 10**9)
