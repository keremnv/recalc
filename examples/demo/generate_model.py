"""Deterministic generator for the LibreCalc demo workbook.

Every cell value is a pure function of (row, column), so each regeneration
produces the same sheet contents. Workbook properties use a fixed timestamp.
Byte-identical output across machines is not promised (ZIP member timestamps
and openpyxl versions may differ); value-identical output is.

Usage:
    python generate_model.py [output.xlsx]   # default: model.xlsx next to this file
"""

from __future__ import annotations

import datetime as dt
import sys
from pathlib import Path

import openpyxl

REGIONS = ["North", "South", "East"]
LEADS = {"North": "A. Rivera", "South": "B. Okafor", "East": "C. Novak"}
PRODUCTS = ["Widget", "Gadget", "Sprocket"]
HEADER = ["Entry", "Region", "Product", "Units", "Price", "Revenue", "Cost", "Margin"]
FIXED_TIME = dt.datetime(2025, 1, 6, 12, 0, 0, tzinfo=dt.UTC)
DATA_ROWS = 400  # rows 2..401; row 402 holds totals, so H402 is the margin total


def data_row(r: int) -> list:
    """Row r (2-based, matching the sheet). Pure function of r."""
    i = r - 2
    units = 100 + ((r * 37 + 11) % 400)
    price = 5 + ((r * 13 + 7) % 40)
    revenue = units * price
    cost = (revenue * 3) // 5
    return [
        f"E{i + 1:04d}",
        REGIONS[i % 3],
        PRODUCTS[(i // 3) % 3],
        units,
        price,
        revenue,
        cost,
        revenue - cost,
    ]


def build(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Forecast"
    ws.append(HEADER)
    for r in range(2, 2 + DATA_ROWS):
        ws.append(data_row(r))
    total_row = 2 + DATA_ROWS  # 402
    totals = ["TOTAL", "", ""]
    for col in range(4, 9):
        totals.append(sum(ws.cell(row=r, column=col).value for r in range(2, total_row)))
    ws.append(totals)

    regions = wb.create_sheet("Regions")
    regions.append(["Region", "Lead", "Target"])
    for name in REGIONS:
        regions.append([name, LEADS[name], 50000])

    props = wb.properties
    props.creator = "librecalc-demo"
    props.lastModifiedBy = "librecalc-demo"
    props.created = FIXED_TIME
    props.modified = FIXED_TIME

    wb.save(path)
    wb.close()
    print(f"wrote {path} ({path.stat().st_size} bytes)")


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).parent / "model.xlsx"
    build(out)
