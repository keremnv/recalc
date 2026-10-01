"""Focused pre-scoring process and assurance comparison for Phase 9."""
from __future__ import annotations

import json
import os
import shutil
import subprocess
from pathlib import Path

from read_engine_phase6 import benchmark as b6
from read_engine_phase6.process_probe import CASES as PHASE6_CASES
from read_engine_phase9 import benchmark as p9

HERE = Path(__file__).resolve().parent
CASES = {
    **PHASE6_CASES,
    "reference_object_identity": (
        "import openpyxl\n"
        "wb=openpyxl.load_workbook('input.xlsx')\n"
        "ws=wb.active\n"
        "print(type(wb).__module__,type(wb).__name__,"
        "type(ws).__module__,type(ws).__name__,"
        "type(ws['A1']).__module__,type(ws['A1']).__name__)\n"
    ),
    "reference_module_identity": (
        "import openpyxl, openpyxl.reader.excel\n"
        "wb=openpyxl.load_workbook('input.xlsx')\n"
        "assert wb.active.title == 'Sheet'\n"
        "print(openpyxl.load_workbook is openpyxl.reader.excel.load_workbook)\n"
    ),
}


def execute(case: str, source: str, arm: str) -> dict:
    base = HERE / "runs/process_semantics" / case / arm.lower()
    work = base / "work"
    work.mkdir(parents=True, exist_ok=True)
    (work / "workload.py").write_text(source)
    (work / "local_helper.py").write_text("VALUE = 17\n")
    from openpyxl import Workbook, load_workbook
    workbook = Workbook()
    workbook.active["A1"] = "before"
    workbook.save(work / "input.xlsx")
    env = {k: v for k, v in os.environ.items()
           if k not in {"PYTHONPATH", "LIBRECALC_CONFIG", "LIBRECALC_RUN_CONTEXT",
                        "READ_ENGINE_PHASE3_CONTEXT", "READ_ENGINE_PHASE6_CONTEXT"}
           and not k.startswith("CANDIDATE_A_")}
    env["XDG_CACHE_HOME"] = str(base / "xdg")
    env["PYTHONDONTWRITEBYTECODE"] = "1"
    env["FIXTURE_KEY"] = "same"
    fd = None
    if case == "inherited_extra_fd":
        fd = os.open(work / "probe_fd.txt", os.O_RDWR | os.O_CREAT | os.O_TRUNC, 0o600)
        os.write(fd, b"Z")
        os.lseek(fd, 0, os.SEEK_SET)
        env["TEST_FD"] = str(fd)
    python = str(b6.VENV)
    if arm == "PY":
        argv = [python, str(work / "workload.py")]
    else:
        root = (p9.ROOT / "read_engine_phase8a/overlay" if arm == "H0"
                else p9.OVERLAY)
        argv = [str(p9.ROOT / "read_engine_phase6/observer"), python,
                str(work / "workload.py"), str(work),
                str(base / "persistent-cache"), str(base / "runs"), str(root)]
    try:
        proc = subprocess.run(argv, cwd=work, env=env, capture_output=True,
                              timeout=30, pass_fds=(fd,) if fd is not None else ())
        code, out, err = proc.returncode, proc.stdout, proc.stderr
    except subprocess.TimeoutExpired as exc:
        code, out, err = None, exc.stdout or b"", exc.stderr or b""
    finally:
        if fd is not None:
            os.close(fd)
    receipt = None
    capture = None
    setup = None
    if arm != "PY":
        last = base / "runs/last_run.json"
        if last.exists():
            run = Path(json.loads(last.read_text())["run_dir"])
            if (run / "observer_receipt.json").exists():
                receipt = json.loads((run / "observer_receipt.json").read_text())
            if (run / "capture.json").exists():
                capture = json.loads((run / "capture.json").read_text())
            if (run / "setup.json").exists():
                setup = json.loads((run / "setup.json").read_text())
    cell = load_workbook(work / "input.xlsx").active["A1"].value
    return {"exit_code": code, "stdout": out.decode(errors="replace"),
            "stderr": err.decode(errors="replace"),
            "workbook_A1": cell,
            "capture_records": len(capture) if capture is not None else None,
            "observer_receipt": receipt,
            "setup": setup,
            "effect_txt": (work / "effect.txt").read_text()
                          if (work / "effect.txt").exists() else None,
            "finally_txt": (work / "finally.txt").read_text()
                           if (work / "finally.txt").exists() else None}


def normalized(text: str, case: str, arm: str) -> str:
    path = HERE / "runs/process_semantics" / case / arm.lower() / "work"
    return text.replace(str(path), "<WORK>")


def main() -> None:
    p9.verify()
    out = HERE / "process_semantics.jsonl"
    if out.exists() and out.stat().st_size:
        raise RuntimeError("Process-semantics ledger already contains rows")
    with out.open("w") as stream:
        for case, source in CASES.items():
            results = {arm: execute(case, source, arm) for arm in p9.ARMS}
            py, h0, h1 = (results[arm] for arm in p9.ARMS)
            output_equal = (
                normalized(py["stdout"], case, "PY")
                == normalized(h1["stdout"], case, "H1"))
            stderr_equal = (
                normalized(py["stderr"], case, "PY")
                == normalized(h1["stderr"], case, "H1"))
            effect_equal = all(
                py[key] == h1[key]
                for key in ("workbook_A1", "effect_txt", "finally_txt"))
            exit_equal = py["exit_code"] == h1["exit_code"]
            receipt_ok = h1["observer_receipt"] is not None
            if case in {"paths_environment", "imports_and_cache"}:
                # The injection context and module cache differ by design.
                passed = exit_equal and effect_equal and receipt_ok
                classification = ("OBSERVABLE_DIFFERENCE_BUT_ACCEPTABLE"
                                  if passed else "PRODUCT_CONTRACT_VIOLATION")
            elif case == "module_identity":
                def selected(x: dict) -> dict:
                    v = json.loads(x["stdout"])
                    return {k: v[k] for k in
                            ("name", "package", "spec", "loader", "main_same")}
                passed = exit_equal and selected(py) == selected(h1) and receipt_ok
                classification = ("EXACT" if passed else "PRODUCT_CONTRACT_VIOLATION")
            elif case == "reference_module_identity":
                passed = (exit_equal and h1["stdout"] == py["stdout"] and receipt_ok
                          and (h1["setup"] or {}).get("fast_path")
                          == "FAST_PATH_PROVEN_REFERENCE")
                classification = ("EXACT" if passed else "PRODUCT_CONTRACT_VIOLATION")
            elif case in {"abrupt_os_exit", "fatal_signal", "atexit_workbook_write"}:
                expected = {"abrupt_os_exit": "abrupt-write",
                            "fatal_signal": "signal-write",
                            "atexit_workbook_write": "atexit-write"}[case]
                passed = (exit_equal and h1["workbook_A1"] == expected
                          and h1["capture_records"] == 1 and receipt_ok
                          and h1["observer_receipt"]["changed_xlsx"] == 1
                          and h1["observer_receipt"]["capture_helper_exit"] == 0)
                classification = ("EXACT" if passed else "PRODUCT_CONTRACT_VIOLATION")
            else:
                passed = exit_equal and output_equal and stderr_equal and effect_equal and receipt_ok
                classification = ("EXACT" if passed else "PRODUCT_CONTRACT_VIOLATION")
            row = {"case": case, "classification": classification,
                   "passed": passed, "PY": py, "H0": h0, "H1": h1,
                   "H1_fast_path": (h1["setup"] or {}).get("fast_path")}
            stream.write(json.dumps(row, sort_keys=True, default=str) + "\n")
            stream.flush()
            print(case, classification, flush=True)
            if not passed:
                raise RuntimeError(f"Process fixture failed: {case}")


if __name__ == "__main__":
    main()
