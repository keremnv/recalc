"""Bounded post-scoring product install, cache, and assurance failure checks."""
from __future__ import annotations

import hashlib
import json
import os
import pathlib
import subprocess
import sys
import tempfile
from concurrent.futures import ThreadPoolExecutor

import openpyxl

from librecalc_agent import runner
from librecalc_agent.config import Config
from librecalc_agent.read_engine import artifact, cache

ROOT = pathlib.Path(__file__).resolve().parents[1]
HERE = pathlib.Path(__file__).resolve().parent
PYTHON = pathlib.Path("/tmp/librecalc-phase10-clean/bin/python")
COMMAND = pathlib.Path("/tmp/librecalc-phase10-clean/bin/librecalc-agent")


def save(name: str, value) -> None:
    (HERE / name).write_text(json.dumps(value, indent=2, sort_keys=True, default=str) + "\n")


def task(base: pathlib.Path, *, write: bool = False) -> tuple[pathlib.Path, pathlib.Path]:
    work = base / "work"
    work.mkdir(parents=True)
    book = openpyxl.Workbook()
    book.active.title = "Sheet1"
    book.active["A1"] = 21
    book.save(work / "input.xlsx")
    script = work / "task.py"
    script.write_text('import openpyxl\n'
                      'wb=openpyxl.load_workbook("input.xlsx")\n'
                      + ('wb["Sheet1"]["A1"]=22\nwb.save("input.xlsx")\n'
                         if write else
                         'ws=wb["Sheet1"]\nprint(ws.cell(row=1,column=1).value)\n'))
    return work, script


def command(work: pathlib.Path, script: pathlib.Path, base: pathlib.Path):
    env = {k: v for k, v in os.environ.items() if k not in {"PYTHONPATH", "LIBRECALC_CONFIG"}}
    env["XDG_CACHE_HOME"] = str(base / "xdg")
    proc = subprocess.run([str(COMMAND), "run", "--workdir", str(work), str(script)],
                          cwd=work, env=env, capture_output=True, timeout=30)
    pointer = base / "xdg/librecalc-agent/runs/last_run.json"
    receipt = {}
    if pointer.exists():
        path = pathlib.Path(json.loads(pointer.read_text())["run_dir"])
        receipt = json.loads((path / "observer_receipt.json").read_text())
        setup = json.loads((path / "setup.json").read_text()) if (path / "setup.json").exists() else {}
        runtime = json.loads((path / "runtime_state.json").read_text()) if (path / "runtime_state.json").exists() else {}
        receipt.update(route=setup.get("route"), artifacts=setup.get("artifacts"),
                       direct_served_loads=(runtime.get("counts") or {}).get("direct_served_loads", 0),
                       fallback_reasons=runtime.get("fallback_reasons", []))
    return {"exit_code": proc.returncode, "stdout": proc.stdout.decode(errors="replace").strip(),
            "stderr_head": proc.stderr.decode(errors="replace")[:250], "receipt": receipt}


def run_checks() -> None:
    install = {"wheel_installed": PYTHON.is_file() and COMMAND.is_file(),
               "command_executable": os.access(COMMAND, os.X_OK),
               "observer_in_wheel": (PYTHON.parent.parent / "lib/python3.13/site-packages/librecalc_agent/native/observer").is_file()}
    install["doctor"] = subprocess.run([str(COMMAND), "doctor", "--json"],
                                        capture_output=True, timeout=20).returncode
    with tempfile.TemporaryDirectory() as d:
        base = pathlib.Path(d)
        example = base / "example"
        install["example"] = subprocess.run([str(COMMAND), "example", str(example)],
                                             capture_output=True, timeout=20).returncode
        install["example_files"] = sorted(x.name for x in example.iterdir())
        work, script = task(base / "task")
        install["cold_run"] = command(work, script, base / "task")
        install["reused_run"] = command(work, script, base / "task")
        env = {**os.environ, "XDG_CACHE_HOME": str(base / "task/xdg")}
        status = subprocess.run([str(COMMAND), "status", "--json"],
                                env=env, capture_output=True, timeout=20)
        install["status"] = status.returncode
        install["status_route"] = json.loads(status.stdout).get("last_run", {}).get("route")
    install["pass"] = (install["wheel_installed"] and install["command_executable"]
                       and install["observer_in_wheel"] and install["doctor"] == 0
                       and install["example"] == 0 and install["status"] == 0
                       and install["cold_run"]["receipt"].get("direct_served_loads") == 1
                       and install["reused_run"]["receipt"].get("direct_served_loads") == 1
                       and install["status_route"] == "DIRECT_RUNTIME")
    save("install_results.json", install)

    failures = []
    with tempfile.TemporaryDirectory() as d:
        base = pathlib.Path(d)
        work, script = task(base)
        first = command(work, script, base)
        digest = cache.sha_file(work / "input.xlsx")
        path, sidecar = cache.paths(base / "xdg/librecalc-agent", digest)
        assert first["receipt"].get("direct_served_loads") == 1
        for name, mutation in (("corrupt_artifact", lambda: path.write_bytes(b"bad")),
                               ("truncated_artifact", lambda: path.write_bytes(path.read_bytes()[:20])),
                               ("missing_artifact", lambda: path.unlink())):
            mutation()
            result = command(work, script, base)
            status = [x.get("status") for x in (result["receipt"].get("artifacts") or {}).values()]
            failures.append({"case": name, "exit": result["exit_code"],
                             "artifact_status": status,
                             "direct_served": result["receipt"].get("direct_served_loads"),
                             "pass": result["exit_code"] == 0 and status == ["BUILT"]
                             and result["receipt"].get("direct_served_loads") == 1})
        book = openpyxl.load_workbook(work / "input.xlsx")
        book.active["A1"] = 22
        book.save(work / "input.xlsx")
        changed = command(work, script, base)
        status = [x.get("status") for x in (changed["receipt"].get("artifacts") or {}).values()]
        failures.append({"case": "stale_source", "artifact_status": status,
                         "stdout": changed["stdout"], "pass": status == ["BUILT"]
                         and changed["stdout"] == "22"})

    with tempfile.TemporaryDirectory() as d:
        base = pathlib.Path(d)
        work, script = task(base)
        for label, directory in (("cache_path_file", base / "cache-file"),
                                 ("permission_denied", pathlib.Path("/proc/librecalc-phase10-denied"))):
            if label == "cache_path_file":
                directory.write_text("not a directory")
            try:
                runner.run(script, [], work, Config(cache_dir=str(directory)), [])
                passed = False; error = None
            except runner.ProductError as exc:
                passed = True; error = str(exc)
            failures.append({"case": label, "pass": passed,
                             "script_not_run": not (work / "count").exists(), "error": error})
        original = runner.observer_binary
        runner.observer_binary = lambda: work / "missing-observer"
        try:
            try:
                runner.run(script, [], work, Config(cache_dir=str(base / "fresh-cache")), [])
                passed = False
            except runner.ProductError:
                passed = True
        finally:
            runner.observer_binary = original
        failures.append({"case": "missing_observer", "pass": passed})

    with tempfile.TemporaryDirectory() as d:
        base = pathlib.Path(d)
        work, script = task(base, write=True)
        cache_root, run_root = base / "cache", base / "runs"
        cache_root.mkdir(); run_root.mkdir()
        observer = runner.observer_binary()
        bootstrap = ROOT / "src/librecalc_agent/_bootstrap"
        bad_helper = base / "missing-helper.py"
        env = {**os.environ, "PYTHONPATH": str(ROOT / "src")}
        proc = subprocess.run([str(observer), str(PYTHON), str(script), str(work),
                               str(cache_root), str(run_root), str(bootstrap), str(bad_helper)],
                              cwd=work, env=env, capture_output=True, timeout=20)
        run_dir = pathlib.Path(json.loads((run_root / "last_run.json").read_text())["run_dir"])
        receipt = json.loads((run_dir / "observer_receipt.json").read_text())
        failures.append({"case": "capture_helper_failure", "exit": proc.returncode,
                         "assurance": receipt["assurance_status"],
                         "pass": proc.returncode == 125 and receipt["assurance_status"] == "FAILED"})

    with tempfile.TemporaryDirectory() as d:
        base = pathlib.Path(d)
        work, script = task(base)
        cache_root, run_root = base / "cache", base / "runs"
        cache_root.mkdir(); run_root.mkdir()
        proc = subprocess.run([str(runner.observer_binary()), str(base / "missing-python"),
                               str(script), str(work), str(cache_root), str(run_root),
                               str(ROOT / "src/librecalc_agent/_bootstrap"),
                               str(ROOT / "src/librecalc_agent/_capture_helper.py")],
                              cwd=work, capture_output=True, timeout=20)
        run_dir = pathlib.Path(json.loads((run_root / "last_run.json").read_text())["run_dir"])
        receipt = json.loads((run_dir / "observer_receipt.json").read_text())
        failures.append({"case": "observer_child_launch_failure", "exit": proc.returncode,
                         "target_exit": receipt["target_exit_code"],
                         "assurance": receipt["assurance_status"],
                         "pass": proc.returncode == 127 and receipt["target_exit_code"] == 127
                         and receipt["assurance_status"] == "PASS"})

    with tempfile.TemporaryDirectory() as d:
        base = pathlib.Path(d)
        work, script = task(base)
        (work / "input.xlsx").write_bytes(b"not an XLSX")
        py = subprocess.run([str(PYTHON), str(script)], cwd=work, capture_output=True)
        prod = command(work, script, base)
        failures.append({"case": "malformed_workbook_reference", "py_exit": py.returncode,
                         "prod_exit": prod["exit_code"],
                         "pass": py.returncode == prod["exit_code"] == 1
                         and prod["receipt"].get("direct_served_loads") == 0})
    save("failure_injection.json", failures)

    versions = []
    with tempfile.TemporaryDirectory() as d:
        base = pathlib.Path(d)
        work, _ = task(base)
        source = work / "input.xlsx"
        digest = cache.sha_file(source)
        root = base / "cache"
        original_path, status, _ = cache.ensure(source, root, digest)
        assert status == "BUILT"
        for field in ("RUNTIME_VERSION", "DECODER", "CONTRACT", "FORMAT"):
            old_cache = getattr(cache, field)
            old_artifact = getattr(artifact, field)
            setattr(cache, field, "NEXT_" + old_cache)
            setattr(artifact, field, "NEXT_" + old_artifact)
            try:
                new_path, new_status, _ = cache.ensure(source, root, digest)
                good = (new_path != original_path and new_status == "BUILT"
                        and cache.validate(new_path, pathlib.Path(str(new_path) + ".json"), digest)[0])
                versions.append({"version_field": field, "old_path": str(original_path),
                                 "new_path": str(new_path), "status": new_status, "pass": good})
            finally:
                setattr(cache, field, old_cache)
                setattr(artifact, field, old_artifact)
    save("upgrade_invalidation.json", versions)

    with tempfile.TemporaryDirectory() as d:
        base = pathlib.Path(d)
        work, script = task(base)
        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(lambda _: command(work, script, base), range(2)))
        digest = cache.sha_file(work / "input.xlsx")
        path, sidecar = cache.paths(base / "xdg/librecalc-agent", digest)
        concurrency = {"exit_codes": [x["exit_code"] for x in results],
                       "artifact_statuses": [[y.get("status") for y in
                                              (x["receipt"].get("artifacts") or {}).values()]
                                             for x in results],
                       "artifact_valid": cache.validate(path, sidecar, digest)[0]}
        concurrency["pass"] = (all(x == 0 for x in concurrency["exit_codes"])
                               and concurrency["artifact_valid"]
                               and all(x["receipt"].get("direct_served_loads") == 1 for x in results))
        save("concurrency_results.json", concurrency)

    print(json.dumps({"install": install["pass"],
                      "failures": [x["case"] for x in failures if not x["pass"]],
                      "versions": all(x["pass"] for x in versions),
                      "concurrency": concurrency["pass"]}, sort_keys=True))


if __name__ == "__main__":
    run_checks()
