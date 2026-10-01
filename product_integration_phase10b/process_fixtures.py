"""Unscored P0–P4 process and genuine-reference-object gate."""
from __future__ import annotations

import hashlib
import json
import os
import signal
import subprocess
import tempfile
from pathlib import Path

import openpyxl

from run_attribution import ARMS, HERE, clean_env, command


FIXTURES = {
    "module_argv_path": ('import os,sys,__main__\n'
                         'print(__name__,__main__.__file__==__file__,sys.argv[1:],'
                         'sys.path[0]==os.path.dirname(__file__),os.getcwd()==os.path.dirname(__file__))\n'),
    "stdout_stderr": 'import sys\nprint("out")\nprint("err",file=sys.stderr)\n',
    "system_exit": 'raise SystemExit(7)\n',
    "uncaught_exception": 'raise RuntimeError("expected")\n',
    "atexit": 'import atexit\natexit.register(lambda: print("atexit-visible"))\n',
    "local_import": 'import helper_local\nimport helper_local\nprint(helper_local.value)\n',
    "subprocess": ('import subprocess,sys\n'
                   'print(subprocess.check_output([sys.executable,"-c","print(123)"],text=True).strip())\n'),
    "inherited_fd": 'import os\nprint(os.read(int(os.environ["P10B_FD"]),8).decode())\n',
    "caught_sigint": ('import os,signal\n'
                      'signal.signal(signal.SIGINT,lambda *_: print("caught"))\n'
                      'os.kill(os.getpid(),signal.SIGINT)\n'),
    "fatal_sigterm": 'import os,signal\nos.kill(os.getpid(),signal.SIGTERM)\n',
    "abrupt_exit_after_write": ('import os,openpyxl\n'
                                'w=openpyxl.load_workbook("input.xlsx")\n'
                                'w.active["A1"]=77\nw.save("input.xlsx")\nos._exit(7)\n'),
    "atexit_write": ('import atexit,openpyxl\n'
                     'def finish():\n'
                     ' w=openpyxl.load_workbook("input.xlsx")\n'
                     ' w.active["A1"]=88\n'
                     ' w.save("input.xlsx")\n'
                     'atexit.register(finish)\n'),
    "reference_objects": ('import openpyxl\n'
                          'w=openpyxl.load_workbook("input.xlsx")\n'
                          's=w.active\nc=s["A1"]\n'
                          'print(type(w).__module__,type(w).__name__,'
                          'type(s).__module__,type(s).__name__,'
                          'type(c).__module__,type(c).__name__,'
                          'openpyxl.load_workbook.__module__,c.parent is s)\n'),
}


def fixture_run(name: str, source: str, arm: str, root: Path) -> dict:
    base = root / arm.lower()
    work = base / "work"
    work.mkdir(parents=True)
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = 1
    workbook.save(work / "input.xlsx")
    (work / "workload.py").write_text(source)
    (work / "helper_local.py").write_text("value = 42\n")
    env = clean_env(base)
    pass_fds = ()
    if name == "inherited_fd":
        rd, wr = os.pipe()
        os.write(wr, b"present")
        env["P10B_FD"] = str(rd)
        pass_fds = (rd,)
    try:
        proc = subprocess.run(command(arm, work, base), cwd=work, env=env,
                              capture_output=True, pass_fds=pass_fds, timeout=20)
    finally:
        if pass_fds:
            os.close(rd); os.close(wr)
    pointer = base / "xdg/librecalc-agent/runs/last_run.json"
    if pointer.exists():
        run_dir = Path(json.loads(pointer.read_text())["run_dir"])
        receipt = json.loads((run_dir / "observer_receipt.json").read_text())
        setup_path = run_dir / "setup.json"
        setup = json.loads(setup_path.read_text()) if setup_path.exists() else {}
    else:
        receipt, setup = {}, {}
    value = openpyxl.load_workbook(work / "input.xlsx").active["A1"].value
    return {"name": name, "arm": arm, "exit_code": proc.returncode,
            "stdout": proc.stdout.decode(errors="replace"),
            "stderr": proc.stderr.decode(errors="replace"),
            "workbook_value": value, "observer": receipt, "setup": setup}


def run_fixtures() -> None:
    output = HERE / "process_semantics.jsonl"
    if output.exists():
        raise RuntimeError("process fixture ledger exists")
    with tempfile.TemporaryDirectory(prefix="phase10b-fixtures-") as temp:
        root = Path(temp)
        for name, source in FIXTURES.items():
            results = {arm: fixture_run(name, source, arm, root / name)
                       for arm in ARMS}
            ref = results["P0"]
            for arm in ARMS:
                result = results[arm]
                exact = (result["exit_code"] == ref["exit_code"]
                         and result["stdout"] == ref["stdout"]
                         and result["workbook_value"] == ref["workbook_value"])
                if name != "uncaught_exception":
                    exact = exact and result["stderr"] == ref["stderr"]
                else:
                    exact = exact and "RuntimeError: expected" in result["stderr"]
                if arm != "P0":
                    exact = exact and result["observer"].get("assurance_status") == "PASS"
                if arm in {"P3", "P4"}:
                    exact = exact and result["setup"].get("route") == "REFERENCE_FAST_PATH"
                    exact = exact and result["setup"].get("artifacts") == {}
                row = {"fixture": name, "arm": arm, "exact": exact,
                       "exit_code": result["exit_code"],
                       "stdout_sha256": hashlib.sha256(result["stdout"].encode()).hexdigest(),
                       "stderr_sha256": hashlib.sha256(result["stderr"].encode()).hexdigest(),
                       "workbook_value": result["workbook_value"],
                       "observer_changed": result["observer"].get("changed_xlsx"),
                       "observer_capture_helper_exit": result["observer"].get("capture_helper_exit"),
                       "route": result["setup"].get("route")}
                with output.open("a") as stream:
                    stream.write(json.dumps(row, sort_keys=True) + "\n")
                if not exact:
                    raise RuntimeError(f"Process fixture failed: {name} {arm}: {result}")
            print(f"fixture {name}: P0–P4 exact", flush=True)


if __name__ == "__main__":
    run_fixtures()
