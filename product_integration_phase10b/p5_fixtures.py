"""Unscored P5 process, reference-object, and direct-route fixtures."""
from __future__ import annotations

import json
import os
import subprocess
import tempfile
from pathlib import Path

import openpyxl

from process_fixtures import FIXTURES
from run_attribution import CONSOLE, HERE, PYTHON, clean_env

P5_VENV = Path("/tmp/librecalc-phase10b-p5-clean")
P5_CONSOLE = P5_VENV / "bin/librecalc-agent"
P5_PYTHON = P5_VENV / "bin/python"
CASES = {**FIXTURES,
         "fatal_sigterm_after_write": ('import os,signal,openpyxl\n'
                                       'w=openpyxl.load_workbook("input.xlsx")\n'
                                       'w.active["A1"]=66\nw.save("input.xlsx")\n'
                                       'os.kill(os.getpid(),signal.SIGTERM)\n'),
         "direct_route_control": ('import openpyxl\n'
                                  'w=openpyxl.load_workbook("input.xlsx")\n'
                                  'print(w["Sheet"].cell(1,1).value)\n')}


def run_one(name: str, arm: str, root: Path) -> dict:
    base = root / name / arm.lower()
    work = base / "work"
    work.mkdir(parents=True)
    book = openpyxl.Workbook()
    book.active["A1"] = 1
    book.save(work / "input.xlsx")
    (work / "workload.py").write_text(CASES[name])
    (work / "helper_local.py").write_text("value = 42\n")
    env = clean_env(base)
    pass_fds = ()
    if name == "inherited_fd":
        rd, wr = os.pipe()
        os.write(wr, b"present")
        env["P10B_FD"] = str(rd)
        pass_fds = (rd,)
    python = PYTHON if arm in {"PY_OLD", "P4"} else P5_PYTHON
    if arm.startswith("PY_"):
        cmd = [str(python), str(work / "workload.py")]
    else:
        console = CONSOLE if arm == "P4" else P5_CONSOLE
        cmd = [str(console), "run", "--workdir", str(work), str(work / "workload.py")]
    try:
        proc = subprocess.run(cmd, cwd=work, env=env, capture_output=True,
                              pass_fds=pass_fds, timeout=20)
    finally:
        if pass_fds:
            os.close(rd); os.close(wr)
    pointer = base / "xdg/librecalc-agent/runs/last_run.json"
    receipt, setup = {}, {}
    if pointer.exists():
        run_dir = Path(json.loads(pointer.read_text())["run_dir"])
        receipt = json.loads((run_dir / "observer_receipt.json").read_text())
        path = run_dir / "setup.json"
        setup = json.loads(path.read_text()) if path.exists() else {}
    return {"exit": proc.returncode, "out": proc.stdout,
            "err": proc.stderr,
            "value": openpyxl.load_workbook(work / "input.xlsx").active["A1"].value,
            "receipt": receipt, "setup": setup}


def main() -> None:
    path = HERE / "p5_process_semantics.jsonl"
    if path.exists():
        raise RuntimeError("P5 process ledger exists")
    with tempfile.TemporaryDirectory(prefix="phase10b-p5-fixtures-") as d:
        root = Path(d)
        for name in CASES:
            values = {arm: run_one(name, arm, root)
                      for arm in ("PY_OLD", "PY_NEW", "P4", "P5")}
            old = values["PY_OLD"]
            for arm, value in values.items():
                good = (value["exit"] == old["exit"] and value["out"] == old["out"]
                        and value["value"] == old["value"])
                if name != "uncaught_exception":
                    good = good and value["err"] == old["err"]
                else:
                    good = good and b"RuntimeError: expected" in value["err"]
                if arm in {"P4", "P5"}:
                    good = good and value["receipt"].get("assurance_status") == "PASS"
                    if name in {"abrupt_exit_after_write", "fatal_sigterm_after_write",
                                "atexit_write"}:
                        good = (good and value["receipt"].get("changed_xlsx") == 1
                                and value["receipt"].get("capture_helper_exit") == 0)
                    if name == "direct_route_control":
                        good = good and value["setup"].get("route") == "DIRECT_RUNTIME"
                    else:
                        good = good and value["setup"].get("route") == "REFERENCE_FAST_PATH"
                row = {"fixture": name, "arm": arm, "valid": good,
                       "exit": value["exit"], "workbook_value": value["value"],
                       "route": value["setup"].get("route"),
                       "negative_proof": value["setup"].get("negative_proof"),
                       "changed_xlsx": value["receipt"].get("changed_xlsx"),
                       "capture_helper_exit": value["receipt"].get("capture_helper_exit")}
                with path.open("a") as stream:
                    stream.write(json.dumps(row, sort_keys=True) + "\n")
                if not good:
                    raise RuntimeError(f"P5 fixture failed: {name} {arm}: {value}")
            print(f"P5 fixture {name}: exact", flush=True)


if __name__ == "__main__":
    main()
