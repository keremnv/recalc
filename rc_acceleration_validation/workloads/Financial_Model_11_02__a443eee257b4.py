
import openpyxl
wb=openpyxl.load_workbook('input.xlsx')
wc=wb['Working Capital']
print('WC8 H-L',[wc.cell(8,c).value for c in range(8,13)])
print('WC30 H-M',[wc.cell(30,c).value for c in range(8,14)])
print('WC6 G',[wc.cell(6,7).value])
cf=wb['CF']
print('CF4',[cf.cell(4,c).value for c in range(15,35)])
print('CF5 O',[cf.cell(5,15).value, cf.cell(5,34).value])
a=wb['Assets Sch.']
print('A8 N-AG',[a.cell(8,c).value for c in range(14,34)])
