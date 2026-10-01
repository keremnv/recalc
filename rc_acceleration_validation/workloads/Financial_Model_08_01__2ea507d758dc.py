import openpyxl
wb = openpyxl.load_workbook('input.xlsx')
ws = wb['Income Statement']
from openpyxl.utils import get_column_letter
# print dates row 5 for I..CB
dates = []
for c in range(9, 81):
    v = ws.cell(5,c).value
    dates.append((get_column_letter(c), str(v)[:10] if v else None))
print(dates)
print("F33 company level:", wb['Assumptions - Company Level']['F33'].value)
print("C8 company level:", wb['Assumptions - Company Level']['C8'].value)
for r in [54,55,57,59,60,62,64,66,68,69]:
    vals = []
    for c in range(9, 81):
        v = ws.cell(r,c).value
        vals.append('X' if v is not None else '.')
    print(r, ws.cell(r,2).value, ''.join(vals))
