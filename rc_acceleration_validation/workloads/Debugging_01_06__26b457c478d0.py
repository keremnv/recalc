import openpyxl
wb=openpyxl.load_workbook('input.xlsx')
ws=wb['Ex 4 - Organic Operating Model']
for r in range(26,43):
    print(r,[ws.cell(row=r,column=c).value for c in [1,2,3,4,5,6,7,14,21,28]])
print('---Ex5---')
ws=wb['Ex 5 - M&A']
print(ws.dimensions)
for r in range(1,ws.max_row+1):
    vals=[(ws.cell(row=r,column=c).coordinate, ws.cell(row=r,column=c).value) for c in range(2,10) if ws.cell(row=r,column=c).value is not None]
    if vals: print(vals)
