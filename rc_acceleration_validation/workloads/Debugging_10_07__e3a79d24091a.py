
import openpyxl
wb = openpyxl.load_workbook('input.xlsx')
ws = wb['FY2024A (Earnings)']
for row in ws.iter_rows(min_row=1, max_row=12, max_col=20):
    for c in row:
        if c.value is not None:
            print(c.coordinate, repr(c.value)[:150])
print('---')
ws2 = wb['FY2027P (IS)']
for row in ws2.iter_rows(min_row=28, max_row=40, max_col=20):
    for c in row:
        if c.value is not None:
            print(c.coordinate, repr(c.value)[:150])
