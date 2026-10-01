
import openpyxl
wb=openpyxl.load_workbook('input.xlsx')
print(wb.sheetnames)
for ws in wb:
    print(ws.title, ws.dimensions, ws.max_row, ws.max_column)
