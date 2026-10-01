
import openpyxl
wb = openpyxl.load_workbook('input.xlsx')
ws = wb['Revenue & COGS Schedule']
for r in [9,10,11,13,14,15,16,17,19,20,21,22,23,24,25]:
    print('ROW', r, repr(ws.cell(row=r,column=2).value))
    for c in range(3,16):
        v = ws.cell(row=r,column=c).value
        if v is not None:
            print('  ', ws.cell(row=r,column=c).coordinate, repr(v))
