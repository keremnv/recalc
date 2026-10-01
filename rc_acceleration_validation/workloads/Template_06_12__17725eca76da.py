
import openpyxl
wb = openpyxl.load_workbook('input.xlsx')
ws = wb['CashFlow_Build']
print('dims', ws.dimensions, ws.max_row, ws.max_column)
for r in range(1, ws.max_row+1):
    for col in range(1, ws.max_column+1):
        c = ws.cell(r,col)
        if c.value is not None:
            print(c.coordinate, repr(c.value))
