"""Pre-registered offline construction, index-ready, and reuse timing."""
from __future__ import annotations

import contextlib
import gc
import hashlib
import json
import os
import random
import shutil
import statistics
import subprocess
import sys
import tempfile
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(HERE))

import openpyxl
from librecalc_agent._frozen import index, substrate
from librecalc_agent._frozen.reads import CandidateALoader, ProxyWorkbook
from librecalc_agent._frozen.runtime import PersistentCompiledSnapshot
from prototype import MemoryBook, SQLiteBook, build_r3, decode_xlsx
from run_correctness import _trace_ops, no_openpyxl_load, verify_frozen

VARIANTS = ("R0", "R1", "R2", "R3")
HORIZONS = (1, 2, 3, 5, 10)


def append(row: dict):
    with (HERE / "raw_timings.jsonl").open("a", encoding="utf-8") as h:
        h.write(json.dumps(row, sort_keys=True, default=str) + "\n")


def verified_stage(source: Path, temp: Path, digest: str) -> Path:
    path = temp / "input.xlsx"
    shutil.copyfile(source, path)
    if hashlib.sha256(path.read_bytes()).hexdigest() != digest:
        raise RuntimeError("Staged workbook differs from frozen bytes")
    return path


def _memory_bytes(root) -> int:
    seen = set()

    def visit(obj):
        if id(obj) in seen:
            return 0
        seen.add(id(obj))
        size = sys.getsizeof(obj)
        if isinstance(obj, dict):
            return size + sum(visit(k) + visit(v) for k, v in obj.items())
        if isinstance(obj, (list, tuple, set)):
            return size + sum(visit(x) for x in obj)
        if hasattr(obj, "__dict__") and obj.__class__.__module__.startswith("prototype"):
            return size + visit(vars(obj))
        return size

    return visit(root)


def cross_process_attach(path: Path, variant: str) -> bool:
    if variant == "R1":
        code = ("import sqlite3,sys; "
                "c=sqlite3.connect('file:'+sys.argv[1]+'?mode=ro',uri=True); "
                "assert c.execute('SELECT * FROM substrate_identity LIMIT 1').fetchone(); c.close()")
    else:
        code = ("import sys; sys.path.insert(0,sys.argv[2]); from prototype import SQLiteBook; "
                "b=SQLiteBook(sys.argv[1],{}); assert b.sheetnames; "
                "s=b[b.sheetnames[0]]; assert s.max_row>=1 and s.max_column>=1; "
                "s.cell(1,1).value; s.cell(1,1).data_type; b.close()")
    result = subprocess.run([sys.executable, "-c", code, str(path.resolve()), str(HERE)],
                            capture_output=True, timeout=20)
    return result.returncode == 0


class Ready:
    def __init__(self, variant: str, path: Path, workdir: Path):
        self.variant = variant
        self.path = path
        self.workdir = workdir
        self.book = None
        self.metadata = None
        self.db_path = None
        self.profile = {}
        self.artifact_bytes = None
        self.memory_bytes = None

    def build(self):
        t0 = time.perf_counter_ns()
        if self.variant == "R0":
            self.book = openpyxl.load_workbook(self.path, data_only=False)
        elif self.variant == "R1":
            index.reset()
            manifest = substrate.prepare(self.workdir, self.workdir / "current_index", "phase1", ("R1",))
            self.metadata = manifest["workbooks"].get(str(self.path.resolve()))
            if self.metadata is None:
                raise RuntimeError(f"R1 construction failed: {manifest['failures']}")
            self.db_path = Path(self.metadata["db_path"])
        elif self.variant == "R2":
            with no_openpyxl_load():
                self.book = decode_xlsx(self.path)
        elif self.variant == "R3":
            self.db_path = self.workdir / "minimal.sqlite"
            with no_openpyxl_load():
                book = build_r3(self.path, self.db_path)
            self.profile = book.profile
        else:
            raise ValueError(self.variant)
        elapsed = time.perf_counter_ns() - t0
        if self.variant == "R1":
            self.artifact_bytes = self.db_path.stat().st_size
            self.profile = {"ensure_s": self.metadata["ensure_s"], "backup_s": self.metadata["backup_s"],
                            "other_R1_components": "UNMEASURED"}
        elif self.variant == "R2":
            self.profile = self.book.profile
            self.memory_bytes = _memory_bytes(self.book)
        elif self.variant == "R3":
            self.artifact_bytes = self.db_path.stat().st_size
            book.close()
        return elapsed

    def acquire(self):
        if self.variant == "R0":
            return openpyxl.load_workbook(self.path, data_only=False)
        if self.variant == "R1":
            loader = CandidateALoader(real_loader=openpyxl.load_workbook)
            snapshot = PersistentCompiledSnapshot(str(self.path), self.metadata)
            return ProxyWorkbook(loader, snapshot)
        if self.variant == "R2":
            return self.book
        if self.variant == "R3":
            return SQLiteBook(self.db_path, {})
        raise ValueError(self.variant)

    def release(self, book):
        if self.variant in ("R0", "R3") and hasattr(book, "close"):
            book.close()
        elif self.variant == "R1" and hasattr(book, "close"):
            book.close()

    def close(self):
        if self.variant == "R0" and self.book is not None:
            self.book.close()
        if self.variant == "R1":
            index.reset()


def execute_one(ready: Ready, trace: dict) -> dict:
    ctx = no_openpyxl_load() if ready.variant in ("R2", "R3") else contextlib.nullcontext()
    with ctx:
        start = time.perf_counter_ns()
        book = ready.acquire()
        acquired = time.perf_counter_ns()
        values, error = _trace_ops(trace, book)
        operated = time.perf_counter_ns()
        ready.release(book)
        stopped = time.perf_counter_ns()
    return {"acquisition_ns": acquired - start, "operations_ns": operated - acquired,
            "release_ns": stopped - operated, "total_ns": stopped - start,
            "compared_values": len(values), "error": error}


def construction(pop):
    for rank, item in enumerate(pop["workbooks"]):
        repeated_order = random.Random(20260925 + rank)
        for repetition in range(3):
            order = list(VARIANTS)
            if repetition == 0:
                order = order[rank % 4:] + order[:rank % 4]
            else:
                repeated_order.shuffle(order)
            for order_index, variant in enumerate(order):
                with tempfile.TemporaryDirectory(prefix="read-engine-build-") as td:
                    temp = Path(td)
                    path = verified_stage(Path(item["path"]), temp, item["sha256"])
                    gc.collect()
                    ready = Ready(variant, path, temp)
                    try:
                        attempted = time.perf_counter_ns()
                        elapsed = ready.build()
                    except Exception as exc:
                        append({"phase": "construction", "workbook_sha256": item["sha256"],
                                "workbook_bytes": item["bytes"], "variant": variant,
                                "repetition": repetition, "order_index": order_index,
                                "cache_regime": "cold-ish first observation" if repetition == 0 else "repeated same OS-cache regime",
                                "status": "failed", "attempt_ns": time.perf_counter_ns() - attempted,
                                "exception_class": type(exc).__name__, "exception_message": str(exc)})
                    else:
                        cross_process = (cross_process_attach(ready.db_path, variant)
                                         if repetition == 0 and variant in ("R1", "R3") else None)
                        row = {"phase": "construction", "workbook_sha256": item["sha256"],
                               "workbook_bytes": item["bytes"], "variant": variant,
                               "repetition": repetition, "cache_regime": "cold-ish first observation" if repetition == 0 else "repeated same OS-cache regime",
                               "order_index": order_index, "status": "ok", "total_ns": elapsed,
                               "artifact_bytes": ready.artifact_bytes,
                               "artifact_source_ratio": ready.artifact_bytes / item["bytes"] if ready.artifact_bytes else None,
                               "approx_memory_bytes": ready.memory_bytes, "profile": ready.profile,
                               "cross_process_readonly_attach": cross_process}
                        append(row)
                    finally:
                        ready.close()
        print(f"construction {rank + 1}/39 {item['sha256'][:12]}", flush=True)


def _get_traces():
    return [json.loads(line) for line in (ROOT / "candidate_a_a1_checkpoint_rerun_01/contact_traces.jsonl").open()]


def index_ready(pop):
    traces = _get_traces()
    assert [x["trace_id"] for x in traces] == [x["trace_id"] for x in pop["traces"]]
    for rank, (meta, trace) in enumerate(zip(pop["traces"], traces)):
        with tempfile.TemporaryDirectory(prefix="read-engine-ready-") as td:
            temp = Path(td)
            path = verified_stage(Path(meta["snapshot_path"]), temp, meta["snapshot_sha256"])
            ready = {v: Ready(v, path, temp) for v in VARIANTS}
            for v in ("R1", "R2", "R3"):
                ready[v].build()
            try:
                for repetition in range(2):
                    order = list(VARIANTS)
                    random.Random(20260925 + rank + 10000 * repetition).shuffle(order)
                    for order_index, variant in enumerate(order):
                        result = execute_one(ready[variant], trace)
                        append({"phase": "index_ready", "trace_id": meta["trace_id"],
                                "snapshot_sha256": meta["snapshot_sha256"], "variant": variant,
                                "repetition": repetition, "order_index": order_index, **result})
            finally:
                for obj in ready.values():
                    obj.close()
        print(f"index-ready {rank + 1}/51 {meta['trace_id']}", flush=True)


def reuse(pop):
    traces = _get_traces()
    for rank, (meta, trace) in enumerate(zip(pop["traces"], traces)):
        order = list(VARIANTS)
        order = order[rank % 4:] + order[:rank % 4]
        for order_index, variant in enumerate(order):
            with tempfile.TemporaryDirectory(prefix="read-engine-reuse-") as td:
                temp = Path(td)
                path = verified_stage(Path(meta["snapshot_path"]), temp, meta["snapshot_sha256"])
                ready = Ready(variant, path, temp)
                build_ns = 0 if variant == "R0" else ready.build()
                cumulative = 0
                for n in range(1, 11):
                    one = execute_one(ready, trace)
                    cumulative += one["total_ns"]
                    if n in HORIZONS:
                        append({"phase": "reuse", "trace_id": meta["trace_id"],
                                "snapshot_sha256": meta["snapshot_sha256"], "variant": variant,
                                "order_index": order_index, "n": n,
                                "build_ns": build_ns, "read_cumulative_ns": cumulative,
                                "total_ns": build_ns + cumulative, "last_read": one,
                                "artifact_bytes": ready.artifact_bytes})
                ready.close()
        print(f"reuse {rank + 1}/51 {meta['trace_id']}", flush=True)


def summarize_profiles():
    rows = [json.loads(x) for x in (HERE / "raw_timings.jsonl").open() if x.strip()]
    result = {"spec_sha256": hashlib.sha256((HERE / "PREREGISTERED_SPEC_v5.md").read_bytes()).hexdigest(),
              "construction": {}, "notes": ["R1 ensure/backup are nested manifest timers; parse, traversal and insertion are UNMEASURED separately.",
                                             "R2/R3 aggregate stage timers are inside construction. Per-cell decode and insertion timers were removed in v4; those components are UNMEASURED individually."]}
    for variant in VARIANTS:
        c = [r for r in rows if r.get("phase") == "construction" and r["variant"] == variant]
        keys = {k for r in c for k, v in r.get("profile", {}).items() if isinstance(v, (int, float))}
        result["construction"][variant] = {k: {"median_s": statistics.median(r["profile"][k] for r in c if k in r.get("profile", {})),
                                                  "n": sum(k in r.get("profile", {}) for r in c)} for k in sorted(keys)}
    (HERE / "profile.json").write_text(json.dumps(result, indent=2, sort_keys=True) + "\n")


def main():
    pop = verify_frozen()
    rows = [json.loads(line) for line in (HERE / "raw_correctness.jsonl").open() if line.strip()]
    by_variant = {v: [x for x in rows if x.get("kind") == "trace" and x.get("variant") == v] for v in VARIANTS}
    if any(len(by_variant[v]) != 51 for v in VARIANTS):
        raise RuntimeError("51-trace correctness gate incomplete")
    if any(not x["exact"] for x in rows if x.get("kind") == "trace"):
        raise RuntimeError("Exact-trace correctness stop rule fired")
    book_summaries = [x for x in rows if x.get("kind") == "workbook_summary"]
    oracle_failures = [x for x in rows if x.get("kind") == "oracle_failure"]
    expected_books = {x["sha256"] for x in pop["workbooks"]}
    failed_books = {x["workbook_sha256"] for x in oracle_failures}
    if not failed_books <= expected_books or len(oracle_failures) != len(failed_books):
        raise RuntimeError("Invalid oracle-failure ledger")
    for variant in ("R2", "R3"):
        actual = [x["workbook_sha256"] for x in book_summaries if x["variant"] == variant]
        if len(actual) != 39 - len(failed_books) or set(actual) != expected_books - failed_books:
            raise RuntimeError(f"Full-cell correctness gate incomplete for {variant}")
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument("--phase", choices=("construction", "index_ready", "reuse", "profile"), required=True)
    args = parser.parse_args()
    if args.phase == "construction":
        construction(pop)
    elif args.phase == "index_ready":
        index_ready(pop)
    elif args.phase == "reuse":
        reuse(pop)
    else:
        summarize_profiles()


if __name__ == "__main__":
    main()
