"""Frozen-population correctness gate for the offline read prototype."""
from __future__ import annotations

import argparse
import contextlib
import hashlib
import json
import os
import shutil
import sqlite3
import sys
import tempfile
from collections import Counter
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(HERE))

import openpyxl
import openpyxl.reader.excel
from openpyxl.utils.cell import coordinate_to_tuple

from librecalc_agent._frozen import index, substrate
from librecalc_agent._frozen.reads import CandidateALoader, ProxyWorkbook
from librecalc_agent._frozen.runtime import PersistentCompiledSnapshot
from prototype import build_r3, canonical, decode_xlsx


def verify_frozen() -> dict:
    expected = (HERE / "PREREGISTERED_SPEC_v5.sha256").read_text().split()[0]
    actual = hashlib.sha256((HERE / "PREREGISTERED_SPEC_v5.md").read_bytes()).hexdigest()
    if actual != expected:
        raise RuntimeError("Preregistered spec changed")
    pop = json.loads((HERE / "population.json").read_text())
    for x in pop["workbooks"]:
        if hashlib.sha256(Path(x["path"]).read_bytes()).hexdigest() != x["sha256"]:
            raise RuntimeError("Workbook identity changed: " + x["path"])
    for x in pop["traces"]:
        if hashlib.sha256(Path(x["snapshot_path"]).read_bytes()).hexdigest() != x["snapshot_sha256"]:
            raise RuntimeError("Trace snapshot identity changed: " + x["trace_id"])
    return pop


@contextlib.contextmanager
def no_openpyxl_load():
    before1, before2 = openpyxl.load_workbook, openpyxl.reader.excel.load_workbook

    def forbidden(*args, **kwargs):
        raise AssertionError("Direct treatment called openpyxl.load_workbook")

    openpyxl.load_workbook = forbidden
    openpyxl.reader.excel.load_workbook = forbidden
    try:
        yield
    finally:
        openpyxl.load_workbook = before1
        openpyxl.reader.excel.load_workbook = before2


def add(path: Path, row: dict):
    with path.open("a", encoding="utf-8") as h:
        h.write(json.dumps(row, sort_keys=True, default=str) + "\n")


def _trace_ops(trace: dict, book: Any) -> tuple[list[dict], dict | None]:
    ws = cell = None
    values = []
    try:
        for event in trace["operations"]:
            op = event.get("operation")
            if op == "load_workbook":
                continue
            if op == "Workbook.sheetnames":
                value = book.sheetnames
            elif op == "Workbook.__getitem__":
                ws = book[event["sheet"]]
                continue
            elif op == "Worksheet.max_row":
                value = ws.max_row
            elif op == "Worksheet.max_column":
                value = ws.max_column
            elif op == "Worksheet.cell":
                row, col = coordinate_to_tuple(event["address"])
                cell = ws.cell(row=row, column=col)
                continue
            elif op == "Cell.value":
                value = cell.value
            elif op == "Cell.data_type":
                value = cell.data_type
            else:
                continue
            values.append({"operation": op, "value": canonical(value)})
        return values, None
    except Exception as exc:
        return values, {"class": type(exc).__name__, "message": str(exc)}


def _r1(path: Path, temp: Path):
    index.reset()
    manifest = substrate.prepare(temp, temp / "state", "phase1", ("R1",))
    entry = manifest["workbooks"].get(str(path.resolve()))
    if entry is None:
        raise RuntimeError(f"R1 build failed: {manifest['failures']}")
    original = openpyxl.load_workbook

    def acquire():
        loader = CandidateALoader(real_loader=original)
        snapshot = PersistentCompiledSnapshot(str(path), entry)
        return ProxyWorkbook(loader, snapshot)

    return acquire, manifest


def _trace_run(trace: dict, ledger: Path):
    source = Path(trace["snapshot_path"])
    with tempfile.TemporaryDirectory(prefix="read-engine-trace-") as td:
        temp = Path(td)
        path = temp / "input.xlsx"
        shutil.copyfile(source, path)
        if hashlib.sha256(path.read_bytes()).hexdigest() != trace["snapshot_sha256"]:
            raise RuntimeError("Staged snapshot hash mismatch")
        acquired = {}
        r0 = openpyxl.load_workbook(path, data_only=False)
        acquired["R0"] = r0
        r1_acquire, manifest = _r1(path, temp)
        acquired["R1"] = r1_acquire()
        with no_openpyxl_load():
            acquired["R2"] = decode_xlsx(path)
            acquired["R3"] = build_r3(path, temp / "minimal.sqlite")
        outputs = {}
        try:
            for variant, book in acquired.items():
                if variant in ("R2", "R3"):
                    with no_openpyxl_load():
                        result, error = _trace_ops(trace, book)
                else:
                    result, error = _trace_ops(trace, book)
                outputs[variant] = (result, error)
                control_result, control_error = outputs["R0"]
                exact = result == control_result and error == control_error
                row = {"kind": "trace", "trace_id": trace["trace_id"], "task_id": trace["task_id"],
                       "snapshot_sha256": trace["snapshot_sha256"], "variant": variant,
                       "operations": len(trace["operations"]), "compared_values": len(result),
                       "exact": exact, "error": error,
                       "first_mismatch": next(({"index": i, "reference": a, "treatment": b}
                                               for i, (a, b) in enumerate(zip(control_result, result)) if a != b), None),
                       "reference_value_count": len(control_result)}
                add(ledger, row)
        finally:
            for book in acquired.values():
                if hasattr(book, "close"):
                    book.close()
            index.reset()


def _category(value: Any, dtype: str) -> str:
    if dtype == "f":
        if value.__class__.__name__ in {"ArrayFormula", "DataTableFormula"}:
            return "array/shared formula"
        return "normal formula"
    if dtype == "d":
        return "date/time"
    if dtype == "e":
        return "error"
    if dtype == "b":
        return "boolean"
    if dtype == "n":
        return "number"
    return "shared/inline string"


def _book_run(item: dict, ledger: Path):
    source = Path(item["path"])
    with tempfile.TemporaryDirectory(prefix="read-engine-book-") as td:
        temp = Path(td)
        with no_openpyxl_load():
            r2 = decode_xlsx(source)
            r3 = build_r3(source, temp / "minimal.sqlite")
        try:
            r0 = openpyxl.load_workbook(source, data_only=False)
        except Exception as exc:
            add(ledger, {"kind": "oracle_failure", "workbook_sha256": item["sha256"],
                         "stage": "normal openpyxl.load_workbook", "exception_class": type(exc).__name__,
                         "exception_message": str(exc), "direct_construction_succeeded": True})
            r3.close()
            return
        try:
            for variant, treatment in (("R2", r2), ("R3", r3)):
                mismatches = Counter()
                checked = 0
                with no_openpyxl_load():
                    if treatment.sheetnames != r0.sheetnames:
                        add(ledger, {"kind": "mismatch", "variant": variant, "workbook_sha256": item["sha256"],
                                     "category": "sheetnames", "reference": r0.sheetnames, "treatment": treatment.sheetnames})
                        mismatches["sheetnames"] += 1
                    for ws in r0.worksheets:
                        tws = treatment[ws.title]
                        if (ws.max_row, ws.max_column, ws.dimensions) != (tws.max_row, tws.max_column, tws.dimensions):
                            add(ledger, {"kind": "mismatch", "variant": variant, "workbook_sha256": item["sha256"],
                                         "category": "bounds/dimension", "sheet": ws.title,
                                         "reference": [ws.max_row, ws.max_column, ws.dimensions],
                                         "treatment": [tws.max_row, tws.max_column, tws.dimensions]})
                            mismatches["bounds/dimension"] += 1
                        for cell in ws._cells.values():
                            if cell.value is None:
                                continue
                            checked += 1
                            actual = tws.cell(cell.row, cell.column)
                            got = (canonical(actual.value), actual.data_type)
                            want = (canonical(cell.value), cell.data_type)
                            if got != want:
                                category = _category(cell.value, cell.data_type)
                                mismatches[category] += 1
                                add(ledger, {"kind": "mismatch", "variant": variant,
                                             "workbook_sha256": item["sha256"], "sheet": ws.title,
                                             "cell": cell.coordinate, "category": category,
                                             "reference": want, "treatment": got})
                add(ledger, {"kind": "workbook_summary", "variant": variant,
                             "workbook_sha256": item["sha256"], "checked_nonempty_cells": checked,
                             "xml_nonempty_cells": item["xml_nonempty_cell_count"],
                             "mismatches": dict(mismatches), "exact": not mismatches})
        finally:
            r0.close()
            r3.close()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--stage", choices=("traces", "books", "all"), default="all")
    args = parser.parse_args()
    pop = verify_frozen()
    ledger = HERE / "raw_correctness.jsonl"
    trace_by_id = {row["trace_id"]: row for row in
                   (json.loads(line) for line in (ROOT / "candidate_a_a1_checkpoint_rerun_01/contact_traces.jsonl").open())}
    if args.stage in ("traces", "all"):
        if ledger.exists():
            raise RuntimeError("Correctness ledger already exists; preserve frozen run")
        for i, meta in enumerate(pop["traces"], 1):
            _trace_run(trace_by_id[meta["trace_id"]], ledger)
            print(f"trace {i}/51 {meta['trace_id']}", flush=True)
    if args.stage in ("books", "all"):
        for i, item in enumerate(pop["workbooks"], 1):
            _book_run(item, ledger)
            print(f"book {i}/39 {item['sha256'][:12]}", flush=True)


if __name__ == "__main__":
    main()
