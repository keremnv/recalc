"""Offline Phase-1 OOXML read prototype. No product imports or mutations.

R2 and R3 share decode_xlsx; only storage differs. Utilities imported from
openpyxl do not open workbooks. The runner guards load_workbook calls.
"""
from __future__ import annotations

import datetime as dt
import json
import math
import os
import re
import sqlite3
import tempfile
import time
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from lxml import etree
from openpyxl.formula.translate import Translator
from openpyxl.styles.numbers import BUILTIN_FORMATS, is_date_format, is_timedelta_format
from openpyxl.utils.datetime import CALENDAR_MAC_1904, WINDOWS_EPOCH, from_ISO8601, from_excel
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKGREL = "http://schemas.openxmlformats.org/package/2006/relationships"
M = "{" + MAIN + "}"
R = "{" + REL + "}"
P = "{" + PKGREL + "}"
_COORD = re.compile(r"^\$?([A-Za-z]+)\$?([1-9][0-9]*)$")


def coord_tuple(coord: str) -> tuple[int, int]:
    match = _COORD.fullmatch(coord)
    if not match:
        raise ValueError(f"{coord} is not a valid coordinate or range")
    col = 0
    for char in match.group(1).upper():
        col = col * 26 + ord(char) - 64
    return int(match.group(2)), col


def coord_name(row: int, col: int) -> str:
    if not isinstance(row, int) or not isinstance(col, int) or row < 1 or col < 1:
        raise ValueError("Row or column values must be at least 1")
    letters = ""
    while col:
        col, rem = divmod(col - 1, 26)
        letters = chr(65 + rem) + letters
    return f"{letters}{row}"


def range_bounds(ref: str) -> tuple[int, int, int, int]:
    first, sep, last = ref.partition(":")
    r1, c1 = coord_tuple(first)
    r2, c2 = coord_tuple(last if sep else first)
    return min(r1, r2), min(c1, c2), max(r1, r2), max(c1, c2)


def _target(target: str) -> str:
    if target.startswith("/"):
        return target.lstrip("/")
    if target.startswith("xl/"):
        return target
    return "xl/" + target.lstrip("/")


def _plain_string(si: etree._Element) -> str:
    # Shared-string <si> and inline <is> have direct <t> or <r><t> children.
    # Phonetic runs are intentionally not included, matching default openpyxl.
    parts: list[str] = []
    for child in si:
        if child.tag == M + "t":
            parts.append(child.text or "")
        elif child.tag == M + "r":
            text = child.find(M + "t")
            if text is not None:
                parts.append(text.text or "")
    return "".join(parts)


def _cast_number(raw: str) -> int | float:
    return float(raw) if any(c in raw for c in (".", "E", "e")) else int(raw)


def _value_to_json(value: Any) -> str:
    if isinstance(value, ArrayFormula):
        obj = {"kind": "array", "ref": value.ref, "text": value.text}
    elif isinstance(value, DataTableFormula):
        obj = {"kind": "datatable", "attrs": dict(value.__dict__)}
    elif isinstance(value, dt.datetime):
        obj = {"kind": "datetime", "value": value.isoformat()}
    elif isinstance(value, dt.date):
        obj = {"kind": "date", "value": value.isoformat()}
    elif isinstance(value, dt.time):
        obj = {"kind": "time", "value": value.isoformat()}
    elif isinstance(value, dt.timedelta):
        obj = {"kind": "timedelta", "seconds": value.total_seconds()}
    else:
        obj = {"kind": "scalar", "value": value}
    return json.dumps(obj, ensure_ascii=False, separators=(",", ":"))


def _value_from_json(raw: str) -> Any:
    obj = json.loads(raw)
    kind = obj["kind"]
    if kind == "array":
        return ArrayFormula(ref=obj["ref"], text=obj["text"])
    if kind == "datatable":
        return DataTableFormula(**obj["attrs"])
    if kind == "datetime":
        return dt.datetime.fromisoformat(obj["value"])
    if kind == "date":
        return dt.date.fromisoformat(obj["value"])
    if kind == "time":
        return dt.time.fromisoformat(obj["value"])
    if kind == "timedelta":
        return dt.timedelta(seconds=obj["seconds"])
    return obj["value"]


@dataclass
class SheetInfo:
    name: str
    min_row: int = 1
    min_col: int = 1
    max_row: int = 1
    max_col: int = 1
    merged: list[tuple[int, int, int, int]] = field(default_factory=list)
    cells: dict[str, tuple[Any, str]] = field(default_factory=dict)

    @property
    def dimensions(self) -> str:
        return f"{coord_name(self.min_row, self.min_col)}:{coord_name(self.max_row, self.max_col)}"


class MemoryBook:
    def __init__(self, sheets: list[SheetInfo], profile: dict[str, float] | None = None):
        self._sheets = {s.name: s for s in sheets}
        self.sheetnames = [s.name for s in sheets]
        self.profile = profile or {}

    def __getitem__(self, name: str):
        if name not in self._sheets:
            raise KeyError(name)
        return ReadSheet(self, self._sheets[name])

    def cell(self, sheet: SheetInfo, coord: str) -> tuple[Any, str]:
        r, c = coord_tuple(coord)
        for r1, c1, r2, c2 in sheet.merged:
            if r1 <= r <= r2 and c1 <= c <= c2 and (r, c) != (r1, c1):
                return None, "n"
        return sheet.cells.get(coord, (None, "n"))

    def close(self):
        pass


class ReadSheet:
    def __init__(self, book: Any, info: SheetInfo):
        self.book, self.info = book, info

    @property
    def max_row(self):
        return self.info.max_row

    @property
    def max_column(self):
        return self.info.max_col

    @property
    def dimensions(self):
        return self.info.dimensions

    def calculate_dimension(self):
        return self.dimensions

    def cell(self, row: int = 1, column: int = 1):
        return ReadCell(self.book, self.info, coord_name(row, column))

    def __getitem__(self, coord: str):
        if not isinstance(coord, str) or ":" in coord:
            raise ValueError("Only literal single-cell coordinates are in the frozen contract")
        coord_tuple(coord)
        return ReadCell(self.book, self.info, coord.upper().replace("$", ""))


class ReadCell:
    def __init__(self, book: Any, info: SheetInfo, coord: str):
        self.book, self.info, self.coordinate = book, info, coord
        self.row, self.column = coord_tuple(coord)

    @property
    def value(self):
        return self.book.cell(self.info, self.coordinate)[0]

    @property
    def data_type(self):
        return self.book.cell(self.info, self.coordinate)[1]


def _decode_cell(el, shared: list[str], date_ids: set[int], delta_ids: set[int],
                 epoch: Any, shared_formulae: dict[str, Translator]) -> tuple[Any, str]:
    dtype = el.get("t", "n")
    coord = el.get("r")
    style_id = int(el.get("s", "0"))
    formula = el.find(M + "f")
    val = el.find(M + "v")
    raw = val.text if val is not None and val.text else None
    if formula is not None:
        text = "=" + (formula.text or "")
        ftype = formula.get("t")
        if ftype == "array":
            return ArrayFormula(ref=formula.get("ref"), text=text), "f"
        if ftype == "shared":
            idx = formula.get("si")
            if idx in shared_formulae:
                text = shared_formulae[idx].translate_formula(coord)
            elif text != "=":
                shared_formulae[idx] = Translator(text, coord)
        if ftype == "dataTable":
            return DataTableFormula(**dict(formula.attrib)), "f"
        return text, "f"
    if raw is not None:
        if dtype == "n":
            value = _cast_number(raw)
            if style_id in date_ids:
                try:
                    return from_excel(value, epoch, timedelta=style_id in delta_ids), "d"
                except (OverflowError, ValueError):
                    return "#VALUE!", "e"
            return value, dtype
        if dtype == "s":
            return shared[int(raw)], dtype
        if dtype == "b":
            return bool(int(raw)), dtype
        if dtype == "str":
            return raw, "s"
        if dtype == "d":
            return from_ISO8601(raw), dtype
        return raw, dtype
    if dtype == "inlineStr":
        inline = el.find(M + "is")
        if inline is not None:
            return _plain_string(inline), "s"
    return None, dtype


def _load_style_ids(z: zipfile.ZipFile) -> tuple[set[int], set[int]]:
    if "xl/styles.xml" not in z.namelist():
        return set(), set()
    root = etree.fromstring(z.read("xl/styles.xml"))
    custom = {int(x.get("numFmtId")): x.get("formatCode", "")
              for x in root.findall(f"{M}numFmts/{M}numFmt")}
    dates, deltas = set(), set()
    for i, xf in enumerate(root.findall(f"{M}cellXfs/{M}xf")):
        nid = int(xf.get("numFmtId", "0"))
        fmt = custom.get(nid, BUILTIN_FORMATS.get(nid, "General"))
        if is_date_format(fmt):
            dates.add(i)
            if is_timedelta_format(fmt):
                deltas.add(i)
    return dates, deltas


def decode_xlsx(path: str | Path, sink: Any = None) -> MemoryBook:
    """Direct OOXML decoder. Sink receives (sheet index, coord, value, dtype)."""
    path = Path(path)
    profile: dict[str, float] = {}
    t_all = time.perf_counter_ns()
    with zipfile.ZipFile(path) as z:
        t = time.perf_counter_ns()
        names = set(z.namelist())
        profile["zip_open_list_s"] = (time.perf_counter_ns() - t) / 1e9
        t = time.perf_counter_ns()
        wb = etree.fromstring(z.read("xl/workbook.xml"))
        relroot = etree.fromstring(z.read("xl/_rels/workbook.xml.rels"))
        rels = {x.get("Id"): _target(x.get("Target", "")) for x in relroot.findall(P + "Relationship")}
        epoch = CALENDAR_MAC_1904 if wb.find(M + "workbookPr") is not None and wb.find(M + "workbookPr").get("date1904") in ("1", "true", "True") else WINDOWS_EPOCH
        sheet_paths = [(x.get("name"), rels[x.get(R + "id")]) for x in wb.findall(f"{M}sheets/{M}sheet")]
        profile["workbook_relationships_s"] = (time.perf_counter_ns() - t) / 1e9
        t = time.perf_counter_ns()
        shared = []
        if "xl/sharedStrings.xml" in names:
            root = etree.fromstring(z.read("xl/sharedStrings.xml"))
            shared = [_plain_string(si) for si in root.findall(M + "si")]
        profile["shared_strings_s"] = (time.perf_counter_ns() - t) / 1e9
        t = time.perf_counter_ns()
        date_ids, delta_ids = _load_style_ids(z)
        profile["styles_dates_s"] = (time.perf_counter_ns() - t) / 1e9
        sheets: list[SheetInfo] = []
        xml_start = time.perf_counter_ns()
        for sheet_idx, (name, xml_path) in enumerate(sheet_paths):
            if xml_path not in names:
                raise KeyError(f"Missing worksheet part {xml_path}")
            info = SheetInfo(name)
            minr = minc = 2**31 - 1
            maxr = maxc = 0
            formulas: dict[str, Translator] = {}
            with z.open(xml_path) as stream:
                for _, el in etree.iterparse(stream, events=("end",), tag=(M + "c", M + "mergeCell")):
                    if el.tag == M + "c":
                        coord = el.get("r")
                        if coord:
                            row, col = coord_tuple(coord)
                            minr, minc = min(minr, row), min(minc, col)
                            maxr, maxc = max(maxr, row), max(maxc, col)
                            value, dtype = _decode_cell(el, shared, date_ids, delta_ids, epoch, formulas)
                            if value is not None or dtype != "n":
                                if sink is None:
                                    info.cells[coord] = (value, dtype)
                                else:
                                    sink.put(sheet_idx, coord, value, dtype)
                    else:
                        ref = el.get("ref")
                        if ref:
                            b = range_bounds(ref)
                            info.merged.append(b)
                            minr, minc = min(minr, b[0]), min(minc, b[1])
                            maxr, maxc = max(maxr, b[2]), max(maxc, b[3])
                    el.clear()
                    while el.getprevious() is not None:
                        del el.getparent()[0]
            if maxr:
                info.min_row, info.min_col, info.max_row, info.max_col = minr, minc, maxr, maxc
            sheets.append(info)
        profile["worksheet_xml_total_s"] = (time.perf_counter_ns() - xml_start) / 1e9
        profile["value_decode_s"] = "UNMEASURED"
        profile["structure_insert_s"] = "UNMEASURED"
    profile["total_decode_s"] = (time.perf_counter_ns() - t_all) / 1e9
    return MemoryBook(sheets, profile)


class SQLiteSink:
    def __init__(self, path: Path):
        self.path = path
        self.con = sqlite3.connect(str(path))
        self.con.execute("CREATE TABLE cells(sheet INTEGER NOT NULL, coord TEXT NOT NULL, payload TEXT NOT NULL, dtype TEXT NOT NULL, PRIMARY KEY(sheet,coord)) WITHOUT ROWID")
        self.con.execute("CREATE TABLE sheets(sheet INTEGER PRIMARY KEY, name TEXT NOT NULL UNIQUE, min_row INTEGER NOT NULL, min_col INTEGER NOT NULL, max_row INTEGER NOT NULL, max_col INTEGER NOT NULL, merged_json TEXT NOT NULL)")

    def put(self, sheet: int, coord: str, value: Any, dtype: str):
        self.con.execute("INSERT INTO cells VALUES (?,?,?,?)", (sheet, coord, _value_to_json(value), dtype))

    def put_sheets(self, sheets: list[SheetInfo]):
        self.con.executemany(
            "INSERT INTO sheets VALUES (?,?,?,?,?,?,?)",
            ((i, s.name, s.min_row, s.min_col, s.max_row, s.max_col,
              json.dumps(s.merged, separators=(",", ":"))) for i, s in enumerate(sheets)),
        )

    def finish(self):
        self.con.commit()
        self.con.close()


class SQLiteBook(MemoryBook):
    def __init__(self, path: Path, profile: dict[str, float]):
        self.path = Path(path)
        self.con = sqlite3.connect(self.path.resolve().as_uri() + "?mode=ro", uri=True)
        rows = self.con.execute(
            "SELECT sheet,name,min_row,min_col,max_row,max_col,merged_json FROM sheets ORDER BY sheet"
        ).fetchall()
        sheets = []
        for sheet_id, name, min_row, min_col, max_row, max_col, merged_json in rows:
            if sheet_id != len(sheets):
                raise ValueError("Noncontiguous sheet identity in persistent artifact")
            sheets.append(SheetInfo(name, min_row, min_col, max_row, max_col,
                                    [tuple(x) for x in json.loads(merged_json)]))
        super().__init__(sheets, profile)
        self.ids = {name: i for i, name in enumerate(self.sheetnames)}

    def cell(self, sheet: SheetInfo, coord: str):
        r, c = coord_tuple(coord)
        for r1, c1, r2, c2 in sheet.merged:
            if r1 <= r <= r2 and c1 <= c <= c2 and (r, c) != (r1, c1):
                return None, "n"
        row = self.con.execute("SELECT payload,dtype FROM cells WHERE sheet=? AND coord=?", (self.ids[sheet.name], coord)).fetchone()
        return (_value_from_json(row[0]), row[1]) if row else (None, "n")

    def close(self):
        self.con.close()


def build_r3(source: str | Path, destination: str | Path) -> SQLiteBook:
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(prefix="ooxml-", suffix=".sqlite.tmp", dir=destination.parent)
    os.close(fd)
    tmp_path = Path(tmp)
    sink = SQLiteSink(tmp_path)
    try:
        book = decode_xlsx(source, sink)
        t = time.perf_counter_ns()
        sink.put_sheets(list(book._sheets.values()))
        book.profile["sheet_metadata_insert_s"] = (time.perf_counter_ns() - t) / 1e9
        t = time.perf_counter_ns()
        sink.finish()
        book.profile["sqlite_commit_s"] = (time.perf_counter_ns() - t) / 1e9
        t = time.perf_counter_ns()
        os.replace(tmp_path, destination)
        book.profile["publication_s"] = (time.perf_counter_ns() - t) / 1e9
        book.profile["total_build_s"] = (book.profile["total_decode_s"] + book.profile["sheet_metadata_insert_s"]
                                         + book.profile["sqlite_commit_s"] + book.profile["publication_s"])
        return SQLiteBook(destination, book.profile)
    except Exception:
        try:
            sink.con.close()
        except Exception:
            pass
        tmp_path.unlink(missing_ok=True)
        raise


def canonical(value: Any) -> Any:
    if isinstance(value, (list, tuple)):
        return {"type": type(value).__name__, "items": [canonical(v) for v in value]}
    if isinstance(value, ArrayFormula):
        return {"type": "ArrayFormula", "ref": value.ref, "text": value.text}
    if isinstance(value, DataTableFormula):
        return {"type": "DataTableFormula", "attrs": dict(value.__dict__)}
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return {"type": type(value).__name__, "value": value.isoformat()}
    if isinstance(value, dt.timedelta):
        return {"type": "timedelta", "seconds": value.total_seconds()}
    if isinstance(value, float) and math.isnan(value):
        return {"type": "float", "value": "NaN"}
    if value is None or isinstance(value, (str, int, float, bool)):
        return {"type": type(value).__name__, "value": value}
    return {"type": type(value).__name__, "repr": repr(value)}
