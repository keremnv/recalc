"""0.2.0 release gates: corrupt | parity | advers. Read-only verification."""
import hashlib
import importlib.util
import json
import os
import shutil
import struct
import subprocess
import sys
import zlib
from pathlib import Path

HERE = Path(__file__).parent
ROOT = HERE.parents[2] if HERE.name == "0.2.0" else HERE.parents[1]
R1 = ROOT / "research" / "execution_surface_census"
CLEAN = R1 / "_staging" / "staged" / "clean"
STAGE = Path("/tmp/rel02_staging")

sys.path.insert(0, str(ROOT / "src"))


def load_pristine():
    if "rc5ctrl.read_engine.artifact" in sys.modules:
        return (sys.modules["rc5ctrl.read_engine.artifact"],
                sys.modules["rc5ctrl.read_engine._identity"])
    dest = STAGE / "pristine"
    if not (dest / "src" / "recalc_agent" / "read_engine" / "artifact.py").exists():
        dest.mkdir(parents=True, exist_ok=True)
        arc = subprocess.run(["git", "archive", "master", "src/recalc_agent"],
                             cwd=ROOT, capture_output=True, check=True).stdout
        subprocess.run(["tar", "-x", "-C", str(dest)], input=arc, check=True)
    root = dest / "src" / "recalc_agent"

    def _load(name, path, is_pkg=False):
        spec = importlib.util.spec_from_file_location(
            name, path,
            submodule_search_locations=([str(path.parent)] if is_pkg else None))
        mod = importlib.util.module_from_spec(spec)
        sys.modules[name] = mod
        spec.loader.exec_module(mod)
        return mod

    _load("rc5ctrl", root / "__init__.py", is_pkg=True)
    _load("rc5ctrl.read_engine", root / "read_engine" / "__init__.py", is_pkg=True)
    _load("rc5ctrl.read_engine._identity", root / "read_engine" / "_identity.py")
    _load("rc5ctrl.read_engine.direct", root / "read_engine" / "direct.py")
    return (_load("rc5ctrl.read_engine.artifact", root / "read_engine" / "artifact.py"),
            sys.modules["rc5ctrl.read_engine._identity"])


def canon_value(v):
    import datetime as _dt
    tn = type(v).__name__
    if tn == "ArrayFormula":
        return ("ArrayFormula", v.ref, v.text)
    if tn == "DataTableFormula":
        return ("DataTableFormula", tuple(sorted((k, repr(x)) for k, x in v.__dict__.items())))
    if isinstance(v, (_dt.datetime, _dt.date, _dt.time)):
        return (tn, v.isoformat())
    if isinstance(v, _dt.timedelta):
        return ("timedelta", v.total_seconds())
    return (tn, repr(v))


def canon_book(book):
    return [(n, (s.min_row, s.min_col, s.max_row, s.max_col), tuple(s.merged),
             tuple(sorted((c, dt, canon_value(v)) for c, (v, dt) in s.cells.items())))
            for n in book.sheetnames for s in [book._sheets[n]]]


def rebuild_app(header, raw, ident_mod):
    from recalc_agent.read_engine.artifact import canonical, sha_bytes
    comp = zlib.compress(raw, level=1)
    header = dict(header)
    header["payload_sha256"] = sha_bytes(raw)
    h = canonical(header)
    body = ident_mod.MAGIC + struct.pack(">I", len(h)) + h \
        + struct.pack(">Q", len(comp)) + comp
    return body + hashlib.sha256(body).digest()


def cmd_corrupt():
    from recalc_agent.read_engine import artifact as art
    from recalc_agent.read_engine import _identity as ident
    assert ident.FORMAT == "JSONZ_MEMORY_V1", ident.FORMAT
    wid = "Template_03_03__c76ea596b408"
    book = sorted((CLEAN / "A" / wid).glob("*.xlsx"))[0]
    digest = ident.sha_file(book)
    apath, _, _ = art.ensure(book, STAGE / "ccache", digest)
    data = apath.read_bytes()
    hlen = struct.unpack(">I", data[8:12])[0]
    header = json.loads(data[12:12 + hlen])
    off = 12 + hlen
    clen = struct.unpack(">Q", data[off:off + 8])[0]
    raw = zlib.decompress(data[off + 8:off + 8 + clen])

    def payload(mut, allow_nan=False):
        obj = json.loads(raw)
        mut(obj)
        raw2 = json.dumps(obj, allow_nan=allow_nan).encode()
        if allow_nan:
            comp = zlib.compress(raw2, level=1)
            h = dict(header)
            h["payload_sha256"] = hashlib.sha256(raw2).hexdigest()
            hb = json.dumps(h, sort_keys=True, separators=(",", ":")).encode()
            body = ident.MAGIC + struct.pack(">I", len(hb)) + hb \
                + struct.pack(">Q", len(comp)) + comp
            return body + hashlib.sha256(body).digest()
        return rebuild_app(header, raw2, ident)

    cases = {
        "truncated_header": data[:20],
        "truncated_payload": data[:-40],
        "wrong_checksum": bytes(bytearray((b ^ 0xFF) if i == len(data) - 1 else b for i, b in enumerate(data))),
        "wrong_source_hash": ("__DIGEST__", "0" * 64),
        "wrong_decoder_version": ("__HEADERFIELD__", ("decoder_sha256", "0" * 64)),
        "wrong_format_version": ("__HEADERFIELD__", ("format_version", "NOPE_V9")),
        "corrupt_compressed_stream": bytes(bytearray((b ^ 0xFF) if i == off + 10 else b for i, b in enumerate(data))),
        "corrupt_serialized_structure": rebuild_app(header, b"{not json", ident),
        "duplicate_coordinates": payload(lambda o: o["sheets"][0]["cells"].append(o["sheets"][0]["cells"][0])),
        "nan_typed_value": payload(lambda o: o["sheets"][0]["cells"].__setitem__(0, [o["sheets"][0]["cells"][0][0], o["sheets"][0]["cells"][0][1], {"kind": "scalar", "value": float("nan")}]), allow_nan=True),
        "impossible_coordinate": payload(lambda o: o["sheets"][0]["cells"][0].__setitem__(0, "ZZZ999999999999")),
        "impossible_typed_kind": payload(lambda o: o["sheets"][0]["cells"][0].__setitem__(2, {"kind": "frobnicate", "value": 1})),
    }
    assert len(cases) == 12
    bad = []
    for name, p in cases.items():
        try:
            if isinstance(p, tuple) and p[0] == "__DIGEST__":
                art.decode(data, p[1])
            elif isinstance(p, tuple):
                hh = dict(header)
                hh[p[1][0]] = p[1][1]
                hb = json.dumps(hh, sort_keys=True, separators=(",", ":")).encode()
                b = ident.MAGIC + struct.pack(">I", len(hb)) + hb + data[off:off + 8 + clen]
                art.decode(b + hashlib.sha256(b).digest(), digest)
            else:
                art.decode(p, digest)
            bad.append(name)
            print("SERVED (BAD):", name)
        except Exception as exc:
            print("REJECTED:", name, type(exc).__name__)
    print("corruption: %d/12 rejected" % (12 - len(bad)))
    return 1 if bad else 0


def cmd_parity():
    from recalc_agent.read_engine import artifact as art
    from recalc_agent.read_engine import _identity as ident
    rc5art, rc5ident = load_pristine()
    man = json.loads((R1 / "WORKLOAD_MANIFEST.json").read_text())
    n = bad = 0
    for pop in ("A", "B"):
        for e in man["populations"][pop]["workloads"]:
            wid = e["workload_id"]
            book = sorted((CLEAN / pop / wid).glob("*.xlsx"))[0]
            d = ident.sha_file(book)
            ap1, _, _ = rc5art.ensure(book, STAGE / "par_rc5", d)
            ap2, _, _ = art.ensure(book, STAGE / "par_rel", d)
            eq = canon_book(rc5art.decode(ap1.read_bytes(), d)) == canon_book(art.decode(ap2.read_bytes(), d))
            n += 1
            if not eq:
                bad += 1
                print("MISMATCH", pop, wid)
    print("canon parity: %d/%d equal" % (n - bad, n))
    return 1 if bad else 0


def cmd_advers():
    sys.path.insert(0, str(R1))
    import aggregate as agg
    blob = subprocess.run(["git", "show",
                           "research/full-cell-iteration-probe:research/full_cell_iteration_probe/certify.py"],
                          cwd=ROOT, capture_output=True, check=True).stdout
    dest = STAGE / "r2cert.py"
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_bytes(blob)
    spec = importlib.util.spec_from_file_location("r2cert", dest)
    r2cert = importlib.util.module_from_spec(spec)
    sys.modules["r2cert"] = r2cert
    spec.loader.exec_module(r2cert)
    stage = STAGE / "adversarial"
    if stage.exists():
        shutil.rmtree(stage)
    stage.mkdir(parents=True)
    books = r2cert.build_books(stage / "books")
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet1"
    wb.active["A1"] = 1
    kp = stage / "books" / "keyerror.xlsx"
    wb.save(kp)
    books["keyerror"] = kp
    cases = r2cert.define_cases()
    cases.append(("keyerror_missing_sheet", "keyerror",
                  "import openpyxl\nwb = openpyxl.load_workbook('input.xlsx')\n"
                  "print(wb['Sheet1'].title)\nprint(wb['Nope'].title)\n",
                  {"route": "DIRECT_RUNTIME"}))
    fails = 0
    for cid, book, script, expect in cases:
        d = stage / "runs" / cid
        (d / "base").mkdir(parents=True)
        (d / "probe").mkdir(parents=True)
        for arm in ("base", "probe"):
            shutil.copy(books[book], d / arm / "input.xlsx")
            (d / arm / "case.py").write_text(script)
        b = r2cert.run([sys.executable, "case.py"], cwd=d / "base", env=dict(os.environ))
        cache = d / "cache"
        cfg = d / "probe" / "rt.toml"
        cfg.write_text('[runtime]\ncache_dir = "%s"\n' % cache)
        penv = dict(os.environ)
        penv["PYTHONPATH"] = os.pathsep.join(
            [str(ROOT / "src")] + ([penv["PYTHONPATH"]] if penv.get("PYTHONPATH") else []))
        penv.pop("RECALC_CONFIG", None)
        penv.pop("RECALC_NO_RUNTIME", None)
        p = r2cert.run([sys.executable, "-m", "recalc_agent", "run",
                        "--config", str(cfg), "--workdir", str(d / "probe"),
                        str(d / "probe" / "case.py")], cwd=d / "probe", env=penv)
        route, events = None, []
        try:
            rd = Path(json.loads((cache / "runs" / "last_run.json").read_text())["run_dir"])
            st = json.loads((rd / "setup.json").read_text())
            route = st.get("route")
            if (rd / "runtime_state.json").exists():
                rt = json.loads((rd / "runtime_state.json").read_text())
                route = rt.get("route") or route
                events = [e.get("event") + ":" + str(e.get("reason", "")) for e in rt.get("events", [])]
        except (OSError, ValueError, KeyError):
            pass
        st_b = {str(x.relative_to(d / "base")): agg.norm_xlsx(x.read_bytes())
                for x in sorted((d / "base").glob("*.xlsx"))}
        st_p = {str(x.relative_to(d / "probe")): agg.norm_xlsx(x.read_bytes())
                for x in sorted((d / "probe").glob("*.xlsx"))}
        ok = (b["exit"] == p["exit"] and r2cert.norm(b["out"]) == r2cert.norm(p["out"])
              and st_b == st_p and not b["timeout"] and not p["timeout"])
        if ok and b["exit"] != 0:
            bl = (b["err"].decode(errors="replace").strip().splitlines() or [""])[-1]
            pl = (p["err"].decode(errors="replace").strip().splitlines() or [""])[-1]
            ok = (bl == pl) if cid == "keyerror_missing_sheet" else (
                bl.split(":")[-1] == pl.split(":")[-1]
                and bl.split(":")[0].split(".")[-1] == pl.split(":")[0].split(".")[-1])
        route_ok = True
        if expect.get("route") == "REFERENCE_FAST_PATH":
            route_ok = (route in ("REFERENCE_FAST_PATH",) or route is None and p["exit"] == b["exit"])
        elif expect.get("route") == "DIRECT_RUNTIME":
            route_ok = (route == "DIRECT_RUNTIME")
        if expect.get("fallback"):
            route_ok = route_ok and any(expect["fallback"] in e for e in events)
        if not (ok and route_ok):
            fails += 1
            print("FAIL", cid, route, b["exit"], p["exit"])
    print("adversarial: %d cases, FAILURES=%d" % (len(cases), fails))
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit({"corrupt": cmd_corrupt, "parity": cmd_parity, "advers": cmd_advers}[sys.argv[1]]())
