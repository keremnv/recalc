#!/usr/bin/env python3
"""Tier 1 Layer R deterministic replay (prereg §9).

Replays each primary run's ordered file-affecting shell/python command
sequence under BASE (plain python3) and RECALC (released v0.2.0 via
recalc_agent.runner.run), paired same-window, pristine inputs.

Layout per (run, arm) workdir W (under ignored _staging/replay/):
  W/data/<relpath>   pristine input copies (container /mnt/spreadsheet_data)
  W/out/             outputs (container /mnt/spreadsheet_output)
  W/tmp/             scratch (container /tmp)
  W/abs/             other absolute paths (container /<...>); also the
                     default step cwd standing in for container /
  W/<input>.xlsx     symlink -> data/... (artifact entries + bare refs)
  W/blocks/          extracted python blocks (+chdir preamble when needed)
  W/cache/           recalc artifact/receipt cache (per run, shared blocks)

Cwd model (verified exhaustive over all 18 trajectories): fresh shell per
step; per-step cwd = last absolute `cd` in its leading chain, else /.
No cross-step cwd state exists in the data. Non-cd steps never use
relative file refs (verified). For recalc-agent (fixed cwd=workdir), the
recorded cd's effect is prepended as `import os; os.chdir(...)` to BOTH
arms' block files (classifier-neutrality verified; identical files).

Emits JSONL rows with ledger keys: RUNTIME_REPLAY.jsonl (per block),
ROUTING_CENSUS.jsonl (per run), plus a validity report per block.
"""
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import time
from pathlib import Path

BASE = Path(__file__).resolve().parent
PROJECT_ROOT = BASE.parents[1]
RUNS = BASE / "_overlay/benchmark-root/benchmark-runs/openrouter"
# Space-free scratch: replay workdirs are transient rebuilds (only JSONL
# ledgers are committed); unquoted shell substitution breaks under spaces.
STAGING = Path(os.environ.get("REPLAY_STAGING", "/tmp/t1replay"))
sys.path.insert(0, str(PROJECT_ROOT / "benchmark"))
sys.path.insert(0, str(PROJECT_ROOT / "src"))
from mine_layerm import ast_counts, dynamic_counts, extract_python  # noqa: E402
from run_openrouter_slice import _task_record  # noqa: E402

BLOCK_TIMEOUT = 300


def _prefix_map(p: str, W: Path):
    """Known-prefix remap; None when no known prefix matches."""
    if p == str(W) or p.startswith(str(W) + "/"):
        return p  # already replay-local: idempotent, never double-remap
    if p.startswith("/mnt/spreadsheet_data"):
        return str(W / "data") + p[len("/mnt/spreadsheet_data"):]
    if p.startswith("/mnt/spreadsheet_output"):
        return str(W / "out") + p[len("/mnt/spreadsheet_output"):]
    if p == "/tmp" or p.startswith("/tmp/"):
        return str(W / "tmp") + p[len("/tmp"):]
    if p == "/dev/null" or p.startswith(("/dev/", "/proc/", "/sys/")):
        return p  # system paths are identical inside/outside replay
    return None


def remap(text: str, W: Path) -> str:
    # Single-pass alternation (W itself lives under /tmp: sequential
    # replaces would double-remap). Path tokens include \-escapes.
    # XML close tags (</c>, </v>, ...) are NOT paths: a token followed
    # by `>` plus a non-word char is left untouched (no genuine
    # `/path>file` redirect occurs in the 18 trajectories).
    def rep(m: re.Match) -> str:
        p = m.group(0)
        after = m.string[m.end():m.end() + 2]
        if after.startswith(">") and (len(after) < 2 or not (after[1].isalnum() or after[1] == "_")):
            return p  # XML close tag, not a path
        known = _prefix_map(p, W)
        if known is not None:
            return known
        return str(W / "abs") + p

    return re.sub(r"(?<![\w/])/(?:[\w.\-]+(?:\\\ |[^\s'\";|&<>])*)", rep, text)


def remap_py(text: str, W: Path) -> str:
    # Python code bodies (inline -c/heredoc + cat-created scripts):
    # ONLY known prefixes map. The generic /foo rule would corrupt
    # XML close tags in regexes (</c>), division/formulas (MONTH(A2)/3),
    # and string checks ('/0' in v). Surveyed: zero genuine
    # other-absolute paths occur in python contexts across all 18 runs.
    def rep(m: re.Match) -> str:
        p = m.group(0)
        known = _prefix_map(p, W)
        return known if known is not None else p

    return re.sub(r"(?<![\w/])/(?:[\w.\-]+(?:\\\ |[^\s'\";|&<>])*)", rep, text)


def step_cwd(action: str, W: Path) -> Path:
    """Remapped cwd for one step: last absolute cd in leading chain else /."""
    cwd = Path("/")  # container default (image WORKDIR=/)
    head = action.split("&&")[0] if "&&" in action.split("\n")[0] else action.split("\n")[0]
    for m in re.finditer(r"\bcd\s+((?:\\\ |[^\s;'\"])+)", head):
        target = m.group(1).replace("\\ ", " ").strip("'\"").rstrip(";")
        if target.startswith("/"):
            cwd = Path(target)
        else:  # relative cd inside one action chains onto previous
            cwd = cwd / target
    if str(cwd) == "/":
        return W / "abs"
    mapped = remap(str(cwd), W)
    return Path(mapped)


def classify_step(action: str) -> str:
    a = action.strip()
    if not a:
        return "skip"
    firstline = a.split("\n")[0]
    first = firstline.split("&&")[0].strip().split()
    first = first[0] if first else ""
    if first in ("view_xlsx", "submit"):
        return "skip"
    # A step is python iff it INVOKES python (heredoc/-c/file-exec),
    # anywhere in the action (incl. after cat-creators in combined steps).
    # Pure cat-creators (never executed) are shell steps.
    if re.search(r"python3?\s+(?:-\s+)?<<|python3?\s+-c\b|python3?\s+\S+\.py", a):
        return "python"
    return "shell"


PYINV = r"python3?\s+(?:-\s+)?<<|python3?\s+-c\b|python3?\s+\S+\.py"


def split_step(action: str) -> list:
    """Split a step into ordered ('shell', text) / ('pycode', code) /
    ('pyfile', path) parts. Heredoc spans (python + cat-creators) are
    atomic; -c bodies end at the quote followed by &&/;/|/newline/end."""
    spans: list = []  # (start, end, kind, payload)

    def overlaps(s, e):
        return any(s < e2 and e > s2 for s2, e2, _, _ in spans)

    for m in re.finditer(
        r"python3?\s+(?:-\s+)?<<\s*['\"]?(\w+)['\"]?\n(.*?)\n\1\b",
        action, re.S,
    ):
        spans.append((m.start(), m.end(), "pycode", m.group(2)))
    for m in re.finditer(CATOM_RE, action, re.S):
        # atomic shell text: blocks inner -c/pyfile false matches.
        # The header line is shell; the heredoc body is file content
        # (python for *.py targets) remapped with prefix-only rules.
        spans.append((m.start(), m.end(), "catshell", ""))
    for m in re.finditer(r"python3?\s+-c\s+(['\"])(.*?)\1(?=\s*(?:2>&1\s*)?(?:&&|;|\||\n|$))", action, re.S):
        if not overlaps(m.start(), m.end()):
            _q, _body = m.group(1), m.group(2)
            if _q == '"':
                # bash double-quote backslash semantics: backslash is
                # special only before $ ` " \ newline (all 50 -c in
                # corpus are double-quoted; heredocs are quoted=literal).
                _body = _body.replace("\\\n", "")
                _body = re.sub(r'\\([$"\\`])', r"\1", _body)
            spans.append((m.start(), m.end(), "pycode", _body))
    for m in re.finditer(r"python3?\s+(\S+\.py\w*)", action):
        if not overlaps(m.start(), m.end()):
            spans.append((m.start(), m.end(), "pyfile", m.group(1)))
    spans.sort()
    parts: list = []
    pos = 0
    for s, e, kind, payload in spans:
        if s > pos:
            parts.append(("shell", action[pos:s]))
        if kind != "catshell":
            parts.append((kind, payload))
        else:
            parts.append(("catshell", action[s:e]))
        pos = e
    if pos < len(action):
        parts.append(("shell", action[pos:]))
    # drop empty shell gaps; detect leading joiner of each shell gap
    out = []
    for kind, text in parts:
        if kind == "shell":
            t = text.strip()
            if not t or t in ("&&", ";", "|"):
                continue
            out.append((kind, text))
        else:
            out.append((kind, text))
    return out


def remap_catom(text: str, W: Path) -> str:
    """Remap a cat-creator: header line as shell, body prefix-only."""
    lines = text.split("\n")
    if len(lines) < 3:
        return remap(text, W)
    header, delim, body = lines[0], lines[-1], "\n".join(lines[1:-1])
    return remap(header, W) + "\n" + remap_py(body, W) + "\n" + delim


CATOM_RE = r"cat\s+>\s*\S+\s+<<\s*['\"]?(\w+)['\"]?\n(.*?)\n\1\b"


def remap_action(text: str, W: Path) -> str:
    """Remap a whole shell action catom-aware (creators' bodies: py rules)."""
    out: list = []
    pos = 0
    for m in re.finditer(CATOM_RE, text, re.S):
        out.append(remap(text[pos:m.start()], W))
        out.append(remap_catom(m.group(0), W))
        pos = m.end()
    out.append(remap(text[pos:], W))
    return "".join(out)


def cds_in(text: str) -> list:
    """Absolute/relative cd targets in order (shell escape aware)."""
    out = []
    for m in re.finditer(r"\bcd\s+((?:\\\ |[^\s;'\"`])+)", text):
        out.append(m.group(1).replace("\\ ", " ").strip("'\"").rstrip(";"))
    return out


def gap_joiner(gap: str) -> str:
    t = gap.strip()
    if t.startswith("&&"):
        return "&&"
    if t.startswith(";"):
        return ";"
    if t.startswith("|"):
        return "|"
    return ""


def norm_output(text: str, W: Path, recorded: bool) -> str:
    if recorded:
        out = text.replace("/mnt/spreadsheet_data", "<DATA>")
        out = out.replace("/mnt/spreadsheet_output", "<OUT>")
        out = re.sub(r"(?<!\S)/tmp(?=/|\s|'|\"|$)", "<TMP>", out)
    else:
        out = text.replace(str(W / "data"), "<DATA>")
        out = out.replace(str(W / "out"), "<OUT>")
        out = out.replace(str(W / "tmp"), "<TMP>")
        out = out.replace(str(W), "<W>")
    return out


class _FdCapture:
    """Capture os-level stdout/stderr (for recalc_run's inherited fds)."""

    def __init__(self, path_out: Path, path_err: Path):
        self.path_out = path_out
        self.path_err = path_err

    def __enter__(self):
        self.sv_out = os.dup(1)
        self.sv_err = os.dup(2)
        self.fo = open(self.path_out, "wb")
        self.fe = open(self.path_err, "wb")
        os.dup2(self.fo.fileno(), 1)
        os.dup2(self.fe.fileno(), 2)
        return self

    def __exit__(self, *exc):
        os.dup2(self.sv_out, 1)
        os.dup2(self.sv_err, 2)
        os.close(self.sv_out)
        os.close(self.sv_err)
        self.fo.close()
        self.fe.close()
        return False


def xlsx_semantic(path: Path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=False)
    sig = [tuple(wb.sheetnames)]
    for ws in wb.worksheets:
        cells = []
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    cells.append((c.coordinate, c.value if not isinstance(c.value, str) or len(c.value) < 500 else c.value[:500], str(c.data_type)))
        sig.append((ws.title, str(ws.dimensions), sorted(cells)))
    wb.close()
    return sig


def run_shell(cmd: str, cwd: Path) -> tuple[int, str, str, float]:
    t0 = time.time()
    try:
        p = subprocess.run(["bash", "-c", cmd], cwd=str(cwd), capture_output=True,
                           text=True, timeout=BLOCK_TIMEOUT)
        return p.returncode, p.stdout, p.stderr, time.time() - t0
    except subprocess.TimeoutExpired as e:
        return 124, e.stdout.decode() if isinstance(e.stdout, bytes) else (e.stdout or ""), \
            (e.stderr.decode() if isinstance(e.stderr, bytes) else (e.stderr or "")) + "\n[TIMEOUT]", BLOCK_TIMEOUT


def main() -> None:
    only = sys.argv[1] if len(sys.argv) > 1 else None
    from recalc_agent import config as recalc_config
    from recalc_agent.runner import run as recalc_run
    assert (PROJECT_ROOT / "src").exists()
    run_dirs = []
    for pat in ("tier1-r*-P-*", "tier1-r*-O-mimo-*"):
        run_dirs.extend(sorted(RUNS.glob(pat)))
    if only:
        run_dirs = [d for d in run_dirs if only in d.name]
    bro = BASE / "_overlay/benchmark-root"
    for run_dir in run_dirs:
        rec = json.loads((run_dir / "ledger.jsonl").read_text().strip().splitlines()[0])
        task = rec.get("task")
        category, task_id = task.split(":")
        record = _task_record(bro, category, task_id)
        traj = json.loads(sorted(run_dir.glob("*/trajectory/*/*.traj"))[0].read_text())
        steps = traj.get("trajectory", [])
        input_src = bro / "data" / category / record["spreadsheet_path"]
        for arm in ("base", "recalc"):
            W = STAGING / run_dir.name / arm
            if W.exists():
                shutil.rmtree(W)
            (W / "data").mkdir(parents=True)
            (W / "out").mkdir(parents=True)
            (W / "tmp").mkdir(parents=True)
            (W / "abs").mkdir(parents=True)
            (W / "blocks").mkdir(parents=True)
            # pristine input at recorded relative path + top-level symlink
            rel = Path(record["spreadsheet_path"])
            dest = W / "data" / rel
            dest.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(input_src, dest)
            link = W / rel.name
            if not link.exists():
                link.symlink_to(dest.relative_to(W))
            cache = W / "cache"
            cfg, issues = recalc_config.load(None, False)
            n_valid = n_total = 0
            block_rows: list = []
            for i, s in enumerate(steps):
                action = s.get("action", "") or ""
                recorded = s.get("observation", "") or ""
                kind = classify_step(action)
                if kind == "skip":
                    continue
                if kind == "shell":
                    # whole action in one real bash at container / (= W/abs)
                    run_shell(remap_action(action, W), W / "abs")
                    continue
                parts = split_step(action)
                step_out: list = []  # accumulated stdout contributions
                step_err: list = []
                last_rc = 0
                skip_and = False
                pend_out = pend_err = ""  # held python output (pipe target)
                prev_py = False  # immediately preceding part was python
                cur = W / "abs"  # tracked cwd; fresh shell starts at /
                has_py = False
                bi = -1
                for kind2, payload in parts:
                    if kind2 in ("shell", "catshell"):
                        gap = remap_catom(payload, W) if kind2 == "catshell" else remap(payload, W)
                        lead = gap_joiner(payload)
                        if lead == "&&" and (skip_and or last_rc != 0):
                            skip_and = True
                            prev_py = False
                            continue
                        if lead == "|" or re.sub(
                                r"^\s*(?:\d*>&?\d+|\d*>\s*\S+|\d*<\s*\S+)(?:\s*(?:\d*>&?\d+|\d*>\s*\S+|\d*<\s*\S+))*\s*",
                                "", gap).startswith("|"):
                            # pipe gap (possibly after redirects: `2>&1 | tail`)
                            frag = re.sub(r"^.*?\|", "", gap, count=1).strip()
                            frag = re.sub(r"2>&1", "", frag).strip()
                            if step_out:
                                step_out.pop()
                            if step_err:
                                step_err.pop()
                            merged = pend_out + (pend_err if "2>&1" in gap else "")
                            p = subprocess.run(["bash", "-c", frag], cwd=str(cur),
                                               input=merged, capture_output=True,
                                               text=True, timeout=BLOCK_TIMEOUT)
                            step_out.append(p.stdout)
                            step_err.append(p.stderr)
                            last_rc = p.returncode
                            pend_out = pend_err = ""
                            prev_py = False
                            skip_and = False
                            continue
                        # leading output redirect diverting python output to a
                        # file: `python3 x.py > f 2>&1; ...` (only r18 s05).
                        rmd = re.match(r"\s*(>>?)\s*(\S+)(.*)$", gap, re.S) if prev_py else None
                        if rmd:
                            if step_out:
                                step_out.pop()
                            if step_err:
                                step_err.pop()
                            content = pend_out + (pend_err if re.match(
                                r"\s*2>&1(\s|;|$)", rmd.group(3)) else "")
                            try:
                                with open(rmd.group(2), "a" if rmd.group(1) == ">>" else "w") as _f:
                                    _f.write(content)
                            except OSError:
                                pass
                            pend_out = pend_err = ""
                            prev_py = False
                            gap = re.sub(r"^\s*2>&1\s*", "", rmd.group(3))
                            if not gap.strip():
                                continue  # rc stays the python's rc
                            lead = gap_joiner(gap)
                            if lead == "&&" and last_rc != 0:
                                skip_and = True
                                continue
                        # strip boundary joiners for execution (gating consumed)
                        exec_text = re.sub(r"^\s*(?:&&|;)\s*", "", gap).rstrip()
                        trail = ""
                        mt = re.search(r"(&&|;)\s*$", exec_text)
                        if mt:
                            trail = mt.group(1)
                            exec_text = exec_text[:mt.start()].rstrip()
                        if not exec_text:
                            prev_py = False
                            continue
                        skip_and = False
                        prev_py = False
                        cur.mkdir(parents=True, exist_ok=True)
                        rc, so, se, _ = run_shell(exec_text, cur)
                        step_out.append(so)
                        step_err.append(se)
                        last_rc = rc
                        # track cds for subsequent parts: first &&-piece of
                        # each ;-segment always runs; later pieces ran iff
                        # the gap succeeded (|| subtleties approximated).
                        if kind2 == "shell":
                            for seg in payload.split(";"):
                                for j, piece in enumerate(seg.split("&&")):
                                    if j > 0 and rc != 0:
                                        break
                                    for t in cds_in(piece):
                                        if t in ("-", "~") or not t:
                                            continue
                                        cur = Path(remap(t, W)) if t.startswith("/") else cur / t
                        if trail == "&&" and rc != 0:
                            skip_and = True
                        continue
                    # python part
                    if skip_and:
                        block_rows.append((i, bi + 1, "", -1, 0.0, "", "", {}))
                        has_py = True
                        bi += 1
                        n_total += 1
                        prev_py = False
                        continue
                    has_py = True
                    bi += 1
                    n_total += 1
                    if kind2 == "pyfile":
                        try:
                            # File content was already remapped when the
                            # creator gap ran; do NOT remap twice.
                            # Relative refs resolve against the part cwd.
                            _p = remap(payload, W)
                            _p = _p if os.path.isabs(_p) else str(cur / _p)
                            body = Path(_p).read_text()
                        except OSError:
                            body = ""
                        remapped = body
                    else:
                        body = payload
                        remapped = remap_py(body, W)
                    pream = f"import os; os.chdir({str(cur)!r})\n"
                    block = W / "blocks" / f"b{i:02d}_{bi}.py"
                    block.write_text(pream + remapped)
                    cur.mkdir(parents=True, exist_ok=True)
                    t0 = time.time()
                    if arm == "base":
                        import shlex
                        rc, so, se, wall = run_shell(
                            f"{shlex.quote(sys.executable)} {shlex.quote(str(block))}", cur)
                        summary = {}
                    else:
                        cap_out, cap_err = W / "blocks" / "cap.out", W / "blocks" / "cap.err"
                        try:
                            with _FdCapture(cap_out, cap_err):
                                rc, summary = recalc_run(
                                    block, [], W, cfg, issues, False)
                            wall = time.time() - t0
                            so = cap_out.read_text(errors="replace")
                            se = "\n".join(
                                ln for ln in cap_err.read_text(errors="replace").splitlines()
                                if not ln.startswith("WARNING: "))
                        except Exception as exc:  # noqa: BLE001
                            rc, so, se, wall = 125, "", f"REPLAY-HARNESS: {exc}", 0.0
                            summary = {}
                    pend_out, pend_err = so, se
                    last_rc = rc
                    prev_py = True
                    # flush unless a pipe gap follows (peek: handled when seen)
                    step_out.append(so)
                    step_err.append(se)
                    block_rows.append((i, bi, body, rc, wall, so, se, summary))
                if not has_py:
                    continue
                combined_out = "".join(step_out)
                combined_err = "".join(step_err)
                validity = "VALID" if norm_output(combined_out, W, False) == norm_output(recorded, W, True) and not combined_err.strip() else "MISMATCH"
                if validity == "MISMATCH":
                    # both-failed: same exception type => replay-consistent
                    m1 = re.search(r"^(\w+(?:Error|Exception|Warning))", combined_err.strip().splitlines()[-1] if combined_err.strip() else "", re.M)
                    m2 = re.search(r"^(\w+(?:Error|Exception|Warning))", recorded.strip().splitlines()[-1] if recorded.strip() else "", re.M)
                    if m1 and m2 and m1.group(1) == m2.group(1) and "Traceback" in combined_err and "Traceback" in recorded:
                        validity = "VALID-both-failed"
                if len(recorded) >= 29000:
                    validity = "VALID-unchecked-truncated" if validity == "VALID" else "MISMATCH-truncated-ref"
                if validity.startswith("MISMATCH") and os.environ.get("REPLAY_DEBUG"):
                    (W / "blocks" / f"debug_s{i:02d}.json").write_text(json.dumps(
                        {"combined_out": combined_out[-8000:], "combined_err": combined_err[-4000:],
                         "recorded": recorded[-8000:]}))
                n_valid += validity.startswith("VALID")
                for (i2, bi2, body2, rc2, wall2, so2, se2, summary2) in [r for r in block_rows if r[0] == i]:
                    times = counts = fb_reasons = None
                    clf = None
                    if arm == "recalc" and body2.strip():
                        try:
                            from recalc_agent._frozen.eligibility import classify as _clf
                            _d = _clf((W / "blocks" / f"b{i2:02d}_{bi2}.py").read_text())
                            clf = {"decision": _d.get("decision"), "reason": _d.get("reason"),
                                   "blockers": [(b.get("reason"), (b.get("detail") or "")[:60]) for b in _d.get("blockers", [])],
                                   "categories": sorted({c.get("category") for c in _d.get("categories", []) if c.get("category")})}
                        except Exception as exc:  # noqa: BLE001
                            clf = {"error": f"{type(exc).__name__}: {exc}"[:120]}
                    if summary2.get("run_dir"):
                        try:
                            rs = json.loads((Path(summary2["run_dir"]) / "runtime_state.json").read_text())
                            times = rs.get("times")
                            counts = rs.get("counts")
                            fb_reasons = rs.get("fallback_reasons")
                        except OSError:
                            pass
                    row = {"ledger": "RUNTIME_REPLAY.jsonl", "run_name": run_dir.name,
                           "task": task, "model": rec.get("model"), "arm": arm,
                           "step": i2, "block": bi2, "bytes": len(body2),
                           "exit": rc2, "wall_s": round(wall2, 3),
                           "stdout_sha": hashlib.sha256(so2.encode()).hexdigest()[:16],
                           "stderr_head": se2[:200],
                           "validity": validity,
                           "ast": ast_counts([body2]), "dynamic": dynamic_counts([body2]),
                           "route": summary2.get("route"), "admitted": summary2.get("admitted"),
                           "admission_reason": summary2.get("admission_reason"),
                           "artifact": summary2.get("artifact"),
                           "direct_served_loads": summary2.get("direct_served_loads"),
                           "fallback": summary2.get("fallback"),
                           "fallback_reasons": fb_reasons,
                           "times": times, "counts": counts,
                           "classifier": clf,
                           "assurance": summary2.get("assurance_status")}
                    print(json.dumps(row), flush=True)
        print(json.dumps({"ledger": "ROUTING_CENSUS.jsonl", "run_name": run_dir.name,
                          "task": task, "model": rec.get("model"),
                          "note": "routes aggregated from RUNTIME_REPLAY rows"}), flush=True)


if __name__ == "__main__":
    main()