"""Tier 1 curated-task evaluator (frozen). Compares agent output vs golden.

Value cells: exact or tolerance-compared. Formula cells: formula presence
(data_type) + computed value after LibreOffice refresh. Structural checks
as specified. No partial-credit invention beyond the frozen spec below.

Usage: python3 evaluate_curated.py <specs.json> <task_id> <output.xlsx> <golden.xlsx>
Prints JSON: {exact_success, n_correct, n_total, failures[]}
"""
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl

EVAL_DIR = Path(__file__).resolve().parents[2] / "benchmark-data" / "SpreadsheetBench-2" / "evaluation"
sys.path.insert(0, str(EVAL_DIR))
from open_spreadsheet import (batch_open_files, start_libreoffice_service,  # noqa: E402
                              stop_libreoffice_service)


def refreshed_values(path):
    """Recalc via LO, return {sheet: {coord: value}} with data_only."""
    with tempfile.TemporaryDirectory() as td:
        work = str(Path(td) / "w.xlsx")
        shutil.copy2(path, work)
        batch_open_files([work])
        wb = openpyxl.load_workbook(work, data_only=True)
        out = {}
        for ws in wb.worksheets:
            out[ws.title] = {}
            for row in ws.iter_rows():
                for c in row:
                    out[ws.title][c.coordinate] = c.value
        return out


def close(a, b, tol=0.01):
    if a is None or b is None:
        return a == b
    if isinstance(a, float) or isinstance(b, float):
        try:
            return abs(float(a) - float(b)) <= tol
        except (TypeError, ValueError):
            return a == b
    return a == b


def check_cells(wb_out, golden, cells, computed):
    """cells: ['Sheet!A1']; formulas in golden checked as formulas+computed."""
    fails, n_ok = [], 0
    for ref in cells:
        sheet, coord = ref.split("!")
        exp = golden[sheet][coord]
        got_wb = wb_out[sheet]
        cell = got_wb[coord]
        if isinstance(exp, str) and exp.startswith("="):
            if cell.data_type != "f":
                fails.append(ref + ": not a formula")
                continue
            if not close(computed[sheet][coord], _golden_computed(ref), tol=0.01):
                fails.append(ref + ": computed mismatch")
                continue
            n_ok += 1
        else:
            if not close(cell.value, exp):
                fails.append(ref + ": %r != %r" % (cell.value, exp))
            else:
                n_ok += 1
    return n_ok, fails


_GOLDEN_COMPUTED = {}


def _golden_computed(ref):
    return _GOLDEN_COMPUTED[ref]


def evaluate(task_id, output, golden_path, specs):
    spec = specs[task_id]
    wb_out = openpyxl.load_workbook(output)
    wb_g = openpyxl.load_workbook(golden_path)
    golden = {ws.title: {c.coordinate: c.value for row in ws.iter_rows() for c in row}
              for ws in wb_g.worksheets}
    need_refresh = spec.get("check") in (
        "dates_iso+quarter_formulas+bold_headers", "unmerged+filled+total_formulas",
        "variance_formulas+top3", "ending_formulas+flags")
    computed = refreshed_values(output) if need_refresh else {}
    global _GOLDEN_COMPUTED
    _GOLDEN_COMPUTED = {}
    if need_refresh:
        gc = refreshed_values(golden_path)
        for sh, cells in gc.items():
            for coord, v in cells.items():
                _GOLDEN_COMPUTED["%s!%s" % (sh, coord)] = v
    fails, n_ok, n_total = [], 0, 0
    if task_id == "T1_R1":
        n_total = 6
        n_ok, fails = check_cells(wb_out, golden, spec["answer_cells"], computed)
    elif task_id == "T1_R2":
        exp = sorted(v for k, v in golden["Findings"].items() if k != "A1" and v is not None)
        got = sorted(v for k, v in
                     {c.coordinate: c.value for row in wb_out["Findings"].iter_rows() for c in row}.items()
                     if k != "A1" and v is not None)
        n_total = len(exp)
        n_ok = len(set(exp) & set(got))
        if set(got) != set(exp):
            fails.append("set mismatch: got=%s exp=%s" % (got, exp))
    elif task_id == "T1_W1":
        months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun"]
        for m in months:
            hdr = wb_out[m]["A1"]
            n_total += 1
            if hdr.font and hdr.font.bold:
                n_ok += 1
            else:
                fails.append(m + " header not bold")
            for r in range(2, 102):
                n_total += 2
                if wb_out[m].cell(row=r, column=1).value == golden[m]["A%d" % r]:
                    n_ok += 1
                else:
                    fails.append("%s!A%d date" % (m, r))
                    break
                c = wb_out[m].cell(row=r, column=3)
                if c.data_type == "f" and close(
                        computed[m]["C%d" % r], _golden_computed("%s!C%d" % (m, r))):
                    n_ok += 1
                else:
                    fails.append("%s!C%d quarter" % (m, r))
                    break
    elif task_id == "T1_W2":
        for sh in ("Dept1", "Dept2"):
            n_total += 1
            if list(wb_out[sh].merged_cells.ranges):
                fails.append(sh + " still merged")
            else:
                n_ok += 1
            for r in range(2, 11):
                n_total += 1
                if wb_out[sh].cell(row=r, column=1).value == golden[sh]["A%d" % r]:
                    n_ok += 1
                else:
                    fails.append("%s!A%d fill" % (sh, r))
                    break
            for col in ("B", "C"):
                n_total += 1
                c = wb_out[sh][col + "11"]
                if c.data_type == "f" and close(
                        computed[sh][col + "11"], _golden_computed("%s!%s11" % (sh, col))):
                    n_ok += 1
                else:
                    fails.append("%s!%s11 total" % (sh, col))
    elif task_id == "T1_M1":
        ws = wb_out["Variance"]
        for r in range(2, 12):
            for col in ("B", "C", "D"):
                n_total += 1
                c = ws[col + str(r)]
                if c.data_type == "f" and close(
                        computed["Variance"][col + str(r)],
                        _golden_computed("Variance!%s%d" % (col, r))):
                    n_ok += 1
                else:
                    fails.append("Variance!%s%d" % (col, r))
                    break
        for r in range(14, 17):
            n_total += 1
            if ws.cell(row=r, column=1).value == golden["Variance"]["A%d" % r]:
                n_ok += 1
            else:
                fails.append("Variance!A%d list" % r)
    elif task_id == "T1_M2":
        ws = wb_out["Inventory"]
        for r in range(2, 17):
            n_total += 2
            c = ws.cell(row=r, column=2)
            if c.data_type == "f" and close(
                    computed["Inventory"]["B%d" % r],
                    _golden_computed("Inventory!B%d" % r)):
                n_ok += 1
            else:
                fails.append("Inventory!B%d ending" % r)
            if ws.cell(row=r, column=3).value == golden["Inventory"]["C%d" % r]:
                n_ok += 1
            else:
                fails.append("Inventory!C%d flag" % r)
    return {"exact_success": not fails, "n_correct": n_ok, "n_total": n_total,
            "failures": fails[:20]}


def main():
    import contextlib
    import io
    specs = json.loads(Path(sys.argv[1]).read_text())
    with contextlib.redirect_stdout(io.StringIO()):
        start_libreoffice_service()
    try:
        with contextlib.redirect_stdout(io.StringIO()):
            res = evaluate(sys.argv[2], sys.argv[3], sys.argv[4], specs)
    finally:
        with contextlib.redirect_stdout(io.StringIO()):
            stop_libreoffice_service()
    print(json.dumps(res, indent=1))


if __name__ == "__main__":
    main()
