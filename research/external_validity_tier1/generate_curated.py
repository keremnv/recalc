"""Tier 1 curated task workbook generator (frozen provenance: authored).

Deterministic (seeded). Produces input.xlsx + golden.xlsx per task.
Usage: python3 generate_curated.py <outdir>
"""
import json
import random
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font


def _save(wb, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def gen_r1(out):
    """Regional sales lookup: top-3 West Q3 products + totals into Report."""
    rng = random.Random(11)
    regions = ["West", "East", "North", "South"]
    products = ["P-%02d" % i for i in range(1, 21)]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Regions"
    ws.append(["Region", "Manager"])
    for i, r in enumerate(regions):
        ws.append([r, "M-%d" % (i + 1)])
    ws = wb.create_sheet("Sales")
    ws.append(["Region", "Product", "Month", "Revenue"])
    for i in range(1500):
        ws.append([rng.choice(regions), rng.choice(products), rng.randint(1, 12),
                   round(rng.uniform(100, 5000), 2)])
    ws = wb.create_sheet("Products")
    ws.append(["Product", "Category"])
    for p in products:
        ws.append([p, rng.choice(["A", "B", "C"])])
    ws = wb.create_sheet("Report")
    ws.append(["Top West Q3 product", "Q3 revenue"])
    _save(wb, out / "T1_R1_input.xlsx")
    # golden: compute top-3
    totals = {}
    for row in wb["Sales"].iter_rows(min_row=2, values_only=True):
        if row[0] == "West" and row[2] in (7, 8, 9):
            totals[row[1]] = totals.get(row[1], 0) + row[3]
    top = sorted(totals.items(), key=lambda kv: -kv[1])[:3]
    rep = wb["Report"]
    for i, (p, t) in enumerate(top, start=2):
        rep.cell(row=i, column=1, value=p)
        rep.cell(row=i, column=2, value=round(t, 2))
    _save(wb, out / "T1_R1_golden.xlsx")
    return {"answer_cells": ["Report!A2", "Report!A3", "Report!A4",
                             "Report!B2", "Report!B3", "Report!B4"]}


def gen_r2(out):
    """Formula audit: find 6 planted defects, list as Sheet!Addr."""
    rng = random.Random(22)
    wb = openpyxl.Workbook()
    wb.active.title = "S1"
    defects = []
    for s in range(1, 5):
        ws = wb["S%d" % s] if s == 1 else wb.create_sheet("S%d" % s)
        ws.append(["Item", "Qty", "Price", "Total"])
        for r in range(2, 42):
            ws.cell(row=r, column=1, value="I-%d-%d" % (s, r))
            ws.cell(row=r, column=2, value=rng.randint(1, 50))
            ws.cell(row=r, column=3, value=round(rng.uniform(1, 99), 2))
            ws.cell(row=r, column=4, value="=B%d*C%d" % (r, r))
    # defects: 2 hardcoded-among-formulas, 2 broken refs, 2 inconsistent formulas
    picks = [("S1", 5), ("S2", 9)]
    for sh, r in picks:
        wb[sh].cell(row=r, column=4, value=1234.5)
        defects.append("%s!D%d" % (sh, r))
    picks = [("S3", 12), ("S4", 20)]
    for sh, r in picks:
        wb[sh].cell(row=r, column=4, value="=#REF!*C%d" % r)
        defects.append("%s!D%d" % (sh, r))
    picks = [("S1", 30), ("S2", 33)]
    for sh, r in picks:
        wb[sh].cell(row=r, column=4, value="=B%d+C%d" % (r, r))
        defects.append("%s!D%d" % (sh, r))
    ws = wb.create_sheet("Findings")
    ws.append(["Defective cell"])
    _save(wb, out / "T1_R2_input.xlsx")
    for i, d in enumerate(sorted(defects), start=2):
        wb["Findings"].cell(row=i, column=1, value=d)
    _save(wb, out / "T1_R2_golden.xlsx")
    return {"answer_cells": ["Findings!A%d" % i for i in range(2, 8)],
            "order_insensitive": True}


def gen_w1(out):
    """Normalize dates + Quarter formulas + bold headers on 6 sheets."""
    rng = random.Random(33)
    months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun"]
    fmts = ["%m/%d/%Y", "%d-%b-%y", "%Y.%m.%d"]
    wb = openpyxl.Workbook()
    first = True
    for m in months:
        ws = wb.active if first else wb.create_sheet(m)
        if first:
            ws.title = m
            first = False
        ws.append(["Date", "Amount", "Quarter"])
        for r in range(2, 102):
            import datetime as dt
            d = dt.date(2024, months.index(m) + 1, rng.randint(1, 28))
            ws.cell(row=r, column=1, value=d.strftime(rng.choice(fmts)))
            ws.cell(row=r, column=2, value=round(rng.uniform(10, 1000), 2))
    _save(wb, out / "T1_W1_input.xlsx")
    bold = Font(bold=True)
    for m in months:
        ws = wb[m]
        for c in range(1, 4):
            ws.cell(row=1, column=c).font = bold
        for r in range(2, 102):
            raw = ws.cell(row=r, column=1).value
            for f in ("%m/%d/%Y", "%d-%b-%y", "%Y.%m.%d"):
                try:
                    import datetime as dt
                    d = dt.datetime.strptime(raw, f).date()
                    break
                except ValueError:
                    continue
            ws.cell(row=r, column=1, value=d.strftime("%Y-%m-%d"))
            ws.cell(row=r, column=3, value='="Q"&INT((MONTH(DATEVALUE(A%d))-1)/3)+1' % r)
    _save(wb, out / "T1_W1_golden.xlsx")
    return {"check": "dates_iso+quarter_formulas+bold_headers"}


def gen_w2(out):
    """Unmerge + fill down + TOTAL SUM rows on 2 sheets."""
    rng = random.Random(44)
    wb = openpyxl.Workbook()
    wb.active.title = "Dept1"
    wb.create_sheet("Dept2")
    for sh in ("Dept1", "Dept2"):
        ws = wb[sh]
        ws.append(["Team", "Q1", "Q2"])
        teams = ["Alpha", "Beta", "Gamma"]
        r = 2
        for t in teams:
            ws.merge_cells(start_row=r, start_column=1, end_row=r + 2, end_column=1)
            ws.cell(row=r, column=1, value=t)
            for i in range(3):
                ws.cell(row=r + i, column=2, value=rng.randint(10, 99))
                ws.cell(row=r + i, column=3, value=rng.randint(10, 99))
            r += 3
    _save(wb, out / "T1_W2_input.xlsx")
    for sh in ("Dept1", "Dept2"):
        ws = wb[sh]
        ws.unmerge_cells("A2:A4")
        ws.unmerge_cells("A5:A7")
        ws.unmerge_cells("A8:A10")
        for t, r0 in (("Alpha", 2), ("Beta", 5), ("Gamma", 8)):
            for r in range(r0, r0 + 3):
                ws.cell(row=r, column=1, value=t)
        ws.cell(row=11, column=1, value="TOTAL")
        ws.cell(row=11, column=2, value="=SUM(B2:B10)")
        ws.cell(row=11, column=3, value="=SUM(C2:C10)")
    _save(wb, out / "T1_W2_golden.xlsx")
    return {"check": "unmerged+filled+total_formulas"}


def gen_m1(out):
    """Budget vs actual: Variance sheet with formulas + top-3 overspend."""
    rng = random.Random(55)
    lines = ["Rent", "Payroll", "Travel", "Software", "Marketing", "Legal",
             "Facilities", "Training", "Insurance", "Misc"]
    wb = openpyxl.Workbook()
    wb.active.title = "Budget"
    wb["Budget"].append(["Line", "Amount"])
    wb.create_sheet("Actuals").append(["Line", "Amount"])
    for i, ln in enumerate(lines, start=2):
        b = round(rng.uniform(5000, 60000), 2)
        a = round(b * rng.uniform(0.7, 1.4), 2)
        wb["Budget"].cell(row=i, column=1, value=ln)
        wb["Budget"].cell(row=i, column=2, value=b)
        wb["Actuals"].cell(row=i, column=1, value=ln)
        wb["Actuals"].cell(row=i, column=2, value=a)
    _save(wb, out / "T1_M1_input.xlsx")
    ws = wb.create_sheet("Variance")
    ws.append(["Line", "Budget", "Actual", "Variance"])
    for i, ln in enumerate(lines, start=2):
        ws.cell(row=i, column=1, value=ln)
        ws.cell(row=i, column=2, value="=Budget!B%d" % i)
        ws.cell(row=i, column=3, value="=Actuals!B%d" % i)
        ws.cell(row=i, column=4, value="=C%d-B%d" % (i, i))
    over = sorted(lines, key=lambda ln: next(
        wb["Actuals"].cell(row=j, column=2).value - wb["Budget"].cell(row=j, column=2).value
        for j in range(2, 12) if wb["Budget"].cell(row=j, column=1).value == ln), reverse=True)[:3]
    ws.cell(row=13, column=1, value="Top overspend")
    for k, ln in enumerate(over):
        ws.cell(row=14 + k, column=1, value=ln)
    _save(wb, out / "T1_M1_golden.xlsx")
    return {"check": "variance_formulas+top3"}


def gen_m2(out):
    """Inventory reconcile: ending inventory formulas + negative flags."""
    rng = random.Random(66)
    skus = ["SKU-%03d" % i for i in range(1, 16)]
    wb = openpyxl.Workbook()
    wb.active.title = "Starting"
    wb["Starting"].append(["SKU", "Qty"])
    wb.create_sheet("Receipts").append(["SKU", "Qty"])
    wb.create_sheet("Shipments").append(["SKU", "Qty"])
    for i, s in enumerate(skus, start=2):
        st = rng.randint(20, 200)
        rc = rng.randint(0, 150)
        sh = rng.randint(0, 300)
        wb["Starting"].cell(row=i, column=1, value=s)
        wb["Starting"].cell(row=i, column=2, value=st)
        wb["Receipts"].cell(row=i, column=1, value=s)
        wb["Receipts"].cell(row=i, column=2, value=rc)
        wb["Shipments"].cell(row=i, column=1, value=s)
        wb["Shipments"].cell(row=i, column=2, value=sh)
    _save(wb, out / "T1_M2_input.xlsx")
    ws = wb.create_sheet("Inventory")
    ws.append(["SKU", "Ending", "Flag"])
    for i, s in enumerate(skus, start=2):
        ws.cell(row=i, column=1, value=s)
        ws.cell(row=i, column=2,
                value="=Starting!B%d+Receipts!B%d-Shipments!B%d" % (i, i, i))
        end = (wb["Starting"].cell(row=i, column=2).value
               + wb["Receipts"].cell(row=i, column=2).value
               - wb["Shipments"].cell(row=i, column=2).value)
        ws.cell(row=i, column=3, value="NEGATIVE" if end < 0 else "OK")
    _save(wb, out / "T1_M2_golden.xlsx")
    return {"check": "ending_formulas+flags"}


INSTRUCTIONS = {
    "T1_R1": ("Using the Sales sheet, find the top 3 products by total Q3 "
              "(months 7-9) revenue in region 'West'. Write the product names "
              "in Report!A2:A4 (rank order) and their Q3 totals in Report!B2:B4. "
              "Do not modify other sheets."),
    "T1_R2": ("Audit formulas on sheets S1-S4: each Total cell (column D) "
              "should equal Qty*Price for its row. List every defective cell "
              "as 'Sheet!Addr' (e.g. S1!D5) in Findings!A2 downward, one per "
              "row. Do not fix the defects; only report them."),
    "T1_W1": ("On each monthly sheet (Jan-Jun): convert all column A dates to "
              "text format YYYY-MM-DD, fill column C with a formula computing "
              "the quarter (Q1-Q4) from the column A date, and make the header "
              "row bold. Preserve amounts."),
    "T1_W2": ("On Dept1 and Dept2: unmerge all merged cells, fill the team "
              "name down into every row of its block, and add a TOTAL row "
              "(row 11) with SUM formulas over B2:B10 and C2:C10. Preserve "
              "all numbers."),
    "T1_M1": ("Build a Variance sheet with columns Line/Budget/Actual/"
              "Variance for all 10 lines, using formulas referencing the "
              "Budget and Actuals sheets. Then list the top 3 overspend lines "
              "(actual minus budget, largest first) in Variance!A14:A16 with "
              "a 'Top overspend' label in A13."),
    "T1_M2": ("Build an Inventory sheet with columns SKU/Ending/Flag for all "
              "15 SKUs. Ending must be a formula: Starting + Receipts - "
              "Shipments for that row. Flag is 'NEGATIVE' when ending is "
              "below zero, else 'OK'."),
}

GENS = {"T1_R1": gen_r1, "T1_R2": gen_r2, "T1_W1": gen_w1,
        "T1_W2": gen_w2, "T1_M1": gen_m1, "T1_M2": gen_m2}


def main(outdir):
    out = Path(outdir)
    specs = {}
    for tid, gen in GENS.items():
        spec = gen(out)
        spec["instruction"] = INSTRUCTIONS[tid]
        specs[tid] = spec
        print(tid, "done")
    (out / "specs.json").write_text(json.dumps(specs, indent=1, sort_keys=True))


if __name__ == "__main__":
    main(sys.argv[1])
