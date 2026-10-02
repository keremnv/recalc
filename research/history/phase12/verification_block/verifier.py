"""Phase-12 frozen verifier: post-edit verification evidence block.

Pure function of (input.xlsx, candidate output.xlsx) + recalc. No gold.
See ../VERIFIER_CONTRACT.md, ../UNIFORMITY_REWRITE_FILTER.md,
../EVIDENCE_BLOCK_FORMAT.md.
"""
from __future__ import annotations

import re
import subprocess
import sys
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

RESEARCH_ROOT = Path(__file__).resolve().parents[4]
sys.path.insert(0, str(RESEARCH_ROOT / "research/history/phase11"))
from mine import (adjacent_family_breaks, blank_refs, derive_workbook,  # noqa: E402
                  expand_ref, rel_fingerprint, structural_diff)

ERRORS = {"#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!",
          "#VALUE!", "#GETTING_DATA", "#SPILL!", "#CALC!", "#FIELD!"}

CONFIG = {
    "unif_scale_cap": 30,
    "rewrite_coverage": 0.50,
    "new_coherence_members": 2,
    "sample_cap": 10,
    "chg_alone_cells": 150,
    "chg_alone_sheets": 5,
    "dep_context_refs": 3,
}


def patch_full_calc(src: Path, dst: Path) -> bool:
    try:
        zin = zipfile.ZipFile(src, "r")
        xml = zin.read("xl/workbook.xml").decode()
    except Exception:
        return False
    if "fullCalcOnLoad" not in xml:
        xml2, n = re.subn(r"<calcPr\b", '<calcPr fullCalcOnLoad="1"',
                          xml, count=1)
        if not n:
            xml2, n = re.subn(r"<workbookPr([^/]*)/>",
                              r'<workbookPr\1/><calcPr fullCalcOnLoad="1"/>',
                              xml, count=1)
        if not n:
            if "</workbook>" not in xml:
                return False
            xml2 = xml.replace("</workbook>",
                               '<calcPr fullCalcOnLoad="1"/></workbook>')
        xml = xml2
    try:
        with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = (xml.encode() if item.filename == "xl/workbook.xml"
                        else zin.read(item.filename))
                zout.writestr(item, data)
    except Exception:
        return False
    finally:
        zin.close()
    return True


def recalc(path: Path, timeout: int = 180) -> Path | None:
    """Recalculate via headless LibreOffice; None on any failure."""
    try:
        tmp = Path(tempfile.mkdtemp(prefix="p12-recalc-"))
        patched = tmp / "patched.xlsx"
        outdir = tmp / "out"
        outdir.mkdir()
        if not patch_full_calc(path, patched):
            return None
        proc = subprocess.run(
            ["soffice", "--headless", "--convert-to",
             "xlsx:Calc MS Excel 2007 XML", str(patched),
             "--outdir", str(outdir)],
            capture_output=True, timeout=timeout)
        if proc.returncode != 0:
            return None
        cand = outdir / "patched.xlsx"
        return cand if cand.is_file() else None
    except Exception:
        return None


def error_cells(path: Path) -> dict[str, str]:
    import openpyxl
    from openpyxl.utils import get_column_letter
    out: dict[str, str] = {}
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception:
        return out
    try:
        for ws in wb.worksheets:
            for r_idx, row in enumerate(ws.iter_rows(), start=1):
                for c_idx, cell in enumerate(row, start=1):
                    v = cell.value
                    if isinstance(v, str) and v in ERRORS:
                        out[f"{ws.title}!{get_column_letter(c_idx)}{r_idx}"] = v
    finally:
        wb.close()
    return out


def rewrite_filter(breaks: list[dict], pre: dict, post: dict,
                   diff: dict) -> list[dict]:
    """UNIF rewrite-aware filter v1. Returns surviving break cells."""
    if len(breaks) > CONFIG["unif_scale_cap"]:
        return []  # rule 3: too broad
    changed = set(diff.get("formulas_added", [])) | set(diff.get("formulas_changed", [])) \
        | set(diff.get("formulas_removed", [])) | set(diff.get("cells_emptied", [])) \
        | set(diff.get("cells_filled", []))
    # affected-block coverage per break cell
    by_row: dict = {}
    by_col: dict = {}
    for coord, f in pre["formulas"].items():
        by_row.setdefault((coord.split("!")[0], f["row"]), []).append(coord)
        by_col.setdefault((coord.split("!")[0], f["col"]), []).append(coord)
    # post bands for added-formula breaks (no pre formula to anchor on)
    post_by_row: dict = {}
    post_by_col: dict = {}
    for coord, f in post["formulas"].items():
        post_by_row.setdefault((coord.split("!")[0], f["row"]), []).append(coord)
        post_by_col.setdefault((coord.split("!")[0], f["col"]), []).append(coord)
    survivors = []
    for b in breaks:
        cell = b["cell"]
        f = pre["formulas"].get(cell)
        sheet = cell.split("!")[0]
        if f is None:
            # added formula: anchor the band on post-edit neighbors
            pf = (post["formulas"].get(cell) or {})
            if not pf:
                continue
            band = [c for c in post_by_row.get((sheet, pf["row"]), []) if c != cell]
            band += [c for c in post_by_col.get((sheet, pf["col"]), []) if c != cell]
        else:
            band = [c for c in by_row.get((sheet, f["row"]), []) if c != cell]
            band += [c for c in by_col.get((sheet, f["col"]), []) if c != cell]
        if band and sum(1 for c in band if c in changed) / len(band) >= CONFIG["rewrite_coverage"]:
            continue  # rule 1: broad rewrite
        # rule 2: new coherence — output fp shared with >=2 changed neighbors
        post_fp = (post["formulas"].get(cell) or {}).get("fp")
        if post_fp:
            mates = [c for c in band
                     if (post["formulas"].get(c) or {}).get("fp") == post_fp
                     and c in changed]
            if len(mates) >= CONFIG["new_coherence_members"]:
                continue
        # rule 4: edge exemption unless directly modified
        ordered = sorted(band)
        if ordered and cell not in changed and cell in (ordered[0], ordered[-1]):
            continue
        survivors.append(b)
    return survivors


def one_hop(cell: str, facts: dict) -> tuple[list[str], list[str]]:
    """(fed_by, feeds_into) one-hop refs for a signal cell."""
    fed: list[str] = []
    for ref in facts["precedents"].get(cell, []):
        fed.extend(expand_ref(ref, cell))
    feeds = [c for c, refs in facts["precedents"].items()
             if any(cell.split("!")[-1] in r or cell in expand_ref(r, c)
                    for r in refs)]
    cap = CONFIG["dep_context_refs"]
    return sorted(set(fed))[:cap], sorted(set(feeds))[:cap]


def verify(input_path: str | Path, candidate_path: str | Path) -> dict:
    """Run the frozen verifier. Never raises on workbook content; returns
    VERIFIER_UNAVAILABLE status dict on infrastructure failure."""
    report: dict = {"positive": None, "status": "OK", "signals": {},
                    "dep_context": [], "model_block": None,
                    "positive_families": []}
    try:
        pre = derive_workbook(str(input_path))
        post = derive_workbook(str(candidate_path))
        if not pre["ok"] or not post["ok"]:
            report.update({"status": "VERIFIER_UNAVAILABLE",
                           "reason": "unreadable workbook"})
            return report
        ri = recalc(Path(input_path))
        ro = recalc(Path(candidate_path))
        if ri is None or ro is None:
            report.update({"status": "VERIFIER_UNAVAILABLE",
                           "reason": "recalc failed"})
            return report
        ie = error_cells(ri)
        oe = error_cells(ro)
        new_err = sorted(set(oe) - set(ie))
        diff = structural_diff(pre, post)
        changed = set(diff["formulas_added"]) | set(diff["formulas_changed"])
        # UNIF: post-edit breaks via output families vs input
        post_breaks = adjacent_family_breaks(post)
        pre_break_cells = {b["cell"] for b in adjacent_family_breaks(pre)}
        fresh = [b for b in post_breaks if b["cell"] not in pre_break_cells]
        unif = rewrite_filter(fresh, pre, post, diff)
        # REF: became-blank
        br_pre, br_post = blank_refs(pre), blank_refs(post)
        ref_items = []
        for coord in sorted(set(br_post)):
            for cell in sorted(set(br_post[coord]) - set(br_pre.get(coord, []))):
                ref_items.append({"formula": coord, "blank": cell})
        # CHG rollup
        changed_cells = (set(diff["formulas_added"]) | set(diff["formulas_changed"])
                         | set(diff["formulas_removed"]) | set(diff["cells_emptied"])
                         | set(diff["cells_filled"]))
        sheets = {c.split("!")[0] for c in changed_cells}
        broad_chg = (len(changed_cells) >= CONFIG["chg_alone_cells"]
                     or len(sheets) >= CONFIG["chg_alone_sheets"])
        report["signals"] = {
            "ERR": {"items": [{"cell": c, "error": oe[c]} for c in new_err],
                    "count": len(new_err),
                    "sheets": len({c.split("!")[0] for c in new_err}),
                    "types": dict(Counter(oe[c] for c in new_err)),
                    "changed_overlap": sum(1 for c in new_err if c in changed)},
            "UNIF": {"items": [{"cell": b["cell"]} for b in unif],
                     "count": len(unif)},
            "REF": {"items": ref_items, "count": len(ref_items)},
            "CHG": {"changed_cells": len(changed_cells),
                    "sheets": len(sheets),
                    "regions": sorted(changed_cells)[:20]},
        }
        pos = []
        if new_err:
            pos.append("ERR")
        if unif:
            pos.append("UNIF")
        if ref_items:
            pos.append("REF")
        if not pos and broad_chg:
            pos.append("CHG")
        report["positive_families"] = pos
        report["positive"] = bool(pos)
        if pos:
            for c in new_err[:CONFIG["sample_cap"]]:
                fed, feeds = one_hop(c, post)
                report["dep_context"].append({"for": c, "fed_by": fed,
                                              "feeds": feeds})
            for b in unif[:CONFIG["sample_cap"]]:
                fed, feeds = one_hop(b["cell"], post)
                report["dep_context"].append({"for": b["cell"], "fed_by": fed,
                                              "feeds": feeds})
            report["model_block"] = render(report)
        return report
    except Exception as exc:  # fail-safe: never fabricate
        return {"positive": None, "status": "VERIFIER_UNAVAILABLE",
                "reason": f"{type(exc).__name__}", "signals": {},
                "dep_context": [], "model_block": None,
                "positive_families": []}


def render(report: dict) -> str:
    s = report["signals"]
    cap = CONFIG["sample_cap"]
    lines = ["Post-edit verification (mechanical check of the current workbook)", "",
             f"Changed cells: {s['CHG']['changed_cells']} across "
             f"{s['CHG']['sheets']} sheet(s).", ""]
    if s["ERR"]["count"]:
        lines.append(f"New formula errors ({s['ERR']['count']}):")
        for it in s["ERR"]["items"][:cap]:
            lines.append(f"- {it['cell']}: {it['error']}")
        if s["ERR"]["count"] > cap:
            extra = s["ERR"]["count"] - cap
            types = ", ".join(f"{t} x{n}" for t, n in s["ERR"]["types"].items())
            lines.append(f"- and {extra} more (grouped: {types})")
        lines.append("")
    if s["UNIF"]["count"]:
        lines.append(f"New formula-pattern breaks ({s['UNIF']['count']}):")
        for it in s["UNIF"]["items"][:cap]:
            lines.append(f"- {it['cell']} differs from the previously "
                         f"consistent row/column pattern.")
        if s["UNIF"]["count"] > cap:
            lines.append(f"- and {s['UNIF']['count'] - cap} more.")
        lines.append("")
    if s["REF"]["count"]:
        lines.append(f"Newly blank referenced cells ({s['REF']['count']}):")
        for it in s["REF"]["items"][:cap]:
            lines.append(f"- {it['formula']} now references blank {it['blank']}")
        if s["REF"]["count"] > cap:
            lines.append(f"- and {s['REF']['count'] - cap} more.")
        lines.append("")
    lines.append("This report is mechanical evidence about the current "
                 "workbook state.")
    lines.append("Review it before deciding whether to submit or revise.")
    return "\n".join(lines)
