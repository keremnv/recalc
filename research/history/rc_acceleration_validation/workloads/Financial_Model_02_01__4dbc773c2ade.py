
import openpyxl
wb = openpyxl.load_workbook('input.xlsx')
ws = wb['Realization & EBIT Per KG']
print(ws.max_row, ws.max_column)
for r in range(1, ws.max_row+1):
    row=[]
    for c in range(1, ws.max_column+1):
        v = ws.cell(row=r,column=c).value
        if v is not None: row.append((ws.cell(row=r,column=c).coordinate, v))
    if row: print('ROW',r, row)
