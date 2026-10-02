import openpyxl
wb = openpyxl.load_workbook('input.xlsx')
ws = wb['Income Statement']
print("dims:", ws.dimensions, ws.max_row, ws.max_column)
for r in range(2, ws.max_row+1):
    b = ws.cell(r,2).value
    c = ws.cell(r,3).value
    i = ws.cell(r,9).value
    print(r, repr(b), '|C:', repr(c)[:60], '|I:', repr(i)[:60])
