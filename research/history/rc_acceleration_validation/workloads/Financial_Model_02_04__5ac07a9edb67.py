
import openpyxl
wb = openpyxl.load_workbook('input.xlsx')
ws = wb['IS,BS,CF']
for r in range(70,95):
    v = ws.cell(row=r,column=2).value
    if v: print(r, repr(v), [ws.cell(row=r,column=c).value for c in range(5,12)])
