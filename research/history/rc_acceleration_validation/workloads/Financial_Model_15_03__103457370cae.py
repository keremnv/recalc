
import openpyxl
wb=openpyxl.load_workbook('input.xlsx')
ws=wb['IS_BS_CF']
print('dims', ws.dimensions, ws.max_row, ws.max_column)
for r in range(1,60):
    a=ws.cell(row=r,column=1).value
    print(r, repr(a))
