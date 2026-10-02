import openpyxl
from openpyxl.utils import get_column_letter as gl
wb=openpyxl.load_workbook('input.xlsx')
ws=wb['Cashflow (Monthly)']
print('maxcol',ws.max_column)
for r in range(12,92):
    b=ws.cell(r,2).value; c=ws.cell(r,3).value; d=ws.cell(r,4).value
    filled=[cc for cc in range(10,ws.max_column+1) if ws.cell(r,cc).value is not None]
    rng=f'{gl(filled[0])}-{gl(filled[-1])} n={len(filled)}' if filled else 'EMPTY'
    print(r,'|',b,'|',c,'|',d,'|',rng)
