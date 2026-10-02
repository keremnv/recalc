
import openpyxl
wb=openpyxl.load_workbook('input.xlsx')
wc=wb['Working Capital']
print('WC6',[wc.cell(6,c).value for c in range(8,35)])
print('WC8',[wc.cell(8,c).value for c in range(15,35)])
print('WC7',[wc.cell(7,c).value for c in range(15,35)])
print('WC29',[wc.cell(29,c).value for c in range(8,14)])
print('WC30',[wc.cell(30,c).value for c in range(2,14)])
ra=wb['Ratio Analysis']
print('RA11',[ra.cell(11,c).value for c in range(2,13)])
bs=wb['BS']
print('BS row5 I-M',[bs.cell(5,c).value for c in range(9,14)])
a=wb['Assets Sch.']
print('A4',[a.cell(4,c).value for c in range(8,35)])
print('A5',[a.cell(5,c).value for c in range(8,14)])
