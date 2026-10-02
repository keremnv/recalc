import openpyxl
wb = openpyxl.load_workbook('input.xlsx', data_only=False)
ws = wb['Capex and Debt Assumptions ']
for r in range(2, 160):
    vals = []
    for c in range(1, 20):
        v = ws.cell(row=r, column=c).value
        if v is not None:
            vals.append((ws.cell(row=r, column=c).coordinate, repr(v)[:80]))
    if vals:
        print(r, vals)
