
import openpyxl
wb=openpyxl.load_workbook('input.xlsx')
ws=wb['ValuationModel']
for r in [6,7,8]:
    print(r, ws.cell(row=r,column=3).value.text)
ws2=wb['RentRoll']
for r in range(14,21):
    print([ws2.cell(row=r,column=c).value for c in range(2,8)])
