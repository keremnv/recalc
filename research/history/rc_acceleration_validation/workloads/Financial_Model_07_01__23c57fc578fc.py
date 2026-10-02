import openpyxl
wb=openpyxl.load_workbook('input.xlsx')
ws=wb['Cashflow (Monthly)']
print(ws.max_row, ws.max_column)
for r in range(1,40):
    vals=[]
    for c in range(1,60):
        v=ws.cell(r,c).value
        if v is not None:
            vals.append((ws.cell(r,c).coordinate,repr(v)[:70]))
    if vals: print(r,vals[:40])
