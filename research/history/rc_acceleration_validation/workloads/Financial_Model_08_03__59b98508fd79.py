import openpyxl
wb = openpyxl.load_workbook('input.xlsx', data_only=False)
ws = wb['DCF Valuation']
for r in range(30, 80):
    vals = []
    for c in range(1, 15):
        v = ws.cell(row=r, column=c).value
        if v is not None:
            vals.append((ws.cell(row=r, column=c).coordinate, repr(v)[:110]))
    if vals:
        print(r, vals)
