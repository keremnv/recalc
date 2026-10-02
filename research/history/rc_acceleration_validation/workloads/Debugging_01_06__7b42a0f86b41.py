import openpyxl
wb=openpyxl.load_workbook('input.xlsx')
ws=wb['Ex 5 - M&A']
for r in range(28,45):
    print(r,[ws.cell(row=r,column=c).value for c in [2,3,4,10,17,24]])
for r in range(74,85):
    print(r,[ws.cell(row=r,column=c).value for c in [2,3,4,10,17,24]])
print('--- Ex10 ---')
ws=wb['Ex 10 - Balance Sheet']
for r in range(1,34):
    print(r,[ws.cell(row=r,column=c).value for c in [2,3,4,5,6,7]])
