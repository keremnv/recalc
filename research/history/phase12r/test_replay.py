"""Mechanical tests, never historical outcome tuning."""
import importlib.util,json,tempfile,unittest,zipfile
from pathlib import Path
import openpyxl
spec=importlib.util.spec_from_file_location('replay',Path(__file__).with_name('replay.py'));r=importlib.util.module_from_spec(spec);spec.loader.exec_module(r)
class ReplayTest(unittest.TestCase):
 def setUp(self):self.t=tempfile.TemporaryDirectory();self.p=Path(self.t.name)
 def tearDown(self):self.t.cleanup()
 def book(self,name):
  p=self.p/name;w=openpyxl.Workbook();w.active['A1']='=1+2';w.active['A2']='="hello"';w.active['A3']='=1/0';w.active['B1']=5;w.save(p);return p
 def cache(self,p,vals):
  with zipfile.ZipFile(p) as z:data={n:z.read(n) for n in z.namelist()}
  root=r.ET.fromstring(data['xl/worksheets/sheet1.xml'])
  for c in root.iter(r.Q+'c'):
   if c.attrib['r'] in vals:
    typ,val=vals[c.attrib['r']];c.set('t',typ);v=c.find(r.Q+'v');v.text=val
  data['xl/worksheets/sheet1.xml']=r.ET.tostring(root)
  with zipfile.ZipFile(p,'w') as z:
   for n,d in data.items():z.writestr(n,d)
 def test_transplant_typed_cache_and_no_semantic_mutation(self):
  a=self.book('a.xlsx');b=self.book('b.xlsx');self.cache(b,{'A1':('n','3'),'A2':('str','hello'),'A3':('e','#DIV/0!')});dst=self.p/'w.xlsx';r.transplant(a,b,dst)
  c0,raw0,_=r.xml_cells(a);cb,_,_=r.xml_cells(b);cw,raww,_=r.xml_cells(dst)
  self.assertEqual(cb,cw);self.assertEqual(raw0,raww);self.assertEqual(r.semantic_snapshot(a),r.semantic_snapshot(dst))
  with zipfile.ZipFile(a) as za,zipfile.ZipFile(dst) as zw:
   for n in za.namelist():
    if not n.startswith('xl/worksheets/'):self.assertEqual(za.read(n),zw.read(n))
 def test_noop_and_empty_string_missing_distinction(self):
  a=self.book('a.xlsx');self.cache(a,{'A2':('str','')});dst=self.p/'w.xlsx';r.transplant(a,a,dst);self.assertEqual(r.xml_cells(a)[:2],r.xml_cells(dst)[:2]);c=r.xml_cells(a)[0];self.assertFalse(c['Sheet!A1']['present']);self.assertTrue(c['Sheet!A2']['present'])
 def test_formula_normalization_preserves_literals_and_absolute_refs(self):
  self.assertEqual(r.norm_function_case('=sum(A1,"sum(foo)")'), '=SUM(A1,"sum(foo)")');self.assertNotEqual(r.norm_function_case('=SUM(A1)'),r.norm_function_case('=SUM($A$1)'));self.assertNotEqual(r.norm_function_case('="a"'),r.norm_function_case('="A"'))
 def test_shared_strings_resolved(self):
  a=self.book('a.xlsx');self.cache(a,{'A2':('s','0')})
  with zipfile.ZipFile(a,'a') as z:z.writestr('xl/sharedStrings.xml',f'<sst xmlns="{r.NS}"><si><t>hello</t></si></sst>')
  c=r.xml_cells(a)[0]['Sheet!A2'];self.assertEqual(c,{'present':True,'type':'str','value':'hello'})
 def test_semantic_change_detected_with_same_formula(self):
  a=self.book('a.xlsx');b=self.book('b.xlsx');w=openpyxl.load_workbook(b);w.active['B1']=6;w.save(b)
  aa=r.semantic_snapshot(a);bb=r.semantic_snapshot(b);self.assertEqual(r.identity(aa,bb,r.xml_cells(a)[1],r.xml_cells(b)[1]),'FORMULAS_BYTE_IDENTICAL');self.assertIn('literals',r.semantic_diff(aa,bb,'Template'))
 def test_volatile_detection_outside_strings(self):
  a=self.book('a.xlsx');w=openpyxl.load_workbook(a);w.active['A1']='="NOW()"';w.save(a);self.assertFalse(r.unsupported(a,r.xml_cells(a)[1],r.semantic_snapshot(a)));w.active['A1']='=NOW()';w.save(a);self.assertIn('VOLATILE_FUNCTION',r.unsupported(a,r.xml_cells(a)[1],r.semantic_snapshot(a)))
 def test_boolean_cache(self):
  a=self.book('a.xlsx');b=self.book('b.xlsx');self.cache(b,{'A1':('b','1')});dst=self.p/'w.xlsx';r.transplant(a,b,dst);self.assertEqual(r.xml_cells(dst)[0]['Sheet!A1']['type'],'b')
if __name__=='__main__':unittest.main()
