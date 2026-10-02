#!/usr/bin/env python3
"""Phase13 analysis: map assessed-root mismatches to loss boundaries.

Checkpointed + task-grouped: input/gold load once per task, released after.
Usage: analyze_footprints.py [fullerr.json] [budget_runs]
Rerun until 'pending: 0'.
"""
import json, sys, os, gc
from collections import Counter
import openpyxl

DATA = 'benchmark-data/SpreadsheetBench-2/data'
DS = {c: {d['id']: d for d in json.load(open(f'{DATA}/{c}/dataset.json'))}
      for c in ('Template', 'Debugging', 'Financial_Model')}
scores = {x['id']: x for x in json.load(open('research/history/phase13/working/popB_scores.json'))}
CKPT = 'research/history/phase13/working/popB_footprints.json'
FT = 'research/history/phase13/working/popB_formulatext_samples.json'

def load_ig(cat, tid):
    d = DS[cat][tid]
    pi = f"{DATA}/{cat}/{d['spreadsheet_path']}"
    pg = f"{DATA}/{cat}/{d['golden_response_path']}"
    return (openpyxl.load_workbook(pi, data_only=False),
            openpyxl.load_workbook(pg, data_only=False))

def main(path, budget):
    det = json.load(open(path))
    roots = json.load(open('research/history/phase13/working/popB_roots.json'))
    out = json.load(open(CKPT)) if os.path.exists(CKPT) else {}
    fts = json.load(open(FT)) if os.path.exists(FT) else []
    bytask = {}
    for id, d in det.items():
        if id not in out:
            bytask.setdefault(d['task'], []).append((id, d))
    done = 0
    for task in sorted(bytask):
        if done >= budget:
            break
        items = bytask[task]
        cat, tid = task.split(':')[0], task.split(':')[1]
        try:
            wi, wg = load_ig(cat, tid)
        except Exception as e:
            for id, d in items:
                out[id] = {'fatal': f'load-ig: {e}'[:100]}
            continue
        goldmod = set()
        for ws in wg.worksheets:
            if ws.title not in wi.sheetnames:
                continue
            wsi = wi[ws.title]
            for row in ws.iter_rows():
                for c in row:
                    if c.value != wsi[c.coordinate].value:
                        goldmod.add((ws.title, c.coordinate))
        for id, d in items:
            if done >= budget:
                break
            if d.get('fatal'):
                out[id] = {'fatal': d['fatal']}
                continue
            try:
                wo = openpyxl.load_workbook(scores[id]['V0'], data_only=False)
            except Exception as e:
                out[id] = {'fatal': f'load-v0: {e}'[:100]}
                continue
            outchg = set()
            for ws in wg.worksheets:
                if ws.title not in wi.sheetnames or ws.title not in wo.sheetnames:
                    continue
                wsi, wso = wi[ws.title], wo[ws.title]
                for row in ws.iter_rows():
                    for c in row:
                        if wso[c.coordinate].value != wsi[c.coordinate].value:
                            outchg.add((ws.title, c.coordinate))
            r = roots[id]
            rootset = set((s, c) for s, c, *_ in r['roots'])
            miss_p, miss_u = 0, 0
            for m in d['mm']:
                if m['cls'] != 'MISSING_OUTPUT_NONE' or (m['sheet'], m['cell']) not in rootset:
                    continue
                rown = ''.join(ch for ch in m['cell'] if ch.isdigit())
                neigh = [g for g in goldmod if g[0] == m['sheet']
                         and ''.join(ch for ch in g[1] if ch.isdigit()) == rown]
                if [g for g in neigh if g in outchg]:
                    miss_p += 1
                else:
                    miss_u += 1
            reg_ov, reg_same = 0, 0
            for m in d['mm']:
                if m['kind'] != 'REG' or (m['sheet'], m['cell']) not in rootset:
                    continue
                if m['cls'] == 'OVERFILL':
                    reg_ov += 1
                    continue
                try:
                    iv = wi[m['sheet']][m['cell']].value
                    ov = wo[m['sheet']][m['cell']].value
                except Exception:
                    continue
                if ov != iv:
                    reg_ov += 1
                else:
                    reg_same += 1
            for m in d['mm']:
                if m['cls'] == 'FORMULA_TEXT' and len(fts) < 40:
                    fts.append((d['task'], m['sheet'], m['cell'], m['ans_f'], m['out_f'], m.get('mode')))
            out[id] = {'task': d['task'], 'n_goldmod': len(goldmod), 'n_outchg': len(outchg),
                       'recall': round(len(goldmod & outchg) / len(goldmod), 4) if goldmod else None,
                       'precision': round(len(goldmod & outchg) / len(outchg), 4) if outchg else None,
                       'miss_partial': miss_p, 'miss_untouched': miss_u,
                       'reg_overwrite': reg_ov, 'reg_same_as_input': reg_same}
            done += 1
            del wo
        del wi, wg, goldmod
        gc.collect()
        json.dump(out, open(CKPT, 'w'))
        json.dump(fts, open(FT, 'w'))
        print(f'task {task}: done, checkpoint {len(out)} runs', flush=True)
    pend = sum(1 for id in det if id not in out)
    print(f'pending: {pend}')
    if pend == 0:
        vals = [v for v in out.values() if isinstance(v, dict)]
        print('miss partial/untouched:',
              sum(v.get('miss_partial', 0) for v in vals),
              sum(v.get('miss_untouched', 0) for v in vals))
        print('reg overwrite/same-as-input:',
              sum(v.get('reg_overwrite', 0) for v in vals),
              sum(v.get('reg_same_as_input', 0) for v in vals))
        rec = sorted(v['recall'] for v in vals if v.get('recall') is not None)
        print('recall q0/25/50/75/100:', rec[0], rec[len(rec)//4], rec[len(rec)//2], rec[3*len(rec)//4], rec[-1])
        pr = sorted(v['precision'] for v in vals if v.get('precision') is not None)
        print('precision q0/25/50/75/100:', pr[0], pr[len(pr)//4], pr[len(pr)//2], pr[3*len(pr)//4], pr[-1])

if __name__ == '__main__':
    main(sys.argv[1] if len(sys.argv) > 1 else 'research/history/phase13/working/popB_fullerr.json',
         int(sys.argv[2]) if len(sys.argv) > 2 else 10**9)
