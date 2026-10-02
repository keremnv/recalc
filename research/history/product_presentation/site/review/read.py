import sys
import time

import openpyxl

t0 = time.perf_counter()
wb = openpyxl.load_workbook("model.xlsx")
print("sheets:", wb.sheetnames)
ws = wb["Forecast"]
print("dimensions:", ws.dimensions)
print("H402 reported total:", ws["H402"].value)
total = 0
for r in range(2, 402):
    total += ws.cell(row=r, column=8).value
print("column H recomputed:", total)
regions = wb["Regions"]
print("North lead:", regions["B2"].value)
wb.close()
t1 = time.perf_counter()
print(f"read phase: {(t1 - t0) * 1000:.1f} ms", file=sys.stderr)
