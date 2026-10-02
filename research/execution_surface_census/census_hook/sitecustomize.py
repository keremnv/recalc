"""Operation-census import hook (measurement only, env-gated).

Active only when RECALC_CENSUS=1. Wraps openpyxl entry points with
count/timing decorators that delegate with identical arguments and return
identical objects. Emits aggregates at exit to RECALC_CENSUS_OUT.

Chains to the recalc bootstrap sitecustomize when RECALC_RUN_CONTEXT is set
(product runs): hooks wrap the original loader first, then the product
bootstrap wraps the wrapped loader, so reference loads are still counted
while direct-served loads bypass openpyxl (correct: they are Recalc-owned).
"""
import atexit
import json
import os
import sys
import time
from collections import defaultdict

_HERE = os.path.dirname(os.path.abspath(__file__))
_OUT = os.environ.get("RECALC_CENSUS_OUT", "")
_MODE = os.environ.get("RECALC_CENSUS_MODE", "time")  # time | count

_counts = defaultdict(int)
_times = defaultdict(float)
_loads = []
_saves = []
_errors = []
# Attribution context: openpyxl's reader/writer touches the same attributes
# the script does. While inside load_workbook/save, families are prefixed so
# parse/serialize-internal traffic never pollutes script-driven counts.
_CTX = []  # stack of "parse_internal" | "serialize_internal" | None


def _bump(family, dt):
    if _CTX and _CTX[-1]:
        family = _CTX[-1] + ":" + family
    _counts[family] += 1
    if _MODE == "time":
        _times[family] += dt


def _wrap_call(family, fn):
    def inner(*a, **k):
        t = time.perf_counter()
        try:
            return fn(*a, **k)
        finally:
            _bump(family, time.perf_counter() - t)
    inner.__name__ = getattr(fn, "__name__", family)
    return inner


def _wrap_prop(cls, name, get_family, set_family=None):
    import inspect
    try:
        orig = inspect.getattr_static(cls, name)
    except Exception:
        return
    if isinstance(orig, property):
        fget, fset, fdel, doc = orig.fget, orig.fset, orig.fdel, orig.__doc__
    elif type(orig).__name__ == "member_descriptor":
        # __slots__ storage: keep the descriptor and delegate to it.
        fget = lambda self: orig.__get__(self, cls)  # noqa: E731
        fset = (lambda self, v: orig.__set__(self, v)) if hasattr(orig, "__set__") else None  # noqa: E731
        fdel, doc = None, None
    else:
        return

    def new_get(self):
        t = time.perf_counter()
        try:
            return fget(self)
        finally:
            _bump(get_family, time.perf_counter() - t)

    def new_set(self, value):
        t = time.perf_counter()
        try:
            fset(self, value)
        finally:
            _bump(set_family or get_family, time.perf_counter() - t)

    try:
        setattr(cls, name, property(new_get, new_set if fset else None,
                                    fdel, doc))
    except Exception as exc:  # never break the wrapped library
        _errors.append({"where": cls.__name__ + "." + name, "error": repr(exc)})


def _wrap_method(cls, name, family):
    try:
        orig = getattr(cls, name, None)
    except Exception:
        return
    if not callable(orig):
        return
    try:
        setattr(cls, name, _wrap_call(family, orig))
    except Exception as exc:
        _errors.append({"where": cls.__name__ + "." + name, "error": repr(exc)})


def _install():
    import openpyxl
    from openpyxl.workbook.workbook import Workbook
    from openpyxl.worksheet.worksheet import Worksheet
    from openpyxl.cell.cell import Cell

    _orig_load = openpyxl.load_workbook

    def load_rec(filename, *a, **k):
        t = time.perf_counter()
        _CTX.append("parse_internal")
        try:
            return _orig_load(filename, *a, **k)
        finally:
            _CTX.pop()
            dt = time.perf_counter() - t
            _bump("load_workbook", dt)
            try:
                _loads.append({"path": str(filename),
                               "kwargs": sorted(str(x) for x in k),
                               "duration_s": dt})
            except Exception:
                pass

    openpyxl.load_workbook = load_rec

    for name, fam in [("__getitem__", "sheet_lookup"),
                      ("close", "close"),
                      ("create_sheet", "sheet_create"),
                      ("remove", "sheet_delete")]:
        _wrap_method(Workbook, name, fam)
    _wrap_prop(Workbook, "sheetnames", "sheet_enumeration")
    _wrap_prop(Workbook, "worksheets", "workbook_iteration")
    _wrap_prop(Workbook, "active", "active_sheet")

    for name, fam in [("__getitem__", "cell_getitem"),
                      ("__setitem__", "cell_assign_getitem"),
                      ("__iter__", "worksheet_iteration"),
                      ("cell", "cell_call"),
                      ("iter_rows", "iter_rows"),
                      ("iter_cols", "iter_cols"),
                      ("calculate_dimension", "dimensions"),
                      ("append", "row_append"),
                      ("delete_rows", "row_ops"),
                      ("delete_cols", "row_ops")]:
        _wrap_method(Worksheet, name, fam)
    _wrap_prop(Worksheet, "values", "values")
    _wrap_prop(Worksheet, "dimensions", "dimensions")
    _wrap_prop(Worksheet, "max_row", "bounds")
    _wrap_prop(Worksheet, "max_column", "bounds")
    _wrap_prop(Worksheet, "merged_cells", "merged_access")
    _wrap_prop(Worksheet, "tables", "tables")
    _wrap_prop(Worksheet, "defined_names", "defined_names")

    _wrap_prop(Cell, "value", "cell_value_get", "cell_value_set")
    _wrap_prop(Cell, "data_type", "data_type_get")

    # Read-only worksheet variants share method names; wrap if present.
    try:
        from openpyxl.worksheet._read_only import ReadOnlyWorksheet
        for name, fam in [("iter_rows", "iter_rows"),
                          ("iter_cols", "iter_cols"),
                          ("calculate_dimension", "dimensions"),
                          ("__getitem__", "cell_getitem")]:
            _wrap_method(ReadOnlyWorksheet, name, fam)
        _wrap_prop(ReadOnlyWorksheet, "values", "values")
        _wrap_prop(ReadOnlyWorksheet, "max_row", "bounds")
        _wrap_prop(ReadOnlyWorksheet, "max_column", "bounds")
    except Exception as exc:
        _errors.append({"where": "ReadOnlyWorksheet", "error": repr(exc)})

    # Save recorder: writer-internal traffic is attributed to serialization,
    # the save call itself to the script.
    _orig_save = Workbook.save

    def save_rec(self, filename, *a, **k):
        t = time.perf_counter()
        _CTX.append("serialize_internal")
        try:
            return _orig_save(self, filename, *a, **k)
        finally:
            _CTX.pop()
            dt = time.perf_counter() - t
            _bump("save", dt)
            try:
                _saves.append({"path": str(filename), "duration_s": dt})
            except Exception:
                pass

    Workbook.save = save_rec


def _emit():
    if not _OUT:
        return
    try:
        payload = {"mode": _MODE,
                   "families": {k: {"count": _counts[k],
                                     "total_s": _times.get(k, 0.0)}
                                for k in sorted(_counts)},
                   "loads": _loads,
                   "saves": _saves,
                   "hook_errors": _errors}
        with open(_OUT, "w") as f:
            json.dump(payload, f)
    except Exception:
        pass


def _chain_recalc_bootstrap():
    ctx = os.environ.get("RECALC_RUN_CONTEXT")
    if not ctx:
        return
    try:
        import importlib.util
        spec = importlib.util.find_spec("recalc_agent")
        if spec is None or not spec.origin:
            return
        boot = os.path.join(os.path.dirname(spec.origin), "_bootstrap",
                            "sitecustomize.py")
        with open(boot) as f:
            code = compile(f.read(), boot, "exec")
        g = {"__name__": "sitecustomize", "__file__": boot}
        exec(code, g)
    except Exception as exc:
        _errors.append({"where": "recalc_bootstrap_chain", "error": repr(exc)})


if os.environ.get("RECALC_CENSUS") == "1":
    try:
        _install()
    except Exception as exc:
        _errors.append({"where": "install", "error": repr(exc)})
    atexit.register(_emit)
    _chain_recalc_bootstrap()
