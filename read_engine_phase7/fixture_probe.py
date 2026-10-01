"""Adversarial merged-child observations, before Phase-7 scored timing."""
from __future__ import annotations

import json
import os
import shutil
import subprocess
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Border, Side, Font, PatternFill, Alignment, Protection

from librecalc_agent._frozen.eligibility import classify
from read_engine_phase7.certificate import certify

ROOT = Path(__file__).resolve().parents[1]
HERE = Path(__file__).resolve().parent
VENV = Path("/tmp/librecalc-hygiene-rc-v55taghk/venv/bin/python")
OBSERVER = ROOT / "read_engine_phase6/observer"
OVERLAY = HERE / "overlay"


CASES = [
    ("covered_value", "print(repr(ws.cell(2,2).value))", "EXACT_DIRECT", 1),
    ("covered_dtype", "print(repr(ws.cell(2,2).data_type))", "EXACT_DIRECT", 1),
    ("covered_coordinate", "print(ws.cell(2,2).coordinate)", "EXACT_DIRECT", 1),
    ("covered_row", "print(ws.cell(2,2).row)", "EXACT_DIRECT", 1),
    ("covered_column", "print(ws.cell(2,2).column)", "EXACT_DIRECT", 1),
    ("scalar_type", "print(type(ws.cell(2,2).value).__name__)", "EXACT_DIRECT", 1),
    ("scalar_repr", "print(repr(ws.cell(2,2).value))", "EXACT_DIRECT", 1),
    ("repeat_same_child", "print(ws.cell(2,2).value, ws.cell(2,2).data_type)", "EXACT_DIRECT", 2),
    ("second_merge", "print(ws.cell(6,5).coordinate, ws.cell(6,5).value)", "EXACT_DIRECT", 2),
    ("anchor_value", "print(ws.cell(1,1).value)", "EXACT_DIRECT", 0),
    ("unmerged_value", "print(ws.cell(8,8).value)", "EXACT_DIRECT", 0),
    ("merged_object_type", "print(type(ws.cell(2,2)).__name__)", "REQUIRES_REFERENCE", 0),
    ("merged_isinstance", "from openpyxl.cell.cell import MergedCell\nprint(isinstance(ws.cell(2,2), MergedCell))", "REQUIRES_REFERENCE", 0),
    ("merged_repr", "print(repr(ws.cell(2,2)))", "REQUIRES_REFERENCE", 0),
    ("merged_str", "print(str(ws.cell(2,2)))", "REQUIRES_REFERENCE", 0),
    ("merged_bool", "print(bool(ws.cell(2,2)))", "REQUIRES_REFERENCE", 0),
    ("merged_equality", "a=ws.cell(2,2)\nb=ws.cell(2,2)\nprint(a==b)", "REQUIRES_REFERENCE", 0),
    ("merged_identity", "a=ws.cell(2,2)\nb=ws.cell(2,2)\nprint(a is b)", "REQUIRES_REFERENCE", 0),
    ("merged_hash", "a=ws.cell(2,2)\nprint(hash(a)==hash(a))", "REQUIRES_REFERENCE", 0),
    ("style_id", "print(ws.cell(2,2).style_id)", "REQUIRES_REFERENCE", 0),
    ("font", "print(ws.cell(2,2).font.name)", "REQUIRES_REFERENCE", 0),
    ("fill", "print(ws.cell(2,2).fill.patternType)", "REQUIRES_REFERENCE", 0),
    ("border", "print(ws.cell(2,2).border.left.style)", "REQUIRES_REFERENCE", 0),
    ("alignment", "print(ws.cell(2,2).alignment.horizontal)", "REQUIRES_REFERENCE", 0),
    ("number_format", "print(ws.cell(2,2).number_format)", "REQUIRES_REFERENCE", 0),
    ("protection", "print(ws.cell(2,2).protection.locked)", "REQUIRES_REFERENCE", 0),
    ("parent_title", "print(ws.cell(2,2).parent.title)", "REQUIRES_REFERENCE", 0),
    ("parent_identity", "print(ws.cell(2,2).parent is ws)", "UNSUPPORTED/UNPROVEN", 0),
    ("arbitrary_attribute", "try:\n print(ws.cell(2,2).made_up_attribute)\nexcept AttributeError:\n print('AttributeError')", "REQUIRES_REFERENCE", 0),
    ("stored_then_observed", "x=ws.cell(2,2)\nprint(x.value)", "REQUIRES_REFERENCE", 0),
    ("list_insertion", "cells=[ws.cell(2,2)]\nprint(cells[0].value)", "REQUIRES_REFERENCE", 0),
    ("dict_insertion", "cells={'x':ws.cell(2,2)}\nprint(cells['x'].value)", "REQUIRES_REFERENCE", 0),
    ("helper_function", "def use(x):\n return x.value\nprint(use(ws.cell(2,2)))", "REQUIRES_REFERENCE", 0),
    ("conditional_object", "if True:\n print(ws.cell(2,2).value)\nelse:\n print(type(ws.cell(2,2)))", "REQUIRES_REFERENCE", 0),
    ("worksheet_subscript", "print(ws['B2'].value)", "REQUIRES_REFERENCE", 0),
    ("method_alias", "cell=ws.cell\nprint(cell(2,2).value)", "REQUIRES_REFERENCE", 0),
    ("dynamic_getattr", "print(getattr(ws.cell(2,2), 'value'))", "REQUIRES_REFERENCE", 0),
]


def make_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Main"
    ws["A1"] = "anchor"
    ws["D5"] = "second"
    ws["H8"] = 42
    ws.merge_cells("A1:C3")
    ws.merge_cells("D5:E6")
    ws["B2"].border = Border(left=Side(style="thin"))
    ws["B2"].font = Font(name="Arial", bold=True)
    ws["B2"].fill = PatternFill(patternType="solid", fgColor="FFFF00")
    ws["B2"].alignment = Alignment(horizontal="center")
    ws["B2"].protection = Protection(locked=False)
    wb.save(path)


def run_case(name: str, body: str, declared: str, count: int, workbook: Path) -> dict:
    source = "import openpyxl\nwb=openpyxl.load_workbook('input.xlsx')\nws=wb['Main']\n" + body + "\n"
    admission = classify(source)
    certificate = certify(source, admission)
    base = HERE / "runs/fixtures" / name
    shutil.rmtree(base, ignore_errors=True)
    results = {}
    for arm in ("PY", "H0", "H1"):
        root = base / arm.lower()
        work = root / "work"
        work.mkdir(parents=True)
        shutil.copyfile(workbook, work / "input.xlsx")
        (work / "workload.py").write_text(source)
        env = {k: v for k, v in os.environ.items()
               if k not in {"PYTHONPATH", "LIBRECALC_CONFIG", "LIBRECALC_RUN_CONTEXT",
                            "READ_ENGINE_PHASE3_CONTEXT", "READ_ENGINE_PHASE6_CONTEXT"}
               and not k.startswith("CANDIDATE_A_")}
        env["XDG_CACHE_HOME"] = str((root / "xdg").resolve())
        env["PYTHONDONTWRITEBYTECODE"] = "1"
        if arm == "PY":
            argv = [str(VENV), str((work / "workload.py").resolve())]
        else:
            selected_root = ROOT if arm == "H0" else OVERLAY
            argv = [str(OBSERVER), str(VENV), str((work / "workload.py").resolve()),
                    str(work.resolve()), str((root / "cache").resolve()),
                    str((root / "runs").resolve()), str(selected_root.resolve())]
        proc = subprocess.run(argv, cwd=work, env=env, capture_output=True, timeout=180)
        events = []
        setup = {}
        receipt = None
        if arm != "PY":
            last = root / "runs/last_run.json"
            if last.exists():
                run = Path(json.loads(last.read_text())["run_dir"])
                event_path = run / "runtime_events.jsonl"
                if event_path.exists():
                    events = [json.loads(x) for x in event_path.read_text().splitlines()]
                if (run / "setup.json").exists():
                    setup = json.loads((run / "setup.json").read_text())
                if (run / "observer_receipt.json").exists():
                    receipt = json.loads((run / "observer_receipt.json").read_text())
        results[arm] = {"exit_code": proc.returncode, "stdout": proc.stdout.decode(errors="replace"),
                        "stderr": proc.stderr.decode(errors="replace"),
                        "events": {key: sum(e["event"] == key for e in events)
                                   for key in ("merged_child_contact", "merged_child_direct",
                                               "merged_child_reference", "reference_parse", "direct_served_load")},
                        "reference_reasons": [e.get("reason") for e in events if e["event"] == "reference_parse"],
                        "artifact_status": [e.get("status") for e in setup.get("artifacts", {}).values()],
                        "certificate": setup.get("merged_certificate"),
                        "capture_ok": receipt is not None and receipt.get("capture_helper_exit") == 0}
    py, h0, h1 = (results[a] for a in ("PY", "H0", "H1"))
    h1_new_difference = (h1["exit_code"], h1["stdout"], h1["stderr"]) != (h0["exit_code"], h0["stdout"], h0["stderr"])
    oracle_equal = (h1["exit_code"], h1["stdout"], h1["stderr"]) == (py["exit_code"], py["stdout"], py["stderr"])
    direct = h1["events"]["merged_child_direct"]
    if declared == "EXACT_DIRECT":
        passed = oracle_equal and certificate["certified"] and direct == count and h1["events"]["reference_parse"] == 0
    elif declared == "REQUIRES_REFERENCE":
        passed = oracle_equal and not certificate["certified"] and direct == 0 and h1["events"]["reference_parse"] >= 1
    else:
        passed = (not certificate["certified"] and direct == 0 and not h1_new_difference
                  and h1["events"]["reference_parse"] >= 1)
    return {"case": name, "classification": declared, "expected_direct_contacts": count,
            "admission": admission["decision"], "certificate": certificate,
            "oracle_equal": oracle_equal, "h0_oracle_equal": (h0["exit_code"], h0["stdout"], h0["stderr"]) ==
            (py["exit_code"], py["stdout"], py["stderr"]),
            "new_h1_difference": h1_new_difference, "passed": passed, "results": results}


def main() -> None:
    ledger = HERE / "merged_cell_fixtures.jsonl"
    if ledger.exists():
        raise RuntimeError("Fixture ledger already exists")
    fixture = HERE / "runs/fixtures/oracle.xlsx"
    fixture.parent.mkdir(parents=True, exist_ok=True)
    make_workbook(fixture)
    with ledger.open("w") as out:
        for idx, (name, body, classification, count) in enumerate(CASES, 1):
            row = run_case(name, body, classification, count, fixture)
            out.write(json.dumps(row, sort_keys=True) + "\n")
            out.flush()
            print(f"fixture {idx}/{len(CASES)} {name} {row['classification']} passed={row['passed']}", flush=True)
            if not row["passed"]:
                raise RuntimeError(f"Semantic fixture failed: {name}")


if __name__ == "__main__":
    main()
