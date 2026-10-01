#!/usr/bin/env python3
"""Counterfactual workbook experiments for the composition-closure probe.

Gold edits are evaluator interventions. They are applied by the harness to
measure what coordination is *necessary*; no model ever sees them.

Each variant is written with the neutral archive-level writer, recalculated in
one batched LibreOffice session, and scored by the official comparison, so a
variant differs from the pristine input only in the cells it deliberately edits.
"""
from __future__ import annotations
import json
import shutil
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
import composition_closure as cc
import end_to_end_composition_probe as old
from xlsx_cell_writer import write_cells

sys.path.insert(0, str(cc.EVAL_DIR))
OUT = old.MECHANICAL / "composition-closure-probe"
WORK = OUT / "variants"


def gold_edit(task: str, cell: cc.Cell) -> dict | None:
    for e in cc.population_E(task):
        if (e["sheet"], e["row"], e["col"]) == cell:
            return e
    return None


def build_variant(task: str, name: str, cells: set[cc.Cell]) -> Path:
    """Write the pristine input with exactly `cells` set to their gold content."""
    inp, _ = cc.workbook_paths(task)
    dest_dir = WORK / task / name
    dest_dir.mkdir(parents=True, exist_ok=True)
    dest = dest_dir / f"{task}_output.xlsx"
    edits = []
    for c in sorted(cells):
        e = gold_edit(task, c)
        if e is None:
            continue
        edits.append({"sheet": e["sheet"], "address": e["address"], "formula": e["golden_payload"]})
    audit = write_cells(inp, dest, edits)
    (dest_dir / "audit.json").write_text(json.dumps(
        {"variant": name, "requested": len(cells), "written": len(edits), "audit": audit}, indent=1))
    return dest


def recalculate(directory: Path) -> None:
    """Recalculate every workbook under `directory` in one LibreOffice run.

    Invoked as the evaluation harness invokes it, as a subprocess over a
    directory: `batch_open_files` alone does not start the UNO service and fails
    with a refused socket, leaving formulas without cached values. A workbook
    whose formulas have no cached values scores as if the cells were empty, so a
    silent failure here would look exactly like a failed experiment.
    """
    import subprocess
    completed = subprocess.run([sys.executable, str(cc.EVAL_DIR / "open_spreadsheet.py"),
                                "--dir_path", str(directory)],
                               cwd=cc.EVAL_DIR, check=False, capture_output=True, text=True)
    if completed.returncode != 0 or "INIT_FAIL" in (completed.stdout + completed.stderr):
        raise RuntimeError(f"recalculation failed: {completed.stdout[-400:]} {completed.stderr[-400:]}")


def score(task: str, path: Path) -> dict:
    from evaluation import compare_workbooks_with_regression
    inp, gold = cc.workbook_paths(task)
    ok, msg, correct, total, reg, mod = compare_workbooks_with_regression(
        str(inp), str(gold), str(path), cc.dataset()[task]["answer_position"])
    return {"exact": ok, "first_error": msg[:160], "regression": reg, "modification": mod,
            "modification_accuracy": round(mod["correct"] / mod["total"], 6) if mod["total"] else None,
            "regression_accuracy": round(reg["correct"] / reg["total"], 6) if reg["total"] else None}


def cell_value(path: Path, cell: cc.Cell):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        return wb[cell[0]].cell(row=cell[1], column=cell[2]).value
    finally:
        wb.close()


def run_variants(task: str, variants: dict[str, set[cc.Cell]], probes: list[cc.Cell]) -> dict:
    paths = {name: build_variant(task, name, cells) for name, cells in variants.items()}
    recalculate(WORK / task)
    for name, path in paths.items():
        if cell_value(path, probes[0]) is None and variants[name]:
            pass  # recorded per variant below; a null probe is itself a datum
    _, gold = cc.workbook_paths(task)
    import openpyxl
    wg = openpyxl.load_workbook(gold, data_only=True)
    out = {}
    for name, path in paths.items():
        row = score(task, path)
        row["n_edits"] = len(variants[name])
        row["probe_cells"] = {}
        for p in probes:
            got = cell_value(path, p)
            want = wg[p[0]].cell(row=p[1], column=p[2]).value
            same = (got == want) or (isinstance(got, (int, float)) and isinstance(want, (int, float))
                                     and abs(got - want) <= max(1e-9, abs(want) * 1e-9))
            row["probe_cells"][f"{p[0]}!{cc.a1(p[1], p[2])}"] = {"value": got, "gold": want, "correct": bool(same)}
        out[name] = row
    wg.close()
    return out
