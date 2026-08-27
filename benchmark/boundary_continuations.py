"""Gold-blind right-edge formula-continuation candidates.

This is an offline measurement helper, not a shipped observation. It deliberately
targets a narrower topology than generic blank ranking: a styled blank immediately
after the rightmost exact horizontal translation run in a row. The motivating example
is Financial_Model 01_02, where Working Capital Schedule J3:L3 is a date sequence and
the golden workbook adds the otherwise unobserved M3 continuation.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from typing import Any

from openpyxl.styles.numbers import is_date_format
from openpyxl.utils.cell import get_column_letter, range_boundaries

from benchmark.blank_ranking import instruction_sheet_names
from librecalc_mcp.domain.formulas import translate_a1_formula


def _is_formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=")


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


@dataclass(frozen=True)
class BoundaryContinuation:
    """One inferable formula immediately beyond a horizontal translation run."""

    sheet: str
    address: str
    inferred_formula: str
    run_start: str
    run_end: str
    run_length: int
    style_match: bool
    date_like: bool
    named_sheet: bool

    def rank_key(self) -> tuple[int, int, int, int, str, str]:
        """Stable strongest-first ranking; no threshold is implied."""

        return (
            -int(self.date_like),
            -int(self.style_match),
            -int(self.named_sheet),
            -self.run_length,
            self.sheet.casefold(),
            self.address,
        )


def _date_like(raw_cell: Any, value_cell: Any) -> bool:
    value = value_cell.value
    return isinstance(value, (date, datetime)) or is_date_format(raw_cell.number_format)


def enumerate_boundary_continuations(
    raw_workbook: Any,
    value_workbook: Any,
    *,
    instruction: str,
    minimum_run: int = 3,
) -> list[BoundaryContinuation]:
    """Return gold-blind formula continuations inside each sheet's used rectangle.

    A candidate must be immediately right of the row's rightmost formula. The trailing
    formulas must form an exact adjacent A1 translation chain. Requiring the candidate
    to remain inside the input used rectangle prevents inventing an unbounded new horizon.
    """

    if minimum_run < 2:
        raise ValueError("minimum_run must be at least 2")
    named = instruction_sheet_names(instruction, [sheet.title for sheet in raw_workbook.worksheets])
    candidates: list[BoundaryContinuation] = []
    for raw_sheet in raw_workbook.worksheets:
        value_sheet = value_workbook[raw_sheet.title]
        min_column, min_row, max_column, max_row = range_boundaries(raw_sheet.calculate_dimension())
        if max_column <= min_column:
            continue
        for row in range(min_row, max_row + 1):
            formula_columns = [
                column
                for column in range(min_column, max_column + 1)
                if _is_formula(raw_sheet.cell(row, column).value)
            ]
            if not formula_columns:
                continue
            run_end_column = max(formula_columns)
            candidate_column = run_end_column + 1
            if candidate_column > max_column:
                continue
            candidate_cell = raw_sheet.cell(row, candidate_column)
            if not _is_blank(candidate_cell.value):
                continue

            run_start_column = run_end_column
            while run_start_column > min_column:
                previous = raw_sheet.cell(row, run_start_column - 1).value
                current = raw_sheet.cell(row, run_start_column).value
                if not (_is_formula(previous) and _is_formula(current)):
                    break
                try:
                    expected = translate_a1_formula(previous, column_offset=1, row_offset=0)
                except ValueError:
                    break
                if expected != current:
                    break
                run_start_column -= 1
            run_length = run_end_column - run_start_column + 1
            if run_length < minimum_run:
                continue
            last_cell = raw_sheet.cell(row, run_end_column)
            try:
                inferred = translate_a1_formula(str(last_cell.value), column_offset=1, row_offset=0)
            except ValueError:
                continue
            address = f"{get_column_letter(candidate_column)}{row}"
            candidates.append(
                BoundaryContinuation(
                    sheet=raw_sheet.title,
                    address=address,
                    inferred_formula=inferred,
                    run_start=f"{get_column_letter(run_start_column)}{row}",
                    run_end=f"{get_column_letter(run_end_column)}{row}",
                    run_length=run_length,
                    style_match=(
                        candidate_cell.style_id != 0
                        and candidate_cell.style_id == last_cell.style_id
                    ),
                    date_like=_date_like(last_cell, value_sheet.cell(row, run_end_column)),
                    named_sheet=raw_sheet.title in named,
                )
            )
    return sorted(candidates, key=BoundaryContinuation.rank_key)


def select_arm(candidates: list[BoundaryContinuation], arm: str) -> list[BoundaryContinuation]:
    """Select one named offline arm."""

    if arm == "run3":
        return candidates
    if arm == "style":
        return [candidate for candidate in candidates if candidate.style_match]
    if arm == "date":
        return [candidate for candidate in candidates if candidate.date_like]
    if arm == "date+style":
        return [
            candidate for candidate in candidates if candidate.date_like and candidate.style_match
        ]
    if arm.startswith("top-"):
        return candidates[: int(arm.removeprefix("top-"))]
    raise ValueError(f"Unknown continuation arm: {arm}")


ARMS = ("run3", "style", "date", "date+style", "top-5", "top-10")

__all__ = [
    "ARMS",
    "BoundaryContinuation",
    "enumerate_boundary_continuations",
    "select_arm",
]
