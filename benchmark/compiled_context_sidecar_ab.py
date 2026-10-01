#!/usr/bin/env python3
"""C0 vs C1 optional compiled-context sidecar A/B on the official coding-agent scaffold.

The autonomous compiled-agent architecture is frozen. This experiment asks whether
Spark on the normal SpreadsheetBench scaffold benefits from optional read-only
calc_query access to the existing compiled workbook world.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import threading
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import openpyxl
import yaml

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / "benchmark"), str(ROOT / "src")]

from calc_query import MAX_ROWS, compile_world, refuse_gold, run_query  # noqa: E402
from run_openrouter_slice import _load_dotenv  # noqa: E402

SLICE = ROOT / "benchmark/slices/compiled-context-sidecar-ten.json"
CONTROL = ROOT / "benchmark/sweagent/spreadsheet-control.yaml"
TREATMENT = ROOT / "benchmark/sweagent/spreadsheet-control-compiled-context.yaml"
TOOL_SCHEMA = ROOT / "benchmark/sweagent/calc_query/config.yaml"
RUNNER = ROOT / "benchmark/run_openrouter_slice.py"
SCORER = ROOT / "benchmark/score_openrouter_run.py"
DATA = ROOT / "benchmark-data/SpreadsheetBench-2/data"
RUNS = ROOT / "benchmark-data/SpreadsheetBench-2/benchmark-runs/openrouter"
ARTIFACT = ROOT / "compiled_context_sidecar_ab"
REPORT = ROOT / "COMPILED_CONTEXT_SIDECAR_AB_REPORT.md"
IDENTITY_SOURCE = (
    ROOT
    / "benchmark-data/SpreadsheetBench-2/benchmark-runs/openrouter"
    / "official-score-spark-1.3-contributor-fm-02_01"
    / "Financial_Model-02_01"
)

MODEL = "meta/muse-spark-1.3-contributor"
CALL_LIMIT = 40
COST_LIMIT = 2.50
TIMEOUT = 3600
EXECUTION_TIMEOUT = 180
WORKERS = 3
WITNESS_TASK = "Financial_Model:13_05"
WITNESS_RANGE = "I20:M20"
CONFLICT_RANGE = "D20:H20"

IDENTITY = {
    "declared_model": MODEL,
    "request_model": f"openrouter/{MODEL}",
    "temperature": 0.0,
    "top_p": 1.0,
    "reasoning_effort": None,
    "tool_choice": "auto",
    "provider": {"allow_fallbacks": True, "require_parameters": True},
    "token_limit_policy": "remaining-budget",
    "cost_limit_usd": COST_LIMIT,
    "call_limit": CALL_LIMIT,
    "execution_timeout": EXECUTION_TIMEOUT,
    "max_requeries": 2,
    "source_run": "official-score-spark-1.3-contributor-fm-02_01",
    "source_call_limit_note": "Source Spark control used 50 calls/$2.50; this probe keeps generation identity and uses the 40/$2.50 envelope.",
}

PRIMARY = {
    "category": "Financial_Model",
    "id": "13_05",
    "role": "primary_mechanism_witness",
    "repeats": 4,
}
GENERALIZATION = [
    {"category": "Financial_Model", "id": "05_01", "role": "generalization", "repeats": 1},
    {"category": "Financial_Model", "id": "02_01", "role": "saturation_non_harm", "repeats": 1},
    {"category": "Financial_Model", "id": "15_05", "role": "negative_structure", "repeats": 1},
    {"category": "Template", "id": "03_03", "role": "generalization", "repeats": 1},
    {"category": "Template", "id": "06_17", "role": "generalization", "repeats": 1},
    {"category": "Template", "id": "16_12", "role": "generalization", "repeats": 1},
    {"category": "Debugging", "id": "01_01", "role": "generalization", "repeats": 1},
    {"category": "Debugging", "id": "09_04", "role": "generalization", "repeats": 1},
    {"category": "Debugging", "id": "10_10", "role": "generalization", "repeats": 1},
]
FORBIDDEN_EXCLUSIONS = (
    "Task IR",
    "Edit Plan",
    "calc_apply",
    "autonomous scheduler",
    "ProgramGroup",
    "semantic formula verifier",
)

A1_TOKEN = re.compile(r"\b(?:[A-Za-z][\w .&,-]*!)?\$?[A-Za-z]{1,3}\$?\d+(?::\$?[A-Za-z]{1,3}\$?\d+)?\b")
CALC_QUERY_RE = re.compile(r"^\s*calc_query\b", re.I)


def now() -> str:
    return datetime.now(UTC).isoformat()


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def sha256_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def sha256_bytes(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def dataset_record(category: str, task_id: str) -> dict[str, Any]:
    records = load_json(DATA / category / "dataset.json")
    matches = [row for row in records if row["id"] == task_id]
    if len(matches) != 1:
        raise ValueError(f"Expected one {category}:{task_id} record, found {len(matches)}")
    return matches[0]


def source_xlsx(category: str, task_id: str) -> Path:
    record = dataset_record(category, task_id)
    return DATA / category / record["spreadsheet_path"]


def gold_xlsx(category: str, task_id: str) -> Path:
    record = dataset_record(category, task_id)
    return DATA / category / record["golden_response_path"]


def _parse_nested_json(path: Path) -> dict[str, Any]:
    text = path.read_text(encoding="utf-8")
    try:
        payload: Any = json.loads(text)
    except json.JSONDecodeError:
        payload = yaml.safe_load(text)
    if isinstance(payload, str):
        payload = json.loads(payload)
    if not isinstance(payload, dict):
        raise TypeError(f"{path} did not contain a JSON object")
    return payload


def identity_source_contract() -> dict[str, Any]:
    contract = load_json(IDENTITY_SOURCE / "run_contract.json")
    traj_config = _parse_nested_json(next((IDENTITY_SOURCE / "trajectory").rglob("config.yaml")))
    model_cfg = traj_config["agent"]["model"]
    kwargs = model_cfg.get("completion_kwargs") or {}
    return {
        "model": contract["model"],
        "model_catalog_id": contract.get("model_catalog_id"),
        "request_model_from_traj": model_cfg["name"],
        "temperature": model_cfg.get("temperature"),
        "top_p": model_cfg.get("top_p"),
        "reasoning_effort": contract.get("reasoning_effort"),
        "tool_choice": kwargs.get("tool_choice"),
        "provider": kwargs.get("provider"),
        "token_limit_policy": "remaining-budget" if kwargs.get("librecalc_budget_max_tokens") else "fixed",
        "cost_limit_usd": model_cfg.get("per_instance_cost_limit"),
        "source_call_limit": model_cfg.get("per_instance_call_limit"),
    }


def prompt_hashes() -> dict[str, Any]:
    c0 = yaml.safe_load(CONTROL.read_text(encoding="utf-8"))
    c1 = yaml.safe_load(TREATMENT.read_text(encoding="utf-8"))
    schema = yaml.safe_load(TOOL_SCHEMA.read_text(encoding="utf-8"))
    return {
        "c0_system_sha256": sha256_text(c0["agent"]["templates"]["system_template"]),
        "c1_system_sha256": sha256_text(c1["agent"]["templates"]["system_template"]),
        "c0_instance_sha256": sha256_text(c0["agent"]["templates"]["instance_template"]),
        "c1_instance_sha256": sha256_text(c1["agent"]["templates"]["instance_template"]),
        "c0_config_sha256": sha256_bytes(CONTROL),
        "c1_config_sha256": sha256_bytes(TREATMENT),
        "c1_tool_schema_sha256": sha256_text(json.dumps(schema, sort_keys=True)),
        "c0_bundles": [item["path"] for item in c0["agent"]["tools"]["bundles"]],
        "c1_bundles": [item["path"] for item in c1["agent"]["tools"]["bundles"]],
    }  # type: ignore[return-value]


def pair_arms(repeat: int, *, even_c0_first: bool = True) -> tuple[str, str]:
    first_c0 = (repeat % 2 == 1) if even_c0_first else (repeat % 2 == 0)
    return ("C0", "C1") if first_c0 else ("C1", "C0")


def build_jobs() -> list[dict[str, Any]]:
    jobs: list[dict[str, Any]] = []
    order = 0
    for repeat in range(1, PRIMARY["repeats"] + 1):
        for arm in pair_arms(repeat, even_c0_first=True):
            order += 1
            jobs.append(_job(PRIMARY, arm, repeat, order))
    for index, spec in enumerate(GENERALIZATION):
        even_c0_first = index % 2 == 0
        for arm in pair_arms(1, even_c0_first=even_c0_first):
            order += 1
            jobs.append(_job(spec, arm, 1, order))
    return jobs


def _job(spec: dict[str, Any], arm: str, repeat: int, order: int) -> dict[str, Any]:
    task = f"{spec['category']}:{spec['id']}"
    run_name = f"spark-sidecar-{arm.lower()}-{spec['category']}-{spec['id']}-r{repeat}"
    return {
        "order": order,
        "task": task,
        "category": spec["category"],
        "id": spec["id"],
        "role": spec["role"],
        "arm": arm,
        "repeat": repeat,
        "run_name": run_name,
        "pair_key": f"{task}:r{repeat}",
    }


def build_pairs(jobs: list[dict[str, Any]]) -> list[list[dict[str, Any]]]:
    grouped: dict[str, list[dict[str, Any]]] = {}
    order: list[str] = []
    for job in jobs:
        key = job["pair_key"]
        if key not in grouped:
            grouped[key] = []
            order.append(key)
        grouped[key].append(job)
    return [grouped[key] for key in order]


def build_groups(jobs: list[dict[str, Any]]) -> list[list[dict[str, Any]]]:
    """Run 13_05 as one sequential chain; other tasks keep their C0/C1 pair order.

    Same-task docker outputs collide, so 13_05 must not occupy both workers.
    """
    witness = [job for job in jobs if job["task"] == WITNESS_TASK]
    others = [job for job in jobs if job["task"] != WITNESS_TASK]
    groups = [witness] if witness else []
    groups.extend(build_pairs(others))
    return groups


def git_commit() -> str:
    result = subprocess.run(
        ["git", "rev-parse", "HEAD"],
        cwd=ROOT,
        check=False,
        capture_output=True,
        text=True,
    )
    return (result.stdout or "").strip()


def first_formula_cell(path: Path) -> tuple[str, str, str] | None:
    workbook = openpyxl.load_workbook(path, data_only=False)
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if isinstance(value, str) and value.startswith("="):
                        return sheet.title, cell.coordinate, value
    finally:
        workbook.close()
    return None


def first_label_cell(path: Path) -> tuple[str, str, str] | None:
    workbook = openpyxl.load_workbook(path, data_only=False)
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if isinstance(value, str) and value.strip() and not value.startswith("="):
                        return sheet.title, cell.coordinate, value
    finally:
        workbook.close()
    return None


def preflight() -> dict[str, Any]:
    hashes = prompt_hashes()
    identity = identity_source_contract()
    if identity["model"] != MODEL:
        raise RuntimeError(f"IDENTITY_SOURCE_MODEL_MISMATCH: {identity['model']}")
    if identity["request_model_from_traj"] != f"openrouter/{MODEL}":
        raise RuntimeError(f"IDENTITY_SOURCE_REQUEST_MISMATCH: {identity['request_model_from_traj']}")
    if identity["temperature"] != 0.0 or identity["top_p"] != 1.0:
        raise RuntimeError("IDENTITY_SOURCE_GENERATION_MISMATCH")
    if identity["tool_choice"] != "auto" or identity["reasoning_effort"] is not None:
        raise RuntimeError("IDENTITY_SOURCE_TOOL_OR_REASONING_MISMATCH")

    slice_data = load_json(SLICE)
    tasks = slice_data["tasks"]
    cache_root = ARTIFACT / "world_cache"
    cache_root.mkdir(parents=True, exist_ok=True)
    rows = []
    for spec in tasks:
        xlsx = source_xlsx(spec["category"], spec["id"])
        gold = gold_xlsx(spec["category"], spec["id"])
        if refuse_gold(xlsx):
            raise RuntimeError(f"SOURCE_LOOKS_LIKE_GOLD: {xlsx}")
        if not refuse_gold(gold):
            raise RuntimeError(f"GOLD_PATH_NOT_REFUSED: {gold}")
        before = sha256_bytes(xlsx)
        dest = cache_root / f"{spec['category']}-{spec['id']}"
        if dest.exists():
            shutil.rmtree(dest)
        compiled = compile_world(xlsx, dest)
        if compiled.get("golden_used"):
            raise RuntimeError(f"GOLD_USED: {spec['id']}")
        inspect_target = first_formula_cell(xlsx) or first_label_cell(xlsx)
        if inspect_target is None:
            raise RuntimeError(f"NO_CELL_FOR_FIDELITY: {spec['id']}")
        sheet, address, value = inspect_target
        inspect = run_query(xlsx, "inspect", sheet=sheet, target=address, cache_root=cache_root)
        search_target = first_label_cell(xlsx)
        search = run_query(
            xlsx,
            "search",
            text=(search_target[2][:24] if search_target else "a"),
            cache_root=cache_root,
        )
        periods = run_query(xlsx, "periods", cache_root=cache_root)
        after = sha256_bytes(xlsx)
        if before != after:
            raise RuntimeError(f"QUERY_MUTATED_SOURCE: {spec['id']}")
        inspect2 = run_query(xlsx, "inspect", sheet=sheet, target=address, cache_root=cache_root)
        if inspect != inspect2:
            raise RuntimeError(f"NONDETERMINISTIC: {spec['id']}")
        formula_ok = True
        if isinstance(value, str) and value.startswith("="):
            got = (inspect.get("cells") or [{}])[0].get("formula")
            formula_ok = got == value
        bounded = all(
            len(json.dumps(payload)) < 400_000
            for payload in (inspect, search, periods)
        ) and not (
            (inspect.get("cells") and len(inspect["cells"]) > MAX_ROWS)
            or (search.get("hits") and len(search["hits"]) > MAX_ROWS)
            or (periods.get("coordinates") and len(periods["coordinates"]) > MAX_ROWS + 8)
        )
        no_authority = all(payload.get("should_edit") is None and payload.get("edit_authority") is False for payload in (inspect, search, periods))
        no_table_body = all("table_body" not in payload and "implicit_table" not in json.dumps(payload) for payload in (inspect, search, periods))
        rows.append(
            {
                "task": f"{spec['category']}:{spec['id']}",
                "source_sha256": before,
                "compiled_db": compiled["db"],
                "gold_refused": True,
                "mutated": False,
                "deterministic": True,
                "formula_fidelity": formula_ok,
                "bounded": bounded,
                "no_edit_authority": no_authority,
                "no_invented_table_body": no_table_body,
                "inspect_status": inspect.get("status"),
                "periods_status": periods.get("status"),
                "period_rows": len(periods.get("coordinates") or []),
            }
        )
    failed = [
        row["task"]
        for row in rows
        if not (row["formula_fidelity"] and row["bounded"] and row["no_edit_authority"] and row["no_invented_table_body"])
    ]
    payload = {
        "passed": not failed,
        "failed_tasks": failed,
        "identity_source": identity,
        "identity": IDENTITY,
        "prompt_hashes": hashes,
        "tasks": rows,
        "exclusions_absent_from_c1_prompt": all(
            token.lower() not in TREATMENT.read_text(encoding="utf-8").lower()
            for token in ("calc_apply", "task ir", "edit plan", "programgroup")
        ),
        "written_at": now(),
    }
    write_json(ARTIFACT / "preflight.json", payload)
    write_json(ARTIFACT / "prompt_hashes.json", hashes)
    write_json(ARTIFACT / "spec.json", experiment_spec(hashes, identity))
    return payload


def experiment_spec(hashes: dict[str, Any], identity: dict[str, Any]) -> dict[str, Any]:
    return {
        "name": "compiled-context-sidecar-ab",
        "hypothesis": "Does a general spreadsheet coding agent perform better when it has optional, well-explained, read-only access to a deterministic compiled structural model of the workbook?",
        "mechanism_question": "When the capability is clearly explained and safe to use, does the agent actually choose to use it?",
        "frozen_architecture": "autonomous compiled-agent (Task IR, Edit Plan, scheduler, ProgramGroups, calc_apply) is not in this experiment",
        "model": MODEL,
        "identity": IDENTITY,
        "identity_source": identity,
        "envelope": {"max_model_calls": CALL_LIMIT, "max_usd": COST_LIMIT},
        "slice": str(SLICE.relative_to(ROOT)),
        "prompt_hashes": hashes,
        "jobs": build_jobs(),
        "git_commit": git_commit(),
    }


def run_name_root(run_name: str) -> Path:
    return RUNS / run_name


def job_task_root(job: dict[str, Any]) -> Path:
    return run_name_root(job["run_name"]) / f"{job['category']}-{job['id']}"


def runner_command(job: dict[str, Any]) -> list[str]:
    flag = "--control" if job["arm"] == "C0" else "--control-compiled-context"
    return [
        sys.executable,
        str(RUNNER),
        "--slice",
        str(SLICE),
        "--run-name",
        job["run_name"],
        "--task",
        job["task"],
        "--model",
        MODEL,
        flag,
        "--call-limit",
        str(CALL_LIMIT),
        "--cost-limit",
        str(COST_LIMIT),
        "--timeout",
        str(TIMEOUT),
        "--execution-timeout",
        str(EXECUTION_TIMEOUT),
        "--max-requeries",
        "2",
        "--skip-existing",
        "--no-score",
    ]


def normalize_response_model(value: str | None) -> str | None:
    """Strip LiteLLM/OpenRouter quoting artifacts from a logged response model.

    Trailing backslashes are a reporting parse artifact from JSON-in-string
    dumps. This does not change requests.
    """
    if value is None:
        return None
    cleaned = str(value).split("|")[0].rstrip("\\").strip()
    return cleaned or None


def classify_failure(returncode: int, task_root: Path) -> str:
    if returncode == 0:
        return "ok"
    logs = list(task_root.rglob("*.log")) if task_root.exists() else []
    blob = ""
    for path in logs[-3:]:
        try:
            blob += path.read_text(encoding="utf-8", errors="replace")[-8000:]
        except OSError:
            continue
    ledger: dict[str, Any] = {}
    ledger_path = task_root.parent / "ledger.jsonl"
    if ledger_path.is_file():
        try:
            ledger = json.loads(ledger_path.read_text(encoding="utf-8").strip().splitlines()[-1])
        except (json.JSONDecodeError, IndexError, OSError):
            ledger = {}
    markers = (
        "Server disconnected",
        "OpenrouterException",
        "litellm.APIError",
        "APIError",
        "Rate limit",
        "Beginning environment shutdown",
        "Failed to start container",
    )
    model_calls = int(ledger.get("model_calls") or 0)
    prompt_tokens = int(ledger.get("prompt_tokens") or 0)
    cost_usd = float(ledger.get("cost_usd") or ledger.get("api_cost") or 0)
    no_model_work = model_calls == 0 and prompt_tokens == 0
    wrapper_crash = calc_query_wrapper_crash(task_root)
    # Envelope exhaustion is not provider infrastructure failure. A task that
    # spent its call or cost budget remains censored for causal comparison,
    # but it is classified separately from Docker/OpenRouter crashes.
    envelope_marker = any(
        token in blob.lower()
        for token in (
            "reached maximum of",
            "reached limit of",
            "cost limit reached",
            "call limit reached",
        )
    )
    if (
        (model_calls >= CALL_LIMIT or cost_usd >= COST_LIMIT or envelope_marker)
        and not wrapper_crash
        and not no_model_work
    ):
        return "envelope_or_resource"
    if (
        any(marker in blob for marker in markers)
        or returncode in {120, 124, 137}
        or no_model_work
        or wrapper_crash
        or (ledger.get("error") == "no output workbook" and model_calls == 0)
    ):
        return "provider_or_infra"
    return "semantic_or_other"


def calc_query_wrapper_crash(task_root: Path) -> bool:
    """True when calc_query died on the Docker-shallow parents[5] IndexError."""
    if not task_root.exists():
        return False
    for traj in task_root.rglob("*.traj"):
        try:
            text = traj.read_text(encoding="utf-8", errors="replace")
        except OSError:
            continue
        if "IndexError: 5" in text and "calc_query" in text:
            return True
    return False


def run_job(job: dict[str, Any]) -> dict[str, Any]:
    task_root = job_task_root(job)
    run_root = run_name_root(job["run_name"])
    if (task_root / "output.xlsx").is_file() and (task_root / "run_contract.json").is_file():
        return {"job": job, "status": "skipped_existing", "returncode": 0, "failure_class": "ok"}
    # Retry only crash/infra attempts with no model work. Semantic/envelope
    # failures (including a 40-call run that produced no workbook) stay frozen.
    if run_root.exists() and not (task_root / "output.xlsx").is_file():
        ledger_path = run_root / "ledger.jsonl"
        model_calls = 0
        if ledger_path.is_file() and ledger_path.read_text(encoding="utf-8").strip():
            try:
                last = json.loads(ledger_path.read_text(encoding="utf-8").strip().splitlines()[-1])
                model_calls = int(last.get("model_calls") or 0)
            except (json.JSONDecodeError, TypeError, ValueError):
                model_calls = 0
        if model_calls > 0 and not calc_query_wrapper_crash(task_root):
            return {
                "job": job,
                "status": "failed",
                "returncode": 1,
                "failure_class": "semantic_or_other",
            }
        shutil.rmtree(run_root)
    cache = ARTIFACT / "job_cache" / job["run_name"]
    cache.mkdir(parents=True, exist_ok=True)
    if job["arm"] == "C1":
        xlsx = source_xlsx(job["category"], job["id"])
        digest = hashlib.sha256(xlsx.read_bytes()).hexdigest()[:24]
        compiled = ARTIFACT / "world_cache" / f"{job['category']}-{job['id']}"
        dest = cache / digest
        if compiled.is_dir() and not (dest / "workbook.sqlite").is_file():
            shutil.copytree(compiled, dest, dirs_exist_ok=True)
    env = os.environ.copy()
    env["CALC_QUERY_HOST_CACHE"] = str(cache)
    started = now()
    result = subprocess.run(runner_command(job), cwd=ROOT, env=env, check=False)
    failure_class = classify_failure(result.returncode, task_root)
    payload = {
        "job": job,
        "status": "completed" if result.returncode == 0 else "failed",
        "returncode": result.returncode,
        "failure_class": failure_class,
        "started_at": started,
        "finished_at": now(),
    }
    write_json(task_root / "sidecar_job.json", payload)
    return payload


def launch(jobs: list[dict[str, Any]], *, workers: int = WORKERS, dry_run: bool = False) -> dict[str, Any]:
    groups = build_groups(jobs)
    write_json(
        ARTIFACT / "manifest.json",
        {
            "jobs": jobs,
            "groups": [[job["run_name"] for job in group] for group in groups],
            "workers": workers,
            "written_at": now(),
        },
    )
    if dry_run:
        return {"dry_run": True, "jobs": len(jobs), "groups": len(groups)}
    _load_dotenv()
    if not os.environ.get("OPENROUTER_API_KEY"):
        raise RuntimeError("OPENROUTER_API_KEY must be set")
    if shutil.which("docker") is None:
        raise RuntimeError("docker is required")
    task_locks = {job["task"]: threading.Lock() for job in jobs}
    results: list[dict[str, Any]] = []

    def run_group(group: list[dict[str, Any]]) -> list[dict[str, Any]]:
        with task_locks[group[0]["task"]]:
            return [run_job(job) for job in group]

    print(f"LAUNCH groups={len(groups)} jobs={len(jobs)} workers={workers}", flush=True)
    with ThreadPoolExecutor(max_workers=workers) as pool:
        futures = {pool.submit(run_group, group): group for group in groups}
        for future in as_completed(futures):
            group_results = future.result()
            results.extend(group_results)
            for item in group_results:
                job = item["job"]
                print(
                    f"JOB {job['arm']} r{job['repeat']} {job['task']} status={item['status']} class={item['failure_class']}",
                    flush=True,
                )
    write_json(ARTIFACT / "launch_results.json", {"results": results, "finished_at": now()})
    print(f"LAUNCH_DONE jobs={len(results)}", flush=True)
    return {"results": results}


def score_runs(jobs: list[dict[str, Any]]) -> dict[str, Any]:
    scores: dict[str, Any] = {}
    for job in jobs:
        run_root = run_name_root(job["run_name"])
        key = f"{job['task']}:{job['arm']}:r{job['repeat']}"
        if not run_root.exists():
            scores[key] = {"status": "missing_run"}
            continue
        result = subprocess.run(
            [sys.executable, str(SCORER), str(run_root), "--write-ledger"],
            cwd=ROOT,
            check=False,
            capture_output=True,
            text=True,
        )
        official = run_root / "official_scores.json"
        payload = load_json(official) if official.is_file() else {}
        task_scores = {}
        if isinstance(payload.get("tasks"), dict):
            task_scores = payload["tasks"].get(job["task"]) or {}
        scores[key] = {
            "returncode": result.returncode,
            "official": task_scores,
            "stderr": (result.stderr or "")[-2000:],
        }
    write_json(ARTIFACT / "official_scores.json", scores)
    return scores


def _find_traj(task_root: Path) -> Path | None:
    matches = sorted(task_root.rglob("*.traj"))
    return matches[-1] if matches else None


def parse_calc_query_calls(traj_path: Path | None) -> dict[str, Any]:
    if traj_path is None or not traj_path.is_file():
        return {"used": False, "classification": "NOT_USED", "calls": []}
    data = json.loads(traj_path.read_text(encoding="utf-8"))
    calls = []
    returned_tokens: list[str] = []
    first_turn = None
    for index, step in enumerate(data.get("trajectory") or [], start=1):
        action = str(step.get("action") or "")
        extra = step.get("extra_info") or {}
        current_calls = []
        if isinstance(extra, dict):
            current_calls.extend(extra.get("tool_calls") or [])
        names = [((item.get("function") or {}).get("name") if isinstance(item, dict) else None) for item in current_calls]
        is_calc = CALC_QUERY_RE.search(action) is not None or "calc_query" in names
        if not is_calc:
            continue
        if first_turn is None:
            first_turn = index
        observation = str(step.get("observation") or "")
        returned_tokens.extend(A1_TOKEN.findall(observation))
        mode = None
        parts = action.split()
        if len(parts) >= 2:
            mode = parts[1]
        arguments = action
        if current_calls:
            fn = current_calls[0].get("function") if isinstance(current_calls[0], dict) else None
            if isinstance(fn, dict):
                arguments = fn.get("arguments") or arguments
                mode = mode or (json.loads(arguments).get("mode") if isinstance(arguments, str) and arguments.startswith("{") else mode)
        calls.append(
            {
                "turn": index,
                "mode": mode,
                "arguments": arguments,
                "result_chars": len(observation),
                "latency_s": step.get("execution_time"),
                "failed": any(
                    marker in observation
                    for marker in ("NOT_AVAILABLE", "GOLD_PATH_REFUSED", "IndexError", "Traceback")
                ),
                "observation_preview": observation[:400],
            }
        )
    later = []
    if calls:
        last_query_turn = calls[-1]["turn"]
        for index, step in enumerate(data.get("trajectory") or [], start=1):
            if index <= last_query_turn:
                continue
            blob = " ".join(
                [
                    str(step.get("action") or ""),
                    str(step.get("thought") or ""),
                    str(step.get("response") or ""),
                ]
            )
            later.append(blob)
        later_text = "\n".join(later)
        uptake = sorted({token for token in returned_tokens if token and token in later_text})
    else:
        uptake = []
    return {
        "used": bool(calls),
        "classification": "USED" if calls else "NOT_USED",
        "first_turn": first_turn,
        "n_calls": len(calls),
        "modes": [item["mode"] for item in calls],
        "repeated_queries": len(calls) - len({(item["mode"], str(item["arguments"])) for item in calls}),
        "calls": calls,
        "evidence_uptake_tokens": uptake[:40],
        "evidence_uptake": bool(uptake),
    }


def ledger_row(task_root: Path) -> dict[str, Any]:
    ledger = task_root.parent / "ledger.jsonl"
    if not ledger.is_file():
        return {}
    rows = [json.loads(line) for line in ledger.read_text(encoding="utf-8").splitlines() if line.strip()]
    return rows[-1] if rows else {}


def identity_audit(job: dict[str, Any]) -> dict[str, Any]:
    task_root = job_task_root(job)
    contract_path = task_root / "run_contract.json"
    if not contract_path.is_file():
        return {"ok": False, "reason": "missing_run_contract"}
    contract = load_json(contract_path)
    traj = _find_traj(task_root)
    traj_model = None
    traj_temp = None
    traj_top_p = None
    traj_tool = None
    response_model = None
    if traj:
        cfg = _parse_nested_json(traj.with_name("config.yaml")) if traj.with_name("config.yaml").is_file() else {}
        model_cfg = (cfg.get("agent") or {}).get("model") or {}
        traj_model = model_cfg.get("name")
        traj_temp = model_cfg.get("temperature")
        traj_top_p = model_cfg.get("top_p")
        traj_tool = (model_cfg.get("completion_kwargs") or {}).get("tool_choice")
        try:
            data = json.loads(traj.read_text(encoding="utf-8"))
            blob = json.dumps(data)
            match = re.search(r"meta/muse-spark-1\.3-contributor[^\"]*", blob)
            if match:
                response_model = normalize_response_model(match.group(0))
        except OSError:
            pass
    declared = contract.get("declared_model") or contract.get("model")
    request = contract.get("request_model") or traj_model
    ok = (
        declared == MODEL
        and (request in {MODEL, f"openrouter/{MODEL}"})
        and contract.get("reasoning_effort") is None
        and (traj_temp in {None, 0, 0.0})
        and (traj_top_p in {None, 1, 1.0})
        and (traj_tool in {None, "auto"})
        and (response_model is None or response_model.startswith(MODEL))
    )
    return {
        "ok": ok,
        "declared_model": declared,
        "request_model": request,
        "response_model": response_model,
        "temperature": traj_temp if traj_temp is not None else contract.get("temperature"),
        "top_p": traj_top_p if traj_top_p is not None else contract.get("top_p"),
        "tool_choice": traj_tool or contract.get("tool_choice"),
        "reasoning_effort": contract.get("reasoning_effort"),
        "effective_generation_parameters": {
            "temperature": traj_temp if traj_temp is not None else 0.0,
            "top_p": traj_top_p if traj_top_p is not None else 1.0,
            "tool_choice": traj_tool or "auto",
            "reasoning_effort": contract.get("reasoning_effort"),
            "provider": contract.get("provider_policy"),
        },
    }


def official_metrics(payload: Any) -> dict[str, Any]:
    if not isinstance(payload, dict):
        return {"modification": None, "regression": None, "exact": None, "usable": None}
    inner = payload.get("official") if isinstance(payload.get("official"), dict) else payload
    if "modification_accuracy" not in inner and isinstance(inner.get("tasks"), dict):
        inner = next(iter(inner["tasks"].values()), {})
    modification = inner.get("modification_accuracy")
    regression = inner.get("regression_accuracy")
    exact = inner.get("accuracy")
    usable = (
        modification is not None
        and regression is not None
        and float(modification) >= 0.99
        and float(regression) >= 0.99
    )
    return {
        "modification": modification,
        "regression": regression,
        "exact": exact == 1.0 if exact is not None else None,
        "exact_value": exact,
        "usable": usable,
        "error_message": inner.get("error_message"),
    }


def _cell_values(path: Path, addresses: list[str]) -> dict[str, Any]:
    if not path.is_file():
        return {}
    workbook = openpyxl.load_workbook(path, data_only=False)
    try:
        best: dict[str, Any] = {}
        for sheet in workbook.worksheets:
            values = {addr: sheet[addr].value for addr in addresses}
            if any(value is not None for value in values.values()):
                return values
            best = values
        return best
    except Exception:
        return {}
    finally:
        workbook.close()


def thirteen_diagnostic(job: dict[str, Any], adoption: dict[str, Any]) -> dict[str, Any]:
    task_root = job_task_root(job)
    output = task_root / "output.xlsx"
    source = source_xlsx(job["category"], job["id"])
    gold = gold_xlsx(job["category"], job["id"])
    queried_periods = any(call.get("mode") == "periods" for call in adoption.get("calls") or [])
    blob = json.dumps(adoption.get("calls") or [])
    exposed_fy = any(token in blob for token in ("2026", "FY26", "fy26", "FY30", "2030"))
    witness_addrs = [f"{col}20" for col in "IJKLM"]
    conflict_addrs = [f"{col}20" for col in "DEFGH"]
    out_w = _cell_values(output, witness_addrs)
    gold_w = _cell_values(gold, witness_addrs)
    src_c = _cell_values(source, conflict_addrs)
    out_c = _cell_values(output, conflict_addrs)
    witness_match = bool(out_w) and bool(gold_w) and all(out_w.get(addr) == gold_w.get(addr) for addr in witness_addrs)
    conflict = bool(out_c) and any(out_c.get(addr) != src_c.get(addr) for addr in conflict_addrs)
    if queried_periods and exposed_fy and witness_match:
        mechanism = "compiler_known_and_queried_and_used"
    elif queried_periods and exposed_fy and not witness_match:
        mechanism = "queried_but_wrong_model_choice"
    elif not queried_periods:
        mechanism = "compiler_known_not_queried"
    else:
        mechanism = "sidecar_exposure_uncertain"
    return {
        "queried_periods": queried_periods,
        "exposed_fy26_fy30": exposed_fy,
        "witness_I20_M20_matches_gold": witness_match,
        "conflict_D20_H20_modified": conflict,
        "mechanism_class": mechanism,
    }


def opportunity_class(job: dict[str, Any], metrics: dict[str, Any], adoption: dict[str, Any], identity: dict[str, Any], failure_class: str) -> str:
    if failure_class == "envelope_or_resource":
        return "ENVELOPE_OR_RESOURCE_CENSORED"
    if failure_class == "provider_or_infra" or not identity.get("ok"):
        return "PROVIDER_OR_RESOURCE_CENSORED"
    if metrics.get("usable") or metrics.get("exact"):
        return "UNRESOLVED"
    if not adoption.get("used"):
        return "AVAILABLE_NOT_QUERIED" if job["role"] != "negative_structure" else "NOT_AVAILABLE_IN_COMPILED_CONTEXT"
    if adoption.get("used") and not metrics.get("usable"):
        return "QUERIED_BUT_WRONG_MODEL_CHOICE"
    return "UNRESOLVED"


def analyze(jobs: list[dict[str, Any]], scores: dict[str, Any], launch_results: list[dict[str, Any]]) -> dict[str, Any]:
    failure_by_run = {
        item["job"]["run_name"]: item.get("failure_class", "ok") for item in launch_results if item.get("job")
    }
    rows = []
    adoption_rows = []
    identity_rows = []
    thirteen_rows = []
    audit_rows = []
    for job in jobs:
        key = f"{job['task']}:{job['arm']}:r{job['repeat']}"
        task_root = job_task_root(job)
        ledger = ledger_row(task_root)
        identity = identity_audit(job)
        identity_rows.append({"key": key, **identity})
        metrics = official_metrics(scores.get(key) or {})
        adoption = {"used": False, "classification": "NOT_USED"}
        if job["arm"] == "C1":
            adoption = parse_calc_query_calls(_find_traj(task_root))
            adoption_rows.append({"key": key, "task": job["task"], "repeat": job["repeat"], **adoption})
        if job["task"] == WITNESS_TASK:
            thirteen_rows.append(
                {
                    "key": key,
                    "arm": job["arm"],
                    "repeat": job["repeat"],
                    **thirteen_diagnostic(job, adoption if job["arm"] == "C1" else {"calls": []}),
                    **metrics,
                }
            )
        failure_class = failure_by_run.get(job["run_name"], "ok")
        if job["arm"] == "C1" and not (metrics.get("usable") or metrics.get("exact")):
            audit_rows.append(
                {
                    "key": key,
                    "task": job["task"],
                    "class": opportunity_class(job, metrics, adoption, identity, failure_class),
                    "error_message": metrics.get("error_message"),
                    "used": adoption.get("used"),
                }
            )
        rows.append(
            {
                "task": job["task"],
                "role": job["role"],
                "arm": job["arm"],
                "repeat": job["repeat"],
                "run_name": job["run_name"],
                "identity_ok": identity.get("ok"),
                "status": ledger.get("status") or failure_class,
                "model_calls": ledger.get("model_calls"),
                "cost_usd": ledger.get("charged_cost_usd"),
                "prompt_tokens": ledger.get("prompt_tokens"),
                "completion_tokens": ledger.get("completion_tokens"),
                "elapsed_seconds": ledger.get("elapsed_seconds"),
                "calc_query_used": adoption.get("used") if job["arm"] == "C1" else None,
                "calc_query_calls": adoption.get("n_calls") if job["arm"] == "C1" else None,
                **metrics,
            }
        )
    paired = pair_summary(rows)
    verdict = choose_verdict(rows, adoption_rows, identity_rows, failure_by_run)
    payload = {
        "verdict": verdict,
        "rows": rows,
        "paired": paired,
        "adoption": adoption_rows,
        "identity": identity_rows,
        "thirteen_05": thirteen_rows,
        "opportunity_audit": audit_rows,
        "written_at": now(),
    }
    write_json(ARTIFACT / "paired_results.json", payload)
    with (ARTIFACT / "paired_results.csv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0].keys()) if rows else ["task"])
        writer.writeheader()
        writer.writerows(rows)
    write_json(ARTIFACT / "calc_query_ledger.json", adoption_rows)
    write_json(ARTIFACT / "tool_adoption.json", {
        "used": sum(1 for row in adoption_rows if row.get("used")),
        "not_used": sum(1 for row in adoption_rows if not row.get("used")),
        "n": len(adoption_rows),
        "by_task": adoption_rows,
    })
    write_json(ARTIFACT / "thirteen_05_mechanism.json", thirteen_rows)
    write_json(ARTIFACT / "opportunity_audit.json", audit_rows)
    write_json(ARTIFACT / "identity_audit.json", identity_rows)
    return payload


def pair_summary(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[tuple[str, int], dict[str, dict[str, Any]]] = {}
    for row in rows:
        grouped.setdefault((row["task"], row["repeat"]), {})[row["arm"]] = row
    out = []
    for (task, repeat), arms in grouped.items():
        c0 = arms.get("C0") or {}
        c1 = arms.get("C1") or {}
        out.append(
            {
                "task": task,
                "repeat": repeat,
                "c0_modification": c0.get("modification"),
                "c1_modification": c1.get("modification"),
                "c0_regression": c0.get("regression"),
                "c1_regression": c1.get("regression"),
                "c0_exact": c0.get("exact"),
                "c1_exact": c1.get("exact"),
                "c0_usable": c0.get("usable"),
                "c1_usable": c1.get("usable"),
                "c0_calls": c0.get("model_calls"),
                "c1_calls": c1.get("model_calls"),
                "c0_cost": c0.get("cost_usd"),
                "c1_cost": c1.get("cost_usd"),
                "c1_used_calc_query": c1.get("calc_query_used"),
                "identity_ok": bool(c0.get("identity_ok") and c1.get("identity_ok")),
            }
        )
    return out


def choose_verdict(
    rows: list[dict[str, Any]],
    adoption: list[dict[str, Any]],
    identity: list[dict[str, Any]],
    failures: dict[str, str],
) -> str:
    if identity and not any(row.get("ok") for row in identity):
        return "PROVIDER_CENSORED"
    infra = sum(1 for value in failures.values() if value == "provider_or_infra")
    if infra >= max(3, len(rows) // 2):
        return "PROVIDER_CENSORED"
    used = [row for row in adoption if row.get("used")]
    paired = pair_summary(rows)
    comparable = [row for row in paired if row.get("identity_ok") and row.get("c0_modification") is not None and row.get("c1_modification") is not None]
    if not used:
        return "NO_ADOPTION_SIGNAL"
    c1_better = 0
    c0_better = 0
    harm = 0
    for row in comparable:
        c0m = float(row["c0_modification"] or 0)
        c1m = float(row["c1_modification"] or 0)
        c0u = bool(row.get("c0_usable"))
        c1u = bool(row.get("c1_usable"))
        if c1m > c0m + 0.02 or (row.get("c1_exact") and not row.get("c0_exact")) or (c1u and not c0u):
            c1_better += 1
        if c0m > c1m + 0.02 or (row.get("c0_exact") and not row.get("c1_exact")):
            c0_better += 1
        if row.get("c0_usable") and not row.get("c1_usable"):
            harm += 1
        if (row.get("c1_regression") or 1) < 0.99 <= (row.get("c0_regression") or 1):
            harm += 1
    if harm >= 2 and c1_better == 0:
        return "COMPILED_CONTEXT_HARMFUL_OR_DISTRACTING"
    if used and c1_better == 0 and c0_better == 0:
        return "COMPILED_CONTEXT_USED_NO_GAIN"
    if used and c1_better >= 1 and harm == 0:
        return "OPTIONAL_COMPILED_CONTEXT_EARNED"
    if used and c1_better < len(comparable) / 2:
        queried_frac = len(used) / max(1, len(adoption))
        if queried_frac < 0.5:
            return "CONTEXT_USEFUL_BUT_ADOPTION_LIMITED"
        return "MIXED_CONTEXT_VALUE"
    if not used:
        return "NO_ADOPTION_SIGNAL"
    return "MIXED_CONTEXT_VALUE"


def write_report(analysis: dict[str, Any], preflight_payload: dict[str, Any]) -> None:
    paired = analysis["paired"]
    lines = [
        "# Compiled-context sidecar A/B",
        "",
        f"Verdict: **{analysis['verdict']}**",
        "",
        "This experiment tests Spark on the normal SpreadsheetBench coding-agent scaffold, with vs without one optional compiled-context capability. It does not test the old LibreCalc harness, the autonomous compiled treatment, Task IR, ProgramGroups, a scheduler, Spark vs GLM, or published-control superiority.",
        "",
        "## Identity",
        "",
        f"- Declared model: `{MODEL}`",
        f"- Request model: `openrouter/{MODEL}`",
        "- Generation: temperature 0.0, top_p 1.0, reasoning_effort unset, tool_choice auto, provider allow_fallbacks + require_parameters",
        "- Envelope: 40 model calls / $2.50 per task (censoring boundary, not a target)",
        f"- Source identity run: `{IDENTITY['source_run']}`",
        "",
        "## Prompt hashes",
        "",
    ]
    hashes = preflight_payload.get("prompt_hashes") or prompt_hashes()
    for key, value in hashes.items():
        lines.append(f"- `{key}`: `{value}`")
    lines += ["", "## Paired official scores", "", "| Task | Repeat | C0 mod | C1 mod | C0 reg | C1 reg | C0 exact | C1 exact | C0 usable | C1 usable | C1 used calc_query |", "| --- | --- | --- | --- | --- | --- | --- | --- | --- | --- | --- |"]
    for row in paired:
        lines.append(
            f"| `{row['task']}` | {row['repeat']} | {row['c0_modification']} | {row['c1_modification']} | "
            f"{row['c0_regression']} | {row['c1_regression']} | {row['c0_exact']} | {row['c1_exact']} | "
            f"{row['c0_usable']} | {row['c1_usable']} | {row['c1_used_calc_query']} |"
        )
    lines += ["", "## Tool adoption (C1)", ""]
    used = sum(1 for row in analysis["adoption"] if row.get("used"))
    lines.append(f"{used}/{len(analysis['adoption'])} C1 trajectories used `calc_query`.")
    lines += ["", "## Financial_Model:13_05 mechanism", "", "| Repeat | Arm | periods queried | FY26–30 exposed | I20:M20 gold match | D20:H20 conflict | exact |", "| --- | --- | --- | --- | --- | --- | --- |"]
    for row in analysis["thirteen_05"]:
        lines.append(
            f"| {row['repeat']} | {row['arm']} | {row.get('queried_periods')} | {row.get('exposed_fy26_fy30')} | "
            f"{row.get('witness_I20_M20_matches_gold')} | {row.get('conflict_D20_H20_modified')} | {row.get('exact')} |"
        )
    lines += ["", "## Post-hoc opportunity audit", ""]
    if not analysis["opportunity_audit"]:
        lines.append("No failed C1 trajectories to audit, or scores were not yet available.")
    for row in analysis["opportunity_audit"]:
        lines.append(f"- `{row['key']}`: `{row['class']}` ({row.get('error_message') or 'no official error message'})")
    lines += [
        "",
        "## Claim boundary",
        "",
        "Causal claims are only from these fresh C0/C1 Spark pairs. Historical GLM or compiled-agent results explain task predeclaration only.",
        "",
        f"Generated at: {now()}",
        "",
    ]
    REPORT.write_text("\n".join(lines), encoding="utf-8")


def update_project_context() -> None:
    path = ROOT / "docs/PROJECT_CONTEXT.md"
    text = path.read_text(encoding="utf-8")
    needle = "## 0. Current handoff"
    paragraph = (
        "The autonomous compiled-agent architecture is frozen. The current hypothesis is whether a default general spreadsheet agent benefits from optional access to a deterministic compiled workbook sidecar.\n\n"
        "**Compiled-context sidecar A/B (2026-09-16).** C0 is the official SpreadsheetBench coding-agent scaffold; C1 adds one optional read-only `calc_query` tool over the existing compiled workbook world. No Task IR, Edit Plan, scheduler, ProgramGroups, `calc_apply`, or hidden structural packet. Spark identity is the successful standard-harness run (`meta/muse-spark-1.3-contributor`, temp 0, top_p 1, tool_choice auto, no reasoning-effort). Envelope 40/$2.50. Details: `COMPILED_CONTEXT_SIDECAR_AB_REPORT.md`.\n\n"
    )
    if "current hypothesis is whether a default general spreadsheet agent benefits" in text:
        return
    if needle not in text:
        return
    # Insert after the section header block's first blank line following the operator brief sentence.
    idx = text.find("§0 wins.")
    if idx == -1:
        path.write_text(paragraph + text, encoding="utf-8")
        return
    insert_at = text.find("\n\n", idx)
    if insert_at == -1:
        return
    path.write_text(text[: insert_at + 2] + paragraph + text[insert_at + 2 :], encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--preflight-only", action="store_true")
    parser.add_argument("--launch", action="store_true")
    parser.add_argument("--score", action="store_true")
    parser.add_argument("--report", action="store_true")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--workers", type=int, default=WORKERS)
    parser.add_argument("--all", action="store_true")
    args = parser.parse_args()
    ARTIFACT.mkdir(parents=True, exist_ok=True)
    jobs = build_jobs()
    if args.preflight_only or args.all or not any([args.launch, args.score, args.report]):
        payload = preflight()
        print(json.dumps({"preflight_passed": payload["passed"], "failed": payload["failed_tasks"]}, indent=2), flush=True)
        if not payload["passed"]:
            return 1
        if args.preflight_only:
            update_project_context()
            return 0
    else:
        payload = load_json(ARTIFACT / "preflight.json") if (ARTIFACT / "preflight.json").is_file() else preflight()
        if not payload.get("passed"):
            raise RuntimeError("Preflight has not passed")
    launch_payload: dict[str, Any] = {"results": []}
    if args.launch or args.all or args.dry_run:
        launch_payload = launch(jobs, workers=args.workers, dry_run=args.dry_run)
        if args.dry_run:
            return 0
    if args.score or args.all:
        scores = score_runs(jobs)
    else:
        scores = load_json(ARTIFACT / "official_scores.json") if (ARTIFACT / "official_scores.json").is_file() else {}
    if args.report or args.all:
        results = launch_payload.get("results") or (load_json(ARTIFACT / "launch_results.json").get("results") if (ARTIFACT / "launch_results.json").is_file() else [])
        analysis = analyze(jobs, scores, results)
        write_report(analysis, payload)
        update_project_context()
        print(f"VERDICT {analysis['verdict']} report={REPORT}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
