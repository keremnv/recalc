"""Offline characterisation of dependency-traversal ranking for Debugging tasks.

Ground truth comes from the OFFICIAL evaluator, not a reimplementation: dataset.json
supplies answer_position per task, and evaluation.classify_cells_by_modification
decides which cells inside those ranges count as modifications. Two earlier hand-rolled
definitions of "truth" were both wrong (ArrayFormula identity, and diffing outside the
scored ranges), which is exactly why this uses their code.

No model calls.
"""

from __future__ import annotations

import json
import sys
from collections import defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "benchmark-data/SpreadsheetBench-2/data/Debugging"
sys.path.insert(0, str(ROOT / "benchmark-data/SpreadsheetBench-2/evaluation"))
sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(ROOT / "src"))

import evaluation as ev
import openpyxl
from xlsx_metadata_repair import install as _install_repair

_install_repair()

from librecalc_mcp.domain.formulas import formula_a1_references

ERRORS = ("#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#N/A", "#NULL!", "#NUM!", "Err:")


def truth_cells(task: dict) -> set[tuple[str, str]]:
    inp = DATA / task["spreadsheet_path"]
    gold = DATA / task["golden_response_path"]
    path = task["spreadsheet_path"]
    with_color = "Color" in path
    with_formula = "Embedded" in path
    wb_i = openpyxl.load_workbook(inp, data_only=not with_formula)
    wb_a = openpyxl.load_workbook(gold, data_only=not with_formula)
    wb_if = openpyxl.load_workbook(inp, data_only=False) if not with_formula else None
    wb_af = openpyxl.load_workbook(gold, data_only=False) if not with_formula else None
    out: set[tuple[str, str]] = set()
    for chunk in ev.parse_answer_position(task["answer_position"]):
        sheet, _, rng = chunk.rpartition("!")
        sheet = sheet.strip().strip("'")
        _reg, mod = ev.classify_cells_by_modification(
            wb_i, wb_a, sheet, rng, with_color, with_formula, wb_if, wb_af
        )
        out.update((sheet, c) for c in mod)
    return out


def normalise(v):
    return getattr(v, "text", None) or v


def load(path: Path, data_only: bool):
    wb = openpyxl.load_workbook(path, data_only=data_only, read_only=True)
    out = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    out[(ws.title, c.coordinate)] = normalise(c.value)
    wb.close()
    return out


def has_err(v) -> bool:
    return isinstance(v, str) and any(t in v for t in ERRORS)


def rank(task: dict):
    inp = DATA / task["spreadsheet_path"]
    formulas = load(inp, data_only=False)
    values = load(inp, data_only=True)
    sheets = {s for s, _ in formulas}
    fwd, rev = defaultdict(set), defaultdict(set)
    for dep, f in formulas.items():
        if not (isinstance(f, str) and f.startswith("=")):
            continue
        for rs, start, end in formula_a1_references(f):
            src = rs or dep[0]
            if src not in sheets or end is not None:
                continue
            fwd[dep].add((src, start))
            rev[(src, start)].add(dep)
    errs = {k for k, v in formulas.items() if has_err(v)} | {
        k for k, v in values.items() if has_err(v)
    }
    roots = {c for c in errs if not (fwd.get(c, set()) & errs)}

    # NEW: roots only, ordered by blast radius.
    root_rank = sorted(roots, key=lambda c: -len(rev.get(c, ())))

    # SHIPPED BASELINE: what formula-anomalies-v1 emits -- error cells collapsed by
    # formula shape, most-repeated shape first, one representative per shape then the rest.
    import re

    def shape(cell):
        f = formulas.get(cell)
        return re.sub(r"\$?[A-Z]{1,3}\$?[0-9]+", "<REF>", f) if isinstance(f, str) else str(f)

    groups = defaultdict(list)
    for c in errs:
        groups[shape(c)].append(c)
    ordered = sorted(groups.items(), key=lambda kv: -len(kv[1]))
    shape_rank = [min(g) for _, g in ordered]
    shape_rank += [c for _, g in ordered for c in sorted(g)[1:]]
    return errs, roots, root_rank, shape_rank


def main() -> int:
    tasks = json.loads((DATA / "dataset.json").read_text())
    only = sys.argv[1:] or None
    print(
        f"{'task':8} {'defect':26} {'truth':>6} {'errs':>6} {'root':>6} "
        f"{'ROOTp@20':>9} {'SHAPEp@20':>10} {'delta':>7}"
    )
    print("-" * 88)
    for t in tasks:
        tid = t["id"]
        if only and tid not in only:
            continue
        defect = Path(t["spreadsheet_path"]).stem.replace("_input", "")
        try:
            truth = truth_cells(t)
            errs, roots, root_rank, shape_rank = rank(t)
        except Exception as exc:  # noqa: BLE001
            print(f"{tid:8} {defect[:26]:26} ERROR {type(exc).__name__}: {str(exc)[:40]}")
            continue
        if not errs:
            print(
                f"{tid:8} {defect[:26]:26} {len(truth):>6} {0:>6} {0:>6}   (no errors: signal silent)"
            )
            continue
        rp = len([c for c in root_rank[:20] if c in truth]) / 20
        sp = len([c for c in shape_rank[:20] if c in truth]) / 20
        print(
            f"{tid:8} {defect[:26]:26} {len(truth):>6} {len(errs):>6} {len(roots):>6} "
            f"{rp:>9.2f} {sp:>10.2f} {rp - sp:>+7.2f}"
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
