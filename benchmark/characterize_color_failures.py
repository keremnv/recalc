#!/usr/bin/env python3
"""Measure what font-color support can recover on Debugging color tasks.

This is an offline golden-side audit. It does not change the official evaluator or expose
goldens to the compiler. For each Inconsistent Color Coding task it separates:

* gold-side modification cells into value-only, color-only, or both;
* the stored run's remaining mismatches into value and color components;
* color regressions introduced on cells the golden preserves;
* gold-blind format-conventions-v1 candidate overlap against color-only gold mods.

The last distinction matters because LibreOffice can materialize workbook theme colors as
direct RGB on save. A task with only color mismatches is not necessarily recoverable by
correcting the intended color targets if the round trip also changed regression colors.
The observation overlap uses openpyxl font fingerprints, not LibreOffice RGB, so it is a
recall check for the census heuristic rather than a byte-identical inspect replay.
"""

from __future__ import annotations

import argparse
import json
import sys
import warnings
from collections import Counter
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "benchmark-data/SpreadsheetBench-2/data/Debugging"
DEFAULT_RUN = (
    ROOT
    / "benchmark-data/SpreadsheetBench-2/benchmark-runs/openrouter"
    / "gpt-5.6-sol-nonvisual-all-medium-1"
)

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "benchmark-data/SpreadsheetBench-2/evaluation"))

import evaluation as ev
from xlsx_metadata_repair import install

from librecalc_mcp.domain.observation import (
    _format_convention_workbook_observation,
    format_convention_requests,
)

install()
warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")


def _arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run", type=Path, default=DEFAULT_RUN)
    parser.add_argument(
        "--observation-overlap",
        action="store_true",
        help=(
            "Also score gold-blind format-conventions-v1 candidate overlap. "
            "This walks every populated input cell and is slow on large workbooks."
        ),
    )
    return parser.parse_args()


def _value_matches(value_a: Any, value_b: Any, raw_a: Any, raw_b: Any) -> bool:
    if ev._has_excel_error(value_a) or ev._has_excel_error(value_b):
        return ev.compare_cell_formula(raw_a, raw_b)
    return ev.compare_cell_value(value_a.value, value_b.value)


def _mismatch_kind(value_matches: bool, color_matches: bool) -> str | None:
    if value_matches and color_matches:
        return None
    if not value_matches and not color_matches:
        return "both"
    return "value" if not value_matches else "color"


def _analyse_task(task: dict[str, Any], run: Path) -> tuple[Counter, Counter, set[tuple[str, str]]]:
    task_id = str(task["id"])
    paths = (
        DATA / task["spreadsheet_path"],
        DATA / task["golden_response_path"],
        run / f"Debugging-{task_id}" / "output.xlsx",
    )
    input_values, golden_values, output_values = [
        openpyxl.load_workbook(path, data_only=True) for path in paths
    ]
    input_raw, golden_raw, output_raw = [
        openpyxl.load_workbook(path, data_only=False) for path in paths
    ]

    gold = Counter()
    output = Counter()
    color_only: set[tuple[str, str]] = set()
    for chunk in ev.parse_answer_position(task["answer_position"]):
        sheet_name, cell_range = chunk.split("!", 1)
        sheet_name = sheet_name.strip("'").strip()
        cell_range = cell_range.strip("'").strip()
        regression, modification = ev.classify_cells_by_modification(
            input_values,
            golden_values,
            sheet_name,
            cell_range,
            True,
            False,
            input_raw,
            golden_raw,
        )
        input_value_sheet = ev._find_sheet(input_values, sheet_name)
        golden_value_sheet = ev._find_sheet(golden_values, sheet_name)
        output_value_sheet = ev._find_sheet(output_values, sheet_name)
        input_raw_sheet = ev._find_sheet(input_raw, sheet_name)
        golden_raw_sheet = ev._find_sheet(golden_raw, sheet_name)
        output_raw_sheet = ev._find_sheet(output_raw, sheet_name)
        assert None not in (
            input_value_sheet,
            golden_value_sheet,
            output_value_sheet,
            input_raw_sheet,
            golden_raw_sheet,
            output_raw_sheet,
        )

        for population, addresses in (("reg", regression), ("mod", modification)):
            for address in addresses:
                if population == "mod":
                    kind = _mismatch_kind(
                        _value_matches(
                            input_value_sheet[address],
                            golden_value_sheet[address],
                            input_raw_sheet[address],
                            golden_raw_sheet[address],
                        ),
                        ev.compare_font_color(
                            input_raw_sheet[address].font,
                            golden_raw_sheet[address].font,
                        ),
                    )
                    assert kind is not None
                    gold[kind] += 1
                    if kind == "color":
                        color_only.add((sheet_name, address))

                kind = _mismatch_kind(
                    _value_matches(
                        golden_value_sheet[address],
                        output_value_sheet[address],
                        golden_raw_sheet[address],
                        output_raw_sheet[address],
                    ),
                    ev.compare_font_color(
                        golden_raw_sheet[address].font,
                        output_raw_sheet[address].font,
                    ),
                )
                if kind is not None:
                    output[f"{population}_{kind}"] += 1
    return gold, output, color_only


def _openpyxl_font_hex(font: Any) -> str | None:
    color = getattr(font, "color", None)
    if color is None or color.type is None:
        return None
    if color.type == "rgb" and color.rgb:
        rgb = str(color.rgb).upper()
        if len(rgb) == 8 and rgb.startswith("FF"):
            rgb = rgb[2:]
        return rgb if rgb.startswith("#") else f"#{rgb}"
    if color.type == "theme":
        return f"theme:{color.theme}:{round(color.tint or 0, 4)}"
    if color.type == "indexed":
        return f"indexed:{color.indexed}"
    return None


def _input_observation(path: Path) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(path, data_only=False)
    sheets: list[dict[str, Any]] = []
    formats: dict[tuple[str, str], dict[str, Any]] = {}
    for worksheet in workbook.worksheets:
        max_row = worksheet.max_row
        max_column = worksheet.max_column
        if not max_row or not max_column:
            continue
        values: list[list[Any]] = []
        formulas: list[list[Any]] = []
        for row in worksheet.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_column):
            value_row: list[Any] = []
            formula_row: list[Any] = []
            for cell in row:
                raw = cell.value
                if isinstance(raw, str) and raw.startswith("="):
                    formula_row.append(raw)
                    value_row.append(None)
                else:
                    formula_row.append(None)
                    value_row.append(raw)
                color = _openpyxl_font_hex(cell.font)
                if color is not None:
                    formats[(worksheet.title, cell.coordinate)] = {"font_color": color}
            values.append(value_row)
            formulas.append(formula_row)
        sheets.append(
            {
                "name": worksheet.title,
                "used_range": f"A1:{get_column_letter(max_column)}{max_row}",
                "result": {"values": values, "formulas": formulas},
            }
        )
    requests, populated = format_convention_requests(sheets)
    formats = {key: formats.get(key, {}) for key in requests}
    return _format_convention_workbook_observation(
        title=path.name,
        url=None,
        sheets=sheets,
        formats=formats,
        populated_cell_count=populated,
    )


def _expand_candidate(sheet: str, cell_range: str) -> set[tuple[str, str]]:
    minimum_column, minimum_row, maximum_column, maximum_row = range_boundaries(cell_range)
    return {
        (sheet, f"{get_column_letter(column)}{row}")
        for row in range(minimum_row, maximum_row + 1)
        for column in range(minimum_column, maximum_column + 1)
    }


def main() -> int:
    args = _arguments()
    tasks = {str(task["id"]): task for task in json.loads((DATA / "dataset.json").read_text())}
    task_ids = [f"{number:02d}_04" for number in range(1, 11)]
    print(
        "task   gold modifications (color/value/both)   "
        "remaining (color/value)   regression-color   intended-color ceiling",
        flush=True,
    )
    intended_recoverable = 0
    oracle_recoverable = 0
    gold_total = Counter()
    output_total = Counter()
    overlap_rows: list[str] = []
    for task_id in task_ids:
        gold, output, color_only = _analyse_task(tasks[task_id], args.run)
        gold_total.update(gold)
        output_total.update(output)
        remaining_color = output["mod_color"] + output["reg_color"]
        remaining_value = sum(
            count for kind, count in output.items() if kind.endswith(("_value", "_both"))
        )
        intended = remaining_value == 0 and output["reg_color"] == 0
        oracle = remaining_value == 0
        intended_recoverable += int(intended)
        oracle_recoverable += int(oracle)
        print(
            f"{task_id}   {gold['color']:5}/{gold['value']:5}/{gold['both']:4}"
            f"                 {remaining_color:5}/{remaining_value:5}"
            f"             {output['reg_color']:5}"
            f"              {'yes' if intended else 'no'}",
            flush=True,
        )
        if args.observation_overlap:
            observation = _input_observation(DATA / tasks[task_id]["spreadsheet_path"])
            candidate_cells = {
                cell
                for candidate in observation["format_outlier_candidates"]
                for cell in _expand_candidate(candidate["sheet"], candidate["range"])
            }
            hits = color_only & candidate_cells
            star_sheets = sum(
                1
                for candidate in observation["format_outlier_candidates"]
                if candidate["sheet"] in {None, "*"}
            )
            overlap_rows.append(
                f"{task_id}   color-only {len(color_only):4}  "
                f"candidates {len(candidate_cells):5}  "
                f"hit {len(hits):4}  "
                f"star-sheets {star_sheets}"
            )

    print(f"\nGold-side modifications: {dict(gold_total)}")
    print(f"Stored-output mismatches: {dict(output_total)}")
    print(f"Recoverable by correcting intended color targets only: {intended_recoverable}/10")
    print(
        "Recoverable if an oracle also repairs color regressions introduced on save: "
        f"{oracle_recoverable}/10"
    )
    if overlap_rows:
        print("\nGold-blind format-conventions-v1 overlap against color-only gold mods:")
        print("\n".join(overlap_rows))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
