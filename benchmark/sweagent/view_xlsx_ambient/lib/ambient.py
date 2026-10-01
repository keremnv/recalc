"""Ambient formula-equivalence appendix for a view_xlsx content window.

The payload describes equivalence classes that already intersect the inspected
cells. It does not rank, recommend, or navigate toward a target.
"""
from __future__ import annotations

import argparse
import shlex
from pathlib import Path
from typing import Any

from fingerprint import a1_address, formula_text
from ranges import compress_cells
from workbook import EquivalenceClass, WorkbookIndex, build_index

OBSERVATION_LIMIT = 10_000


def parse_view_xlsx_argv(argv: list[str]) -> argparse.Namespace:
    """Match the official view_xlsx positional/flag parser."""
    parser = argparse.ArgumentParser(add_help=False)
    parser.add_argument("file_path")
    parser.add_argument("--mode", choices=["list", "content"], dest="mode_opt")
    parser.add_argument("--sheet", dest="sheet_opt")
    parser.add_argument("--start_row", type=int, dest="start_row_opt")
    parser.add_argument("--end_row", type=int, dest="end_row_opt")
    parser.add_argument("mode", nargs="?")
    parser.add_argument("sheet", nargs="?")
    parser.add_argument("start_row", nargs="?", type=int)
    parser.add_argument("end_row", nargs="?", type=int)
    args = parser.parse_args(argv)
    raw_mode = args.mode_opt or args.mode
    sheet = args.sheet_opt or args.sheet
    start_row = args.start_row_opt if args.start_row_opt is not None else args.start_row
    end_row = args.end_row_opt if args.end_row_opt is not None else args.end_row
    if raw_mode in (None, "list", "content"):
        mode = raw_mode or "content"
    else:
        mode = "content"
        if sheet is None:
            sheet = raw_mode
    args.mode = mode
    args.sheet = sheet
    args.start_row = start_row
    args.end_row = end_row
    return args


def parse_view_xlsx_action(action: str) -> argparse.Namespace | None:
    text = str(action or "").strip()
    if not text.startswith("view_xlsx"):
        return None
    try:
        parts = shlex.split(text)
    except ValueError:
        return None
    if not parts:
        return None
    try:
        return parse_view_xlsx_argv(parts[1:])
    except SystemExit:
        return None


def data_bounds(sheet) -> tuple[int | None, int | None, int | None, int | None]:
    if sheet.max_row is None or sheet.max_column is None:
        return None, None, None, None
    min_row = max_row = min_col = max_col = None
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is not None:
                if min_row is None or cell.row < min_row:
                    min_row = cell.row
                if max_row is None or cell.row > max_row:
                    max_row = cell.row
                if min_col is None or cell.column < min_col:
                    min_col = cell.column
                if max_col is None or cell.column > max_col:
                    max_col = cell.column
    return min_row, max_row, min_col, max_col


def match_sheet(workbook, sheet_name: str | None):
    hidden_states = ("hidden", "veryHidden")
    visible = [
        name for name in workbook.sheetnames if workbook[name].sheet_state not in hidden_states
    ]
    if not sheet_name:
        if not visible:
            raise ValueError("No visible sheets in workbook")
        return workbook[visible[0]]
    if sheet_name in workbook.sheetnames:
        return workbook[sheet_name]
    stripped = sheet_name.strip()
    if stripped in workbook.sheetnames:
        return workbook[stripped]
    for actual in workbook.sheetnames:
        if actual.strip() == stripped:
            return workbook[actual]
    raise ValueError(f"Sheet {sheet_name!r} not found")


def resolve_content_window(
    path: Path | str,
    *,
    sheet: str | None = None,
    start_row: int | None = None,
    end_row: int | None = None,
) -> dict[str, Any]:
    from xlsx_metadata_repair import install as install_repair

    install_repair()
    import openpyxl

    workbook = openpyxl.load_workbook(path, data_only=False)
    try:
        target = match_sheet(workbook, sheet)
        min_row, max_row, min_col, max_col = data_bounds(target)
        if min_row is None:
            return {
                "sheet": target.title,
                "empty": True,
                "c1": 1,
                "r1": 1,
                "c2": 1,
                "r2": 1,
            }
        actual_start = start_row if start_row is not None else min_row
        actual_end = end_row if end_row is not None else max_row
        if actual_start < 1:
            actual_start = 1
        if actual_end < actual_start:
            actual_end = actual_start
        c1 = min_col if min_col else 1
        c2 = max_col if max_col else target.max_column
        return {
            "sheet": target.title,
            "empty": False,
            "c1": int(c1),
            "r1": int(actual_start),
            "c2": int(c2),
            "r2": int(actual_end),
        }
    finally:
        workbook.close()


def _visible_cells(
    record: EquivalenceClass, sheet: str, c1: int, r1: int, c2: int, r2: int
) -> list[tuple[int, int]]:
    cells = [
        (col, row)
        for col, row in record.members.get(sheet, set())
        if c1 <= col <= c2 and r1 <= row <= r2
    ]
    cells.sort(key=lambda item: (item[1], item[0]))
    return cells


def member_range_lines(record: EquivalenceClass) -> list[str]:
    lines: list[str] = []
    for sheet in sorted(record.members):
        for spec in compress_cells(record.members[sheet]):
            lines.append(f"    {sheet}!{spec}")
    return lines


def _class_block(
    record: EquivalenceClass,
    *,
    visible_source: str,
    formula: str,
    members: bool,
    formula_limit: int | None = None,
) -> str:
    shown = formula
    if formula_limit is not None and len(shown) > formula_limit:
        shown = shown[: formula_limit - 1] + "…"
    flags = []
    if record.opaque:
        flags.append("opaque=true")
        if record.reason:
            flags.append(f"reason={record.reason}")
    header = f"eq {record.eq_id}"
    if flags:
        header += " " + " ".join(flags)
    lines = [
        header,
        f"  visible_source: {visible_source}",
        f"  formula: {shown}",
        f"  instances: {record.n}",
    ]
    if members:
        lines.append("  members:")
        lines.extend(member_range_lines(record) or ["    (none)"])
    return "\n".join(lines) + "\n"


def render_ambient(
    records: list[tuple[EquivalenceClass, str, str]],
    *,
    sheet: str,
    c1: int,
    r1: int,
    c2: int,
    r2: int,
    limit: int = OBSERVATION_LIMIT,
) -> str:
    col_start = a1_address(c1, 1)[:-1]
    col_end = a1_address(c2, 1)[:-1]
    header_core = (
        f"STRUCTURAL INDEX sheet={sheet} rows={r1}:{r2} "
        f"cols={col_start}:{col_end} classes={len(records)}"
    )

    def wrap(body: str, truncated: bool, omitted: int) -> str:
        return (
            f"{header_core} truncated={'true' if truncated else 'false'} "
            f"omitted_member_ranges={omitted}\n{body}"
        )

    if not records:
        return wrap("", False, 0)

    full_body = "".join(
        _class_block(record, visible_source=source, formula=formula, members=True)
        for record, source, formula in records
    )
    full = wrap(full_body, False, 0)
    if limit > 0 and len(full) <= limit:
        return full
    if limit <= 0:
        limit = 1

    included_members = len(records) - 1
    while included_members >= 0:
        parts: list[str] = []
        omitted = 0
        for index, (record, source, formula) in enumerate(records):
            if index < included_members:
                parts.append(
                    _class_block(record, visible_source=source, formula=formula, members=True)
                )
            else:
                omitted += len(member_range_lines(record))
                parts.append(
                    _class_block(record, visible_source=source, formula=formula, members=False)
                )
        text = wrap("".join(parts), True, omitted)
        if len(text) <= limit:
            return text
        included_members -= 1

    for formula_limit in (200, 80, 24, 8):
        omitted = sum(len(member_range_lines(record)) for record, _s, _f in records)
        body = "".join(
            _class_block(
                record,
                visible_source=source,
                formula=formula,
                members=False,
                formula_limit=formula_limit,
            )
            for record, source, formula in records
        )
        text = wrap(body, True, omitted)
        if len(text) <= limit:
            return text

    omitted = sum(len(member_range_lines(record)) for record, _s, _f in records)
    compact = "".join(
        f"eq {record.eq_id} instances={record.n}"
        f"{' opaque=true' if record.opaque else ''}\n"
        for record, _s, _f in records
    )
    text = wrap(compact, True, omitted)
    if len(text) <= limit:
        return text
    return text


def visible_class_records(
    index: WorkbookIndex,
    formulas: dict[tuple[str, int, int], str],
    *,
    sheet: str,
    c1: int,
    r1: int,
    c2: int,
    r2: int,
) -> list[tuple[EquivalenceClass, str, str]]:
    spec = f"{a1_address(c1, r1)}:{a1_address(c2, r2)}"
    records = []
    for record in index.classes_in_range(sheet, spec):
        visible = _visible_cells(record, sheet, c1, r1, c2, r2)
        if not visible:
            continue
        col, row = visible[0]
        source = f"{sheet}!{a1_address(col, row)}"
        formula = formulas.get((sheet, col, row)) or record.exemplar_formula
        records.append((record, source, formula))
    return records


def formulas_from_workbook(path: Path | str) -> dict[tuple[str, int, int], str]:
    from xlsx_metadata_repair import install as install_repair

    install_repair()
    import openpyxl

    workbook = openpyxl.load_workbook(path, data_only=False)
    out: dict[tuple[str, int, int], str] = {}
    try:
        for sheet in workbook.worksheets:
            for cell in sheet._cells.values():
                text = formula_text(cell.value)
                if text:
                    out[(sheet.title, int(cell.column), int(cell.row))] = text
    finally:
        workbook.close()
    return out


def build_ambient_payload(
    path: Path | str,
    *,
    sheet: str | None = None,
    start_row: int | None = None,
    end_row: int | None = None,
    limit: int = OBSERVATION_LIMIT,
    index: WorkbookIndex | None = None,
    formulas: dict[tuple[str, int, int], str] | None = None,
) -> dict[str, Any]:
    window = resolve_content_window(
        path, sheet=sheet, start_row=start_row, end_row=end_row
    )
    if index is None:
        index = build_index(path)
    if formulas is None:
        formulas = formulas_from_workbook(path)
    if window["empty"]:
        text = render_ambient(
            [],
            sheet=window["sheet"],
            c1=1,
            r1=1,
            c2=1,
            r2=1,
            limit=limit,
        )
        return {"window": window, "eq_ids": [], "text": text, "classes": 0}
    records = visible_class_records(
        index,
        formulas,
        sheet=window["sheet"],
        c1=window["c1"],
        r1=window["r1"],
        c2=window["c2"],
        r2=window["r2"],
    )
    text = render_ambient(
        records,
        sheet=window["sheet"],
        c1=window["c1"],
        r1=window["r1"],
        c2=window["c2"],
        r2=window["r2"],
        limit=limit,
    )
    return {
        "window": window,
        "eq_ids": [record.eq_id for record, _s, _f in records],
        "text": text,
        "classes": len(records),
        "opaque_classes": sum(1 for record, _s, _f in records if record.opaque),
    }
