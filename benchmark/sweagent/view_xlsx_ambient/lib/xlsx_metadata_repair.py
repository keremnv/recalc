"""Make openpyxl tolerant of workbooks whose docProps metadata has an unbound prefix.

Financial_Model tasks 06_01..06_05 ship a docProps/core.xml whose root element declares
only xmlns:cp, while <dc:creator> and <cp:lastModifiedBy> near the end use prefixes that
are bound nowhere -- dc: is declared locally on <dc:language> and goes out of scope. Python's
expat rejects the document, so openpyxl.load_workbook raises ParseError and the whole file is
unreadable. The grid is intact; only the metadata part is malformed. LibreOffice opens these
files without complaint, which is why the product never saw the problem and the official
evaluator (evaluation.py:362) cannot score those five tasks at all.

install() makes load_workbook retry once against a repaired copy: unbound prefixes in the
failing part get their standard OOXML namespace bound on the root element, in a temp file.
Nothing under benchmark-data/ is written to -- the dataset stays exactly as distributed.
"""

from __future__ import annotations

import atexit
import os
import re
import shutil
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

import openpyxl

# The prefixes OOXML metadata parts use. Binding one of these cannot change cell data:
# every part we repair lives under docProps/ or _rels/.
NAMESPACES = {
    "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
    "dc": "http://purl.org/dc/elements/1.1/",
    "dcterms": "http://purl.org/dc/terms/",
    "dcmitype": "http://purl.org/dc/dcmitype/",
    "xsi": "http://www.w3.org/2001/XMLSchema-instance",
    "vt": "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes",
}

_PREFIX = re.compile(r"<\s*/?\s*([A-Za-z_][\w.-]*):|\s([A-Za-z_][\w.-]*):[\w.-]+\s*=")
_ROOT = re.compile(r"<\s*([A-Za-z_][\w.:-]*)((?:\s[^<>]*?)?)(/?)>", re.DOTALL)
_original = None
_scratch: Path | None = None


def _repair_part(xml: str) -> str | None:
    """Bind every prefix the part uses on its root element. None if nothing to add."""
    used = {m.group(1) or m.group(2) for m in _PREFIX.finditer(xml)} - {None, "xmlns", "xml"}
    match = _ROOT.search(xml)
    if match is None:
        return None
    declared = set(re.findall(r"xmlns:([\w.-]+)\s*=", match.group(0)))
    missing = sorted(p for p in used if p not in declared and p in NAMESPACES)
    if not missing:
        return None
    additions = "".join(f' xmlns:{p}="{NAMESPACES[p]}"' for p in missing)
    start, end = match.span()
    repaired = f"<{match.group(1)}{match.group(2)}{additions}{match.group(3)}>"
    return xml[:start] + repaired + xml[end:]


def repair(path: Path, out_dir: Path) -> Path | None:
    """Write a copy of `path` with malformed metadata parts repaired. None if unrepairable."""
    repairs: dict[str, str] = {}
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            # This compatibility shim is intentionally metadata-only. A parse failure in
            # workbook content must remain visible rather than being silently rewritten.
            if not name.startswith("docProps/") or not name.endswith(".xml"):
                continue
            raw = archive.read(name)
            try:
                ET.fromstring(raw)
                continue
            except ET.ParseError:
                pass
            fixed = _repair_part(raw.decode("utf-8", "replace"))
            if fixed is None:
                return None
            try:
                ET.fromstring(fixed)
            except ET.ParseError:
                return None
            repairs[name] = fixed
        if not repairs:
            return None
        out_dir.mkdir(parents=True, exist_ok=True)
        target = out_dir / path.name
        with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as out:
            for item in archive.infolist():
                data = archive.read(item.filename)
                if item.filename in repairs:
                    data = repairs[item.filename].encode("utf-8")
                out.writestr(item, data)
    return target


def install() -> None:
    """Patch openpyxl.load_workbook to retry once through repair(). Idempotent."""
    global _original, _scratch
    if _original is not None:
        return
    _original = openpyxl.load_workbook
    cache: dict[str, Path] = {}
    _scratch = Path(tempfile.mkdtemp(prefix="xlsx-metadata-repair-"))

    def load_workbook(filename, *args, **kwargs):
        try:
            return _original(filename, *args, **kwargs)
        except ET.ParseError:
            if not isinstance(filename, (str, os.PathLike)):
                raise
            source = Path(filename)
            key = str(source.resolve())
            if key not in cache:
                target = repair(source, _scratch)
                if target is None:
                    raise
                cache[key] = target
            return _original(cache[key], *args, **kwargs)

    openpyxl.load_workbook = load_workbook
    # The evaluator calls openpyxl.load_workbook through the module attribute, so this is
    # sufficient and can be fully undone. Do not mutate arbitrary imported modules.
    atexit.register(uninstall)


def uninstall() -> None:
    global _original, _scratch
    if _original is not None:
        openpyxl.load_workbook = _original
        _original = None
    if _scratch is not None:
        shutil.rmtree(_scratch, ignore_errors=True)
        _scratch = None


__all__ = ["install", "repair", "uninstall"]
