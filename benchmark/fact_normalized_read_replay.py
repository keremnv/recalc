#!/usr/bin/env python3
"""Zero-model fact-normalized read-side replay.

This is an evidence harness only. It does not change prompts, expose a new
helper, or run model inference. It reconstructs the facts visible in the
frozen H0 observations, checks whether the explicit pre-observation request
can mechanically produce them, and compares representations of that same H.
"""
from __future__ import annotations

import ast
import datetime as dt
import hashlib
import json
import math
import re
import statistics
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from benchmark.inspection_helpers import index as index_module
from benchmark.read_side_v2_adjudication import (
    _address_key,
    _facts_for_requests,
    _line_value_facts,
    _sheet_variables,
    compact,
    extract_sheets,
    historical_records,
    replay_v2,
)

OUT = PROJECT_ROOT / "research/history/fact_normalized_read_replay"
WORK = PROJECT_ROOT / "research/history/thin_architecture_checkpoint" / "work"
RANKING = PROJECT_ROOT / "research/history/thin_architecture_checkpoint" / "exposure_ranking.json"
POPULATION = PROJECT_ROOT / "research/history/thin_architecture_checkpoint" / "population.json"
MAX_CELLS = 10_000
TASKS = [
    "Financial_Model:08_01",
    "Debugging:10_04",
    "Financial_Model:08_03",
    "Financial_Model:07_01",
    "Debugging:10_10",
    "Financial_Model:15_04",
]

ROW_LIST_RE = re.compile(r"^\s*(\d+)\s+(\[.*\])\s*$")
CELL_LINE_RE = re.compile(
    r"^\s*(?:[^\s]+\s+)?([A-Z]{1,3}\d+)\s*:?\s*(.+?)\s*$"
)
OBS_HEAD_RE = re.compile(r"<observation_head>\n(.*?)(?:\n</observation_head>|\n<elided_chars>)", re.DOTALL)
OBS_TAIL_RE = re.compile(r"<observation_tail>\n(.*?)(?:\n</observation_tail>|\n</observation>)", re.DOTALL)


def write_json(name: str, value: Any) -> None:
    (OUT / name).write_text(
        json.dumps(value, ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )


def write_jsonl(name: str, rows: list[dict[str, Any]]) -> None:
    with (OUT / name).open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False, sort_keys=True, default=str) + "\n")


def sha256_bytes(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def norm_scalar(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return str(value)
    return str(value)


def parse_literal(raw: str) -> Any:
    raw = raw.strip()
    try:
        return ast.literal_eval(raw)
    except (SyntaxError, ValueError):
        try:
            return eval(
                raw,
                {"__builtins__": {}, "datetime": dt.datetime, "date": dt.date},
            )
        except (AttributeError, NameError, SyntaxError, TypeError, ValueError):
            return raw


def raw_fact(sheet: str, address: str, value: Any) -> dict[str, Any] | None:
    if value is None:
        return None
    if isinstance(value, str) and value.startswith("="):
        return {
            "sheet": sheet,
            "address": address,
            "formula": value,
            "visible_fields": ["sheet", "address", "formula"],
        }
    return {
        "sheet": sheet,
        "address": address,
        "value": norm_scalar(value),
        "visible_fields": ["sheet", "address", "value"],
    }


def observation_body(observation: str) -> str:
    # A normal observation is already complete. For capped output retain only
    # the visible head and tail; hidden cells never enter H.
    head = OBS_HEAD_RE.search(observation)
    tail = OBS_TAIL_RE.search(observation)
    if head or tail:
        return "\n".join(part.group(1) for part in (head, tail) if part)
    return observation


def range_bounds(request: dict[str, str]) -> tuple[int, int, int, int]:
    from openpyxl.utils.cell import range_boundaries

    min_col, min_row, max_col, max_row = range_boundaries(request["range"])
    return min_col, min_row, max_col, max_row


def view_facts(operation: dict[str, Any]) -> tuple[list[dict[str, Any]], int, str]:
    request = operation["requests"][0]
    min_col, _, _, _ = range_bounds(request)
    from openpyxl.utils import get_column_letter

    facts: list[dict[str, Any]] = []
    empty_slots = 0
    body = observation_body(operation["observation"])
    for line in body.splitlines():
        match = re.match(r"\s*Row\s+(\d+):\s+(\[.*\])\s*$", line)
        if not match:
            continue
        values = parse_literal(match.group(2))
        if not isinstance(values, list):
            continue
        row = int(match.group(1))
        for offset, value in enumerate(values):
            address = f"{get_column_letter(min_col + offset)}{row}"
            if value is None:
                empty_slots += 1
                continue
            fact = raw_fact(request["sheet"], address, value)
            if fact:
                facts.append(fact)
    if not facts:
        status = "UNRESOLVED"
    elif operation["historical_truncated"]:
        status = "CAPPED_VISIBLE"
    else:
        status = "FULL_VISIBLE"
    return facts, empty_slots, status


def _request_for_row(requests: list[dict[str, str]], row: int, width: int) -> dict[str, str] | None:
    candidates = []
    for request in requests:
        min_col, min_row, max_col, max_row = range_bounds(request)
        if min_row <= row <= max_row and max_col - min_col + 1 >= width:
            candidates.append((max_col - min_col + 1, request))
    return min(candidates, key=lambda item: item[0])[1] if candidates else None


def bash_facts(operation: dict[str, Any]) -> tuple[list[dict[str, Any]], int, str, str]:
    command = str(operation["args"].get("command", ""))
    requests = operation["requests"]
    sheets = []
    for request in requests:
        if request["sheet"] not in sheets:
            sheets.append(request["sheet"])
    if not sheets:
        sheets = extract_sheets(command)
    current_sheet = sheets[0] if len(sheets) == 1 else None
    facts: list[dict[str, Any]] = []
    empty_slots = 0
    body = observation_body(operation["observation"])
    for line in body.splitlines():
        stripped = line.strip()
        if stripped.startswith("==="):
            label = stripped.strip("= ")
            if label in sheets:
                current_sheet = label
            else:
                for sheet in sheets:
                    if label.startswith(sheet + " "):
                        current_sheet = sheet
                        break
        # Python prints such as ``[("B4", "=..."), ("C4", 12)]`` are the
        # historical visible representation. Parse the complete line so a
        # repr-truncated formula remains truncated in H rather than being
        # silently replaced by the full workbook formula.
        tuple_source = stripped
        list_start = stripped.find("[(")
        if list_start > 0:
            tuple_source = stripped[list_start:]
        parsed_line = parse_literal(tuple_source) if tuple_source.startswith(("[(", "(", "[")) else None
        tuple_items: list[Any] = []
        if isinstance(parsed_line, tuple) and len(parsed_line) == 2:
            tuple_items = [parsed_line]
        elif isinstance(parsed_line, list):
            tuple_items = [item for item in parsed_line if isinstance(item, (tuple, list)) and len(item) == 2]
        if tuple_items and current_sheet:
            for address, value in tuple_items:
                if isinstance(address, str) and re.fullmatch(r"[A-Z]{1,3}\d+", address):
                    fact = raw_fact(current_sheet, address, value)
                    if fact:
                        facts.append(fact)
            continue
        row_match = ROW_LIST_RE.match(line)
        if row_match and current_sheet:
            row = int(row_match.group(1))
            values = parse_literal(row_match.group(2))
            if isinstance(values, list):
                request = _request_for_row(requests, row, len(values))
                if request:
                    min_col, _, _, _ = range_bounds(request)
                    from openpyxl.utils import get_column_letter

                    for offset, value in enumerate(values):
                        if value is None:
                            empty_slots += 1
                            continue
                        fact = raw_fact(
                            request["sheet"],
                            f"{get_column_letter(min_col + offset)}{row}",
                            value,
                        )
                        if fact:
                            facts.append(fact)
            continue
        for address, raw_value in _line_value_facts(line, current_sheet):
            value = parse_literal(raw_value)
            sheet = current_sheet
            if not sheet:
                continue
            fact = raw_fact(sheet, address, value)
            if fact:
                facts.append(fact)
    # Stable de-duplication: a line may be visible in both the capped head and
    # tail, or appear in a tuple and a coordinate-prefixed line.
    unique = {(f["sheet"], f["address"], compact(f)): f for f in facts}
    facts = sorted(unique.values(), key=lambda f: (f["sheet"], _address_key(f["address"])))
    if not facts:
        status = "METADATA_ONLY" if requests else "UNRESOLVED"
    elif operation["historical_truncated"]:
        status = "CAPPED_VISIBLE"
    elif "[('" in body or "[(\"" in body:
        status = "RANGE_TABLE"
    else:
        status = "VALUE_ONLY"
    return facts, empty_slots, status, command


def event_facts(operation: dict[str, Any]) -> tuple[list[dict[str, Any]], int, str, str]:
    if operation["tool"] == "view_xlsx":
        facts, empty_slots, status = view_facts(operation)
        return facts, empty_slots, status, "observation-row-parser"
    facts, empty_slots, status, command = bash_facts(operation)
    return facts, empty_slots, status, f"observation-text-parser:{command[:80]}"


def fact_key(fact: dict[str, Any]) -> tuple[Any, ...]:
    visible = tuple(f for f in ("value", "formula") if f in fact)
    return tuple([fact.get("sheet"), fact.get("address")] + [fact.get(f) for f in visible])


def projected_fact(fact: dict[str, Any], visible_fields: list[str]) -> dict[str, Any]:
    out = {field: fact[field] for field in visible_fields if field in fact}
    out["visible_fields"] = visible_fields
    return out


def canonical_fact_bytes(facts: list[dict[str, Any]]) -> int:
    return len(compact(facts).encode("utf-8"))


def verbose_bytes(facts: list[dict[str, Any]]) -> int:
    rows = [{key: value for key, value in fact.items() if key != "visible_fields"} for fact in facts]
    return len(compact({"results": rows, "truncated": False}).encode("utf-8"))


def compact_exact_h_bytes(facts: list[dict[str, Any]]) -> int:
    if not facts:
        return len(compact({"fields": [], "cells": []}).encode("utf-8"))
    order = ["sheet", "address", "row", "col", "value", "formula", "dtype"]
    fields = [field for field in order if any(field in fact for fact in facts)]
    cells = [[fact.get(field) for field in fields] for fact in facts]
    return len(compact({"fields": fields, "cells": cells}).encode("utf-8"))


def source_fact_map(path: Path, requests: list[dict[str, str]]) -> dict[tuple[str, str], dict[str, Any]]:
    return {
        (fact["sheet"], fact["address"]): fact
        for fact in _facts_for_requests(path, requests)
    }


def explicit_query_addresses(path: Path, requests: list[dict[str, str]]) -> set[tuple[str, str]]:
    return set(source_fact_map(path, requests))


def query_realizability(
    operation: dict[str, Any],
    facts: list[dict[str, Any]],
    status: str,
    path: Path,
) -> tuple[str, str]:
    if not facts:
        return ("AMBIGUOUS" if operation["requests"] else "UNRESOLVED", "no_visible_cell_facts")
    if operation["historical_truncated"]:
        return "REQUIRES_HINDSIGHT", "visible_prefix_is_capped"
    h_addresses = {(fact["sheet"], fact["address"]) for fact in facts}
    q_addresses = explicit_query_addresses(path, operation["requests"])
    command = str(operation["args"].get("command", ""))
    if h_addresses == q_addresses:
        return "DIRECTLY_REALIZABLE", "explicit_request_determines_visible_fact_entities"
    if operation["tool"] == "view_xlsx":
        return "AMBIGUOUS", "full_range_visible_entity_set_does_not_match_request_envelope"
    # Only code-visible deterministic filters qualify for projection. The H
    # address set itself is never used as a filter key.
    mechanical_markers = (
        " if ", "if ", "[:", "[-", "cols[:", "cols[-", "startswith(",
        "endswith(", "==", " in [",
    )
    if any(marker in command for marker in mechanical_markers):
        return "REALIZABLE_WITH_MECHANICAL_PROJECTION", "explicit_code_filter_or_slice"
    return "REQUIRES_HINDSIGHT", "historical_visible_subset_not_in_query"


def normalize_candidate_fact(fact: dict[str, Any]) -> dict[str, Any]:
    out = {"sheet": fact.get("sheet"), "address": fact.get("address")}
    if fact.get("formula") is not None:
        out["formula"] = str(fact["formula"])
    elif fact.get("value") is not None:
        out["value"] = norm_scalar(fact["value"])
    return out


def compare_visible_facts(
    historical: list[dict[str, Any]],
    compiled: list[dict[str, Any]],
) -> tuple[str, list[dict[str, Any]], list[dict[str, Any]]]:
    compiled_map = {
        (fact.get("sheet"), fact.get("address")): normalize_candidate_fact(fact)
        for fact in compiled
    }
    missing = []
    wrong = []
    for fact in historical:
        key = (fact["sheet"], fact["address"])
        expected = {k: fact[k] for k in ("sheet", "address", "value", "formula") if k in fact}
        actual = compiled_map.get(key)
        if actual is None:
            missing.append(expected)
        elif actual != expected:
            wrong.append({"expected": expected, "actual": actual})
    if missing:
        classification = "MISSING_VISIBLE_FACTS"
    elif wrong:
        classification = "WRONG_VISIBLE_FACTS"
    else:
        classification = "EXACT_VISIBLE_FACT_MATCH"
    return classification, missing[:100], wrong[:100]


def event_id(operation: dict[str, Any]) -> str:
    return f"{operation['task_id']}:{operation['sequence_id']}:{operation['operation_index']}"


def request_signature(requests: list[dict[str, str]]) -> tuple[tuple[str, str], ...]:
    return tuple(sorted((item["sheet"], item["range"]) for item in requests))


def mechanical_candidate_requests(
    operation: dict[str, Any], path: Path
) -> tuple[list[dict[str, str]], str]:
    """Recover only finite ranges explicitly present in the old command.

    This is deliberately conservative. It handles bounded ``iter_rows`` and
    row-slice idioms plus literal cell references. It does not turn an
    unbounded scan, a later output, or a visible address set into a query.
    """
    if operation["requests"]:
        return operation["requests"], "historical_request_reconstruction"
    command = str(operation["args"].get("command", ""))
    bindings = _sheet_variables(command)
    if not bindings:
        return [], "no_literal_sheet_binding"
    from openpyxl.utils import get_column_letter

    requests: list[dict[str, str]] = []
    for variable, sheets in bindings.items():
        for match in re.finditer(
            rf"{re.escape(variable)}\.iter_rows\(([^)]*)\)", command
        ):
            kwargs = {
                key: int(value)
                for key, value in re.findall(
                    r"(min_row|max_row|min_col|max_col)\s*=\s*(\d+)",
                    match.group(1),
                )
            }
            if not {"min_row", "max_row"}.issubset(kwargs):
                continue
            min_col = kwargs.get("min_col", 1)
            max_col = kwargs.get("max_col")
            if max_col is None:
                continue
            cell_range = (
                f"{get_column_letter(min_col)}{kwargs['min_row']}:"
                f"{get_column_letter(max_col)}{kwargs['max_row']}"
            )
            for sheet in sheets:
                requests.append({"sheet": sheet, "range": cell_range})

        loop_rows = {
            name: (int(start), int(end) - 1)
            for name, start, end in re.findall(
                r"for\s+(\w+)\s+in\s+range\(\s*(\d+)\s*,\s*(\d+)\s*\)",
                command,
            )
        }
        for row_var, (start, end) in loop_rows.items():
            for match in re.finditer(
                rf"{re.escape(variable)}\[\s*{re.escape(row_var)}\s*\]\s*\[:\s*(\d+)\s*\]",
                command,
            ):
                max_col = int(match.group(1))
                cell_range = f"A{start}:{get_column_letter(max_col)}{end}"
                for sheet in sheets:
                    requests.append({"sheet": sheet, "range": cell_range})

        for match in re.finditer(
            rf"{re.escape(variable)}\[\s*['\"]([A-Z]{{1,3}}\d+)['\"]\s*\]",
            command,
        ):
            for sheet in sheets:
                address = match.group(1)
                requests.append({"sheet": sheet, "range": f"{address}:{address}"})

    unique = {(item["sheet"], item["range"]): item for item in requests}
    if not unique:
        return [], "no_finite_explicit_range_recovered"
    return list(unique.values()), "mechanically_recovered_explicit_command_ranges"


def extract_rows_for_cycles(operations: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for op in operations:
        grouped[op["sequence_id"]].append(op)
    output = {}
    for sequence, ops in grouped.items():
        requests = [request for op in ops for request in op["requests"]]
        output[sequence] = {
            "task_id": ops[0]["task_id"],
            "sequence_id": sequence,
            "historical_model_tool_cycles": len({op["turn"] for op in ops}),
            "historical_tool_executions": len(ops),
            "historical_python_executions": sum(op["tool"] == "bash" for op in ops),
            "historical_view_xlsx_executions": sum(op["tool"] == "view_xlsx" for op in ops),
            "historical_ranges_internally_printed": len(requests),
            "minimum_compiled_calls_one_per_historical_execution": len(ops),
            "cycle_reduction_classification": "NO_CYCLE_REDUCTION",
            "cycle_reduction_reason": "range consolidation inside one execution is not a model/tool-cycle saving",
        }
    return output


def consolidation_rows(operations: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = []
    by_sequence: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for op in operations:
        by_sequence[op["sequence_id"]].append(op)
    for sequence, ops in by_sequence.items():
        seen: dict[tuple[tuple[str, str], ...], dict[str, Any]] = {}
        for op in ops:
            sig = request_signature(op["requests"])
            if not sig:
                continue
            if sig in seen:
                rows.append({
                    "task_id": op["task_id"],
                    "sequence_id": sequence,
                    "event_id": event_id(op),
                    "prior_event_id": event_id(seen[sig]),
                    "kind": "EXACT_REPEATED_EXPLICIT_QUERY",
                    "mechanically_predictable": True,
                    "cycles_avoidable_ceiling": 1,
                })
            else:
                seen[sig] = op
    return rows


def median_or_none(values: list[float]) -> float | None:
    return round(statistics.median(values), 2) if values else None


def task_summary(rows: list[dict[str, Any]]) -> dict[str, Any]:
    eligible = [row for row in rows if row["reconstruction_status"] == "ELIGIBLE"]
    realizable = [
        row for row in eligible
        if row["realizability"] in {"DIRECTLY_REALIZABLE", "REALIZABLE_WITH_MECHANICAL_PROJECTION"}
    ]
    matched = [row for row in realizable if row["fact_classification"] in {
        "EXACT_VISIBLE_FACT_MATCH", "NORMALIZED_VISIBLE_FACT_MATCH"
    }]
    comparable = [row for row in eligible if row.get("raw_compiled_fact_classification")]
    raw_reductions = [row["raw_payload_reduction_pct"] for row in matched]
    normalized_reductions = [row["normalized_payload_reduction_pct"] for row in matched]
    facts_total = sum(row["historical_fact_count"] for row in eligible)
    comparable_facts_total = sum(row["historical_fact_count"] for row in comparable)
    comparable_facts_matched = sum(row["matched_fact_count"] for row in comparable)
    return {
        "task_id": rows[0]["task_id"] if rows else None,
        "eligible_observations": len(eligible),
        "realizable_observations": len(realizable),
        "realizability_rate": round(len(realizable) / len(eligible), 4) if eligible else None,
        "matched_observations": len(matched),
        # Fidelity and realizability are separate gates. This rate asks
        # whether the compiled candidate matches H whenever a candidate query
        # was actually reconstructed; the separate realizability rate asks
        # whether that query was available without hindsight.
        "eligible_visible_fact_count": facts_total,
        "comparable_visible_fact_count": comparable_facts_total,
        "comparable_matched_fact_count": comparable_facts_matched,
        "query_unavailable_fact_count": facts_total - comparable_facts_total,
        "fact_match_rate": round(comparable_facts_matched / comparable_facts_total, 4) if comparable_facts_total else None,
        "fact_match_rate_over_all_eligible": round(comparable_facts_matched / facts_total, 4) if facts_total else None,
        "historical_raw_bytes": sum(row["historical_raw_bytes"] for row in matched),
        "historical_normalized_fact_bytes": sum(row["historical_normalized_fact_bytes"] for row in matched),
        "compact_exact_h_bytes": sum(row["compact_exact_h_bytes"] for row in matched),
        "raw_payload_reduction_pct_median": median_or_none(raw_reductions),
        "normalized_payload_reduction_pct_median": median_or_none(normalized_reductions),
        "raw_payload_reduction_range_pct": [min(raw_reductions), max(raw_reductions)] if raw_reductions else None,
        "normalized_payload_reduction_range_pct": [min(normalized_reductions), max(normalized_reductions)] if normalized_reductions else None,
        "formatting_overhead_bytes": sum(row["historical_formatting_overhead_bytes"] for row in matched),
        "compiled_overfetch_facts": sum(row["compiled_overfetch_fact_count"] for row in matched),
        "compiled_full_minus_projected_bytes": sum(row.get("compiled_full_minus_projected_bytes", 0) for row in matched),
        "compiled_overfetch_bytes_upper_bound": sum(row.get("compiled_overfetch_bytes_upper_bound", 0) for row in matched),
        "eligible_class_counts": dict(sorted(defaultdict(int, {
            row["historical_class"]: sum(1 for item in eligible if item["historical_class"] == row["historical_class"])
            for row in eligible
        }).items())),
    }


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    population = {
        "selection_rule": "Frozen exact six-task population from the prior v2 adjudication; no reselection.",
        "tasks": TASKS,
        "source_population": str(POPULATION.relative_to(PROJECT_ROOT)),
        "source_ranking": str(RANKING.relative_to(PROJECT_ROOT)),
    }
    write_json("population.json", population)

    all_operations: list[dict[str, Any]] = []
    historical_rows = []
    for task in TASKS:
        _, operations = historical_records(task)
        all_operations.extend(operations)
        for op in operations:
            historical_rows.append({
                "event_id": event_id(op),
                "task_id": task,
                "trajectory_id": op["trajectory_id"],
                "sequence_id": op["sequence_id"],
                "operation_index": op["operation_index"],
                "turn": op["turn"],
                "tool": op["tool"],
                "historical_class_candidate": "CAPPED_VISIBLE" if op["historical_truncated"] else "FULL_VISIBLE",
                "historical_raw_bytes": op["historical_observation_bytes"],
                "historical_observation_tokens_estimate": op["historical_observation_tokens_estimate"],
                "historical_truncated": op["historical_truncated"],
                "observation_sha256": sha256_bytes(op["observation"]),
                "explicit_requests": op["requests"],
            })
    write_jsonl("historical_observations.jsonl", historical_rows)

    visible_rows = []
    query_rows = []
    equivalence_rows = []
    size_rows = []
    task_events: dict[str, list[dict[str, Any]]] = defaultdict(list)
    freshness_path = OUT / "freshness.jsonl"
    if freshness_path.exists():
        freshness_path.unlink()
    import os

    os.environ["AB_FRESHNESS_LOG"] = str(freshness_path)
    indexed_task = None
    for op in all_operations:
        path = WORK / op["task_id"].replace(":", "_") / "input.xlsx"
        facts, empty_slots, historical_class, parse_mode = event_facts(op)
        candidate_requests, candidate_request_source = mechanical_candidate_requests(op, path)
        reconstruction_status = "ELIGIBLE" if facts else "UNRESOLVED"
        realizability, realizability_reason = query_realizability(
            op, facts, historical_class, path
        )
        if (
            facts
            and not candidate_requests
            and realizability in {"DIRECTLY_REALIZABLE", "REALIZABLE_WITH_MECHANICAL_PROJECTION"}
        ):
            realizability = "REQUIRES_HINDSIGHT"
            realizability_reason = "mechanical_filter_seen_but_no_finite_compiled_query_form_recovered"
        fact_row = {
            "event_id": event_id(op),
            "task_id": op["task_id"],
            "sequence_id": op["sequence_id"],
            "turn": op["turn"],
            "tool": op["tool"],
            "historical_class": historical_class,
            "parse_mode": parse_mode,
            "reconstruction_status": reconstruction_status,
            "historical_fact_count": len(facts),
            "historical_empty_slots_visible": empty_slots,
            "historical_raw_bytes": op["historical_observation_bytes"],
            "historical_truncated": op["historical_truncated"],
            "explicit_requests": op["requests"],
            "compiled_candidate_requests": candidate_requests,
            "compiled_candidate_request_source": candidate_request_source,
        }
        visible_rows.append({**fact_row, "facts": facts})
        query_rows.append({
            "event_id": event_id(op),
            "task_id": op["task_id"],
            "tool": op["tool"],
            "turn": op["turn"],
            "explicit_requests": op["requests"],
            "compiled_candidate_requests": candidate_requests,
            "compiled_candidate_request_source": candidate_request_source,
            "query_command_sha256": sha256_bytes(str(op["args"].get("command", ""))),
            "query_command_excerpt": str(op["args"].get("command", ""))[:500],
            "available_before_observation": True,
        })
        result = {
            **fact_row,
            "realizability": realizability,
            "realizability_reason": realizability_reason,
        }
        fact_row["realizability"] = realizability
        fact_row["realizability_reason"] = realizability_reason
        if facts and candidate_requests:
            if indexed_task != op["task_id"]:
                index_module.reset()
                indexed_task = op["task_id"]
            response, compiled_facts = replay_v2(path, candidate_requests)
            classification, missing, wrong = compare_visible_facts(facts, compiled_facts)
            if realizability not in {"DIRECTLY_REALIZABLE", "REALIZABLE_WITH_MECHANICAL_PROJECTION"}:
                primary_classification = "NOT_QUERY_REALIZABLE"
            else:
                primary_classification = classification
            compiled_map = {
                (fact.get("sheet"), fact.get("address")): normalize_candidate_fact(fact)
                for fact in compiled_facts
            }
            projected = []
            for fact in facts:
                actual = compiled_map.get((fact["sheet"], fact["address"]))
                if actual:
                    projected.append(projected_fact(actual, fact["visible_fields"]))
            historical_norm = canonical_fact_bytes(facts)
            verbose = verbose_bytes(facts)
            compact_exact = compact_exact_h_bytes(projected)
            full_compact = len(compact(response).encode("utf-8"))
            matched_count = len(facts) - len(missing) - len(wrong)
            fact_row.update({
                "fact_classification": primary_classification,
                "raw_compiled_fact_classification": classification,
                "missing_visible_facts": missing,
                "wrong_visible_facts": wrong,
                "matched_fact_count": matched_count,
                "compiled_fact_count": len(compiled_facts),
                "compiled_overfetch_fact_count": max(0, len(compiled_facts) - len(facts)),
                "historical_normalized_fact_bytes": historical_norm,
                "verbose_exact_h_bytes": verbose,
                "compact_exact_h_bytes": compact_exact,
                "compiled_full_bytes": full_compact,
                "historical_formatting_overhead_bytes": fact_row["historical_raw_bytes"] - historical_norm,
                "historical_raw_minus_normalized_fact_bytes": fact_row["historical_raw_bytes"] - historical_norm,
                "compiled_full_minus_projected_bytes": full_compact - compact_exact,
                "compiled_overfetch_bytes_upper_bound": max(0, full_compact - compact_exact),
                "raw_payload_reduction_pct": round((1 - compact_exact / fact_row["historical_raw_bytes"]) * 100, 2) if fact_row["historical_raw_bytes"] else None,
                "normalized_payload_reduction_pct": round((1 - compact_exact / historical_norm) * 100, 2) if historical_norm else None,
                "verbose_to_compact_reduction_pct": round((1 - compact_exact / verbose) * 100, 2) if verbose else None,
                "compiled_full_to_projected_reduction_pct": round((1 - compact_exact / full_compact) * 100, 2) if full_compact else None,
                "compiled_index_generation": response.get("index_generation"),
                "compiled_truncated": bool(response.get("truncated")) or any(
                    item.get("truncated", False)
                    for item in response.get("results", [])
                    if isinstance(item, dict)
                ),
            })
            equivalence_rows.append({
                "event_id": event_id(op),
                "task_id": op["task_id"],
                "realizability": realizability,
                "compiled_candidate_request_source": candidate_request_source,
                "compiled_candidate_requests": candidate_requests,
                "classification": primary_classification,
                "raw_compiled_classification": classification,
                "historical_fact_count": len(facts),
                "compiled_fact_count": len(compiled_facts),
                "matched_fact_count": matched_count,
                "missing_visible_facts": missing,
                "wrong_visible_facts": wrong,
                "compiled_truncated": fact_row["compiled_truncated"],
            })
            size_rows.append({
                "event_id": event_id(op),
                "task_id": op["task_id"],
                "realizability": realizability,
                "fact_classification": primary_classification,
                "facts": len(facts),
                "historical_raw_bytes": fact_row["historical_raw_bytes"],
                "historical_normalized_fact_bytes": historical_norm,
                "verbose_structured_bytes": verbose,
                "compact_exact_h_bytes": compact_exact,
                "compiled_full_bytes": full_compact,
                "raw_payload_reduction_pct": fact_row["raw_payload_reduction_pct"],
                "normalized_payload_reduction_pct": fact_row["normalized_payload_reduction_pct"],
                "verbose_to_compact_reduction_pct": fact_row["verbose_to_compact_reduction_pct"],
                "formatting_overhead_bytes": fact_row["historical_formatting_overhead_bytes"],
                "overfetch_fact_count": fact_row["compiled_overfetch_fact_count"],
                "compiled_full_minus_projected_bytes": fact_row["compiled_full_minus_projected_bytes"],
                "compiled_overfetch_bytes_upper_bound": fact_row["compiled_overfetch_bytes_upper_bound"],
            })
        else:
            fact_row.update({
                "fact_classification": "UNRESOLVED",
                "matched_fact_count": 0,
                "compiled_truncated": False,
            })
            if facts:
                equivalence_rows.append({
                    "event_id": event_id(op),
                    "task_id": op["task_id"],
                    "realizability": realizability,
                    "classification": "NOT_QUERY_REALIZABLE",
                    "raw_compiled_classification": None,
                    "historical_fact_count": len(facts),
                    "compiled_fact_count": None,
                    "matched_fact_count": 0,
                    "missing_visible_facts": [],
                    "wrong_visible_facts": [],
                    "compiled_truncated": False,
                    "compiled_candidate_request_source": candidate_request_source,
                })
        task_events[op["task_id"]].append(fact_row)

    write_jsonl("visible_fact_sets.jsonl", visible_rows)
    write_jsonl("query_inputs.jsonl", query_rows)
    write_jsonl("query_realizability.jsonl", [
        {
            "event_id": row["event_id"],
            "task_id": row["task_id"],
            "historical_class": row["historical_class"],
            "reconstruction_status": row["reconstruction_status"],
            "realizability": row.get("realizability", "UNRESOLVED"),
            "reason": row.get("realizability_reason", "no_visible_cell_facts"),
            "explicit_requests": row["explicit_requests"],
            "compiled_candidate_requests": row.get("compiled_candidate_requests", []),
            "compiled_candidate_request_source": row.get("compiled_candidate_request_source"),
        }
        for rows in task_events.values() for row in rows
    ])
    write_jsonl("fact_equivalence.jsonl", equivalence_rows)
    write_jsonl("representation_sizes.jsonl", size_rows)

    task_results = {
        task: task_summary(rows) for task, rows in task_events.items()
    }
    write_json("task_results.json", task_results)

    cycle_by_sequence = extract_rows_for_cycles(all_operations)
    write_jsonl("cycle_analysis.jsonl", list(cycle_by_sequence.values()))
    consolidation = consolidation_rows(all_operations)
    write_jsonl("consolidation_opportunities.jsonl", consolidation)

    matched_rows = [
        row for rows in task_events.values() for row in rows
        if row.get("fact_classification") in {"EXACT_VISIBLE_FACT_MATCH", "NORMALIZED_VISIBLE_FACT_MATCH"}
        and row.get("realizability") in {"DIRECTLY_REALIZABLE", "REALIZABLE_WITH_MECHANICAL_PROJECTION"}
    ]
    task_raw_ceiling = [task_results[task]["raw_payload_reduction_pct_median"] for task in TASKS if task_results[task]["raw_payload_reduction_pct_median"] is not None]
    task_norm_ceiling = [task_results[task]["normalized_payload_reduction_pct_median"] for task in TASKS if task_results[task]["normalized_payload_reduction_pct_median"] is not None]
    historical_raw = sum(row["historical_raw_bytes"] for row in matched_rows)
    compact_bytes = sum(row["compact_exact_h_bytes"] for row in matched_rows)
    opportunity = {
        "payload_ceiling": {
            "matched_realizable_observations": len(matched_rows),
            "historical_raw_bytes": historical_raw,
            "counterfactual_compact_exact_h_bytes": compact_bytes,
            "absolute_bytes_avoided": historical_raw - compact_bytes,
            "pooled_reduction_pct": round((1 - compact_bytes / historical_raw) * 100, 2) if historical_raw else None,
            "per_task_median_reduction_pct": median_or_none(task_raw_ceiling),
            "per_task_range_pct": [min(task_raw_ceiling), max(task_raw_ceiling)] if task_raw_ceiling else None,
            "normalized_historical_fact_bytes": sum(row["historical_normalized_fact_bytes"] for row in matched_rows),
            "normalized_fact_pooled_reduction_pct": round((1 - compact_bytes / sum(row["historical_normalized_fact_bytes"] for row in matched_rows)) * 100, 2) if matched_rows else None,
            "per_task_normalized_median_reduction_pct": median_or_none(task_norm_ceiling),
        },
        "round_trip_ceiling": {
            "historical_model_tool_cycles": sum(row["historical_model_tool_cycles"] for row in cycle_by_sequence.values()),
            "minimum_compiled_cycles": sum(row["minimum_compiled_calls_one_per_historical_execution"] for row in cycle_by_sequence.values()),
            "cycles_avoidable_by_within_execution_batching": 0,
            "cycles_avoidable_by_exact_repeated_query_reuse": len(consolidation),
            "note": "Range consolidation inside one Python/view execution is not a model/tool-cycle saving.",
        },
    }
    write_json("opportunity_ceiling.json", opportunity)

    freshness_rows = []
    if freshness_path.exists():
        freshness_rows = [json.loads(line) for line in freshness_path.read_text(encoding="utf-8").splitlines() if line.strip()]
    no_freshness_failure = all(
        not row.get("stale", False)
        and not row.get("freshness_failure", False)
        and not row.get("truncated", False)
        for row in freshness_rows
    )
    primary_wrong_rows = sum(
        row.get("fact_classification") == "WRONG_VISIBLE_FACTS"
        for rows in task_events.values() for row in rows
    )
    primary_wrong_fact_count = sum(
        len(row.get("wrong_visible_facts", []))
        for rows in task_events.values() for row in rows
    )
    no_freshness_or_fidelity_failure = no_freshness_failure and primary_wrong_rows == 0
    gate_rows = []
    for task in TASKS:
        result = task_results[task]
        gate_rows.append({
            "task_id": task,
            "fact_preservation_pass": (result["fact_match_rate"] or 0) >= 0.95,
            "query_realizability_pass": (result["realizability_rate"] or 0) >= 0.75,
            "payload_materiality_pass": (result["raw_payload_reduction_pct_median"] or -math.inf) >= 25,
            "fact_match_rate": result["fact_match_rate"],
            "fact_match_rate_over_all_eligible": result["fact_match_rate_over_all_eligible"],
            "eligible_visible_fact_count": result["eligible_visible_fact_count"],
            "comparable_visible_fact_count": result["comparable_visible_fact_count"],
            "query_unavailable_fact_count": result["query_unavailable_fact_count"],
            "realizability_rate": result["realizability_rate"],
            "raw_payload_reduction_pct_median": result["raw_payload_reduction_pct_median"],
        })
    gate = {
        "tasks": 6,
        "task_gate_rows": gate_rows,
        "tasks_fact_preservation_ge_95_pct": sum(row["fact_preservation_pass"] for row in gate_rows),
        "tasks_query_realizability_ge_75_pct": sum(row["query_realizability_pass"] for row in gate_rows),
        "tasks_payload_median_reduction_ge_25_pct": sum(row["payload_materiality_pass"] for row in gate_rows),
        "no_stale_or_fidelity_failure": no_freshness_or_fidelity_failure,
        "no_stale_or_silent_truncation": no_freshness_failure,
        "primary_wrong_fact_rows": primary_wrong_rows,
        "primary_wrong_fact_count": primary_wrong_fact_count,
        "freshness_log_rows": len(freshness_rows),
        "pass": (
            sum(row["fact_preservation_pass"] for row in gate_rows) >= 4
            and sum(row["query_realizability_pass"] for row in gate_rows) >= 4
            and sum(row["payload_materiality_pass"] for row in gate_rows) >= 4
            and no_freshness_or_fidelity_failure
        ),
        "failed_components": [
            component for component, condition in (
                ("FACT_PRESERVATION", sum(row["fact_preservation_pass"] for row in gate_rows) >= 4),
                ("QUERY_REALIZABILITY", sum(row["query_realizability_pass"] for row in gate_rows) >= 4),
                ("PAYLOAD_MATERIALITY", sum(row["payload_materiality_pass"] for row in gate_rows) >= 4),
                ("FRESHNESS_OR_FIDELITY", no_freshness_or_fidelity_failure),
            ) if not condition
        ],
        "verdict": "FACT_NORMALIZED_READ_EFFICIENCY_MECHANICALLY_SUPPORTED" if all([
            sum(row["fact_preservation_pass"] for row in gate_rows) >= 4,
            sum(row["query_realizability_pass"] for row in gate_rows) >= 4,
            sum(row["payload_materiality_pass"] for row in gate_rows) >= 4,
            no_freshness_or_fidelity_failure,
        ]) else "FACT_NORMALIZED_READ_EFFICIENCY_NOT_GENERAL",
        "rule": ">=4/6 tasks at >=95% facts, >=75% realizability, and median >=25% raw historical payload reduction; zero freshness/fidelity failures.",
    }
    write_json("gate.json", gate)
    write_json("corrected_evidence_ledger.json", {
        "CAPABILITY PRESERVATION": "SUPPORTED",
        "TRANSPARENT RUNTIME": "SUPPORTED",
        "V1 READ-SIDE MECHANISM": "EARNED",
        "V1 END-TO-END EFFICIENCY": "CONFOUNDED",
        "COMPACT ENCODING": "SUPPORTED_NARROWLY",
        "FACT-NORMALIZED PROJECTION": "SUPPORTED_NARROWLY" if gate["tasks_fact_preservation_ge_95_pct"] >= 4 else "NOT_ESTABLISHED",
        "QUERY REALIZABILITY": "SUPPORTED_NARROWLY" if gate["tasks_query_realizability_ge_75_pct"] >= 4 else "NOT_ESTABLISHED",
        "ROUND-TRIP CONSOLIDATION": "SUPPORTED_NARROWLY" if len(consolidation) else "NOT_ESTABLISHED",
        "V2 BEHAVIORAL SUBSTITUTION": "UNTESTED",
        "V2 END-TO-END EFFICIENCY": "UNTESTED",
        "FULL-BENCHMARK JUSTIFICATION": "NOT_ESTABLISHED",
    })
    write_json("next_experiment.json", {
        "live_v1_v2_justified": bool(gate["pass"]),
        "gate_verdict": gate["verdict"],
        "next_experiment": "Run the already specified identical-documentation V1/V2 discriminator only if this gate passes; otherwise do not run live inference.",
        "model_calls_in_this_replay": 0,
    })
    print(json.dumps(gate, indent=2))


if __name__ == "__main__":
    main()
