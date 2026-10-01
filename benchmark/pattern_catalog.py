#!/usr/bin/env python3
"""Load and validate the gold-blind pattern catalog. Extract case facts from a stored run."""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Any

import yaml
from openpyxl.utils import column_index_from_string, get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "benchmark"))

from inspect_decision import _decode
from paired_autopsy import (
    DATA_DIR,
    RUNS_DIR,
    WRITE_TOOLS,
    _datasets,
    _first_miss,
    _input_kind,
    _inspect_mentions,
    _json_blobs,
    _load_steps,
    _split_address,
    _write_covers,
    autopsy_trajectory,
)

from librecalc_mcp.domain.commit_checks import expand_range

CATALOG = ROOT / "benchmark" / "pattern-catalog"
_A1 = re.compile(r"^([A-Z]+)([0-9]+)$")


def catalog_dir() -> Path:
    return CATALOG


def load_surface() -> dict[str, Any]:
    return yaml.safe_load((CATALOG / "surface.yaml").read_text(encoding="utf-8"))


def load_cases() -> list[dict[str, Any]]:
    cases = []
    for path in sorted((CATALOG / "cases").glob("*.yaml")):
        payload = yaml.safe_load(path.read_text(encoding="utf-8"))
        payload["_path"] = str(path.relative_to(ROOT))
        cases.append(payload)
    return cases


def _formula_text(value: Any) -> str | None:
    if value is None:
        return None
    if type(value).__name__ == "ArrayFormula":
        return str(getattr(value, "text", value))
    if isinstance(value, str):
        return value
    return repr(value)


def _cell_snapshot(path: Path, sheet: str, cell: str) -> dict[str, Any]:
    import openpyxl

    workbook = openpyxl.load_workbook(path, data_only=False)
    try:
        if sheet not in workbook.sheetnames:
            return {"missing_sheet": True, "sheets": workbook.sheetnames[:16]}
        value = workbook[sheet][cell].value
        kind = "blank"
        if isinstance(value, str) and value.startswith("="):
            kind = "formula"
        elif type(value).__name__ == "ArrayFormula":
            kind = "array_formula"
        elif value is not None and value != "":
            kind = "constant"
        return {"kind": kind, "value": _formula_text(value)}
    finally:
        workbook.close()


def _neighborhood(path: Path, sheet: str, cell: str, span: int = 3) -> list[str]:
    import openpyxl

    match = _A1.match(cell.upper())
    if match is None:
        return []
    column = column_index_from_string(match.group(1))
    row = int(match.group(2))
    workbook = openpyxl.load_workbook(path, data_only=False, read_only=True)
    try:
        if sheet not in workbook.sheetnames:
            return [f"missing_sheet:{sheet}"]
        worksheet = workbook[sheet]
        lines = []
        for row_i in range(max(1, row - span), row + span + 1):
            parts = []
            for col_i in range(max(1, column - span), column + span + 1):
                address = f"{get_column_letter(col_i)}{row_i}"
                mark = "*" if address == cell.upper() else " "
                parts.append(f"{mark}{address}={_formula_text(worksheet[address].value)!s}"[:70])
            lines.append(" | ".join(parts))
        return lines
    finally:
        workbook.close()


def _inspect_selected(observation: str) -> list[dict[str, Any]]:
    selected: list[dict[str, Any]] = []
    for blob in _json_blobs(observation):
        if not isinstance(blob, dict):
            continue
        errors = blob.get("formula_errors") or {}
        for item in errors.get("selected_cells") or []:
            if isinstance(item, dict):
                selected.append(
                    {
                        "sheet": item.get("sheet"),
                        "address": item.get("address"),
                        "error": item.get("error") or item.get("kind"),
                        "formula": str(item.get("formula") or "")[:120],
                    }
                )
    return selected[:16]


def _boundary_hit(observation: str, sheet: str, cell: str) -> dict[str, Any] | None:
    for blob in _json_blobs(observation):
        if not isinstance(blob, dict):
            continue
        group = blob.get("boundary_continuations") or {}
        if not isinstance(group, dict):
            continue
        for item in group.get("candidates") or group.get("selected_candidates") or []:
            if (
                isinstance(item, dict)
                and item.get("sheet") == sheet
                and str(item.get("address") or "").upper() == cell
            ):
                return {
                    "inferred_formula": item.get("inferred_formula"),
                    "run": item.get("run") or f"{item.get('run_start')}:{item.get('run_end')}",
                    "note": group.get("note"),
                }
    return None


def _covering_writes(action: str, sheet: str, cell: str) -> list[dict[str, Any]]:
    import characterize_commit_checks as cc

    decoded = _decode(action)
    if decoded.split()[0] not in WRITE_TOOLS:
        return []
    hits = []
    for payload in cc._decoded_payloads(action):
        items = payload if isinstance(payload, list) else [payload]
        for item in items:
            if not isinstance(item, dict) or item.get("sheet") != sheet:
                continue
            cell_range = str(item.get("range") or item.get("cell_range") or "")
            covers = False
            if cell_range:
                try:
                    covers = ("_", cell) in {
                        ("_", addr) for _, addr in expand_range("_", cell_range)
                    }
                except ValueError:
                    covers = cell in cell_range.upper()
            if covers:
                hits.append(
                    {
                        "range": cell_range,
                        "formula": str(item.get("formula") or item.get("value") or "")[:160],
                    }
                )
    return hits


def extract_facts(run_dir: Path, task: str, address: str | None = None) -> dict[str, Any]:
    """Facts only. No pattern claims. No golden."""

    scores = json.loads((run_dir / "official_scores.json").read_text(encoding="utf-8"))["tasks"]
    score = scores[task]
    category, _, task_id = task.partition(":")
    kind, scored_address = _first_miss(score)
    address = address or scored_address
    if address is None:
        raise ValueError(f"no first-miss address for {task}")
    sheet, cell = _split_address(address)
    entry = _datasets()[(category, task_id)]
    input_path = DATA_DIR / category / entry["spreadsheet_path"]
    output_path = run_dir / "submission" / "outputs" / category / f"{task_id}_output.xlsx"
    steps = _load_steps(run_dir, task)
    view = autopsy_trajectory(steps, address)
    inspect_obs = ""
    tools: list[str] = []
    writes: list[dict[str, Any]] = []
    read_errors: list[str] = []
    thoughts: list[str] = []
    compare_summaries: list[str] = []
    for step in steps:
        action = (step.get("action") or "").strip()
        tool = action.split()[0] if action else ""
        tools.append(tool)
        observation = step.get("observation") if isinstance(step.get("observation"), str) else ""
        thought = " ".join(str(step.get("thought") or "").split())
        if thought:
            thoughts.append(thought[:400])
        if tool == "calc_inspect" and not inspect_obs:
            inspect_obs = observation
        failed_read = '"ok": false' in observation or '"ok":false' in observation
        if failed_read and ("exhausted" in observation or "covers" in observation):
            read_errors.append(observation[:180].replace("\n", " "))
        if _write_covers(action, sheet, cell):
            writes.extend(_covering_writes(action, sheet, cell))
        if tool == "calc_compare":
            blob = next(
                (
                    item
                    for item in _json_blobs(observation)
                    if isinstance(item, dict) and "summary" in item
                ),
                None,
            )
            if blob:
                compare_summaries.append(json.dumps(blob.get("summary"), default=str)[:400])
    inspect = _inspect_mentions(inspect_obs, sheet, cell)
    facts = {
        "run": run_dir.name,
        "task": task,
        "spreadsheet": Path(entry["spreadsheet_path"]).name,
        "instruction": entry["instruction"].replace("\n", " ").strip()[:800],
        "first_miss": {
            "kind": kind,
            "address": address,
            "error_message": (score.get("error_message") or "")[:240],
            "evaluator_leak": True,
        },
        "scores": {
            "mod": score.get("modification_accuracy"),
            "reg": score.get("regression_accuracy"),
            "exact": bool(score.get("accuracy")),
        },
        "input": {
            **_cell_snapshot(input_path, sheet, cell),
            "input_kind": _input_kind(category, task_id, sheet, cell),
            "neighborhood": _neighborhood(input_path, sheet, cell),
        },
        "output": _cell_snapshot(output_path, sheet, cell) if output_path.is_file() else None,
        "autopsy": {
            "named": view["named"],
            "named_sources": view.get("named_sources") or [],
            "named_rank": view.get("named_rank"),
            "read": view["read"],
            "read_attempted": view.get("read_attempted") or False,
            "written": view["written"],
            "drowned": view["drowned"],
            "bucket": view["bucket"],
            "selected_count": view["selected_count"],
            "error_cell_count": view["error_cell_count"],
            "compared": view["compared"],
            "after_compare": view["after_compare"],
        },
        "inspect": {
            "named": inspect["named"],
            "named_sources": inspect.get("named_sources") or [],
            "rank": inspect.get("rank"),
            "selected_count": inspect.get("selected_count"),
            "error_cell_count": inspect.get("error_cell_count"),
            "selected_cells": _inspect_selected(inspect_obs),
            "boundary": _boundary_hit(inspect_obs, sheet, cell),
        },
        "trajectory": {
            "n_steps": len(steps),
            "tools": tools,
            "writes_covering_miss": writes,
            "read_errors": read_errors[:6],
            "compare_summaries": compare_summaries,
            "thoughts": thoughts[:8],
        },
    }
    return facts


def _arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    extract = sub.add_parser("extract-facts", help="Print gold-blind facts for one task")
    extract.add_argument("run")
    extract.add_argument("task")
    extract.add_argument("--address")
    extract.add_argument("--runs-dir", type=Path, default=RUNS_DIR)
    sub.add_parser("validate", help="Load surface.yaml + cases and check ids")
    return parser.parse_args()


def validate() -> list[str]:
    errors: list[str] = []
    surface = load_surface()
    patterns = {item["id"]: item for item in surface.get("patterns") or []}
    levers = set(yaml.safe_load((CATALOG / "arms.yaml").read_text(encoding="utf-8"))["levers"])
    case_ids: set[str] = set()
    for case in load_cases():
        case_id = case.get("id")
        if not case_id:
            errors.append(f"{case.get('_path')}: missing id")
            continue
        if case_id in case_ids:
            errors.append(f"duplicate case id {case_id}")
        case_ids.add(case_id)
        for required in ("run", "task", "first_miss", "facts", "claims"):
            if required not in case:
                errors.append(f"{case_id}: missing {required}")
        for pattern_id in case.get("claims", {}).get("patterns") or []:
            if pattern_id not in patterns:
                errors.append(f"{case_id}: unknown pattern {pattern_id}")
        if case.get("facts", {}).get("opened_golden"):
            errors.append(f"{case_id}: opened_golden must be false")
        claims = case.get("claims") or {}
        lever = claims.get("lever")
        if lever is None:
            errors.append(f"{case_id}: claims.lever is unset; the gate cannot judge it")
        elif lever not in levers:
            errors.append(f"{case_id}: unknown lever {lever!r}, expected one of {sorted(levers)}")
        if not claims.get("mechanism"):
            errors.append(f"{case_id}: claims.mechanism is unset")
    for pattern in surface.get("patterns") or []:
        if pattern.get("status") != "established":
            continue
        missing = [cid for cid in pattern.get("establishing_cases") or [] if cid not in case_ids]
        if missing:
            errors.append(f"pattern {pattern['id']}: missing cases {missing}")
    return errors


def main() -> int:
    args = _arguments()
    if args.command == "extract-facts":
        facts = extract_facts(args.runs_dir / args.run, args.task, args.address)
        print(json.dumps(facts, indent=2, default=str))
        return 0
    errors = validate()
    if errors:
        print("\n".join(errors))
        return 1
    print(f"ok {len(load_cases())} cases, {len(load_surface().get('patterns') or [])} patterns")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
