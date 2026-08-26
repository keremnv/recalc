#!/usr/bin/env python3
"""Measure the shipped blank-dependency bridge observation against direct value targets.

This calls the product observation function itself, not a parallel implementation. For
each workbook every sheet is treated as selected. Candidate generation is gold-blind;
offline golden data scores the candidates afterward. Real scoped inspection may expose
fewer candidates.
"""

from __future__ import annotations

import json
import sys
from collections import Counter
from pathlib import Path

from openpyxl.utils.cell import get_column_letter, range_boundaries

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "benchmark-data/SpreadsheetBench-2/evaluation"))
sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(ROOT / "src"))

import evaluation as ev
import openpyxl
from target_classification import classify_cache_robust_targets
from xlsx_metadata_repair import install as _install_repair

from librecalc_mcp.domain.observation import _blank_dependency_bridges

_install_repair()


def _normalise(value):
    return getattr(value, "text", value)


def _sheet_payloads(input_values, input_raw) -> list[tuple[str, str, dict]]:
    payloads = []
    for raw_sheet in input_raw.worksheets:
        value_sheet = ev._find_sheet(input_values, raw_sheet.title)
        if value_sheet is None:
            continue
        used_range = raw_sheet.calculate_dimension()
        min_col, min_row, max_col, max_row = range_boundaries(used_range)
        values = []
        formulas = []
        for row in range(min_row, max_row + 1):
            values.append(
                [value_sheet.cell(row, column).value for column in range(min_col, max_col + 1)]
            )
            formulas.append(
                [
                    _normalise(raw_sheet.cell(row, column).value)
                    for column in range(min_col, max_col + 1)
                ]
            )
        normalised_range = (
            f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"
        )
        payloads.append(
            (raw_sheet.title, normalised_range, {"values": values, "formulas": formulas})
        )
    return payloads


def analyse(task: dict, data_dir: Path) -> Counter:
    source = data_dir / task["spreadsheet_path"]
    golden = data_dir / task["golden_response_path"]
    input_values = openpyxl.load_workbook(source, data_only=True)
    answer_values = openpyxl.load_workbook(golden, data_only=True)
    input_raw = openpyxl.load_workbook(source, data_only=False)
    answer_raw = openpyxl.load_workbook(golden, data_only=False)

    direct_blank_targets: set[tuple[str, str]] = set()
    for chunk in ev.parse_answer_position(task["answer_position"]):
        sheet_name, _, cell_range = chunk.rpartition("!")
        sheet_name = sheet_name.strip().strip("'")
        result = classify_cache_robust_targets(
            ev,
            input_values,
            answer_values,
            input_raw,
            answer_raw,
            sheet_name,
            cell_range,
        )
        raw_sheet = ev._find_sheet(input_raw, sheet_name)
        if raw_sheet is None:
            continue
        for address in result.modification:
            value = _normalise(raw_sheet[address].value)
            if value is None or (isinstance(value, str) and not value.strip()):
                direct_blank_targets.add((raw_sheet.title, address))

    bridge = _blank_dependency_bridges(
        _sheet_payloads(input_values, input_raw),
        selected_sheets=set(input_raw.sheetnames),
    )
    visible = {
        (candidate["sheet"], candidate["address"]) for candidate in bridge["selected_candidates"]
    }
    hits = visible & direct_blank_targets
    counts: Counter = Counter(
        {
            "blank targets": len(direct_blank_targets),
            "bridge candidates": bridge["candidate_count"],
            "visible candidates": len(visible),
            "hits": len(hits),
        }
    )
    for workbook in {input_values, answer_values, input_raw, answer_raw}:
        workbook.close()
    return counts


def main() -> int:
    category = sys.argv[1] if len(sys.argv) > 1 else "Financial_Model"
    limit = int(sys.argv[2]) if len(sys.argv) > 2 else 0
    data_dir = ROOT / "benchmark-data/SpreadsheetBench-2/data" / category
    tasks = json.loads((data_dir / "dataset.json").read_text())
    if limit:
        tasks = tasks[:limit]

    per_task: list[Counter] = []
    for task in tasks:
        try:
            counts = analyse(task, data_dir)
        except Exception as exc:  # noqa: BLE001
            print(f"{task['id']:8} ERROR {type(exc).__name__}: {str(exc)[:80]}", file=sys.stderr)
            continue
        per_task.append(counts)
        print(
            f"{task['id']:8} blanks={counts['blank targets']:5} "
            f"candidates={counts['bridge candidates']:4} visible={counts['visible candidates']:3} "
            f"hits={counts['hits']:3}"
        )

    total: Counter = Counter()
    for counts in per_task:
        total.update(counts)
    print(f"\n=== {category}: {len(per_task)} tasks, all sheets selected ===")
    print(f"direct blank value targets {total['blank targets']:8,}")
    print(f"bridge candidates          {total['bridge candidates']:8,}")
    print(f"visible capped candidates  {total['visible candidates']:8,}")
    print(f"visible target hits        {total['hits']:8,}")
    print(f"pooled visible precision   {total['hits'] / max(1, total['visible candidates']):8.1%}")
    print(f"pooled blank-target recall {total['hits'] / max(1, total['blank targets']):8.1%}")
    reached = sum(1 for counts in per_task if counts["hits"])
    print(f"tasks with >=1 visible hit {reached}/{len(per_task)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
