#!/usr/bin/env python3
"""Offline evaluator-side integration autopsy. No provider calls are permitted.

Runtime modules never import this file. Gold is used only to measure archived
artifacts. Every missing observation is explicit rather than counted as a loss.
"""
from __future__ import annotations

import argparse
import gc
from collections import Counter, defaultdict, namedtuple
import json
from pathlib import Path
import statistics
import sys
from types import MethodType

sys.path.insert(0, str(Path(__file__).resolve().parent))
import matched_compiled_treatment as m
import openpyxl
from fingerprint import relative_fingerprint

OUT = m.RUN_ROOT / "integration_autopsy"
FLOOR = OUT / "zero_edit_floor"
INERT_THRESHOLD = 0.001


def ratio(a, b):
    return a / b if b else None


def cellkey(t):
    if "row" in t:
        return (t["sheet"], t["row"], t["col"])
    r, c = openpyxl.utils.cell.coordinate_to_tuple(t["address"])
    return t["sheet"], r, c


def label(c):
    return f"{c[0]}!{m.closure.a1(c[1], c[2])}"


def cells(path):
    wb = openpyxl.load_workbook(path, data_only=False)
    try:
        result = {}
        for ws in wb:
            for c in ws._cells.values():
                v = c.value
                if hasattr(v, "text"):
                    v = v.text
                if v is None or isinstance(v, str) and not v.strip():
                    continue
                result[(ws.title, c.row, c.column)] = v
        return result
    finally:
        wb.close()


def changed(a, b):
    return {c for c in a.keys() | b.keys() if a.get(c) != b.get(c)}


def formula(v):
    return isinstance(v, str) and v.startswith("=")


def exact(a, b):
    return formula(a) and formula(b) and m.program_group.canonical(a) == m.program_group.canonical(b)


def fp(v, c):
    if not formula(v):
        return None
    p = relative_fingerprint(v, c[2], c[1], sheet=c[0])
    return p.get("fingerprint") if isinstance(p, dict) else str(p)


def prepare_floor():
    OUT.mkdir(parents=True, exist_ok=True)
    m.write_json(OUT / "predeclared_analysis.json", {
        "live_model_calls": 0, "inert_absolute_modification_gain_threshold": INERT_THRESHOLD,
        "zero_edit_floor": "original input, metadata-only tolerance, identical official refresh/scorer path",
        "diagnosis": "earliest supported loss; low authority recall (<0.5) precedes scheduling; semantic retention not inferred from mere field presence",
        "phase_b": "BLOCKED_UNTIL_ALL_REPAIR_GATES_PASS",
    })
    for row in m.task_rows():
        gc.collect()
        m.write_cells(m.task_source(row["task_key"]), FLOOR / row["task_key"].replace(":", "-") / "output.xlsx", [])
    return m.scorer(FLOOR, "integration-autopsy-zero-edit-floor", refresh=True)


def detailed_score(row, outputs):
    """Same cell population/comparator as official evaluation; retain exact cells."""
    sys.path.insert(0, str(m.ROOT / "benchmark-data/SpreadsheetBench-2/evaluation"))
    import evaluation as ev
    with_font_color = row["category"] == "Debugging" and "Color" in row["input_path"]
    with_formula = row["category"] == "Debugging" and "Embedded" in row["input_path"]
    loader = sparse_workbook if with_font_color else compact_workbook
    wi = loader(row["input_path"], data_only=True)
    wif = loader(row["input_path"], data_only=False)
    wg = loader(row["gold_path"], data_only=True)
    wgf = loader(row["gold_path"], data_only=False)
    population = []
    for rng in ev.parse_answer_position(row["answer_position"]):
        sn, cr = rng.split("!", 1) if "!" in rng else (wg.sheetnames[0], rng)
        sn, cr = sn.strip("'").strip(), cr.strip("'").strip()
        regs, mods = ev.classify_cells_by_modification(wif if with_formula else wi, wgf if with_formula else wg, sn, cr, with_font_color, with_formula, wb_input_formula=wif, wb_answer_formula=wgf)
        for kind, names in (("regression", regs), ("modification", mods)):
            population.extend((kind, sn, n) for n in names)
    result = {}
    for name, path in outputs.items():
        wo = loader(path, data_only=True)
        wof = loader(path, data_only=False)
        acc = {k: {"total": 0, "official_correct": 0, "value_correct": 0, "official_cells": [], "value_cells": []} for k in ("modification", "regression")}
        for kind, sn, n in population:
            a = acc[kind]; a["total"] += 1
            wa, wu = ev._find_sheet(wg, sn), ev._find_sheet(wo, sn)
            if wu is None:
                continue
            ca, co = wa[n], wu[n]
            value_ok = ev._compare_cells(ca, co, False, False)
            fallback = not with_formula and (ev._has_excel_error(ca) or ev._has_excel_error(co))
            official_ok = ev.compare_cell_formula(ev._find_sheet(wgf, sn)[n], ev._find_sheet(wof, sn)[n]) if fallback else ev._compare_cells(ev._find_sheet(wgf, sn)[n] if with_formula else ca, ev._find_sheet(wof, sn)[n] if with_formula else co, with_font_color, with_formula)
            for metric, ok in (("official", official_ok), ("value", value_ok)):
                if ok:
                    a[metric + "_correct"] += 1
                    a[metric + "_cells"].append(f"{sn}!{n}")
        for a in acc.values():
            for metric in ("official", "value"):
                a[metric + "_accuracy"] = ratio(a[metric + "_correct"], a["total"])
        result[name] = acc
        wo.close(); wof.close()
    for w in (wi, wif, wg, wgf):
        w.close()
    return result


def sparse_workbook(path, *, data_only):
    """Read-only diagnostics: avoid retaining millions of synthetic blank Cells.

    The official comparator indexes every address in answer ranges. openpyxl's
    default accessor stores a new Cell for each absent address in six workbooks.
    Returning the identical default Cell without storing it preserves value,
    formula and font comparisons and prevents evaluation from exhausting RAM.
    """
    from openpyxl.cell.cell import Cell
    wb = openpyxl.load_workbook(path, data_only=data_only)
    def get_cell(ws, row, column):
        c = ws._cells.get((row, column))
        return c if c is not None else Cell(ws, row=row, column=column)
    for ws in wb:
        ws._get_cell = MethodType(get_cell, ws)
    return wb


ValueCell = namedtuple("ValueCell", "value")


class ValueSheet(dict):
    def __missing__(self, key):
        return ValueCell(None)


class ValueWorkbook(dict):
    @property
    def sheetnames(self):
        return list(self)

    def close(self):
        pass


def compact_workbook(path, *, data_only):
    """Stream workbook values/formulas for official modes that ignore fonts."""
    wb = openpyxl.load_workbook(path, data_only=data_only, read_only=True)
    result = ValueWorkbook()
    try:
        for ws in wb:
            dest = ValueSheet()
            for cells_row in ws.iter_rows():
                for c in cells_row:
                    if c.value is not None:
                        dest[c.coordinate] = ValueCell(c.value)
            result[ws.title] = dest
    finally:
        wb.close()
    return result


def request_target(call):
    try:
        user = json.loads(call["request_body"]["messages"][-1]["content"])
    except (KeyError, ValueError, TypeError):
        return None
    return user.get("TARGET") or (user.get("SESSION_STATE") or {}).get("target")


def analyze_task(row):
    key = row["task_key"]; d = m.LIVE / key.replace(":", "-")
    r = m.read_json(d / "result.json"); state = m.read_json(d / "state.json")
    source, gold, output = cells(m.task_source(key)), cells(Path(row["gold_path"])), cells(d / "output.xlsx")
    g = changed(source, gold); gf = {c for c in g if formula(gold.get(c))}
    out = changed(source, output); written_gold = out & g
    correct_write = {c for c in out if exact(output.get(c), gold.get(c)) or not formula(gold.get(c)) and output.get(c) == gold.get(c)}
    plan = dict(r["edit_plan"]); plan["spine"] = m.spine_for(key)
    auth_meta, cids = m.authorised_cells(plan); auth = set(auth_meta)
    sched = r["schedule"]; props = sched.get("canonical_decisions", [])
    sessions = {tuple(p["seed"]): p for p in props}
    scheduled = set(sessions)
    calls = [m.read_json(p) | {"artifact": str(p)} for p in sorted((d / "calls").glob("*.json"))]
    attempts = [c for c in calls if c.get("failure_class") not in ("TASK_MODEL_CALL_LIMIT", "TASK_COST_LIMIT")]
    call_rows = []; proposal_rows = []; group_rows = []; unit_rows = []; conditional = []
    group_members = {tuple(c) for u in sched.get("groups", []) for gr in u["groups"] for c in gr["member_cells"]}
    writes_sent = {cellkey(e) for e in r["write_audit"].get("applied", [])}
    hard_accepted = {tuple(p["seed"]) for p in props if (p.get("validation") or {}).get("hard_verifier_result") == "HARD_ACCEPT"}
    parseable = set(); correct_proposal = set(); retrieval_targets = set()
    for call in calls:
        t = request_target(call); c = cellkey(t) if t else None
        if call["stage"] == "retrieval" and c:
            retrieval_targets.add(c)
        stage = {"task_ir": "Task IR", "edit_plan": "Edit Plan", "retrieval": "retrieval/query selection"}.get(call["stage"], "other")
        if call["stage"] == "synthesis":
            stage = "canonical synthesis" if c in {tuple(u["seed"]) for u in sched.get("groups", []) if any(tuple(u["seed"]) in [tuple(x) for x in gr["member_cells"]] for gr in u["groups"])} else "ungrouped-cell synthesis"
        call_rows.append({"task": key, "category": row["category"], "stage": stage, "runtime_stage": call["stage"], "index": call["call_index_within_task"], "provider_attempt": call in attempts, "failure": call.get("failure_class"), "target": label(c) if c else None, "artifact": call["artifact"]})
        parsed = call.get("parsed_response") or {}
        f = parsed.get("formula") if isinstance(parsed, dict) else None
        if not formula(f):
            continue
        if not c:
            proposal_rows.append({"task": key, "index": call["call_index_within_task"], "classification": "UNKNOWN", "reason": "target absent from persisted request", "formula": f}); continue
        parseable.add(c)
        ok = exact(f, gold.get(c)); match = fp(f, c) is not None and fp(f, c) == fp(gold.get(c), c)
        if ok: correct_proposal.add(c)
        p = sessions.get(c, {}); validation = p.get("validation") or {}
        if exact(output.get(c), f) and c in out:
            cls = "WRITTEN_AS_PROPOSED"
        elif c in out and c in group_members:
            cls = "SUPERSEDED_BY_GROUP"
        elif validation.get("hard_verifier_result") == "HARD_REJECT":
            cls = "REJECTED_HARD"
        elif c not in auth:
            cls = "DROPPED_AUTHORITY"
        elif c in writes_sent and c not in out:
            cls = "NEVER_ACTUATED_OTHER" if output.get(c) == source.get(c) == f else "DROPPED_WRITER"
        elif c not in sessions:
            cls = "DROPPED_STATE"
        elif parsed.get("status") != "PROPOSED":
            cls = "NEVER_ACTUATED_OTHER"
        elif c in hard_accepted:
            cls = "DROPPED_SCHEDULER"
        else:
            cls = "NEVER_ACTUATED_OTHER"
        proposal_rows.append({"task": key, "category": row["category"], "index": call["call_index_within_task"], "stage": call["stage"], "target": label(c), "formula": f, "gold_target": c in gf, "exact_formula": ok, "fingerprint_match": match, "classification": cls, "validation": validation.get("hard_verifier_result"), "output_formula": output.get(c), "semantic_noop": output.get(c) == source.get(c) == f})
    for c, p in sessions.items():
        if c not in gf: continue
        sess = p["session"]; ws = set(sess.get("working_set_ids", [])); needed = m.proposal_precedents(gold[c], c)
        world = m.World(m.db_for(key))
        sheet_ids = {s["name"]: s["sheet_id"].split(":")[-1] for s in world.sheets.values()}
        world.close()
        needed_ids = {f"cell:{sheet_ids[x[0]]}:r{x[1]}:c{x[2]}" for x in needed if x[0] in sheet_ids}
        existing = any(fp(v, x) == fp(gold[c], c) for x, v in source.items() if formula(v))
        conditional.append({"task": key, "target": label(c), "retrieval_complete": needed_ids <= ws, "completeness_definition": "all evaluator gold direct reference cells in working-set IDs; not semantic sufficiency", "required_reference_cells": len(needed_ids), "retrieved_reference_cells": len(needed_ids & ws), "existing": existing, "proposal": bool(p.get("formula")), "exact_formula": exact(p.get("formula"), gold[c]), "fingerprint_match": fp(p.get("formula"), c) is not None and fp(p.get("formula"), c) == fp(gold[c], c)})
    for u in sched.get("groups", []):
        seed = tuple(u["seed"]); members = {tuple(c) for c in u["execution_members"]}
        p = sessions.get(seed, {}); f = p.get("formula")
        needed = m.proposal_precedents(gold.get(seed), seed) & gf if seed in gf else set()
        unit_rows.append({"task": key, "seed": label(seed), "seed_correct": exact(f, gold.get(seed)), "c1_executed": True, "closure_size": len(members), "members": sorted(map(label, members)), "required_gold_members": len(needed), "required_outside_authority": len(needed-auth), "recovered_required_members": len(needed & members), "members_scheduled": len(members & scheduled), "members_solved": len(members & correct_proposal), "members_written": len(members & out), "authority_invariant": members <= auth, "unsolved_marked_assigned_same_operation": sorted(label(c) for c in members - scheduled - out if auth_meta.get(c, {}).get("operation_id") == auth_meta.get(seed, {}).get("operation_id"))})
        for gi, gr in enumerate(u["groups"]):
            members = {tuple(c) for c in gr["member_cells"]}
            translated = {c: m.program_group.translate(f, seed, c) for c in members}
            group_rows.append({"task": key, "group": f"{label(seed)}:{gi}", "members": sorted(map(label, members)), "canonical": gr["canonical_member"], "actual_origin": label(seed), "origin_is_member": seed in members, "canonical_proposal": f, "canonical_correct": exact(f, gold.get(seed)) if seed in members else False, "translations": {label(c): v for c, v in translated.items()}, "translations_written": sum(exact(output.get(c), v) and c in out for c, v in translated.items()), "translated_member_correctness": sum(exact(output.get(c), gold.get(c)) and c in out for c in members), "grouped_gold_cells": len(members & gf), "stochastic_decisions_avoided": len(members - scheduled), "members_resynthesized": sorted(label(c) for c in members & scheduled if c != seed)})
    missed = []
    for gr in (row.get("programgroup") or {}).get("groups", []):
        members = {tuple(c) for c in gr["member_cells"]}
        if members <= group_members: continue
        if not members <= auth: cause = "target not authorised"
        elif not members <= scheduled | group_members: cause = "member not scheduled"
        elif not members & parseable: cause = "canonical never proposed"
        else: cause = "scheduler skipped group"
        missed.append({"task": key, "members": sorted(map(label, members)), "cause": cause})
    authority_recall = ratio(len(auth & g), len(g))
    compiler = r["compiler"]
    if not compiler.get("obligations"): diagnosis = "F0_TASK_SPEC"
    elif plan["status"] == "SESSION_RESOURCE_LIMIT": diagnosis = "OPERATIONAL_FAILURE"
    elif plan["status"] not in ("VALID_PLAN", "EMPTY_EXPANSION") or not auth or (authority_recall is not None and authority_recall < .5): diagnosis = "F2_EDIT_PLAN_AUTHORITY"
    elif len(scheduled & gf) < len(auth & gf): diagnosis = "F3_TARGET_SCHEDULING"
    elif any(p["classification"] == "DROPPED_SCHEDULER" for p in proposal_rows): diagnosis = "F8_PROPOSAL_ACTUATION"
    elif gf - correct_proposal: diagnosis = "F5_SYNTHESIS"
    else: diagnosis = "NO_SUPPORTED_DIAGNOSIS"
    totals = {"task": key, "category": row["category"], "gold_edit_cells": len(g), "gold_formula_cells": len(gf), "gold_independent_programs": len({fp(gold[c], c) or label(c) for c in gf}), "output_semantic_edits": len(out), "output_edits_intersect_gold": len(written_gold), "output_edits_outside_gold": len(out-g), "gold_write_recall": ratio(len(written_gold), len(g)), "write_precision": ratio(len(written_gold), len(out)), "correct_submitted_writes": len(correct_write), "operation_count": len((plan.get("parsed") or {}).get("operations", [])), "authority_cells": len(auth), "authority_intersect_gold": len(auth & g), "authority_gold_recall": authority_recall, "authority_precision": ratio(len(auth & g), len(auth)), "formula_target_recall": ratio(len(auth & gf), len(gf)), "scheduled_targets": len(scheduled), "scheduled_gold_formula_targets": len(scheduled & gf), "parseable_proposals": len(proposal_rows), "correct_proposals": sum(p.get("exact_formula", False) for p in proposal_rows), "model_calls": len(attempts), "state_model_calls": r["state"]["model_call_count"], "budget_block_records": len(calls)-len(attempts), "diagnosis": diagnosis, "plan_status": plan["status"], "admissibility_blocks": int(plan["status"] == "ADMISSIBILITY_LIMIT"), "model_calls_before_first_useful_gold_write": len(attempts) if correct_write & g else None, "calls_per_submitted_write": ratio(len(attempts), len(out)), "calls_per_gold_write": ratio(len(attempts), len(written_gold)), "calls_per_correct_canonical": ratio(len(attempts), sum(x["canonical_correct"] for x in group_rows)), "first_write_note": "writer batch is after all task model calls; no useful write means undefined"}
    stage_sets = {
        "G0": g, "G1": gf, "P1": auth, "P2": auth & g, "P3": auth,
        "S0": scheduled, "S1": scheduled, "S2": retrieval_targets, "S3": {tuple(p["seed"]) for p in props if p.get("formula")}, "S4": parseable, "S5": correct_proposal,
        "E0": hard_accepted, "E2": {tuple(c) for u in sched.get("groups", []) for c in u["execution_members"]}, "E3": group_members,
        "E4": {tuple(u["seed"]) for u in sched.get("groups", []) if u["groups"]}, "E5": group_members,
        "W0": {tuple(c["cell"]) for c in sched.get("translated_formula_instances", [])}, "W1": writes_sent, "W2": writes_sent, "W3": out,
    }
    stages = {}; prior_cells = g
    for name, cs in stage_sets.items():
        # Branches are not a single nesting chain: report intersection survival.
        stages[name] = {"count_entering": len(prior_cells), "count_surviving": len(prior_cells & cs), "stage_count": len(cs), "survival_percent": 100 * len(prior_cells & cs) / len(prior_cells) if prior_cells else None, "gold_target_recall": ratio(len(cs & g), len(g)), "precision": ratio(len(cs & g), len(cs)), "cells": sorted(map(label, cs))}
        prior_cells = cs
    stages.update({"G2": {"stage_count": totals["gold_independent_programs"], "unit": "independent programs"}, "T0": {"stage_count": len(compiler.get("obligations", [])), "retention": None, "missing_event": "no independent clause annotation in frozen run"}, "T1": {"stage_count": sum(bool(o.get("locus")) + bool(o.get("subject")) + bool(o.get("scope")) for o in compiler.get("obligations", [])), "retention": None, "missing_event": "field presence cannot establish semantic requirement retention"}, "P0": {"stage_count": int(plan["status"] == "VALID_PLAN"), "status": plan["status"]}, "E1": {"stage_count": len(unit_rows), "unit": "ExecutionUnits"}, "R0": {"stage_count": int((m.LIVE / "submission/outputs" / row["category"] / f"{row['task']}_output.xlsx").exists()), "unit": "submitted workbook with scorer refresh"}})
    return {"totals": totals, "stages": stages, "calls": call_rows, "proposals": proposal_rows, "groups": group_rows, "units": unit_rows, "conditional": conditional, "missed_groups": missed, "raw_content_edits": {"gold": sorted(map(label,g)), "output": sorted(map(label,out)), "intersection": sorted(map(label,written_gold)), "outside": sorted(map(label,out-g))}}


def analyze():
    OUT.mkdir(parents=True, exist_ok=True)
    results = []
    for row in m.task_rows():
        gc.collect()
        path = OUT / "tasks" / (row["task_key"].replace(":", "-") + ".json")
        if path.exists():
            item = m.read_json(path)
        else:
            print("AUTOPSY", row["task_key"], flush=True)
            item = analyze_task(row)
            m.write_json(path, item)
        results.append(item)
    rows = [x["totals"] for x in results]
    m.write_csv(OUT / "per_task_funnel.csv", rows)
    for name, field in (("call_stage_counts", "calls"), ("proposal_actuation", "proposals"), ("programgroup_cashout", "groups"), ("executionunit_cashout", "units"), ("conditional_synthesis", "conditional"), ("missed_programgroups", "missed_groups")):
        m.write_csv(OUT / (name + ".csv"), [v for x in results for v in x[field]])
    summary = {}
    for category in ("ALL", "Template", "Financial_Model", "Debugging"):
        rs = [r for r in rows if category == "ALL" or r["category"] == category]
        sums = {k: sum(r[k] for r in rs) for k in ("gold_edit_cells", "gold_formula_cells", "output_semantic_edits", "output_edits_intersect_gold", "output_edits_outside_gold", "correct_submitted_writes", "model_calls", "state_model_calls", "parseable_proposals", "correct_proposals", "authority_cells", "authority_intersect_gold")}
        sums.update({"tasks": len(rs), "diagnosis_counts": dict(Counter(r["diagnosis"] for r in rs)), "gold_write_recall": ratio(sums["output_edits_intersect_gold"], sums["gold_edit_cells"]), "write_precision": ratio(sums["output_edits_intersect_gold"], sums["output_semantic_edits"]), "authority_gold_recall": ratio(sums["authority_intersect_gold"], sums["gold_edit_cells"]), "authority_precision": ratio(sums["authority_intersect_gold"], sums["authority_cells"]), "authority_recall_distribution": [r["authority_gold_recall"] for r in rs]})
        summary[category] = sums
    allcalls = [v for x in results for v in x["calls"]]
    payload = {"live_model_calls": 0, "summary": summary, "call_stage_totals": dict(Counter(x["stage"] for x in allcalls if x["provider_attempt"])), "proposal_dispositions": dict(Counter(p["classification"] for x in results for p in x["proposals"])), "tasks": results}
    m.write_json(OUT / "funnel.json", payload)
    print(json.dumps({k: v for k, v in payload.items() if k != "tasks"}, indent=2))


def score_details():
    results = []
    for row in m.task_rows():
        gc.collect()
        key = row["task_key"]; path = OUT / "scorer_cells" / (key.replace(":", "-") + ".json")
        if path.exists(): result = m.read_json(path)
        else:
            rel = Path("submission/outputs") / row["category"] / f"{row['task']}_output.xlsx"
            result = detailed_score(row, {"treatment": m.LIVE / rel, "zero_edit_floor": FLOOR / rel})
            m.write_json(path, result)
        tm, zm = result["treatment"]["modification"], result["zero_edit_floor"]["modification"]
        gain = tm["official_accuracy"] - zm["official_accuracy"] if tm["total"] else 0
        results.append({"task": key, "modification_gain": gain, "value_only_gain": tm["value_accuracy"] - zm["value_accuracy"] if tm["total"] else 0, "effectively_inert": abs(gain) <= INERT_THRESHOLD, "correct_scorer_cells_gained": len(set(tm["official_cells"]) - set(zm["official_cells"])), "correct_value_cells_gained": len(set(tm["value_cells"]) - set(zm["value_cells"])), "treatment_modification": tm["official_accuracy"], "zero_edit_floor": zm["official_accuracy"], "treatment_value_only": tm["value_accuracy"], "zero_edit_floor_value_only": zm["value_accuracy"]})
        print("SCORE_DETAILS", key, gain, flush=True)
    m.write_json(OUT / "zero_edit_floor_scores.json", {"threshold": INERT_THRESHOLD, "inert_count": sum(r["effectively_inert"] for r in results), "rows": results})
    official = {}
    for row in m.task_rows():
        detail = m.read_json(OUT / "scorer_cells" / (row["task_key"].replace(":", "-") + ".json"))["zero_edit_floor"]
        reg = round(detail["regression"]["official_accuracy"] or 0, 4)
        reg = 1.0 if reg >= .998 else reg
        mod = round(detail["modification"]["official_accuracy"] or 0, 4)
        official[row["task_key"]] = {"id": row["task"], "regression_accuracy": reg, "modification_accuracy": mod, "accuracy": float(reg == mod == 1)}
    m.write_json(FLOOR / "official_scores.json", {"evaluation_runtime": "metadata-tolerant-local-v1", "scorer_path": "official classify_cells_by_modification and cell comparators after official open_spreadsheet.py refresh", "tasks": official, "scored": len(official), "exact": sum(r["accuracy"] for r in official.values())})


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=["floor", "analyze", "score-details"])
    args = parser.parse_args()
    {"floor": prepare_floor, "analyze": analyze, "score-details": score_details}[args.command]()
