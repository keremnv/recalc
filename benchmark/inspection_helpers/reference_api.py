"""H0 reference backend for model-facing inspection helpers.

Ordinary/openpyxl implementation of the exact factual contract of
benchmark.inspection_helpers.api (periods/search/inspect/inspect_ranges).
Fresh direct workbook read per call; no persistent index, no SQLite, no
substrate contact. Final retained helper backend for both arms; the historical substrate
implementation remains available only for evidence reproduction.

Output shapes match api.py exactly except index_generation (always 0: no
index exists) and an added backend marker inside the log record only
(never in returned values).
"""
from __future__ import annotations

import hashlib
import json
import os
import time
from pathlib import Path
from typing import Any

from benchmark.inspection_helpers.api import (
    CAP,
    CELL_CAP,
    _compact,
    _page,
    _styles,
)
from benchmark.inspection_helpers.index import PERIOD_RES

_counter = 0


def _qid() -> int:
    global _counter
    _counter += 1
    return _counter


def _log(call: dict) -> None:
    # Telemetry must never break the helper contract: all writes fail-safe.
    call = {**call, "backend": "reference_openpyxl"}
    log = os.environ.get("AB_FRESHNESS_LOG")
    if log:
        try:
            with open(log, "a", encoding="utf-8") as fh:
                fh.write(json.dumps(call) + "\n")
        except OSError:
            pass
    tel = os.environ.get("REP_HELPER_TELEMETRY")
    if tel:
        try:
            with open(tel, "a", encoding="utf-8") as fh:
                fh.write(json.dumps(call) + "\n")
        except OSError:
            pass


def _timed(helper: str, fn, *args, **kwargs):
    t0 = time.perf_counter_ns()
    try:
        return fn(*args, **kwargs), None
    except Exception as exc:  # noqa: BLE001 - parity: same exceptions surface
        return None, exc
    finally:
        pass


def _emit_time(helper: str, workbook: str, started_ns: int) -> None:
    tel = os.environ.get("REP_HELPER_TELEMETRY")
    if not tel:
        return
    role = {"periods": "HELPER_PERIODS", "search": "HELPER_SEARCH",
            "inspect": "HELPER_INSPECT",
            "inspect_ranges": "HELPER_INSPECT"}.get(helper, "OTHER")
    try:
        with open(tel, "a", encoding="utf-8") as fh:
            fh.write(json.dumps({
                "event": "helper_backend", "helper": helper,
                "consumer_role": role, "backend": "reference_openpyxl",
                "workbook": str(Path(workbook).name),
                "duration_ns": time.perf_counter_ns() - started_ns,
            }) + "\n")
    except OSError:
        pass


def _file_hash(path: str) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def _iter_cells(workbook: str, sheet: str | None = None):
    """Yield (ws_title, coord, row, col, v, f, dtype) exactly like index.build_index."""
    import re as _re

    from openpyxl.reader.excel import load_workbook

    wb = load_workbook(workbook, data_only=False, read_only=True)
    try:
        if sheet and sheet not in wb.sheetnames:
            return
        sheets = [wb[sheet]] if sheet else list(wb.worksheets)
        for ws in sheets:
            for row in ws.iter_rows():
                for c in row:
                    v, f = c.value, None
                    if c.data_type == "f":
                        if isinstance(c.value, str):
                            f = c.value
                        elif hasattr(c.value, "text"):
                            f = c.value.text
                        else:
                            f = str(c.value)
                        v = None
                    if v is None and f is None:
                        continue
                    yield (ws.title, c.coordinate, c.row, c.column, v, f,
                           c.data_type)
    finally:
        wb.close()


def periods(workbook: str, sheet: str | None = None, *,
            limit: int = CAP, offset: int = 0) -> dict:
    started = time.perf_counter_ns()
    try:
        limit, offset = _page(limit, offset, CAP)
        rows = []
        for ws_title, coord, r, c, v, f, _ in _iter_cells(workbook, sheet):
            if not isinstance(v, str):
                continue
            for rx in PERIOD_RES:
                for m in rx.finditer(v):
                    rows.append({"label": m.group(0), "sheet": ws_title,
                                 "address": coord, "row": r, "col": c})
        rows.sort(key=lambda d: (d["sheet"], d["row"], d["col"]))
        # NOTE: api.py temporal table preserves insertion order within
        # (sheet,row,col) across regexes (rowid order); our loop appends in
        # identical nested order (cell-major, regex order), then stable-sort
        # keeps that relative order. Verified by differential test.
        raw = rows[offset:offset + limit + 1]
        truncated = len(raw) > limit
        out = raw[:limit]
        _log({"helper": "periods", "workbook": str(Path(workbook).name),
              "sheet": sheet, "workbook_generation": _file_hash(workbook)[:12],
              "index_generation": 0, "query_generation": _qid(),
              "rebuilt": False, "n_results": len(out),
              "truncated": truncated, "offset": offset, "limit": limit})
        return {"results": out, "truncated": truncated,
                "next_offset": offset + limit if truncated else None,
                "index_generation": 0}
    finally:
        _emit_time("periods", workbook, started)


def search(workbook: str, pattern: str, regex: bool = False,
           sheet: str | None = None, *, limit: int = CAP,
           offset: int = 0) -> dict:
    import re as _re

    started = time.perf_counter_ns()
    try:
        limit, offset = _page(limit, offset, CAP)
        out_all = []
        if not regex:
            want = pattern.lower()
            for ws_title, coord, r, c, v, f, _ in _iter_cells(workbook,
                                                              sheet):
                # api.py JOIN emits one row per matching anchor token:
                # replicate occurrence multiplicity exactly (verified:
                # cells with the token twice yield duplicate rows).
                n = 0
                for text in ((str(v) if v is not None else None), f):
                    if not text:
                        continue
                    for tok in _re.findall(r"[A-Za-z0-9%\$]+",
                                           text.lower()):
                        if tok == want:
                            n += 1
                for _ in range(n):
                    out_all.append(
                        {"sheet": ws_title, "address": coord,
                         "value": None if v is None else str(v),
                         "formula": f})
        else:
            rx = _re.compile(pattern)
            for ws_title, coord, r, c, v, f, _ in _iter_cells(workbook,
                                                              sheet):
                vs = str(v) if v is not None else None
                if (vs and rx.search(vs)) or (f and rx.search(f)):
                    out_all.append({"sheet": ws_title, "address": coord,
                                    "value": vs, "formula": f,
                                    "_r": r, "_c": c})
        # api.py ORDER BY sheet,row,col (name order, not worksheet order)
        if regex:
            out_all.sort(key=lambda d: (d["sheet"], d["_r"], d["_c"]))
            for d in out_all:
                d.pop("_r", None)
                d.pop("_c", None)
        else:
            keyed = []
            for d in out_all:
                keyed.append(d)
            # rebuild with coordinates for exact ordering
            by_addr = {}
            for ws_title, coord, r, c, v, f, _ in _iter_cells(workbook,
                                                              sheet):
                by_addr[(ws_title, coord)] = (r, c)
            out_all.sort(key=lambda d: (
                d["sheet"], *by_addr[(d["sheet"], d["address"])]))
        raw = out_all[offset:offset + limit + 1]
        truncated = len(raw) > limit
        out = raw[:limit]
        _log({"helper": "search", "workbook": str(Path(workbook).name),
              "pattern": pattern[:80], "regex": regex, "sheet": sheet,
              "workbook_generation": _file_hash(workbook)[:12],
              "index_generation": 0, "query_generation": _qid(),
              "rebuilt": False, "n_results": len(out),
              "truncated": truncated, "offset": offset, "limit": limit})
        return {"results": out, "truncated": truncated,
                "next_offset": offset + limit if truncated else None,
                "index_generation": 0}
    finally:
        _emit_time("search", workbook, started)


def _range_cells(workbook: str, sheet: str,
                 cell_range: str) -> list[dict[str, Any]]:
    from openpyxl.utils import range_boundaries

    minc, minr, maxc, maxr = range_boundaries(cell_range)
    cells = []
    for ws_title, coord, r, c, v, f, d in _iter_cells(workbook, sheet):
        if minr <= r <= maxr and minc <= c <= maxc:
            cells.append({"address": coord, "row": r, "col": c,
                          "value": None if v is None else str(v),
                          "formula": f, "dtype": d})
    cells.sort(key=lambda d: (d["row"], d["col"]))
    return cells


def inspect(workbook: str, sheet: str, cell_range: str,
            with_styles: bool = False, *, limit: int = CELL_CAP,
            offset: int = 0, compact: bool = False) -> dict:
    started = time.perf_counter_ns()
    try:
        limit, offset = _page(limit, offset, CELL_CAP)
        cells = _range_cells(workbook, sheet, cell_range)
        raw = cells[offset:offset + limit + 1]
        truncated = len(raw) > limit
        cells = raw[:limit]
        styles = _styles(workbook, sheet, cell_range) if with_styles else None
        _log({"helper": "inspect", "workbook": str(Path(workbook).name),
              "sheet": sheet, "range": cell_range,
              "with_styles": with_styles,
              "workbook_generation": _file_hash(workbook)[:12],
              "index_generation": 0, "query_generation": _qid(),
              "rebuilt": False, "n_results": len(cells),
              "truncated": truncated, "offset": offset, "limit": limit})
        return {"results": _compact(cells) if compact else cells,
                "styles": styles, "truncated": truncated,
                "next_offset": offset + limit if truncated else None,
                "index_generation": 0}
    finally:
        _emit_time("inspect", workbook, started)


def inspect_ranges(workbook: str, ranges: list[dict[str, str] | tuple],
                   *, with_styles: bool = False,
                   max_cells: int = CELL_CAP) -> dict:
    started = time.perf_counter_ns()
    try:
        if not isinstance(ranges, list) or not ranges:
            raise ValueError("ranges must be a non-empty list")
        max_cells, _ = _page(max_cells, 0, CELL_CAP)
        out: list[dict[str, Any]] = []
        remaining = max_cells
        for request in ranges:
            if isinstance(request, dict):
                sheet = request.get("sheet")
                cell_range = request.get("range")
            elif isinstance(request, (tuple, list)) and len(request) == 2:
                sheet, cell_range = request
            else:
                raise ValueError(
                    "each range must be {'sheet': ..., 'range': ...}")
            if not isinstance(sheet, str) or not isinstance(cell_range,
                                                             str):
                raise TypeError("each range requires string sheet and range")
            if remaining <= 0:
                out.append({"sheet": sheet, "range": cell_range,
                            **_compact([]), "truncated": True,
                            "next_offset": 0})
                continue
            cells = _range_cells(workbook, sheet, cell_range)
            raw = cells[:remaining + 1]
            truncated = len(raw) > remaining
            cells = raw[:remaining]
            item: dict[str, Any] = {
                "sheet": sheet, "range": cell_range,
                **_compact(cells),
                "truncated": truncated,
                "next_offset": len(cells) if truncated else None,
            }
            if with_styles:
                item["styles"] = _styles(workbook, sheet, cell_range)
            out.append(item)
            remaining -= len(cells)
        _log({"helper": "inspect_ranges",
              "workbook": str(Path(workbook).name),
              "ranges": [{"sheet": r["sheet"], "range": r["range"]}
                         for r in out],
              "with_styles": with_styles,
              "workbook_generation": _file_hash(workbook)[:12],
              "index_generation": 0, "query_generation": _qid(),
              "rebuilt": False, "n_ranges": len(out),
              "n_results": sum(len(r["cells"]) for r in out),
              "max_cells": max_cells})
        return {"results": out, "index_generation": 0,
                "max_cells": max_cells}
    finally:
        _emit_time("inspect_ranges", workbook, started)
