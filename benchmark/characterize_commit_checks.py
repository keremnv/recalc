#!/usr/bin/env python3
"""Replay the commit-time checks against every stored run, before spending on a model.

Each check is generated gold-blind from the input workbook, the produced output workbook, and
the agent's own declared write targets recovered from its trajectory. Goldens are opened only
to score a finding afterwards: a finding is a true positive when the golden agrees the cell
should not have been touched (or should not have been left blank), and a false positive when
the golden wanted exactly what the agent did.

This is the gate before the interface changes. A check that floods here dies here, for free.
"""

from __future__ import annotations

import argparse
import json
import sys
import urllib.parse
from collections import Counter
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "src"))

import openpyxl

from librecalc_mcp.domain.boundary_continuations import payload_from_xlsx
from librecalc_mcp.domain.commit_checks import (
    CommitFinding,
    broken_check_cells,
    declared_targets,
    expand_range,
    new_formula_errors,
    unextended_continuations,
    unrequested_writes,
)

DATA_DIR = ROOT / "benchmark-data/SpreadsheetBench-2/data"
RUNS_DIR = ROOT / "benchmark-data/SpreadsheetBench-2/benchmark-runs/openrouter"
WRITE_TOOLS = ("calc_write", "calc_fill_formulas", "calc_program")
DATA_TABLE_SENTINEL = "<excel-data-table>"


def _arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--runs-dir", type=Path, default=RUNS_DIR)
    parser.add_argument("--limit", type=int, default=0, help="Analyse at most N task outputs.")
    parser.add_argument("--run", action="append", help="Restrict to these run names.")
    parser.add_argument("--category", action="append", help="Restrict to these categories.")
    parser.add_argument("--json", type=Path, help="Write the per-finding detail here.")
    return parser.parse_args()


def _datasets() -> dict[tuple[str, str], dict[str, Any]]:
    entries: dict[tuple[str, str], dict[str, Any]] = {}
    for category_dir in sorted(DATA_DIR.iterdir()):
        dataset = category_dir / "dataset.json"
        if not dataset.is_file():
            continue
        for task in json.loads(dataset.read_text()):
            entries[(category_dir.name, task["id"])] = task
    return entries


def _decoded_payloads(action: str) -> list[Any]:
    """Recover the JSON arguments a tool call carried, which arrive percent-encoded."""
    payloads: list[Any] = []
    for token in action.split():
        if not token.startswith(("%5B", "%7B", "[", "{")):
            continue
        try:
            payloads.append(json.loads(urllib.parse.unquote(token)))
        except (ValueError, UnicodeDecodeError):
            continue
    return payloads


def _declared_from_trajectory(trajectory_path: Path) -> set[tuple[str, str]]:
    document = json.loads(trajectory_path.read_text())
    requests: list[dict[str, Any]] = []
    for step in document.get("trajectory", []):
        action = (step.get("action") or "").strip()
        if not action.startswith(WRITE_TOOLS):
            continue
        for payload in _decoded_payloads(action):
            if isinstance(payload, dict):
                payload = [payload]
            if not isinstance(payload, list):
                continue
            requests.extend(item for item in payload if isinstance(item, dict))
    return declared_targets(requests)


def _cell_map(path: Path) -> dict[tuple[str, str], tuple[Any, Any]]:
    values = openpyxl.load_workbook(path, data_only=True)
    raw = openpyxl.load_workbook(path, data_only=False)
    cells: dict[tuple[str, str], tuple[Any, Any]] = {}
    data_table_cells: set[tuple[str, str]] = set()
    try:
        for sheet_name in raw.sheetnames:
            raw_sheet = raw[sheet_name]
            value_sheet = values[sheet_name] if sheet_name in values.sheetnames else None
            for row in raw_sheet.iter_rows():
                for cell in row:
                    raw_value = cell.value
                    # openpyxl returns ArrayFormula / DataTableFormula objects rather than a
                    # string. An Excel-authored input carries the object where any
                    # LibreOffice-saved file carries the equivalent "=..." text, so comparing
                    # them naively reports every such cell as a rewrite.
                    formula_text = getattr(raw_value, "text", None)
                    if type(raw_value).__name__ == "DataTableFormula":
                        # Only the master cell carries the descriptor; the rest of `ref` holds
                        # bare cached results. Mark the whole region so a recalculated result
                        # is not mistaken for a write.
                        reference = getattr(raw_value, "ref", None)
                        if isinstance(reference, str):
                            data_table_cells |= expand_range(sheet_name, reference)
                        # An Excel what-if data table. openpyxl exposes it as a structural
                        # descriptor with no formula text, while any LibreOffice-saved copy
                        # carries the equivalent "=TABLE(...)" string. The two are the same
                        # cell and cannot be compared, so both sides normalise to a sentinel.
                        # Live commit checks never see this: both workbooks come from UNO.
                        formula = DATA_TABLE_SENTINEL
                        value = None
                    elif isinstance(formula_text, str):
                        formula = formula_text if formula_text.startswith("=") else f"={formula_text}"
                        value = None
                    elif isinstance(raw_value, str) and raw_value.upper().startswith("=TABLE("):
                        formula = DATA_TABLE_SENTINEL
                        value = None
                    elif isinstance(raw_value, str) and raw_value.startswith("="):
                        formula = raw_value
                        value = raw_value
                    else:
                        formula = None
                        value = raw_value
                    if formula is not None and value_sheet is not None:
                        value = value_sheet[cell.coordinate].value
                    if formula is None and (value is None or (isinstance(value, str) and not value.strip())):
                        continue
                    cells[(sheet_name, cell.coordinate)] = (value, formula)
    finally:
        values.close()
        raw.close()
    for key in data_table_cells:
        if key in cells:
            cells[key] = (None, DATA_TABLE_SENTINEL)
    return cells


def _score(
    findings: list[CommitFinding],
    before: dict[tuple[str, str], tuple[Any, Any]],
    golden: dict[tuple[str, str], tuple[Any, Any]],
) -> Counter:
    """Each check has its own notion of a correct finding; score them separately."""
    tally: Counter = Counter()
    for finding in findings:
        key = (finding.sheet, finding.address)
        golden_content = golden.get(key, (None, None))
        golden_value, _ = golden_content
        if finding.check == "unrequested_write":
            # Correct when the golden kept the input content the agent overwrote.
            correct = before.get(key, (None, None)) == golden_content
        elif finding.check == "unextended_continuation":
            # Correct when the golden puts something in the cell the agent left blank.
            correct = golden_value is not None
        elif finding.check == "broken_check_cell":
            # Correct when a right answer leaves the tie-out satisfied.
            from librecalc_mcp.domain.commit_checks import _is_passing_check

            correct = _is_passing_check(golden_value)
        elif finding.check == "new_formula_error":
            from librecalc_mcp.domain.commit_checks import error_kind

            golden_formula = golden_content[1]
            if golden_formula is not None and golden_value is None:
                # The golden carries a formula with no cached result. openpyxl cannot say what
                # it evaluates to, so "the golden has no error here" is vacuously true and
                # would score as a win. Offline replay cannot judge this cell.
                tally[f"{finding.check}:unscorable"] += 1
                continue
            correct = error_kind(golden_content) is None
        else:
            correct = False
        tally[f"{finding.check}:{'true_positive' if correct else 'false_positive'}"] += 1
    return tally


def main() -> int:
    args = _arguments()
    entries = _datasets()
    tally: Counter = Counter()
    detail: list[dict[str, Any]] = []
    analysed = 0

    run_dirs = sorted(d for d in args.runs_dir.iterdir() if d.is_dir())
    if args.run:
        wanted = set(args.run)
        run_dirs = [d for d in run_dirs if d.name in wanted]

    for run_dir in run_dirs:
        for task_dir in sorted(d for d in run_dir.iterdir() if d.is_dir()):
            output_path = task_dir / "output.xlsx"
            if not output_path.is_file() or "-" not in task_dir.name:
                continue
            category, _, task_id = task_dir.name.partition("-")
            if args.category and category not in set(args.category):
                continue
            task = entries.get((category, task_id))
            if task is None:
                continue
            trajectories = list(task_dir.glob("trajectory/*/*.traj"))
            if not trajectories:
                continue
            if args.limit and analysed >= args.limit:
                break

            try:
                declared = _declared_from_trajectory(trajectories[0])
                before = _cell_map(DATA_DIR / category / task["spreadsheet_path"])
                after = _cell_map(output_path)
                golden = _cell_map(DATA_DIR / category / task["golden_response_path"])
            except Exception as error:  # noqa: BLE001 - a stored artifact may be unreadable
                tally["unreadable"] += 1
                detail.append({"run": run_dir.name, "task": task_dir.name, "error": str(error)[:200]})
                continue

            continuation_payload = payload_from_xlsx(str(DATA_DIR / category / task["spreadsheet_path"]))
            candidates = (continuation_payload or {}).get("candidates", [])
            findings = (
                unrequested_writes(before, after, declared)
                + new_formula_errors(before, after)
                + broken_check_cells(before, after)
                + unextended_continuations(after, candidates)
            )
            analysed += 1
            tally["tasks"] += 1
            tally["declared_cells"] += len(declared)
            if findings:
                tally["tasks_with_a_finding"] += 1
            for finding in findings:
                tally[finding.check] += 1
            tally.update(_score(findings, before, golden))
            detail.append(
                {
                    "run": run_dir.name,
                    "task": task_dir.name,
                    "declared_cells": len(declared),
                    "findings": [finding.as_dict() for finding in findings[:40]],
                    "finding_count": len(findings),
                }
            )

    print(f"analysed task outputs: {tally['tasks']}")
    print(f"tasks with at least one finding: {tally['tasks_with_a_finding']}")
    for check in (
        "unrequested_write",
        "new_formula_error",
        "broken_check_cell",
        "unextended_continuation",
    ):
        total = tally[check]
        true_positive = tally[f"{check}:true_positive"]
        false_positive = tally[f"{check}:false_positive"]
        scored = true_positive + false_positive
        precision = f"{true_positive / scored:.1%}" if scored else "n/a"
        print(
            f"  {check}: {total} findings, precision {precision} "
            f"(tp {true_positive} / fp {false_positive} / "
            f"unscorable {tally[f'{check}:unscorable']})"
        )
    if tally["unreadable"]:
        print(f"unreadable artifacts skipped: {tally['unreadable']}")

    if args.json:
        args.json.write_text(json.dumps(detail, indent=2))
        print(f"detail written to {args.json}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
