"""Gold-blind TEMPORAL_SPINE_V2: atoms and axis coordinates.

Goldens are never used to construct facts. Discovery freezes which
mechanisms are enabled; this module only executes those mechanisms.
"""
from __future__ import annotations

import calendar
import datetime as dt
import re
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any

_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(_ROOT / "benchmark"))
sys.path.insert(0, str(_ROOT / "benchmark/sweagent/formula_index/lib"))
sys.path.insert(0, str(_ROOT / "src"))

from fingerprint import a1_address, column_number, formula_text
from formula_schema import parse_period, period_key
from xlsx_metadata_repair import install

install()
import openpyxl  # noqa: E402
from openpyxl.styles.numbers import is_date_format
from openpyxl.utils.datetime import from_excel

ALL_FEATURES = (
    "DATE_SERIAL_DECODING",
    "MULTIROW_HEADER_COMPOSITION",
    "MERGED_HEADER_PROPAGATION",
    "BLOCK_INHERITANCE",
    "FORMULA_DERIVED_DATE",
    "ROW_AXIS_SUPPORT",
    "FY_TOKEN_NORMALIZATION",
    "ACTUAL_FORECAST_MARKER",
)

_WS = re.compile(r"\s+")
_MARKER = re.compile(
    r"^(a|e|f|b|act|est|actual|estimate|estimated|forecast|budget)$",
    re.I,
)
_FY_LOOSE = re.compile(
    r"^(?:fy|fye|cy)\s*'?((?:19|20)\d{2}|\d{2})(?:\s*[-/]\s*((?:19|20)?\d{2}))?[a-z]?$",
    re.I,
)
_YEAR_RANGE = re.compile(r"^((?:19|20)\d{2})\s*[-/]\s*((?:19|20)?\d{2})[a-z]?$", re.I)
_EDATE = re.compile(
    r"^\s*(?:_xlfn\.)?(EDATE|EOMONTH)\(\s*(\$?[A-Za-z]{1,3}\$?[1-9][0-9]*)\s*,\s*(-?\d+)\s*\)\s*$",
    re.I,
)
_DATE_FN = re.compile(
    r"^\s*(?:_xlfn\.)?DATE\(\s*(\d{4})\s*,\s*(\d{1,2})\s*,\s*(\d{1,2})\s*\)\s*$",
    re.I,
)
_DATE_FORMULA_HINT = re.compile(r"(?:_xlfn\.)?(EDATE|EOMONTH|\bDATE)\s*\(", re.I)
_ADD_DAYS = re.compile(
    r"^\s*(\$?[A-Za-z]{1,3}\$?[1-9][0-9]*)\s*\+\s*(\d+)\s*$",
    re.I,
)
_A1 = re.compile(r"^\$?([A-Za-z]{1,3})\$?([1-9][0-9]*)$")
_MONTH_NAME = {name.lower(): i for i, name in enumerate(calendar.month_name) if name}
_MONTH_ABBR = {name.lower(): i for i, name in enumerate(calendar.month_abbr) if name}


def _expand_year(token: str) -> int | None:
    digits = re.sub(r"\D", "", token)
    if len(digits) == 4:
        year = int(digits)
        return year if 1990 <= year <= 2100 else None
    if len(digits) == 2:
        year = 2000 + int(digits)
        return year if 1990 <= year <= 2100 else None
    return None


def _parse_a1(addr: str) -> tuple[int, int] | None:
    match = _A1.match(addr.replace("$", ""))
    if not match:
        return None
    return column_number(match.group(1)), int(match.group(2))


def snapshot_cell(cell: Any, *, merge_origin: dict[tuple[int, int], tuple[int, int]]) -> dict[str, Any]:
    col, row = int(cell.column), int(cell.row)
    value = cell.value
    kind = "blank"
    formula = formula_text(value)
    number_format = str(getattr(cell, "number_format", "") or "")
    origin = merge_origin.get((col, row), (col, row))
    rec: dict[str, Any] = {
        "col": col,
        "row": row,
        "address": a1_address(col, row),
        "number_format": number_format,
        "is_date_format": bool(number_format and is_date_format(number_format)),
        "formula": formula,
        "merged": (col, row) in merge_origin,
        "merge_origin": {"col": origin[0], "row": origin[1]},
        "python_type": type(value).__name__,
        "parse_period_v1": parse_period(value) if not formula else None,
    }
    if formula:
        rec["kind"] = "formula"
        rec["raw"] = formula[:240]
    elif value is None or (isinstance(value, str) and not value.strip()):
        rec["kind"] = "blank"
        rec["raw"] = None
    elif hasattr(value, "year") and hasattr(value, "month"):
        rec["kind"] = "datetime"
        rec["raw"] = value.isoformat()
        rec["dt"] = {"year": int(value.year), "month": int(value.month), "day": int(getattr(value, "day", 1))}
    elif isinstance(value, bool):
        rec["kind"] = "bool"
        rec["raw"] = str(value)
    elif isinstance(value, (int, float)):
        rec["kind"] = "number"
        rec["raw"] = float(value)
    else:
        rec["kind"] = "text"
        rec["raw"] = str(value).strip()[:240]
    return rec


def _marker_from_text(text: str) -> str | None:
    token = _WS.sub(" ", text.strip()).strip()
    if not token or not _MARKER.match(token):
        return None
    return token.upper()[:1] if len(token) <= 2 else token.lower()


def _fy_loose(text: str) -> dict[str, int] | None:
    compact = _WS.sub(" ", text.strip())
    match = _FY_LOOSE.match(compact) or _FY_LOOSE.match(compact.replace(" ", ""))
    if not match:
        match = _YEAR_RANGE.match(compact.replace(" ", ""))
        if match:
            year = _expand_year(match.group(1))
            return {"year": year} if year else None
        return None
    year = _expand_year(match.group(1))
    return {"year": year} if year else None


def atoms_from_snapshot(
    snap: dict[str, Any],
    *,
    features: set[str],
    epoch: Any = None,
) -> list[dict[str, Any]]:
    """Gold-blind atomic temporal evidence from one cell snapshot."""
    atoms: list[dict[str, Any]] = []
    raw = snap.get("raw")
    formula = snap.get("formula")

    def emit(kind: str, components: dict[str, Any], source: str) -> None:
        atoms.append(
            {
                "source_cell": snap["address"],
                "row": snap["row"],
                "col": snap["col"],
                "raw_value": raw if not isinstance(raw, float) else round(raw, 6),
                "kind": kind,
                "components": {k: v for k, v in components.items() if v is not None},
                "confidence_source": source,
                "formula": formula,
            }
        )

    if snap.get("kind") == "datetime" and snap.get("dt"):
        emit(
            "DATE_FORMATTED_VALUE",
            {"year": snap["dt"]["year"], "month": snap["dt"]["month"], "day": snap["dt"].get("day")},
            "python_datetime",
        )
        return atoms

    v1 = snap.get("parse_period_v1")
    if v1:
        emit("V1_PARSE_PERIOD", dict(v1), "parse_period")
        if "year" in v1 and isinstance(raw, str) and re.search(r"(fy|fye|cy)", raw, re.I):
            emit("YEAR_LITERAL", {"year": v1["year"]}, "fy_token_v1")
        if isinstance(raw, str):
            marker = None
            if re.search(r"(19|20)\d{2}\s*[aefb]\b", raw, re.I) or re.search(r"[aefb]$", raw.strip(), re.I):
                tail = raw.strip()[-1].upper()
                if tail in {"A", "E", "F", "B"} and "ACTUAL_FORECAST_MARKER" in features:
                    marker = tail
            if marker:
                emit("ACTUAL_MARKER" if marker == "A" else "FORECAST_MARKER" if marker in {"E", "F"} else "BUDGET_MARKER", {"marker": marker}, "lexical_suffix")
        return atoms

    if formula and "FORMULA_DERIVED_DATE" in features:
        return atoms

    if snap.get("kind") == "number" and "DATE_SERIAL_DECODING" in features:
        if snap.get("is_date_format") and isinstance(raw, (int, float)) and 1 <= float(raw) <= 80000:
            try:
                parsed = from_excel(float(raw), epoch=epoch) if epoch is not None else from_excel(float(raw))
            except Exception:
                parsed = None
            if parsed is not None and hasattr(parsed, "year"):
                emit(
                    "DATE_SERIAL",
                    {"year": int(parsed.year), "month": int(parsed.month), "day": int(getattr(parsed, "day", 1))},
                    "excel_serial+date_format",
                )
                return atoms

    if snap.get("kind") == "text" and isinstance(raw, str):
        if "FY_TOKEN_NORMALIZATION" in features:
            fy = _fy_loose(raw)
            if fy:
                emit("YEAR_LITERAL", fy, "fy_loose")
                return atoms
        marker = _marker_from_text(raw) if "ACTUAL_FORECAST_MARKER" in features else None
        if marker:
            kind = "ACTUAL_MARKER" if marker in {"A", "actual", "act"} or str(marker).lower().startswith("act") else (
                "BUDGET_MARKER" if str(marker).lower().startswith("b") else "FORECAST_MARKER"
            )
            emit(kind, {"marker": marker}, "marker_cell")
            return atoms
        lower = raw.strip().lower()
        if lower in _MONTH_NAME or lower in _MONTH_ABBR:
            month = _MONTH_NAME.get(lower) or _MONTH_ABBR.get(lower)
            emit("MONTH_LITERAL", {"month": month}, "month_name")
            return atoms
    return atoms


def _merge_map(sheet: Any) -> tuple[dict[tuple[int, int], tuple[int, int]], list[tuple[int, int, int, int]]]:
    origin: dict[tuple[int, int], tuple[int, int]] = {}
    spans: list[tuple[int, int, int, int]] = []
    for rng in sheet.merged_cells.ranges:
        span = (int(rng.min_col), int(rng.min_row), int(rng.max_col), int(rng.max_row))
        spans.append(span)
        for col in range(span[0], span[2] + 1):
            for row in range(span[1], span[3] + 1):
                origin[(col, row)] = (span[0], span[1])
    return origin, spans


def _components_union(atoms: list[dict[str, Any]]) -> dict[str, Any]:
    out: dict[str, Any] = {}
    for atom in atoms:
        for key, value in (atom.get("components") or {}).items():
            if key == "formula":
                continue
            if key not in out and value is not None:
                out[key] = value
    return out


def _apply_formula_atoms(
    snaps: dict[tuple[int, int], dict[str, Any]],
    atoms_at: dict[tuple[int, int], list[dict[str, Any]]],
    *,
    features: set[str],
) -> None:
    if "FORMULA_DERIVED_DATE" not in features:
        return
    formula_keys = sorted(
        (key for key, snap in snaps.items() if snap.get("formula")),
        key=lambda item: (item[1], item[0]),
    )
    for _ in range(8):
        progressed = False
        for (col, row) in formula_keys:
            if atoms_at.get((col, row)):
                continue
            snap = snaps[(col, row)]
            formula = snap.get("formula") or ""
            if not formula:
                continue
            body = formula[1:] if formula.startswith("=") else formula
            match = _EDATE.match(body)
            if match:
                ref = _parse_a1(match.group(2))
                months = int(match.group(3))
                if not ref or not atoms_at.get(ref):
                    continue
                base = _components_union(atoms_at[ref])
                if "year" not in base or "month" not in base:
                    continue
                month0 = int(base["month"]) + months - 1
                year = int(base["year"]) + month0 // 12
                month = month0 % 12 + 1
                atoms_at[(col, row)] = [
                    {
                        "source_cell": snap["address"],
                        "row": row,
                        "col": col,
                        "raw_value": formula,
                        "kind": "FORMULA_DERIVED_DATE",
                        "components": {"year": year, "month": month},
                        "confidence_source": f"{match.group(1).upper()}({match.group(2)},{months})",
                        "formula": formula,
                        "depends_on": a1_address(*ref),
                    }
                ]
                progressed = True
                continue
            match = _DATE_FN.match(body)
            if match:
                year, month, day = int(match.group(1)), int(match.group(2)), int(match.group(3))
                if not (1990 <= year <= 2100 and 1 <= month <= 12):
                    continue
                atoms_at[(col, row)] = [
                    {
                        "source_cell": snap["address"],
                        "row": row,
                        "col": col,
                        "raw_value": formula,
                        "kind": "FORMULA_DERIVED_DATE",
                        "components": {"year": year, "month": month, "day": day},
                        "confidence_source": "DATE()",
                        "formula": formula,
                    }
                ]
                progressed = True
                continue
            match = _ADD_DAYS.match(body)
            if match:
                ref = _parse_a1(match.group(1))
                days = int(match.group(2))
                if not ref or not atoms_at.get(ref) or days not in {1, 7, 28, 29, 30, 31, 365, 366}:
                    continue
                base = _components_union(atoms_at[ref])
                if "year" not in base or "month" not in base:
                    continue
                day = int(base.get("day") or 1)
                try:
                    stamp = dt.date(int(base["year"]), int(base["month"]), day) + dt.timedelta(days=days)
                except ValueError:
                    continue
                atoms_at[(col, row)] = [
                    {
                        "source_cell": snap["address"],
                        "row": row,
                        "col": col,
                        "raw_value": formula,
                        "kind": "FORMULA_DERIVED_DATE",
                        "components": {"year": stamp.year, "month": stamp.month, "day": stamp.day},
                        "confidence_source": f"add_days({days})",
                        "formula": formula,
                        "depends_on": a1_address(*ref),
                    }
                ]
                progressed = True
        if not progressed:
            break


def _header_rows(atoms_at: dict[tuple[int, int], list[dict[str, Any]]], min_row: int, max_row: int) -> set[int]:
    by_row: dict[int, int] = defaultdict(int)
    for (col, row), atoms in atoms_at.items():
        if atoms:
            by_row[row] += 1
    chosen = {row for row, n in by_row.items() if n >= 2}
    chosen.update(row for row, n in by_row.items() if n >= 1 and row <= min_row + 12)
    if not chosen:
        chosen.update(row for row, n in by_row.items() if n >= 1)
    return {row for row in chosen if min_row <= row <= max_row}


def _header_cols(atoms_at: dict[tuple[int, int], list[dict[str, Any]]], min_col: int, max_col: int) -> set[int]:
    by_col: dict[int, int] = defaultdict(int)
    for (col, row), atoms in atoms_at.items():
        if atoms:
            by_col[col] += 1
    chosen = {col for col, n in by_col.items() if n >= 2}
    chosen.update(col for col, n in by_col.items() if n >= 1 and col <= min_col + 8)
    if not chosen:
        chosen.update(col for col, n in by_col.items() if n >= 1)
    return {col for col in chosen if min_col <= col <= max_col}


def _classes_for(atoms: list[dict[str, Any]], *, axis: str, inherited: list[str]) -> list[str]:
    classes = []
    kinds = {a["kind"] for a in atoms}
    rows = {a["row"] for a in atoms}
    if "DATE_SERIAL" in kinds:
        classes.append("DATE_SERIAL_DECODING")
    if "FORMULA_DERIVED_DATE" in kinds or "FORMULA_HEADER" in kinds:
        classes.append("FORMULA_DERIVED_DATE")
    if len(rows) > 1:
        classes.append("MULTIROW_HEADER_COMPOSITION")
    if "merge" in inherited:
        classes.append("MERGED_HEADER_PROPAGATION")
    if "block" in inherited:
        classes.append("BLOCK_INHERITANCE")
    if axis == "row":
        classes.append("ROW_AXIS_SUPPORT")
    if any(a.get("confidence_source") in {"fy_loose", "fy_token_v1"} for a in atoms):
        classes.append("FY_TOKEN_NORMALIZATION")
    if any("MARKER" in a["kind"] for a in atoms):
        classes.append("ACTUAL_FORECAST_MARKER")
    if not classes:
        classes.append("V1_PARSE_PERIOD")
    return classes


HEADER_ROW_CAP = 48
HEADER_COL_CAP = 12


def _keep_as_temporal_candidate(col: int, row: int, value: object) -> bool:
    if row <= HEADER_ROW_CAP or col <= HEADER_COL_CAP:
        return True
    text = formula_text(value)
    return bool(text and _DATE_FORMULA_HINT.search(text))


def compile_temporal_sheet(
    sheet: Any,
    *,
    sheet_index: int,
    features: set[str],
    epoch: Any = None,
) -> dict[str, Any]:
    merge_origin, merge_spans = _merge_map(sheet)
    snaps: dict[tuple[int, int], dict[str, Any]] = {}
    min_col = min_row = None
    max_col = max_row = None
    for cell in sheet._cells.values():
        col, row = int(cell.column), int(cell.row)
        min_col = col if min_col is None else min(min_col, col)
        max_col = col if max_col is None else max(max_col, col)
        min_row = row if min_row is None else min(min_row, row)
        max_row = row if max_row is None else max(max_row, row)
        if not _keep_as_temporal_candidate(col, row, cell.value):
            continue
        snaps[(col, row)] = snapshot_cell(cell, merge_origin=merge_origin)
    if min_col is None:
        return {
            "sheet_title": sheet.title,
            "sheet_index": sheet_index,
            "atoms": [],
            "coordinates": [],
            "n_snaps": 0,
        }
    atoms_at: dict[tuple[int, int], list[dict[str, Any]]] = {}
    for key, snap in snaps.items():
        found = atoms_from_snapshot(snap, features=features, epoch=epoch)
        if found:
            atoms_at[key] = found

    if "MERGED_HEADER_PROPAGATION" in features:
        for span in merge_spans:
            if span[1] > HEADER_ROW_CAP and span[0] > HEADER_COL_CAP:
                continue
            origin = (span[0], span[1])
            if not atoms_at.get(origin):
                continue
            for col in range(span[0], span[2] + 1):
                for row in range(span[1], span[3] + 1):
                    if (col, row) == origin:
                        continue
                    cloned = []
                    for atom in atoms_at[origin]:
                        item = dict(atom)
                        item["inherited_from"] = a1_address(*origin)
                        item["propagation_rule"] = "merged_range"
                        item["stopping_boundary"] = a1_address(span[2], span[3])
                        item["row"] = row
                        item["col"] = col
                        cloned.append(item)
                    atoms_at.setdefault((col, row), []).extend(cloned)

    _apply_formula_atoms(snaps, atoms_at, features=features)

    header_rows = _header_rows(atoms_at, min_row, max_row)
    header_cols = _header_cols(atoms_at, min_col, max_col)

    if "BLOCK_INHERITANCE" in features:
        for row in sorted(header_rows):
            carry = None
            carry_from = None
            for col in range(min_col, max_col + 1):
                local = atoms_at.get((col, row)) or []
                comps = _components_union(local)
                if comps.get("year"):
                    carry = comps["year"]
                    carry_from = a1_address(col, row)
                    continue
                stacked = []
                for hr in header_rows:
                    stacked.extend(atoms_at.get((col, hr)) or [])
                stacked_c = _components_union(stacked)
                if carry and (stacked_c.get("month") or stacked_c.get("quarter")) and "year" not in stacked_c:
                    atoms_at.setdefault((col, row), []).append(
                        {
                            "source_cell": carry_from,
                            "row": row,
                            "col": col,
                            "raw_value": None,
                            "kind": "YEAR_LITERAL",
                            "components": {"year": carry},
                            "confidence_source": "block_inherit_year",
                            "inherited_from": carry_from,
                            "propagation_rule": "year_until_next_explicit",
                            "stopping_boundary": "next_year_atom_or_sheet_edge",
                        }
                    )

    coordinates: list[dict[str, Any]] = []
    sid = f"sheet:s{sheet_index:02d}"

    def emit_coord(axis: str, index: int, atoms: list[dict[str, Any]], row: int, col: int) -> None:
        comps = _components_union(atoms)
        if not any(k in comps for k in ("year", "month", "quarter")):
            return
        inherited = []
        if any(a.get("propagation_rule") == "merged_range" for a in atoms):
            inherited.append("merge")
        if any(a.get("propagation_rule") == "year_until_next_explicit" for a in atoms):
            inherited.append("block")
        classes = _classes_for(atoms, axis=axis, inherited=inherited)
        if axis == "row" and "ROW_AXIS_SUPPORT" not in features:
            return
        if "MULTIROW_HEADER_COMPOSITION" not in features and "MULTIROW_HEADER_COMPOSITION" in classes:
            same_row = [a for a in atoms if a["row"] == row]
            comps = _components_union(same_row)
            atoms = same_row
            if not any(k in comps for k in ("year", "month", "quarter")):
                return
            classes = _classes_for(atoms, axis=axis, inherited=inherited)
        coord_id = f"tcoord:s{sheet_index:02d}:{axis[0]}:{index}"
        coordinates.append(
            {
                "id": coord_id,
                "axis": axis,
                "sheet_id": sid,
                "sheet_title": sheet.title,
                "row": row,
                "col": col,
                "row_id": f"row:s{sheet_index:02d}:r{row}",
                "col_id": f"col:s{sheet_index:02d}:c{col}",
                "cell_id": f"cell:s{sheet_index:02d}:r{row}:c{col}",
                "address": a1_address(col, row),
                "components": comps,
                "period": {k: comps[k] for k in ("year", "month", "quarter") if k in comps},
                "period_key": period_key({k: comps[k] for k in ("year", "month", "quarter") if k in comps}),
                "encoding_class": classes,
                "evidence_cells": sorted({a.get("source_cell") or a1_address(a["col"], a["row"]) for a in atoms}),
                "inherited_from": [a.get("inherited_from") for a in atoms if a.get("inherited_from")],
                "propagation_rule": [a.get("propagation_rule") for a in atoms if a.get("propagation_rule")],
                "header_text": next((str(a.get("raw_value")) for a in atoms if a.get("raw_value") is not None), None),
                "provenance": [
                    {
                        "cell": a.get("source_cell"),
                        "kind": a["kind"],
                        "source": a.get("confidence_source"),
                    }
                    for a in atoms
                ],
            }
        )

    col_atoms: dict[int, list[dict[str, Any]]] = defaultdict(list)
    for (col, row), atoms in atoms_at.items():
        if row in header_rows:
            col_atoms[col].extend(atoms)
    for col, atoms in col_atoms.items():
        primary_row = min((a["row"] for a in atoms), default=min_row)
        emit_coord("column", col, atoms, primary_row, col)

    if "ROW_AXIS_SUPPORT" in features:
        row_atoms: dict[int, list[dict[str, Any]]] = defaultdict(list)
        for (col, row), atoms in atoms_at.items():
            if col in header_cols:
                row_atoms[row].extend(atoms)
        for row, atoms in row_atoms.items():
            primary_col = min((a["col"] for a in atoms), default=min_col)
            emit_coord("row", row, atoms, row, primary_col)

    flat_atoms = [a for group in atoms_at.values() for a in group]
    return {
        "sheet_title": sheet.title,
        "sheet_index": sheet_index,
        "bounds": {"min_col": min_col, "max_col": max_col, "min_row": min_row, "max_row": max_row},
        "n_snaps": len(snaps),
        "n_atoms": len(flat_atoms),
        "header_rows": sorted(header_rows)[:40],
        "header_cols": sorted(header_cols)[:40],
        "coordinates": coordinates,
        "atoms_sample": flat_atoms[:80],
    }


def compile_temporal_workbook(
    path: Path,
    *,
    features: set[str] | None = None,
    closure: bool = False,
) -> dict[str, Any]:
    features = set(features or ALL_FEATURES)
    try:
        workbook = openpyxl.load_workbook(path, data_only=False, read_only=False)
    except Exception as exc:
        return {
            "readable": False,
            "error": f"{type(exc).__name__}: {exc}",
            "features": sorted(features),
            "n_sheets": 0,
            "n_coordinates": 0,
            "sheets": [],
            "coordinates": [],
            "closure": None,
            "golden_used": False,
        }
    try:
        epoch = getattr(workbook, "epoch", None)
        sheets = []
        coordinates = []
        for index, sheet in enumerate(workbook.worksheets):
            compiled = compile_temporal_sheet(sheet, sheet_index=index, features=features, epoch=epoch)
            sheets.append({k: compiled[k] for k in compiled if k != "coordinates"})
            for coord in compiled["coordinates"]:
                coord.setdefault("derivation_status", "LOCAL_TEMPORAL_COORDINATE")
                coordinates.append(coord)
        closure_stats = None
        conflicts: list[dict[str, Any]] = []
        if closure:
            from temporal_provenance import close_temporal_workbook

            closed = close_temporal_workbook(
                workbook, coordinates, features=features, epoch=epoch
            )
            coordinates = closed["coordinates"]
            conflicts = closed.get("conflicts") or []
            closure_stats = closed.get("stats")
        return {
            "readable": True,
            "error": None,
            "features": sorted(features),
            "n_sheets": len(sheets),
            "n_coordinates": len(coordinates),
            "n_local_coordinates": closure_stats["n_local"] if closure_stats else len(coordinates),
            "n_propagated_coordinates": closure_stats["n_propagated"] if closure_stats else 0,
            "sheets": sheets,
            "coordinates": coordinates,
            "temporal_conflicts": conflicts,
            "closure": closure_stats,
            "golden_used": False,
        }
    except Exception as exc:
        return {
            "readable": False,
            "error": f"{type(exc).__name__}: {exc}",
            "features": sorted(features),
            "n_sheets": 0,
            "n_coordinates": 0,
            "sheets": [],
            "coordinates": [],
        }
    finally:
        workbook.close()


def coordinates_as_periods(compiled: dict[str, Any]) -> list[dict[str, Any]]:
    """View V2 coordinates as retrieve_scope period records."""
    out = []
    for coord in compiled.get("coordinates") or []:
        period = coord.get("period") or {}
        if not period:
            continue
        out.append(
            {
                "id": coord["id"],
                "cell_id": coord["cell_id"],
                "sheet_id": coord["sheet_id"],
                "row_id": coord["row_id"],
                "col_id": coord["col_id"],
                "row": coord["row"],
                "col": coord["col"],
                "address": coord["address"],
                "period": period,
                "period_key": coord.get("period_key"),
                "header_text": coord.get("header_text"),
                "axis": coord.get("axis"),
                "encoding_class": coord.get("encoding_class"),
            }
        )
    return out


def overlay_periods(spine: dict[str, Any], compiled: dict[str, Any]) -> dict[str, Any]:
    """Replace spine periods with TEMPORAL_COORDINATE views. Retrieval rules unchanged."""
    from workbook_grounding_spine import rebuild_id_set

    over = dict(spine)
    over["periods"] = coordinates_as_periods(compiled)
    rebuild_id_set(over)
    return over


def neighborhood_strip(
    sheet: Any,
    *,
    gold_row: int,
    gold_col: int,
    row_span: int = 8,
    col_span: int = 4,
    merge_origin: dict[tuple[int, int], tuple[int, int]] | None = None,
) -> dict[str, Any]:
    """Bounded input-only neighborhood. Gold locates the coordinate only."""
    if merge_origin is None:
        merge_origin, _spans = _merge_map(sheet)
    # Do not scan the used range. Inspect a header band at the top of the
    # sheet and a local window above the gold cell.
    rows = set(range(1, row_span + 1))
    rows.update(range(max(1, gold_row - row_span), gold_row + 1))
    cols = range(max(1, gold_col - col_span), gold_col + col_span + 1)
    cells = []
    for row in sorted(rows)[:24]:
        for col in cols:
            cell = sheet.cell(row=row, column=col)
            cells.append(snapshot_cell(cell, merge_origin=merge_origin))
    return {
        "gold_row": gold_row,
        "gold_col": gold_col,
        "header_rows": sorted(rows)[:24],
        "header_cols": list(cols),
        "cells": cells,
    }
