"""Experimental live bootstrap for the frozen Candidate-A surface.

This module is loaded only through ``sitecustomize`` in the live treatment.
It is not production runtime code and adds no model-visible API.
"""
from __future__ import annotations

import json
import os
import sqlite3
import time
from pathlib import Path
from typing import Any, Callable

from benchmark.inspection_helpers import index

from benchmark.candidate_a_shadow_interposition import (
    CandidateALoader,
    CompiledSnapshot,
    ProxyWorkbook,
    WorkbookMetadata,
)


class PersistentCompiledSnapshot(CompiledSnapshot):
    """Read-only view of the common substrate prepared by the parent runner."""

    def __init__(self, path: str, entry: dict[str, Any]):
        self.path = str(path)
        failed = index._failures.get(self.path)
        if failed and failed["workbook_generation"] in (None, entry.get("workbook_hash")):
            raise index.SubstrateDisabled(failed)
        self.handle = {"valid": False, "db": None}
        stage = "freshness_initialization"
        try:
            if index.workbook_hash(path) != entry["workbook_hash"]:
                raise ValueError("stale substrate generation")
            stage = "workbook_metadata_parse"
            self.meta = WorkbookMetadata(path)
            stage = "substrate_availability"
            # mode=ro cannot silently create a missing/empty database.
            con = sqlite3.connect(Path(entry["db_path"]).resolve().as_uri() + "?mode=ro", uri=True)
            self.handle["db"] = con
            con.execute("SELECT sheet,addr,row,col,value,formula,dtype FROM cells LIMIT 1").fetchall()
            identity = con.execute("SELECT workbook_hash,index_generation FROM substrate_identity").fetchall()
            if identity != [(entry["workbook_hash"], entry["index_generation"])]:
                raise ValueError("incomplete or wrong-generation database")
            stage = "freshness_verification"
            if index.workbook_hash(path) != entry["workbook_hash"]:
                raise ValueError("workbook changed during snapshot initialization")
            self.handle.update(workbook_hash=entry["workbook_hash"],
                               index_generation=entry["index_generation"], valid=True)
            self.rebuilt = False
            index._snapshot_handles.setdefault(self.path, []).append(self.handle)
        except Exception as exc:
            if self.handle["db"] is not None:
                self.handle["db"].close()
            event = index.disable(path, entry.get("workbook_hash"), stage, exc,
                                  partial=self.handle["db"] is not None,
                                  substrate_generation=entry.get("index_generation"))
            raise index.SubstrateDisabled(event) from exc


class TelemetryList(list):
    def __init__(self, callback: Callable[[str], None]):
        super().__init__()
        self.callback = callback

    def append(self, value: str) -> None:
        super().append(value)
        self.callback(value)


def emit(payload: dict[str, Any]) -> None:
    path = os.environ.get("CANDIDATE_A_EXEC_TELEMETRY")
    if not path:
        return
    row = {
        "task_id": os.environ.get("CANDIDATE_A_TASK"),
        "run_id": os.environ.get("CANDIDATE_A_RUN_ID"),
        "arm": os.environ.get("CANDIDATE_A_ARM"),
        "python_pid": os.getpid(),
        "timestamp_ns": time.perf_counter_ns(),
        **payload,
    }
    try:
        with Path(path).open("a", encoding="utf-8") as handle:
            handle.write(json.dumps(row, sort_keys=True, default=str) + "\n")
    except OSError:
        # Telemetry failure must never change the agent's Python behavior.
        pass


def observable_result(value: Any) -> dict[str, Any]:
    """Small JSON-safe timing/trace value; never affects the returned value."""
    import datetime as _dt
    if value is None or isinstance(value, (str, int, float, bool)):
        return {"type": type(value).__name__, "value": value}
    if isinstance(value, (_dt.datetime, _dt.date, _dt.time)):
        return {"type": type(value).__name__, "value": value.isoformat()}
    return {"type": type(value).__name__, "repr": repr(value)}


def install() -> None:
    import openpyxl

    original = openpyxl.load_workbook
    manifest_path = os.environ.get("CANDIDATE_A_SUBSTRATE_MANIFEST")
    manifest: dict[str, Any] = {}
    if manifest_path and not os.environ.get("CANDIDATE_A_SUBSTRATE_DISABLED"):
        try:
            manifest = json.loads(Path(manifest_path).read_text(encoding="utf-8"))
            if not isinstance(manifest.get("workbooks"), dict):
                raise ValueError("invalid substrate manifest")
        except Exception as exc:
            emit(index.disable(manifest_path, None, "manifest_initialization", exc))
            manifest = {}

    class SharedLoader(CandidateALoader):
        def load_workbook(self, filename: str, *args: Any, **kwargs: Any) -> Any:
            started_ns = time.perf_counter_ns()
            data_only = bool(kwargs.get("data_only", False))
            read_only = bool(kwargs.get("read_only", False))
            write_only = bool(kwargs.get("write_only", False))
            supported = (
                isinstance(filename, (str, os.PathLike))
                and not args and not data_only and not read_only and not write_only
                and Path(filename).suffix.lower() in {".xlsx", ".xlsm"}
                and set(kwargs) <= {"data_only", "read_only", "write_only"}
                and not os.environ.get("CANDIDATE_A_FORCE_REAL")
            )
            resolved = str(Path(filename).resolve()) if isinstance(filename, (str, os.PathLike)) else None
            reason = "forced real path" if os.environ.get("CANDIDATE_A_FORCE_REAL") else "unsupported load mode/options"
            if supported:
                entry = manifest.get("workbooks", {}).get(resolved)
                reason = "shared substrate unavailable"
                if resolved not in self.disabled:
                    try:
                        if not entry:
                            prior = next((e for e in manifest.get("failures", [])
                                          if e.get("path") == resolved), None)
                            if prior:
                                raise index.SubstrateDisabled(prior)
                            raise FileNotFoundError("shared substrate unavailable")
                        snapshot = self.snapshots.get(resolved)
                        # Loads are freshness boundaries; existing proxy objects
                        # retain their loaded-workbook semantics until retired.
                        if snapshot is not None and index.workbook_hash(resolved) != snapshot.handle["workbook_hash"]:
                            snapshot.handle["valid"] = False
                            raise ValueError("stale substrate generation")
                        if snapshot is None:
                            snapshot = PersistentCompiledSnapshot(resolved, entry)
                            self.snapshots[resolved] = snapshot
                        if not snapshot.handle["valid"]:
                            raise ValueError("retired substrate")
                        result = ProxyWorkbook(self, snapshot)
                    except Exception as exc:
                        event = (exc.event if isinstance(exc, index.SubstrateDisabled) else
                                 index.disable(resolved, entry.get("workbook_hash") if entry else None,
                                               "snapshot_initialization", exc,
                                               partial=resolved in self.snapshots,
                                               served=resolved in self.served,
                                               substrate_generation=entry.get("index_generation") if entry else None))
                        self.disabled.add(resolved)
                        old = self.snapshots.pop(resolved, None)
                        if old is not None:
                            old.handle["valid"] = False
                            old.handle["db"].close()
                        self.emit("substrate_fallback", **{k: v for k, v in event.items() if k != "event"})
                    else:
                        snapshot.handle["accelerated_served"] = True
                        self.served.add(resolved)
                        self.emit("candidate_operation", operation="load_workbook", path=resolved, status="ACCELERATED", workbook_generation=entry["workbook_hash"], substrate_generation=entry["index_generation"], rebuilt=False, parse_performed=False, substrate_hit=True, duration_ns=time.perf_counter_ns() - started_ns)
                        return result
                reason = "substrate disabled for workbook in this Python process"
            # Record before calling reference: its own exception must propagate
            # unchanged, and still leave an observable fallback witness.
            self.emit("candidate_operation", operation="load_workbook", path=resolved, status="PREDECLARED_FALLBACK", fallback_reason=reason, data_only=data_only, read_only=read_only, write_only=write_only, parse_performed=True, substrate_hit=False, duration_ns=time.perf_counter_ns() - started_ns)
            self.fallback_reasons.append(reason)
            return original(filename, *args, **kwargs)

    loader = SharedLoader(real_loader=original, event_sink=lambda row: emit(row))
    loader.fallback_reasons = TelemetryList(
        lambda reason: emit({
            "event": "candidate_operation",
            "operation": "fallback",
            "status": "RUNTIME_FALLBACK",
            "fallback_reason": reason,
        })
    )
    openpyxl.load_workbook = loader.load_workbook
    emit({"event": "bootstrap", "status": "INSTALLED", "candidate": "A"})


def install_h0_spy() -> None:
    import openpyxl

    original = openpyxl.load_workbook

    def load_workbook(filename: str, *args: Any, **kwargs: Any) -> Any:
        started_ns = time.perf_counter_ns()
        data_only = bool(kwargs.get("data_only", False))
        read_only = bool(kwargs.get("read_only", False))
        source_decision = os.environ.get("CANDIDATE_A_A1_DECISION")
        safe = (
            not args
            and not data_only
            and not read_only
            and not kwargs.get("write_only", False)
            and set(kwargs) <= {"data_only", "read_only", "write_only"}
        )
        if source_decision:
            safe = source_decision == "A1_ADMIT"
        result = original(filename, *args, **kwargs)
        emit({
            "event": "candidate_operation",
            "operation": "load_workbook",
            "status": "H0_COUNTERFACTUAL_ELIGIBLE" if safe else "H0_COUNTERFACTUAL_REJECTED",
            "data_only": data_only,
            "read_only": read_only,
            "fallback_reason": None if safe else os.environ.get("CANDIDATE_A_A1_REASON", "unsupported load mode/options"),
            "parse_performed": True,
            "substrate_hit": False,
            "duration_ns": time.perf_counter_ns() - started_ns,
        })
        return result

    openpyxl.load_workbook = load_workbook
    emit({"event": "bootstrap", "status": "INSTALLED", "candidate": "H0_SPY"})


def main() -> None:
    import openpyxl
    original = openpyxl.load_workbook
    try:
        if os.environ.get("CANDIDATE_A_ARM") == "H1":
            install()
        elif os.environ.get("CANDIDATE_A_ARM") == "H0":
            install_h0_spy()
    except Exception as exc:
        partial = openpyxl.load_workbook is not original
        openpyxl.load_workbook = original
        emit(index.disable("<runtime>", None, "runtime_installation", exc,
                           partial=partial, served=False))


main()
