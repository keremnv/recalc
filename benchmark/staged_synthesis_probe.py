#!/usr/bin/env python3
"""Gold-blind staged construction probe over the four clean canonicals.

F0 is the current one-shot synthesis.  F1 persists a model-generated
    computation sketch, then asks a second call to bind its holes.
F2 persists a model-generated closed-world operand set, then asks a second call
    to construct the computation over that set.

Gold is loaded only by the post-run evaluator.  Primary prompts are built from
the stored sessions and frozen working sets, and every request/response body is
retained so this probe does not repeat the previous reconstruction limitation.
"""
from __future__ import annotations

import hashlib
import json
import os
import re
import sys
import time
import urllib.error
import urllib.request
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parent))
import canonical_choice_probe as cp
import composition_counterfactual as ccf
import end_to_end_composition_probe as old
import execution_unit_probe as P
import execution_unit_score as ES
import instrumentation_freeze as instr
import program_group as pg
import synthesis_decomposition as sd
import workbook_grounding as wg
from xlsx_cell_writer import write_cells


OUT = old.MECHANICAL / "staged-synthesis-probe"
CALLS = OUT / "calls"
VARIANTS = OUT / "variants"
MODEL = old.MODEL
TEMPERATURE = old.TEMPERATURE
REASONING = old.REASONING
MAX_TOKENS = old.SYNTHESIS_MAX_TOKENS

# Deliberately fixed before reading any model outcome.  C68/D68 are retained as
# two cells because the frozen sessions/evidence differ, but the report groups
# them into one independent program under deterministic translation.
POPULATION = (
    ("08_04", "Working Capital!C68"),
    ("08_04", "Working Capital!D68"),
    ("08_05", "Working Capital!C66"),
    ("08_03", "Working Capital!J50"),
)

FROZEN = (
    "benchmark/end_to_end_composition_probe.py",
    "benchmark/execution_unit_probe.py",
    "benchmark/program_group.py",
    "benchmark/xlsx_cell_writer.py",
    "benchmark/instrumentation_freeze.py",
    "benchmark/composition_counterfactual.py",
    "benchmark/execution_unit_score.py",
)

OPS = {"ADD": "+", "SUB": "-", "MUL": "*", "DIV": "/"}
ALLOWED_CALLS = {"SUM", "AVERAGE", "MIN", "MAX", "IF"}
PREC = {"ADD": 1, "SUB": 1, "MUL": 2, "DIV": 2, "NEG": 3}

# Keep the frozen relational schema/evidence contract, but do not carry either
# the retrieval-action protocol or the one-shot formula-output transition into
# staged prompts.  Both are instructions for other stages and otherwise
# compete with the stage schema.
_spine_start = old.SYNTHESIS_SYSTEM.index("# Workbook spine relational contract")
_spine_end = old.SYNTHESIS_SYSTEM.index("ARCHITECTURE V1.1 SYNTHESIS PROTOCOL")
COMMON = old.SYNTHESIS_SYSTEM[_spine_start:_spine_end] + "\n\n"
F1A_SYSTEM = COMMON + """

STAGED CONSTRUCTION PROTOCOL — F1a PROGRAM SKETCH
Return JSON only.  Infer the computation, but do not choose workbook
references.  The sketch is an AST using only these nodes:
  {"op":"ADD|SUB|MUL|DIV|NEG", "args":[...]}
  {"op":"SUM|AVERAGE|MIN|MAX|IF", "args":[...]}
  {"op":"REF_HOLE", "name":"ARG1", "type":"SCALAR|RANGE"}
  {"op":"CONST", "value": number}
Use numbered ARG holes.  Do not put sheet names, A1 addresses, stable IDs, or
formula strings in the sketch.  You may abstain.
Return exactly:
{"status":"PROPOSED","sketch":<AST>} or {"status":"ABSTAIN","sketch":null}
"""
F1B_SYSTEM = COMMON + """

STAGED CONSTRUCTION PROTOCOL — F1b BIND MODEL SKETCH
The MODEL_GENERATED_SKETCH in the user payload is the only prior-stage artifact.
Keep the computation unchanged and bind every REF_HOLE to a closed-world entity
ID from WORKING_SET_ENTITY_IDS.  Use only cell:<...> or range:<...> IDs present
there.  Do not invent addresses or formulas.  Bind each hole exactly once.
You may abstain.
Return exactly:
{"status":"PROPOSED","bindings":{"ARG1":"cell:..."}} or
{"status":"ABSTAIN","bindings":null}
"""
F2A_SYSTEM = COMMON + """

STAGED CONSTRUCTION PROTOCOL — F2a MODEL OPERAND SET
Identify the workbook entities needed to compute the target, but do not choose
operators, functions, ordering, or a formula.  Select only closed-world cell:
or range: IDs present in WORKING_SET_ENTITY_IDS.  Return a set, not a formula.
You may abstain.
Return exactly:
{"status":"PROPOSED","operand_ids":["cell:..."]} or
{"status":"ABSTAIN","operand_ids":null}
"""
F2B_SYSTEM = COMMON + """

STAGED CONSTRUCTION PROTOCOL — F2b STRUCTURE OVER MODEL OPERANDS
The MODEL_SELECTED_OPERAND_SET in the user payload is the only prior-stage
artifact.  Construct the computation using only those selected entity IDs.
Return an AST using:
  {"op":"ADD|SUB|MUL|DIV|NEG", "args":[...]}
  {"op":"SUM|AVERAGE|MIN|MAX|IF", "args":[...]}
  {"op":"REF","id":"cell:..."} or {"op":"CONST","value":number}
An operand may be reused; a selected operand may be omitted.  Do not introduce
a workbook ID outside the selected set.  You may abstain.
Return exactly:
{"status":"PROPOSED","structure":<AST>} or
{"status":"ABSTAIN","structure":null}
"""


def sha(value: Any) -> str:
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, separators=(",", ":"), sort_keys=True).encode()).hexdigest()


def write(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(value if isinstance(value, str) else json.dumps(value, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def load(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def units() -> dict[tuple[str, str], dict]:
    return sd.units_by_key()


def session_for(task: str, member: str) -> dict:
    u = units()[(task, member)]
    return load(sd.session_for(u))


def raw_task_for(task: str) -> str:
    return P.authorised(task)[0]["compiler"]["raw_task"]


def base_payload(task: str, member: str, session: dict) -> tuple[str, dict]:
    # This is the same frozen transition used by the current one-shot arm.  No
    # gold formula or evaluator label is present.
    working = set(session["working_set_ids"])
    evidence = old.materialize_working_set(task, working)
    transition = {
        "SYNTHESIS_TRANSITION": old.SYNTHESIS_TRANSITION,
        "RAW_TASK": raw_task_for(task),
        "TARGET": session["target"],
        "OBLIGATION": session["obligation"],
        "WORKING_SET_COUNTS": dict(Counter(x.split(":", 1)[0] for x in working)),
        "WORKING_SET_ENTITY_IDS": sorted(working),
        "WORKING_SET_EVIDENCE": evidence,
    }
    return json.dumps(transition, ensure_ascii=False, separators=(",", ":")), transition


def stage_payload(base: dict, stage_instruction: str, artifact_key: str | None = None, artifact: Any = None) -> str:
    out = dict(base)
    out["SYNTHESIS_TRANSITION"] = stage_instruction
    if artifact_key is not None:
        out[artifact_key] = artifact
    return json.dumps(out, ensure_ascii=False, separators=(",", ":"))


def call_model(name: str, stage: str, system: str, user: str, key: str, prior: Any = None) -> dict:
    path = CALLS / f"{name}.json"
    if path.exists():
        return load(path)
    body = {
        "model": MODEL,
        "temperature": TEMPERATURE,
        "reasoning": {"effort": REASONING},
        "max_tokens": MAX_TOKENS,
        "response_format": {"type": "json_object"},
        "messages": [{"role": "system", "content": system}, {"role": "user", "content": user}],
    }
    started = time.perf_counter()
    raw_body = ""
    status_code = None
    payload = None
    failure = None
    try:
        req = urllib.request.Request(
            old.relational.OPENROUTER_URL,
            data=json.dumps(body).encode(),
            headers={"Authorization": f"Bearer {key}", "Content-Type": "application/json", "X-Title": "librecalc-staged-synthesis"},
            method="POST",
        )
        with urllib.request.urlopen(req, timeout=300) as response:
            status_code = response.status
            raw_body = response.read().decode(errors="replace")
        payload = json.loads(raw_body)
    except urllib.error.HTTPError as exc:
        status_code = exc.code
        raw_body = exc.read().decode(errors="replace")
        failure = "MODEL_ACCESS_FAILURE"
    except Exception as exc:
        raw_body = f"{type(exc).__name__}: {exc}"
        failure = "MODEL_ACCESS_FAILURE"
    choices = (payload or {}).get("choices") or []
    message = (choices[0].get("message") or {}) if choices else {}
    text = message.get("content") or ""
    finish = choices[0].get("finish_reason") if choices else None
    parsed = old.extract_json_object(text) if payload is not None else None
    usage = (payload or {}).get("usage") or {}
    at_budget = int(usage.get("completion_tokens") or 0) == MAX_TOKENS
    trunc = None
    if failure is None and at_budget and parsed is None:
        trunc = "TRUNCATED_NO_CONTENT" if not text else "TRUNCATED_AT_BUDGET"
    rec = {
        "name": name, "stage": stage, "request": body,
        "request_sha256": sha(body), "system_prompt": system, "user_payload": user,
        "system_sha256": hashlib.sha256(system.encode()).hexdigest(),
        "user_payload_sha256": hashlib.sha256(user.encode()).hexdigest(),
        "prior_stage_artifact": prior,
        "prior_stage_artifact_sha256": sha(prior) if prior is not None else None,
        "response_http_status": status_code, "response_raw_body": raw_body,
        "response_body": payload, "raw_response_text": text,
        "parsed_response": parsed, "finish_reason": finish, "usage": usage,
        "truncation_class": trunc, "failure_class": failure,
        "input_tokens": int(usage.get("prompt_tokens") or 0),
        "output_tokens": int(usage.get("completion_tokens") or 0),
        "wall_seconds": round(time.perf_counter() - started, 3),
    }
    write(path, rec)
    return rec


# ------------------------------ AST parsing, validation and serialization

def _as_number(v: Any) -> str | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return str(v)
    if isinstance(v, str) and re.fullmatch(r"[-+]?\d+(?:\.\d+)?", v.strip()):
        return v.strip()
    return None


def normalize_ast(node: Any, mode: str, allowed_refs: set[str] | None = None) -> tuple[Any | None, str | None]:
    if not isinstance(node, dict):
        return None, "INVALID_SKETCH"
    op = str(node.get("op") or "").upper()
    if op == "REF_HOLE":
        if mode != "F1":
            return None, "INVALID_SKETCH"
        name = node.get("name")
        typ = str(node.get("type") or "SCALAR").upper()
        if not isinstance(name, str) or not re.fullmatch(r"ARG\d+", name) or typ not in {"SCALAR", "RANGE"}:
            return None, "INVALID_SKETCH"
        return ("HOLE", name, typ), None
    if op == "REF" or "ref" in node:
        if mode != "F2":
            return None, "INVALID_SKETCH"
        rid = node.get("id", node.get("ref"))
        if not isinstance(rid, str) or not re.fullmatch(r"(?:cell|range):[^\s]+", rid):
            return None, "INVALID_ENTITY_ID"
        if allowed_refs is not None and rid not in allowed_refs:
            return None, "INVALID_ENTITY_ID"
        return ("REF", rid), None
    if op == "CONST":
        number = _as_number(node.get("value"))
        if number is None:
            return None, "INVALID_SKETCH"
        return ("CONST", number), None
    if op == "NEG":
        args = node.get("args")
        if not isinstance(args, list) or len(args) != 1:
            return None, "SKETCH_WRONG_ARITY"
        child, err = normalize_ast(args[0], mode, allowed_refs)
        return ("NEG", child), err if child is None else None
    if op in OPS or op in ALLOWED_CALLS:
        args = node.get("args")
        if not isinstance(args, list) or not args:
            return None, "SKETCH_WRONG_ARITY"
        if op in {"SUB", "DIV"} and len(args) != 2:
            return None, "SKETCH_WRONG_ARITY"
        children = []
        for a in args:
            child, err = normalize_ast(a, mode, allowed_refs)
            if child is None:
                return None, err or "INVALID_SKETCH"
            children.append(child)
        return (op, tuple(children)), None
    return None, "SKETCH_WRONG_OPERATOR"


TOKEN = re.compile(r"\s*(?:(\d+(?:\.\d+)?)|([A-Za-z_][A-Za-z0-9_ ]*!)?\$?([A-Z]{1,3})\$?(\d+)|([+\-*/(),]))")


def formula_tokens(formula: str) -> list[tuple[str, str]]:
    s = formula.lstrip("=")
    pos, out = 0, []
    while pos < len(s):
        m = TOKEN.match(s, pos)
        if not m:
            raise ValueError(f"unsupported formula token near {s[pos:pos+20]!r}")
        if m.group(1): out.append(("CONST", m.group(1)))
        elif m.group(2) or m.group(3): out.append(("REF", (m.group(2) or "") + m.group(3) + m.group(4)))
        else: out.append((m.group(5), m.group(5)))
        pos = m.end()
    return out


def gold_ast(formula: str) -> Any:
    ts, i = formula_tokens(formula), 0
    def expr(minp=0):
        nonlocal i
        if i >= len(ts): raise ValueError("missing expression")
        typ, val = ts[i]; i += 1
        if typ == "CONST": left = ("CONST", val)
        elif typ == "REF": left = ("REF", val)
        elif typ == "-": left = ("NEG", expr(3))
        elif typ == "(":
            left = expr(0)
            if i >= len(ts) or ts[i][0] != ")": raise ValueError("missing )")
            i += 1
        else: raise ValueError(f"unexpected {typ}")
        while i < len(ts) and ts[i][0] in {"+", "-", "*", "/"}:
            op = ts[i][0]; p = 1 if op in {"+", "-"} else 2
            if p < minp: break
            i += 1
            right = expr(p + 1)
            left = ({"+":"ADD", "-":"SUB", "*":"MUL", "/":"DIV"}[op], (left, right))
        return left
    result = expr(0)
    if i != len(ts): raise ValueError("trailing formula tokens")
    return result


def abstract(ast: Any) -> Any:
    if ast[0] in {"REF", "HOLE"}: return ("HOLE",)
    if ast[0] == "CONST": return ast
    if ast[0] == "NEG": return ("NEG", abstract(ast[1]))
    return (ast[0], tuple(abstract(x) for x in ast[1]))


def render(ast: Any, bindings: dict[str, str], id_to_a1: dict[str, str], parent=0) -> str:
    kind = ast[0]
    if kind == "HOLE":
        rid = bindings[ast[1]]
        return id_to_a1[rid]
    if kind == "REF":
        return id_to_a1[ast[1]]
    if kind == "CONST": return ast[1]
    if kind == "NEG":
        text = "-" + render(ast[1], bindings, id_to_a1, PREC["NEG"])
        return f"({text})" if parent > PREC["NEG"] else text
    if kind in OPS:
        p = PREC[kind]
        a, b = ast[1]
        left = render(a, bindings, id_to_a1, p)
        right = render(b, bindings, id_to_a1, p + (1 if kind in {"SUB", "DIV"} else 0))
        text = left + OPS[kind] + right
        return f"({text})" if parent > p else text
    raise ValueError(f"cannot render {kind}")


def ast_holes(ast: Any) -> list[str]:
    if ast[0] == "HOLE": return [ast[1]]
    if ast[0] in {"CONST", "REF"}: return []
    if ast[0] == "NEG": return ast_holes(ast[1])
    out = []
    for x in ast[1]: out.extend(ast_holes(x))
    return out


def id_to_a1(task: str, ids: set[str]) -> dict[str, str]:
    # Lower only entities present in the frozen workbook spine.  Formula IDs
    # are not accepted as operand bindings, but cell/range IDs are retained.
    con = sd.oa._conn(task)
    try:
        cells = {r["cell_id"]: r for r in con.execute("SELECT cell_id,sheet_id,address FROM cells").fetchall()}
        sheets = {r["sheet_id"]: r["name"] for r in con.execute("SELECT sheet_id,name FROM sheets").fetchall()}
        ranges = {r["range_id"]: r for r in con.execute("SELECT range_id,sheet_id,r1,c1,r2,c2 FROM ranges").fetchall()}
        out = {}
        for rid in ids:
            if rid.startswith("cell:") and rid in cells:
                c = cells[rid]; sheet = sheets[c["sheet_id"]]
                out[rid] = (f"'{sheet}'!{c['address']}" if " " in sheet else f"{sheet}!{c['address']}")
            elif rid.startswith("range:") and rid in ranges:
                r = ranges[rid]; sheet = sheets[r["sheet_id"]]
                from openpyxl.utils import get_column_letter
                addr = f"{get_column_letter(r['c1'])}{r['r1']}:{get_column_letter(r['c2'])}{r['r2']}"
                out[rid] = (f"'{sheet}'!{addr}" if " " in sheet else f"{sheet}!{addr}")
        return out
    finally:
        con.close()


def formula_from_ast(ast: Any, bindings: dict[str, str], task: str, working: set[str], home_sheet: str) -> tuple[str | None, str | None]:
    ids = set(bindings.values()) | {x[1] for x in collect_refs(ast)}
    if any(x not in working for x in ids):
        return None, "INVALID_ENTITY_ID"
    lowered = id_to_a1(task, ids)
    try:
        if any(x not in lowered for x in ids): return None, "INVALID_ENTITY_ID"
        # Same-sheet references use the canonical formula spelling used by the
        # frozen workbook and scorer; cross-sheet references retain the sheet.
        local = {k: (v.split("!", 1)[1] if "!" in v and v.split("!", 1)[0].strip("'") == home_sheet else v)
                 for k, v in lowered.items()}
        return "=" + render(ast, bindings, local), None
    except Exception as exc:
        return None, f"ASSEMBLY_FAILURE:{type(exc).__name__}"


def collect_refs(ast: Any) -> list[tuple[str, str]]:
    if ast[0] == "REF": return [ast]
    if ast[0] in {"CONST", "HOLE"}: return []
    if ast[0] == "NEG": return collect_refs(ast[1])
    out = []
    for x in ast[1]: out.extend(collect_refs(x))
    return out


# ------------------------------ population and primary execution

def population_manifest() -> list[dict]:
    out = []
    for task, member in POPULATION:
        u = units()[(task, member)]
        s = session_for(task, member)
        working = sorted(s["working_set_ids"])
        _, base = base_payload(task, member, s)
        out.append({
            "task": task, "canonical_member": member, "canonical_cell": u["canonical_cell"],
            "canonical_cell_id": u["canonical_cell_id"], "obligation_id": u["obligation_id"],
            "operation_id": u["operation_id"], "axis": u["axis"], "members": u["members"],
            "member_cells": u["member_cells"], "availability_class": u["availability_class"],
            "working_set_ids": working, "working_set_sha256": sha(working),
            "evidence_sha256": sha(base["WORKING_SET_EVIDENCE"]),
            "base_payload_sha256": hashlib.sha256(json.dumps(base, ensure_ascii=False, separators=(",", ":")).encode()).hexdigest(),
            "working_set_size": len(working),
        })
    return out


def run_primary() -> dict:
    instr.check(); old.guard(MODEL); old.load_dotenv()
    key = os.environ["OPENROUTER_API_KEY"]
    pop = population_manifest()
    manifest = {
        "model": MODEL, "temperature": TEMPERATURE, "reasoning": REASONING,
        "retries": 0, "max_tokens": MAX_TOKENS, "instrumentation_generation": instr.GENERATION,
        "frozen_sha256": {n: hashlib.sha256((old.ROOT / n).read_bytes()).hexdigest() for n in FROZEN},
        "population": pop, "arms": ["F0", "F1a", "F1b", "F2a", "F2b"],
        "gold_in_primary_prompts": False,
    }
    write(OUT / "freeze.json", manifest)
    rows = []
    for meta in pop:
        task, member = meta["task"], meta["canonical_member"]
        s = session_for(task, member)
        base_user, base = base_payload(task, member, s)
        working = set(s["working_set_ids"])
        prefix = f'{task}__{member.replace("!", "_").replace(" ", "_")}'
        f0 = call_model(prefix + "__F0", "F0", old.SYNTHESIS_SYSTEM, base_user, key)
        f0["working_set_sha256"] = meta["working_set_sha256"]
        f1a_user = stage_payload(base, "Follow the F1a staged protocol in the system message; do not emit a spreadsheet formula.")
        f1a = call_model(prefix + "__F1a", "F1a", F1A_SYSTEM, f1a_user, key)
        sketch_art = f1a.get("parsed_response")
        f1b = None
        if isinstance(sketch_art, dict) and sketch_art.get("status") == "PROPOSED" and sketch_art.get("sketch") is not None:
            f1b_user = stage_payload(base, "Follow the F1b staged protocol in the system message; do not emit a spreadsheet formula.", "MODEL_GENERATED_SKETCH", sketch_art)
            f1b = call_model(prefix + "__F1b", "F1b", F1B_SYSTEM, f1b_user, key, prior=sketch_art)
        f2a_user = stage_payload(base, "Follow the F2a staged protocol in the system message; do not emit a spreadsheet formula.")
        f2a = call_model(prefix + "__F2a", "F2a", F2A_SYSTEM, f2a_user, key)
        operand_art = f2a.get("parsed_response")
        f2b = None
        if isinstance(operand_art, dict) and operand_art.get("status") == "PROPOSED" and operand_art.get("operand_ids") is not None:
            f2b_user = stage_payload(base, "Follow the F2b staged protocol in the system message; do not emit a spreadsheet formula.", "MODEL_SELECTED_OPERAND_SET", operand_art)
            f2b = call_model(prefix + "__F2b", "F2b", F2B_SYSTEM, f2b_user, key, prior=operand_art)
        rows.append({"meta": meta, "F0": f0, "F1a": f1a, "F1b": f1b, "F2a": f2a, "F2b": f2b})
        print(json.dumps({"task": task, "canonical_member": member,
                          "F0": (f0.get("parsed_response") or {}).get("status"),
                          "F1a": (f1a.get("parsed_response") or {}).get("status"),
                          "F1b": (f1b or {}).get("parsed_response", {}).get("status") if f1b else "UPSTREAM_ABSTAIN",
                          "F2a": (f2a.get("parsed_response") or {}).get("status"),
                          "F2b": (f2b or {}).get("parsed_response", {}).get("status") if f2b else "UPSTREAM_ABSTAIN"}), flush=True)
    write(OUT / "primary.json", rows)
    return {"population": pop, "rows": rows}


# ------------------------------ evaluator and workbook consequence

def gold_for(task: str, member: str) -> str:
    rows = sd.audit_rows()
    row = next(r for r in rows if r["task"] == task and r["canonical_member"] == member)
    return row["gold_canonical_formula"]


def gold_refs(task: str, member: str, formula: str) -> list[str]:
    home = member.split("!", 1)[0]
    ops = sd.oa.operands(formula, home)
    con = sd.oa._conn(task)
    try:
        sheets = sd.oa._sheet_ids(con); cells = sd.oa._cell_index(con); out = []
        for op in ops:
            if op["kind"] != "POINT": continue
            sid = sheets.get(op["sheet"]); c = cells.get((sid, op["start"])) if sid else None
            if c: out.append(c["cell_id"])
        return out
    finally: con.close()


def classify_shape(ast: Any, gold: Any) -> str:
    if ast is None: return "SKETCH_ABSTAIN"
    if abstract(ast) == abstract(gold): return "SKETCH_CORRECT"
    if ast[0] != gold[0]: return "SKETCH_WRONG_OPERATOR"
    if ast[0] == "CONST" or gold[0] == "CONST": return "SKETCH_WRONG_CONSTANT"
    if ast[0] in {"HOLE", "REF"} or gold[0] in {"HOLE", "REF"}: return "SKETCH_WRONG_OPERATOR"
    aa, gg = ast[1], gold[1]
    if len(aa) != len(gg): return "SKETCH_WRONG_ARITY"
    return "SKETCH_WRONG_NESTING"


def classify_binding(bindings: Any, ast: Any, expected: list[str]) -> str:
    if not isinstance(bindings, dict): return "BINDING_ABSTAIN"
    holes = ast_holes(ast)
    if any(h not in bindings for h in holes): return "BINDING_MISSING_REFERENCE"
    if any(h not in holes for h in bindings): return "BINDING_EXTRA_REFERENCE"
    got = [bindings[h] for h in holes]
    if got == expected: return "BINDING_CORRECT"
    if set(got) - set(expected): return "BINDING_EXTRA_REFERENCE"
    return "BINDING_WRONG_REFERENCE"


def operand_class(operand_ids: Any, expected: set[str], working: set[str]) -> str:
    if not isinstance(operand_ids, list): return "OPERANDS_ABSTAIN"
    got = set(x for x in operand_ids if isinstance(x, str))
    if not got <= working: return "INVALID_ENTITY_ID"
    if got == expected: return "OPERANDS_EXACT"
    if got < expected: return "OPERANDS_MISSING"
    if got > expected: return "OPERANDS_EXTRA"
    return "OPERANDS_WRONG"


def parse_stage(row: dict) -> dict:
    task, member = row["meta"]["task"], row["meta"]["canonical_member"]
    gold = gold_for(task, member)
    g_ast = gold_ast(gold)
    expected_refs = gold_refs(task, member, gold)
    working = set(row["meta"]["working_set_ids"])
    id_map = id_to_a1(task, working)
    out = {"task": task, "canonical_member": member, "program_key": sd.program_key(next(r for r in sd.audit_rows() if r["task"] == task and r["canonical_member"] == member)), "gold_formula": gold}
    p0 = (row["F0"].get("parsed_response") or {}) if row.get("F0") else {}
    out["F0_formula"] = p0.get("formula") if p0.get("status") == "PROPOSED" else None
    out["F0_exact"] = ES._same(out["F0_formula"], gold)

    p1a = (row["F1a"].get("parsed_response") or {}) if row.get("F1a") else {}
    ast1, err1 = normalize_ast(p1a.get("sketch"), "F1") if p1a.get("status") == "PROPOSED" else (None, "SKETCH_ABSTAIN")
    out["F1a_sketch"] = p1a.get("sketch")
    out["F1a_sketch_error"] = err1
    out["F1a_class"] = classify_shape(ast1, g_ast)
    p1b = (row["F1b"].get("parsed_response") or {}) if row.get("F1b") else {}
    bindings = p1b.get("bindings") if p1b.get("status") == "PROPOSED" else None
    out["F1b_bindings"] = bindings
    out["F1b_bindings_unconditional_class"] = classify_binding(bindings, ast1, expected_refs) if ast1 is not None else "NOT_PARSEABLE"
    out["F1b_class_given_correct_sketch"] = classify_binding(bindings, ast1, expected_refs) if out["F1a_class"] == "SKETCH_CORRECT" else "NOT_CONDITIONED_ON_CORRECT_SKETCH"
    home_sheet = member.split("!", 1)[0]
    f1_formula, f1err = (formula_from_ast(ast1, bindings or {}, task, working, home_sheet) if ast1 is not None and isinstance(bindings, dict) else (None, "ASSEMBLY_FAILURE"))
    out["F1_assembled_formula"] = f1_formula
    out["F1_assembly_error"] = f1err
    out["F1_assembled_exact"] = ES._same(f1_formula, gold)

    p2a = (row["F2a"].get("parsed_response") or {}) if row.get("F2a") else {}
    operand_ids = p2a.get("operand_ids") if p2a.get("status") == "PROPOSED" else None
    expected_set = set(expected_refs)
    out["F2a_operand_ids"] = operand_ids
    out["F2a_class"] = operand_class(operand_ids, expected_set, working)
    p2b = (row["F2b"].get("parsed_response") or {}) if row.get("F2b") else {}
    ast2, err2 = normalize_ast(p2b.get("structure"), "F2", set(operand_ids or [])) if p2b.get("status") == "PROPOSED" else (None, "STRUCTURE_ABSTAIN")
    out["F2b_structure"] = p2b.get("structure")
    out["F2b_structure_error"] = err2
    out["F2b_structure_unconditional_exact"] = bool(ast2 is not None and abstract(ast2) == abstract(g_ast))
    out["F2b_class_given_sufficient_operands"] = ("STRUCTURE_CORRECT" if ast2 is not None and abstract(ast2) == abstract(g_ast) else ("STRUCTURE_WRONG_NESTING" if ast2 is not None else "STRUCTURE_ABSTAIN")) if out["F2a_class"] == "OPERANDS_EXACT" else "NOT_CONDITIONED_ON_SUFFICIENT_OPERANDS"
    f2_formula, f2err = (formula_from_ast(ast2, {}, task, working, home_sheet) if ast2 is not None else (None, "ASSEMBLY_FAILURE"))
    out["F2_assembled_formula"] = f2_formula
    out["F2_assembly_error"] = f2err
    out["F2_assembled_exact"] = ES._same(f2_formula, gold)
    target_addr = member.split("!", 1)[1]
    def mechanical_formula_check(formula):
        if not formula:
            return {"valid": True, "reason": "NO_FORMULA"}
        try:
            from openpyxl.formula import Tokenizer
            tokens = Tokenizer(formula).items
            target = target_addr.replace("$", "").upper()
            for token in tokens:
                value, kind, subtype = token.value, token.type, token.subtype
                if kind == "OPERAND" and subtype == "RANGE":
                    ref = value.replace("$", "").upper()
                    if "!" not in ref and ref == target:
                        return {"valid": False, "reason": "DIRECT_SELF_REFERENCE"}
            return {"valid": True, "reason": "OK"}
        except Exception as exc:
            return {"valid": False, "reason": f"PARSE_FAILURE:{type(exc).__name__}"}
    out["F0_mechanical_validation"] = mechanical_formula_check(out["F0_formula"])
    out["F1_mechanical_validation"] = mechanical_formula_check(f1_formula)
    out["F2_mechanical_validation"] = mechanical_formula_check(f2_formula)
    out["working_set_sha256"] = row["meta"]["working_set_sha256"]
    return out


def end_to_end(task: str, member: str, unit: dict, stage: str, formula: str | None) -> dict:
    name = f"{task}__{member.replace('!', '_').replace(' ', '_')}__{stage}"
    dest_dir = VARIANTS / task / name
    dest = dest_dir / f"{task}_output.xlsx"
    edits = []
    applied = {"formulas": {}, "failures": {}}
    if formula:
        applied = pg.apply_program(formula, unit)
        for addr, f in applied["formulas"].items():
            sheet, address = addr.split("!", 1)
            edits.append({"sheet": sheet, "address": address, "formula": f})
    audit = write_cells(ccf.cc.workbook_paths(task)[0], dest, edits)
    write(dest_dir / "audit.json", {"stage": stage, "edits": edits, "translation": applied, "audit": audit})
    return {"path": str(dest), "edits": edits, "translation_failures": applied.get("failures", {}), "audit": audit}


def batched_scores(task: str, paths: dict[str, Path]) -> dict[str, tuple[dict, dict]]:
    """Official/value-only scoring with input and gold loaded once per task.

    This is the same evaluator comparison and error-value fallback as the
    frozen scorer; only workbook loading is shared across the three staged
    variants.  It keeps the end-to-end consequence measurement from becoming
    an accidental repeated-I/O experiment.
    """
    import openpyxl
    sys.path.insert(0, str(ccf.cc.EVAL_DIR))
    from evaluation import (cell_level_compare_with_classification,
                            classify_cells_by_modification, compare_cell_value,
                            compare_cell_formula, _find_sheet, _has_excel_error,
                            parse_answer_position)
    inp, gold = ccf.cc.workbook_paths(task)
    wi = openpyxl.load_workbook(inp, data_only=True)
    wg = openpyxl.load_workbook(gold, data_only=True)
    wif = openpyxl.load_workbook(inp, data_only=False)
    wgf = openpyxl.load_workbook(gold, data_only=False)
    classes = []
    for spec in parse_answer_position(ccf.cc.dataset()[task]["answer_position"]):
        sn, cr = (spec.split("!", 1) if "!" in spec else (wg.sheetnames[0], spec))
        sn, cr = sn.strip("'").strip(), cr.strip("'").strip()
        reg, mod = classify_cells_by_modification(wi, wg, sn, cr, False, False,
                                                  wb_input_formula=wif, wb_answer_formula=wgf)
        classes.append((sn, reg, mod))
    out = {}
    for name, path in paths.items():
        wo = openpyxl.load_workbook(path, data_only=True)
        wof = openpyxl.load_workbook(path, data_only=False)
        reg_stats = {"correct": 0, "total": 0}
        mod_stats = {"correct": 0, "total": 0}
        value_stats = {"regression": {"correct": 0, "total": 0},
                       "modification": {"correct": 0, "total": 0}}
        fallback = {"regression": 0, "modification": 0}
        fallback_only = {"regression": 0, "modification": 0}
        errors = []
        for sn, reg, mod in classes:
            rc, rt, mc, mt, msgs = cell_level_compare_with_classification(
                wg, wo, sn, reg, mod, False, False,
                wb_answer_formula=wgf, wb_output_formula=wof)
            reg_stats["correct"] += rc; reg_stats["total"] += rt
            mod_stats["correct"] += mc; mod_stats["total"] += mt
            errors.extend(msgs)
            ws_g, ws_o = _find_sheet(wg, sn), _find_sheet(wo, sn)
            ws_gf, ws_of = _find_sheet(wgf, sn), _find_sheet(wof, sn)
            for label, cells in (("regression", reg), ("modification", mod)):
                for cell_name in cells:
                    cg, co = ws_g[cell_name], ws_o[cell_name]
                    by_value = compare_cell_value(cg.value, co.value)
                    value_stats[label]["total"] += 1
                    value_stats[label]["correct"] += bool(by_value)
                    has_error = _has_excel_error(cg) or _has_excel_error(co)
                    if has_error:
                        fallback[label] += 1
                        official = compare_cell_formula(ws_gf[cell_name], ws_of[cell_name])
                        fallback_only[label] += bool(official and not by_value)
        total = reg_stats["total"] + mod_stats["total"]
        correct = reg_stats["correct"] + mod_stats["correct"]
        official = {
            "exact": correct == total,
            "first_error": errors[0] if errors else "",
            "regression": {**reg_stats, "accuracy": round(reg_stats["correct"] / reg_stats["total"], 6) if reg_stats["total"] else None},
            "modification": {**mod_stats, "accuracy": round(mod_stats["correct"] / mod_stats["total"], 6) if mod_stats["total"] else None},
        }
        strict = {}
        for label in ("regression", "modification"):
            v = value_stats[label]
            o = official[label]
            strict[label] = {
                "total": o["total"], "official_correct": o["correct"],
                "value_correct": v["correct"], "took_error_fallback": fallback[label],
                "correct_only_via_error_fallback": fallback_only[label],
                "official_accuracy": o["accuracy"],
                "value_only_accuracy": round(v["correct"] / v["total"], 6) if v["total"] else None,
            }
        out[name] = (official, strict)
        wo.close(); wof.close()
    wi.close(); wg.close(); wif.close(); wgf.close()
    return out


def evaluate() -> dict:
    rows = load(OUT / "primary.json")
    evaluated = [parse_stage(r) for r in rows]
    # Build variants per task, then recalculate in the same batched evaluator
    # used by the frozen execution probe.  No gold formula is written.
    by_task = defaultdict(list)
    for r, e in zip(rows, evaluated):
        task, member = e["task"], e["canonical_member"]
        u = units()[(task, member)]
        for stage, key in (("F0", "F0_formula"), ("F1", "F1_assembled_formula"), ("F2", "F2_assembled_formula")):
            e[f"{stage}_end_to_end"] = end_to_end(task, member, u, stage, e[key])
        by_task[task].extend([e])
    for task in by_task:
        ccf.recalculate(VARIANTS / task)
    for task, task_rows in by_task.items():
        paths = {f"{e['canonical_member']}__{stage}": Path(e[f"{stage}_end_to_end"]["path"])
                 for e in task_rows for stage in ("F0", "F1", "F2")}
        scored = batched_scores(task, paths)
        for e in task_rows:
            for stage in ("F0", "F1", "F2"):
                key = f"{e['canonical_member']}__{stage}"
                e[f"{stage}_score"], e[f"{stage}_strict_score"] = scored[key]
    for e in evaluated:
        for stage in ("F0", "F1", "F2"):
            formulas = e[f"{stage}_end_to_end"]["edits"]
            def member_gold(cell):
                edit = ccf.gold_edit(e["task"], tuple(cell))
                return edit["golden_payload"] if edit else None
            e[f"{stage}_all_member_exact"] = all(
                ES._same(next((x["formula"] for x in formulas
                               if x["sheet"] == c[0] and x["address"] == sd.cc.a1(c[1], c[2])), None),
                         member_gold(c))
                for c in units()[(e["task"], e["canonical_member"])] ["member_cells"]
            )
    # Explicit corruption counters: first-stage artifact correct but final
    # assembled formula wrong.
    for e in evaluated:
        e["F1_staged_composition_loss"] = e["F1a_class"] == "SKETCH_CORRECT" and e["F1_assembled_exact"] is not True
        e["F2_staged_composition_loss"] = e["F2a_class"] == "OPERANDS_EXACT" and e["F2_assembled_exact"] is not True
    write(OUT / "evaluated.json", evaluated)
    return {"rows": evaluated}


def main() -> None:
    if (OUT / "primary.json").exists():
        print("primary.json exists; use evaluate-only by deleting nothing and run report separately")
    else:
        run_primary()
    evaluate()


if __name__ == "__main__":
    main()
