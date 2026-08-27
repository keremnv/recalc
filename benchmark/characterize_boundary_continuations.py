#!/usr/bin/env python3
"""Measure narrow right-edge formula continuations against direct blank targets.

Candidate generation is gold-blind. Golden workbooks are used only to score address
precision and whether the inferred formula itself matches the direct target.
"""

from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from concurrent.futures import ProcessPoolExecutor
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "benchmark-data/SpreadsheetBench-2/evaluation"))
sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(ROOT / "src"))

import evaluation as ev
import openpyxl

from benchmark.boundary_continuations import (
    ARMS,
    enumerate_boundary_continuations,
    select_arm,
)
from benchmark.target_classification import classify_cache_robust_targets
from benchmark.xlsx_metadata_repair import install as _install_repair

_install_repair()


def _arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("category", nargs="?", default="Financial_Model")
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--task", action="append")
    parser.add_argument("--workers", type=int, default=1)
    parser.add_argument("--false-positive-samples", type=int, default=12)
    return parser.parse_args()


def _blank_targets(
    task: dict[str, Any], input_values: Any, answer_values: Any, input_raw: Any, answer_raw: Any
) -> set[tuple[str, str]]:
    targets: set[tuple[str, str]] = set()
    with_color = "Color" in task["spreadsheet_path"]
    with_formula = "Embedded" in task["spreadsheet_path"]
    for chunk in ev.parse_answer_position(task["answer_position"]):
        sheet_name, _, cell_range = chunk.rpartition("!")
        sheet_name = sheet_name.strip().strip("'")
        classified = classify_cache_robust_targets(
            ev,
            input_values,
            answer_values,
            input_raw,
            answer_raw,
            sheet_name,
            cell_range,
            with_font_color=with_color,
            with_formula=with_formula,
        )
        sheet = ev._find_sheet(input_raw, sheet_name)
        if sheet is None:
            continue
        for address in classified.modification:
            value = sheet[address].value
            if value is None or (isinstance(value, str) and not value.strip()):
                targets.add((sheet.title, address))
    return targets


def _formula_matches(inferred: str, golden: Any) -> bool:
    golden = getattr(golden, "text", golden)
    if not isinstance(golden, str) or not golden.startswith("="):
        return False
    return inferred.replace("$", "").upper() == golden.replace("$", "").upper()


def analyse(task: dict[str, Any], data_dir: Path) -> tuple[dict[str, Counter], list[dict]]:
    input_values = openpyxl.load_workbook(data_dir / task["spreadsheet_path"], data_only=True)
    answer_values = openpyxl.load_workbook(data_dir / task["golden_response_path"], data_only=True)
    input_raw = openpyxl.load_workbook(data_dir / task["spreadsheet_path"], data_only=False)
    answer_raw = openpyxl.load_workbook(data_dir / task["golden_response_path"], data_only=False)
    try:
        truth = _blank_targets(task, input_values, answer_values, input_raw, answer_raw)
        candidates = enumerate_boundary_continuations(
            input_raw, input_values, instruction=task["instruction"]
        )
        results: dict[str, Counter] = {}
        details: list[dict] = []
        for arm in ARMS:
            selected = select_arm(candidates, arm)
            counts = Counter(
                {
                    "tasks": 1,
                    "blank_targets": len(truth),
                    "pool": len(selected),
                    "tasks_with_candidates": int(bool(selected)),
                }
            )
            for candidate in selected:
                key = (candidate.sheet, candidate.address)
                hit = key in truth
                golden_sheet = ev._find_sheet(answer_raw, candidate.sheet)
                formula_hit = bool(
                    hit
                    and golden_sheet is not None
                    and _formula_matches(
                        candidate.inferred_formula, golden_sheet[candidate.address].value
                    )
                )
                counts["hits"] += int(hit)
                counts["formula_hits"] += int(formula_hit)
                if key == ("Working Capital Schedule", "M3"):
                    counts["watch:M3:present"] += 1
                    counts["watch:M3:hit"] += int(hit)
                    counts["watch:M3:formula_hit"] += int(formula_hit)
                if hit:
                    counts["tasks_with_hit"] = 1
                # Only this arm is rendered below. Keeping every candidate from every
                # arm retains hundreds of thousands of duplicate dictionaries on the
                # full category without changing any reported aggregate.
                if arm == "date+style":
                    details.append(
                        {
                            "arm": arm,
                            "task": task["id"],
                            "sheet": candidate.sheet,
                            "address": candidate.address,
                            "run": f"{candidate.run_start}:{candidate.run_end}",
                            "run_length": candidate.run_length,
                            "style_match": candidate.style_match,
                            "date_like": candidate.date_like,
                            "named_sheet": candidate.named_sheet,
                            "hit": hit,
                            "formula_hit": formula_hit,
                            "inferred_formula": candidate.inferred_formula,
                        }
                    )
            results[arm] = counts
        return results, details
    finally:
        for workbook in (input_values, answer_values, input_raw, answer_raw):
            workbook.close()


def _analyse_entry(
    payload: tuple[dict[str, Any], Path],
) -> tuple[str, dict[str, Counter] | None, list[dict], str | None]:
    task, data_dir = payload
    try:
        results, details = analyse(task, data_dir)
    except Exception as exc:  # noqa: BLE001
        return task["id"], None, [], f"{type(exc).__name__}: {str(exc)[:100]}"
    return task["id"], results, details, None


def _report(category: str, measured: int, totals: dict[str, Counter]) -> None:
    print(f"\n=== {category}: {measured} tasks measured ===")
    print(
        f"{'arm':12} {'pool':>8} {'hits':>7} {'formula':>8} {'precision':>10} "
        f"{'target-share':>12} {'hit-tasks':>10}"
    )
    for arm in ARMS:
        counts = totals[arm]
        pool, hits = counts["pool"], counts["hits"]
        print(
            f"{arm:12} {pool:8,} {hits:7,} {counts['formula_hits']:8,} "
            f"{hits / pool if pool else 0:10.1%} "
            f"{hits / counts['blank_targets'] if counts['blank_targets'] else 0:12.1%} "
            f"{counts['tasks_with_hit']:3}/{counts['tasks']:<6}"
        )
    print("\nM3 survival:")
    for arm in ARMS:
        counts = totals[arm]
        print(
            f"  {arm:12} address {counts['watch:M3:hit']}/{counts['watch:M3:present']} "
            f"formula {counts['watch:M3:formula_hit']}/{counts['watch:M3:present']}"
        )


def main() -> int:
    args = _arguments()
    data_dir = ROOT / "benchmark-data/SpreadsheetBench-2/data" / args.category
    tasks = json.loads((data_dir / "dataset.json").read_text())
    if args.task:
        selected = set(args.task)
        tasks = [task for task in tasks if task["id"] in selected]
    if args.limit:
        tasks = tasks[: args.limit]
    totals = {arm: Counter() for arm in ARMS}
    details: list[dict] = []
    measured = 0
    payloads = [(task, data_dir) for task in tasks]
    if args.workers == 1:
        entries = map(_analyse_entry, payloads)
    else:
        # A large task can peak near 500 MB while four workbook variants are open.
        # openpyxl does not reliably return all arenas to a long-lived worker, so
        # recycle after every task to keep a full-category sweep bounded.
        executor = ProcessPoolExecutor(max_workers=args.workers, max_tasks_per_child=1)
        entries = executor.map(_analyse_entry, payloads)
    try:
        for task_id, results, task_details, error in entries:
            if error is not None or results is None:
                print(f"{task_id:8} ERROR {error}", file=sys.stderr)
                continue
            measured += 1
            for arm, counts in results.items():
                totals[arm].update(counts)
            details.extend(task_details)
            date_style = results["date+style"]
            print(
                f"{task_id:8} pool={date_style['pool']:4} hits={date_style['hits']:3} "
                f"formula={date_style['formula_hits']:3}",
                flush=True,
            )
    finally:
        if args.workers != 1:
            executor.shutdown()
    _report(args.category, measured, totals)
    true_positives = [item for item in details if item["hit"]]
    if true_positives:
        print("\ndate+style true positives:")
        for item in true_positives:
            print(
                f"  {item['task']} {item['sheet']}!{item['address']} "
                f"run={item['run']} formula_hit={item['formula_hit']} "
                f"inferred={item['inferred_formula']}"
            )
    false_positives = [item for item in details if item["arm"] == "date+style" and not item["hit"]]
    if false_positives:
        print("\ndate+style false-positive representatives:")
        for item in false_positives[: args.false_positive_samples]:
            print(
                f"  {item['task']} {item['sheet']}!{item['address']} "
                f"run={item['run']} inferred={item['inferred_formula']}"
            )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
