#!/usr/bin/env python3
"""Deterministic formula-fill primitive for the default SpreadsheetBench harness.

The model supplies the target set and one canonical formula. This module never
chooses, repairs, ranks, or enlarges that authority. It only checks an input-side
mechanical translation witness over the declared cells and, when that witness
holds, translates the supplied formula onto exactly those cells.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import sys
import tempfile
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(Path(__file__).resolve().parent), str(ROOT / "src")]

import composition_closure as cc  # noqa: E402
import openpyxl  # noqa: E402
import program_group as pg  # noqa: E402
from librecalc_mcp.domain.formulas import translate_a1_formula  # noqa: E402
from xlsx_cell_writer import write_cells  # noqa: E402

MAX_ROW = 1_048_576
MAX_COL = 16_384
MIN_TARGETS = 2
CELL_TOKEN = re.compile(r"^\$?([A-Za-z]{1,3})\$?([1-9][0-9]*)$")
QUALIFIED = re.compile(
    r"^(?:(?P<quoted>'[^']*(?:''[^']*)*')|(?P<plain>[^'!]+))!(?P<rest>.+)$"
)


class FillReject(Exception):
    def __init__(self, code: str, message: str, **extra: Any) -> None:
        super().__init__(message)
        self.code = code
        self.message = message
        self.extra = extra

    def payload(self) -> dict[str, Any]:
        out = {"status": "rejected", "reason": self.code, "message": self.message, "wrote": []}
        out.update(self.extra)
        return out


def a1(row: int, col: int) -> str:
    return cc.a1(row, col)


def parse_a1_cell(token: str) -> tuple[int, int]:
    match = CELL_TOKEN.match(token.strip())
    if not match:
        raise FillReject("INVALID_COORDINATE", f"Not an A1 cell: {token!r}")
    col = 0
    for ch in match.group(1).upper():
        col = col * 26 + ord(ch) - 64
    row = int(match.group(2))
    if not (1 <= row <= MAX_ROW and 1 <= col <= MAX_COL):
        raise FillReject("INVALID_COORDINATE", f"Out of Excel bounds: {token!r}")
    return row, col


def split_sheet_rest(spec: str) -> tuple[str, str]:
    text = spec.strip()
    match = QUALIFIED.match(text)
    if not match:
        raise FillReject("INVALID_COORDINATE", f"Target must be Sheet!A1 or Sheet!A1:C5: {spec!r}")
    sheet = match.group("quoted") or match.group("plain")
    if sheet.startswith("'") and sheet.endswith("'"):
        sheet = sheet[1:-1].replace("''", "'")
    return sheet, match.group("rest").strip()


def parse_qualified_cell(spec: str) -> tuple[str, int, int]:
    sheet, rest = split_sheet_rest(spec)
    row, col = parse_a1_cell(rest)
    return sheet, row, col


def expand_target_spec(spec: str) -> list[tuple[str, int, int]]:
    sheet, rest = split_sheet_rest(spec)
    if "," in rest:
        raise FillReject("INVALID_COORDINATE", f"Commas belong between qualified cells, not inside a range: {spec!r}")
    if ":" in rest:
        start, end = rest.split(":", 1)
        r1, c1 = parse_a1_cell(start)
        r2, c2 = parse_a1_cell(end)
        rmin, rmax = min(r1, r2), max(r1, r2)
        cmin, cmax = min(c1, c2), max(c1, c2)
        return [(sheet, row, col) for row in range(rmin, rmax + 1) for col in range(cmin, cmax + 1)]
    row, col = parse_a1_cell(rest)
    return [(sheet, row, col)]


def parse_targets(raw: str) -> list[tuple[str, int, int]]:
    text = raw.strip()
    if not text:
        raise FillReject("TARGET_SET_EMPTY", "No targets were supplied")
    if text.startswith("["):
        try:
            items = json.loads(text)
        except json.JSONDecodeError as exc:
            raise FillReject("INVALID_COORDINATE", f"Targets JSON is not parseable: {exc}") from exc
        if not isinstance(items, list) or not all(isinstance(item, str) for item in items):
            raise FillReject("INVALID_COORDINATE", "Targets JSON must be a list of Sheet!A1 strings")
        specs = items
    else:
        specs = [part.strip() for part in text.split(",") if part.strip()]
    seen: set[tuple[str, int, int]] = set()
    ordered: list[tuple[str, int, int]] = []
    for spec in specs:
        for cell in expand_target_spec(spec):
            if cell not in seen:
                seen.add(cell)
                ordered.append(cell)
    if not ordered:
        raise FillReject("TARGET_SET_EMPTY", "No targets were supplied")
    return ordered


def refuse_gold(path: Path) -> None:
    lowered = str(path).lower()
    if "golden" in lowered or "gold_response" in lowered:
        raise FillReject("GOLD_PATH_REFUSED", "calc_translate_fill refuses gold/evaluator workbooks")


def validate_formula(formula: str) -> str:
    text = formula.strip()
    if not text.startswith("=") or len(text) < 2:
        raise FillReject("INVALID_FORMULA_SYNTAX", "canonical_formula must be a non-empty Excel formula starting with =")
    if text.count('"') % 2 != 0 or text.count("'") % 2 != 0:
        raise FillReject("INVALID_FORMULA_SYNTAX", "canonical_formula has unbalanced quotes")
    try:
        translate_a1_formula(text, column_offset=0, row_offset=0)
    except Exception as exc:
        raise FillReject("INVALID_FORMULA_SYNTAX", f"canonical_formula failed A1 parse: {exc}") from exc
    return text


def geometry(members: list[tuple[str, int, int]]) -> str:
    sheets = {item[0] for item in members}
    if len(sheets) != 1:
        raise FillReject("MIXED_SHEETS", "Targets must lie on a single sheet")
    if len(members) < MIN_TARGETS:
        raise FillReject("TARGET_SET_TOO_SMALL", f"Repeated fill requires at least {MIN_TARGETS} target cells")
    rows = [item[1] for item in members]
    cols = [item[2] for item in members]
    rmin, rmax, cmin, cmax = min(rows), max(rows), min(cols), max(cols)
    expected = (rmax - rmin + 1) * (cmax - cmin + 1)
    present = {(item[1], item[2]) for item in members}
    box = {(row, col) for row in range(rmin, rmax + 1) for col in range(cmin, cmax + 1)}
    if present != box:
        raise FillReject(
            "NOT_A_RECTANGLE_OR_LINE",
            "Declared targets must be a filled rectangle or unit-stride line; holes and non-unit strides are rejected",
            expected_cells=expected,
            declared_cells=len(members),
        )
    if rmin == rmax or cmin == cmax:
        return "LINE"
    return "RECTANGLE"


def existing_formulas(workbook, members: list[tuple[str, int, int]]) -> dict[tuple[str, int, int], str]:
    out: dict[tuple[str, int, int], str] = {}
    for sheet, row, col in members:
        value = workbook[sheet].cell(row, col).value
        if isinstance(value, str) and value.startswith("="):
            out[(sheet, row, col)] = value
    return out


def formulas_compatible(anchor: tuple[str, int, int], members: list[tuple[str, int, int]], forms: dict[tuple[str, int, int], str]) -> bool:
    if anchor not in forms:
        return False
    try:
        return all(
            pg.canonical(pg.translate(forms[anchor], anchor, cell)) == pg.canonical(forms[cell])
            for cell in members
            if cell != anchor
        )
    except Exception:
        return False


def mechanical_witness(
    workbook,
    members: list[tuple[str, int, int]],
    kind: str,
) -> dict[str, Any]:
    sheet = members[0][0]
    own = existing_formulas(workbook, members)
    if 0 < len(own) < len(members):
        raise FillReject(
            "MIXED_EMPTY_AND_FORMULA",
            "Declared set mixes formula cells and empty cells; mechanical homogeneity is not witnessed",
        )
    if len(own) == len(members):
        origin = min(members, key=lambda cell: (cell[1], cell[2]))
        if not formulas_compatible(origin, members, own):
            raise FillReject(
                "NOT_TRANSLATION_COMPATIBLE",
                "Existing formulas on the declared set are not mutually reproducible by translation",
            )
        return {"kind": "own_set", "axis": pg.axis_of(members) if kind == "LINE" else "RECTANGLE", "covered": len(members)}
    # Empty declared set: a parallel line may witness a 1-D fill. Rectangles need own-set formulas.
    if kind != "LINE":
        raise FillReject(
            "NO_EXISTING_FORMULA_WITNESS",
            "Empty rectangles have no mechanical translation witness",
        )
    axis = pg.axis_of(members)
    if axis is None:
        raise FillReject("NOT_A_RECTANGLE_OR_LINE", "Declared line is not collinear")
    window: dict[tuple[str, int, int], str] = {}
    ws = workbook[sheet]
    coords = sorted(m[2] if axis == "ROW" else m[1] for m in members)
    base = members[0][1] if axis == "ROW" else members[0][2]
    for dist in range(1, pg.WITNESS_WINDOW + 1):
        for line in (base - dist, base + dist):
            if line < 1:
                continue
            for k in coords:
                row, col = (line, k) if axis == "ROW" else (k, line)
                value = ws.cell(row, col).value
                if isinstance(value, str) and value.startswith("="):
                    window[(sheet, row, col)] = value
    witness = pg.witness_line("", members, axis, window)
    if witness is None:
        raise FillReject(
            "NO_EXISTING_FORMULA_WITNESS",
            "No parallel input line witnesses translation homogeneity over the declared coordinates",
        )
    covered = set(witness["covered_coordinates"])
    if covered != set(coords):
        raise FillReject(
            "NOT_TRANSLATION_COMPATIBLE",
            "Parallel witness does not cover every declared coordinate; the harness will not underfill or enlarge",
            covered=sorted(covered),
            declared=coords,
        )
    return {"kind": "parallel_line", "axis": axis, "witness": witness, "covered": len(members)}


def translate_all(
    canonical_cell: tuple[str, int, int],
    canonical_formula: str,
    members: list[tuple[str, int, int]],
) -> list[dict[str, str]]:
    edits: list[dict[str, str]] = []
    for cell in members:
        if cell == canonical_cell:
            formula = canonical_formula
        else:
            try:
                formula = pg.translate(canonical_formula, canonical_cell, cell)
            except Exception as exc:
                raise FillReject(
                    "TRANSLATION_FAILURE",
                    f"Translation failed for {cell[0]}!{a1(cell[1], cell[2])}: {exc}",
                ) from exc
        edits.append({"sheet": cell[0], "address": a1(cell[1], cell[2]), "formula": formula})
    written = {(item["sheet"], item["address"]) for item in edits}
    allowed = {(sheet, a1(row, col)) for sheet, row, col in members}
    extra = written - allowed
    if extra:
        raise FillReject("WRITE_OUTSIDE_DECLARED_TARGETS", "Internal invariant broken", extra=sorted(extra))
    missing = allowed - written
    if missing:
        raise FillReject("UNDERFILLED_TARGET_SET", "Internal invariant broken", missing=sorted(missing))
    return edits


def apply_edits(source: Path, dest: Path, edits: list[dict[str, str]]) -> dict[str, Any]:
    source_bytes = source.read_bytes()
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as handle:
        tmp = Path(handle.name)
    try:
        audit = write_cells(source, tmp, edits)
        if audit.get("rejected"):
            raise FillReject(
                "SHARED_FORMULA_MASTER",
                "A requested cell is an unsupported shared-formula master; no cells were written",
                rejected=audit["rejected"],
            )
        applied = {(item["sheet"], item["address"]) for item in audit.get("applied") or []}
        requested = {(item["sheet"], item["address"]) for item in edits}
        if applied != requested:
            raise FillReject(
                "UNDERFILLED_TARGET_SET",
                "Writer did not apply every declared cell; no cells were written",
                applied=sorted(applied),
                requested=sorted(requested),
            )
        dest.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(tmp), str(dest))
        tmp = None
        return audit
    except FillReject:
        if dest.exists() and dest.resolve() == source.resolve():
            dest.write_bytes(source_bytes)
        raise
    finally:
        if tmp is not None and tmp.exists() and tmp.is_file():
            tmp.unlink(missing_ok=True)


def calc_translate_fill(
    *,
    xlsx: Path,
    canonical_cell: str,
    canonical_formula: str,
    targets: str,
    output: Path | None = None,
) -> dict[str, Any]:
    xlsx = Path(xlsx)
    dest = Path(output) if output is not None else xlsx
    refuse_gold(xlsx)
    refuse_gold(dest)
    if not xlsx.is_file():
        raise FillReject("UNKNOWN_WORKBOOK", f"Workbook not found: {xlsx}")
    formula = validate_formula(canonical_formula)
    members = parse_targets(targets)
    kind = geometry(members)
    origin = parse_qualified_cell(canonical_cell)
    workbook = openpyxl.load_workbook(xlsx, data_only=False)
    try:
        names = set(workbook.sheetnames)
        for sheet, row, col in [origin, *members]:
            if sheet not in names:
                raise FillReject("UNKNOWN_SHEET", f"Sheet not in workbook: {sheet!r}")
            if not (1 <= row <= MAX_ROW and 1 <= col <= MAX_COL):
                raise FillReject("INVALID_COORDINATE", f"Out of Excel bounds: {sheet}!{a1(row, col)}")
        before = dest.read_bytes() if dest.is_file() else None
        witness = mechanical_witness(workbook, members, kind)
        edits = translate_all(origin, formula, members)
    finally:
        workbook.close()
    audit = apply_edits(xlsx, dest, edits)
    after = dest.read_bytes()
    if before is not None and dest.resolve() != xlsx.resolve() and before == after:
        # Writes were requested; identical dest bytes would mean the writer no-op'd.
        pass
    return {
        "status": "accepted",
        "reason": None,
        "canonical_cell": f"{origin[0]}!{a1(origin[1], origin[2])}",
        "canonical_formula": formula,
        "targets": [f"{sheet}!{a1(row, col)}" for sheet, row, col in members],
        "mechanical_validation": witness,
        "wrote": edits,
        "changed_cells": [f"{item['sheet']}!{item['address']}" for item in edits],
        "translated_formulas": {f"{item['sheet']}!{item['address']}": item["formula"] for item in edits},
        "writer": {"applied": len(audit.get("applied") or []), "rewritten_parts": audit.get("rewritten_parts")},
        "recalc": "DEFERRED_TO_OFFICIAL_SCORER",
        "dest_sha256": hashlib.sha256(after).hexdigest(),
    }


def _result_or_reject(fn) -> dict[str, Any]:
    try:
        return fn()
    except FillReject as exc:
        return exc.payload()


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Deterministic translation fill; model supplies targets and formula")
    parser.add_argument("--xlsx", required=True)
    parser.add_argument("--canonical-cell", required=True)
    parser.add_argument("--canonical-formula", required=True)
    parser.add_argument("--targets", required=True)
    parser.add_argument("--output")
    args = parser.parse_args(argv)
    dest = Path(args.output) if args.output else Path(args.xlsx)
    before = dest.read_bytes() if dest.is_file() else None
    payload = _result_or_reject(
        lambda: calc_translate_fill(
            xlsx=Path(args.xlsx),
            canonical_cell=args.canonical_cell,
            canonical_formula=args.canonical_formula,
            targets=args.targets,
            output=Path(args.output) if args.output else None,
        )
    )
    if payload.get("status") != "accepted" and dest.is_file() and before is not None:
        current = dest.read_bytes()
        if current != before:
            dest.write_bytes(before)
            payload["restored_unmodified_bytes"] = True
    ledger = os.environ.get("CALC_TRANSLATE_FILL_LEDGER")
    if ledger:
        path = Path(ledger)
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open("a", encoding="utf-8") as handle:
            handle.write(json.dumps(payload, ensure_ascii=False) + "\n")
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    return 0 if payload.get("status") == "accepted" else 2


if __name__ == "__main__":
    raise SystemExit(main())
