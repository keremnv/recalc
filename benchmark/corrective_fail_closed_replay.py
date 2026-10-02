"""Only the two identified parser cases. No model calls, scoring or benchmark."""
from __future__ import annotations

import hashlib
import json
import shutil
import sys
import tempfile
from pathlib import Path
from unittest.mock import patch

import openpyxl

from benchmark.candidate_a_live_runtime import install
from benchmark.inspection_helpers import index
from benchmark.representative_checkpoint import prepare_substrate, task_source

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "research/history/final_architecture_freeze"
REAL = openpyxl.load_workbook


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def outcome(loader, path):
    try:
        wb = loader(path)
        result = {"status": "OPENED", "sheetnames": wb.sheetnames}
        wb.close()
        return result
    except Exception as exc:  # noqa: BLE001 - records exact reference exception for parity
        return {"status": "REFERENCE_EXCEPTION", "class": type(exc).__name__, "message": str(exc)}


def replay(task, original_record, inject=False):
    source = task_source(task)
    with tempfile.TemporaryDirectory(prefix="corrective_fail_closed_") as temp:
        work = Path(temp)
        path = work / "input.xlsx"
        shutil.copy2(source, path)
        index.reset()
        if inject:
            error = UnicodeDecodeError("utf-8", b"x" * 4133 + b"\x8d", 4133, 4134, "invalid start byte")
            with patch.object(index, "_open_db", side_effect=error):
                prepared = prepare_substrate(work, work / "shared", "CORRECTIVE_FAIL_CLOSED_REPLAY", ["CANDIDATE_A"])
        else:
            prepared = prepare_substrate(work, work / "shared", "CORRECTIVE_FAIL_CLOSED_REPLAY", ["CANDIDATE_A"])
        reference = outcome(REAL, path)
        events = []
        with patch.dict("os.environ", {"CANDIDATE_A_SUBSTRATE_MANIFEST": str(work / "shared/manifest.json")}), patch("benchmark.candidate_a_live_runtime.emit", events.append), patch.object(openpyxl, "load_workbook", REAL):
            install()
            treatment = outcome(openpyxl.load_workbook, path)
        assert reference == treatment
        if prepared["failures"]:
            assert not prepared["workbooks"]
            assert not any(e.get("status") == "ACCELERATED" for e in events)
            assert any(e.get("status") == "PREDECLARED_FALLBACK" for e in events)
        result = {"task_id": task, "label": "CORRECTIVE_FAIL_CLOSED_REPLAY",
                  "kind": "INJECTED_EXCEPTION_BOUNDARY_CHECK" if inject else "EXACT_SOURCE_INPUT_MECHANICAL_REPLAY",
                  "input": str(source.relative_to(ROOT.parent)) if source.is_relative_to(ROOT.parent) else str(source),
                  "input_sha256": sha(source), "original_record": original_record,
                  "original_record_sha256": sha(ROOT / original_record),
                  "original_error": json.loads((ROOT / original_record).read_text())["error"],
                  "H0": reference, "H1": treatment, "reference_equal": True,
                  "runner_preparation_proceeded": True, "published_workbooks": len(prepared["workbooks"]),
                  "fallback_occurred": bool(prepared["failures"]),
                  "preparation_failures": prepared["failures"], "candidate_events": events}
        # Temporary paths describe execution identity, not persistent artifacts.
        index.reset()
        return result


def main():
    preserved = [p for p in (ROOT / "research/history/representative_architecture_checkpoint").rglob("*")
                 if p.is_file() and p.suffix in {".json", ".jsonl", ".md"}]
    before = {str(p.relative_to(ROOT)): sha(p) for p in preserved}
    det = "research/history/representative_architecture_checkpoint/reps/Financial_Model_06_01_H1_primary_25/run_record.json"
    trans = "research/history/representative_architecture_checkpoint/reps_preserved_infra_2/Template_06_08_H1_primary_10/run_record.json"
    cases = [replay("Financial_Model:06_01", det), replay("Template:06_08", trans), replay("Template:06_08", trans, inject=True)]
    unchanged = all(sha(ROOT / p) == digest for p, digest in before.items())
    assert unchanged
    OUT.mkdir(exist_ok=True)
    (OUT / "preserved_checkpoint_hashes.json").write_text(json.dumps(before, indent=2) + "\n")
    result = {"label": "CORRECTIVE_FAIL_CLOSED_REPLAY", "model_calls": 0, "scoring_runs": 0,
              "new_treatment_evidence": False, "python": sys.version, "openpyxl": openpyxl.__version__, "lxml": openpyxl.LXML, "primary_results_preserved": unchanged,
              "preserved_artifacts": len(before), "cases": cases,
              "transient_interpretation": "Exact source input did not reproduce the archived UnicodeDecodeError. Injected exception tests the boundary only; origin/stack of the historical transient is not established.",
              "verdict": "PASS_REFERENCE_PARITY_NO_SUBSTRATE_VETO",
              "scope": "Mechanical runner preparation and load boundary only; no new model trajectory or success/score claim."}
    (OUT / "corrective_replay.json").write_text(json.dumps(result, indent=2) + "\n")
    print(json.dumps({"verdict": result["verdict"], "cases": [{"task": c["task_id"], "kind": c["kind"], "fallback": c["fallback_occurred"], "H0": c["H0"]["status"], "H1": c["H1"]["status"]} for c in cases]}))


if __name__ == "__main__":
    main()
