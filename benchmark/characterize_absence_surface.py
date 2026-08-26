#!/usr/bin/env python3
"""How much of the required work is *absent* state, and how much of it is reachable?

Section 9.1 says absence is data. This measures the surface: across a whole category,
what fraction of the cells an agent must modify are blank in the input, and of those,
how many could be nominated by a signal that generalises from populated peers.

Classification of each blank target:
  row-peer     a cell in the same row within 8 columns holds a formula
  col-peer     a cell in the same column within 8 rows holds a formula
  block-peer   both of the above
  isolated     neither -- no peer-based signal can reach it

Valuation!G59 (the recurring 09_04 miss) is block-peer, not isolated: H59 and I59 carry
formulas along the row and G60/G61 along the column. Reaching it was never a detection
problem. Working Capital Schedule!G4 (01_03) is col-peer -- its whole row is empty --
which is why only the carry-chain bridge nominates it.

Aggregates over a whole category are dominated by tasks whose entire answer region is
blank ("build this sheet"), where absence carries no signal because everything is absent.
Those are reported separately: the topology mix inverts between the two populations.

Ground truth is the official evaluator's own classify_cells_by_modification, restricted
to each task's answer_position. No model calls.
"""

from __future__ import annotations

import json
import statistics
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "benchmark-data/SpreadsheetBench-2/evaluation"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import evaluation as ev
import openpyxl
from xlsx_metadata_repair import install as _install_repair

_install_repair()

PEER_DISTANCE = 8


def _norm(v):
    return getattr(v, "text", None) or v


def _is_formula(v) -> bool:
    return isinstance(v, str) and v.startswith("=")


def _split(addr: str) -> tuple[int, int]:
    letters = "".join(c for c in addr if c.isalpha())
    digits = "".join(c for c in addr if c.isdigit())
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n, int(digits)


def _col(n: int) -> str:
    out = ""
    while n:
        n, rem = divmod(n - 1, 26)
        out = chr(65 + rem) + out
    return out


def analyse(task: dict, data_dir: Path) -> Counter:
    inp = data_dir / task["spreadsheet_path"]
    gold = data_dir / task["golden_response_path"]
    with_color = "Color" in task["spreadsheet_path"]
    with_formula = "Embedded" in task["spreadsheet_path"]

    wb_i = openpyxl.load_workbook(inp, data_only=not with_formula)
    wb_a = openpyxl.load_workbook(gold, data_only=not with_formula)
    wb_if = openpyxl.load_workbook(inp, data_only=False) if not with_formula else None
    wb_af = openpyxl.load_workbook(gold, data_only=False) if not with_formula else None
    formulas = openpyxl.load_workbook(inp, data_only=False)

    counts: Counter = Counter()
    for chunk in ev.parse_answer_position(task["answer_position"]):
        sheet_name, _, rng = chunk.rpartition("!")
        sheet_name = sheet_name.strip().strip("'")
        _reg, mod = ev.classify_cells_by_modification(
            wb_i, wb_a, sheet_name, rng, with_color, with_formula, wb_if, wb_af
        )
        if not mod:
            continue
        ws_i = ev._find_sheet(wb_i, sheet_name)
        ws_f = ev._find_sheet(formulas, sheet_name)
        if ws_i is None or ws_f is None:
            continue
        for cell in mod:
            counts["targets"] += 1
            # A cell is present if it holds anything at all -- a literal, or a formula that
            # was never calculated. Many of these workbooks ship with no cached results, so
            # testing the data_only workbook alone reports live formulas as absent and
            # inflates the blank count by an order of magnitude.
            if ws_i[cell].value is not None or ws_f[cell].value is not None:
                counts["populated"] += 1
                continue
            counts["blank"] += 1
            c, r = _split(cell)
            row_peer = any(
                _is_formula(_norm(ws_f[f"{_col(cc)}{r}"].value))
                for cc in range(max(1, c - PEER_DISTANCE), c + PEER_DISTANCE + 1)
                if cc != c
            )
            col_peer = any(
                _is_formula(_norm(ws_f[f"{_col(c)}{rr}"].value))
                for rr in range(max(1, r - PEER_DISTANCE), r + PEER_DISTANCE + 1)
                if rr != r
            )
            if row_peer and col_peer:
                counts["block-peer"] += 1
            elif row_peer:
                counts["row-peer"] += 1
            elif col_peer:
                counts["col-peer"] += 1
            else:
                counts["isolated"] += 1
    return counts


def main() -> int:
    category = sys.argv[1] if len(sys.argv) > 1 else "Financial_Model"
    limit = int(sys.argv[2]) if len(sys.argv) > 2 else 0
    data_dir = ROOT / "benchmark-data/SpreadsheetBench-2/data" / category
    tasks = json.loads((data_dir / "dataset.json").read_text())
    if limit:
        tasks = tasks[:limit]
    per_task: list[Counter] = []
    for task in tasks:
        try:
            counts = analyse(task, data_dir)
        except Exception as exc:  # noqa: BLE001
            print(f"{task['id']:8} ERROR {type(exc).__name__}: {str(exc)[:50]}", file=sys.stderr)
            continue
        per_task.append(counts)
        if counts["blank"]:
            print(
                f"{task['id']:8} targets={counts['targets']:6} blank={counts['blank']:6} "
                f"block={counts['block-peer']:5} row={counts['row-peer']:4} "
                f"col={counts['col-peer']:4} isolated={counts['isolated']:4}"
            )
    _report(f"{category}: all {len(per_task)} tasks measured", per_task)
    partial = [c for c in per_task if c["blank"] < c["targets"]]
    whole = [c for c in per_task if c["blank"] == c["targets"]]
    if whole and partial:
        print(
            f"\n{len(whole)} task(s) have an entirely blank answer region "
            f"({sum(c['targets'] for c in whole):,} targets). Absence carries no signal "
            f"there -- every target is absent -- so they are excluded below."
        )
        _report(f"{category}: {len(partial)} tasks with a partially populated region", partial)
    return 0


def _report(title: str, per_task: list[Counter]) -> None:
    total: Counter = Counter()
    for counts in per_task:
        total.update(counts)
    t, b = total["targets"], total["blank"]
    print(f"\n=== {title} ===")
    print(f"modification targets        {t:8,}")
    print(
        f"  already populated         {total['populated']:8,}  ({total['populated'] / max(1, t):.1%})"
    )
    print(f"  BLANK in input            {b:8,}  ({b / max(1, t):.1%})")
    for key in ("block-peer", "row-peer", "col-peer", "isolated"):
        print(f"    {key:22} {total[key]:8,}  ({total[key] / max(1, b):.1%} of blanks)")
    # Pooled shares follow whichever task has the most cells. Per-task medians say what
    # a typical task looks like, which is what a detector actually faces.
    blank_fr = [c["blank"] / c["targets"] for c in per_task if c["targets"]]
    if blank_fr:
        print(f"  per-task blank fraction   median {statistics.median(blank_fr):7.1%}")
    for key in ("block-peer", "row-peer", "col-peer", "isolated"):
        shares = [c[key] / c["blank"] for c in per_task if c["blank"]]
        if shares:
            print(f"    {key:22} median {statistics.median(shares):7.1%} of a task's blanks")


if __name__ == "__main__":
    raise SystemExit(main())
