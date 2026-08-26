#!/usr/bin/env python3
"""Can gold-blind blank signals form a smaller pool without destroying recall?

This does not ship a detector. It scores named shortlist arms against cache-robust
direct blank targets from the official evaluator, then reports pool size, recall,
precision, and whether the known 09_04 / 01_03 cells survive each cut.

Candidate generation uses the input workbook and instruction only.
"""

from __future__ import annotations

import json
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "benchmark-data/SpreadsheetBench-2/evaluation"))
sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(ROOT / "src"))

import evaluation as ev
import openpyxl
from blank_ranking import ARMS, enumerate_blank_candidates, shortlist
from target_classification import classify_cache_robust_targets
from xlsx_metadata_repair import install as _install_repair

_install_repair()

WATCH = {
    ("Valuation", "G59"),
    ("Working Capital Schedule", "G4"),
    ("Working Capital Schedule", "M3"),
}


def _blank_targets(task: dict, input_values, answer_values, input_raw, answer_raw) -> set[tuple[str, str]]:
    targets: set[tuple[str, str]] = set()
    for chunk in ev.parse_answer_position(task["answer_position"]):
        sheet_name, _, cell_range = chunk.rpartition("!")
        sheet_name = sheet_name.strip().strip("'")
        classified = classify_cache_robust_targets(
            ev,
            input_values,
            answer_values,
            input_raw,
            answer_raw,
            sheet_name,
            cell_range,
        )
        raw_sheet = ev._find_sheet(input_raw, sheet_name)
        if raw_sheet is None:
            continue
        for address in classified.modification:
            value = raw_sheet[address].value
            if value is None or (isinstance(value, str) and not value.strip()):
                targets.add((raw_sheet.title, address))
    return targets


def analyse(task: dict, data_dir: Path) -> dict[str, Counter]:
    source = data_dir / task["spreadsheet_path"]
    golden = data_dir / task["golden_response_path"]
    input_values = openpyxl.load_workbook(source, data_only=True)
    answer_values = openpyxl.load_workbook(golden, data_only=True)
    input_raw = openpyxl.load_workbook(source, data_only=False)
    answer_raw = openpyxl.load_workbook(golden, data_only=False)
    try:
        targets = _blank_targets(task, input_values, answer_values, input_raw, answer_raw)
        candidates = enumerate_blank_candidates(
            input_raw, instruction=task["instruction"]
        )
        per_arm: dict[str, Counter] = {}
        truth = targets
        for arm in ARMS:
            selected = {(item.sheet, item.address) for item in shortlist(candidates, arm)}
            hits = selected & truth
            counts = Counter(
                {
                    "blank_targets": len(truth),
                    "pool": len(selected),
                    "hits": len(hits),
                    "tasks": 1,
                    "tasks_with_hit": int(bool(hits)),
                    "tasks_with_targets": int(bool(truth)),
                }
            )
            for cell in WATCH:
                if cell in truth:
                    counts[f"watch:{cell[0]}!{cell[1]}:present"] += 1
                    counts[f"watch:{cell[0]}!{cell[1]}:hit"] += int(cell in selected)
            per_arm[arm] = counts
        return per_arm
    finally:
        for workbook in (input_values, answer_values, input_raw, answer_raw):
            workbook.close()


def _report(title: str, totals: dict[str, Counter]) -> None:
    print(f"\n=== {title} ===")
    print(
        f"{'arm':22} {'pool':>8} {'hits':>8} {'targets':>8} "
        f"{'prec':>8} {'recall':>8} {'hit-tasks':>10}"
    )
    for arm in ARMS:
        total = totals[arm]
        pool = total["pool"]
        hits = total["hits"]
        targets = total["blank_targets"]
        tasks = total["tasks_with_targets"]
        precision = hits / pool if pool else 0.0
        recall = hits / targets if targets else 0.0
        print(
            f"{arm:22} {pool:8,} {hits:8,} {targets:8,} "
            f"{precision:8.1%} {recall:8.1%} "
            f"{total['tasks_with_hit']:3}/{tasks:<6}"
        )
    print("\nknown cells (present in this sample / survived the cut):")
    for sheet, address in sorted(WATCH):
        key = f"watch:{sheet}!{address}"
        bits = [
            f"{arm} {totals[arm][key + ':hit']}/{totals[arm][key + ':present']}"
            for arm in ARMS
            if totals[arm][key + ":present"]
        ]
        if bits:
            print(f"  {sheet}!{address}: " + "; ".join(bits))


def main() -> int:
    category = sys.argv[1] if len(sys.argv) > 1 else "Financial_Model"
    limit = int(sys.argv[2]) if len(sys.argv) > 2 else 0
    only = {item.strip() for item in sys.argv[3].split(",") if item.strip()} if len(sys.argv) > 3 else set()
    data_dir = ROOT / "benchmark-data/SpreadsheetBench-2/data" / category
    tasks = json.loads((data_dir / "dataset.json").read_text())
    if only:
        tasks = [task for task in tasks if task["id"] in only]
    if limit:
        tasks = tasks[:limit]
    totals = {arm: Counter() for arm in ARMS}
    measured = 0
    for task in tasks:
        try:
            per_arm = analyse(task, data_dir)
        except Exception as exc:  # noqa: BLE001
            print(f"{task['id']:8} ERROR {type(exc).__name__}: {str(exc)[:80]}", file=sys.stderr)
            continue
        measured += 1
        for arm, counts in per_arm.items():
            totals[arm].update(counts)
        print(
            f"{task['id']:8} targets={per_arm['block-peer']['blank_targets']:5} "
            f"named-block-label pool={per_arm['named-block-label']['pool']:5} "
            f"hits={per_arm['named-block-label']['hits']:4}",
            flush=True,
        )
    _report(f"{category}: {measured} tasks measured", totals)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
