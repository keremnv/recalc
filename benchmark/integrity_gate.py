#!/usr/bin/env python3
"""Offline experiment-integrity gate.

This gate makes no provider request.  It validates request identity mechanics,
rebuilds one isolated 06_01 source->spine->SQLite->evidence lineage, and
records the frozen O3 evidence-access baseline.
"""
from __future__ import annotations

import copy
import hashlib
import json
import sqlite3
import sys
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / "benchmark"), str(ROOT / "src"), str(ROOT / "benchmark/sweagent/formula_index/lib")]

import openpyxl  # noqa: E402

import evidence_access_baseline as evidence_baseline  # noqa: E402
import formula_projection_preflight as projection  # noqa: E402
import matched_compiled_treatment as treatment  # noqa: E402
import workbook_grounding_spine as spine_compiler  # noqa: E402
from experiment_config import (  # noqa: E402
    AUTHORITATIVE_EXPERIMENT_CONFIG,
    archived_mismatch_probe,
    request_identity,
    response_identity,
)
from temporal_spine import compile_temporal_workbook  # noqa: E402
from workbook_grounding import project_obligation  # noqa: E402
from workbook_spine_sqlite import build_database  # noqa: E402


TASK_KEY = "Financial_Model:06_01"
SOURCE = ROOT / "benchmark-data/SpreadsheetBench-2/benchmark-runs/matched-glm-compiled-sixty/prep/repaired_inputs/06_01_DigiMark_input.xlsx"
ARCHIVE_LIVE = ROOT / "benchmark-data/SpreadsheetBench-2/benchmark-runs/matched-glm-compiled-sixty/resource_feasibility/clean_integrated/live/Financial_Model-06_01"
OLD_SPINE = ROOT / "benchmark-data/SpreadsheetBench-2/benchmark-runs/matched-glm-compiled-sixty/prep/spines/Financial_Model-06_01.json"
OLD_DB = ROOT / "benchmark-data/SpreadsheetBench-2/benchmark-runs/matched-glm-compiled-sixty/prep/db/Financial_Model-06_01.sqlite"
REPAIRED_OLD_DB = ROOT / "benchmark-data/SpreadsheetBench-2/benchmark-runs/matched-glm-compiled-sixty/integration_autopsy/repaired_db/Financial_Model-06_01.sqlite"
OUT = ROOT / "benchmark-data/SpreadsheetBench-2/benchmark-runs/mechanical/integrity-gate"
RESULT = OUT / "integrity_gate.json"


def digest(value: Any) -> str:
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()).hexdigest()


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def address_cell_id(spine: dict[str, Any], sheet: str, address: str) -> str:
    target = next((x for x in spine.get("sheets", []) if x.get("title") == sheet), None)
    if target is None:
        raise AssertionError(f"SHEET_NOT_IN_SPINE: {sheet}")
    from openpyxl.utils.cell import coordinate_to_tuple, column_index_from_string
    row, col = coordinate_to_tuple(address)
    return f"cell:s{int(target['index']):02d}:r{row}:c{col}"


def source_facts(source: Path, samples: list[dict[str, str]]) -> dict[str, dict[str, Any]]:
    workbook = openpyxl.load_workbook(source, data_only=False, read_only=True)
    try:
        out = {}
        for sample in samples:
            cell = workbook[sample["sheet"]][sample["address"]]
            value = cell.value
            kind = "blank" if value is None or (isinstance(value, str) and not value.strip()) else "formula" if isinstance(value, str) and value.startswith("=") else "value"
            out[sample["cell_id"]] = {
                "sheet": sample["sheet"], "address": sample["address"],
                "kind": kind, "value": value, "number_format": cell.number_format,
                "data_type": cell.data_type,
            }
        return out
    finally:
        workbook.close()


def scan_value_availability(source: Path) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(source, data_only=False, read_only=True)
    negative: list[str] = []
    booleans: list[str] = []
    try:
        for ws in workbook.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    value = cell.value
                    if value is None:
                        continue
                    from openpyxl.utils import get_column_letter
                    label = f"{ws.title}!{get_column_letter(cell.column)}{cell.row}"
                    if isinstance(value, bool):
                        booleans.append(label)
                    elif isinstance(value, (int, float)) and value < 0:
                        negative.append(label)
    finally:
        workbook.close()
    return {
        "negative_numeric": {"available": bool(negative), "cells": negative[:20], "count": len(negative)},
        "boolean": {"available": bool(booleans), "cells": booleans[:20], "count": len(booleans)},
    }


def query_cell(db: Path, cell_id: str) -> dict[str, Any] | None:
    conn = sqlite3.connect(db)
    conn.row_factory = sqlite3.Row
    try:
        row = conn.execute("SELECT cell_id,kind,raw_value,display_value FROM cells WHERE cell_id=?", (cell_id,)).fetchone()
        return dict(row) if row else None
    finally:
        conn.close()


def evidence_cell(evidence: dict[str, Any], cell_id: str) -> dict[str, Any] | None:
    table = (evidence.get("entities") or {}).get("cells") or {}
    columns = table.get("columns") or []
    for row in table.get("rows") or []:
        value = dict(zip(columns, row))
        if value.get("cell_id") == cell_id:
            return value
    return None


def spine_cell(spine: dict[str, Any], cell_id: str) -> dict[str, Any] | None:
    return next((x for x in spine.get("occupied", []) if x.get("id") == cell_id), None)


def build_fresh_lineage() -> tuple[dict[str, Any], Path, dict[str, Any], dict[str, Any], dict[str, Any]]:
    if not SOURCE.exists():
        raise RuntimeError(f"EXACT_EFFECTIVE_SOURCE_MISSING: {SOURCE}")
    source_sha = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    lineage_dir = OUT / f"fresh_lineage_{source_sha[:12]}"
    lineage_dir.mkdir(parents=True, exist_ok=True)
    spine = spine_compiler.compile_spine(SOURCE, workbook_key=TASK_KEY)
    if spine.get("readable") is not True or not spine.get("lineage", {}).get("payload_sha256"):
        raise AssertionError("FRESH_SPINE_LINEAGE_INCOMPLETE")
    spine_path = lineage_dir / "spine.json"
    temporal_path = lineage_dir / "temporal.json"
    db_path = lineage_dir / "world.sqlite"
    write_json(spine_path, spine)
    temporal = compile_temporal_workbook(SOURCE, closure=True)
    write_json(temporal_path, temporal)
    db_counts = build_database(spine_path, temporal_path, db_path)
    return spine, db_path, temporal, db_counts, {"directory": str(lineage_dir), "spine": str(spine_path), "temporal": str(temporal_path), "database": str(db_path), "source_sha256": source_sha, "database_metadata": treatment._database_metadata(db_path)}


def request_gate() -> dict[str, Any]:
    body = {**AUTHORITATIVE_EXPERIMENT_CONFIG.request_fields(), "messages": [{"role": "system", "content": "s"}, {"role": "user", "content": "u"}]}
    wire = request_identity(body, stage="integrity_gate")
    response = response_identity({"model": AUTHORITATIVE_EXPERIMENT_CONFIG.model}, stage="integrity_gate")
    mismatch = archived_mismatch_probe()
    return {
        "status": "PASS",
        "authoritative_config": AUTHORITATIVE_EXPERIMENT_CONFIG.metadata(),
        "current_request": {**wire, "request_provider": AUTHORITATIVE_EXPERIMENT_CONFIG.provider},
        "matching_response": response,
        "archived_mismatch_regression": mismatch,
        "provider_attempts": 0,
        "stages": {stage: {"request_constructor": "matched_compiled_treatment.request_body", "validator": "experiment_config.request_identity", "timeout_seconds": timeout} for stage, timeout in AUTHORITATIVE_EXPERIMENT_CONFIG.timeouts.items()},
        "legacy_override_policy": "stage-specific request bodies and runtime request-identity overrides are rejected; historical helper paths are not valid experiment entry points",
    }


def fidelity_gate() -> dict[str, Any]:
    spine, db_path, temporal, db_counts, lineage = build_fresh_lineage()
    task_ir = json.loads((ARCHIVE_LIVE / "calls/001_task_ir.json").read_text(encoding="utf-8"))["parsed_response"]
    obligation = next(x for x in task_ir["obligations"] if x["id"] == "O3")
    packet = project_obligation(spine, obligation)
    packet_ids = projection.packet_ids(packet)
    target = treatment.target_from_id(spine, "cell:s01:r49:c3", TASK_KEY, "O3")
    bootstrap = treatment.prior.compile_bootstrap(spine, packet, obligation, target)
    bootstrap_ids = set(bootstrap["bootstrap_entity_ids"])
    accepted = json.loads((ROOT / "research/history/resource_demand_autopsy/accepted_edit_slices.json").read_text(encoding="utf-8"))
    c49 = next(x for x in accepted if x["cell_id"] == "cell:s01:r49:c3")
    working_ids = set(c49["full_evidence_ids"])
    bootstrap_evidence = treatment.materialize_complete(db_path, bootstrap_ids | {target["cell_id"]})
    final_evidence = treatment.materialize_complete(db_path, working_ids)
    samples_raw = [
        ("DCF", "C47", "positive_numeric_percentage"),
        ("DCF", "C48", "archived_witness_numeric"),
        ("DCF", "C51", "archived_witness_numeric"),
        ("DCF", "C44", "zero"),
        ("DCF", "B49", "text"),
        ("DCF", "C50", "formula"),
        ("DCF", "C49", "blank"),
    ]
    samples = []
    for sheet, address, category in samples_raw:
        samples.append({"sheet": sheet, "address": address, "category": category, "cell_id": address_cell_id(spine, sheet, address)})
    sources = source_facts(SOURCE, samples)
    facts = []
    for sample in samples:
        cid = sample["cell_id"]
        source = sources[cid]
        compiled = spine_cell(spine, cid)
        db_value = query_cell(db_path, cid)
        boot_value = evidence_cell(bootstrap_evidence, cid)
        final_value = evidence_cell(final_evidence, cid)
        expected_serialized = None if source["kind"] == "blank" else str(source["value"])
        source_value_ok = (compiled or {}).get("payload") == source["value"] if source["kind"] != "blank" else (compiled or {}).get("kind") == "blank" and "payload" not in (compiled or {})
        db_value_ok = bool(db_value) and db_value.get("raw_value") == expected_serialized and db_value.get("display_value") == expected_serialized
        packet_present = cid in packet_ids
        bootstrap_ok = bool(boot_value) and boot_value.get("raw_value") == expected_serialized and boot_value.get("display_value") == expected_serialized
        final_ok = bool(final_value) and final_value.get("raw_value") == expected_serialized and final_value.get("display_value") == expected_serialized
        bootstrap_stage_ok = (cid not in bootstrap_ids) or bootstrap_ok
        facts.append({
            **sample,
            "source": source,
            "compiled_spine": {"present": compiled is not None, "kind": (compiled or {}).get("kind"), "payload": (compiled or {}).get("payload"), "value_ok": source_value_ok},
            "grounding_projection": {"present": packet_present, "packet_id_count": len(packet_ids), "projected_value": (compiled or {}).get("payload") if packet_present else None},
            "bootstrap_evidence": {"present": bool(boot_value), "required_by_bootstrap": cid in bootstrap_ids, "value": boot_value, "value_ok": bootstrap_ok, "stage_ok": bootstrap_stage_ok},
            "retrieval_working_set": {"present": cid in working_ids, "working_set_id_count": len(working_ids)},
            "final_synthesis_packet": {"present": bool(final_value), "value": final_value, "value_ok": final_ok},
            "stage_pass": bool(source_value_ok and db_value_ok and bootstrap_stage_ok and cid in working_ids and final_ok),
        })
    stale_rows = {cid: {"prep_db": query_cell(OLD_DB, cid), "repaired_db": query_cell(REPAIRED_OLD_DB, cid), "stale_spine": spine_cell(json.loads(OLD_SPINE.read_text(encoding="utf-8")), cid)} for cid in ("cell:s01:r47:c3", "cell:s01:r48:c3", "cell:s01:r51:c3")}
    archived_environment = treatment.PREP / "environment.json"
    stale_environment_rejected = False
    if archived_environment.exists():
        environment = json.loads(archived_environment.read_text(encoding="utf-8"))
        archived_record = next((x for x in environment.get("records", []) if x.get("task_key") == TASK_KEY), None)
        stale_environment_rejected = archived_record is not None and not treatment._environment_record_current(archived_record)
    lineage_guard = {"archived_environment": str(archived_environment), "stale_archived_environment_rejected": stale_environment_rejected}
    earliest = {
        "classification": "stale database/cache lineage",
        "boundary": "archived effective source -> stale compiled spine/database artifact",
        "witnesses": stale_rows,
        "fresh_lineage_repaired": all(x["stage_pass"] for x in facts),
        "lineage_guard": lineage_guard,
        "explanation": "The effective source workbook contains the witness numerics; fresh compile_spine preserves payloads and fresh SQLite materialization carries them into synthesis evidence. The archived spine omitted those payloads and both archived databases stored null raw/display values.",
    }
    unavailable = scan_value_availability(SOURCE)
    return {
        "status": "PASS" if all(x["stage_pass"] for x in facts) and stale_environment_rejected else "FAIL",
        "source": str(SOURCE),
        "source_sha256": lineage["source_sha256"],
        "fresh_isolated_lineage": lineage,
        "database_counts": db_counts,
        "task_ir_obligation": obligation,
        "grounding_projection": {"packet_id_count": len(packet_ids), "target_candidate_count": len(packet.get("target_cell_ids") or []), "bootstrap_id_count": len(bootstrap_ids), "retrieval_working_set_id_count": len(working_ids), "final_synthesis_packet_id_count": len(final_evidence.get("entity_ids") or [])},
        "facts": facts,
        "availability": unavailable,
        "archived_numeric_witness_loss": earliest,
    }


def main() -> dict[str, Any]:
    request = request_gate()
    evidence_baseline_result = evidence_baseline.main()
    fidelity = fidelity_gate()
    status = "EXPERIMENT_INTEGRITY_VERIFIED" if request["status"] == "PASS" and fidelity["status"] == "PASS" and evidence_baseline_result["status"] == "PASS" else "MULTIPLE_INTEGRITY_DEFECTS_REMAIN"
    result = {"status": status, "provider_calls": 0, "request_configuration": request, "evidence_access_baseline": evidence_baseline_result, "evidence_fidelity": fidelity, "benchmark_comparison_launched": False, "forbidden_next_actions_observed": False}
    write_json(RESULT, result)
    return result


if __name__ == "__main__":
    print(json.dumps(main(), ensure_ascii=False, indent=2))
