"""Gold-blind S0/S1 workbook-native label and schema extraction.

No goldens, no embeddings, no LLM label interpretation, no S2 synonyms.
False-negative relations are preferred to guessed ones.
"""
from __future__ import annotations

import calendar
import re
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from fingerprint import a1_address, formula_text
from xlsx_metadata_repair import install

install()
import openpyxl  # noqa: E402

LEFT_SPAN = 12
UP_STACK = 12
HEADER_UP = 24
JACCARD_THRESHOLDS = (1.0, 0.80, 0.60, 0.40)

_PUNCT = re.compile(r"[^\w\s]+", re.UNICODE)
_WS = re.compile(r"\s+")
_PAREN = re.compile(r"\([^)]*\)")
_NUM_SUFFIX = re.compile(r"^(.*?)(?:\s*[-_/]?\s*\d+)$")
_FY4 = re.compile(r"^(?:fy|fye|cy)?\s*'?(19|20)\d{2}[a-z]?$", re.I)
_FY2 = re.compile(r"^(?:fy|fye|cy)\s*'?(\d{2})[a-z]?$", re.I)
_YEAR = re.compile(r"^(19|20)\d{2}$")
_Q_ONLY = re.compile(r"^(?:q\s*([1-4])|([1-4])\s*q)$", re.I)
_Q_YEAR = re.compile(
    r"^(?:q\s*([1-4])|([1-4])\s*q)\s*[-/' ]\s*(?:fy)?\s*'?((?:19|20)\d{2}|\d{2})$",
    re.I,
)
_YEAR_Q = re.compile(r"^((?:19|20)\d{2}|\d{2})\s*[-/' ]\s*q\s*([1-4])$", re.I)
_MONTH_NAME = {name.lower(): i for i, name in enumerate(calendar.month_name) if name}
_MONTH_ABBR = {name.lower(): i for i, name in enumerate(calendar.month_abbr) if name}

DEFINITIONS = {
    "golden_in_generation": False,
    "S0": (
        "Mechanical schema facts from workbook geometry and cell contents: "
        "text-bearing cells, candidate row labels within 12 columns left, "
        "merged-header membership, upward label stacks of 12 rows, headers "
        "within 24 rows above, literal year/quarter/month/integer periods, "
        "exact normalized text equality, and cross-sheet label repetition. "
        "No synonym inference."
    ),
    "S1": (
        "Deterministic lexical relations on preserved original strings: "
        "lowercase, punctuation stripped, parenthetical kept and stripped "
        "as two views, token Jaccard, ordered-token equality, substring "
        "containment. Thresholds exact / 0.80 / 0.60 / 0.40 are predeclared. "
        "No embeddings or language models."
    ),
    "S2_not_implemented": "Semantic synonymy (revenue↔sales) is not used.",
    "row_label_canonical": (
        "Nearest non-empty textual cell to the left of C within 12 columns "
        "that is not period-only, has at least one letter, and is not a "
        "formula. Merged cells intersecting the row count at their leftmost "
        "column. Ties: nearest distance, then lowest column."
    ),
    "not_used": (
        "Goldens, task text, S2 semantics, agents, official scores, and "
        "formula reconstruction are not used in extraction."
    ),
}


def is_formula(value: object) -> bool:
    return formula_text(value) is not None


def as_text(value: object) -> str | None:
    if value is None or is_formula(value):
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, str):
        text = value.strip()
        return text or None
    if hasattr(value, "year") and hasattr(value, "month"):
        return None
    if isinstance(value, (int, float)):
        return None
    text = str(value).strip()
    return text or None


def _expand_year(token: str) -> int | None:
    if len(token) == 4:
        year = int(token)
        return year if 1990 <= year <= 2100 else None
    year = 2000 + int(token)
    return year if 1990 <= year <= 2100 else None


def parse_period(value: object) -> dict[str, int] | None:
    """Return a mechanical PERIOD dict, or None if not directly parseable."""
    if value is None or is_formula(value):
        return None
    if hasattr(value, "year") and hasattr(value, "month"):
        return {"year": int(value.year), "month": int(value.month)}
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)) and float(value).is_integer():
        number = int(value)
        if 1990 <= number <= 2100:
            return {"year": number}
        return None
    text = str(value).strip()
    if not text:
        return None
    compact = _WS.sub(" ", text.replace("_", " ")).strip()
    if _FY4.match(compact.replace(" ", "")):
        digits = re.search(r"(19|20)\d{2}", compact)
        if digits:
            return {"year": int(digits.group(0))}
    match = _FY2.match(compact.replace(" ", ""))
    if match:
        year = _expand_year(match.group(1))
        if year is not None:
            return {"year": year}
    if _YEAR.match(compact):
        return {"year": int(compact)}
    match = _Q_YEAR.match(compact)
    if match:
        quarter = int(match.group(1) or match.group(2))
        year = _expand_year(match.group(3))
        if year is not None:
            return {"year": year, "quarter": quarter}
    match = _YEAR_Q.match(compact)
    if match:
        year = _expand_year(match.group(1))
        if year is not None:
            return {"year": year, "quarter": int(match.group(2))}
    match = _Q_ONLY.match(compact)
    if match:
        return {"quarter": int(match.group(1) or match.group(2))}
    lower = compact.lower()
    for table in (_MONTH_NAME, _MONTH_ABBR):
        for name, month in table.items():
            if lower == name:
                return {"month": month}
            prefix = name + " "
            if lower.startswith(prefix) or lower.startswith(name + "-") or lower.startswith(name + "/"):
                rest = re.sub(r"^" + re.escape(name) + r"[\s\-/' ]*", "", lower)
                rest = rest.replace("fy", "").strip(" '")
                year = None
                if rest.isdigit():
                    year = _expand_year(rest)
                else:
                    found = re.search(r"(19|20)\d{2}", rest)
                    if found:
                        year = int(found.group(0))
                out: dict[str, int] = {"month": month}
                if year is not None:
                    out["year"] = year
                return out
    return None


def period_key(period: dict[str, int] | None) -> str | None:
    if not period:
        return None
    parts = []
    if "year" in period:
        parts.append(f"year={period['year']}")
    if "quarter" in period:
        parts.append(f"quarter={period['quarter']}")
    if "month" in period:
        parts.append(f"month={period['month']}")
    if "index" in period:
        parts.append(f"index={period['index']}")
    return "PERIOD(" + ", ".join(parts) + ")" if parts else None


def is_period_only(value: object) -> bool:
    text = as_text(value)
    if text is None:
        return parse_period(value) is not None
    if parse_period(text) is None:
        return False
    letters = re.sub(r"[^a-z]", "", text.lower())
    allowed = set("fyqcyjanfebmaraprmayjunjulaugsepoctnovdeceomth")
    return not letters or set(letters) <= allowed


def has_letter(text: str) -> bool:
    return any(ch.isalpha() for ch in text)


def normalize_text(text: str, *, strip_paren: bool = False) -> str:
    value = text.strip().lower()
    if strip_paren:
        value = _PAREN.sub(" ", value)
    value = _PUNCT.sub(" ", value)
    return _WS.sub(" ", value).strip()


def tokens(text: str, *, strip_paren: bool = False) -> tuple[str, ...]:
    norm = normalize_text(text, strip_paren=strip_paren)
    if not norm:
        return ()
    return tuple(norm.split(" "))


def strip_numeric_suffix_tokens(toks: tuple[str, ...]) -> tuple[str, ...]:
    if len(toks) >= 2 and toks[-1].isdigit():
        return toks[:-1]
    return toks


def jaccard(a: tuple[str, ...], b: tuple[str, ...]) -> float | None:
    if not a and not b:
        return 1.0
    if not a or not b:
        return 0.0
    left, right = set(a), set(b)
    union = left | right
    if not union:
        return None
    return round(len(left & right) / len(union), 4)


def lexical_relation(a: str | None, b: str | None) -> dict[str, Any]:
    empty = {
        "exact_norm": False,
        "ordered_equal": False,
        "ordered_equal_noparen": False,
        "substring": False,
        "jaccard": None,
        "jaccard_noparen": None,
        "jaccard_stem": None,
        "bands": {str(th): False for th in JACCARD_THRESHOLDS},
    }
    if not a or not b:
        return empty
    ta, tb = tokens(a), tokens(b)
    na, nb = tokens(a, strip_paren=True), tokens(b, strip_paren=True)
    sa, sb = strip_numeric_suffix_tokens(na), strip_numeric_suffix_tokens(nb)
    jac = jaccard(ta, tb)
    jac_np = jaccard(na, nb)
    jac_st = jaccard(sa, sb)
    best = max(v for v in (jac, jac_np, jac_st) if v is not None)
    return {
        "exact_norm": normalize_text(a) == normalize_text(b),
        "ordered_equal": ta == tb and bool(ta),
        "ordered_equal_noparen": na == nb and bool(na),
        "substring": normalize_text(a) in normalize_text(b) or normalize_text(b) in normalize_text(a),
        "jaccard": jac,
        "jaccard_noparen": jac_np,
        "jaccard_stem": jac_st,
        "bands": {str(th): best >= th for th in JACCARD_THRESHOLDS},
    }


@dataclass
class SheetIndex:
    title: str
    values: dict[tuple[int, int], object]
    merges: list[tuple[int, int, int, int]]
    merge_origin: dict[tuple[int, int], tuple[int, int]]
    max_col: int
    max_row: int
    text_by_row: dict[int, list[tuple[int, str]]]
    text_by_col: dict[int, list[tuple[int, str]]]


def load_workbook_index(path: Path) -> dict[str, SheetIndex]:
    workbook = openpyxl.load_workbook(path, data_only=False, read_only=False)
    out: dict[str, SheetIndex] = {}
    try:
        for sheet in workbook.worksheets:
            values: dict[tuple[int, int], object] = {}
            max_col = max_row = 1
            for cell in sheet._cells.values():
                col, row = int(cell.column), int(cell.row)
                values[(col, row)] = cell.value
                max_col = max(max_col, col)
                max_row = max(max_row, row)
            merges = [
                (int(rng.min_col), int(rng.min_row), int(rng.max_col), int(rng.max_row))
                for rng in sheet.merged_cells.ranges
            ]
            origin: dict[tuple[int, int], tuple[int, int]] = {}
            for c1, r1, c2, r2 in merges:
                for col in range(c1, c2 + 1):
                    for row in range(r1, r2 + 1):
                        origin[(col, row)] = (c1, r1)
            text_by_row: dict[int, list[tuple[int, str]]] = defaultdict(list)
            text_by_col: dict[int, list[tuple[int, str]]] = defaultdict(list)
            for (col, row), value in values.items():
                text = as_text(value)
                if text is None or not has_letter(text):
                    continue
                text_by_row[row].append((col, text))
                text_by_col[col].append((row, text))
            for row in text_by_row:
                text_by_row[row].sort()
            for col in text_by_col:
                text_by_col[col].sort()
            out[sheet.title] = SheetIndex(
                title=sheet.title,
                values=values,
                merges=merges,
                merge_origin=origin,
                max_col=max_col,
                max_row=max_row,
                text_by_row=dict(text_by_row),
                text_by_col=dict(text_by_col),
            )
    finally:
        workbook.close()
    return out


def _display_value(index: SheetIndex, col: int, row: int) -> object:
    origin = index.merge_origin.get((col, row), (col, row))
    return index.values.get(origin)


def _merged_span(index: SheetIndex, col: int, row: int) -> tuple[int, int, int, int] | None:
    origin = index.merge_origin.get((col, row))
    if origin is None:
        return None
    for c1, r1, c2, r2 in index.merges:
        if (c1, r1) == origin:
            return (c1, r1, c2, r2)
    return None


def candidate_row_labels(index: SheetIndex, col: int, row: int) -> list[dict[str, Any]]:
    lo = max(1, col - LEFT_SPAN)
    seen: set[tuple[int, int]] = set()
    out: list[dict[str, Any]] = []
    for c in range(col - 1, lo - 1, -1):
        origin = index.merge_origin.get((c, row), (c, row))
        if origin in seen:
            continue
        value = index.values.get(origin)
        text = as_text(value)
        if text is None:
            continue
        seen.add(origin)
        span = _merged_span(index, origin[0], origin[1])
        precedes = any(
            index.values.get((x, row)) not in (None, "")
            for x in range(origin[0] + 1, min(index.max_col, origin[0] + 8) + 1)
        )
        out.append(
            {
                "sheet": index.title,
                "col": origin[0],
                "row": origin[1],
                "address": a1_address(origin[0], origin[1]),
                "raw": text,
                "norm": normalize_text(text),
                "distance": col - origin[0],
                "merged": span is not None,
                "merged_span": span,
                "period_only": is_period_only(text),
                "has_letter": has_letter(text),
                "precedes_occupied": precedes,
                "provenance": "same_row_left" if span is None else "merged_intersecting_row",
            }
        )
    return out


def canonical_row_label(candidates: list[dict[str, Any]]) -> dict[str, Any] | None:
    eligible = [
        item
        for item in candidates
        if item["has_letter"] and not item["period_only"]
    ]
    if not eligible:
        return None
    eligible.sort(key=lambda item: (item["distance"], item["col"]))
    return eligible[0]


def row_label_stack(
    index: SheetIndex,
    col: int,
    row: int,
    label_col: int | None,
) -> list[dict[str, Any]]:
    parents: list[dict[str, Any]] = []
    lo = max(1, row - UP_STACK)
    align_cols = {label_col} if label_col else set()
    align_cols.add(1)
    for r in range(row - 1, lo - 1, -1):
        found = None
        for c in sorted(align_cols):
            if c is None:
                continue
            origin = index.merge_origin.get((c, r), (c, r))
            text = as_text(index.values.get(origin))
            if text is None or not has_letter(text) or is_period_only(text):
                continue
            span = _merged_span(index, origin[0], origin[1])
            overlaps = False
            if span is not None and span[0] <= col <= span[2]:
                overlaps = True
            if origin[0] in align_cols or overlaps:
                found = {
                    "sheet": index.title,
                    "col": origin[0],
                    "row": origin[1],
                    "address": a1_address(origin[0], origin[1]),
                    "raw": text,
                    "norm": normalize_text(text),
                    "distance": row - origin[1],
                    "merged": span is not None,
                    "merged_span": span,
                    "provenance": "upward_align_or_overlap",
                }
                break
        if found is None:
            for c, text in index.text_by_row.get(r, []):
                if c > col:
                    continue
                span = _merged_span(index, c, r)
                if span is not None and span[0] <= col <= span[2] and has_letter(text) and not is_period_only(text):
                    found = {
                        "sheet": index.title,
                        "col": c,
                        "row": r,
                        "address": a1_address(c, r),
                        "raw": text,
                        "norm": normalize_text(text),
                        "distance": row - r,
                        "merged": True,
                        "merged_span": span,
                        "provenance": "upward_merge_overlap",
                    }
                    break
        if found is not None:
            parents.append(found)
    return parents


def header_stack(index: SheetIndex, col: int, row: int) -> dict[str, Any]:
    headers: list[dict[str, Any]] = []
    lo = max(1, row - HEADER_UP)
    for r in range(row - 1, lo - 1, -1):
        origin = index.merge_origin.get((col, r), (col, r))
        value = index.values.get(origin)
        text = as_text(value)
        period = parse_period(value if text is None else text)
        if text is None and period is None:
            continue
        span = _merged_span(index, origin[0], origin[1])
        headers.append(
            {
                "sheet": index.title,
                "col": origin[0],
                "row": origin[1],
                "address": a1_address(origin[0], origin[1]),
                "raw": text,
                "norm": normalize_text(text) if text else None,
                "distance": row - origin[1],
                "merged": span is not None,
                "merged_span": span,
                "period": period,
                "period_key": period_key(period),
            }
        )
    nearest = headers[0] if headers else None
    parents = [item for item in headers[1:] if item.get("merged")]
    period_hit = next((item for item in headers if item["period"]), None)
    return {
        "headers": headers[:8],
        "nearest": nearest,
        "parent_merged": parents[:4],
        "period": None if period_hit is None else period_hit["period"],
        "period_key": None if period_hit is None else period_hit["period_key"],
        "period_address": None if period_hit is None else period_hit["address"],
    }


def schema_for_cell(index: SheetIndex, col: int, row: int) -> dict[str, Any]:
    candidates = candidate_row_labels(index, col, row)
    canonical = canonical_row_label(candidates)
    label_col = None if canonical is None else canonical["col"]
    stack = row_label_stack(index, col, row, label_col)
    headers = header_stack(index, col, row)
    return {
        "sheet": index.title,
        "col": col,
        "row": row,
        "address": a1_address(col, row),
        "immediate": canonical,
        "label_candidates": candidates[:8],
        "stack": stack[:6],
        "headers": headers,
        "has_immediate": canonical is not None,
        "has_stack": bool(stack),
        "has_header": headers["nearest"] is not None,
        "has_period": headers["period"] is not None,
        "label_norm": None if canonical is None else canonical["norm"],
        "label_raw": None if canonical is None else canonical["raw"],
        "period_key": headers["period_key"],
        "parent_norms": [item["norm"] for item in stack],
    }


def compact_schema(schema: dict[str, Any]) -> dict[str, Any]:
    imm = schema.get("immediate")
    nearest = (schema.get("headers") or {}).get("nearest")
    return {
        "sheet": schema["sheet"],
        "address": schema["address"],
        "label_raw": schema.get("label_raw"),
        "label_norm": schema.get("label_norm"),
        "label_addr": None if not imm else f"{imm['sheet']}!{imm['address']}",
        "parents": [item["raw"] for item in schema.get("stack") or []],
        "header_raw": None if not nearest else nearest.get("raw"),
        "period_key": schema.get("period_key"),
        "has_immediate": schema.get("has_immediate"),
        "has_stack": schema.get("has_stack"),
        "has_header": schema.get("has_header"),
        "has_period": schema.get("has_period"),
    }


@dataclass
class WorkbookSchema:
    sheets: dict[str, SheetIndex]
    label_rows: dict[str, list[tuple[str, int]]] = field(default_factory=dict)
    period_cols: dict[str, dict[int, str]] = field(default_factory=dict)

    def cell(self, sheet: str, col: int, row: int) -> dict[str, Any] | None:
        index = self.sheets.get(sheet)
        if index is None:
            return None
        return schema_for_cell(index, col, row)


def build_workbook_schema(path: Path) -> WorkbookSchema:
    sheets = load_workbook_index(path)
    label_rows: dict[str, list[tuple[str, int]]] = defaultdict(list)
    period_cols: dict[str, dict[int, str]] = {}
    for title, index in sheets.items():
        periods: dict[int, str] = {}
        for row, items in index.text_by_row.items():
            eligible = [
                (col, text)
                for col, text in items
                if has_letter(text) and not is_period_only(text)
            ]
            if not eligible:
                continue
            _col, text = eligible[0]
            label_rows[normalize_text(text)].append((title, row))
        for col in range(1, min(index.max_col, 80) + 1):
            for row in range(1, min(index.max_row, HEADER_UP) + 1):
                period = parse_period(_display_value(index, col, row))
                key = period_key(period)
                if key:
                    periods[col] = key
                    break
        period_cols[title] = periods
    return WorkbookSchema(sheets=sheets, label_rows=dict(label_rows), period_cols=period_cols)


def exact_label_counterparts(schema: WorkbookSchema, label_norm: str | None) -> list[tuple[str, int]]:
    if not label_norm:
        return []
    return list(schema.label_rows.get(label_norm, []))
