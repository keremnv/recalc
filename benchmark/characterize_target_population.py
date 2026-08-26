#!/usr/bin/env python3
"""Separate direct, downstream, and dynamic-only benchmark target populations.

This offline diagnostic leaves the official evaluator and score unchanged.
See target_classification.py for the precise populations reported here.
"""

from __future__ import annotations

import json
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "benchmark-data/SpreadsheetBench-2/evaluation"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import evaluation as ev
import openpyxl
from target_classification import classify_cache_robust_targets
from xlsx_metadata_repair import install as _install_repair

_install_repair()


def analyse(task: dict, data_dir: Path) -> Counter:
    source = data_dir / task["spreadsheet_path"]
    golden = data_dir / task["golden_response_path"]
    with_color = "Color" in task["spreadsheet_path"]
    with_formula = "Embedded" in task["spreadsheet_path"]
    input_values = openpyxl.load_workbook(source, data_only=not with_formula)
    answer_values = openpyxl.load_workbook(golden, data_only=not with_formula)
    input_raw = (
        openpyxl.load_workbook(source, data_only=False) if not with_formula else input_values
    )
    answer_raw = (
        openpyxl.load_workbook(golden, data_only=False) if not with_formula else answer_values
    )

    counts: Counter = Counter()
    for chunk in ev.parse_answer_position(task["answer_position"]):
        sheet_name, _, cell_range = chunk.rpartition("!")
        sheet_name = sheet_name.strip().strip("'")
        result = classify_cache_robust_targets(
            ev,
            input_values,
            answer_values,
            input_raw,
            answer_raw,
            sheet_name,
            cell_range,
            with_font_color=with_color,
            with_formula=with_formula,
        )
        counts["official regression"] += len(result.regression) - len(
            result.unchanged_formula_value_differences
        )
        counts["official modification"] += len(result.modification) + len(
            result.unchanged_formula_value_differences
        )
        counts["direct value targets"] += len(result.modification)
        counts["unchanged-formula"] += len(result.unchanged_formula_value_differences)
        counts["dynamic-only"] += len(result.value_equivalent_formula_differences)
        counts["indeterminate"] += len(result.indeterminate_uncached_formula_differences)

    for workbook in {input_values, answer_values, input_raw, answer_raw}:
        workbook.close()
    return counts


def main() -> int:
    category = sys.argv[1] if len(sys.argv) > 1 else "Financial_Model"
    limit = int(sys.argv[2]) if len(sys.argv) > 2 else 0
    data_dir = ROOT / "benchmark-data/SpreadsheetBench-2/data" / category
    tasks = json.loads((data_dir / "dataset.json").read_text())
    if limit:
        tasks = tasks[:limit]

    total: Counter = Counter()
    measured = 0
    for task in tasks:
        try:
            counts = analyse(task, data_dir)
        except Exception as exc:  # noqa: BLE001
            print(f"{task['id']:8} ERROR {type(exc).__name__}: {str(exc)[:80]}", file=sys.stderr)
            continue
        measured += 1
        total.update(counts)
        print(
            f"{task['id']:8} official_mod={counts['official modification']:7} "
            f"unchanged_formula={counts['unchanged-formula']:7} "
            f"direct={counts['direct value targets']:6} "
            f"dynamic_only={counts['dynamic-only']:6} indeterminate={counts['indeterminate']:5}"
        )

    official = total["official modification"]
    print(f"\n=== {category}: {measured} tasks ===")
    print(f"official modification population {official:10,}")
    print(
        f"  unchanged-formula value differences {total['unchanged-formula']:7,} "
        f"({total['unchanged-formula'] / max(1, official):.1%})"
    )
    print(f"cache-robust direct value targets {total['direct value targets']:7,}")
    print(f"dynamic-only formula differences {total['dynamic-only']:7,}")
    print(f"  indeterminate (both uncached) {total['indeterminate']:8,}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
