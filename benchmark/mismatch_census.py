#!/usr/bin/env python3
"""Census every scored cell mismatch in stored benchmark outputs.

SpreadsheetBench 2's result JSON retains only the first error per task. This diagnostic replays
the vendored evaluator's cell comparisons over every answer range, then adds an orthogonal causal
classification from :mod:`target_classification`:

* a direct blank/populated target was left unchanged or edited incorrectly;
* an unchanged formula is wrong because an upstream edit is missing (a downstream cascade);
* a downstream formula cell was itself over-edited; or
* an official regression cell was over-edited.

The official regression/modification partition is preserved verbatim. The additional direct-target
axis is offline analysis only and never changes official scores or exposes goldens to an agent.
Known-invalid harness attempts are excluded from aggregates by default.
"""

from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(ROOT / "benchmark-data/SpreadsheetBench-2/evaluation"))

import evaluation as ev
import openpyxl
from experiment_validity import invalid_reason
from target_classification import classify_cache_robust_targets
from xlsx_metadata_repair import install

install()

DEFAULT_RUNS_ROOT = ROOT / "benchmark-data/SpreadsheetBench-2/benchmark-runs/openrouter"
DEFAULT_DATA_ROOT = ROOT / "benchmark-data/SpreadsheetBench-2/data"


@dataclass(frozen=True)
class WorkbookViews:
    """Calculated-value and raw-formula views of one workbook."""

    values: Any
    raw: Any


@dataclass(frozen=True)
class CellMismatch:
    sheet: str
    address: str
    official_partition: str
    target_kind: str
    causal_class: str
    input_raw: Any
    golden_raw: Any
    output_raw: Any
    golden_value: Any
    output_value: Any


@dataclass(frozen=True)
class AttemptCensus:
    run_name: str
    task: str
    official_exact: bool
    stored_regression_accuracy: float | None
    stored_modification_accuracy: float | None
    recomputed_regression_accuracy: float
    recomputed_modification_accuracy: float
    score_matches_stored: bool
    regression_correct: int
    regression_total: int
    modification_correct: int
    modification_total: int
    mismatches: tuple[CellMismatch, ...]


def _display(value: Any) -> Any:
    """Return a compact JSON-safe representation of a cell value."""

    value = getattr(value, "text", value)
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _load_views(path: Path, *, with_formula: bool) -> WorkbookViews:
    if with_formula:
        raw = openpyxl.load_workbook(path, data_only=False)
        return WorkbookViews(raw, raw)
    return WorkbookViews(
        openpyxl.load_workbook(path, data_only=True),
        openpyxl.load_workbook(path, data_only=False),
    )


def _close_views(views: WorkbookViews) -> None:
    views.values.close()
    if views.raw is not views.values:
        views.raw.close()


def _sheet_name_and_range(workbook: Any, chunk: str) -> tuple[str, str]:
    if "!" not in chunk:
        return workbook.sheetnames[0], chunk.strip("'").strip()
    sheet_name, _, cell_range = chunk.rpartition("!")
    return sheet_name.strip("'").strip(), cell_range.strip("'").strip()


def _official_cells_match(
    answer_value_cell,
    output_value_cell,
    answer_raw_cell,
    output_raw_cell,
    *,
    with_font_color: bool,
    with_formula: bool,
) -> bool:
    if (
        not with_formula
        and (ev._has_excel_error(answer_value_cell) or ev._has_excel_error(output_value_cell))
    ):
        return ev.compare_cell_formula(answer_raw_cell, output_raw_cell)
    return ev._compare_cells(
        answer_value_cell,
        output_value_cell,
        with_font_color,
        with_formula,
    )


def _content_matches_input(input_raw_cell, output_raw_cell, *, with_font_color: bool) -> bool:
    content_matches = ev.compare_cell_formula(input_raw_cell, output_raw_cell)
    colors_match = not with_font_color or ev.compare_font_color(
        input_raw_cell.font, output_raw_cell.font
    )
    return content_matches and colors_match


def _rounded_accuracy(correct: int, total: int, *, clamp_regression: bool) -> float:
    value = round(correct / total, 4) if total else 0.0
    if clamp_regression and value >= 0.998:
        return 1.0
    return value


def _score_agrees(stored: float | None, recomputed: float) -> bool:
    return stored is None or abs(float(stored) - recomputed) < 1e-9


def census_workbooks(
    *,
    run_name: str,
    task_key: str,
    answer_position: str,
    input_views: WorkbookViews,
    golden_views: WorkbookViews,
    output_views: WorkbookViews,
    stored_score: dict[str, Any],
    with_font_color: bool = False,
    with_formula: bool = False,
) -> AttemptCensus:
    """Return every official cell mismatch plus its offline causal classification."""

    mismatches: list[CellMismatch] = []
    regression_correct = regression_total = 0
    modification_correct = modification_total = 0

    for chunk in ev.parse_answer_position(answer_position):
        sheet_name, cell_range = _sheet_name_and_range(golden_views.values, chunk)
        official_regression, official_modification = ev.classify_cells_by_modification(
            input_views.values,
            golden_views.values,
            sheet_name,
            cell_range,
            with_font_color,
            with_formula,
            input_views.raw if not with_formula else None,
            golden_views.raw if not with_formula else None,
        )
        classified = classify_cache_robust_targets(
            ev,
            input_views.values,
            golden_views.values,
            input_views.raw,
            golden_views.raw,
            sheet_name,
            cell_range,
            with_font_color=with_font_color,
            with_formula=with_formula,
        )
        direct_targets = set(classified.modification)
        downstream_targets = set(classified.unchanged_formula_value_differences)
        dynamic_only = set(classified.value_equivalent_formula_differences)

        input_values = ev._find_sheet(input_views.values, sheet_name)
        input_raw = ev._find_sheet(input_views.raw, sheet_name)
        golden_values = ev._find_sheet(golden_views.values, sheet_name)
        golden_raw = ev._find_sheet(golden_views.raw, sheet_name)
        output_values = ev._find_sheet(output_views.values, sheet_name)
        output_raw = ev._find_sheet(output_views.raw, sheet_name)

        regression_total += len(official_regression)
        modification_total += len(official_modification)
        if None in (input_values, input_raw, golden_values, golden_raw):
            continue

        for official_partition, addresses in (
            ("regression", official_regression),
            ("modification", official_modification),
        ):
            for address in addresses:
                if output_values is None or output_raw is None:
                    matched = False
                else:
                    matched = _official_cells_match(
                        golden_values[address],
                        output_values[address],
                        golden_raw[address],
                        output_raw[address],
                        with_font_color=with_font_color,
                        with_formula=with_formula,
                    )
                if matched:
                    if official_partition == "regression":
                        regression_correct += 1
                    else:
                        modification_correct += 1
                    continue

                if output_values is None or output_raw is None:
                    target_kind = (
                        "direct"
                        if address in direct_targets
                        else "downstream"
                        if address in downstream_targets
                        else "dynamic-only"
                        if address in dynamic_only
                        else "regression"
                    )
                    causal_class = "structural: output sheet missing"
                    output_raw_value = None
                    output_value = None
                else:
                    unchanged = _content_matches_input(
                        input_raw[address],
                        output_raw[address],
                        with_font_color=with_font_color,
                    )
                    blank_input = (
                        input_values[address].value is None and input_raw[address].value is None
                    )
                    if address in direct_targets:
                        target_kind = "direct"
                        population = "blank" if blank_input else "populated"
                        action = "unchanged" if unchanged else "wrong edit"
                        causal_class = f"direct {population} target: {action}"
                    elif address in downstream_targets:
                        target_kind = "downstream"
                        causal_class = (
                            "downstream cascade at unchanged formula"
                            if unchanged
                            else "downstream cell over-edited"
                        )
                    elif address in dynamic_only:
                        target_kind = "dynamic-only"
                        causal_class = "dynamic-only regression cell over-edited"
                    else:
                        target_kind = "regression"
                        causal_class = "regression cell over-edited"
                    output_raw_value = output_raw[address].value
                    output_value = output_values[address].value

                mismatches.append(
                    CellMismatch(
                        sheet=sheet_name,
                        address=address,
                        official_partition=official_partition,
                        target_kind=target_kind,
                        causal_class=causal_class,
                        input_raw=_display(input_raw[address].value),
                        golden_raw=_display(golden_raw[address].value),
                        output_raw=_display(output_raw_value),
                        golden_value=_display(golden_values[address].value),
                        output_value=_display(output_value),
                    )
                )

    recomputed_regression = _rounded_accuracy(
        regression_correct, regression_total, clamp_regression=True
    )
    recomputed_modification = _rounded_accuracy(
        modification_correct, modification_total, clamp_regression=False
    )
    stored_regression = stored_score.get("regression_accuracy")
    stored_modification = stored_score.get("modification_accuracy")
    return AttemptCensus(
        run_name=run_name,
        task=task_key,
        official_exact=stored_score.get("accuracy") == 1.0,
        stored_regression_accuracy=stored_regression,
        stored_modification_accuracy=stored_modification,
        recomputed_regression_accuracy=recomputed_regression,
        recomputed_modification_accuracy=recomputed_modification,
        score_matches_stored=(
            _score_agrees(stored_regression, recomputed_regression)
            and _score_agrees(stored_modification, recomputed_modification)
        ),
        regression_correct=regression_correct,
        regression_total=regression_total,
        modification_correct=modification_correct,
        modification_total=modification_total,
        mismatches=tuple(mismatches),
    )


def _dataset(data_root: Path, category: str) -> dict[str, dict[str, Any]]:
    rows = json.loads((data_root / category / "dataset.json").read_text(encoding="utf-8"))
    return {str(row["id"]): row for row in rows}


def discover_attempts(
    runs_root: Path,
    *,
    include_known_invalid: bool,
    run_names: list[str] | None = None,
) -> tuple[list[tuple[Path, str, str, dict[str, Any]]], list[dict[str, str]]]:
    """Return stored scored attempts and separately documented exclusions."""

    allowed = set(run_names) if run_names else None
    attempts: list[tuple[Path, str, str, dict[str, Any]]] = []
    exclusions: list[dict[str, str]] = []
    for score_path in sorted(runs_root.glob("*/official_scores.json")):
        payload = json.loads(score_path.read_text(encoding="utf-8"))
        run_name = str(payload.get("run_name") or score_path.parent.name)
        if allowed is not None and run_name not in allowed:
            continue
        for task_key, score in sorted(payload.get("tasks", {}).items()):
            if not isinstance(score, dict):
                continue
            reason = invalid_reason(run_name, task_key)
            if reason and not include_known_invalid:
                exclusions.append({"run_name": run_name, "task": task_key, "reason": reason})
                continue
            attempts.append((score_path.parent, run_name, task_key, score))
    return attempts, exclusions


def run_census(
    *,
    runs_root: Path,
    data_root: Path,
    include_known_invalid: bool = False,
    run_names: list[str] | None = None,
) -> tuple[list[AttemptCensus], list[dict[str, str]]]:
    attempts, exclusions = discover_attempts(
        runs_root,
        include_known_invalid=include_known_invalid,
        run_names=run_names,
    )
    datasets: dict[str, dict[str, dict[str, Any]]] = {}
    results: list[AttemptCensus] = []
    for run_root, run_name, task_key, stored_score in attempts:
        category, task_id = task_key.split(":", 1)
        if category not in datasets:
            datasets[category] = _dataset(data_root, category)
        task = datasets[category][task_id]
        task_root = data_root / category
        input_path = task_root / task["spreadsheet_path"]
        golden_path = task_root / task["golden_response_path"]
        output_path = run_root / f"{category}-{task_id}" / "output.xlsx"
        if not output_path.is_file():
            continue

        with_font_color = category == "Debugging" and "Color" in task["spreadsheet_path"]
        with_formula = category == "Debugging" and "Embedded" in task["spreadsheet_path"]
        input_views = _load_views(input_path, with_formula=with_formula)
        golden_views = _load_views(golden_path, with_formula=with_formula)
        output_views = _load_views(output_path, with_formula=with_formula)
        try:
            results.append(
                census_workbooks(
                    run_name=run_name,
                    task_key=task_key,
                    answer_position=task["answer_position"],
                    input_views=input_views,
                    golden_views=golden_views,
                    output_views=output_views,
                    stored_score=stored_score,
                    with_font_color=with_font_color,
                    with_formula=with_formula,
                )
            )
        finally:
            _close_views(input_views)
            _close_views(golden_views)
            _close_views(output_views)
    return results, exclusions


def _print_summary(results: list[AttemptCensus], exclusions: list[dict[str, str]]) -> None:
    mismatches = [cell for attempt in results for cell in attempt.mismatches]
    distinct = {
        (attempt.task, cell.sheet, cell.address, cell.causal_class)
        for attempt in results
        for cell in attempt.mismatches
    }
    print(
        f"{len(results)} valid scored attempts; "
        f"{sum(attempt.official_exact for attempt in results)} official exact"
    )
    print(
        f"{len(mismatches)} mismatch occurrences; "
        f"{len(distinct)} distinct task/cell/class signatures"
    )
    print(f"{len(exclusions)} known-invalid attempt(s) excluded\n")

    print("=== mismatch occurrences by causal class ===")
    for name, count in Counter(cell.causal_class for cell in mismatches).most_common():
        print(f"{count:7}  {count / len(mismatches):6.1%}  {name}")

    print("\n=== attempts with the most mismatches ===")
    for attempt in sorted(results, key=lambda row: len(row.mismatches), reverse=True)[:15]:
        print(
            f"{len(attempt.mismatches):7}  {attempt.run_name:42} {attempt.task}"
        )

    disagreement = [attempt for attempt in results if not attempt.score_matches_stored]
    print(f"\nscore replay disagreements: {len(disagreement)}/{len(results)}")
    for attempt in disagreement:
        print(
            f"  {attempt.run_name} {attempt.task}: "
            f"stored {attempt.stored_regression_accuracy}/{attempt.stored_modification_accuracy}, "
            f"recomputed {attempt.recomputed_regression_accuracy}/"
            f"{attempt.recomputed_modification_accuracy}"
        )


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--runs-root", type=Path, default=DEFAULT_RUNS_ROOT)
    parser.add_argument("--data-root", type=Path, default=DEFAULT_DATA_ROOT)
    parser.add_argument("--output-json", type=Path)
    parser.add_argument(
        "--run-name",
        action="append",
        default=None,
        help="restrict to these run directory names; repeatable",
    )
    parser.add_argument(
        "--include-known-invalid",
        action="store_true",
        help="include attempts with documented harness/output invalidity",
    )
    args = parser.parse_args()
    results, exclusions = run_census(
        runs_root=args.runs_root,
        data_root=args.data_root,
        include_known_invalid=args.include_known_invalid,
        run_names=args.run_name,
    )
    _print_summary(results, exclusions)
    if args.output_json:
        args.output_json.parent.mkdir(parents=True, exist_ok=True)
        args.output_json.write_text(
            json.dumps(
                {
                    "schema_version": 1,
                    "attempts": [asdict(result) for result in results],
                    "excluded_attempts": exclusions,
                },
                indent=2,
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        print(f"\nJSON {args.output_json}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
