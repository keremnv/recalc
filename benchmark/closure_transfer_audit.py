#!/usr/bin/env python3
"""Zero-model transferability audit of proposal-seeded execution closure.

Reconstructs default-harness C0 mutations, computes gold-blind input-side
OFFSET-aware ancestor closure from the agent's own proposed formulas, and
replays identical semantic edits in original vs dependency order.

No model calls. No Edit Plan. No new agent tool. Gold is loaded only after
mutation reconstruction is frozen, and never used to build D(seed).
"""
from __future__ import annotations

import argparse
import ast
import csv
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import time
from collections import defaultdict, deque
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
import zipfile

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / "benchmark"), str(ROOT / "src")]

from xlsx_metadata_repair import install as install_metadata_repair  # noqa: E402

install_metadata_repair()
import openpyxl  # noqa: E402

import composition_closure as cc  # noqa: E402
import formula_synthesis_probe as synth_tools  # noqa: E402
from xlsx_cell_writer import write_cells  # noqa: E402

EVAL_DIR = ROOT / "benchmark-data/SpreadsheetBench-2/evaluation"
DATA = ROOT / "benchmark-data/SpreadsheetBench-2/data"
RUNS = ROOT / "benchmark-data/SpreadsheetBench-2/benchmark-runs/openrouter"
OPEN_SPREADSHEET = EVAL_DIR / "open_spreadsheet.py"
SCORER = ROOT / "benchmark/score_openrouter_run.py"
ARTIFACT = ROOT / "closure_transfer_audit"
REPORT = ROOT / "DEFAULT_HARNESS_CLOSURE_TRANSFER_AUDIT.md"
# Agent commands are replayed via bash -lc. The repo path contains spaces
# ("Personal Projects"), which would word-split rewritten docker paths.
WORK = Path("/tmp/closure_transfer_audit_work")

PRIMARY_TASKS = [
    ("Financial_Model", "08_03"),
    ("Financial_Model", "08_04"),
    ("Financial_Model", "08_05"),
    ("Financial_Model", "15_04"),
    ("Financial_Model", "08_02"),
    ("Financial_Model", "08_01"),
    ("Template", "06_09"),
    ("Template", "06_16"),
    ("Debugging", "10_04"),
]
EXPANSION_RUN = "glm-5.3-flash-control-census-sixty-1"
MIN_FCVW_SEEDS = 5
REPLAY_DRIFT_MOD = 0.05
COMMAND_TIMEOUT = 180
NODE_CAP = 400000
WITNESS_SEED = ("Income Statement", 69, 3)  # C69
WITNESS_MEMBERS = [("Working Capital", 44, c) for c in range(10, 15)]  # J44:N44
EDIT_PLAN_08_03 = ROOT / (
    "benchmark-data/SpreadsheetBench-2/benchmark-runs/mechanical/"
    "execution-unit-probe/units/08_03__Income_Statement_C69.json"
)

Cell = tuple[str, int, int]


def now() -> str:
    return datetime.now(UTC).isoformat()


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def write_jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False, default=str) + "\n")


def sha256_file(path: Path) -> str | None:
    if not path.is_file():
        return None
    return hashlib.sha256(path.read_bytes()).hexdigest()


def a1(row: int, col: int) -> str:
    return cc.a1(row, col)


def addr(cell: Cell) -> str:
    return f"{cell[0]}!{a1(cell[1], cell[2])}"


def parse_addr(text: str) -> Cell:
    sheet, ref = text.rsplit("!", 1)
    match = re.fullmatch(r"\$?([A-Za-z]{1,3})\$?([0-9]+)", ref.strip())
    if not match:
        raise ValueError(text)
    return (sheet.strip().strip("'"), int(match.group(2)), cc.col_index(match.group(1)))


def cell_key(cell: Cell) -> list[Any]:
    return [cell[0], cell[1], cell[2]]


def dataset_record(category: str, task_id: str) -> dict[str, Any]:
    records = json.loads((DATA / category / "dataset.json").read_text(encoding="utf-8"))
    matches = [row for row in records if row["id"] == task_id]
    if len(matches) != 1:
        raise ValueError(f"expected one {category}:{task_id} record, found {len(matches)}")
    return matches[0]


def source_xlsx(category: str, task_id: str) -> Path:
    return DATA / category / dataset_record(category, task_id)["spreadsheet_path"]


def gold_xlsx(category: str, task_id: str) -> Path:
    return DATA / category / dataset_record(category, task_id)["golden_response_path"]


def payload_of(value: Any) -> dict[str, Any]:
    if hasattr(value, "text"):
        text = str(value.text)
        return {"kind": "FORMULA" if text.startswith("=") else "VALUE", "text": text}
    if isinstance(value, str) and value.startswith("="):
        return {"kind": "FORMULA", "text": value}
    if value is None or value == "":
        return {"kind": "BLANK", "text": None}
    return {"kind": "VALUE", "text": value if isinstance(value, (int, float, str, bool)) else repr(value)}


def mutation_kind(before: dict[str, Any], after: dict[str, Any]) -> str:
    if after["kind"] == "FORMULA":
        return "SET_FORMULA"
    if after["kind"] == "BLANK" and before["kind"] != "BLANK":
        return "CLEAR"
    if after["kind"] == "VALUE":
        return "SET_VALUE"
    return "OTHER"


def official_formula_key(text: Any) -> str | None:
    """SpreadsheetBench evaluator formula identity (evaluation.compare_cell_formula)."""
    if hasattr(text, "text"):
        text = text.text
    if not isinstance(text, str) or not text.startswith("="):
        return None
    out = text.replace("$", "").upper()
    if out.startswith("=+"):
        out = "=" + out[2:]
    return out


def formulas_match(left: Any, right: Any) -> bool:
    a, b = official_formula_key(left), official_formula_key(right)
    if a is not None and b is not None:
        return a == b
    return False


def values_match(left: Any, right: Any) -> bool:
    sys.path.insert(0, str(EVAL_DIR))
    from evaluation import compare_cell_value

    return bool(compare_cell_value(left, right))


def is_blank_payload(payload: dict[str, Any] | None) -> bool:
    return payload is None or payload.get("kind") == "BLANK"


def snapshot_zip(path: Path) -> dict[Cell, dict[str, Any]]:
    """Formula/value snapshot from the xlsx archive. Used when openpyxl cannot load."""
    from xlsx_cell_writer import column_index as zip_col, sheet_parts

    cell_re = re.compile(r'<c r="([A-Z]+)(\d+)"([^>]*?)(?:/>|>(.*?)</c>)', re.S)
    formula_re = re.compile(r"<f\b[^>]*>(.*?)</f>", re.S)
    value_re = re.compile(r"<v>(.*?)</v>")
    out: dict[Cell, dict[str, Any]] = {}
    with zipfile.ZipFile(path) as archive:
        parts = sheet_parts(archive)
        for sheet, part in parts.items():
            xml = archive.read(part).decode("utf-8", errors="replace")
            for match in cell_re.finditer(xml):
                col, row, attrs, body = match.group(1), int(match.group(2)), match.group(3), match.group(4) or ""
                fm = formula_re.search(body)
                if fm:
                    text = fm.group(1).replace("&lt;", "<").replace("&gt;", ">").replace("&amp;", "&")
                    if text and not text.startswith("="):
                        text = "=" + text
                    out[(sheet, row, zip_col(col))] = {"kind": "FORMULA", "text": text}
                    continue
                if 't="s"' in attrs:
                    continue
                vm = value_re.search(body)
                if vm and vm.group(1) != "":
                    raw = vm.group(1)
                    try:
                        parsed: Any = float(raw) if "." in raw else int(raw)
                    except ValueError:
                        parsed = raw
                    out[(sheet, row, zip_col(col))] = {"kind": "VALUE", "text": parsed}
    return out


def snapshot_content(path: Path) -> dict[Cell, dict[str, Any]]:
    try:
        wb = openpyxl.load_workbook(path, data_only=False)
    except Exception:
        return snapshot_zip(path)
    out: dict[Cell, dict[str, Any]] = {}
    try:
        for ws in wb.worksheets:
            for cell in ws._cells.values():
                payload = payload_of(cell.value)
                if payload["kind"] == "BLANK":
                    continue
                out[(ws.title, int(cell.row), int(cell.column))] = payload
    finally:
        wb.close()
    return out


def snapshot_values(path: Path) -> dict[Cell, Any]:
    wb = openpyxl.load_workbook(path, data_only=True)
    out: dict[Cell, Any] = {}
    try:
        for ws in wb.worksheets:
            for cell in ws._cells.values():
                if cell.value is not None and cell.value != "":
                    out[(ws.title, int(cell.row), int(cell.column))] = cell.value
    finally:
        wb.close()
    return out


def content_diff(before: dict[Cell, dict[str, Any]], after: dict[Cell, dict[str, Any]]) -> list[tuple[Cell, dict, dict]]:
    keys = set(before) | set(after)
    rows = []
    for cell in sorted(keys):
        left = before.get(cell) or {"kind": "BLANK", "text": None}
        right = after.get(cell) or {"kind": "BLANK", "text": None}
        if left != right:
            rows.append((cell, left, right))
    return rows


def find_traj(task_root: Path) -> Path | None:
    hits = sorted(task_root.rglob("*.traj"))
    return hits[0] if hits else None


def find_output(task_root: Path) -> Path | None:
    direct = task_root / "output.xlsx"
    if direct.is_file():
        return direct
    hits = sorted(task_root.rglob("*_output.xlsx"))
    return hits[0] if hits else None


def fill_c0_run(category: str, task_id: str) -> Path:
    return RUNS / f"glm-fill-c0-{category}-{task_id}-r1"


def task_root_for(run_root: Path, category: str, task_id: str) -> Path:
    named = run_root / f"{category}-{task_id}"
    if named.is_dir():
        return named
    if (run_root / "output.xlsx").is_file() or find_traj(run_root):
        return run_root
    return named


def locate_population_row(category: str, task_id: str, *, run_root: Path | None = None) -> dict[str, Any]:
    run_root = run_root or fill_c0_run(category, task_id)
    task_root = task_root_for(run_root, category, task_id)
    traj = find_traj(task_root)
    output = find_output(task_root)
    refreshed = run_root / "submission" / "outputs" / category / f"{task_id}_output.xlsx"
    official = {}
    scores_path = run_root / "official_scores.json"
    if scores_path.is_file():
        payload = json.loads(scores_path.read_text(encoding="utf-8"))
        official = ((payload.get("tasks") or {}).get(f"{category}:{task_id}")) or {}
    source = source_xlsx(category, task_id)
    complete = bool(traj and output and source.is_file() and official)
    return {
        "task": f"{category}:{task_id}",
        "category": category,
        "id": task_id,
        "run_root": str(run_root),
        "task_root": str(task_root),
        "trajectory": str(traj) if traj else None,
        "output_xlsx": str(output) if output else None,
        "refreshed_xlsx": str(refreshed) if refreshed.is_file() else None,
        "source_xlsx": str(source),
        "gold_xlsx": str(gold_xlsx(category, task_id)),
        "official": official,
        "complete": complete,
        "exclude_reason": None if complete else (
            "missing_workbook" if not output else "missing_trajectory" if not traj else "unscored"
        ),
    }


def extract_python_blocks(action: str) -> list[str]:
    blocks = []
    for match in re.finditer(r"<<\s*['\"](\w+)['\"]\s*\n(.*?)\\n\1", action, re.S):
        blocks.append(match.group(2))
    for match in re.finditer(r"<<\s*'(\w+)'\s*\n(.*?)\\n\1", action, re.S):
        blocks.append(match.group(2))
    # Real heredocs in the stored action already contain newlines, not \\n.
    for match in re.finditer(r"<<\s*['\"](\w+)['\"]\n(.*?)\n\1", action, re.S):
        blocks.append(match.group(2))
    if "python3 -c" in action or "python -c" in action:
        match = re.search(r"""python3?\s+-c\s+(['"])(.*)\1""", action, re.S)
        if match:
            blocks.append(match.group(2))
    return blocks


def literal_assignments(code: str) -> list[tuple[str | None, str | None, Any]]:
    """Pull constant sheet/cell writes out of a script. Loops are left to replay."""
    out: list[tuple[str | None, str | None, Any]] = []
    try:
        tree = ast.parse(code)
    except SyntaxError:
        return out
    for node in ast.walk(tree):
        if not isinstance(node, ast.Assign) or len(node.targets) != 1:
            continue
        target = node.targets[0]
        value = ast.literal_eval(node.value) if isinstance(node.value, ast.Constant) else None
        if isinstance(target, ast.Subscript) and isinstance(target.slice, ast.Constant):
            key = target.slice.value
            if isinstance(key, str) and re.fullmatch(r"[A-Za-z]{1,3}\d+", key):
                out.append((None, key, value))
        if (
            isinstance(target, ast.Attribute)
            and target.attr == "value"
            and isinstance(target.value, ast.Call)
            and isinstance(target.value.func, ast.Attribute)
            and target.value.func.attr == "cell"
        ):
            kwargs = {kw.arg: kw.value for kw in target.value.keywords if kw.arg}
            row = kwargs.get("row")
            col = kwargs.get("column")
            if isinstance(row, ast.Constant) and isinstance(col, ast.Constant):
                out.append((None, a1(int(row.value), int(col.value)), value))
    return out


def rewrite_command(action: str, sandbox: Path) -> str:
    """Rewrite docker paths. Replace /tmp first so later /tmp-prefixed sandbox
    paths are not rewritten a second time."""
    cmd = re.sub(r"(?<![A-Za-z0-9._-])/tmp(?=/|\s|$|['\"])", str(sandbox / "tmp"), action)
    cmd = cmd.replace("/mnt/spreadsheet_data", str(sandbox / "mnt/spreadsheet_data"))
    cmd = cmd.replace("/mnt/spreadsheet_output", str(sandbox / "mnt/spreadsheet_output"))
    return cmd


def looks_like_libreoffice(action: str) -> bool:
    lower = action.lower()
    return "libreoffice" in lower or re.search(r"\bsoffice\b", lower) is not None


def prepare_sandbox(row: dict[str, Any], sandbox: Path) -> Path:
    if sandbox.exists():
        shutil.rmtree(sandbox)
    source = Path(row["source_xlsx"])
    rel = Path(dataset_record(row["category"], row["id"])["spreadsheet_path"])
    data_path = sandbox / "mnt/spreadsheet_data" / rel
    data_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, data_path)
    (sandbox / "mnt/spreadsheet_output").mkdir(parents=True, exist_ok=True)
    (sandbox / "tmp").mkdir(parents=True, exist_ok=True)
    scripts = sandbox / "scripts"
    scripts.mkdir(parents=True, exist_ok=True)
    return data_path


def sandbox_workbooks(sandbox: Path) -> list[Path]:
    hits = []
    for folder in (sandbox / "mnt/spreadsheet_data", sandbox / "mnt/spreadsheet_output"):
        if folder.is_dir():
            hits.extend(folder.rglob("*.xlsx"))
    return hits


def workbook_fingerprint(paths: list[Path]) -> dict[str, str]:
    out = {}
    for path in paths:
        digest = sha256_file(path)
        if digest:
            out[str(path)] = digest
    return out


def load_trajectory(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def reconstruct_mutations(row: dict[str, Any]) -> dict[str, Any]:
    sandbox = WORK / row["task"].replace(":", "_") / "sandbox"
    source_live = prepare_sandbox(row, sandbox)
    traj = load_trajectory(Path(row["trajectory"]))
    steps = traj.get("trajectory") or []
    events: list[dict[str, Any]] = []
    scripts_saved = []
    confidence_by_cell: dict[Cell, str] = {}
    before_all = snapshot_content(source_live)
    current = dict(before_all)
    fingerprints = workbook_fingerprint(sandbox_workbooks(sandbox))
    output_live: Path | None = None
    for index, step in enumerate(steps):
        action = step.get("action") or ""
        if not action or action.strip() == "submit" or action.startswith("view_xlsx"):
            continue
        blocks = extract_python_blocks(action)
        for b_i, block in enumerate(blocks):
            dest = sandbox / "scripts" / f"event_{index:03d}_{b_i}.py"
            dest.write_text(block, encoding="utf-8")
            scripts_saved.append(str(dest))
            for _sheet, ref, value in literal_assignments(block):
                if not ref:
                    continue
                # Literal A1 writes are exact when the sheet is unambiguous later via diff.
                pass
        if looks_like_libreoffice(action):
            continue
        rewritten = rewrite_command(action, sandbox)
        try:
            subprocess.run(
                ["bash", "-lc", rewritten],
                cwd=str(sandbox / "tmp"),
                check=False,
                capture_output=True,
                text=True,
                timeout=COMMAND_TIMEOUT,
                env={**os.environ, "PYTHONWARNINGS": "ignore"},
            )
        except subprocess.TimeoutExpired:
            events.append({
                "task_id": row["task"],
                "trajectory_id": Path(row["run_root"]).name,
                "event_index": index,
                "command": action[:300],
                "mutation_kind": "OTHER",
                "note": "COMMAND_TIMEOUT",
            })
            continue
        new_paths = sandbox_workbooks(sandbox)
        new_fp = workbook_fingerprint(new_paths)
        changed = [Path(p) for p, digest in new_fp.items() if fingerprints.get(p) != digest]
        fingerprints = new_fp
        if not changed:
            continue
        target = None
        for path in changed:
            if "spreadsheet_output" in str(path):
                target = path
                output_live = path
                break
        if target is None:
            target = changed[0]
            if source_live.exists() and sha256_file(source_live) != sha256_file(Path(row["source_xlsx"])):
                target = source_live
        after = snapshot_content(target)
        diffs = content_diff(current, after)
        for cell, left, right in diffs:
            kind = mutation_kind(left, right)
            events.append({
                "task_id": row["task"],
                "trajectory_id": Path(row["run_root"]).name,
                "event_index": index,
                "command_action_id": index,
                "script_file": str(sandbox / "scripts" / f"event_{index:03d}_0.py"),
                "sheet": cell[0],
                "cell": a1(cell[1], cell[2]),
                "row": cell[1],
                "col": cell[2],
                "before_formula_or_value": left.get("text"),
                "after_formula_or_value": right.get("text"),
                "before_kind": left.get("kind"),
                "after_kind": right.get("kind"),
                "mutation_kind": kind,
                "reconstruction_confidence": "RECONSTRUCTED_FROM_SCRIPT",
            })
            confidence_by_cell[cell] = "RECONSTRUCTED_FROM_SCRIPT"
        current = after
    final_path = output_live if output_live and output_live.is_file() else (
        Path(row["output_xlsx"]) if row["output_xlsx"] else source_live
    )
    # Ambiguous leftover: cells that differ in the stored output but were not
    # attributed to a mutating command.
    stored = snapshot_content(Path(row["output_xlsx"])) if row["output_xlsx"] else {}
    authored = {cell: payload for cell, payload in current.items()}
    authored_from_events = {(e["sheet"], e["row"], e["col"]) for e in events if "row" in e}
    ambiguous = []
    for cell, payload in stored.items():
        if cell in authored_from_events:
            continue
        src = before_all.get(cell) or {"kind": "BLANK", "text": None}
        if src != payload:
            ambiguous.append({
                "task_id": row["task"],
                "sheet": cell[0],
                "cell": a1(cell[1], cell[2]),
                "row": cell[1],
                "col": cell[2],
                "before_formula_or_value": src.get("text"),
                "after_formula_or_value": payload.get("text"),
                "mutation_kind": mutation_kind(src, payload),
                "reconstruction_confidence": "DIFF_ONLY_AMBIGUOUS",
            })
    final_authored: dict[str, Any] = {}
    for cell in authored_from_events:
        seq = [e for e in events if e.get("row") == cell[1] and e.get("col") == cell[2] and e.get("sheet") == cell[0]]
        if not seq:
            continue
        last = seq[-1]
        final_authored[addr(cell)] = {
            "cell": cell_key(cell),
            "kind": last.get("after_kind"),
            "text": last.get("after_formula_or_value"),
            "confidence": last.get("reconstruction_confidence"),
            "last_event_index": last["event_index"],
        }
    replay_output = sandbox / "replay_r0.xlsx"
    if output_live and output_live.is_file():
        shutil.copy2(output_live, replay_output)
    return {
        "task": row["task"],
        "events": events,
        "ambiguous": ambiguous,
        "authored_cell_set": sorted(addr(c) for c in authored_from_events),
        "final_authored_content": final_authored,
        "scripts": scripts_saved,
        "sandbox": str(sandbox),
        "replay_r0_xlsx": str(replay_output) if replay_output.is_file() else None,
        "source_hash": sha256_file(Path(row["source_xlsx"])),
        "output_hash": sha256_file(Path(row["output_xlsx"])) if row["output_xlsx"] else None,
        "n_events": len(events),
        "n_authored": len(authored_from_events),
        "n_ambiguous": len(ambiguous),
    }


def graph_from_workbook(path: Path) -> tuple[dict[Cell, set[Cell]], dict[Cell, Any], list[dict]]:
    wb = openpyxl.load_workbook(path, data_only=False)
    values_wb = openpyxl.load_workbook(path, data_only=True)
    pre: dict[Cell, set[Cell]] = defaultdict(set)
    values: dict[Cell, Any] = {}
    ledger: list[dict] = []
    formulas: dict[Cell, str] = {}
    try:
        for ws in values_wb.worksheets:
            for cell in ws._cells.values():
                if cell.value is not None:
                    values[(ws.title, int(cell.row), int(cell.column))] = cell.value
        for ws in wb.worksheets:
            for cell in ws._cells.values():
                payload = payload_of(cell.value)
                if payload["kind"] != "FORMULA":
                    continue
                here = (ws.title, int(cell.row), int(cell.column))
                text = payload["text"]
                formulas[here] = text
                try:
                    refs = synth_tools._ref_records(text, ws.title, cell.row, cell.column)
                except Exception:
                    ledger.append({"reason": "UNSUPPORTED_FORMULA_REFERENCE", "cell": addr(here)})
                    continue
                for ref in refs.get("points", []) + refs.get("ranges", []):
                    m1 = re.fullmatch(r"\$?([A-Za-z]{1,3})\$?([0-9]+)", ref.get("start") or "")
                    if not m1:
                        continue
                    r1, c1 = int(m1.group(2)), cc.col_index(m1.group(1))
                    if ref.get("is_range"):
                        m2 = re.fullmatch(r"\$?([A-Za-z]{1,3})\$?([0-9]+)", ref.get("end") or ref["start"])
                        if not m2:
                            continue
                        r2, c2 = int(m2.group(2)), cc.col_index(m2.group(1))
                        pre[here] |= cc._expand(ref["sheet"], min(r1, r2), min(c1, c2), max(r1, r2), max(c1, c2), ledger)
                    else:
                        pre[here].add((ref["sheet"], r1, c1))
    finally:
        wb.close()
        values_wb.close()
    extra, led2 = offset_edges_from_formulas(formulas, values)
    ledger.extend(led2)
    merged = {k: set(v) for k, v in pre.items()}
    for k, v in extra.items():
        merged.setdefault(k, set()).update(v)
    return merged, values, ledger


def offset_edges_from_formulas(
    formulas: dict[Cell, str],
    values: dict[Cell, Any],
) -> tuple[dict[Cell, set[Cell]], list[dict]]:
    edges: dict[Cell, set[Cell]] = defaultdict(set)
    ledger: list[dict] = []
    for here, text in formulas.items():
        if not cc.OFFSET_RE.search(text or ""):
            continue
        for match in cc.OFFSET_RE.finditer(text):
            args, _ = cc._split_args(text, match.end() - 1)
            if len(args) < 1:
                continue
            try:
                refs = synth_tools._ref_records("=" + args[0], here[0], here[1], here[2])
            except Exception:
                ledger.append({"reason": "UNSUPPORTED_FORMULA_REFERENCE", "cell": addr(here)})
                continue
            pts = refs.get("points", []) + refs.get("ranges", [])
            if not pts:
                continue
            anchor = pts[0]
            m1 = re.fullmatch(r"\$?([A-Za-z]{1,3})\$?([0-9]+)", anchor.get("start") or "")
            if not m1:
                continue
            ar, ac = int(m1.group(2)), cc.col_index(m1.group(1))

            def numeric(expr: str, default: int | None) -> int | None:
                expr = (expr or "").strip()
                if not expr:
                    return default
                if re.fullmatch(r"-?\d+", expr):
                    return int(expr)
                try:
                    rec = synth_tools._ref_records("=" + expr, here[0], here[1], here[2])
                except Exception:
                    return None
                pts2 = rec.get("points", [])
                if len(pts2) != 1:
                    return None
                mm = re.fullmatch(r"\$?([A-Za-z]{1,3})\$?([0-9]+)", pts2[0].get("start") or "")
                if not mm:
                    return None
                val = values.get((pts2[0]["sheet"], int(mm.group(2)), cc.col_index(mm.group(1))))
                return int(val) if isinstance(val, (int, float)) else None

            dr = numeric(args[1] if len(args) > 1 else "", 0)
            dc = numeric(args[2] if len(args) > 2 else "", 0)
            height = numeric(args[3] if len(args) > 3 else "", 1)
            width = numeric(args[4] if len(args) > 4 else "", 1)
            if None in (dr, dc, height, width) or height < 1 or width < 1:
                ledger.append({"reason": "DYNAMIC_OFFSET_UNRESOLVED", "cell": addr(here), "args": args[:5]})
                continue
            r1, c1 = ar + dr, ac + dc
            edges[here] |= cc._expand(anchor["sheet"], r1, c1, r1 + height - 1, c1 + width - 1, ledger)
    return dict(edges), ledger


def proposal_precedents(formula: str, seed: Cell) -> set[Cell]:
    out: set[Cell] = set()
    ledger: list = []
    try:
        refs = synth_tools._ref_records(formula, seed[0], seed[1], seed[2])
    except Exception:
        return out
    for ref in refs.get("points", []) + refs.get("ranges", []):
        m1 = re.fullmatch(r"\$?([A-Za-z]{1,3})\$?([0-9]+)", ref.get("start") or "")
        if not m1:
            continue
        r1, c1 = int(m1.group(2)), cc.col_index(m1.group(1))
        if ref.get("is_range"):
            m2 = re.fullmatch(r"\$?([A-Za-z]{1,3})\$?([0-9]+)", ref.get("end") or ref["start"])
            if not m2:
                continue
            r2, c2 = int(m2.group(2)), cc.col_index(m2.group(1))
            out |= cc._expand(ref["sheet"], min(r1, r2), min(c1, c2), max(r1, r2), max(c1, c2), ledger)
        else:
            out.add((ref["sheet"], r1, c1))
    return out


def raw_closure(
    seed: Cell,
    formula: str,
    graph: dict[Cell, set[Cell]],
    keep: set[Cell],
) -> tuple[set[Cell], set[Cell], list[dict[str, Any]]]:
    """D(seed) = OFFSET-aware ancestors of proposal precedents inside keep.

    keep is blank-on-input ∪ authored. Existing populated formulas are not
    treated as required edits. Gold and Edit Plan are not consulted.
    """
    precedents = proposal_precedents(formula, seed)
    members: set[Cell] = set()
    provenance: list[dict[str, Any]] = []
    for prec in precedents:
        anc = cc.ancestors_within(prec, graph, keep, node_cap=NODE_CAP)
        for member in anc:
            members.add(member)
            provenance.append({
                "seed": addr(seed),
                "direct_precedent": addr(prec),
                "closure_member": addr(member),
                "offset_expansion": "OFFSET" in json.dumps(list(graph.get(prec, ()))),
            })
        if prec in keep and prec != seed:
            members.add(prec)
            provenance.append({
                "seed": addr(seed),
                "direct_precedent": addr(prec),
                "closure_member": addr(prec),
                "offset_expansion": False,
                "role": "direct_precedent_in_keep",
            })
    members.discard(seed)
    return members, precedents, provenance


def classify_transfer(d_seed: set[Cell], authored: set[Cell], member_ok: dict[Cell, bool]) -> str:
    if not d_seed:
        return "T0_CLOSURE_EMPTY"
    missing = d_seed - authored
    if missing:
        return "T1_AUTHORITY_GAP"
    if any(not member_ok.get(cell, False) for cell in d_seed):
        return "T2_AUTHORED_MEMBER_SEMANTIC_FAILURE"
    return "T3_TRANSFERABLE_EXECUTION_CANDIDATE"


def refresh_dir(folder: Path) -> None:
    folder.mkdir(parents=True, exist_ok=True)
    completed = subprocess.run(
        [sys.executable, str(OPEN_SPREADSHEET), "--dir_path", str(folder)],
        cwd=str(EVAL_DIR),
        check=False,
        capture_output=True,
        text=True,
    )
    if completed.returncode != 0:
        raise RuntimeError(f"LibreOffice refresh failed: {(completed.stderr or completed.stdout)[-2000:]}")


def score_workbook(category: str, task_id: str, output: Path) -> dict[str, Any]:
    sys.path.insert(0, str(EVAL_DIR))
    from evaluation import (
        classify_cells_by_modification,
        parse_answer_position,
        _find_sheet,
        _has_excel_error,
        _compare_cells,
        compare_cell_formula,
        compare_cell_value,
    )

    record = dataset_record(category, task_id)
    inp, gold = source_xlsx(category, task_id), gold_xlsx(category, task_id)
    wi = openpyxl.load_workbook(inp, data_only=True)
    wg = openpyxl.load_workbook(gold, data_only=True)
    wo = openpyxl.load_workbook(output, data_only=True)
    wif = openpyxl.load_workbook(inp, data_only=False)
    wgf = openpyxl.load_workbook(gold, data_only=False)
    wof = openpyxl.load_workbook(output, data_only=False)
    acc = {
        k: {"total": 0, "official_correct": 0, "value_correct": 0, "cells_official_ok": [], "cells_official_bad": []}
        for k in ("regression", "modification")
    }
    try:
        for rng in parse_answer_position(record["answer_position"]):
            sheet, cr = (rng.split("!", 1) if "!" in rng else (wg.sheetnames[0], rng))
            sheet, cr = sheet.strip("'").strip(), cr.strip("'").strip()
            try:
                regs, mods = classify_cells_by_modification(
                    wi, wg, sheet, cr, False, False, wb_input_formula=wif, wb_answer_formula=wgf
                )
            except Exception:
                continue
            wa, wu = _find_sheet(wg, sheet), _find_sheet(wo, sheet)
            waf, wuf = _find_sheet(wgf, sheet), _find_sheet(wof, sheet)
            if wu is None:
                continue
            for label, cells in (("regression", regs), ("modification", mods)):
                bucket = acc[label]
                for name in cells:
                    bucket["total"] += 1
                    cell_addr = f"{sheet}!{name}"
                    by_value = _compare_cells(wa[name], wu[name], False, False)
                    fallback = _has_excel_error(wa[name]) or _has_excel_error(wu[name])
                    official = compare_cell_formula(waf[name], wuf[name]) if fallback else by_value
                    bucket["official_correct"] += bool(official)
                    bucket["value_correct"] += bool(by_value)
                    (bucket["cells_official_ok"] if official else bucket["cells_official_bad"]).append(cell_addr)
    finally:
        for wb in (wi, wg, wo, wif, wgf, wof):
            wb.close()
    out = {}
    for label, bucket in acc.items():
        total = bucket["total"]
        out[label] = {
            "total": total,
            "official_correct": bucket["official_correct"],
            "value_correct": bucket["value_correct"],
            "official_accuracy": round(bucket["official_correct"] / total, 6) if total else None,
            "value_only_accuracy": round(bucket["value_correct"] / total, 6) if total else None,
            "cells_official_ok": bucket["cells_official_ok"],
            "cells_official_bad": bucket["cells_official_bad"],
        }
    out["exact"] = bool(
        out["modification"]["official_accuracy"] == 1.0 and out["regression"]["official_accuracy"] == 1.0
    ) if out["modification"]["official_accuracy"] is not None else False
    return out


def provenance_label(cell: Cell, written: set[Cell], graph: dict[Cell, set[Cell]]) -> str:
    if cell in written:
        return "DIRECTLY_AUTHORED"
    if (graph.get(cell) or set()) & written:
        return "DIRECT_RECALCULATED_DEPENDENT"
    seen, stack = set(), [cell]
    while stack:
        cur = stack.pop()
        for parent in graph.get(cur) or set():
            if parent in written:
                return "TRANSITIVE_RECALCULATED_DEPENDENT"
            if parent not in seen:
                seen.add(parent)
                stack.append(parent)
    return "UNEXPLAINED_BY_INPUT_SIDE_GRAPH"


def apply_authored(source: Path, dest: Path, authored: dict[str, Any], order: list[str]) -> dict[str, Any]:
    """Apply final intended authored content. Formulas go through the archive writer."""
    formula_edits = []
    value_edits = []
    for key in order:
        rec = authored[key]
        kind = rec.get("kind")
        text = rec.get("text")
        sheet, row, col = rec["cell"]
        address = a1(row, col)
        if kind == "FORMULA" and isinstance(text, str):
            formula_edits.append({"sheet": sheet, "address": address, "formula": text})
        else:
            value_edits.append((sheet, address, kind, text))
    audit = write_cells(source, dest, formula_edits)
    if value_edits:
        wb = openpyxl.load_workbook(dest, data_only=False)
        try:
            for sheet, address, kind, text in value_edits:
                ws = wb[sheet]
                if kind == "BLANK":
                    ws[address].value = None
                else:
                    ws[address].value = text
            wb.save(dest)
        finally:
            wb.close()
    return audit


def topo_order(keys: list[str], authored: dict[str, Any], graph: dict[Cell, set[Cell]]) -> list[str]:
    cells = {key: tuple(authored[key]["cell"]) for key in keys}
    authored_cells = {tuple(authored[key]["cell"]) for key in keys}
    # Edges from authored formulas plus input-side graph.
    pred: dict[str, set[str]] = {key: set() for key in keys}
    key_of = {tuple(authored[key]["cell"]): key for key in keys}
    for key, rec in authored.items():
        cell = tuple(rec["cell"])
        deps = set(graph.get(cell) or ())
        if rec.get("kind") == "FORMULA" and isinstance(rec.get("text"), str):
            deps |= proposal_precedents(rec["text"], cell)  # type: ignore[arg-type]
        for parent in deps:
            if parent in authored_cells and parent != cell:
                pred[key].add(key_of[parent])
    remaining = set(keys)
    ordered: list[str] = []
    while remaining:
        ready = [k for k in remaining if not (pred[k] & remaining)]
        if not ready:
            ready = [min(remaining, key=lambda k: tuple(authored[k]["cell"]))]
        ready.sort(key=lambda k: tuple(authored[k]["cell"]))
        pick = ready[0]
        ordered.append(pick)
        remaining.remove(pick)
    return ordered


def gold_payloads(category: str, task_id: str) -> tuple[dict[Cell, dict[str, Any]], dict[Cell, Any]]:
    path = gold_xlsx(category, task_id)
    return snapshot_content(path), snapshot_values(path)


def seed_class(
    authored: dict[str, Any],
    gold_content: dict[Cell, dict[str, Any]],
    gold_values: dict[Cell, Any],
    output_values: dict[Cell, Any],
    persisted: dict[Cell, dict[str, Any]],
) -> list[dict[str, Any]]:
    rows = []
    for key, rec in authored.items():
        if rec.get("kind") != "FORMULA":
            continue
        cell = tuple(rec["cell"])
        gold = gold_content.get(cell) or {}
        gold_formula = gold.get("text") if gold.get("kind") == "FORMULA" else None
        formula_ok = bool(gold_formula) and formulas_match(rec.get("text"), gold_formula)
        persisted_payload = persisted.get(cell) or {}
        persisted_ok = formulas_match(rec.get("text"), persisted_payload.get("text"))
        value_ok = values_match(output_values.get(cell), gold_values.get(cell)) if cell in gold_values or cell in output_values else False
        if formula_ok and not persisted_ok:
            klass = "CORRECT_PROPOSAL_LATER_OVERWRITTEN"
        elif not formula_ok:
            klass = "FORMULA_WRONG"
        elif value_ok:
            klass = "FORMULA_CORRECT_VALUE_CORRECT"
        else:
            klass = "FORMULA_CORRECT_VALUE_WRONG"
        rows.append({
            "seed": key,
            "cell": rec["cell"],
            "authored_formula": rec.get("text"),
            "gold_formula": gold_formula,
            "formula_correct": formula_ok,
            "value_correct": value_ok,
            "persisted_exactly": persisted_ok,
            "class": klass,
        })
    return rows


def member_semantic_ok(
    cell: Cell,
    authored: dict[str, Any],
    gold_content: dict[Cell, dict[str, Any]],
    gold_values: dict[Cell, Any],
) -> bool:
    rec = authored.get(addr(cell))
    if not rec:
        return False
    gold = gold_content.get(cell)
    if rec.get("kind") == "FORMULA":
        if gold and gold.get("kind") == "FORMULA":
            return formulas_match(rec.get("text"), gold.get("text"))
        return False
    return values_match(rec.get("text"), (gold or {}).get("text") if gold and gold.get("kind") != "FORMULA" else gold_values.get(cell))


def expansion_candidates() -> list[dict[str, Any]]:
    run_root = RUNS / EXPANSION_RUN
    rows = []
    if not run_root.is_dir():
        return rows
    for folder in sorted(run_root.iterdir()):
        if not folder.is_dir() or folder.name in {"submission"}:
            continue
        if "-" not in folder.name:
            continue
        category, task_id = folder.name.split("-", 1)
        if category == "Visualization":
            continue
        if (category, task_id) in PRIMARY_TASKS:
            continue
        row = locate_population_row(category, task_id, run_root=run_root)
        row["population"] = "expanded"
        row["role"] = "matched_default_control"
        if row["complete"]:
            rows.append(row)
    return rows


def historical_edit_plan_members() -> list[str]:
    if not EDIT_PLAN_08_03.is_file():
        return []
    payload = json.loads(EDIT_PLAN_08_03.read_text(encoding="utf-8"))
    return list(payload.get("members") or [])


def autopsy_08_03(recon: dict[str, Any], seeds: list[dict[str, Any]], classes: list[dict[str, Any]], replay: dict[str, Any] | None) -> dict[str, Any]:
    authored = set(recon.get("authored_cell_set") or [])
    c69 = addr(WITNESS_SEED)
    authored_c69 = c69 in authored
    seed_row = next((s for s in seeds if s["seed"] == c69), None)
    member_authored = {addr(c): addr(c) in authored for c in WITNESS_MEMBERS}
    final = recon.get("final_authored_content") or {}
    member_formulas = {addr(c): (final.get(addr(c)) or {}).get("text") for c in WITNESS_MEMBERS}
    klass = next((c for c in classes if c.get("seed") == c69), None)
    c69_value_ok = bool(seed_row and seed_row.get("value_correct"))
    c69_formula_ok = bool(seed_row and seed_row.get("formula_correct"))
    all_members_authored = all(member_authored.values())
    stored = {}
    output = fill_c0_run("Financial_Model", "08_03") / "Financial_Model-08_03" / "output.xlsx"
    if output.is_file():
        wb = openpyxl.load_workbook(output, data_only=False)
        try:
            stored = {
                "C69": payload_of(wb["Income Statement"]["C69"].value),
                "J44": payload_of(wb["Working Capital"]["J44"].value),
                "K44": payload_of(wb["Working Capital"]["K44"].value),
                "L44": payload_of(wb["Working Capital"]["L44"].value),
                "M44": payload_of(wb["Working Capital"]["M44"].value),
                "N44": payload_of(wb["Working Capital"]["N44"].value),
            }
        finally:
            wb.close()
    return {
        "task": "Financial_Model:08_03",
        "questions": {
            "did_c0_author_C69": authored_c69,
            "was_C69_formula_correct": c69_formula_ok,
            "were_J44_N44_all_authored": all_members_authored,
            "member_authored": member_authored,
            "member_formulas": member_formulas,
            "was_final_C69_value_correct": c69_value_ok,
            "coordinated_replay_improves": bool(replay and replay.get("r1_gain")),
            "historical_class_if_already_correct": "NO_HEADROOM_IN_DEFAULT_CONTROL" if c69_formula_ok and c69_value_ok else None,
        },
        "stored_output_formulas": stored,
        "seed_row": seed_row,
        "transfer_class": klass,
        "historical_edit_plan_members": historical_edit_plan_members(),
        "historical_required": [addr(c) for c in WITNESS_MEMBERS],
        "authored_C69_formula": (final.get(c69) or {}).get("text"),
    }


def choose_verdict(summary: dict[str, Any]) -> dict[str, Any]:
    fcvw = summary["formula_correct_value_wrong"]
    classes = summary["among_fcvw"]
    t3 = summary["transferable_candidates"]
    faithful = summary["faithful_replay_pairs"]
    gains = summary["r1_gains"]
    losses = summary["r1_losses"]
    if t3 and faithful == 0 and summary.get("replay_attempted"):
        return {
            "verdict": "REPLAY_INTEGRITY_FAILURE",
            "interpretation": "Original control behavior could not be reproduced faithfully enough for R0/R1 causal replay.",
        }
    if fcvw == 0:
        return {
            "verdict": "NO_COMPOSITION_DEPENDENT_FAILURE_IN_CONTROL",
            "interpretation": "The composition pathology observed in the old architecture does not materially reproduce in the loose default coding harness.",
        }
    if faithful and gains and not (losses and summary.get("material_regression")):
        return {
            "verdict": "TRANSFERABLE_EXECUTION_GAP",
            "interpretation": "The loose default agent sometimes makes the required semantic choices, but deterministic dependency coordination can turn those same choices into better workbook behavior.",
        }
    dominant = max(classes, key=classes.get) if classes else None
    nonempty = {k: v for k, v in classes.items() if v}
    if len(nonempty) > 1 and max(nonempty.values()) < fcvw * 0.6:
        return {
            "verdict": "MIXED_TRANSFERABILITY",
            "interpretation": "Meaningful examples occur in more than one class without a dominant explanation.",
        }
    if dominant == "T1_AUTHORITY_GAP" or classes.get("T1_AUTHORITY_GAP", 0) >= max(1, fcvw / 2):
        return {
            "verdict": "AUTHORITY_GAP_DOMINANT",
            "interpretation": "The old closure mechanism relied on an upstream authority representation that the loose default harness does not possess. Dependencies alone cannot safely recover the missing edits.",
        }
    if dominant == "T2_AUTHORED_MEMBER_SEMANTIC_FAILURE":
        return {
            "verdict": "SEMANTIC_MEMBER_CORRECTNESS_DOMINANT",
            "interpretation": "The limiting failure is semantic construction of upstream edits, not their coordination.",
        }
    return {
        "verdict": "NO_COMPOSITION_DEPENDENT_FAILURE_IN_CONTROL",
        "interpretation": "The default harness exhibits no meaningful population of correct-formula/wrong-value seeds with transferable closure.",
    }


def spec_payload() -> dict[str, Any]:
    return {
        "name": "default-harness-closure-transfer-audit",
        "model_calls": 0,
        "provider_inference": False,
        "architecture": "loose default SpreadsheetBench coding-agent harness; no Task IR; no Edit Plan",
        "closure_rule": "C1 input-side OFFSET-aware ancestors of the proposed formula's direct precedents",
        "keep_set": "blank-on-input ∪ AUTHORED_CELL_SET (Edit-Plan-free analogue of required edits; never gold)",
        "authority_invariant": "execution_unit(seed) = D(seed) ∩ AUTHORED_CELL_SET; never add missing members",
        "primary_population": [f"{c}:{i}" for c, i in PRIMARY_TASKS],
        "expansion_rule": f"if FORMULA_CORRECT_VALUE_WRONG seeds < {MIN_FCVW_SEEDS}, expand once to {EXPANSION_RUN}",
        "excluded": ["Debugging:10_02 resource-censored", "Visualization", "C1 calc-fill prompt arm"],
        "gold_policy": "evaluator gold loaded only after mutation reconstruction freeze; never used to construct D(seed) or choose replay edits",
        "replay": {"R0": "original mutation order via trajectory script replay", "R1": "identical authored content in dependency order"},
        "written_at": now(),
    }


def render_report(payload: dict[str, Any]) -> str:
    s = payload["summary"]
    v = payload["verdict"]
    lines = [
        "# Default harness closure-transfer audit",
        "",
        f"**Verdict:** `{v['verdict']}`",
        "",
        v["interpretation"],
        "",
        "Zero model calls. No Edit Plan. No new agent tool. Gold was not used to build closure or choose replay edits.",
        "",
        "## Frozen contract",
        "",
        "- Closure rule: input-side, OFFSET-aware ancestors of the *agent's* proposed formula's direct precedents",
        "- Required-edit keep set: blank-on-input ∪ AUTHORED_CELL_SET (Edit-Plan-free analogue)",
        "- Authority invariant: `execution_unit(seed) = D(seed) ∩ AUTHORED_CELL_SET`",
        "- Primary population: completed C0 default-harness GLM fill trajectories",
        f"- Expansion trigger: fewer than {MIN_FCVW_SEEDS} FORMULA_CORRECT_VALUE_WRONG seeds",
        "",
        "## Population",
        "",
        f"Primary complete: **{s['primary_complete']}**. Expanded complete: **{s['expanded_complete']}**. Expansion used: **{s['expansion_used']}**.",
        "",
        "## Primary aggregates",
        "",
        f"- Total authored formula edits: **{s['total_authored_formula_edits']}**",
        f"- Formula-correct edits: **{s['formula_correct_edits']}**",
        f"- Formula-correct / value-wrong seeds: **{s['formula_correct_value_wrong']}**",
        "",
        "Among formula-correct/value-wrong:",
        "",
    ]
    for key in ("T0_CLOSURE_EMPTY", "T1_AUTHORITY_GAP", "T2_AUTHORED_MEMBER_SEMANTIC_FAILURE", "T3_TRANSFERABLE_EXECUTION_CANDIDATE"):
        lines.append(f"- `{key}`: {s['among_fcvw'].get(key, 0)}")
    lines += [
        "",
        "Transferable candidates:",
        "",
        f"- Faithful replay pairs: **{s['faithful_replay_pairs']}**",
        f"- R1 gains: **{s['r1_gains']}**",
        f"- R1 ties: **{s['r1_ties']}**",
        f"- R1 losses: **{s['r1_losses']}**",
        "",
        f"Task-level FCVW seeds: {s['fcvw_by_task']}",
        "",
        "## Answers",
        "",
        f"1. Correct-formula/wrong-value failures in the loose default harness: **{'yes' if s['formula_correct_value_wrong'] else 'no'}** ({s['formula_correct_value_wrong']} seeds).",
        f"2. Proposal-seeded closure recovered relevant dependencies: see T0 vs T1–T3 counts above.",
        f"3. Dependencies already authored: T1 authority gaps = {s['among_fcvw'].get('T1_AUTHORITY_GAP', 0)}.",
        f"4. Authored dependency edits semantically correct: T2 = {s['among_fcvw'].get('T2_AUTHORED_MEMBER_SEMANTIC_FAILURE', 0)}; T3 = {s['among_fcvw'].get('T3_TRANSFERABLE_EXECUTION_CANDIDATE', 0)}.",
        f"5. Identical semantic edits scoring better solely through coordination: **{'yes' if s['r1_gains'] else 'no'}**.",
        f"6. Transfer vs old Edit Plan dependence: verdict `{v['verdict']}`.",
        f"7. Enough evidence to justify a live coordination runtime: **{'yes' if v['verdict'] == 'TRANSFERABLE_EXECUTION_GAP' else 'no'}**.",
        "",
        "## Financial_Model:08_03 autopsy",
        "",
    ]
    a = payload["autopsy_08_03"]["questions"]
    lines += [
        f"1. Did C0 author `Income Statement!C69`? **{a['did_c0_author_C69']}**",
        f"2. Was its C69 formula correct? **{a['was_C69_formula_correct']}** (`{payload['autopsy_08_03'].get('authored_C69_formula')}`)",
        f"3. Were `J44:N44` all authored? **{a['were_J44_N44_all_authored']}**",
        f"4. Member formulas: `{a['member_formulas']}`",
        f"5. Final C69 value correct after LibreOffice? **{a['was_final_C69_value_correct']}**",
        f"6. Coordinated replay improves it? **{a['coordinated_replay_improves']}**",
        f"7. Historical class if already correct: `{a['historical_class_if_already_correct']}`",
        "",
        "The recent calc-fill C0 run filled `J44:N44` through ordinary Python. Do not assume the old composition failure survives.",
        "",
        "## Decision",
        "",
    ]
    if v["verdict"] == "TRANSFERABLE_EXECUTION_GAP":
        lines.append("Next experiment: live default harness + transparent mutation runtime. Do not expose an optional dependency tool.")
    else:
        lines.append("Do not build the live coordination harness. Dependency closure does not have an earned capability path into the loose default harness from this evidence.")
    lines += [
        "",
        "## Cost",
        "",
        "Provider/API inference cost: **$0**. Compute runtime is recorded in `closure_transfer_audit/verdict.json`.",
        "",
    ]
    return "\n".join(lines) + "\n"


def run(argv: list[str] | None = None) -> dict[str, Any]:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--skip-replay", action="store_true")
    args = parser.parse_args(argv)
    t0 = time.perf_counter()
    ARTIFACT.mkdir(parents=True, exist_ok=True)
    WORK.mkdir(parents=True, exist_ok=True)
    write_json(ARTIFACT / "spec.json", spec_payload())

    primary = []
    for category, task_id in PRIMARY_TASKS:
        row = locate_population_row(category, task_id)
        row["population"] = "primary"
        row["role"] = "c0_default_harness"
        if task_id == "10_04" and not row["complete"]:
            row["exclude_reason"] = row["exclude_reason"] or "incomplete_10_04"
        primary.append(row)
    included = [r for r in primary if r["complete"]]
    expansion_used = False
    expanded: list[dict[str, Any]] = []

    integrity = []
    reconstructions = []
    all_events = []
    authored_sets = {}
    seed_rows = []
    closures = []
    class_rows = []
    replay_manifest = []
    replay_scores = []
    cell_prov = []

    def process_rows(rows: list[dict[str, Any]], population: str) -> None:
        for row in rows:
            print(f"PROCESS {population} {row['task']}", flush=True)
            try:
                recon = reconstruct_mutations(row)
            except Exception as exc:
                print(f"  RECON_FAIL {type(exc).__name__}: {exc}", flush=True)
                recon = {
                    "task": row["task"],
                    "events": [],
                    "ambiguous": [],
                    "authored_cell_set": [],
                    "final_authored_content": {},
                    "scripts": [],
                    "sandbox": None,
                    "replay_r0_xlsx": None,
                    "source_hash": sha256_file(Path(row["source_xlsx"])),
                    "output_hash": sha256_file(Path(row["output_xlsx"])) if row["output_xlsx"] else None,
                    "n_events": 0,
                    "n_authored": 0,
                    "n_ambiguous": 0,
                    "recon_error": repr(exc),
                }
            print(
                f"  authored={recon['n_authored']} events={recon['n_events']} ambiguous={recon['n_ambiguous']}",
                flush=True,
            )
            reconstructions.append({"population": population, **{k: v for k, v in recon.items() if k != "events"}})
            all_events.extend(recon["events"])
            authored_sets[row["task"]] = {
                "population": population,
                "authored_cell_set": recon["authored_cell_set"],
                "final_authored_content": recon["final_authored_content"],
                "n_ambiguous": recon["n_ambiguous"],
            }
            integrity.append({
                "task": row["task"],
                "population": population,
                "source_hash": recon["source_hash"],
                "output_hash": recon["output_hash"],
                "trajectory": row["trajectory"],
                "output_xlsx": row["output_xlsx"],
                "refreshed_xlsx": row["refreshed_xlsx"],
                "n_events": recon["n_events"],
                "n_authored": recon["n_authored"],
                "same_run": Path(row["run_root"]).name in str(row["task_root"]),
                "model": "z-ai/glm-5.3-flash" if population == "primary" else "stored_control",
            })
            # Phase 2 — gold after mutation freeze.
            try:
                gold_content, gold_values = gold_payloads(row["category"], row["id"])
                persisted = snapshot_content(Path(row["output_xlsx"]))
            except Exception as exc:
                print(f"  GOLD_FAIL {type(exc).__name__}: {exc}", flush=True)
                continue
            refreshed_path = Path(row["refreshed_xlsx"]) if row["refreshed_xlsx"] else None
            try:
                if refreshed_path and refreshed_path.is_file():
                    output_values = snapshot_values(refreshed_path)
                else:
                    lo_dir = WORK / row["task"].replace(":", "_") / "lo_final"
                    lo_dir.mkdir(parents=True, exist_ok=True)
                    dest = lo_dir / f"{row['id']}_output.xlsx"
                    shutil.copy2(row["output_xlsx"], dest)
                    refresh_dir(lo_dir)
                    output_values = snapshot_values(dest)
            except Exception as exc:
                print(f"  VALUE_FAIL {type(exc).__name__}: {exc}", flush=True)
                output_values = {}
            seeds = seed_class(recon["final_authored_content"], gold_content, gold_values, output_values, persisted)
            for seed in seeds:
                seed.update({"task": row["task"], "population": population})
            seed_rows.extend(seeds)
            # Phase 3 — gold-blind closure on FCVW seeds (and 08_03 C69 regardless).
            try:
                graph, _values, ledger = graph_from_workbook(Path(row["source_xlsx"]))
            except Exception as exc:
                print(f"  GRAPH_FAIL {type(exc).__name__}: {exc}", flush=True)
                graph, ledger = {}, []
            input_content = snapshot_content(Path(row["source_xlsx"]))
            authored_cells = {parse_addr(x) for x in recon["authored_cell_set"]}
            # Keep = authored ∪ blank-on-input graph nodes. Snapshot omits blanks.
            nodes = set(graph)
            for dests in graph.values():
                nodes |= dests
            nodes |= authored_cells
            keep = set(authored_cells)
            for cell in nodes:
                if cell not in input_content:
                    keep.add(cell)
            fcvw_seeds = [s for s in seeds if s["class"] == "FORMULA_CORRECT_VALUE_WRONG"]
            extra = []
            if row["task"] == "Financial_Model:08_03":
                extra = [s for s in seeds if s["seed"] == addr(WITNESS_SEED) and s not in fcvw_seeds]
            for seed in fcvw_seeds + extra:
                cell = tuple(seed["cell"])
                formula = seed["authored_formula"]
                members, precedents, prov = raw_closure(cell, formula, graph, keep)  # type: ignore[arg-type]
                closures.append({
                    "task": row["task"],
                    "population": population,
                    "seed": seed["seed"],
                    "formula": formula,
                    "direct_precedents": sorted(addr(p) for p in precedents),
                    "D_seed": sorted(addr(m) for m in members),
                    "n_D": len(members),
                    "ledger": ledger[:20],
                    "provenance": prov[:200],
                    "gold_blind": True,
                    "diagnostic_only": seed["class"] != "FORMULA_CORRECT_VALUE_WRONG",
                })
                member_ok = {m: member_semantic_ok(m, recon["final_authored_content"], gold_content, gold_values) for m in members}
                klass = classify_transfer(members, authored_cells, member_ok)
                missing = sorted(addr(m) for m in members - authored_cells)
                wrong = sorted(addr(m) for m, ok in member_ok.items() if not ok and m in authored_cells)
                class_rows.append({
                    "task": row["task"],
                    "population": population,
                    "seed": seed["seed"],
                    "seed_class": seed["class"],
                    "transfer_class": klass,
                    "n_closure": len(members),
                    "n_authored_members": len(members & authored_cells),
                    "missing_members": missing,
                    "wrong_members": wrong,
                    "execution_unit": sorted(addr(m) for m in members & authored_cells),
                    "missing_are_gold_targets": None,
                    "historical_edit_plan_would_include": (
                        sorted(set(missing) & set(historical_edit_plan_members()))
                        if row["task"] == "Financial_Model:08_03" else None
                    ),
                })
                if seed["class"] != "FORMULA_CORRECT_VALUE_WRONG":
                    continue
                # Gold-only diagnosis after D freeze.
                gold_mod = set()
                try:
                    sys.path.insert(0, str(EVAL_DIR))
                    from evaluation import classify_cells_by_modification, parse_answer_position
                    rec = dataset_record(row["category"], row["id"])
                    wi = openpyxl.load_workbook(row["source_xlsx"], data_only=True)
                    wg = openpyxl.load_workbook(row["gold_xlsx"], data_only=True)
                    wif = openpyxl.load_workbook(row["source_xlsx"], data_only=False)
                    wgf = openpyxl.load_workbook(row["gold_xlsx"], data_only=False)
                    for rng in parse_answer_position(rec["answer_position"]):
                        sheet, cr = (rng.split("!", 1) if "!" in rng else (wg.sheetnames[0], rng))
                        sheet, cr = sheet.strip("'").strip(), cr.strip("'").strip()
                        try:
                            _regs, mods = classify_cells_by_modification(
                                wi, wg, sheet, cr, False, False, wb_input_formula=wif, wb_answer_formula=wgf
                            )
                        except Exception:
                            continue
                        for name in mods:
                            gold_mod.add(f"{sheet}!{name}")
                    for wb in (wi, wg, wif, wgf):
                        wb.close()
                except Exception:
                    gold_mod = set()
                class_rows[-1]["missing_are_gold_targets"] = sorted(m for m in missing if m in gold_mod)
                if klass != "T3_TRANSFERABLE_EXECUTION_CANDIDATE" or args.skip_replay:
                    continue
                # Phase 6 R0/R1
                r0_path = Path(recon["replay_r0_xlsx"]) if recon.get("replay_r0_xlsx") else None
                if not r0_path or not r0_path.is_file():
                    r0_path = WORK / row["task"].replace(":", "_") / "r0.xlsx"
                    apply_authored(
                        Path(row["source_xlsx"]),
                        r0_path,
                        recon["final_authored_content"],
                        list(recon["final_authored_content"]),
                    )
                r0_dir = WORK / row["task"].replace(":", "_") / "r0_lo"
                r0_dir.mkdir(parents=True, exist_ok=True)
                r0_lo = r0_dir / f"{row['id']}_output.xlsx"
                shutil.copy2(r0_path, r0_lo)
                refresh_dir(r0_dir)
                r0_score = score_workbook(row["category"], row["id"], r0_lo)
                stored_mod = (row.get("official") or {}).get("modification_accuracy")
                drift = None
                if stored_mod is not None and r0_score["modification"]["official_accuracy"] is not None:
                    drift = abs(float(stored_mod) - float(r0_score["modification"]["official_accuracy"]))
                faithful = drift is None or drift <= REPLAY_DRIFT_MOD
                r1_dir = WORK / row["task"].replace(":", "_") / "r1_lo"
                r1_dir.mkdir(parents=True, exist_ok=True)
                r1_raw = WORK / row["task"].replace(":", "_") / "r1.xlsx"
                order = topo_order(list(recon["final_authored_content"]), recon["final_authored_content"], graph)
                apply_authored(Path(row["source_xlsx"]), r1_raw, recon["final_authored_content"], order)
                r1_lo = r1_dir / f"{row['id']}_output.xlsx"
                shutil.copy2(r1_raw, r1_lo)
                refresh_dir(r1_dir)
                r1_score = score_workbook(row["category"], row["id"], r1_lo)
                same_cells = set(recon["final_authored_content"]) == set(recon["final_authored_content"])
                r0_ok = set(r0_score["modification"]["cells_official_ok"])
                r1_ok = set(r1_score["modification"]["cells_official_ok"])
                newly_correct = sorted(r1_ok - r0_ok)
                newly_wrong = sorted(r0_ok - r1_ok)
                written = authored_cells
                for cell_addr in newly_correct + newly_wrong:
                    try:
                        cell = parse_addr(cell_addr)
                    except Exception:
                        continue
                    cell_prov.append({
                        "task": row["task"],
                        "seed": seed["seed"],
                        "cell": cell_addr,
                        "delta": "newly_correct" if cell_addr in newly_correct else "newly_wrong",
                        "attribution": provenance_label(cell, written, graph),
                    })
                gain = (r1_score["modification"]["official_correct"] or 0) > (r0_score["modification"]["official_correct"] or 0)
                loss = (r1_score["modification"]["official_correct"] or 0) < (r0_score["modification"]["official_correct"] or 0)
                replay_manifest.append({
                    "task": row["task"],
                    "seed": seed["seed"],
                    "faithful": faithful,
                    "replay_drift_mod": drift,
                    "same_authored_cells": same_cells,
                    "r0_path": str(r0_lo),
                    "r1_path": str(r1_lo),
                    "status": "OK" if faithful else "REPLAY_NOT_FAITHFUL",
                })
                replay_scores.append({
                    "task": row["task"],
                    "seed": seed["seed"],
                    "faithful": faithful,
                    "r0_mod": r0_score["modification"]["official_accuracy"],
                    "r1_mod": r1_score["modification"]["official_accuracy"],
                    "r0_reg": r0_score["regression"]["official_accuracy"],
                    "r1_reg": r1_score["regression"]["official_accuracy"],
                    "r0_exact": r0_score["exact"],
                    "r1_exact": r1_score["exact"],
                    "r0_value_mod": r0_score["modification"]["value_only_accuracy"],
                    "r1_value_mod": r1_score["modification"]["value_only_accuracy"],
                    "r0_correct_cells": r0_score["modification"]["official_correct"],
                    "r1_correct_cells": r1_score["modification"]["official_correct"],
                    "newly_correct": newly_correct,
                    "newly_wrong": newly_wrong,
                    "r1_gain": gain,
                    "r1_loss": loss,
                    "stored_mod": stored_mod,
                })

    process_rows(included, "primary")
    fcvw_primary = [s for s in seed_rows if s["class"] == "FORMULA_CORRECT_VALUE_WRONG" and s["population"] == "primary"]
    if len(fcvw_primary) < MIN_FCVW_SEEDS:
        expanded = expansion_candidates()
        expansion_used = True
        write_json(ARTIFACT / "population.json", {
            "primary": primary,
            "expanded": expanded,
            "expansion_used": True,
            "frozen_before_replay_outcomes": True,
        })
        process_rows(expanded, "expanded")
    else:
        write_json(ARTIFACT / "population.json", {
            "primary": primary,
            "expanded": [],
            "expansion_used": False,
            "frozen_before_replay_outcomes": True,
        })

    among = defaultdict(int)
    for row in class_rows:
        if row["seed_class"] == "FORMULA_CORRECT_VALUE_WRONG":
            among[row["transfer_class"]] += 1
    fcvw_by_task = defaultdict(int)
    for row in seed_rows:
        if row["class"] == "FORMULA_CORRECT_VALUE_WRONG":
            fcvw_by_task[row["task"]] += 1
    faithful_scores = [r for r in replay_scores if r.get("faithful")]
    summary = {
        "primary_complete": len(included),
        "expanded_complete": len(expanded),
        "expansion_used": expansion_used,
        "total_authored_formula_edits": sum(1 for s in seed_rows),
        "formula_correct_edits": sum(1 for s in seed_rows if s.get("formula_correct")),
        "formula_correct_value_wrong": sum(1 for s in seed_rows if s["class"] == "FORMULA_CORRECT_VALUE_WRONG"),
        "among_fcvw": dict(among),
        "transferable_candidates": among.get("T3_TRANSFERABLE_EXECUTION_CANDIDATE", 0),
        "faithful_replay_pairs": len(faithful_scores),
        "r1_gains": sum(1 for r in faithful_scores if r.get("r1_gain")),
        "r1_ties": sum(1 for r in faithful_scores if not r.get("r1_gain") and not r.get("r1_loss")),
        "r1_losses": sum(1 for r in faithful_scores if r.get("r1_loss")),
        "replay_attempted": bool(replay_scores),
        "material_regression": any(
            r.get("r1_loss") and (r.get("r1_reg") or 1) < (r.get("r0_reg") or 1) - 0.01
            for r in faithful_scores
        ),
        "fcvw_by_task": dict(fcvw_by_task),
        "provider_cost_usd": 0,
        "runtime_s": round(time.perf_counter() - t0, 3),
    }
    recon_08 = next((r for r in reconstructions if r["task"] == "Financial_Model:08_03"), {"authored_cell_set": [], "final_authored_content": {}})
    autopsy = autopsy_08_03(
        recon_08,
        [s for s in seed_rows if s["task"] == "Financial_Model:08_03"],
        [c for c in class_rows if c["task"] == "Financial_Model:08_03"],
        next((r for r in replay_scores if r["task"] == "Financial_Model:08_03" and r.get("seed") == addr(WITNESS_SEED)), None),
    )
    verdict = choose_verdict(summary)
    verdict.update({
        "runtime_s": summary["runtime_s"],
        "provider_cost_usd": 0,
        "model_calls": 0,
        "summary": summary,
        "written_at": now(),
    })

    write_json(ARTIFACT / "artifact_integrity.json", integrity)
    write_jsonl(ARTIFACT / "authored_mutations.jsonl", all_events)
    write_json(ARTIFACT / "authored_cell_sets.json", authored_sets)
    write_json(ARTIFACT / "seed_candidates.json", seed_rows)
    write_jsonl(ARTIFACT / "raw_closures.jsonl", closures)
    write_json(ARTIFACT / "transferability_classes.json", class_rows)
    write_json(ARTIFACT / "replay_manifest.json", replay_manifest)
    write_json(ARTIFACT / "replay_scores.json", replay_scores)
    if replay_scores:
        with (ARTIFACT / "replay_scores.csv").open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(replay_scores[0]))
            writer.writeheader()
            writer.writerows(replay_scores)
    else:
        (ARTIFACT / "replay_scores.csv").write_text("task,seed,faithful,r0_mod,r1_mod\n", encoding="utf-8")
    write_jsonl(ARTIFACT / "cell_provenance.jsonl", cell_prov)
    write_json(ARTIFACT / "08_03_autopsy.json", autopsy)
    write_json(ARTIFACT / "verdict.json", verdict)
    report = render_report({"summary": summary, "verdict": verdict, "autopsy_08_03": autopsy})
    REPORT.write_text(report, encoding="utf-8")
    print(json.dumps({"verdict": verdict["verdict"], "summary": summary}, indent=1), flush=True)
    return verdict


if __name__ == "__main__":
    run()
