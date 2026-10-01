"""Pre-timing semantic and invalidation gate for JSONZ_MEMORY_V1."""
from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import struct
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "read_engine_phase1"))
from prototype import canonical, decode_xlsx  # noqa: E402
from read_engine_phase3 import safe_artifact  # noqa: E402
from openpyxl.utils.cell import coordinate_to_tuple  # noqa: E402

PHASE1_POP = ROOT / "read_engine_phase1/population.json"
TRACE_FILE = ROOT / "candidate_a_a1_checkpoint_rerun_01/contact_traces.jsonl"
PHASE1_POP_SHA = "6492d1cbc2163ea252c0d34cf346171d332a923bd5f05fc1fd54c3e851097afa"
TRACE_SHA = "665244f6150dfbb74ae647dfc8a7b8b034c02794f6de9cf7419f6d56da41696f"


def trace_values(trace: dict, book) -> list:
    ws = cell = None
    values = []
    for event in trace["operations"]:
        op = event.get("operation")
        if op == "load_workbook":
            continue
        if op == "Workbook.sheetnames":
            value = book.sheetnames
        elif op == "Workbook.__getitem__":
            ws = book[event["sheet"]]
            continue
        elif op == "Worksheet.max_row":
            value = ws.max_row
        elif op == "Worksheet.max_column":
            value = ws.max_column
        elif op == "Worksheet.cell":
            r, c = coordinate_to_tuple(event["address"])
            cell = ws.cell(row=r, column=c)
            continue
        elif op == "Cell.value":
            value = cell.value
        elif op == "Cell.data_type":
            value = cell.data_type
        else:
            continue
        values.append({"operation": op, "value": canonical(value)})
    return values


def digest(values) -> str:
    return hashlib.sha256(json.dumps(values, sort_keys=True, separators=(",", ":")).encode()).hexdigest()


def child(args) -> None:
    trace = json.loads(Path(args.trace_json).read_text())
    with safe_artifact.no_openpyxl_load():
        book = safe_artifact.load(Path(args.artifact), args.source_sha)
        values = trace_values(trace, book)
    print(json.dumps({"trace_id": trace["trace_id"], "digest": digest(values),
                      "value_count": len(values)}))


def run(append) -> bool:
    if safe_artifact.sha_file(PHASE1_POP) != PHASE1_POP_SHA or safe_artifact.sha_file(TRACE_FILE) != TRACE_SHA:
        raise RuntimeError("Phase-1 population or archived traces changed")
    population = json.loads(PHASE1_POP.read_text())
    traces = {x["trace_id"]: x for x in map(json.loads, TRACE_FILE.read_text().splitlines())}
    okay = True
    with tempfile.TemporaryDirectory(prefix="phase3-artifact-gate-") as td:
        root = Path(td)
        for idx, item in enumerate(population["workbooks"]):
            source = Path(item["path"])
            assert safe_artifact.sha_file(source) == item["sha256"]
            with safe_artifact.no_openpyxl_load():
                reference = decode_xlsx(source)
                artifact, witness, _ = safe_artifact.ensure(source, root / f"book-{idx}", item["sha256"])
                reopened = safe_artifact.load(artifact, item["sha256"])
            exact = reference.sheetnames == reopened.sheetnames
            checked = 0
            first = None
            for name in reference.sheetnames:
                a, b = reference._sheets[name], reopened._sheets[name]
                if (a.min_row, a.min_col, a.max_row, a.max_col, a.merged) != (b.min_row, b.min_col, b.max_row, b.max_col, b.merged):
                    exact = False; first = first or {"sheet": name, "category": "metadata"}
                if set(a.cells) != set(b.cells):
                    exact = False; first = first or {"sheet": name, "category": "coordinates"}
                for coord in a.cells:
                    checked += 1
                    if coord not in b.cells or (canonical(a.cells[coord][0]), a.cells[coord][1]) != (canonical(b.cells[coord][0]), b.cells[coord][1]):
                        exact = False; first = first or {"sheet": name, "cell": coord, "category": "value_or_type"}
            append({"kind": "artifact_book", "workbook_sha256": item["sha256"],
                    "witness": witness, "exact": exact, "checked_cells": checked,
                    "first_mismatch": first, "artifact_bytes": artifact.stat().st_size})
            okay &= exact and witness == "BUILT"
            if (idx + 1) % 10 == 0:
                print(f"artifact books {idx+1}/39", flush=True)
        snapshot_artifacts = {}
        snapshot_books = {}
        for item in population["traces"]:
            trace = traces[item["trace_id"]]
            source_sha = item["snapshot_sha256"]
            source = Path(item["snapshot_path"])
            assert safe_artifact.sha_file(source) == source_sha
            if source_sha not in snapshot_artifacts:
                with safe_artifact.no_openpyxl_load():
                    artifact, _, _ = safe_artifact.ensure(source, root / "trace-artifacts", source_sha)
                    snapshot_books[source_sha] = decode_xlsx(source)
                snapshot_artifacts[source_sha] = artifact
            with safe_artifact.no_openpyxl_load():
                expected = digest(trace_values(trace, snapshot_books[source_sha]))
            trace_path = root / "trace.json"
            trace_path.write_text(json.dumps(trace))
            child_run = subprocess.run([sys.executable, str(Path(__file__).resolve()), "child",
                                        "--artifact", str(snapshot_artifacts[source_sha]),
                                        "--source-sha", source_sha, "--trace-json", str(trace_path)],
                                       capture_output=True, text=True)
            if child_run.returncode:
                append({"kind": "artifact_trace", "trace_id": item["trace_id"],
                        "exact": False, "error": child_run.stderr[-1000:]})
                okay = False
                continue
            got = json.loads(child_run.stdout)
            exact = got["digest"] == expected
            append({"kind": "artifact_trace", "trace_id": item["trace_id"],
                    "snapshot_sha256": source_sha, "exact": exact,
                    "value_count": got["value_count"], "fresh_process": True})
            okay &= exact
        selected = [x for x in population["workbooks"] if x["sha256"] in {
            "1768dcc784d3c9bd2319dabeaa3a7c27d4bc0ab048a63c7084450abab3f3a1a2",
            "1b64bd5ac59b13beceb019da28ee25a8e12d479e14d625ce3ead88ccc344c741"}]
        for idx, item in enumerate(selected):
            source = Path(item["path"])
            cache = root / f"invalidation-{idx}"
            stage = root / f"staged-{idx}" / "input.xlsx"
            stage.parent.mkdir()
            shutil.copyfile(source, stage)
            artifact, witness, _ = safe_artifact.ensure(stage, cache, item["sha256"])
            _, warm, _ = safe_artifact.ensure(stage, cache, item["sha256"])
            original = artifact.read_bytes()
            artifact.write_bytes(original[:max(1, len(original)//3)])
            rejected, reason = safe_artifact.validate(artifact, Path(str(artifact)+".json"), item["sha256"])
            _, rebuilt, _ = safe_artifact.ensure(stage, cache, item["sha256"])
            hlen = struct.unpack(">I", original[8:12])[0]
            header = json.loads(original[12:12+hlen])
            header["format_version"] = "JSONZ_MEMORY_V0"
            changed_header = safe_artifact.canonical(header)
            body = original[:8] + struct.pack(">I", len(changed_header)) + changed_header + original[12+hlen:-32]
            tampered_version = body + hashlib.sha256(body).digest()
            version_rejected = False
            try:
                safe_artifact.decode(tampered_version, item["sha256"])
            except safe_artifact.ArtifactError:
                version_rejected = True
            other = selected[1-idx]
            shutil.copyfile(other["path"], stage)
            h2 = safe_artifact.sha_file(stage)
            _, changed, _ = safe_artifact.ensure(stage, cache, h2)
            shutil.copyfile(source, stage)
            _, restored, _ = safe_artifact.ensure(stage, cache, item["sha256"])
            exact = (witness == "BUILT" and warm == "REUSED" and not rejected and rebuilt == "BUILT"
                     and version_rejected and changed == "BUILT" and restored == "REUSED")
            append({"kind": "artifact_invalidation", "workbook_sha256": item["sha256"],
                    "exact": exact, "sequence": [witness, warm, rebuilt, changed, restored],
                    "h2_sha256": h2, "corrupt_rejection_reason": reason,
                    "version_mismatch_rejected": version_rejected})
            okay &= exact
    return okay


def main() -> None:
    parser = argparse.ArgumentParser()
    sub = parser.add_subparsers(dest="mode", required=True)
    c = sub.add_parser("child")
    c.add_argument("--artifact", required=True)
    c.add_argument("--source-sha", required=True)
    c.add_argument("--trace-json", required=True)
    args = parser.parse_args()
    if args.mode == "child":
        child(args)


if __name__ == "__main__":
    main()
