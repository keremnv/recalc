"""Offline, schema-checked native MemoryBook cache. No pickle or product writes."""
from __future__ import annotations

import contextlib
import hashlib
import json
import os
import re
import struct
import sys
import tempfile
import time
import zlib
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "read_engine_phase1"))
from prototype import MemoryBook, SheetInfo, _value_from_json, _value_to_json, decode_xlsx  # noqa: E402

MAGIC = b"LCRE3JZ1"
FORMAT = "JSONZ_MEMORY_V1"
CONTRACT = "PHASE1_NARROW_V5"
DECODER = "a6e95f1503ec6090d6176065471372924bf447f907a2a428faceee12a2f72473"
MAX_HEADER = 16 * 1024
MAX_COMPRESSED = 256 * 1024 * 1024
MAX_UNCOMPRESSED = 512 * 1024 * 1024
MAX_SHEETS = 512
MAX_CELLS = 10_000_000
MAX_STRING = 2 * 1024 * 1024
COORD = re.compile(r"^[A-Z]{1,4}[1-9][0-9]{0,7}$")
HEXDIGEST = re.compile(r"^[0-9a-f]{64}$")


class ArtifactError(ValueError):
    pass


def now() -> int:
    return time.perf_counter_ns()


def sha_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def canonical(obj: Any) -> bytes:
    return json.dumps(obj, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False).encode("utf-8")


def identity(source_sha: str) -> dict[str, str]:
    if not HEXDIGEST.fullmatch(source_sha):
        raise ArtifactError("invalid source hash")
    return {"source_sha256": source_sha, "decoder_sha256": DECODER,
            "contract_version": CONTRACT, "format_version": FORMAT}


def artifact_key(source_sha: str) -> str:
    return sha_bytes("|".join(identity(source_sha).values()).encode())


def paths(cache_root: Path, source_sha: str) -> tuple[Path, Path]:
    artifact = cache_root / "read-engine" / f"{artifact_key(source_sha)}.r3jz"
    return artifact, Path(str(artifact) + ".json")


def _typed(value: Any) -> dict:
    obj = json.loads(_value_to_json(value))
    _check_typed(obj)
    return obj


def _check_typed(obj: Any) -> None:
    if not isinstance(obj, dict) or obj.get("kind") not in {
        "scalar", "array", "datatable", "datetime", "date", "time", "timedelta"
    }:
        raise ArtifactError("invalid typed value kind")
    kind = obj["kind"]
    if kind == "scalar":
        if set(obj) != {"kind", "value"} or type(obj["value"]) not in (type(None), bool, int, float, str):
            raise ArtifactError("invalid scalar")
        if isinstance(obj["value"], str) and len(obj["value"]) > MAX_STRING:
            raise ArtifactError("scalar too long")
    elif kind == "array":
        if set(obj) != {"kind", "ref", "text"} or not isinstance(obj["ref"], (str, type(None))) or not isinstance(obj["text"], str):
            raise ArtifactError("invalid array formula")
        if len(obj["text"]) > MAX_STRING:
            raise ArtifactError("formula too long")
    elif kind == "datatable":
        if set(obj) != {"kind", "attrs"} or not isinstance(obj["attrs"], dict) or len(obj["attrs"]) > 32:
            raise ArtifactError("invalid data-table formula")
        if any(not isinstance(k, str) or type(v) not in (str, int, bool, type(None)) or len(str(v)) > MAX_STRING for k, v in obj["attrs"].items()):
            raise ArtifactError("invalid data-table attributes")
    elif kind == "timedelta":
        if set(obj) != {"kind", "seconds"} or type(obj["seconds"]) not in (int, float):
            raise ArtifactError("invalid timedelta")
    else:
        if set(obj) != {"kind", "value"} or not isinstance(obj["value"], str) or len(obj["value"]) > 128:
            raise ArtifactError("invalid temporal value")


def _payload(book: MemoryBook) -> bytes:
    if not isinstance(book.sheetnames, list) or len(book.sheetnames) > MAX_SHEETS:
        raise ArtifactError("too many sheets")
    records = []
    count = 0
    for name in book.sheetnames:
        sheet = book._sheets[name]
        if not isinstance(name, str) or len(name) > 255:
            raise ArtifactError("invalid sheet name")
        merged = [list(x) for x in sheet.merged]
        cells = []
        for coord in sorted(sheet.cells):
            value, dtype = sheet.cells[coord]
            cells.append([coord, dtype, _typed(value)])
            count += 1
            if count > MAX_CELLS:
                raise ArtifactError("too many cells")
        records.append({"name": name, "bounds": [sheet.min_row, sheet.min_col, sheet.max_row, sheet.max_col],
                        "merged": merged, "cells": cells})
    raw = canonical({"sheets": records})
    if len(raw) > MAX_UNCOMPRESSED:
        raise ArtifactError("payload too large")
    return raw


def _book(raw: bytes) -> MemoryBook:
    try:
        obj = json.loads(raw)
    except (ValueError, UnicodeDecodeError) as exc:
        raise ArtifactError("invalid payload JSON") from exc
    if not isinstance(obj, dict) or set(obj) != {"sheets"} or not isinstance(obj["sheets"], list) or len(obj["sheets"]) > MAX_SHEETS:
        raise ArtifactError("invalid sheet list")
    sheets = []
    seen = set()
    count = 0
    for entry in obj["sheets"]:
        if not isinstance(entry, dict) or set(entry) != {"name", "bounds", "merged", "cells"}:
            raise ArtifactError("invalid sheet record")
        name, bounds, merged, cells = entry["name"], entry["bounds"], entry["merged"], entry["cells"]
        if not isinstance(name, str) or len(name) > 255 or name in seen:
            raise ArtifactError("invalid or duplicate sheet name")
        seen.add(name)
        if not isinstance(bounds, list) or len(bounds) != 4 or any(type(x) is not int or x < 1 or x > 2**31 - 1 for x in bounds):
            raise ArtifactError("invalid bounds")
        if bounds[0] > bounds[2] or bounds[1] > bounds[3]:
            raise ArtifactError("inverted bounds")
        if not isinstance(merged, list) or not isinstance(cells, list):
            raise ArtifactError("invalid merged/cells list")
        if len(merged) > MAX_CELLS:
            raise ArtifactError("too many merged ranges")
        ranges = []
        for item in merged:
            if not isinstance(item, list) or len(item) != 4 or any(type(x) is not int or x < 1 or x > 2**31 - 1 for x in item):
                raise ArtifactError("invalid merged range")
            if item[0] > item[2] or item[1] > item[3]:
                raise ArtifactError("inverted merged range")
            ranges.append(tuple(item))
        store = {}
        for item in cells:
            if not isinstance(item, list) or len(item) != 3:
                raise ArtifactError("invalid cell record")
            coord, dtype, typed = item
            if not isinstance(coord, str) or len(coord) > 12 or not COORD.fullmatch(coord) or coord in store:
                raise ArtifactError("invalid or duplicate coordinate")
            if not isinstance(dtype, str) or dtype not in {"n", "s", "b", "e", "f", "d", "inlineStr", "str"}:
                raise ArtifactError("invalid data type")
            _check_typed(typed)
            try:
                value = _value_from_json(canonical(typed).decode())
            except (TypeError, ValueError, KeyError, OverflowError) as exc:
                raise ArtifactError("invalid typed value") from exc
            store[coord] = (value, dtype)
            count += 1
            if count > MAX_CELLS:
                raise ArtifactError("too many cells")
        sheets.append(SheetInfo(name, *bounds, ranges, store))
    return MemoryBook(sheets)


def encode(book: MemoryBook, source_sha: str) -> bytes:
    raw = _payload(book)
    compressed = zlib.compress(raw, level=1)
    if len(compressed) > MAX_COMPRESSED:
        raise ArtifactError("compressed payload too large")
    header = {**identity(source_sha), "sheet_count": len(book.sheetnames),
              "payload_sha256": sha_bytes(raw)}
    h = canonical(header)
    if len(h) > MAX_HEADER:
        raise ArtifactError("header too large")
    body = MAGIC + struct.pack(">I", len(h)) + h + struct.pack(">Q", len(compressed)) + compressed
    return body + hashlib.sha256(body).digest()


def decode(data: bytes, expected_source_sha: str) -> MemoryBook:
    if len(data) < 52 or data[:8] != MAGIC:
        raise ArtifactError("bad magic or truncated artifact")
    if hashlib.sha256(data[:-32]).digest() != data[-32:]:
        raise ArtifactError("artifact checksum mismatch")
    hlen = struct.unpack(">I", data[8:12])[0]
    if hlen > MAX_HEADER or 12 + hlen + 8 + 32 > len(data):
        raise ArtifactError("invalid header length")
    try:
        header = json.loads(data[12:12 + hlen])
    except (ValueError, UnicodeDecodeError) as exc:
        raise ArtifactError("invalid header JSON") from exc
    if not isinstance(header, dict) or set(header) != set(identity(expected_source_sha)) | {"sheet_count", "payload_sha256"}:
        raise ArtifactError("invalid header shape")
    if any(header[k] != v for k, v in identity(expected_source_sha).items()):
        raise ArtifactError("artifact identity mismatch")
    if type(header["sheet_count"]) is not int or not 0 <= header["sheet_count"] <= MAX_SHEETS or not isinstance(header["payload_sha256"], str) or not HEXDIGEST.fullmatch(header["payload_sha256"]):
        raise ArtifactError("invalid header fields")
    offset = 12 + hlen
    clen = struct.unpack(">Q", data[offset:offset + 8])[0]
    if clen > MAX_COMPRESSED or offset + 8 + clen + 32 != len(data):
        raise ArtifactError("invalid payload length")
    try:
        dec = zlib.decompressobj()
        raw = dec.decompress(data[offset + 8:offset + 8 + clen], MAX_UNCOMPRESSED + 1)
        raw += dec.flush()
    except zlib.error as exc:
        raise ArtifactError("invalid compressed payload") from exc
    if not dec.eof or dec.unused_data or len(raw) > MAX_UNCOMPRESSED or sha_bytes(raw) != header["payload_sha256"]:
        raise ArtifactError("payload integrity mismatch")
    book = _book(raw)
    if len(book.sheetnames) != header["sheet_count"]:
        raise ArtifactError("sheet count mismatch")
    return book


def _sync_dir(directory: Path) -> None:
    fd = os.open(directory, os.O_RDONLY)
    try:
        os.fsync(fd)
    finally:
        os.close(fd)


def _atomic(path: Path, data: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(dir=path.parent, prefix="phase3-", suffix=".tmp")
    temp = Path(name)
    try:
        with os.fdopen(fd, "wb") as stream:
            stream.write(data)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temp, path)
        _sync_dir(path.parent)
    finally:
        temp.unlink(missing_ok=True)


def validate(artifact: Path, sidecar: Path, source_sha: str) -> tuple[bool, str]:
    try:
        meta = json.loads(sidecar.read_bytes())
        expected = identity(source_sha)
        if not isinstance(meta, dict) or any(meta.get(k) != v for k, v in expected.items()):
            return False, "manifest_identity"
        if meta.get("artifact_key") != artifact_key(source_sha) or not isinstance(meta.get("artifact_sha256"), str) or not HEXDIGEST.fullmatch(meta["artifact_sha256"]):
            return False, "manifest_fields"
        if not artifact.is_file() or artifact.stat().st_size != meta.get("artifact_bytes") or artifact.stat().st_size > MAX_COMPRESSED + MAX_HEADER + 52:
            return False, "artifact_missing_or_size"
        blob = artifact.read_bytes()
        if sha_bytes(blob) != meta["artifact_sha256"]:
            return False, "artifact_hash"
        decode(blob, source_sha)
        return True, "valid"
    except (OSError, ValueError, TypeError, KeyError, OverflowError) as exc:
        return False, f"invalid:{type(exc).__name__}"


@contextlib.contextmanager
def no_openpyxl_load():
    import openpyxl
    import openpyxl.reader.excel
    before = openpyxl.load_workbook, openpyxl.reader.excel.load_workbook
    def forbidden(*args, **kwargs):
        raise AssertionError("direct treatment invoked openpyxl.load_workbook")
    openpyxl.load_workbook = forbidden
    openpyxl.reader.excel.load_workbook = forbidden
    try:
        yield
    finally:
        openpyxl.load_workbook, openpyxl.reader.excel.load_workbook = before


def ensure(source: Path, cache_root: Path, source_sha: str) -> tuple[Path, str, dict]:
    """Validate existing state or build/publish from direct OOXML."""
    artifact, sidecar = paths(cache_root, source_sha)
    profile: dict[str, Any] = {}
    t = now()
    good, reason = validate(artifact, sidecar, source_sha)
    profile["artifact_discovery_validation_ns"] = now() - t
    if good:
        profile["artifact_bytes"] = artifact.stat().st_size
        return artifact, "REUSED", profile
    profile["rebuild_reason"] = reason
    t = now()
    with no_openpyxl_load():
        book = decode_xlsx(source)
    profile["direct_decode_ns"] = now() - t
    profile["decoder_profile"] = book.profile
    t = now()
    blob = encode(book, source_sha)
    profile["serialization_ns"] = now() - t
    t = now()
    _atomic(artifact, blob)
    meta = {**identity(source_sha), "artifact_key": artifact_key(source_sha),
            "artifact_bytes": len(blob), "artifact_sha256": sha_bytes(blob)}
    _atomic(sidecar, canonical(meta))
    profile["publication_ns"] = now() - t
    profile["artifact_bytes"] = len(blob)
    return artifact, "BUILT", profile


def load(artifact: Path, source_sha: str) -> MemoryBook:
    with no_openpyxl_load():
        return decode(artifact.read_bytes(), source_sha)
