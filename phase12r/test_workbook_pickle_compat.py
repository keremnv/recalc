import pickle,unittest
import openpyxl
import workbook_pickle_compat
class GraphParity(unittest.TestCase):
 def test_nested_graph(self):
  wb=openpyxl.Workbook();ws=wb.active;ws['A1']='=1+2';ws['B2']=17
  ws.column_dimensions['A'].width=23;ws.row_dimensions[2].height=19
  for _ in range(4):
   wb=pickle.loads(pickle.dumps(wb,protocol=5));ws=wb.active
   self.assertEqual(ws['A1'].value,'=1+2');self.assertEqual(ws['B2'].value,17)
   self.assertEqual(ws.column_dimensions['A'].width,23);self.assertEqual(ws.row_dimensions[2].height,19)
   self.assertIs(ws.column_dimensions.worksheet,ws);self.assertIs(ws.row_dimensions.worksheet,ws)
   self.assertEqual(ws.column_dimensions['Z'].index,'Z')
 def test_absent_factory(self):
  wb=openpyxl.Workbook();wb.active.column_dimensions.default_factory=None
  for _ in range(3):
   wb=pickle.loads(pickle.dumps(wb,protocol=5));self.assertIsNone(wb.active.column_dimensions.default_factory)
if __name__=='__main__':unittest.main()
