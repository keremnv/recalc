"""Unit tests for gold-blind formula-completion certificates."""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "benchmark"))

from formula_completion_certs import (  # noqa: E402
    DEFINITIONS,
    certificates_for_workbook,
    load_grids,
    lr_certificate,
    ud_certificate,
)


def _save(tmp_path: Path, fill) -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Model"
    fill(sheet)
    path = tmp_path / "input.xlsx"
    workbook.save(path)
    workbook.close()
    return path


def test_definitions_forbid_golden_and_semantics() -> None:
    assert DEFINITIONS["golden_in_generation"] is False
    assert DEFINITIONS["same_sheet_only"] is True
    assert "style" not in DEFINITIONS["LR"].lower()
    assert "golden" not in DEFINITIONS["K_peer"].lower()


def test_lr_agreement_example(tmp_path: Path) -> None:
    def fill(sheet) -> None:
        sheet["G5"] = 10
        sheet["G7"] = 3
        sheet["G10"] = "=G5-G7"
        sheet["I5"] = 10
        sheet["I7"] = 3
        sheet["I10"] = "=I5-I7"

    path = _save(tmp_path, fill)
    grid = load_grids(path)[0]
    cert = lr_certificate(grid, 8, 10)
    assert cert is not None
    assert cert.candidate == "=H5-H7"
    assert cert.agreeing_sources == 2
    assert cert.adjacent is True
    assert cert.same_equivalence_class is True
    assert {source.address for source in cert.sources} == {"G10", "I10"}


def test_lr_does_not_skip_ineligible_formula(tmp_path: Path) -> None:
    def fill(sheet) -> None:
        sheet["E10"] = "=E5-E7"
        sheet["G10"] = '=INDIRECT("G5")'
        sheet["I10"] = "=I5-I7"

    path = _save(tmp_path, fill)
    grid = load_grids(path)[0]
    assert lr_certificate(grid, 8, 10) is None


def test_ud_agreement(tmp_path: Path) -> None:
    def fill(sheet) -> None:
        sheet["H8"] = "=E8-G8"
        sheet["H12"] = "=E12-G12"

    path = _save(tmp_path, fill)
    grid = load_grids(path)[0]
    cert = ud_certificate(grid, 8, 10)
    assert cert is not None
    assert cert.candidate == "=E10-G10"
    assert cert.axes == ["vertical"]


def test_cross_and_k_nested(tmp_path: Path) -> None:
    def fill(sheet) -> None:
        sheet["G10"] = "=G5-G7"
        sheet["I10"] = "=I5-I7"
        sheet["H8"] = "=H3-H5"
        sheet["H12"] = "=H7-H9"

    path = _save(tmp_path, fill)
    payload = certificates_for_workbook(path)
    lr = [item for item in payload["certificates"]["LR"] if item["address"] == "H10"]
    ud = [item for item in payload["certificates"]["UD"] if item["address"] == "H10"]
    cross = [item for item in payload["certificates"]["CROSS"] if item["address"] == "H10"]
    k2 = payload["certificates"]["K2"]
    assert len(lr) == 1
    assert len(ud) == 1
    assert len(cross) == 1
    assert lr[0]["candidate"] == "=H5-H7"
    assert ud[0]["candidate"] == "=H5-H7"
    assert cross[0]["candidate"] == "=H5-H7"
    assert cross[0]["agreeing_sources"] == 4
    assert any(item["address"] == "H10" for item in k2)


def test_k_conflict_when_two_formulas_each_have_two_peers(tmp_path: Path) -> None:
    def fill(sheet) -> None:
        sheet["A10"] = "=SUM(A1:A9)"
        sheet["B10"] = "=SUM(B1:B9)"
        sheet["D10"] = "=D5-D7"
        sheet["E10"] = "=E5-E7"

    path = _save(tmp_path, fill)
    payload = certificates_for_workbook(path)
    k2 = [item for item in payload["certificates"]["K2"] if item["address"] == "C10"]
    conflicts = [
        item
        for item in payload["conflicts"]
        if item["address"] == "C10" and item["rule"] == "K2"
    ]
    assert k2 == []
    assert len(conflicts) == 1
    assert conflicts[0]["n_groups"] == 2


def test_lr_ud_disagree_is_not_cross(tmp_path: Path) -> None:
    def fill(sheet) -> None:
        sheet["G10"] = "=G5-G7"
        sheet["I10"] = "=I5-I7"
        sheet["H8"] = "=A8"
        sheet["H12"] = "=A12"

    path = _save(tmp_path, fill)
    payload = certificates_for_workbook(path)
    assert payload["certificates"]["CROSS"] == []
    assert payload["lr_ud_disagree"]
    assert payload["lr_ud_disagree"][0]["lr"] == "=H5-H7"
    assert payload["lr_ud_disagree"][0]["ud"] == "=A10"


def test_generation_does_not_require_a_golden(tmp_path: Path) -> None:
    def fill(sheet) -> None:
        sheet["A1"] = "=B1"
        sheet["C1"] = "=D1"

    path = _save(tmp_path, fill)
    payload = certificates_for_workbook(path)
    cert = payload["certificates"]["LR"][0]
    assert cert["address"] == "B1"
    assert cert["candidate"] == "=C1"


def test_apply_writes_blanks_only(tmp_path: Path) -> None:
    from formula_completion_certs import apply_certificates

    def fill(sheet) -> None:
        sheet["A1"] = "=B1"
        sheet["C1"] = "=D1"
        sheet["A2"] = 5

    path = _save(tmp_path, fill)
    dest = tmp_path / "out.xlsx"
    stats = apply_certificates(
        path,
        dest,
        [
            {"sheet": "Model", "col": 2, "row": 1, "candidate": "=C1"},
            {"sheet": "Model", "col": 1, "row": 2, "candidate": "=B2"},
        ],
    )
    assert stats["written"] == 1
    assert stats["skipped_occupied"] == 1
    workbook = openpyxl.load_workbook(dest)
    assert workbook["Model"]["B1"].value == "=C1"
    assert workbook["Model"]["A2"].value == 5
    workbook.close()


def test_apply_skips_merged_cells(tmp_path: Path) -> None:
    from formula_completion_certs import apply_certificates

    def fill(sheet) -> None:
        sheet.merge_cells("B1:C1")
        sheet["A1"] = "=Z1"

    path = _save(tmp_path, fill)
    dest = tmp_path / "out.xlsx"
    stats = apply_certificates(
        path,
        dest,
        [{"sheet": "Model", "col": 3, "row": 1, "candidate": "=Y1"}],
    )
    assert stats["written"] == 0
    assert stats["skipped_merged"] == 1


def test_apply_copies_without_opening_when_there_are_no_certificates(tmp_path: Path) -> None:
    from formula_completion_certs import apply_certificates

    path = tmp_path / "input.xlsx"
    path.write_bytes(b"not-a-workbook")
    dest = tmp_path / "out.xlsx"
    stats = apply_certificates(path, dest, [])
    assert stats["written"] == 0
    assert dest.read_bytes() == b"not-a-workbook"
