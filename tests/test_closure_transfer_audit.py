from __future__ import annotations

import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "benchmark"))

import closure_transfer_audit as audit  # noqa: E402


def test_official_formula_policy_strips_abs_and_unary_plus() -> None:
    assert audit.formulas_match("=C68/C$6", "=C68/C6")
    assert audit.formulas_match("=+J35+J26+J15+J6", "=J6+J15+J26+J35") is False
    assert audit.formulas_match("=+J35+J26+J15+J6", "=J35+J26+J15+J6")
    assert audit.official_formula_key("=c68/c$6") == "=C68/C6"


def test_mutation_kinds() -> None:
    blank = {"kind": "BLANK", "text": None}
    formula = {"kind": "FORMULA", "text": "=A1"}
    value = {"kind": "VALUE", "text": 3}
    assert audit.mutation_kind(blank, formula) == "SET_FORMULA"
    assert audit.mutation_kind(blank, value) == "SET_VALUE"
    assert audit.mutation_kind(formula, blank) == "CLEAR"


def test_transfer_classes() -> None:
    authored = {("S", 1, 2)}
    assert audit.classify_transfer(set(), authored, {}) == "T0_CLOSURE_EMPTY"
    assert audit.classify_transfer({("S", 1, 2), ("S", 1, 3)}, authored, {("S", 1, 2): True}) == "T1_AUTHORITY_GAP"
    assert audit.classify_transfer({("S", 1, 2)}, authored, {("S", 1, 2): False}) == "T2_AUTHORED_MEMBER_SEMANTIC_FAILURE"
    assert audit.classify_transfer({("S", 1, 2)}, authored, {("S", 1, 2): True}) == "T3_TRANSFERABLE_EXECUTION_CANDIDATE"


def test_blank_keep_offset_closure(tmp_path: Path) -> None:
    path = tmp_path / "in.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = 1
    ws["B1"] = "=OFFSET(A1,1,0)"
    ws["A2"] = None
    wb.save(path)
    wb.close()
    graph, _values, _ledger = audit.graph_from_workbook(path)
    seed = ("Sheet1", 3, 2)  # B3
    keep = {("Sheet1", 2, 1)}  # A2 blank
    members, precedents, _prov = audit.raw_closure(seed, "=B1", graph, keep)
    assert ("Sheet1", 1, 2) in precedents  # B1
    # A2 is the OFFSET sweep of B1 and is blank, so it is a required member.
    assert ("Sheet1", 2, 1) in members


def test_no_provider_inference_in_module() -> None:
    text = Path(audit.__file__).read_text(encoding="utf-8")
    assert "OPENROUTER_API_KEY" not in text
    assert "chat.completions" not in text.lower()
    assert "urllib.request" not in text
    assert "openai" not in text.lower()


def test_rewrite_command_does_not_double_replace_tmp(tmp_path: Path) -> None:
    sandbox = Path("/tmp/closure_transfer_audit_work/demo/sandbox")
    cmd = "mkdir -p /mnt/spreadsheet_output && cat > /tmp/process.py << 'EOF'\nprint(1)\nEOF\npython3 /tmp/process.py"
    out = audit.rewrite_command(cmd, sandbox)
    assert "mkdir -p /mnt/spreadsheet_output" not in out
    assert str(sandbox / "mnt/spreadsheet_output") in out
    assert "python3 /tmp/process.py" not in out
    assert out.count(str(sandbox / "tmp")) >= 2
    assert "sandbox/tmp/closure_transfer_audit_work" not in out


def test_verdict_no_fcvw() -> None:
    v = audit.choose_verdict({
        "formula_correct_value_wrong": 0,
        "among_fcvw": {},
        "transferable_candidates": 0,
        "faithful_replay_pairs": 0,
        "r1_gains": 0,
        "r1_losses": 0,
        "replay_attempted": False,
    })
    assert v["verdict"] == "NO_COMPOSITION_DEPENDENT_FAILURE_IN_CONTROL"

