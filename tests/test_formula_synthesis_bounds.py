"""Used-range checks must not assume read-only worksheets expose integer bounds."""
from __future__ import annotations

import json
import re
import sys
import zipfile
from pathlib import Path

import openpyxl
import pytest
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "benchmark"))
sys.path.insert(0, str(ROOT / "src"))

import formula_synthesis_probe as synth
import matched_compiled_treatment as m

ARCHIVED_INPUT = (
    ROOT
    / "benchmark-data/SpreadsheetBench-2/benchmark-runs/matched-glm-compiled-sixty"
    / "prep/repaired_inputs/Financial_Model-04_01_input.xlsx"
)
ARCHIVED_PROPOSAL = "=F10/E10-1"
ARCHIVED_TARGET = {"sheet": "Financials", "row": 11, "col": 6, "address": "F11"}
ARCHIVED_SPINE = {"title_to_index": {"Financials": 3}}
ARCHIVED_CALL = (
    ROOT
    / "research/history/scheduler_live_validation/Financial_Model-04_01/calls/056_synthesis.json"
)


def _strip_dimension(path: Path) -> None:
    tmp = path.with_suffix(".unsized.xlsx")
    with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(tmp, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet") and item.filename.endswith(".xml"):
                data = re.sub(rb"<dimension[^/]*/>", b"", data)
            zout.writestr(item, data)
    tmp.replace(path)


def _unsized_workbook(path: Path, title: str, cells: dict[str, object]) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    for address, value in cells.items():
        sheet[address] = value
    workbook.save(path)
    workbook.close()
    _strip_dimension(path)
    return path


def _readonly_maxima(path: Path, sheet: str) -> tuple[object, object]:
    workbook = openpyxl.load_workbook(path, data_only=False, read_only=True)
    try:
        ws = workbook[sheet]
        return ws.max_column, ws.max_row
    finally:
        workbook.close()


def _validate(path: Path, formula: str, target: dict, spine: dict) -> dict:
    row = {"task": "bounds-fixture", "input_path": str(path), "target": target}
    return synth._validate_formula(row, formula, {}, spine, None, None)


def test_archived_accepted_proposal_does_not_crash_on_unsized_readonly_sheet():
    if not ARCHIVED_INPUT.exists():
        pytest.skip("frozen Financial_Model:04_01 input is not present")
    if ARCHIVED_CALL.exists():
        persisted = json.loads(ARCHIVED_CALL.read_text(encoding="utf-8"))
        parsed = persisted["parsed_response"]
        assert parsed["status"] == "PROPOSED"
        assert parsed["formula"] == ARCHIVED_PROPOSAL
        assert parsed["target_id"] == "cell:s03:r11:c6"
    max_col, max_row = _readonly_maxima(ARCHIVED_INPUT, "Financials")
    assert max_col is None
    assert max_row is None
    result = _validate(ARCHIVED_INPUT, ARCHIVED_PROPOSAL, ARCHIVED_TARGET, ARCHIVED_SPINE)
    assert result["parser_ok"] is True
    assert result["invalid_sheet"] is False
    assert result["invalid_address"] is False
    refs = {(host, start, end) for host, start, end in synth.formula_a1_references(ARCHIVED_PROPOSAL)}
    assert refs == {(None, "F10", None), (None, "E10", None)}


def test_archived_proposal_survives_scheduler_validate_formula_entry():
    if not ARCHIVED_INPUT.exists():
        pytest.skip("frozen Financial_Model:04_01 input is not present")
    result = m.validate_formula(
        "Financial_Model:04_01",
        ARCHIVED_TARGET,
        ARCHIVED_PROPOSAL,
        {},
        ARCHIVED_SPINE,
    )
    assert result["hard_verifier_result"] in {"HARD_ACCEPT", "HARD_REJECT"}
    assert result["invalid_address"] is False


def test_unsized_sheet_still_rejects_addresses_beyond_used_range(tmp_path: Path):
    path = _unsized_workbook(
        tmp_path / "unsized.xlsx",
        "Financials",
        {"E10": 1, "F10": 2},
    )
    assert _readonly_maxima(path, "Financials") == (None, None)
    spine = {"title_to_index": {"Financials": 0}}
    target = {"sheet": "Financials", "row": 11, "col": 6}
    inside = _validate(path, "=F10/E10-1", target, spine)
    assert inside["invalid_address"] is False
    beyond = _validate(path, "=Z99", target, spine)
    assert beyond["invalid_address"] is True
    beyond_range = _validate(path, "=SUM(E10:ZZ10)", target, spine)
    assert beyond_range["invalid_address"] is True
