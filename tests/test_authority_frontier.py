"""Compiled facts must reach planning through the same world used for expansion."""
import datetime as dt
import json
import os
import subprocess
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'benchmark'))

import fm_resource_feasibility as feasibility
import matched_compiled_treatment as treatment
import openpyxl
import pytest
from edit_plan import World
from temporal_spine import compile_temporal_workbook
from workbook_grounding import parse_scope_spec, project_obligation
from workbook_grounding_spine import compile_spine
from workbook_spine_sqlite import build_database


@pytest.mark.parametrize('rate', ['10%', '30%', '12.5%', '1,000%', '100%', '12 %', '12,500.5%'])
def test_percentage_is_not_a_year_even_when_scope_also_has_dates(rate):
    scope = {'scope': [{'text': f'{rate} FY26, 30% FY27–FY30'}]}
    assert parse_scope_spec(scope)['years'] == list(range(2026, 2031))
    assert parse_scope_spec({'scope': [{'text': rate}]})['years'] == []
    assert parse_scope_spec(scope)['raw'] == [scope['scope'][0]['text']]


@pytest.fixture
def compiled_world(tmp_path, monkeypatch):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'Schedule'
    sheet['A2'] = 'Output'
    sheet['B1'] = '=Calendar!B1'
    sheet['C1'] = '=EDATE(B1,12)'
    sheet['C3'] = 1  # implicit blank targets B2:C2 inside bounds
    calendar = workbook.create_sheet('Calendar')
    calendar['B1'] = dt.date(2026, 1, 1)
    calendar['B1'].number_format = 'yyyy'
    source = tmp_path / 'input.xlsx'
    workbook.save(source)
    workbook.close()
    spine = compile_spine(source, workbook_key='test')
    temporal = compile_temporal_workbook(source, closure=True)
    spine_path, temporal_path, database = (tmp_path / p for p in ('spine.json', 'temporal.json', 'world.sqlite'))
    spine_path.write_text(json.dumps(spine))
    temporal_path.write_text(json.dumps(temporal))
    build_database(spine_path, temporal_path, database)
    monkeypatch.setattr(treatment, 'spine_for', lambda k: spine)
    monkeypatch.setattr(treatment, 'db_for', lambda k: database)
    return spine, database


def obligation():
    return {'id': 'O1', 'locus': {'text': 'Schedule'}, 'subject': {'text': 'Output'},
            'scope': [{'text': 'FY26–FY27'}], 'then_after': []}


def test_formula_linked_temporal_facts_reach_grounding_and_expand_in_same_world(compiled_world):
    legacy, database = compiled_world
    before = json.dumps(legacy, sort_keys=True)
    assert project_obligation(legacy, obligation())['scope'] == []
    world = World(database)
    try:
        projected = treatment.planning_spine_for('test', world)
        packet = project_obligation(projected, obligation())
        assert {p['period']['year'] for p in packet['scope']} == {2026, 2027}
        assert set(packet['target_cell_ids']) == {'cell:s00:r2:c2', 'cell:s00:r2:c3'}
        coords = sorted(packet['scope'], key=lambda p: p['period']['year'])
        plan = {'operations': [{'operation_id': 'op1', 'obligation_id': 'O1', 'operation_kind': 'SET_FORMULA',
                 'target_set': {'kind': 'TEMPORAL_INTERVAL', 'sheet_id': 'sheet:s00', 'axis': 'column',
                                'start_coordinate': coords[0]['id'], 'end_coordinate': coords[-1]['id'],
                                'row_constraint': {'r1': 2, 'r2': 2}}}]}
        expanded = treatment.expand_edit_plan(plan, world, {'O1'}, id_contract='V2')
        assert set(expanded['cell_ids']) == set(packet['target_cell_ids'])
        assert {p['id'] for p in projected['periods']} == set(world.temporal)
    finally:
        world.close()
    assert json.dumps({k: v for k, v in legacy.items() if not k.startswith('_')}, sort_keys=True) == before


@pytest.mark.parametrize('mode', ['monolithic', 'sharded_old', 'sharded_projected'])
def test_each_planner_mode_receives_compiled_temporal_evidence(compiled_world, monkeypatch, tmp_path, mode):
    requests = []
    def retained_call(*args, **kwargs):
        requests.append(json.loads(args[4]))
        return {'parsed_response': None}
    monkeypatch.setattr(treatment, 'call_or_stub', retained_call)
    monkeypatch.setattr(treatment, 'ACTIVE_FRONTEND_MODE', mode)
    compiler = {'raw_task': 'Calculate Output in Schedule for FY26–FY27.', 'obligations': [obligation()]}
    plan = treatment.plan_task('test', {}, compiler, {}, tmp_path, stub=True)
    assert requests
    assert {p['period']['year'] for p in plan['packets']['O1']['scope']} == {2026, 2027}
    assert 'tcoord:s00:c:2' in json.dumps(requests)


@pytest.mark.parametrize('raises', [False, True])
def test_repaired_runner_routes_to_repaired_db_and_restores_globals(tmp_path, monkeypatch, raises):
    (tmp_path / 'Financial_Model-04_01.sqlite').touch()
    monkeypatch.setattr(feasibility, 'REPAIRED_DATABASES', tmp_path)
    old_db, old_request = treatment.DATABASES, treatment.request_body
    def run_one(*args, **kwargs):
        assert treatment.DATABASES == tmp_path
        # Request construction is authoritative and cannot be monkeypatched by
        # a feasibility profile.
        assert treatment.request_body is old_request
        if raises:
            raise RuntimeError('test failure')
        return {'status': 'test-only'}
    monkeypatch.setattr(treatment, 'run_one_task', run_one)
    if raises:
        with pytest.raises(RuntimeError, match='test failure'):
            feasibility.run_treatment_task('04_01')
    else:
        assert feasibility.run_treatment_task('04_01')['status'] == 'test-only'
    assert treatment.DATABASES == old_db
    assert treatment.request_body is old_request


def test_repaired_runner_does_not_silently_use_placeholder_db(tmp_path, monkeypatch):
    monkeypatch.setattr(feasibility, 'REPAIRED_DATABASES', tmp_path)
    monkeypatch.setattr(treatment, 'run_one_task', lambda *a, **k: pytest.fail('No provider work without repaired world'))
    with pytest.raises(FileNotFoundError, match='Repaired compiled database required'):
        feasibility.run_treatment_task('04_01')


def test_bounded_dependency_packet_is_independent_of_python_hash_seed(tmp_path):
    workbook = openpyxl.Workbook()
    workbook.active.title = 'Schedule'
    workbook.active['B1'] = 'FY26'
    for row in range(2, 262):
        workbook.active.cell(row, 1, 'Output')
        workbook.active.cell(row, 3, f'=B{row}')
    source = tmp_path / 'input.xlsx'
    workbook.save(source)
    workbook.close()
    path = tmp_path / 'spine.json'
    path.write_text(json.dumps(compile_spine(source, workbook_key='test')))
    code = '''
import json,sys
sys.path.insert(0, sys.argv[1])
import matched_compiled_treatment
from workbook_grounding import project_obligation
spine = json.load(open(sys.argv[2]))
ob = {'id':'O1','locus':{'text':'Schedule'},'subject':{'text':'Output'},'scope':[{'text':'FY26'}]}
packet = project_obligation(spine, ob)
print(json.dumps(packet['dependency_facts'], sort_keys=True))
'''
    results = [subprocess.check_output([sys.executable, '-c', code, str(Path(__file__).resolve().parents[1] / 'benchmark'), str(path)],
               env={**os.environ, 'PYTHONHASHSEED': seed}, text=True, timeout=30) for seed in ('1', '2')]
    assert len(json.loads(results[0])) == 80
    assert results[0] == results[1]
