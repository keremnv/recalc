"""Process and external-effect contracts of the packaged observer command."""
from __future__ import annotations

import json
import os
import signal
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

from librecalc_agent.runner import observer_binary

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def work(tmp_path):
    directory = tmp_path / "work"
    directory.mkdir()
    book = openpyxl.Workbook()
    book.active["A1"] = 1
    book.save(directory / "input.xlsx")
    return directory


def observed(work: Path, script: str, args: list[str] | None = None,
             env_extra: dict[str, str] | None = None,
             pass_fds: tuple[int, ...] = ()):
    target = work / "fixture.py"
    target.write_text(script)
    cache = work.parent / "cache"
    runs = cache / "runs"
    cache.mkdir(exist_ok=True)
    runs.mkdir(exist_ok=True)
    command = [str(observer_binary()), sys.executable, str(target), str(work),
               str(cache), str(runs), str(ROOT / "src/librecalc_agent/_bootstrap"),
               str(ROOT / "src/librecalc_agent/_capture_helper.py"), *(args or [])]
    env = {**os.environ, "PYTHONPATH": str(ROOT / "src"), **(env_extra or {})}
    result = subprocess.run(command, cwd=work, env=env, capture_output=True,
                            pass_fds=pass_fds, timeout=20)
    pointer = json.loads((runs / "last_run.json").read_text())
    run_dir = Path(pointer["run_dir"])
    return result, json.loads((run_dir / "observer_receipt.json").read_text()), run_dir


def observed_raw(work: Path, script: str):
    """Run the observer without reading a receipt (for pre-launch failures)."""
    target = work / "fixture.py"
    target.write_text(script)
    cache = work.parent / "cache"
    runs = cache / "runs"
    cache.mkdir(exist_ok=True)
    runs.mkdir(exist_ok=True)
    command = [str(observer_binary()), sys.executable, str(target), str(work),
               str(cache), str(runs), str(ROOT / "src/librecalc_agent/_bootstrap"),
               str(ROOT / "src/librecalc_agent/_capture_helper.py")]
    env = {**os.environ, "PYTHONPATH": str(ROOT / "src")}
    return subprocess.run(command, cwd=work, env=env, capture_output=True,
                          timeout=60)


def test_observer_child_launch_failure_is_target_status(work):
    target = work / "fixture.py"
    target.write_text("print('must not run')\n")
    cache = work.parent / "cache"
    runs = cache / "runs"
    cache.mkdir(exist_ok=True)
    runs.mkdir(exist_ok=True)
    command = [str(observer_binary()), "/nonexistent/python", str(target),
               str(work), str(cache), str(runs),
               str(ROOT / "src/librecalc_agent/_bootstrap"),
               str(ROOT / "src/librecalc_agent/_capture_helper.py")]
    env = {**os.environ, "PYTHONPATH": str(ROOT / "src")}
    result = subprocess.run(command, cwd=work, env=env, capture_output=True,
                            timeout=20)
    assert result.returncode == 127
    pointer = json.loads((runs / "last_run.json").read_text())
    receipt = json.loads((Path(pointer["run_dir"]) / "observer_receipt.json").read_text())
    assert receipt["target_exit_code"] == 127
    assert receipt["target_signal"] == 0
    assert receipt["assurance_status"] == "PASS"


def test_snapshot_file_count_bound_names_the_limit(work):
    for i in range(1001):
        (work / f"fill_{i:04d}.xlsx").write_bytes(b"")
    result = observed_raw(work, "print('must not run')\n")
    assert result.returncode == 125
    assert b"too many XLSX files (limit 1000)" in result.stderr
    assert b"must not run" not in result.stdout


def test_snapshot_aggregate_bound_names_the_limit(work):
    with open(work / "huge.xlsx", "wb") as stream:
        stream.truncate(600 * 1024 * 1024)
    result = observed_raw(work, "print('must not run')\n")
    assert result.returncode == 125
    assert b"exceeds 512 MiB aggregate limit" in result.stderr


def test_snapshot_unreadable_file_names_the_file(work):
    if os.geteuid() == 0:
        pytest.skip("root can read mode-000 files")
    locked = work / "locked.xlsx"
    locked.write_bytes(b"")
    locked.chmod(0o000)
    try:
        result = observed_raw(work, "print('must not run')\n")
    finally:
        locked.chmod(0o644)
    assert result.returncode == 125
    assert b"cannot read XLSX file 'locked.xlsx'" in result.stderr


@pytest.mark.parametrize("ending,expected", [
    ("", 0), ("raise SystemExit(7)", 7), ("raise RuntimeError('expected')", 1),
])
def test_script_process_metadata_and_exit(work, ending, expected):
    script = ('import sys,os,__main__\n'
              'print(__name__,__main__.__file__==__file__,sys.argv[1:],'
              'sys.path[0]==os.path.dirname(__file__),os.getcwd())\n' + ending + '\n')
    target = work / "fixture.py"
    target.write_text(script)
    direct = subprocess.run([sys.executable, str(target), "arg"], cwd=work,
                            capture_output=True, timeout=20)
    product, receipt, _ = observed(work, script, ["arg"])
    assert direct.returncode == product.returncode == expected
    assert direct.stdout == product.stdout
    if expected == 0:
        assert direct.stderr == product.stderr
    else:
        assert b"expected" in product.stderr or expected == 7
    assert receipt["target_exit_code"] == expected
    assert receipt["assurance_status"] == "PASS"


@pytest.mark.parametrize("ending,signal_number", [
    ("os._exit(7)", 0),
    ("os.kill(os.getpid(), signal.SIGTERM)", signal.SIGTERM),
])
def test_abrupt_exit_preserves_changed_workbook_observation(work, ending, signal_number):
    script = ('import openpyxl,os,signal\n'
              'w=openpyxl.load_workbook("input.xlsx")\n'
              'w.active["A1"]=99\nw.save("input.xlsx")\n' + ending + '\n')
    product, receipt, run_dir = observed(work, script)
    assert product.returncode == (-signal_number if signal_number else 7)
    assert receipt["target_signal"] == signal_number
    assert receipt["changed_xlsx"] == 1
    assert receipt["capture_helper_exit"] == 0
    assert receipt["assurance_status"] == "PASS"
    capture = json.loads((run_dir / "capture_state.json").read_text())
    assert capture["validation_passed"] is True
    assert openpyxl.load_workbook(work / "input.xlsx").active["A1"].value == 99


def test_atexit_effect_is_observed_after_hook(work):
    script = ('import atexit,openpyxl\n'
              'def finish():\n'
              ' w=openpyxl.load_workbook("input.xlsx")\n'
              ' w.active["A1"]=77\n'
              ' w.save("input.xlsx")\n'
              'atexit.register(finish)\n')
    product, receipt, run_dir = observed(work, script)
    assert product.returncode == 0
    assert receipt["changed_xlsx"] == 1 and receipt["assurance_status"] == "PASS"
    assert json.loads((run_dir / "capture_state.json").read_text())["validation_passed"]
    assert openpyxl.load_workbook(work / "input.xlsx").active["A1"].value == 77


def test_subprocess_and_file_descriptor(work):
    script = ('import os,subprocess,sys\n'
              'print(subprocess.check_output([sys.executable,"-c","print(123)"],'
              'text=True).strip())\n'
              'print(os.environ.get("PRODUCT_FD_TEST"))\n')
    target = work / "fixture.py"
    target.write_text(script)
    env = {**os.environ, "PRODUCT_FD_TEST": "visible"}
    direct = subprocess.run([sys.executable, str(target)], cwd=work, env=env,
                            capture_output=True)
    product, receipt, _ = observed(work, script)
    assert direct.returncode == product.returncode == 0
    assert receipt["assurance_status"] == "PASS"
    assert b"123" in product.stdout


def test_inherited_descriptor_matches_direct_python(work):
    script = ('import os\n'
              'fd=int(os.environ["PRODUCT_TEST_FD"])\n'
              'print(os.read(fd, 8).decode())\n')
    target = work / "fixture.py"
    target.write_text(script)
    for arm in ("direct", "product"):
        read_fd, write_fd = os.pipe()
        try:
            os.write(write_fd, b"present")
            env = {**os.environ, "PRODUCT_TEST_FD": str(read_fd)}
            if arm == "direct":
                result = subprocess.run([sys.executable, str(target)], cwd=work,
                                        env=env, capture_output=True,
                                        pass_fds=(read_fd,), timeout=20)
            else:
                result, receipt, _ = observed(work, script,
                                              env_extra={"PRODUCT_TEST_FD": str(read_fd)},
                                              pass_fds=(read_fd,))
                assert receipt["assurance_status"] == "PASS"
            assert result.returncode == 0
            assert result.stdout == b"present\n"
        finally:
            os.close(read_fd)
            os.close(write_fd)


def test_caught_sigint_matches_direct_python(work):
    script = ('import os,signal\n'
              'signal.signal(signal.SIGINT, lambda *_: print("caught"))\n'
              'os.kill(os.getpid(), signal.SIGINT)\n')
    target = work / "fixture.py"
    target.write_text(script)
    direct = subprocess.run([sys.executable, str(target)], cwd=work,
                            capture_output=True, timeout=20)
    product, receipt, _ = observed(work, script)
    assert direct.returncode == product.returncode == 0
    assert direct.stdout == product.stdout == b"caught\n"
    assert receipt["target_signal"] == 0
    assert receipt["assurance_status"] == "PASS"
