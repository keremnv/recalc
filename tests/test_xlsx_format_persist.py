"""set_format must survive as RGB the official evaluator can read.

UNO CharColor can be live-true while the saved xlsx still has the input font xf.
SpreadsheetBench compare_font_color uses openpyxl, never LibreOffice.
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Color, Font

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "benchmark-data/SpreadsheetBench-2/evaluation"))

import evaluation as ev

from librecalc_mcp.backend.uno import _a1_addresses, _persist_xlsx_formats


def test_a1_addresses_expand_ranges_in_row_major_order() -> None:
    assert _a1_addresses("E4") == ["E4"]
    assert _a1_addresses("B2:C3") == ["B2", "C2", "B3", "C3"]


def test_persist_writes_font_rgb_compare_font_color_accepts(tmp_path: Path) -> None:
    source = tmp_path / "input.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Model"
    sheet["E4"] = 0
    sheet["E4"].font = Font(name="Times New Roman", color="FF000000")
    sheet["V3"] = 1
    sheet["V3"].font = Font(name="Times New Roman", color="FF000000")
    workbook.save(source)

    _persist_xlsx_formats(
        str(source),
        [
            ("Model", "E4", {"font_color": "#0000FF"}),
            ("Model", "V3", {"font_color": "#0000FF"}),
        ],
    )

    golden = Workbook()
    golden.active.title = "Model"
    golden.active["E4"].font = Font(color="FF0000FF")
    processed = load_workbook(source)
    assert ev.compare_font_color(golden.active["E4"].font, processed["Model"]["E4"].font)
    assert ev.compare_font_color(golden.active["E4"].font, processed["Model"]["V3"].font)
    assert processed["Model"]["E4"].value == 0
    assert processed["Model"]["E4"].font.name == "Times New Roman"


def test_persist_writes_fill_and_bold_without_touching_other_cells(tmp_path: Path) -> None:
    source = tmp_path / "fills.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet["B2"] = "keep"
    sheet["C2"] = "paint"
    sheet["C2"].font = Font(color="FF000000", name="Calibri")
    workbook.save(source)

    _persist_xlsx_formats(
        str(source),
        [
            (
                "Sheet1",
                "C2",
                {
                    "font_color": "#00B050",
                    "background_color": "#FFF2CC",
                    "background_transparent": False,
                    "font_weight": 150,
                },
            )
        ],
    )

    processed = load_workbook(source)
    painted = processed["Sheet1"]["C2"]
    untouched = processed["Sheet1"]["B2"]
    assert painted.font.color.rgb[-6:] == "00B050"
    assert painted.font.bold is True
    assert painted.fill.fgColor.rgb[-6:] == "FFF2CC"
    assert untouched.font.bold is not True
    assert untouched.value == "keep"


def test_persist_restores_unpatched_theme_fonts_from_source(tmp_path: Path) -> None:
    """Container LO writes workbook-theme RGB; the evaluator maps Office theme 8 + tint."""

    source = tmp_path / "themed.xlsx"
    dest = tmp_path / "flattened.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "LBO"
    sheet["C6"] = "label"
    sheet["C6"].font = Font(name="Calibri", color=Color(theme=8, tint=-0.25))
    sheet["J27"] = 1
    sheet["J27"].font = Font(name="Calibri", color="FF000000")
    workbook.save(source)

    flattened = Workbook()
    flat_sheet = flattened.active
    flat_sheet.title = "LBO"
    flat_sheet["C6"] = "label"
    flat_sheet["C6"].font = Font(name="Calibri", color="FF2F5597")
    flat_sheet["J27"] = 1
    flat_sheet["J27"].font = Font(name="Calibri", color="FF000000")
    flattened.save(dest)

    _persist_xlsx_formats(
        str(dest),
        [("LBO", "J27", {"font_color": "#0000FF"})],
        source_path=str(source),
    )

    processed = load_workbook(dest)
    golden = Workbook()
    golden.active["C6"].font = Font(color=Color(theme=8, tint=-0.25))
    golden.active["J27"].font = Font(color="FF0000FF")
    assert ev.compare_font_color(golden.active["C6"].font, processed["LBO"]["C6"].font)
    assert ev.compare_font_color(golden.active["J27"].font, processed["LBO"]["J27"].font)
    assert processed["LBO"]["C6"].font.color.type == "theme"
    assert processed["LBO"]["C6"].font.color.theme == 8
