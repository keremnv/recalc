import copy
import json
from pathlib import Path
import sqlite3
import sys
from types import SimpleNamespace

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "benchmark"))
import compiled_scheduler
import matched_compiled_treatment as m


def fixture_runtime(tmp_path, *, groups=None, precedents=None):
    members = [("S", 10, c) for c in (3, 4, 5)]
    meta = {c: {"operation_id": "op", "obligation_id": "o", "operation_kind": "SET_FORMULA"} for c in members}
    cids = {c: f"cell:s00:r{c[1]}:c{c[2]}" for c in members}
    calls = []
    def session(*args, **kwargs):
        target = args[3]; calls.append(target["cell_id"])
        return {"synthesis": {"parsed": {"status": "PROPOSED", "formula": "=A10"}}, "calls": []}
    pg = SimpleNamespace(groups_for=lambda *args, **kw: (groups or [], []), translate=m.program_group.translate)
    runtime = SimpleNamespace(spine_for=lambda k: {}, authorised_cells=lambda p: (meta, cids), task_source=lambda k: tmp_path / "input.xlsx", formula_forms=lambda p: {}, input_kinds=lambda p: {}, db_for=lambda k: None, db_precedent_graph=lambda p: {}, program_group=pg, closure=m.closure, TaskBudget=m.TaskBudget, write_json=m.write_json, target_from_id=lambda sp, cid, k, oid: {"cell_id": cid, "sheet": "S", "address": "C10"}, retrieval_synthesis=session, validate_formula=lambda *args: {"hard_verifier_result": "HARD_ACCEPT"}, proposal_precedents=lambda f, c: set(precedents or []))
    return runtime, members, calls


def run(runtime, tmp_path, state=None):
    return compiled_scheduler.schedule(runtime, "task", {}, {"obligations": [{"id": "o"}]}, {"expansion": {"operations": [{"operation_id": "op"}]}, "packets": {"o": {}}}, state if state is not None else {}, tmp_path)


def test_closure_members_are_solved_not_marked_complete(tmp_path):
    runtime, members, calls = fixture_runtime(tmp_path, precedents=[("S", 10, 4), ("OTHER", 1, 1)])
    result = run(runtime, tmp_path)
    assert len(calls) == 3
    assert len(result["edits"]) == 3
    assert all(set(map(tuple, u["execution_members"])) <= set(members) for u in result["groups"])


def test_seed_not_translated_into_unrelated_prerequisite_group(tmp_path):
    group = {"canonical_cell": ["S", 10, 4], "member_cells": [["S", 10, 4], ["S", 10, 5]]}
    runtime, _, calls = fixture_runtime(tmp_path, groups=[group], precedents=[("S", 10, 4), ("S", 10, 5)])
    result = run(runtime, tmp_path)
    assert calls == ["cell:s00:r10:c3", "cell:s00:r10:c4"]
    assert len(result["edits"]) == 3
    assert result["translated_count"] == 1
    assert result["groups"][0]["groups"] == []


def test_translated_members_never_resynthesized_and_resume_is_durable(tmp_path):
    group = {"canonical_cell": ["S", 10, 3], "member_cells": [["S", 10, c] for c in (3, 4, 5)]}
    runtime, _, calls = fixture_runtime(tmp_path, groups=[group])
    first = run(runtime, tmp_path)
    resumed_state = json.loads((tmp_path / "state.json").read_text())
    second = run(runtime, tmp_path, resumed_state)
    assert len(calls) == 1
    assert first["edits"] == second["edits"]
    assert len(first["edits"]) == 3
    assert first["unresolved_authorised_targets"] == []


def test_hard_reject_and_abstain_have_explicit_dispositions(tmp_path):
    runtime, _, _ = fixture_runtime(tmp_path)
    runtime.validate_formula = lambda *args: {"hard_verifier_result": "HARD_REJECT"}
    result = run(runtime, tmp_path)
    assert result["edits"] == []
    assert {v["status"] for v in result["dispositions"].values()} == {"REJECTED_HARD"}


def test_persisted_request_resume_recounts_cost_without_provider_call(tmp_path, monkeypatch):
    request = m.request_body("system", "user")
    record = {"request_sha256": m.digest(request), "stage": "synthesis", "call_index_within_task": 1, "provider_cost_usd": .25, "raw_response_body": {"choices": []}}
    m.persist_call(tmp_path, record)
    monkeypatch.setattr(m, "model_call", lambda *a, **k: pytest.fail("new model call"))
    state = {"_resume_mode": True, "model_call_count": 0, "provider_cost_usd": 0}
    assert m.call_or_stub(tmp_path, "t", "synthesis", "system", "user", state, stub=False) == record
    assert state["model_call_count"] == 1
    assert state["provider_cost_usd"] == .25
    with pytest.raises(RuntimeError, match="IMMUTABLE_CALL_COLLISION"):
        m.persist_call(tmp_path, record)


def test_budget_block_is_not_a_provider_call_artifact(tmp_path, monkeypatch):
    state = {"model_call_count": m.MAX_MODEL_CALLS}
    monkeypatch.setattr(m, "model_call", lambda *a, **k: pytest.fail("provider called after budget"))
    result = m.call_or_stub(tmp_path, "t", "synthesis", "s", "u", state, stub=False)
    assert result["provider_attempt"] is False
    assert not list((tmp_path / "calls").glob("*"))


def test_synthesis_transition_has_no_retrieval_prohibition():
    assert "Do not synthesize a formula now" not in m.SYNTHESIS_SYSTEM
    assert "ARCHITECTURE V1.1 SYNTHESIS PROTOCOL" in m.SYNTHESIS_SYSTEM


def test_materializer_preserves_more_than_sql_read_limit(tmp_path):
    from workbook_spine_sqlite import SCHEMA_SQL
    db = tmp_path / "world.sqlite"
    conn = sqlite3.connect(db)
    conn.executescript(SCHEMA_SQL)
    n = m.relational.RESULT_ROW_LIMIT + 1
    for i in range(n):
        conn.execute("INSERT INTO formula_classes VALUES (?,?)", (f"formula_class:{i}", "=1"))
    conn.commit(); conn.close()
    result = m.materialize(db, {f"formula_class:{i}" for i in range(n)})
    assert len(result["entities"]["formula_classes"]["rows"]) == n


def test_materializer_includes_authorized_implicit_blank(tmp_path):
    from workbook_spine_sqlite import SCHEMA_SQL
    db = tmp_path / "world.sqlite"
    conn = sqlite3.connect(db); conn.executescript(SCHEMA_SQL)
    conn.execute("INSERT INTO sheets VALUES (?,?,?,?,?,?,?,?,?,?)", ("sheet:s00", "wb", 0, "S", "s", "visible", 1, 1, 10, 10))
    conn.commit(); conn.close()
    result = m.materialize(db, {"cell:s00:r8:c3"})
    table = result["entities"]["cells"]
    row = dict(zip(table["columns"], table["rows"][0]))
    assert row["kind"] == "blank"
    assert row["implicit_blank"] is True


def test_sparse_evaluator_reads_preserve_blank_value_formula_and_style(tmp_path):
    import openpyxl
    import integration_autopsy as a
    wb = openpyxl.Workbook()
    wb.active["A1"] = "=1+2"
    path = tmp_path / "input.xlsx"
    wb.save(path)
    regular = openpyxl.load_workbook(path)
    sparse = a.sparse_workbook(path, data_only=False)
    before = len(sparse.active._cells)
    for address in ("A1", "B2", "XFD1000"):
        left, right = regular.active[address], sparse.active[address]
        assert left.value == right.value
        assert left.data_type == right.data_type
        assert left.font == copy.copy(right.font)
        assert left.number_format == right.number_format
    assert len(sparse.active._cells) == before
    sparse.close(); regular.close()


def test_atomic_state_write_keeps_previous_state_if_replace_fails(tmp_path, monkeypatch):
    path = tmp_path / "state.json"
    m.write_json(path, {"completed": ["C10"]})
    def fail(*args):
        raise OSError("simulated interruption")
    monkeypatch.setattr(m.os, "replace", fail)
    with pytest.raises(OSError):
        m.write_json(path, {"completed": ["C10", "D10"]})
    assert m.read_json(path) == {"completed": ["C10"]}


def test_target_content_comes_from_compiled_payload_and_preserves_zero():
    spine = {"sheets": [{"index": 0, "title": "S"}], "occupied": [{"id": "cell:s00:r1:c1", "payload": 0}]}
    target = m.target_from_id(spine, "cell:s00:r1:c1", "unused", "o")
    assert target["current_input_content"] == 0
    assert target["current_input_kind"] == "value"


def test_compact_score_matches_official_comparator(tmp_path):
    import openpyxl
    import integration_autopsy as a
    sys.path.insert(0, str(m.ROOT / "benchmark-data/SpreadsheetBench-2/evaluation"))
    import evaluation as ev
    wb = openpyxl.Workbook()
    wb.active["A1"] = "=SUM(B1:C1)"
    wb.active["B1"] = 0
    wb.active["C1"] = "#REF!"
    wb.active["D1"] = False
    path = tmp_path / "workbook.xlsx"
    wb.save(path)
    normal = openpyxl.load_workbook(path)
    compact = a.compact_workbook(path, data_only=False)
    for addr in ("A1", "B1", "C1", "D1", "E1"):
        assert ev.compare_cell_formula(normal.active[addr], compact["Sheet"][addr])
        assert ev._compare_cells(normal.active[addr], compact["Sheet"][addr], False, False)
    normal.close()


@pytest.mark.parametrize('existing_placeholder', [False, True])
def test_environment_materializes_temporal_coordinates_instead_of_placeholder(tmp_path, monkeypatch, existing_placeholder):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active["A1"] = "Item"
    wb.active["B1"] = 45100
    wb.active["B1"].number_format = "MMM-YY"
    wb.active["A2"] = "Revenue"
    source = tmp_path / "source.xlsx"
    wb.save(source)
    for name, folder in (("PREP", "prep"), ("INPUTS", "inputs"), ("SPINES", "spines"), ("TEMPORAL", "temporal"), ("DATABASES", "db")):
        monkeypatch.setattr(m, name, tmp_path / folder)
    if existing_placeholder:
        m.write_json(tmp_path / 'temporal/Financial_Model-test.json', {'coordinates': [], 'n_coordinates': 0})
    monkeypatch.setattr(m, "task_rows", lambda: [{"task_key": "Financial_Model:test", "category": "Financial_Model", "task": "test", "input_path": str(source)}])
    result = m.build_environment()
    assert result["all_inputs_present"]
    conn = sqlite3.connect(tmp_path / "db/Financial_Model-test.sqlite")
    assert conn.execute("SELECT COUNT(*) FROM temporal_coordinates WHERE year=2023").fetchone()[0] > 0
    conn.close()


def test_noop_seed_is_not_sent_to_writer_but_still_coordinates_group(tmp_path):
    group = {"canonical_cell": ["S", 10, 3], "member_cells": [["S", 10, c] for c in (3, 4, 5)]}
    runtime, _, calls = fixture_runtime(tmp_path, groups=[group])
    runtime.formula_forms = lambda p: {("S", 10, 3): "=A10"}
    result = run(runtime, tmp_path)
    assert len(calls) == 1
    assert result["dispositions"]["cell:s00:r10:c3"]["status"] == "NO_SEMANTIC_CHANGE"
    assert {e["address"] for e in result["edits"]} == {"D10", "E10"}


def test_last_model_opportunity_is_reserved_for_synthesis(tmp_path, monkeypatch):
    stages = []
    monkeypatch.setattr(m, "spine_for", lambda k: {})
    monkeypatch.setattr(m, "db_for", lambda k: None)
    monkeypatch.setattr(m.prior, "compile_bootstrap", lambda *a: {"bootstrap_entity_ids": []})
    monkeypatch.setattr(m, "ReadOnlySqlite", lambda *a, **k: None)
    monkeypatch.setattr(m, "materialize", lambda *a: {})
    def call(*args, **kwargs):
        stages.append(args[2])
        args[5]["model_call_count"] += 1
        return {"stage": args[2], "parsed_response": {"status": "ABSTAIN"}}
    monkeypatch.setattr(m, "call_or_stub", call)
    target = {"cell_id": "cell:s00:r10:c3", "sheet": "S", "address": "C10"}
    state = {"model_call_count": m.MAX_MODEL_CALLS - 1}
    result = m.retrieval_synthesis("task", {"instruction": "test"}, {}, target, {}, state, tmp_path, stub=False)
    assert stages == ["synthesis"]
    assert result["sql_calls"] == 0


def test_retrieval_continuation_receives_actual_entity_delta(tmp_path, monkeypatch):
    sent = []
    monkeypatch.setattr(m, "spine_for", lambda k: {})
    monkeypatch.setattr(m, "db_for", lambda k: None)
    monkeypatch.setattr(m.prior, "compile_bootstrap", lambda *a: {"bootstrap_entity_ids": []})
    entity = "cell:s00:r10:c1"
    executor = SimpleNamespace(execute=lambda sql: {"status": "OK", "rows": [{"cell_id": entity}], "row_count": 1})
    monkeypatch.setattr(m, "ReadOnlySqlite", lambda *a, **k: executor)
    monkeypatch.setattr(m, "materialize", lambda *a: {})
    def call(*args, **kwargs):
        sent.append((args[2], json.loads(args[4])))
        args[5]["model_call_count"] = args[5].get("model_call_count", 0) + 1
        action = {"action": "execute_sql", "sql": "SELECT cell_id FROM cells"} if len(sent) == 1 else {"action": "final", "status": "ENOUGH_EVIDENCE"} if args[2] == "retrieval" else {"status": "ABSTAIN"}
        return {"stage": args[2], "parsed_response": action}
    monkeypatch.setattr(m, "call_or_stub", call)
    target = {"cell_id": "cell:s00:r10:c3", "sheet": "S", "address": "C10"}
    m.retrieval_synthesis("task", {"instruction": "test"}, {}, target, {}, {}, tmp_path, stub=False)
    assert sent[1][1]["NEW_SINCE_LAST_TURN"] == {"count": 1, "entity_ids": [entity]}
