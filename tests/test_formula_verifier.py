"""Unit tests for gold-blind formula candidate verification."""
from __future__ import annotations

import inspect
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "benchmark"))
sys.path.insert(0, str(ROOT / "benchmark/sweagent/formula_index/lib"))

import formula_verifier_probe  # noqa: E402
from formula_dependency_selection import build_graph  # noqa: E402
from formula_operational import build_view  # noqa: E402
from formula_verifier import (  # noqa: E402
    DEFINITIONS,
    apply_policy,
    candidate_kind,
    operator_skeleton,
    scc_stats,
    verify_formula,
)


def _save(tmp_path: Path, fill) -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Assumptions"
    fill(sheet, workbook)
    path = tmp_path / "in.xlsx"
    workbook.save(path)
    workbook.close()
    return path


def _by_id(results: list[dict]) -> dict[str, dict]:
    return {row["checker_id"]: row for row in results}


def test_definitions_and_extract_are_gold_blind() -> None:
    assert DEFINITIONS["golden_in_generation"] is False
    source = inspect.getsource(sys.modules["formula_verifier"])
    assert "_golden_path" not in source
    assert "eval_golden_formula" not in source
    freeze_src = inspect.getsource(formula_verifier_probe.cmd_extract)
    assert "_golden_path" not in freeze_src
    assert "eval_golden_formula" not in freeze_src


def test_absence_of_violation_is_not_automatic_pass(tmp_path: Path) -> None:
    def fill(sheet, _wb) -> None:
        sheet["K6"] = None
        sheet["J6"] = 1

    view = build_view(build_graph(_save(tmp_path, fill)))
    out = verify_formula(view, ("Assumptions", 11, 6), "=J6*(1+K27)")
    b2 = _by_id(out["results"])["B2_k3"]
    assert b2["verdict"] == "ABSTAIN"
    assert b2["applicable"] is False


def test_direct_cycle_rejects_wrong_and_allows_gold(tmp_path: Path) -> None:
    def fill(sheet, wb) -> None:
        other = wb.create_sheet("Financials")
        sheet["K163"] = None
        other["K82"] = "=Assumptions!K163"
        sheet["J163"] = "=Financials!J82"
        other["J82"] = 10
        sheet["H163"] = "=Financials!H82"
        other["H82"] = 10
        sheet["I163"] = "=Financials!I82"
        other["I82"] = 10
        sheet["K164"] = 0.05
        sheet["K18"] = 100

    view = build_view(build_graph(_save(tmp_path, fill)))
    target = ("Assumptions", 11, 163)
    wrong = verify_formula(view, target, "=Financials!K82")
    gold = verify_formula(view, target, "=K164*K18")
    w = _by_id(wrong["results"])
    g = _by_id(gold["results"])
    assert w["A2"]["verdict"] == "REJECT"
    assert g["A2"]["verdict"] == "PASS"
    h_w = {p["policy"]: p["verdict"] for p in wrong["policies"]}
    h_g = {p["policy"]: p["verdict"] for p in gold["policies"]}
    assert h_w["H"] == "REJECT"
    assert h_g["H"] != "REJECT"


def test_sum_versus_plus_does_not_reject() -> None:
    sk_sum = operator_skeleton("=SUM(C8:C11)")
    sk_plus = operator_skeleton("=C8+C9+C10+C11")
    assert sk_sum != sk_plus
    assert candidate_kind("=SUM(C8:C11)") == "REDUCTION_CANDIDATE"
    assert candidate_kind("=C8+C9+C10+C11") == "REDUCTION_CANDIDATE"


def test_d2_abstains_on_sum_versus_explicit_plus(tmp_path: Path) -> None:
    def fill(sheet, _wb) -> None:
        sheet["C12"] = None
        for col, letter in enumerate("DEFGH", start=4):
            sheet.cell(12, col, f"=SUM({letter}8:{letter}11)")
            for row in range(8, 12):
                sheet.cell(row, col, row)

    view = build_view(build_graph(_save(tmp_path, fill)))
    target = ("Assumptions", 3, 12)
    plus = verify_formula(view, target, "=C8+C9+C10+C11")
    summed = verify_formula(view, target, "=SUM(C8:C11)")
    d2_plus = _by_id(plus["results"])["D2_k3"]
    d2_sum = _by_id(summed["results"])["D2_k3"]
    assert d2_plus["verdict"] != "REJECT"
    assert d2_sum["verdict"] in {"PASS", "ABSTAIN"}
    if d2_plus["verdict"] == "ABSTAIN":
        assert "sum_vs_explicit_plus" in d2_plus["reason"] or d2_plus["reason"] in {
            "insufficient_reduction_peers",
            "peers_not_unanimous",
            "sum_vs_explicit_plus_equivalent_syntax",
        }


def test_self_reference_is_hard_reject(tmp_path: Path) -> None:
    def fill(sheet, _wb) -> None:
        sheet["A1"] = None

    view = build_view(build_graph(_save(tmp_path, fill)))
    out = verify_formula(view, ("Assumptions", 1, 1), "=A1+1")
    assert _by_id(out["results"])["A1"]["verdict"] == "REJECT"


def test_policy_h_ignores_empirical_reject(tmp_path: Path) -> None:
    def fill(sheet, _wb) -> None:
        sheet["K6"] = None
        sheet["J6"] = "=I6*(1+J27)"
        sheet["I6"] = 1
        sheet["J27"] = 0.1
        sheet["H6"] = "=G6*(1+H27)"
        sheet["G6"] = 1
        sheet["H27"] = 0.1
        sheet["L6"] = "=K6*(1+L27)"
        sheet["L27"] = 0.1

    view = build_view(build_graph(_save(tmp_path, fill)))
    out = verify_formula(view, ("Assumptions", 11, 6), "=J6*(1+K12)")
    results = out["results"]
    # Force an empirical reject into policy mix without a hard reject.
    empirical = [r for r in results if r["checker_id"] == "B2_k3"]
    hard = [r for r in results if r["checker_id"] in {"A1", "A2", "A3", "A4", "A5"}]
    assert apply_policy(hard + empirical, "H")["verdict"] != "REJECT" or any(
        r["verdict"] == "REJECT" for r in hard
    )
    h = {p["policy"]: p for p in out["policies"]}
    if all(r["verdict"] != "REJECT" for r in hard):
        assert h["H"]["verdict"] != "REJECT"


def test_acyclic_scc_on_simple_graph(tmp_path: Path) -> None:
    def fill(sheet, _wb) -> None:
        sheet["B1"] = "=A1"
        sheet["A1"] = 1

    view = build_view(build_graph(_save(tmp_path, fill)))
    stats = scc_stats(view)
    assert stats["acyclic"] is True
    assert stats["n_nontrivial_scc"] == 0


def test_historical_label_fingerprint_and_sum_plus() -> None:
    assert formula_verifier_probe._label_historical("=A1+B1", "=A1+B1", "S", 1, 2) == "CORRECT"
    assert formula_verifier_probe._label_historical("=C8+C9+C10+C11", "=SUM(C8:C11)", "DCF", 3, 12) == (
        "EQUIVALENT_CORRECT"
    )
    assert formula_verifier_probe._label_historical("=K164*K18", "=Financials!K82", "Assumptions", 11, 163) == "WRONG"
    assert formula_verifier_probe._label_historical("=A1", None, "S", 1, 1) == "WRONG"
