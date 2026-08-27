import importlib.util
import os
import tempfile
from pathlib import Path

import pytest

from librecalc_mcp.domain.models import CalcOperation


def _calc_tool_module():
    path = Path(__file__).parents[1] / "benchmark/sweagent/librecalc/lib/calc_tool.py"
    spec = importlib.util.spec_from_file_location("benchmark_calc_tool", path)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_calc_tool_read_budget_blocks_second_read(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    calc_tool = _calc_tool_module()
    monkeypatch.setenv("LIBRECALC_READ_BUDGET_ENABLED", "1")
    monkeypatch.setenv("LIBRECALC_READ_BUDGET_PATH", str(tmp_path / "budget.json"))
    calc_tool._read_budget().reset_read_budget()
    calc_tool._read_budget().consume_read_budget(successful=True)
    assert calc_tool._read_budget().read_budget_error() is not None


@pytest.mark.skipif(os.environ.get("LIBRECALC_RUN_UNO") != "1", reason="requires LibreOffice UNO")
def test_uno_chart_xlsx_roundtrip_openpyxl_count() -> None:
    openpyxl = pytest.importorskip("openpyxl")
    from librecalc_mcp.backend.uno import UnoCalcBackend

    with tempfile.TemporaryDirectory() as temporary:
        root = Path(temporary)
        seed = root / "seed.xlsx"
        source = root / "chart_source.xlsx"
        workbook = openpyxl.Workbook()
        workbook.active.title = "Sheet1"
        workbook.save(seed)
        backend = UnoCalcBackend()
        backend.execute_program(
            [
                CalcOperation(
                    op="write_range",
                    sheet="Sheet1",
                    range="A1:B4",
                    values=[
                        ["Month", "Sales"],
                        ["Jan", 10],
                        ["Feb", 20],
                        ["Mar", 30],
                    ],
                ),
                CalcOperation.from_dict(
                    {
                        "op": "upsert_chart",
                        "chart": {
                            "id": "SalesLine",
                            "sheet": "Sheet1",
                            "chart_type": "line",
                            "category_range": "A2:A4",
                            "series": [{"name": "Sales", "values_range": "B2:B4"}],
                            "title": "Sales",
                            "anchor": "D2",
                        },
                    }
                ),
            ],
            path=str(seed),
            output_path=str(source),
        )
        workbook = openpyxl.load_workbook(source)
        chart_count = sum(len(getattr(sheet, "_charts", []) or []) for sheet in workbook.worksheets)
        assert chart_count >= 1


@pytest.mark.skipif(os.environ.get("LIBRECALC_RUN_UNO") != "1", reason="requires LibreOffice UNO")
def test_uno_chart_axis_titles_data_labels_and_point_colors() -> None:
    from librecalc_mcp.backend.uno import UnoCalcBackend

    with tempfile.TemporaryDirectory() as temporary:
        root = Path(temporary)
        source = root / "in.xlsx"
        output = root / "styled.xlsx"
        # Minimal xlsx seed so UNO has a path to open.
        try:
            from openpyxl import Workbook

            wb = Workbook()
            wb.active.title = "Sheet1"
            wb.save(source)
        except Exception:
            pytest.skip("openpyxl required to seed workbook")

        backend = UnoCalcBackend()
        backend.execute_program(
            [
                CalcOperation(
                    op="write_range",
                    sheet="Sheet1",
                    range="A1:B5",
                    values=[
                        ["Cat", "Val"],
                        ["A", 10],
                        ["B", 20],
                        ["C", 30],
                        ["D", 40],
                    ],
                ),
                CalcOperation.from_dict(
                    {
                        "op": "upsert_chart",
                        "chart": {
                            "id": "Styled",
                            "sheet": "Sheet1",
                            "chart_type": "column",
                            "category_range": "A2:A5",
                            "series": [
                                {
                                    "name": "Val",
                                    "values_range": "B2:B5",
                                    "color": "#3366FF",
                                    "point_colors": ["#00AA00", "#FF0000", "#0000FF", "#FFAA00"],
                                }
                            ],
                            "title": "Styled Columns",
                            "legend": False,
                            "data_labels": True,
                            "x_axis": {"title": "Category"},
                            "y_axis": {"title": "Value"},
                            "anchor": "D2",
                            "width": 10,
                            "height": 12,
                        },
                    }
                ),
            ],
            path=str(source),
            output_path=str(output),
        )

        backend._connect()
        uno = backend._uno()
        doc = backend._desktop.loadComponentFromURL(
            uno.systemPathToFileUrl(str(output.resolve())), "_blank", 0, ()
        )
        try:
            chart_doc = doc.Sheets.getByName("Sheet1").getCharts().getByIndex(0).getEmbeddedObject()
            diagram = chart_doc.getDiagram()
            assert getattr(diagram, "HasXAxisTitle", False) is True
            assert diagram.XAxisTitle.String == "Category"
            assert getattr(diagram, "HasYAxisTitle", False) is True
            assert diagram.YAxisTitle.String == "Value"
            assert diagram.getDataRowProperties(0).DataCaption == 4  # TEXT

            series = (
                chart_doc.getFirstDiagram()
                .getCoordinateSystems()[0]
                .getChartTypes()[0]
                .getDataSeries()[0]
            )
            assert series.getDataPointByIndex(0).FillColor == 0x00AA00
            assert series.getDataPointByIndex(1).FillColor == 0xFF0000
            assert series.getDataPointByIndex(0).Label.ShowCategoryName is True
        finally:
            doc.close(True)


@pytest.mark.skipif(os.environ.get("LIBRECALC_RUN_UNO") != "1", reason="requires LibreOffice UNO")
def test_uno_bubble_round_trips_x_y_size_labels_id_and_axis_format() -> None:
    from openpyxl import Workbook

    from librecalc_mcp.backend.uno import UnoCalcBackend

    with tempfile.TemporaryDirectory() as temporary:
        root = Path(temporary)
        source = root / "in.xlsx"
        output = root / "bubble.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Strategy"
        rows = [
            ["Product", "Revenue", "Growth", "Market Size"],
            ["A", 10, 0.1, 500],
            ["B", 20, 0.2, 300],
            ["C", 30, 0.3, 1000],
        ]
        for row_index, row in enumerate(rows, start=1):
            for column_index, value in enumerate(row, start=1):
                sheet.cell(row_index, column_index, value)
        workbook.save(source)

        backend = UnoCalcBackend()
        result = backend.execute_program(
            [
                CalcOperation.from_dict(
                    {
                        "op": "upsert_chart",
                        "chart": {
                            "id": "Portfolio",
                            "sheet": "Strategy",
                            "chart_type": "bubble",
                            "category_range": "A2:A4",
                            "series": [
                                {
                                    "name": "Growth",
                                    "x_values_range": "B2:B4",
                                    "values_range": "C2:C4",
                                    "bubble_size_range": "D2:D4",
                                    "point_colors": ["#3366FF", "#FF0000", "#00AA00"],
                                }
                            ],
                            "data_labels": True,
                            "x_axis": {"number_format": "0.0"},
                        },
                    }
                )
            ],
            path=str(source),
            output_path=str(output),
        )
        operation = result["operations"][0]
        assert operation["dropped"] == []

        chart = backend.inspect_charts(str(output))[0]
        assert chart["id"] == "Portfolio"
        assert chart["chart_type"] == "bubble"
        assert chart["category_labels"] == ["A", "B", "C"]
        assert chart["series"] == [
            {
                "name": "Growth",
                "name_range": "$Strategy.$C$1",
                "values_range": "$Strategy.$C$2:$C$4",
                "x_values_range": "$Strategy.$B$2:$B$4",
                "bubble_size_range": "$Strategy.$D$2:$D$4",
                "point_colors": ["#3366FF", "#FF0000", "#00AA00"],
            }
        ]
        assert chart["data_labels"] is True
        assert chart["x_axis"]["number_format"] == "0.0"

        replacement = root / "replacement.xlsx"
        backend.execute_program(
            [
                CalcOperation.from_dict(
                    {
                        "op": "upsert_chart",
                        "chart": {
                            "id": "Portfolio",
                            "sheet": "Strategy",
                            "chart_type": "bubble",
                            "category_range": "A2:A4",
                            "series": [
                                {
                                    "name": "Growth",
                                    "x_values_range": "B2:B4",
                                    "values_range": "C2:C4",
                                    "bubble_size_range": "D2:D4",
                                }
                            ],
                            "title": "Replaced",
                        },
                    }
                )
            ],
            path=str(output),
            output_path=str(replacement),
        )
        replacement_charts = backend.inspect_charts(str(replacement))
        assert len(replacement_charts) == 1
        assert replacement_charts[0]["id"] == "Portfolio"
        assert replacement_charts[0]["title"] == "Replaced"

        deleted = root / "deleted.xlsx"
        backend.execute_program(
            [CalcOperation.from_dict({"op": "delete_chart", "name": "Portfolio"})],
            path=str(replacement),
            output_path=str(deleted),
        )
        assert backend.inspect_charts(str(deleted)) == []


@pytest.mark.skipif(os.environ.get("LIBRECALC_RUN_UNO") != "1", reason="requires LibreOffice UNO")
def test_uno_sliced_series_name_range_round_trips() -> None:
    from openpyxl import Workbook

    from librecalc_mcp.backend.uno import UnoCalcBackend

    with tempfile.TemporaryDirectory() as temporary:
        root = Path(temporary)
        source = root / "in.xlsx"
        dropped_output = root / "dropped.xlsx"
        linked_output = root / "linked.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sheet1"
        sheet["A1"] = "Unique ID"
        sheet["B1"] = "Foo"
        sheet["C1"] = "Faa"
        for row in range(2, 42):
            sheet.cell(row, 1, row - 1)
            sheet.cell(row, 2, 1)
            sheet.cell(row, 3, 2)
        workbook.save(source)

        backend = UnoCalcBackend()
        dropped = backend.execute_program(
            [
                CalcOperation.from_dict(
                    {
                        "op": "upsert_chart",
                        "chart": {
                            "id": "slice",
                            "sheet": "Sheet1",
                            "chart_type": "column",
                            "category_range": "A22:A41",
                            "series": [
                                {"name": "Foo", "values_range": "B22:B41"},
                                {"name": "Faa", "values_range": "C22:C41"},
                            ],
                            "title": "Unique ID 21-40 Analysis",
                        },
                    }
                )
            ],
            path=str(source),
            output_path=str(dropped_output),
        )
        assert dropped["operations"][0]["dropped"] == [
            "series[0].name (literal series names require a matching header cell or name_range)",
            "series[1].name (literal series names require a matching header cell or name_range)",
        ]
        assert [
            item["name"] for item in backend.inspect_charts(str(dropped_output))[0]["series"]
        ] == [
            None,
            None,
        ]

        linked = backend.execute_program(
            [
                CalcOperation.from_dict(
                    {
                        "op": "upsert_chart",
                        "chart": {
                            "id": "slice",
                            "sheet": "Sheet1",
                            "chart_type": "column",
                            "category_range": "A22:A41",
                            "series": [
                                {
                                    "name": "Foo",
                                    "name_range": "B1",
                                    "values_range": "B22:B41",
                                },
                                {
                                    "name": "Faa",
                                    "name_range": "C1",
                                    "values_range": "C22:C41",
                                },
                            ],
                            "title": "Unique ID 21-40 Analysis",
                        },
                    }
                )
            ],
            path=str(source),
            output_path=str(linked_output),
        )
        assert linked["operations"][0]["dropped"] == []
        assert "series[0].name" in linked["operations"][0]["applied"]
        assert "series[1].name" in linked["operations"][0]["applied"]

        series = backend.inspect_charts(str(linked_output))[0]["series"]
        assert [
            (item["name"], item.get("name_range"), item["values_range"]) for item in series
        ] == [
            ("Foo", "$Sheet1.$B$1", "$Sheet1.$B$22:$B$41"),
            ("Faa", "$Sheet1.$C$1", "$Sheet1.$C$22:$C$41"),
        ]


@pytest.mark.skipif(os.environ.get("LIBRECALC_RUN_UNO") != "1", reason="requires LibreOffice UNO")
def test_uno_explicit_row_series_do_not_collapse_into_columns() -> None:
    from openpyxl import Workbook

    from librecalc_mcp.backend.uno import UnoCalcBackend

    with tempfile.TemporaryDirectory() as temporary:
        root = Path(temporary)
        source = root / "in.xlsx"
        output = root / "four-series.xlsx"
        workbook = Workbook()
        workbook.active.title = "Sheet1"
        workbook.save(source)

        backend = UnoCalcBackend()
        result = backend.execute_program(
            [
                CalcOperation(
                    op="write_range",
                    sheet="Sheet1",
                    range="A1:G5",
                    values=[
                        [None, None, "Q1", "Q2", "Q3", "Q4", None],
                        [2022, "Type A", 0, 1, 0, 1, "2022 Type A"],
                        [None, "Type B", 1, 3, 3, 3, "2022 Type B"],
                        [2023, "Type A", 2, 7, 3, 2, "2023 Type A"],
                        [None, "Type B", 6, 6, 2, 7, "2023 Type B"],
                    ],
                ),
                CalcOperation.from_dict(
                    {
                        "op": "upsert_chart",
                        "chart": {
                            "id": "sales_chart",
                            "sheet": "Sheet1",
                            "chart_type": "stacked_column",
                            "category_range": "C1:F1",
                            "series": [
                                {
                                    "name": f"{year} Type {kind}",
                                    "name_range": f"G{row}",
                                    "values_range": f"C{row}:F{row}",
                                }
                                for row, year, kind in (
                                    (2, 2022, "A"),
                                    (3, 2022, "B"),
                                    (4, 2023, "A"),
                                    (5, 2023, "B"),
                                )
                            ],
                        },
                    }
                ),
            ],
            path=str(source),
            output_path=str(output),
        )

        assert result["operations"][1]["dropped"] == []
        chart = backend.inspect_charts(str(output))[0]
        assert chart["category_range"] == "$Sheet1.$C$1:$F$1"
        assert [item["values_range"] for item in chart["series"]] == [
            "$Sheet1.$C$2:$F$2",
            "$Sheet1.$C$3:$F$3",
            "$Sheet1.$C$4:$F$4",
            "$Sheet1.$C$5:$F$5",
        ]
        assert [item["name"] for item in chart["series"]] == [
            "2022 Type A",
            "2022 Type B",
            "2023 Type A",
            "2023 Type B",
        ]
