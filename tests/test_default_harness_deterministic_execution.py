from __future__ import annotations

import hashlib
import json
import sys
from pathlib import Path

import openpyxl
import pytest
import yaml

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "benchmark"))
sys.path.insert(0, str(ROOT / "tests"))

from test_openrouter_runner import _runner_module  # noqa: E402

import calc_translate_fill as fill  # noqa: E402
import default_harness_deterministic_execution as exp  # noqa: E402

CONTROL = ROOT / "benchmark/sweagent/spreadsheet-control.yaml"
TREATMENT = ROOT / "benchmark/sweagent/spreadsheet-control-translate-fill.yaml"


def _book(tmp_path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = 1
    ws["B1"] = 2
    ws["A2"] = "=A1+1"
    ws["B2"] = "=B1+1"
    ws["C2"] = "=C1+1"
    ws["A3"] = 10
    other = wb.create_sheet("Other Sheet")
    other["A1"] = 5
    path = tmp_path / "in.xlsx"
    wb.save(path)
    wb.close()
    return path


def test_relative_absolute_mixed_and_cross_sheet_translation(tmp_path: Path) -> None:
    path = _book(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Sheet1"]
    ws["D2"] = "=$A$1+A2+$B1+A$1"
    ws["E2"] = "=$A$1+B2+$B1+B$1"
    ws["F2"] = "=$A$1+C2+$B1+C$1"
    ws["G2"] = "='Other Sheet'!A1+A2"
    ws["H2"] = "='Other Sheet'!B1+B2"
    wb.save(path)
    wb.close()
    dest = tmp_path / "out.xlsx"
    result = fill.calc_translate_fill(
        xlsx=path,
        canonical_cell="Sheet1!D2",
        canonical_formula="=$A$1+A2+$B1+A$1",
        targets="Sheet1!D2:F2",
        output=dest,
    )
    assert result["status"] == "accepted"
    assert result["translated_formulas"]["Sheet1!E2"] == "=$A$1+B2+$B1+B$1"
    assert result["translated_formulas"]["Sheet1!F2"] == "=$A$1+C2+$B1+C$1"
    cross = fill.calc_translate_fill(
        xlsx=path,
        canonical_cell="Sheet1!G2",
        canonical_formula="='Other Sheet'!A1+A2",
        targets="Sheet1!G2:H2",
        output=tmp_path / "cross.xlsx",
    )
    assert cross["status"] == "accepted"
    assert cross["translated_formulas"]["Sheet1!H2"] == "='Other Sheet'!B1+B2"


def test_never_writes_outside_declared_targets(tmp_path: Path) -> None:
    path = _book(tmp_path)
    dest = tmp_path / "out.xlsx"
    result = fill.calc_translate_fill(
        xlsx=path,
        canonical_cell="Sheet1!A2",
        canonical_formula="=A1+1",
        targets="Sheet1!A2:B2",
        output=dest,
    )
    assert result["status"] == "accepted"
    assert set(result["changed_cells"]) == {"Sheet1!A2", "Sheet1!B2"}
    wb = openpyxl.load_workbook(dest)
    assert wb["Sheet1"]["C2"].value == "=C1+1"
    wb.close()


def test_rejection_does_not_mutate_workbook(tmp_path: Path) -> None:
    path = _book(tmp_path)
    dest = tmp_path / "out.xlsx"
    dest.write_bytes(path.read_bytes())
    before = dest.read_bytes()
    payload = fill.main(
        [
            "--xlsx",
            str(dest),
            "--canonical-cell",
            "Sheet1!A2",
            "--canonical-formula",
            "=A1+1",
            "--targets",
            "Sheet1!A2,Sheet1!C2",
        ]
    )
    assert payload == 2
    assert dest.read_bytes() == before


def test_non_unit_stride_is_rejected(tmp_path: Path) -> None:
    path = _book(tmp_path)
    with pytest.raises(fill.FillReject) as caught:
        fill.calc_translate_fill(
            xlsx=path,
            canonical_cell="Sheet1!A2",
            canonical_formula="=A1+1",
            targets="Sheet1!A2,Sheet1!C2",
        )
    assert caught.value.code == "NOT_A_RECTANGLE_OR_LINE"


def test_wrong_formula_is_copied_not_repaired(tmp_path: Path) -> None:
    path = _book(tmp_path)
    dest = tmp_path / "out.xlsx"
    result = fill.calc_translate_fill(
        xlsx=path,
        canonical_cell="Sheet1!A2",
        canonical_formula="=A1+99",
        targets="Sheet1!A2:C2",
        output=dest,
    )
    assert result["status"] == "accepted"
    assert result["translated_formulas"]["Sheet1!B2"] == "=B1+99"


def test_gold_path_refused(tmp_path: Path) -> None:
    src = _book(tmp_path)
    gold = tmp_path / "golden_response.xlsx"
    gold.write_bytes(src.read_bytes())
    with pytest.raises(fill.FillReject) as caught:
        fill.calc_translate_fill(
            xlsx=gold,
            canonical_cell="Sheet1!A2",
            canonical_formula="=A1+1",
            targets="Sheet1!A2:B2",
        )
    assert caught.value.code == "GOLD_PATH_REFUSED"


def test_deterministic_twice(tmp_path: Path) -> None:
    path = _book(tmp_path)
    kwargs = dict(
        xlsx=path,
        canonical_cell="Sheet1!A2",
        canonical_formula="=A1+1",
        targets="Sheet1!A2:C2",
    )
    a = fill.calc_translate_fill(**kwargs, output=tmp_path / "a.xlsx")
    b = fill.calc_translate_fill(**kwargs, output=tmp_path / "b.xlsx")
    assert a["translated_formulas"] == b["translated_formulas"]
    assert (tmp_path / "a.xlsx").read_bytes() == (tmp_path / "b.xlsx").read_bytes()


def test_c0_c1_differ_only_by_translate_fill() -> None:
    c0 = yaml.safe_load(CONTROL.read_text(encoding="utf-8"))
    c1 = yaml.safe_load(TREATMENT.read_text(encoding="utf-8"))
    hashes = exp.prompt_hashes()
    assert hashes["c0_system_sha256"] == hashes["c1_system_sha256"]
    assert hashes["c0_instance_sha256"] != hashes["c1_instance_sha256"]
    assert c0["agent"]["templates"]["system_template"] == c1["agent"]["templates"]["system_template"]
    for key in (
        "next_step_template",
        "next_step_no_output_template",
        "next_step_truncated_observation_template",
        "command_cancelled_timeout_template",
        "max_observation_length",
    ):
        assert c0["agent"]["templates"][key] == c1["agent"]["templates"][key]
    prompt = c1["agent"]["templates"]["instance_template"]
    assert "calc_translate_fill" in prompt
    assert "never chooses targets or formulas" in prompt
    assert "never expands your requested target set" in prompt
    lowered = prompt.lower()
    assert "you must use `calc_translate_fill`" not in lowered
    assert "always call" not in lowered
    assert "semantically correct" not in lowered
    assert "task ir" not in lowered
    assert "edit plan" not in lowered
    assert [item["path"] for item in c0["agent"]["tools"]["bundles"]] == ["tools/submit", "tools/view_xlsx"]
    assert [item["path"] for item in c1["agent"]["tools"]["bundles"]] == [
        "tools/submit",
        "tools/view_xlsx",
        "../../../benchmark/sweagent/calc_translate_fill",
    ]


def test_control_translate_fill_flag_and_mounts(tmp_path, monkeypatch) -> None:
    runner = _runner_module()
    monkeypatch.setattr(
        sys,
        "argv",
        ["run_openrouter_slice.py", "--slice", "slice.json", "--run-name", "probe", "--control-translate-fill"],
    )
    arguments = runner._arguments()
    runner._apply_arm_config(arguments)
    assert arguments.config == runner.CONTROL_TRANSLATE_FILL_CONFIG
    assert runner._arm_name(arguments) == "control-translate-fill"
    staged = runner._stage_tool_policy(
        source_config=runner.CONTROL_TRANSLATE_FILL_CONFIG,
        sweagent_root=ROOT / "benchmark-data/SpreadsheetBench-2/SWE-agent",
        temporary_root=tmp_path,
        read_policy="progressive",
        execution="semantic-program-v1",
        control_translate_fill=True,
    )
    config = yaml.safe_load(staged.read_text(encoding="utf-8"))
    bundles = [Path(item["path"]).name for item in config["agent"]["tools"]["bundles"]]
    assert bundles == ["submit", "view_xlsx", "calc_translate_fill"]
    joined = " ".join(config["env"]["deployment"]["docker_args"])
    assert "/opt/librecalc/src:ro" in joined
    assert "/opt/librecalc/benchmark:ro" in joined
    assert "calc_query_cache" not in joined


def test_legacy_witness_ids_are_fully_qualified() -> None:
    ids = [f"{row['category']}:{row['id']}" for row in exp.LEGACY_WITNESSES]
    assert ids == [
        "Financial_Model:08_03",
        "Financial_Model:08_04",
        "Financial_Model:08_05",
        "Financial_Model:15_04",
    ]


def test_executor_hash_is_stable() -> None:
    assert hashlib.sha256(exp.EXECUTOR.read_bytes()).hexdigest() == exp.prompt_hashes()["executor_sha256"]
