"""Tests for the three harness fixes and the execution-admissibility gate."""
import shutil
import sqlite3
import sys
import zipfile
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "benchmark"))
from workbook_spine_sqlite import SCHEMA_SQL
from edit_plan import World, PlanError, expand_edit_plan
from xlsx_cell_writer import write_cells, SharedMasterError, patch_sheet, column_index
import edit_plan_replication as rep


def _probe():
    return rep.old


@pytest.fixture
def world(tmp_path):
    path = tmp_path / "world.sqlite"
    db = sqlite3.connect(path)
    db.executescript(SCHEMA_SQL)
    db.execute("INSERT INTO workbooks (workbook_id,readable) VALUES ('wb:t',1)")
    db.execute("INSERT INTO sheets VALUES ('sheet:s00','wb:t',0,'Test','test','visible',1,1,4,4)")
    for r, c, k in [(1, 1, 'text'), (2, 2, 'formula'), (3, 2, 'numeric'), (5, 3, 'numeric')]:
        db.execute("INSERT INTO cells VALUES (?,?,?,?,?,?,?,?,?,?)",
                   (f'cell:s00:r{r}:c{c}', 'sheet:s00', f'row:s00:r{r}', f'col:s00:c{c}', r, c, f'B{r}', k, None, None))
    db.execute("INSERT INTO rows (row_id,sheet_id,row_idx) VALUES ('row:s00:r1','sheet:s00',1)")
    db.execute("INSERT INTO columns (col_id,sheet_id,col_idx) VALUES ('col:s00:c1','sheet:s00',1)")
    db.execute("INSERT INTO formulas VALUES ('formula:s00:r2:c2','cell:s00:r2:c2','=1','fp1',0)")
    db.execute("INSERT INTO text_anchors (anchor_id,cell_id,exact_text,normalized_text,row_id,col_id) VALUES "
               "('text:s00:r1:c1','cell:s00:r1:c1','Revenue','revenue','row:s00:r1','col:s00:c1')")
    # One heading cell carrying both a row-axis and a column-axis coordinate.
    for tid, axis, idx in [('tcoord:s00:c:2', 'column', 2), ('tcoord:s00:r:1', 'row', 1), ('tcoord:s00:c:4', 'column', 4)]:
        db.execute("INSERT INTO temporal_coordinates (temporal_id,sheet_id,axis,axis_index,year,cell_id) VALUES (?,?,?,?,?,?)",
                   (tid, 'sheet:s00', axis, idx, 2023 + idx, 'cell:s00:r1:c1' if idx != 4 else 'cell:s00:r1:c4'))
    db.commit()
    db.close()
    w = World(path)
    yield w
    w.close()


def plan(expr, **extra):
    return {'operations': [{'operation_id': 'op1', 'obligation_id': 'O1', 'operation_kind': 'SET_FORMULA', 'target_set': expr, **extra}]}


RECT = {'kind': 'RECTANGLE', 'sheet_id': 'sheet:s00', 'r1': 1, 'c1': 1, 'r2': 2, 'c2': 2}


# --- FIX 1: identity contract -------------------------------------------------

@pytest.mark.parametrize('eid', ['text:s00:r1:c1', 'row:s00:r1', 'col:s00:c1', 'wb:t', 'period:s00:r1:c1'])
def test_v1_rejects_identities_the_context_exposes(world, eid):
    p = plan(RECT, source_relation={'entity_ids': [eid]})
    with pytest.raises(PlanError) as exc:
        expand_edit_plan(p, world, {'O1'}, 'V1')
    assert exc.value.category == 'INVALID_ENTITY'


@pytest.mark.parametrize('eid', ['text:s00:r1:c1', 'row:s00:r1', 'col:s00:c1', 'wb:t', 'period:s00:r1:c1', 'cell:s00:r2:c2', 'sheet:s00'])
def test_v2_accepts_every_exposed_identity(world, eid):
    p = plan(RECT, source_relation={'entity_ids': [eid]})
    assert expand_edit_plan(p, world, {'O1'}, 'V2')['status'] == 'VALID_PLAN'


def test_v2_still_rejects_identities_absent_from_the_world(world):
    for eid in ('text:s00:r9:c9', 'row:s09:r1', 'nonsense:1'):
        with pytest.raises(PlanError):
            expand_edit_plan(plan(RECT, source_relation={'entity_ids': [eid]}), world, {'O1'}, 'V2')


def test_contract_does_not_change_which_cells_expand(world):
    p = plan(RECT, source_relation={'entity_ids': ['cell:s00:r2:c2']})
    assert expand_edit_plan(p, world, {'O1'}, 'V1')['cell_ids'] == expand_edit_plan(p, world, {'O1'}, 'V2')['cell_ids']


def test_period_endpoints_resolve_on_the_declared_axis(world):
    # The same heading cell is both tcoord:s00:c:2 and tcoord:s00:r:1.
    assert world.temporal_endpoint('period:s00:r1:c1', 'column', 'V2') == 'tcoord:s00:c:2'
    assert world.temporal_endpoint('period:s00:r1:c1', 'row', 'V2') == 'tcoord:s00:r:1'
    assert world.temporal_endpoint('period:s00:r1:c1', 'column', 'V1') == 'period:s00:r1:c1'


def test_v2_expands_a_temporal_interval_named_by_period_ids(world):
    expr = {'kind': 'TEMPORAL_INTERVAL', 'sheet_id': 'sheet:s00', 'axis': 'column',
            'start_coordinate': 'period:s00:r1:c1', 'end_coordinate': 'period:s00:r1:c4',
            'row_constraint': {'r1': 2, 'r2': 2}}
    with pytest.raises(PlanError):
        expand_edit_plan(plan(expr), world, {'O1'}, 'V1')
    # Only columns carrying a temporal coordinate in range are members; c3 has none.
    assert set(expand_edit_plan(plan(expr), world, {'O1'}, 'V2')['cell_ids']) == {
        'cell:s00:r2:c2', 'cell:s00:r2:c4'}


# --- GATE: execution admissibility -------------------------------------------

def test_gate_blocks_whole_sheet_selection_without_touching_metrics(world):
    expansion = expand_edit_plan(plan({'kind': 'SHEET', 'sheet_id': 'sheet:s00'}), world, {'O1'}, 'V2')
    audit = rep.expansion_audit(expansion, world)
    assert audit['operations'][0]['max_sheet_fraction'] == 1.0
    assert not audit['plan_admissible'] and not audit['executable']
    # The raw expansion is untouched: the gate never edits the target set.
    assert len(expansion['cell_ids']) == 16


def test_gate_admits_an_ordinary_region(world):
    audit = rep.expansion_audit(expand_edit_plan(plan(RECT), world, {'O1'}, 'V2'), world)
    assert audit['plan_admissible'] and audit['executable'] and audit['operations'][0]['reasons'] == []


def test_gate_thresholds_are_predeclared():
    assert rep.ADMISSIBILITY['max_operation_cells'] == 5000
    assert rep.ADMISSIBILITY['max_plan_cells'] == 20000
    assert rep.ADMISSIBILITY['max_sheet_fraction'] == 0.5


# --- FIX 3: writer neutrality -------------------------------------------------

@pytest.fixture
def book(tmp_path):
    path = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet One"
    ws["A1"] = 1
    ws["B1"] = 2
    ws["D4"] = "=A1"
    wb.create_sheet("Second")
    wb.save(path)
    return path


def test_zero_write_is_byte_identical(book, tmp_path):
    dest = tmp_path / "out.xlsx"
    audit = write_cells(book, dest, [])
    assert audit["byte_identical_to_source"]
    assert dest.read_bytes() == book.read_bytes()


def test_single_edit_rewrites_only_its_own_sheet_part(book, tmp_path):
    dest = tmp_path / "out.xlsx"
    audit = write_cells(book, dest, [{"sheet": "Sheet One", "address": "D4", "formula": "=A1+B1"}])
    assert audit["applied"] and not audit["rejected"]
    with zipfile.ZipFile(book) as a, zipfile.ZipFile(dest) as b:
        assert a.namelist() == b.namelist()
        assert [n for n in a.namelist() if a.read(n) != b.read(n)] == audit["rewritten_parts"]
        assert len(audit["rewritten_parts"]) == 1
    wb = openpyxl.load_workbook(dest)
    assert wb["Sheet One"]["D4"].value == "=A1+B1"
    wb.close()


@pytest.mark.parametrize("address", ["C1", "A9", "Z40"])
def test_writes_to_absent_cells_and_rows(book, tmp_path, address):
    dest = tmp_path / f"out{address}.xlsx"
    write_cells(book, dest, [{"sheet": "Sheet One", "address": address, "formula": "=A1*2"}])
    wb = openpyxl.load_workbook(dest)
    assert wb["Sheet One"][address].value == "=A1*2"
    assert wb["Sheet One"]["A1"].value == 1
    wb.close()


def test_unknown_sheet_is_rejected_not_written(book, tmp_path):
    dest = tmp_path / "out.xlsx"
    audit = write_cells(book, dest, [{"sheet": "Nope", "address": "A1", "formula": "=1"}])
    assert audit["rejected"][0]["reason"] == "UNKNOWN_SHEET" and not audit["applied"]


def test_formula_text_is_xml_escaped(book, tmp_path):
    dest = tmp_path / "out.xlsx"
    write_cells(book, dest, [{"sheet": "Sheet One", "address": "C3", "formula": '=IF(A1<B1,"a&b","c")'}])
    wb = openpyxl.load_workbook(dest)
    assert wb["Sheet One"]["C3"].value == '=IF(A1<B1,"a&b","c")'
    wb.close()


def test_overwriting_a_shared_formula_master_is_refused():
    xml = '<sheetData><row r="1"><c r="A1"><f t="shared" ref="A1:C1" si="0">B1+1</f><v>2</v></c></row></sheetData>'
    with pytest.raises(SharedMasterError):
        patch_sheet(xml, {"A1": "=99"})


def test_column_index_handles_multi_letter_columns():
    assert [column_index(a) for a in ("A1", "Z9", "AA1", "CQ2")] == [1, 26, 27, 95]


def test_all_edits_refused_falls_back_to_byte_identical(book, tmp_path, monkeypatch):
    """A refused edit must not leave a re-zipped archive behind."""
    import xlsx_cell_writer as w
    monkeypatch.setattr(w, "_patch_cell", lambda *a, **k: (_ for _ in ()).throw(w.SharedMasterError("refused")))
    dest = tmp_path / "out.xlsx"
    audit = w.write_cells(book, dest, [{"sheet": "Sheet One", "address": "D4", "formula": "=A1"}])
    assert audit["rejected"] and not audit["applied"]
    assert audit["byte_identical_to_source"] and audit["rewritten_parts"] == []
    assert dest.read_bytes() == book.read_bytes()


# --- delta working-set serialization and the per-session resource bound ---

def _delta_fixture():
    target = {"cell_id": "cell:s01:r1:c1", "sheet": "S", "address": "A1", "target_id": "cell:s01:r1:c1"}
    obligation = {"id": "O1", "text": "t"}
    working = {"cell:s01:r%d:c1" % i for i in range(400)} | {"sheet:s01"}
    history = [{"q": 1, "status": "OK", "row_count": 3, "new_ids": 2, "query_kind": "cells"}]
    latest = {"status": "OK", "row_count": 1, "rows": [{"cell_id": "cell:s01:r9:c1"}]}
    return target, obligation, working, history, latest


def test_delta_serialization_drops_the_id_dump_but_keeps_exact_counts():
    old = _probe()
    target, obligation, working, history, latest = _delta_fixture()
    full = old.session_summary(target, obligation, working, history, latest)
    delta = old.session_summary(target, obligation, working, history, latest, {"cell:s01:r9:c1"}, "ws-abc12345")
    assert "cell:s01:r399:c1" in full
    assert "cell:s01:r399:c1" not in delta          # older identities are no longer re-sent
    assert "cell:s01:r9:c1" in delta                # the turn's own additions still are
    assert str(len(working)) in delta               # the exact total survives
    assert "ws-abc12345" in delta
    assert len(delta) < len(full) / 10


def test_full_serialization_is_unchanged_when_no_handle_is_given():
    old = _probe()
    target, obligation, working, history, latest = _delta_fixture()
    assert old.session_summary(target, obligation, working, history, latest) == \
        old.session_summary(target, obligation, working, history, latest, None, None)


def test_session_cap_stops_retrieval_explicitly_and_still_synthesizes(monkeypatch):
    """The bound must be visible as a record, never a silent context drop."""
    old = _probe()
    calls = {"n": 0}

    def fake_call_glm(key, system, content, tag, max_tokens):
        calls["n"] += 1
        if tag.endswith("retrieval"):
            return {"http_ok": True, "text": '{"action":"execute_sql","sql":"SELECT 1"}',
                    "usage": {"prompt_tokens": 900, "completion_tokens": 5}}
        return {"http_ok": True, "text": '{"status":"ABSTAIN"}', "usage": {"prompt_tokens": 10, "completion_tokens": 2}}

    monkeypatch.setattr(old, "call_glm", fake_call_glm)
    monkeypatch.setattr(old, "compile_bootstrap", lambda *a: {"bootstrap_entity_ids": []})
    monkeypatch.setattr(old, "sql_query", lambda ex, sql: {"status": "OK", "row_count": 0, "rows": []})
    monkeypatch.setattr(old, "materialize_working_set", lambda task, working: {})
    monkeypatch.setattr(old, "ReadOnlySqlite", lambda *a, **k: object())
    target = {"cell_id": "cell:s01:r1:c1", "sheet": "S", "address": "A1"}
    r = old.run_target("t", "raw", {"id": "O1"}, target, {}, {}, "k",
                       synthesis_system="s", retrieval_system="r", delta_state=True, session_input_cap=1000)
    limit = [c for c in r["calls"] if c.get("failure_class") == "SESSION_RESOURCE_LIMIT"]
    assert len(limit) == 1 and limit[0]["result"]["cap"] == 1000
    assert r["session_resource_limited"] is True
    assert r["working_set_serialization"] == "DELTA"
    assert r["status"] == "PROPOSAL_RETURNED"          # synthesis still happens
    assert r["synthesis"]["parsed"] == {"status": "ABSTAIN"}
    # The bound is checked against tokens already spent, so a session can
    # overshoot by at most one call. Here q1 and q2 cost 900 each; the check
    # before q3 sees 1800 >= 1000 and stops. Two retrievals, then synthesis.
    assert calls["n"] == 3
    assert limit[0]["q"] == 3 and limit[0]["result"]["retrieval_input_tokens"] == 1800


def test_uncapped_session_records_no_resource_limit(monkeypatch):
    old = _probe()
    monkeypatch.setattr(old, "call_glm", lambda *a, **k: {"http_ok": True, "text": '{"action":"final","status":"ENOUGH_EVIDENCE"}', "usage": {}})
    monkeypatch.setattr(old, "compile_bootstrap", lambda *a: {"bootstrap_entity_ids": []})
    monkeypatch.setattr(old, "materialize_working_set", lambda task, working: {})
    monkeypatch.setattr(old, "ReadOnlySqlite", lambda *a, **k: object())
    r = old.run_target("t", "raw", {"id": "O1"}, {"cell_id": "cell:s01:r1:c1", "sheet": "S", "address": "A1"},
                       {}, {}, "k", session_input_cap=1000)
    assert r["session_resource_limited"] is False
    assert r["working_set_serialization"] == "FULL"
    assert not [c for c in r["calls"] if c.get("failure_class") == "SESSION_RESOURCE_LIMIT"]


# --- post-hoc parsing and labelling of stored synthesis responses ---

import report_edit_plan_replay as replay_report  # noqa: E402


def test_frozen_extractor_drops_a_correct_answer_that_was_repeated():
    """The defect this post-hoc parse exists to correct.

    The live extractor has since been repaired, so the pre-repair behaviour is
    pinned in the report module: a historical comparison must not improve just
    because the instrument was fixed afterwards.
    """
    text = '{"status":"ABSTAIN","formula":null}\n\n\n{"status":"ABSTAIN","formula":null}\n'
    assert replay_report.frozen_extract_json_object(text) is None
    obj, how = replay_report.posthoc_parse(text)
    assert obj == {"status": "ABSTAIN", "formula": None}
    assert how == "RECOVERED_DUPLICATE_EMISSION"


def test_posthoc_parse_refuses_to_pick_between_disagreeing_objects():
    """Recovering a repeated answer is parsing; choosing between two different
    answers would be a semantic repair, so it stays unparseable."""
    obj, how = replay_report.posthoc_parse('{"status":"ABSTAIN"}\n{"status":"PROPOSED","formula":"=A1"}')
    assert obj is None and how == "CONFLICTING_OBJECTS"


def test_posthoc_parse_defers_to_the_frozen_extractor_when_it_succeeds():
    for text in ('{"status":"PROPOSED"}', '```json\n{"status":"PROPOSED"}\n```'):
        obj, how = replay_report.posthoc_parse(text)
        assert obj == {"status": "PROPOSED"} and how == "PARSED_BY_FROZEN_EXTRACTOR"


def test_posthoc_parse_reports_empty_and_objectless_text_distinctly():
    assert replay_report.posthoc_parse("") == (None, "EMPTY")
    assert replay_report.posthoc_parse("no json here") == (None, "NO_JSON_OBJECT")


def test_exhausted_output_budget_is_labelled_apart_from_unparseable():
    """max_tokens is shared between reasoning and content, so an empty response
    at exactly the budget is a truncation, not a malformed answer."""
    truncated = {"synthesis": {"parsed": None, "response": {
        "text": "", "usage": {"completion_tokens": replay_report.SYNTH_MAX_TOKENS,
                               "completion_tokens_details": {"reasoning_tokens": 2200}}}}}
    assert replay_report.rep_label(truncated) == "TRUNCATED_NO_CONTENT"
    garbage = {"synthesis": {"parsed": None, "response": {"text": "sorry", "usage": {"completion_tokens": 12}}}}
    assert replay_report.rep_label(garbage) == "UNPARSEABLE"
    good = {"synthesis": {"parsed": {"status": "PROPOSED"}, "response": {"text": "{}", "usage": {}}}}
    assert replay_report.rep_label(good) == "PROPOSED"
    assert replay_report.rep_label({}) == "NOT_RUN"


def test_leaked_retrieval_status_is_not_hidden_behind_its_status_field():
    """A leaked {"action":"final","status":"ENOUGH_EVIDENCE"} must be labelled as
    leakage, not as its retrieval status. Checking status before action would
    silently count it as a distinct synthesis outcome."""
    leaked = {"synthesis": {"parsed": {"action": "final", "status": "ENOUGH_EVIDENCE"}, "response": {"text": "{}"}}}
    assert replay_report.rep_label(leaked) == "RETRIEVAL_ACTION_FINAL"
    statusless = {"synthesis": {"parsed": {"status": "UNRESOLVED"}, "response": {"text": "{}"}}}
    assert replay_report.rep_label(statusless) == "RETRIEVAL_STATUS_UNRESOLVED"


def test_label_ordering_matches_the_replication_report_for_shared_cases():
    import report_edit_plan_replication as rep_report
    for parsed in ({"status": "PROPOSED", "formula": "=A1"}, {"status": "ABSTAIN"},
                   {"action": "final", "status": "ENOUGH_EVIDENCE"}, {"action": "execute_sql"}):
        mine = replay_report.rep_label({"synthesis": {"parsed": parsed, "response": {"text": "{}"}}})
        assert mine == rep_report.transition_label(parsed), parsed


# --- harness repairs for defects 5 and 6 (measurement, not architecture) ---

def test_shared_extractor_now_tolerates_a_repeated_valid_emission():
    """Defect 6 repaired at the source shared by every probe."""
    from task_obligation_compile import extract_json_object
    assert extract_json_object('{"status":"PROPOSED","formula":"=A1"}\n\n\n{"status":"PROPOSED","formula":"=A1"}\n') \
        == {"status": "PROPOSED", "formula": "=A1"}


def test_shared_extractor_still_refuses_to_choose_between_conflicting_answers():
    from task_obligation_compile import extract_json_object
    assert extract_json_object('{"status":"ABSTAIN"}\n{"status":"PROPOSED","formula":"=A1"}') is None


def test_shared_extractor_behaviour_is_unchanged_on_everything_else():
    from task_obligation_compile import extract_json_object
    assert extract_json_object('{"a":1}') == {"a": 1}
    assert extract_json_object('```json\n{"a":1}\n```') == {"a": 1}
    assert extract_json_object('prose {"a":1} more') == {"a": 1}
    assert extract_json_object("") is None
    assert extract_json_object("no json here") is None
    assert extract_json_object("[1,2,3]") is None


def test_synthesis_gets_its_own_larger_output_budget():
    """Defect 5: reasoning shares the output budget, and synthesis reasons longest."""
    old = _probe()
    assert old.SYNTHESIS_MAX_TOKENS > old.RETRIEVAL_MAX_TOKENS


def test_exhausted_synthesis_budget_is_recorded_by_the_harness(monkeypatch):
    """An empty body at full budget must be labelled, not left to look like a bad answer."""
    old = _probe()

    def fake_call_glm(key, system, content, tag, max_tokens):
        if tag.endswith("retrieval"):
            return {"http_ok": True, "text": '{"action":"final","status":"ENOUGH_EVIDENCE"}', "usage": {}}
        return {"http_ok": True, "text": "", "usage": {"completion_tokens": old.SYNTHESIS_MAX_TOKENS,
                                                       "completion_tokens_details": {"reasoning_tokens": old.SYNTHESIS_MAX_TOKENS}}}

    monkeypatch.setattr(old, "call_glm", fake_call_glm)
    monkeypatch.setattr(old, "compile_bootstrap", lambda *a: {"bootstrap_entity_ids": []})
    monkeypatch.setattr(old, "materialize_working_set", lambda task, working: {})
    monkeypatch.setattr(old, "ReadOnlySqlite", lambda *a, **k: object())
    r = old.run_target("t", "raw", {"id": "O1"}, {"cell_id": "cell:s01:r1:c1", "sheet": "S", "address": "A1"}, {}, {}, "k")
    assert r["synthesis_truncated"] is True
    assert r["failure_class"] == "TRUNCATED_NO_CONTENT"


def test_a_normal_synthesis_is_not_flagged_as_truncated(monkeypatch):
    old = _probe()

    def fake_call_glm(key, system, content, tag, max_tokens):
        if tag.endswith("retrieval"):
            return {"http_ok": True, "text": '{"action":"final","status":"ENOUGH_EVIDENCE"}', "usage": {}}
        return {"http_ok": True, "text": '{"status":"ABSTAIN"}', "usage": {"completion_tokens": 40}}

    monkeypatch.setattr(old, "call_glm", fake_call_glm)
    monkeypatch.setattr(old, "compile_bootstrap", lambda *a: {"bootstrap_entity_ids": []})
    monkeypatch.setattr(old, "materialize_working_set", lambda task, working: {})
    monkeypatch.setattr(old, "ReadOnlySqlite", lambda *a, **k: object())
    r = old.run_target("t", "raw", {"id": "O1"}, {"cell_id": "cell:s01:r1:c1", "sheet": "S", "address": "A1"}, {}, {}, "k")
    assert r["synthesis_truncated"] is False
    assert r["failure_class"] is None
