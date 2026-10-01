"""Stage-B helper tests: mechanical results + freshness fail-closed."""
import json

import openpyxl
import pytest


def _wb(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "FY24"
    ws["B1"] = "FY25"
    ws["A2"] = 10
    ws["B2"] = "=A2*2"
    ws["A3"] = "Revenue total"
    wb.save(path)
    wb.close()


def test_helpers_mechanical(tmp_path, monkeypatch):
    from benchmark.inspection_helpers import index as I
    from benchmark.inspection_helpers.api import inspect, periods, search
    I.reset()
    log = tmp_path / "fresh.jsonl"
    monkeypatch.setenv("AB_FRESHNESS_LOG", str(log))
    p = str(tmp_path / "b.xlsx")
    _wb(p)
    pr = periods(p)
    assert any(r["label"] == "FY24" and r["address"] == "A1" for r in pr["results"])
    sr = search(p, "revenue")
    assert any(r["address"] == "A3" for r in sr["results"])
    ir = inspect(p, "Data", "A1:B2")
    by_addr = {c["address"]: c for c in ir["results"]}
    assert by_addr["B2"]["formula"] == "=A2*2"
    assert by_addr["A2"]["value"] == "10"
    # generations logged on every call
    rows = [json.loads(l) for l in log.read_text().splitlines()]
    assert len(rows) == 3 and all("index_generation" in r and "query_generation" in r for r in rows)


def test_freshness_rebuilds_after_mutation(tmp_path, monkeypatch):
    from benchmark.inspection_helpers import index as I
    from benchmark.inspection_helpers.api import search
    I.reset()
    log = tmp_path / "fresh.jsonl"
    monkeypatch.setenv("AB_FRESHNESS_LOG", str(log))
    p = str(tmp_path / "b.xlsx")
    _wb(p)
    g1 = search(p, "revenue")["index_generation"]
    wb = openpyxl.load_workbook(p)
    wb["Data"]["A4"] = "revenue Q1"
    wb.save(p)
    wb.close()
    out = search(p, "revenue")
    assert out["index_generation"] == g1 + 1
    assert any(r["address"] == "A4" for r in out["results"])
    rows = [json.loads(l) for l in log.read_text().splitlines()]
    assert rows[-1]["rebuilt"] is True
    # same generation reused when file untouched
    out2 = search(p, "revenue")
    assert out2["index_generation"] == g1 + 1


def test_freshness_fail_closed_missing(tmp_path):
    from benchmark.inspection_helpers import index as I
    from benchmark.inspection_helpers.api import periods
    I.reset()
    with pytest.raises(FileNotFoundError):
        periods(str(tmp_path / "nope.xlsx"))


def test_inspect_reports_truncation_and_supports_paging(tmp_path):
    from benchmark.inspection_helpers import index as I
    from benchmark.inspection_helpers.api import inspect

    I.reset()
    path = tmp_path / "many.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for row in range(1, 6):
        for col in range(1, 3):
            ws.cell(row, col).value = f"{row}:{col}"
    wb.save(path)
    wb.close()

    first = inspect(str(path), "Data", "A1:B3", limit=3)
    second = inspect(str(path), "Data", "A1:B3", limit=3, offset=3)
    assert first["truncated"] is True
    assert first["next_offset"] == 3
    assert len(first["results"]) == 3
    assert second["truncated"] is False
    assert len(second["results"]) == 3
    assert first["results"][-1]["address"] != second["results"][0]["address"]


def test_inspect_ranges_is_compact_and_single_generation(tmp_path, monkeypatch):
    from benchmark.inspection_helpers import index as I
    from benchmark.inspection_helpers.api import inspect_ranges

    I.reset()
    log = tmp_path / "fresh.jsonl"
    monkeypatch.setenv("AB_FRESHNESS_LOG", str(log))
    path = tmp_path / "many.xlsx"
    _wb(path)
    result = inspect_ranges(str(path), [
        {"sheet": "Data", "range": "A1:B2"},
        ("Data", "A3:A3"),
    ])
    assert len(result["results"]) == 2
    assert result["results"][0]["fields"] == [
        "address", "row", "col", "value", "formula", "dtype"
    ]
    assert result["results"][0]["cells"][1][0] == "B1"
    assert result["results"][1]["cells"][0][3] == "Revenue total"
    rows = [json.loads(line) for line in log.read_text().splitlines()]
    assert len(rows) == 1
    assert rows[0]["helper"] == "inspect_ranges"
