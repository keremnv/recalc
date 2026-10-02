"""Maintained product behavior. Frozen research remains independently archived."""
from __future__ import annotations

import hashlib
import json
import os
import subprocess
import sys
from dataclasses import replace
from pathlib import Path

import openpyxl
import pytest

from recalc_agent import diagnostics, runner
from recalc_agent.config import Config, load
from recalc_agent.read_engine import artifact, cache

ROOT = Path(__file__).resolve().parents[1]
DIRECT = ('import openpyxl\nwb=openpyxl.load_workbook("input.xlsx")\n'
          'ws=wb["Sheet1"]\nprint(ws.cell(row=1,column=1).value)\nwb.close()\n')
REFERENCE = ('import openpyxl\nwb=openpyxl.load_workbook("input.xlsx")\n'
             'print(wb.active["A1"].value)\nwb.close()\n')
WRITE = ('import openpyxl\nwb=openpyxl.load_workbook("input.xlsx")\n'
         'wb["Sheet1"]["B2"]="=A1*2"\nwb.save("output.xlsx")\nwb.close()\n')


@pytest.fixture
def task(tmp_path):
    work = tmp_path / "work"
    work.mkdir()
    book = openpyxl.Workbook()
    book.active.title = "Sheet1"
    book.active["A1"] = 21
    book.save(work / "input.xlsx")
    script = work / "task.py"
    script.write_text(DIRECT)
    return work, script, Config(cache_dir=str(tmp_path / "cache"))


def test_frozen_extraction_sources_unchanged():
    for row in json.loads((ROOT / "research/history/product_hygiene/extraction_manifest.json").read_text()):
        assert hashlib.sha256((ROOT / row["source"]).read_bytes()).hexdigest() == row["source_sha256"]


def test_no_research_imports():
    import ast
    import recalc_agent
    for path in Path(recalc_agent.__file__).parent.rglob("*.py"):
        for node in ast.walk(ast.parse(path.read_text())):
            if isinstance(node, ast.ImportFrom):
                assert not (node.module or "").startswith(("benchmark", "librecalc_mcp")), path
            if isinstance(node, ast.Import):
                assert all(not n.name.startswith(("benchmark", "librecalc_mcp", "openai"))
                           for n in node.names), path


def test_direct_build_and_reuse(task, capfd):
    work, script, config = task
    states = []
    for _ in range(2):
        code, receipt = runner.run(script, [], work, config, [])
        assert code == 0
        assert receipt["route"] == "DIRECT_RUNTIME"
        assert receipt["direct_served_loads"] == 1
        assert receipt["assurance_status"] == "PASS"
        states.append(receipt["artifact"])
    assert states == [["BUILT"], ["REUSED"]]
    assert capfd.readouterr().out.count("21") == 2


def test_reference_branch_does_not_publish_artifact(task):
    work, script, config = task
    script.write_text(REFERENCE)
    code, receipt = runner.run(script, [], work, config, [])
    assert code == 0
    assert receipt["route"] == "REFERENCE_FAST_PATH"
    assert receipt["artifact"] == ["NOT_APPLICABLE"]
    assert not (Path(config.cache_dir) / "read-engine").exists()


def test_fallback_after_contact(task, capfd):
    work, script, config = task
    script.write_text('import openpyxl\nwb=openpyxl.load_workbook("input.xlsx")\n'
                      'ws=wb["Sheet1"]\nfor row in ws:\n print(row[0].value)\nwb.close()\n')
    code, receipt = runner.run(script, [], work, config, [])
    assert code == 0
    assert receipt["route"] == "DIRECT_WITH_FALLBACK"
    assert receipt["admitted"] is True
    assert any(f.get("reason") == "proxy_operation_escape" for f in receipt["fallback"])
    assert "21" in capfd.readouterr().out


def test_merged_terminal_child_served_directly(task, capfd):
    import json
    work, script, config = task
    book = openpyxl.load_workbook(work / "input.xlsx")
    book.active.merge_cells("A1:A2")
    book.save(work / "input.xlsx")
    script.write_text('import openpyxl\nwb=openpyxl.load_workbook("input.xlsx")\n'
                      'ws=wb["Sheet1"]\nprint(ws.cell(row=2,column=1).value)\n')
    code, receipt = runner.run(script, [], work, config, [])
    assert code == 0
    assert receipt["route"] == "DIRECT_RUNTIME"
    assert capfd.readouterr().out.strip() == "None"
    setup = json.loads((Path(receipt["run_dir"]) / "setup.json").read_text())
    assert setup["merged_certificate"]["certified"] is True


def test_certified_iter_rows_served_directly(task, capfd):
    import subprocess
    import sys
    work, script, config = task
    book = openpyxl.load_workbook(work / "input.xlsx")
    book.active["B2"] = "x"
    book.save(work / "input.xlsx")
    script.write_text('import openpyxl\nwb=openpyxl.load_workbook("input.xlsx")\n'
                      'ws=wb["Sheet1"]\nfor row in ws.iter_rows(min_row=1,max_row=2,min_col=1,max_col=2):\n'
                      ' print([(c.coordinate,c.value,c.row,c.column,c.data_type) for c in row])\nwb.close()\n')
    ref = subprocess.run([sys.executable, str(script)], cwd=work,
                         capture_output=True, text=True)
    assert ref.returncode == 0
    code, receipt = runner.run(script, [], work, config, [])
    assert code == 0
    assert receipt["route"] == "DIRECT_RUNTIME"
    assert capfd.readouterr().out == ref.stdout


def test_values_only_iter_rows_stays_reference(task):
    work, script, config = task
    script.write_text('import openpyxl\nwb=openpyxl.load_workbook("input.xlsx")\n'
                      'ws=wb["Sheet1"]\nfor row in ws.iter_rows(values_only=True):\n print(row)\nwb.close()\n')
    code, receipt = runner.run(script, [], work, config, [])
    assert code == 0
    assert receipt["route"] == "REFERENCE_FAST_PATH"


def test_malformed_workbook_matches_reference(task):
    import subprocess
    import sys
    work, script, config = task
    (work / "input.xlsx").write_bytes(b"not a zip file")
    script.write_text('import openpyxl\nopenpyxl.load_workbook("input.xlsx")\n')
    direct = subprocess.run([sys.executable, str(script)], cwd=work,
                            capture_output=True)
    code, receipt = runner.run(script, [], work, config, [])
    assert code == direct.returncode != 0
    assert receipt["target_status"]["exit_code"] == direct.returncode


def test_missing_sheet_keyerror_matches_reference(task, capfd):
    import subprocess
    import sys
    work, script, config = task
    script.write_text('import openpyxl\nwb=openpyxl.load_workbook("input.xlsx")\n'
                      'print(wb["Sheet1"].title)\nprint(wb["Nope"].title)\n')
    direct = subprocess.run([sys.executable, str(script)], cwd=work,
                            capture_output=True)
    code, receipt = runner.run(script, [], work, config, [])
    assert receipt["route"] == "DIRECT_RUNTIME"
    assert code == direct.returncode != 0
    # Existing-sheet lookup unaffected; stdout identical up to the raise.
    assert direct.stdout == b"Sheet1\n"
    assert capfd.readouterr().out == "Sheet1\n"
    # Exception type/args/message identical to pinned openpyxl.
    direct_err = direct.stderr.decode().strip().splitlines()[-1]
    assert direct_err == "KeyError: 'Worksheet Nope does not exist.'"
    from recalc_agent.read_engine.direct import decode_xlsx
    from recalc_agent.read_engine.runtime import ProxyWorkbook
    book = decode_xlsx(work / "input.xlsx")
    proxy = ProxyWorkbook.__new__(ProxyWorkbook)
    object.__setattr__(proxy, "_book", book)
    try:
        proxy["Nope"]
    except KeyError as exc:
        assert exc.args == ("Worksheet Nope does not exist.",)
        assert "Worksheet Nope does not exist." in str(exc)
    else:
        raise AssertionError("expected KeyError")


def test_version_mismatch_sidecar_rebuilds(task):
    import json
    work, script, config = task
    assert runner.run(script, [], work, config, [])[1]["artifact"] == ["BUILT"]
    digest = cache.sha_file(work / "input.xlsx")
    _, sidecar = cache.paths(Path(config.cache_dir), digest)
    meta = json.loads(sidecar.read_text())
    meta["runtime_version"] = "0.0.0-old"
    sidecar.write_text(json.dumps(meta))
    code, receipt = runner.run(script, [], work, config, [])
    assert code == 0 and receipt["artifact"] == ["BUILT"]
    assert receipt["direct_served_loads"] == 1


def test_changed_workbook_capture(task):
    work, script, config = task
    script.write_text(WRITE)
    code, receipt = runner.run(script, [], work, config, [])
    assert code == 0
    assert receipt["route"] == "REFERENCE_FAST_PATH"
    assert receipt["capture_records"] == 1
    assert receipt["effect_capture_status"] == receipt["validation_status"] == "PASS"
    assert openpyxl.load_workbook(work / "output.xlsx")["Sheet1"]["B2"].value == "=A1*2"


def test_disabled_runtime_uses_direct_python(task):
    work, script, config = task
    code, receipt = runner.run(script, [], work, replace(config, enabled=False), [])
    assert code == 0
    assert receipt["route"] == "REFERENCE_FAST_PATH"
    assert receipt["assurance_status"] == "NOT_REQUESTED"


def test_disabled_capture_is_truthful(task):
    work, script, config = task
    script.write_text(WRITE)
    code, receipt = runner.run(script, [], work, replace(config, capture=False), [])
    assert code == 0
    assert receipt["effect_capture_status"] == "NOT_REQUESTED"
    assert receipt["assurance_status"] == "NOT_REQUESTED"


def test_corrupt_cache_rebuilds(task):
    work, script, config = task
    assert runner.run(script, [], work, config, [])[1]["artifact"] == ["BUILT"]
    digest = cache.sha_file(work / "input.xlsx")
    path, _ = cache.paths(Path(config.cache_dir), digest)
    path.write_bytes(b"corrupt")
    code, receipt = runner.run(script, [], work, config, [])
    assert code == 0 and receipt["artifact"] == ["BUILT"]
    assert receipt["direct_served_loads"] == 1


def test_source_change_invalidates(task):
    work, script, config = task
    assert runner.run(script, [], work, config, [])[1]["artifact"] == ["BUILT"]
    book = openpyxl.load_workbook(work / "input.xlsx")
    book.active["A1"] = 22
    book.save(work / "input.xlsx")
    code, receipt = runner.run(script, [], work, config, [])
    assert code == 0 and receipt["artifact"] == ["BUILT"]


def test_version_change_changes_artifact_key(monkeypatch):
    from recalc_agent.read_engine import _identity
    digest = "a" * 64
    original = cache.artifact_key(digest)
    for attr, value in [("CONTRACT", "NEXT_CONTRACT"), ("DECODER", "b" * 64),
                        ("FORMAT", "NEXT_FORMAT"),
                        ("RUNTIME_VERSION", "next-runtime")]:
        monkeypatch.setattr(_identity, attr, value)
        assert cache.artifact_key(digest) != original
        assert artifact.artifact_key(digest) != original
        monkeypatch.undo()
    assert cache.artifact_key(digest) == original


def test_artifact_key_stable_and_single_sourced():
    from recalc_agent.read_engine import _identity
    digest = "a" * 64
    # Fixed vector locks the key format across the consolidation; any drift
    # between layers or across refactors fails here. Update only with an
    # intentional, documented version bump (which orphans old artifacts).
    # Vector rotated for the intentional 0.2.0rc5 bump (CHANGELOG):
    # the key mixes in RUNTIME_VERSION, so rc4 artifacts rebuild once.
    expected = "aeaa250e4a03582890663c66c1b2988a3f5746ae72e7f12bdb9d46b8cdf94375"
    assert _identity.artifact_key(digest) == expected
    assert cache.artifact_key is _identity.artifact_key
    assert artifact.artifact_key is _identity.artifact_key
    assert cache.identity is _identity.identity
    assert artifact.identity is _identity.identity
    assert cache.sha_file is _identity.sha_file
    assert artifact.sha_file is _identity.sha_file
    assert cache.paths is _identity.paths
    assert artifact.paths is _identity.paths


def test_artifact_round_trip_and_corruption(task):
    work, _, _ = task
    from recalc_agent.read_engine.direct import decode_xlsx
    book = decode_xlsx(work / "input.xlsx")
    digest = artifact.sha_file(work / "input.xlsx")
    raw = artifact.encode(book, digest)
    assert artifact.decode(raw, digest)["Sheet1"].cell(1, 1).value == 21
    with pytest.raises(artifact.ArtifactError):
        artifact.decode(raw[:-1], digest)


def test_script_exit_and_arguments(task):
    work, script, config = task
    script.write_text('import sys\nfrom pathlib import Path\nPath("count").write_text("x")\n'
                      'assert sys.argv[1:]==["hello world"]\nraise SystemExit(7)\n')
    code, receipt = runner.run(script, ["hello world"], work, config, [])
    assert code == 7 and (work / "count").read_text() == "x"
    assert receipt["target_status"]["exit_code"] == 7


def test_reads_key_and_deprecated_aliases(task):
    work, _, config = task
    path = work / "reads.toml"
    path.write_text(f'[runtime]\nreads=false\ncache_dir="{config.cache_dir}"\n')
    loaded, issues = load(str(path))
    assert not issues and not loaded.reads_effective
    assert loaded.assurance  # capture independent of reads
    old = work / "old.toml"
    old.write_text(f'[runtime]\nsubstrate=false\ncache_dir="{config.cache_dir}"\n')
    loaded, issues = load(str(old))
    assert not loaded.reads_effective
    assert any(i["status"] == "WARNING" and "deprecated" in i["message"] for i in issues)
    both = work / "both.toml"
    both.write_text(f'[runtime]\nreads=true\ncandidate_a=false\ncache_dir="{config.cache_dir}"\n')
    loaded, issues = load(str(both))
    assert not loaded.reads_effective  # aliases AND with reads during the window
    assert any("candidate_a" in i["message"] for i in issues)


def test_bad_config_disables_runtime(task):
    work, script, config = task
    path = work / "bad.toml"
    path.write_text("[runtime]\nunknown=true\n")
    loaded, issues = load(str(path))
    assert not loaded.enabled and issues
    code, receipt = runner.run(script, [], work, replace(loaded, cache_dir=config.cache_dir), issues)
    assert code == 0 and receipt["assurance_status"] == "NOT_REQUESTED"


def test_missing_observer_fails_before_script(task, monkeypatch):
    work, script, config = task
    monkeypatch.setattr(runner, "observer_binary", lambda: work / "missing-observer")
    with pytest.raises(runner.ProductError, match="no task was run"):
        runner.run(script, [], work, config, [])


def test_receipt_admission_shape(task):
    work, script, config = task
    _, direct = runner.run(script, [], work, config, [])
    assert direct["admitted"] is True
    assert direct["admission_reason"] == "admitted"
    assert "read_gate" not in direct
    script.write_text(REFERENCE)
    _, ref = runner.run(script, [], work, config, [])
    assert ref["admitted"] is False
    assert ref["admission_reason"] == "not-admitted"
    assert "read_gate" not in ref
    _, disabled = runner.run(script, [], work, replace(config, enabled=False), [])
    assert disabled["admitted"] is False
    assert disabled["admission_reason"] == "runtime-disabled"


def test_status_records_receipt(task):
    work, script, config = task
    runner.run(script, [], work, config, [])
    report = diagnostics.status(config, [])
    assert report["last_run"]["route"] == "DIRECT_RUNTIME"
    assert report["last_run"]["direct_served_loads"] == 1


def test_nested_interpreter_does_not_inherit_interposition(task, capfd):
    work, script, config = task
    script.write_text('import subprocess,sys\nsubprocess.run([sys.executable,"-c",'
                      '"import openpyxl; print(openpyxl.load_workbook.__module__)"],check=True)\n')
    code, _ = runner.run(script, [], work, config, [])
    assert code == 0 and "openpyxl.reader.excel" in capfd.readouterr().out


def test_concurrent_publication(task):
    work, script, config = task
    env = {**os.environ, "PYTHONPATH": str(ROOT / "src")}
    command = [sys.executable, "-m", "recalc_agent.cli", "run", "--workdir", str(work), str(script)]
    cfg = work / "runtime.toml"
    cfg.write_text(f'[runtime]\ncache_dir="{config.cache_dir}"\n')
    command[4:4] = ["--config", str(cfg)]
    procs = [subprocess.Popen(command, cwd=work, env=env, stdout=subprocess.PIPE,
                              stderr=subprocess.PIPE) for _ in range(2)]
    results = [proc.communicate(timeout=20) for proc in procs]
    assert all(proc.returncode == 0 and out.strip() == b"21" and not err
               for proc, (out, err) in zip(procs, results))
    digest = cache.sha_file(work / "input.xlsx")
    path, sidecar = cache.paths(Path(config.cache_dir), digest)
    assert cache.validate(path, sidecar, digest)[0]


def test_capture_rejects_entity_expansion_xml():
    import io
    import zipfile
    from recalc_agent._frozen.delta import derive_delta
    from recalc_agent._frozen.validate import _safe_fromstring, validate_mechanical
    bomb = (b'<?xml version="1.0"?><!DOCTYPE r [<!ENTITY x "y">'
            b'<!ENTITY x2 "&x;&x;&x;&x;&x;&x;&x;&x;">'
            b'<!ENTITY xxe SYSTEM "file:///etc/passwd">]><r>&x2;</r>')
    with pytest.raises(ValueError, match="DOCTYPE"):
        _safe_fromstring(bomb, "[Content_Types].xml")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as archive:
        archive.writestr("[Content_Types].xml", bomb)
    post = buf.getvalue()
    report = validate_mechanical(None, post, post, derive_delta(None, post))
    assert report.passed is False
    assert report.checks["serialization_valid"] is False
    assert "DOCTYPE" in report.detail["serialization_valid"]


def test_cache_summary_reports_use(task):
    work, script, config = task
    runner.run(script, [], work, config, [])
    summary = diagnostics.cache_summary(Path(config.cache_dir))
    assert summary["exists"] is True
    assert summary["artifact_count"] == 1
    assert summary["run_count"] == 1
    assert summary["total_bytes"] > 0
    assert summary["truncated"] is False
    assert "safe to delete" in summary["guidance"]
    report = diagnostics.check(config, [])
    assert report["cache_summary"]["artifact_count"] == 1
    missing = diagnostics.cache_summary(Path(config.cache_dir) / "nope")
    assert missing["exists"] is False and missing["file_count"] == 0


def test_cli_example_and_doctor(tmp_path):
    from recalc_agent.cli import main
    assert main(["example", str(tmp_path / "example")]) == 0
    assert (tmp_path / "example/update.py").is_file()
    assert main(["example", str(tmp_path / "example")]) == 2
    assert main(["doctor", "--json"]) == 0


def _direct(typed):
    from recalc_agent.read_engine.artifact import _check_typed, _direct_value
    _check_typed(typed)
    return _direct_value(typed)


def test_direct_value_scalars():
    assert _direct({"kind": "scalar", "value": None}) is None
    assert _direct({"kind": "scalar", "value": True}) is True
    assert _direct({"kind": "scalar", "value": 7}) == 7
    assert _direct({"kind": "scalar", "value": 2.5}) == 2.5
    assert _direct({"kind": "scalar", "value": "x"}) == "x"
    assert _direct({"kind": "scalar", "value": "=A1*2"}) == "=A1*2"
    assert _direct({"kind": "scalar", "value": "#DIV/0!"}) == "#DIV/0!"


def test_direct_value_temporals():
    import datetime as dt
    assert _direct({"kind": "datetime", "value": "2024-05-01T12:30:00"}) == dt.datetime(2024, 5, 1, 12, 30)
    assert _direct({"kind": "date", "value": "2024-05-01"}) == dt.date(2024, 5, 1)
    assert _direct({"kind": "time", "value": "12:30:00"}) == dt.time(12, 30)
    assert _direct({"kind": "timedelta", "seconds": 90}) == dt.timedelta(seconds=90)
    assert _direct({"kind": "timedelta", "seconds": 1.5}) == dt.timedelta(seconds=1.5)


def test_direct_value_formulas():
    from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula
    a = _direct({"kind": "array", "ref": "A1:A2", "text": "SUM(A1:A2)"})
    assert isinstance(a, ArrayFormula) and (a.ref, a.text) == ("A1:A2", "SUM(A1:A2)")
    d = _direct({"kind": "datatable", "attrs": {"ref": "A1:B2"}})
    assert isinstance(d, DataTableFormula)


def test_direct_value_matches_legacy_reconstruction():
    import datetime as dt
    import json
    from recalc_agent.read_engine.direct import _value_from_json, _value_to_json
    from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula
    values = [None, True, 0, -3, 2.5, "", "s", "=F(1)",
              dt.datetime(2024, 5, 1, 12, 30), dt.date(2024, 5, 1),
              dt.time(12, 30), dt.timedelta(seconds=90)]
    for v in values:
        raw = _value_to_json(v)
        assert _direct(json.loads(raw)) == _value_from_json(raw)
    for v in (ArrayFormula(ref="A1:A2", text="SUM(A1:A2)"),
              DataTableFormula(ref="A1:B2")):
        raw = _value_to_json(v)
        new, old = _direct(json.loads(raw)), _value_from_json(raw)
        assert type(new) is type(old) and vars(new) == vars(old)


def test_direct_value_rejects_non_finite():
    import math
    from recalc_agent.read_engine.artifact import ArtifactError
    for bad in (math.nan, math.inf, -math.inf):
        with pytest.raises(ArtifactError):
            _direct({"kind": "scalar", "value": bad})
        with pytest.raises(ArtifactError):
            _direct({"kind": "timedelta", "seconds": bad})


def test_direct_value_rejects_unknown_kind():
    from recalc_agent.read_engine.artifact import ArtifactError, _direct_value
    with pytest.raises(ArtifactError):
        _direct_value({"kind": "frobnicate", "value": 1})
