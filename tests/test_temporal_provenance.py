"""Unit tests for frozen temporal provenance walking."""
from __future__ import annotations

import datetime as dt
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "benchmark"))
sys.path.insert(0, str(ROOT / "benchmark/sweagent/formula_index/lib"))
sys.path.insert(0, str(ROOT / "src"))

from temporal_provenance import (  # noqa: E402
    classify_coordinate,
    classify_formula,
    classify_seed,
    compact_path,
    header_candidates,
    index_coordinates,
    project_walk,
    walk_from,
)
from temporal_spine import _merge_map, snapshot_cell  # noqa: E402


def _save(tmp_path: Path, build) -> tuple[Path, openpyxl.Workbook]:
    workbook = openpyxl.Workbook()
    build(workbook)
    path = tmp_path / "in.xlsx"
    workbook.save(path)
    workbook.close()
    loaded = openpyxl.load_workbook(path, data_only=False)
    return path, loaded


def test_classify_direct_and_transforms() -> None:
    assert classify_formula("='Income Statement'!C5")["kind"] == "E1"
    assert classify_formula("=+C4")["kind"] == "E1"
    assert classify_formula("=EOMONTH(O5,1)")["kind"] == "E2"
    assert classify_formula("=EDATE('Assumptions'!M33,-1)")["transform"] == "EDATE"
    assert classify_formula("=C5+7")["kind"] == "E3"
    assert classify_formula("=C5-1")["offset"] == -1
    assert classify_formula("=DATE(2023,8,1)")["kind"] == "LITERAL_DATE"


def test_classify_opaque_unsupported() -> None:
    assert classify_formula("=OFFSET(C5,0,1)")["kind"] == "OPAQUE"
    assert classify_formula("=INDIRECT(\"C5\")")["kind"] == "OPAQUE"
    assert classify_formula("=IF(A1>0,B1,C1)")["kind"] == "OPAQUE"
    assert classify_formula("=VLOOKUP(A1,B:C,2,FALSE)")["kind"] == "OPAQUE"
    assert classify_formula("=INDEX(A1:A10,MATCH(1,B1:B10,0))")["kind"] == "OPAQUE"
    assert classify_formula("=StartDate")["kind"] == "OPAQUE"
    assert classify_formula("=A1+B1")["kind"] == "OPAQUE"
    assert classify_formula("=TEXT(A1,\"mmm-yy\")")["kind"] == "OPAQUE"


def test_integer_twelve_is_not_a_year_seed(tmp_path: Path) -> None:
    def build(wb):
        sheet = wb.active
        sheet.title = "Balance Sheet"
        sheet["C4"] = 12
        sheet["C5"] = "='Income Statement'!C5"

    _path, wb = _save(tmp_path, build)
    try:
        sheet = wb["Balance Sheet"]
        snap = snapshot_cell(sheet["C4"], merge_origin={})
        assert classify_seed(snap, sheet_title="Balance Sheet") is None
    finally:
        wb.close()


def test_cross_sheet_direct_reaches_datetime(tmp_path: Path) -> None:
    def build(wb):
        income = wb.active
        income.title = "Income Statement"
        income["C5"] = dt.datetime(2023, 12, 31)
        income["C5"].number_format = "mmm-yy"
        balance = wb.create_sheet("Balance Sheet")
        balance["C5"] = "='Income Statement'!C5"
        balance["C5"].number_format = "mmm-yy"

    _path, wb = _save(tmp_path, build)
    try:
        result = walk_from(
            wb,
            start_sheet=wb["Balance Sheet"],
            start_col=3,
            start_row=5,
            coord_index={},
            merge_maps={},
        )
        assert result["status"] == "REACHABLE_TEMPORAL_SEED"
        best = result["paths"][0]
        assert best["path_type"] == "DIRECT_CROSS_SHEET"
        assert best["seed"]["seed_type"] == "DATETIME_SEED"
        assert best["depth"] == 1
    finally:
        wb.close()


def test_existing_coordinate_seed_preferred(tmp_path: Path) -> None:
    def build(wb):
        income = wb.active
        income.title = "Income Statement"
        income["C5"] = dt.datetime(2023, 12, 31)
        balance = wb.create_sheet("Balance Sheet")
        balance["C5"] = "='Income Statement'!C5"

    _path, wb = _save(tmp_path, build)
    try:
        coords = [
            {
                "id": "tcoord:s00:c:3",
                "sheet_title": "Income Statement",
                "col": 3,
                "row": 5,
                "period": {"year": 2023, "month": 12},
                "evidence_cells": ["C5"],
            }
        ]
        result = walk_from(
            wb,
            start_sheet=wb["Balance Sheet"],
            start_col=3,
            start_row=5,
            coord_index=index_coordinates(coords),
            merge_maps={},
        )
        assert result["paths"][0]["seed"]["seed_type"] == "EXISTING_COORDINATE"
    finally:
        wb.close()


def test_eomonth_chain_same_sheet(tmp_path: Path) -> None:
    def build(wb):
        sheet = wb.active
        sheet.title = "IS"
        sheet["C5"] = dt.datetime(2023, 8, 31)
        sheet["C5"].number_format = "mmm-yy"
        sheet["D5"] = "=EOMONTH(C5,1)"
        sheet["D5"].number_format = "mmm-yy"
        sheet["E5"] = "=EOMONTH(D5,1)"
        sheet["E5"].number_format = "mmm-yy"

    _path, wb = _save(tmp_path, build)
    try:
        result = walk_from(
            wb,
            start_sheet=wb["IS"],
            start_col=5,
            start_row=5,
            coord_index={},
            merge_maps={},
        )
        assert result["status"] == "REACHABLE_TEMPORAL_SEED"
        best = result["paths"][0]
        assert best["path_type"] == "SAME_SHEET_TRANSFORM"
        assert best["depth"] == 2
        assert best["seed"]["seed_type"] == "DATETIME_SEED"
    finally:
        wb.close()


def test_depth_limit_does_not_continue(tmp_path: Path) -> None:
    def build(wb):
        sheet = wb.active
        sheet["A1"] = dt.datetime(2020, 1, 31)
        prev = "A"
        for col in range(2, 10):
            letter = chr(ord("A") + col - 1)
            sheet[f"{letter}1"] = f"=EOMONTH({prev}1,1)"
            prev = letter

    _path, wb = _save(tmp_path, build)
    try:
        result = walk_from(
            wb,
            start_sheet=wb.active,
            start_col=8,
            start_row=1,
            coord_index={},
            merge_maps={},
            max_depth=6,
        )
        assert result["status"] == "DEPTH_LIMIT"
        assert result["paths"] == []
        reachable = walk_from(
            wb,
            start_sheet=wb.active,
            start_col=7,
            start_row=1,
            coord_index={},
            merge_maps={},
            max_depth=6,
        )
        assert reachable["status"] == "REACHABLE_TEMPORAL_SEED"
        assert reachable["paths"][0]["depth"] == 6
    finally:
        wb.close()


def test_cycle_detected(tmp_path: Path) -> None:
    def build(wb):
        sheet = wb.active
        sheet["A1"] = "=B1"
        sheet["B1"] = "=A1"
        sheet["A1"].number_format = "mmm-yy"
        sheet["B1"].number_format = "mmm-yy"

    _path, wb = _save(tmp_path, build)
    try:
        result = walk_from(
            wb,
            start_sheet=wb.active,
            start_col=1,
            start_row=1,
            coord_index={},
            merge_maps={},
        )
        assert result["status"] == "CYCLE"
    finally:
        wb.close()


def test_offset_header_is_opaque(tmp_path: Path) -> None:
    def build(wb):
        sheet = wb.active
        sheet["B1"] = dt.datetime(2024, 1, 31)
        sheet["C1"] = "=OFFSET(B1,0,0)"
        sheet["C1"].number_format = "mmm-yy"

    _path, wb = _save(tmp_path, build)
    try:
        classified = classify_coordinate(
            wb,
            sheet=wb.active,
            axis="column",
            col=3,
            row=None,
            coord_index={},
            merge_maps={},
        )
        assert classified["primary"] == "OPAQUE_PATH"
    finally:
        wb.close()


def test_no_header_candidate(tmp_path: Path) -> None:
    def build(wb):
        sheet = wb.active
        sheet["A50"] = "Revenue"
        sheet["C50"] = "=SUM(C10:C40)"

    _path, wb = _save(tmp_path, build)
    try:
        classified = classify_coordinate(
            wb,
            sheet=wb.active,
            axis="column",
            col=3,
            row=None,
            coord_index={},
            merge_maps={},
        )
        assert classified["primary"] == "NO_HEADER_CANDIDATE"
        assert header_candidates(wb.active, axis="column", col=3, merge_origin={}) == []
    finally:
        wb.close()


def test_cross_sheet_then_transform(tmp_path: Path) -> None:
    def build(wb):
        assumptions = wb.active
        assumptions.title = "Assumptions"
        assumptions["M33"] = dt.datetime(2023, 8, 31)
        assumptions["M33"].number_format = "mmm-yy"
        wc = wb.create_sheet("Working Capital")
        wc["J4"] = "='Assumptions'!M33"
        wc["J4"].number_format = "mmm-yy"
        wc["K4"] = "=EOMONTH(J4,1)"
        wc["K4"].number_format = "mmm-yy"

    _path, wb = _save(tmp_path, build)
    try:
        result = walk_from(
            wb,
            start_sheet=wb["Working Capital"],
            start_col=11,
            start_row=4,
            coord_index={},
            merge_maps={},
        )
        assert result["status"] == "REACHABLE_TEMPORAL_SEED"
        assert result["paths"][0]["path_type"] == "TRANSFORM_THEN_CROSS_SHEET"
    finally:
        wb.close()


def test_deeper_cap_recovers_eomonth_chain(tmp_path: Path) -> None:
    def build(wb):
        sheet = wb.active
        sheet["A1"] = dt.datetime(2020, 1, 31)
        prev = "A"
        for col in range(2, 12):
            letter = chr(ord("A") + col - 1)
            sheet[f"{letter}1"] = f"=EOMONTH({prev}1,1)"
            prev = letter

    _path, wb = _save(tmp_path, build)
    try:
        shallow = walk_from(
            wb,
            start_sheet=wb.active,
            start_col=11,
            start_row=1,
            coord_index={},
            merge_maps={},
            max_depth=6,
        )
        deep = walk_from(
            wb,
            start_sheet=wb.active,
            start_col=11,
            start_row=1,
            coord_index={},
            merge_maps={},
            max_depth=12,
        )
        closure = walk_from(
            wb,
            start_sheet=wb.active,
            start_col=11,
            start_row=1,
            coord_index={},
            merge_maps={},
            max_depth=10**9,
            max_visited=256,
        )
        assert shallow["status"] == "DEPTH_LIMIT"
        assert deep["status"] == "REACHABLE_TEMPORAL_SEED"
        assert deep["paths"][0]["depth"] == 10
        assert closure["status"] == "REACHABLE_TEMPORAL_SEED"
        projected = project_walk(closure, max_depth=6)
        assert projected["status"] == "DEPTH_LIMIT"
        projected12 = project_walk(closure, max_depth=12)
        assert projected12["status"] == "REACHABLE_TEMPORAL_SEED"
    finally:
        wb.close()


def test_visited_cap_is_closure_limit_not_silent(tmp_path: Path) -> None:
    def build(wb):
        sheet = wb.active
        sheet["A1"] = dt.datetime(2020, 1, 31)
        prev = "A"
        for col in range(2, 20):
            letter = chr(ord("A") + col - 1)
            sheet[f"{letter}1"] = f"=EOMONTH({prev}1,1)"
            prev = letter

    _path, wb = _save(tmp_path, build)
    try:
        result = walk_from(
            wb,
            start_sheet=wb.active,
            start_col=19,
            start_row=1,
            coord_index={},
            merge_maps={},
            max_depth=10**9,
            max_visited=4,
        )
        assert result["status"] == "SAFE_CLOSURE_LIMIT"
        assert result["paths"] == []
    finally:
        wb.close()


def test_compact_path_collapses_eomonth_run() -> None:
    path = {
        "start": "IS!P5",
        "end": "IS!I5",
        "seed": {"seed_type": "DATETIME_SEED", "period": {"year": 2023, "month": 8}},
        "steps": [
            {
                "kind": "E2",
                "transform": "EOMONTH",
                "from": f"IS!{chr(ord('P') - i)}5",
                "to": f"IS!{chr(ord('O') - i)}5",
                "cross_sheet": False,
            }
            for i in range(6)
        ]
        + [
            {
                "kind": "E1",
                "transform": None,
                "from": "IS!I5",
                "to": "Assumptions!F33",
                "cross_sheet": True,
            }
        ],
    }
    lines = compact_path(path)
    assert any("6 EOMONTH hops" in line for line in lines)
    assert any("cross-sheet" in line for line in lines)
    assert lines[-1].startswith("seed DATETIME_SEED")


def test_derive_eomonth_shifts_month() -> None:
    from temporal_provenance import derive_period_from_path

    seed = {"period": {"year": 2023, "month": 1}}
    steps = [
        {"kind": "E2", "offset": 1},
        {"kind": "E2", "offset": 1},
        {"kind": "E1", "offset": None},
    ]
    derived = derive_period_from_path(seed, steps)
    assert derived["year"] == 2023
    assert derived["month"] == 3


def test_closure_materializes_cross_sheet_eomonth(tmp_path: Path) -> None:
    from temporal_spine import compile_temporal_workbook

    def build(wb):
        assumptions = wb.active
        assumptions.title = "Assumptions"
        assumptions["L6"] = dt.datetime(2023, 1, 31)
        assumptions["L6"].number_format = "mmm-yy"
        income = wb.create_sheet("Income Statement")
        income["I5"] = "='Assumptions'!L6"
        income["I5"].number_format = "mmm-yy"
        income["J5"] = "=EOMONTH(I5,1)"
        income["J5"].number_format = "mmm-yy"
        income["P5"] = "=EOMONTH(O5,1)"
        income["O5"] = "=EOMONTH(N5,1)"
        income["N5"] = "=EOMONTH(M5,1)"
        income["M5"] = "=EOMONTH(L5,1)"
        income["L5"] = "=EOMONTH(K5,1)"
        income["K5"] = "=EOMONTH(J5,1)"
        for addr in ("O5", "N5", "M5", "L5", "K5", "P5"):
            income[addr].number_format = "mmm-yy"

    path, _wb = _save(tmp_path, build)
    _wb.close()
    local = compile_temporal_workbook(path, features=set(["FORMULA_DERIVED_DATE", "ROW_AXIS_SUPPORT", "FY_TOKEN_NORMALIZATION", "MERGED_HEADER_PROPAGATION", "DATE_SERIAL_DECODING"]), closure=False)
    closed = compile_temporal_workbook(path, features=set(local["features"]), closure=True)
    assert closed["golden_used"] is False
    assert closed["n_propagated_coordinates"] >= 1
    cols = {
        c["col"]: c["period"]
        for c in closed["coordinates"]
        if c["sheet_title"] == "Income Statement" and c["axis"] == "column"
    }
    assert cols.get(16, {}).get("year") == 2023
    assert cols.get(16, {}).get("month") == 8
    p_coord = next(c for c in closed["coordinates"] if c["sheet_title"] == "Income Statement" and c["col"] == 16 and c["axis"] == "column")
    assert p_coord["derivation_status"] == "PROPAGATED_TEMPORAL_COORDINATE"
    assert p_coord["path_length"] >= 7
    assert p_coord["provenance"][0]["compact"]


def test_closure_conflict_is_not_exposed(tmp_path: Path) -> None:
    from temporal_spine import compile_temporal_workbook

    def build(wb):
        sheet = wb.active
        sheet["A1"] = dt.datetime(2020, 1, 31)
        sheet["B1"] = dt.datetime(2021, 6, 30)
        sheet["C1"] = "=A1"
        sheet["C2"] = "=B1"
        sheet["C1"].number_format = "mmm-yy"
        sheet["C2"].number_format = "mmm-yy"

    path, _wb = _save(tmp_path, build)
    _wb.close()
    closed = compile_temporal_workbook(path, features=set(["FORMULA_DERIVED_DATE", "ROW_AXIS_SUPPORT", "DATE_SERIAL_DECODING"]), closure=True)
    col_c = [c for c in closed["coordinates"] if c["axis"] == "column" and c["col"] == 3 and c.get("derivation_status") == "PROPAGATED_TEMPORAL_COORDINATE"]
    assert col_c == []
    assert closed["temporal_conflicts"]
