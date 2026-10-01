"""Corrective regression suite: existing Candidate-A surface only, no model runs."""
import io
import json
import sqlite3
import zipfile
from pathlib import Path

import openpyxl
import pytest

from benchmark.candidate_a_shadow_interposition import (
    CandidateALoader,
    ProxyWorkbook,
    create_fixture,
)
from benchmark.inspection_helpers import index, substrate

REAL = openpyxl.load_workbook


@pytest.fixture(autouse=True)
def clean_index(monkeypatch):
    index.reset()
    for key in ("CANDIDATE_A_ARM", "CANDIDATE_A_FORCE_REAL", "CANDIDATE_A_SUBSTRATE_DISABLED"):
        monkeypatch.delenv(key, raising=False)
    yield
    index.reset()


@pytest.fixture
def workbook(tmp_path):
    p = tmp_path / "input.xlsx"
    create_fixture(p)
    return p


def runtime(monkeypatch, tmp_path, manifest):
    from benchmark import candidate_a_live_runtime as live
    p = tmp_path / "manifest.json"
    p.write_text(json.dumps(manifest))
    events = []
    monkeypatch.setenv("CANDIDATE_A_SUBSTRATE_MANIFEST", str(p))
    monkeypatch.setattr(live, "emit", events.append)
    monkeypatch.setattr(openpyxl, "load_workbook", REAL)
    live.install()
    return events


@pytest.mark.parametrize("content", [b"untrusted bytes", b"PK\x03\x04truncated"])
def test_malformed_bytes_fall_back_with_reference_exception(tmp_path, monkeypatch, content):
    p = tmp_path / "input.xlsx"
    p.write_bytes(content)
    manifest = substrate.prepare(tmp_path, tmp_path / "shared")
    assert not manifest["workbooks"]
    assert manifest["failures"][0]["exception_class"] == "BadZipFile"
    events = runtime(monkeypatch, tmp_path, manifest)
    with pytest.raises(zipfile.BadZipFile) as ref:
        REAL(p)
    with pytest.raises(type(ref.value)) as actual:
        openpyxl.load_workbook(p)
    assert str(actual.value) == str(ref.value)
    assert not any(e.get("status") == "ACCELERATED" for e in events)
    assert any(e.get("status") == "PREDECLARED_FALLBACK" for e in events)


def test_malformed_package_xml(tmp_path):
    p = tmp_path / "input.xlsx"
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("[Content_Types].xml", "<broken")
    result = substrate.prepare(tmp_path, tmp_path / "shared")
    assert result["failures"][0]["stage"] == "workbook_index_parse"
    assert not result["workbooks"]


@pytest.mark.parametrize("stage", ["parser", "database"])
def test_initial_failure_is_sticky_and_reference_succeeds(workbook, monkeypatch, stage):
    if stage == "parser":
        def fail(*args, **kwargs):
            raise UnicodeDecodeError("utf-8", b"\x8d", 0, 1, "injected parser defect")
        monkeypatch.setattr(openpyxl, "load_workbook", fail)
    else:
        monkeypatch.setattr(index, "_open_db", lambda: (_ for _ in ()).throw(sqlite3.OperationalError("injected")))
    result = substrate.prepare(workbook.parent, workbook.parent / "shared")
    assert not result["workbooks"]
    assert result["failures"][0]["exception_class"] == ("UnicodeDecodeError" if stage == "parser" else "OperationalError")
    events = []
    loader = CandidateALoader(real_loader=REAL, event_sink=events.append)
    for _ in range(2):
        wb = loader.load_workbook(workbook)
        assert not isinstance(wb, ProxyWorkbook)
        assert wb["Data"]["C1"].value == 7
        wb.close()
    assert not any(e.get("status") == "ACCELERATED" for e in events)


def test_partial_build_cleanup(workbook, monkeypatch):
    con = sqlite3.connect(":memory:")
    con.execute("CREATE TABLE cells(sheet,addr,row,col,value,formula,dtype)")
    monkeypatch.setattr(index, "_open_db", lambda: con)
    with pytest.raises(index.SubstrateDisabled) as exc:
        index.ensure_fresh(workbook)
    assert exc.value.event["partially_initialized"] is True
    assert str(workbook) not in index._state
    with pytest.raises(sqlite3.ProgrammingError):
        con.execute("SELECT 1")


def test_failed_refresh_retires_old_handle_and_new_generation_rebuilds(workbook, monkeypatch):
    old, _ = index.ensure_fresh(workbook)
    wb = REAL(workbook); wb["Data"]["C1"] = 8; wb.save(workbook); wb.close()
    original = index._open_db
    monkeypatch.setattr(index, "_open_db", lambda: (_ for _ in ()).throw(RuntimeError("build failed")))
    with pytest.raises(index.SubstrateDisabled):
        index.ensure_fresh(workbook)
    assert old["valid"] is False
    monkeypatch.setattr(index, "_open_db", original)
    with pytest.raises(index.SubstrateDisabled):
        index.ensure_fresh(workbook)  # same bytes cannot revive
    wb = REAL(workbook); wb["Data"]["C1"] = 9; wb.save(workbook); wb.close()
    new, rebuilt = index.ensure_fresh(workbook)
    assert rebuilt and new["valid"] and new["index_generation"] == 2
    assert new["db"].execute("SELECT value FROM cells WHERE sheet='Data' AND addr='C1'").fetchone() == ("9",)


@pytest.mark.parametrize("failure", ["missing", "corrupt", "stale", "manifest"])
def test_persistent_substrate_unavailable_reference_proceeds(workbook, monkeypatch, tmp_path, failure):
    manifest = substrate.prepare(tmp_path, tmp_path / "shared")
    entry = manifest["workbooks"][str(workbook)]
    if failure == "missing":
        Path(entry["db_path"]).unlink()
    elif failure == "corrupt":
        Path(entry["db_path"]).write_bytes(b"not sqlite")
    elif failure == "stale":
        entry["workbook_hash"] = "old"
    else:
        manifest = {"workbooks": []}
    events = runtime(monkeypatch, tmp_path, manifest)
    wb = openpyxl.load_workbook(workbook)
    assert not isinstance(wb, ProxyWorkbook)
    assert wb["Data"]["C1"].value == 7
    assert not any(e.get("status") == "ACCELERATED" for e in events)
    assert any(e.get("event") == "substrate_fallback" for e in events)
    wb.close()


def test_unsupported_chartsheet_uses_reference(workbook):
    from openpyxl.chart import BarChart, Reference
    wb = REAL(workbook)
    chart = BarChart(); chart.add_data(Reference(wb["Data"], min_col=3, min_row=1, max_row=2))
    wb.create_chartsheet("Chart").add_chart(chart)
    wb.save(workbook); wb.close()
    loader = CandidateALoader(real_loader=REAL)
    result = loader.load_workbook(workbook)
    assert not isinstance(result, ProxyWorkbook)
    assert result.sheetnames[-1] == "Chart"
    assert index.failure_events[-1]["exception_class"] == "NotImplementedError"
    result.close()


def test_valid_candidate_surface_remains_exact(workbook, monkeypatch, tmp_path):
    manifest = substrate.prepare(tmp_path, tmp_path / "shared")
    events = runtime(monkeypatch, tmp_path, manifest)
    proxy = openpyxl.load_workbook(workbook)
    real = REAL(workbook)
    assert isinstance(proxy, ProxyWorkbook)
    assert proxy.sheetnames == real.sheetnames
    for name in ["Data", "ユニコード", "EmptyMetadata"]:
        assert proxy[name].dimensions == real[name].dimensions
        assert list(proxy[name].iter_rows(values_only=True)) == list(real[name].iter_rows(values_only=True))
    assert any(e.get("status") == "ACCELERATED" for e in events)
    assert not any(e.get("event") == "substrate_fallback" for e in events)
    real.close(); proxy.close()


def test_retired_proxy_and_sql_failure_use_reference(workbook):
    events = []
    loader = CandidateALoader(real_loader=REAL, event_sink=events.append)
    wb = loader.load_workbook(workbook)
    cell = wb["Data"]["C1"]
    wb._snapshot.handle["db"].close()
    assert cell.value == 7
    n = len(events)
    assert cell.value == 7
    assert wb.sheetnames == ["EmptyMetadata", "Data", "ユニコード"]
    assert not any(e.get("status") == "ACCELERATED" for e in events[n:])
    assert index.failure_events[-1]["accelerated_operation_already_served"]
    wb.close()


def test_backup_failure_never_publishes_partial(workbook, monkeypatch, tmp_path):
    original = substrate.os.replace
    def fail(src, dst):
        if str(dst).endswith(".sqlite"):
            raise OSError("publication failed")
        return original(src, dst)
    monkeypatch.setattr(substrate.os, "replace", fail)
    result = substrate.prepare(tmp_path, tmp_path / "shared")
    assert not result["workbooks"]
    assert result["failures"][0]["stage"] == "compiled_database_publication"
    assert not list((tmp_path / "shared").glob("*.tmp"))
    assert not json.loads((tmp_path / "shared/manifest.json").read_text())["workbooks"]


def test_scope_is_per_workbook(workbook, tmp_path):
    (tmp_path / "bad.xlsx").write_bytes(b"invalid")
    result = substrate.prepare(tmp_path, tmp_path / "shared")
    assert list(result["workbooks"]) == [str(workbook)]
    assert len(result["failures"]) == 1


def test_freshness_initialization_failure_is_sticky(workbook, monkeypatch):
    monkeypatch.setattr(index, "workbook_hash", lambda p: (_ for _ in ()).throw(OSError("unavailable")))
    result = CandidateALoader(real_loader=REAL).load_workbook(workbook)
    assert not isinstance(result, ProxyWorkbook)
    assert index.failure_events[-1]["stage"] == "freshness_initialization"
    result.close()


def test_reference_helpers_no_index_contact(workbook, monkeypatch):
    from benchmark.inspection_helpers import lx_helpers, reference_api
    from benchmark.representative_checkpoint import SHIM_H0_TEXT, SHIM_H1_TEXT
    monkeypatch.setattr(index, "ensure_fresh", lambda p: pytest.fail("helper used index"))
    assert lx_helpers.inspect is reference_api.inspect
    assert SHIM_H0_TEXT == SHIM_H1_TEXT
    assert lx_helpers.inspect(str(workbook), "Data", "C1:D1") == reference_api.inspect(str(workbook), "Data", "C1:D1")


def test_filelike_fallback_passes_original_object(workbook):
    stream = io.BytesIO(workbook.read_bytes())
    seen = []
    def original(filename, *args, **kwargs):
        seen.append(filename)
        return REAL(filename, *args, **kwargs)
    wb = CandidateALoader(real_loader=original).load_workbook(stream)
    assert seen == [stream]
    wb.close()


def test_partial_database_without_commit_marker_rejected(workbook, monkeypatch, tmp_path):
    manifest = substrate.prepare(tmp_path, tmp_path / "shared")
    entry = manifest["workbooks"][str(workbook)]
    with sqlite3.connect(entry["db_path"]) as con:
        con.execute("DROP TABLE substrate_identity")
    events = runtime(monkeypatch, tmp_path, manifest)
    wb = openpyxl.load_workbook(workbook)
    assert not isinstance(wb, ProxyWorkbook)
    assert not any(e.get("status") == "ACCELERATED" for e in events)
    wb.close()


def test_manifest_storage_unavailable_disables_runner_scope(workbook, monkeypatch, tmp_path):
    monkeypatch.setattr(substrate, "_write_manifest", lambda *a: (_ for _ in ()).throw(PermissionError("unavailable")))
    result = substrate.prepare(tmp_path, tmp_path / "shared")
    assert not result["workbooks"]
    assert result["failures"][0]["exception_class"] == "PermissionError"
    events = runtime(monkeypatch, tmp_path, result)
    wb = openpyxl.load_workbook(workbook)
    assert not isinstance(wb, ProxyWorkbook)
    assert not any(e.get("status") == "ACCELERATED" for e in events)
    wb.close()


def test_new_load_after_mutation_falls_back_until_parent_refresh(workbook, monkeypatch, tmp_path):
    manifest = substrate.prepare(tmp_path, tmp_path / "shared")
    events = runtime(monkeypatch, tmp_path, manifest)
    old = openpyxl.load_workbook(workbook)
    assert isinstance(old, ProxyWorkbook)
    wb = REAL(workbook); wb["Data"]["C1"] = 42; wb.save(workbook); wb.close()
    new = openpyxl.load_workbook(workbook)
    assert not isinstance(new, ProxyWorkbook)
    assert new["Data"]["C1"].value == 42
    n = len(events)
    assert old["Data"]["C1"].value == 42
    assert not any(e.get("status") == "ACCELERATED" for e in events[n:])
    new.close(); old.close()
    monkeypatch.setattr(openpyxl, "load_workbook", REAL)
    manifest = substrate.prepare(tmp_path, tmp_path / "shared")
    runtime(monkeypatch, tmp_path, manifest)  # next Python generation
    rebuilt = openpyxl.load_workbook(workbook)
    assert isinstance(rebuilt, ProxyWorkbook)
    assert rebuilt["Data"]["C1"].value == 42
    rebuilt.close()


def test_helpers_bypass_candidate_loader_including_styles(workbook, monkeypatch):
    from benchmark.inspection_helpers import lx_helpers
    monkeypatch.setattr(openpyxl, "load_workbook", lambda *a, **k: pytest.fail("helper contacted Candidate A"))
    result = lx_helpers.inspect(str(workbook), "Data", "C1:D1", with_styles=True)
    assert result["results"][0]["value"] == "7"


def test_runtime_initialization_restores_original_loader(monkeypatch):
    from benchmark import candidate_a_live_runtime as live
    events = []
    monkeypatch.setenv("CANDIDATE_A_ARM", "H1")
    monkeypatch.setattr(live, "emit", events.append)
    monkeypatch.setattr(openpyxl, "load_workbook", REAL)
    def broken_install():
        openpyxl.load_workbook = lambda *a, **k: None
        raise RuntimeError("bootstrap failure")
    monkeypatch.setattr(live, "install", broken_install)
    live.main()
    assert openpyxl.load_workbook is REAL
    assert events[-1]["stage"] == "runtime_installation"
    assert events[-1]["partially_initialized"] is True


def test_failed_build_retires_existing_persistent_snapshot(workbook, monkeypatch, tmp_path):
    manifest = substrate.prepare(tmp_path, tmp_path / "shared")
    events = runtime(monkeypatch, tmp_path, manifest)
    proxy = openpyxl.load_workbook(workbook)
    assert isinstance(proxy, ProxyWorkbook)
    monkeypatch.setattr(openpyxl, "load_workbook", REAL)
    wb = REAL(workbook); wb["Data"]["C1"] = 99; wb.save(workbook); wb.close()
    monkeypatch.setattr(index, "_open_db", lambda: (_ for _ in ()).throw(RuntimeError("refresh failure")))
    with pytest.raises(index.SubstrateDisabled):
        index.ensure_fresh(workbook)
    assert proxy._snapshot.handle["valid"] is False
    n = len(events)
    assert proxy["Data"]["C1"].value == 99
    assert not any(e.get("status") == "ACCELERATED" for e in events[n:])
    assert index.failure_events[-1]["accelerated_operation_already_served"] is True
    proxy.close()


def test_new_loader_cannot_revive_failed_generation(workbook, monkeypatch, tmp_path):
    manifest = substrate.prepare(tmp_path, tmp_path / "shared")
    entry = manifest["workbooks"][str(workbook)]
    index.disable(str(workbook), entry["workbook_hash"], "snapshot_initialization", RuntimeError("failed"))
    events = runtime(monkeypatch, tmp_path, manifest)
    wb = openpyxl.load_workbook(workbook)
    assert not isinstance(wb, ProxyWorkbook)
    assert not any(e.get("status") == "ACCELERATED" for e in events)
    wb.close()
