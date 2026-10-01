"""Unit tests for the GLM dependency localization probe."""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl

from collections import Counter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "benchmark"))

from formula_dependency_glm_probe import (  # noqa: E402
    CONDITIONS,
    DEPENDENCY_PREFACE,
    QUESTION_C0,
    QUESTION_C1,
    QUESTION_C2,
    SYSTEM_CLASSIFY,
    SYSTEM_ORACLE,
    build_prompt,
    classification_metrics,
    extract_window,
    even_space,
    paired_transitions,
    parse_classify,
    parse_oracle,
    pick_bucket,
    pick_stratum,
    point_consumers,
    point_stratum,
    render_window,
    scan_forbidden,
    tag_reason,
)
from formula_dependency_selection import build_graph  # noqa: E402


def test_even_space_is_deterministic_and_spans() -> None:
    items = list(range(10))
    assert even_space(items, 1) == [0]
    assert even_space(items, 2) == [0, 9]
    assert even_space(items, 3) == [0, 4, 9]
    assert even_space(items, 10) == items
    assert even_space(items, 0) == []


def test_point_stratum_mutually_exclusive() -> None:
    assert (
        point_stratum({"n_point": 2, "point_same_sheet": True, "point_cross_sheet": False})
        == "SAME_SHEET_ONLY"
    )
    assert (
        point_stratum({"n_point": 1, "point_same_sheet": False, "point_cross_sheet": True})
        == "CROSS_SHEET_ONLY"
    )
    assert (
        point_stratum({"n_point": 3, "point_same_sheet": True, "point_cross_sheet": True})
        == "BOTH"
    )
    assert point_stratum({"n_point": 0, "point_same_sheet": False, "point_cross_sheet": False}) is None


def test_pick_stratum_caps_workbook_across_labels() -> None:
    true_cands = []
    blank_cands = []
    for task_id in ("14_05", "15_04", "08_01", "20_04", "13_03", "03_01", "17_03"):
        for index in range(6):
            true_cands.append(
                {
                    "cell_id": f"{task_id}:T!A{index + 1}",
                    "task_id": task_id,
                    "n_point_eq": 1 + index,
                    "n_point": 1 + index,
                    "sheet": "T",
                    "col": 1,
                    "row": index + 1,
                    "eval_role": "TRUE_TARGET",
                }
            )
            blank_cands.append(
                {
                    "cell_id": f"{task_id}:B!A{index + 1}",
                    "task_id": task_id,
                    "n_point_eq": 1 + index,
                    "n_point": 1 + index,
                    "sheet": "B",
                    "col": 1,
                    "row": index + 1,
                    "eval_role": "INTENTIONAL_BLANK",
                }
            )
    true_sel, blank_sel, meta = pick_stratum(
        true_cands,
        blank_cands,
        ["14_05", "15_04", "08_01", "20_04", "13_03", "03_01", "17_03"],
        k_each=6,
        max_per_book=2,
    )
    selected = true_sel + blank_sel
    assert len(true_sel) == 6
    assert len(blank_sel) == 6
    counts = Counter(row["task_id"] for row in selected)
    assert max(counts.values()) <= 2
    assert meta["relaxed"] is False
    assert len({row["n_point"] for row in selected}) >= 3


def test_pick_bucket_is_deterministic_and_respects_book_cap() -> None:
    candidates = []
    for task_id, n in (("14_05", 8), ("17_03", 5), ("03_01", 2), ("15_04", 9), ("08_01", 7)):
        for index in range(n):
            candidates.append(
                {
                    "cell_id": f"{task_id}:S!A{index + 1}",
                    "task_id": task_id,
                    "n_point_eq": 1 + (index % 3),
                    "n_point": index + 1,
                    "sheet": "S",
                    "col": 1,
                    "row": index + 1,
                }
            )
    picked, meta = pick_bucket(
        candidates,
        ["09_05", "14_05", "17_03", "13_03", "20_04", "03_01", "15_04", "08_01"],
        k=6,
        max_per_book=2,
    )
    assert len(picked) == 6
    counts = {}
    for row in picked:
        counts[row["task_id"]] = counts.get(row["task_id"], 0) + 1
        assert counts[row["task_id"]] <= 2
    assert len(counts) >= 4
    assert meta["relaxed"] is False
    assert len({row["n_point"] for row in picked}) >= 2
    again, _ = pick_bucket(
        candidates,
        ["09_05", "14_05", "17_03", "13_03", "20_04", "03_01", "15_04", "08_01"],
        k=6,
        max_per_book=2,
    )
    assert [row["cell_id"] for row in again] == [row["cell_id"] for row in picked]


def test_window_clips_at_origin_and_reports_merges(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "M"
    sheet["A1"] = "Year"
    sheet["B1"] = 2024
    sheet.merge_cells("B1:D1")
    sheet["A2"] = None
    path = tmp_path / "in.xlsx"
    workbook.save(path)
    workbook.close()
    loaded = openpyxl.load_workbook(path)
    ws = loaded.active
    window = extract_window(ws, 1, 1, radius=3)
    loaded.close()
    assert window["clipped_left"] is True
    assert window["clipped_top"] is True
    assert window["clipped_right"] is False
    assert window["used"][0] == 1
    assert window["used"][1] == 1
    assert window["used"][2] == 4
    assert window["used"][3] == 4
    target = next(cell for cell in window["cells"] if cell["is_target"])
    assert target["display"] == "TARGET [BLANK]"
    assert any(item["range"] == "B1:D1" for item in window["merged"])
    text = render_window(window)
    assert "TARGET [BLANK]" in text
    assert "clipped_left" in text


def test_window_includes_target_past_used_range(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "M"
    sheet["A1"] = 1
    path = tmp_path / "in.xlsx"
    workbook.save(path)
    workbook.close()
    loaded = openpyxl.load_workbook(path)
    ws = loaded.active
    window = extract_window(ws, 62, 12, radius=3)
    loaded.close()
    assert window["beyond_used_range"] is True
    target = next(cell for cell in window["cells"] if cell["is_target"])
    assert target["address"] == "BJ12"
    assert "TARGET [BLANK]" in render_window(window)
    assert window["used"] == [59, 9, 65, 15]


def test_prompts_are_neutral_and_omit_golden() -> None:
    cell = {
        "instruction": "Complete the financial model based on the provided assumptions.",
        "category": "Financial_Model",
        "task_id": "13_03",
        "spreadsheet_path": "spreadsheet/13_Project Logistic/13_03_Logistic_input.xlsx",
        "sheet": "P&L",
        "address": "C10",
        "eval_role": "TRUE_TARGET",
        "eval_golden_formula": "=Assumptions!C10*Rates!C10",
        "n_point": 3,
        "n_point_same": 2,
        "n_point_cross": 1,
        "n_point_eq": 2,
        "consumers_truncated": 0,
        "window": {"radius": 3},
        "window_text": "row   C\n10    TARGET [BLANK]",
        "consumers": [
            {
                "sheet": "P&L",
                "address": "C20",
                "relation": "same-sheet",
                "formula": "=C10+C11",
                "consumer_class_id": "EQ01",
            },
            {
                "sheet": "BS",
                "address": "C4",
                "relation": "cross-sheet",
                "formula": "=P&L!C10",
                "consumer_class_id": "EQ02",
            },
        ],
    }
    probe_text = "\n".join(
        [SYSTEM_CLASSIFY, SYSTEM_ORACLE, QUESTION_C0, QUESTION_C1, QUESTION_C2, DEPENDENCY_PREFACE]
    )
    assert scan_forbidden(probe_text) == []
    c0 = build_prompt(cell, "C0")
    c1 = build_prompt(cell, "C1")
    c2 = build_prompt(cell, "C2")
    assert cell["eval_golden_formula"] not in c0
    assert cell["eval_golden_formula"] not in c1
    assert cell["eval_golden_formula"] not in c2
    assert "TRUE_TARGET" not in c0
    assert "TARGET [BLANK]" in c0
    assert "DEPENDENCY FACTS FOR TARGET" not in c0
    assert DEPENDENCY_PREFACE in c1
    assert "=C10+C11" not in c1
    assert "=C10+C11" in c2
    assert "P&L!C20" in c1
    assert "No ranking" in c2
    assert scan_forbidden(DEPENDENCY_PREFACE) == []
    assert scan_forbidden(SYSTEM_CLASSIFY) == []
    assert scan_forbidden(QUESTION_C0) == []
    assert scan_forbidden(QUESTION_C1) == []
    assert scan_forbidden(QUESTION_C2) == []
    c1_block = c1.split("DEPENDENCY FACTS FOR TARGET", 1)[1]
    assert scan_forbidden(c1_block) == []


def test_parse_classify_and_oracle_from_fences() -> None:
    parsed = parse_classify(
        '```json\n{"label": "MISSING_FORMULA", "probability_missing_formula": 0.72, "reason": "Row total is blank."}\n```'
    )
    assert parsed["label"] == "MISSING_FORMULA"
    assert parsed["probability_missing_formula"] == 0.72
    percent = parse_classify(
        '{"label": "INTENTIONAL_BLANK", "probability_missing_formula": "8%", "reason": "Optional assumption."}'
    )
    assert percent["label"] == "INTENTIONAL_BLANK"
    assert percent["probability_missing_formula"] == 0.08
    oracle = parse_oracle('{"formula": "C10+1", "reason": "Add one."}')
    assert oracle["formula"] == "=C10+1"


def test_metrics_and_paired_transitions() -> None:
    def row(cell_id: str, role: str, label: str, p: float) -> dict:
        return {
            "cell_id": cell_id,
            "eval_role": role,
            "label": label,
            "probability_missing_formula": p,
            "valid": True,
        }

    c0 = {
        "a": row("a", "TRUE_TARGET", "INTENTIONAL_BLANK", 0.2),
        "b": row("b", "INTENTIONAL_BLANK", "INTENTIONAL_BLANK", 0.1),
        "c": row("c", "TRUE_TARGET", "MISSING_FORMULA", 0.8),
    }
    c1 = {
        "a": row("a", "TRUE_TARGET", "MISSING_FORMULA", 0.7),
        "b": row("b", "INTENTIONAL_BLANK", "MISSING_FORMULA", 0.6),
        "c": row("c", "TRUE_TARGET", "MISSING_FORMULA", 0.9),
    }
    metrics = classification_metrics(list(c0.values()))
    assert metrics["tp"] == 1
    assert metrics["fn"] == 1
    assert metrics["tn"] == 1
    assert metrics["fp"] == 0
    trans = paired_transitions(c0, c1)
    assert trans["wrong_to_right"] == 1
    assert trans["right_to_wrong"] == 1
    assert trans["unchanged_right"] == 1
    assert trans["net_correct"] == 0


def test_rationale_tags_use_explicit_reason_only() -> None:
    assert "LOCAL_SEMANTICS" in tag_reason("The header row looks like a totals line.", "C0")
    assert "DEPENDENCY_USED" in tag_reason(
        "Three formulas on another sheet reference this cell.", "C1"
    )
    c2 = tag_reason("The consumer formula =P&L!C10 pulls this value into assets.", "C2")
    assert "CONSUMER_FORMULA_USED" in c2
    assert "BLANK_AS_ZERO_ASSUMPTION" in tag_reason(
        "Referenced blanks are treated as zero.", "C1"
    )
    assert tag_reason("Not sure.", "C0") == ["UNSUPPORTED_GUESS"]


def test_point_consumers_cap_and_exclude_ranges(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "M"
    other = workbook.create_sheet("N")
    sheet["B1"] = None
    for index, addr in enumerate(["A1", "A2", "A3"], 1):
        sheet[addr] = f"=B1+{index}"
    sheet["A10"] = "=SUM(B1:B20)"
    other["C1"] = "=M!B1"
    path = tmp_path / "in.xlsx"
    workbook.save(path)
    workbook.close()
    graph = build_graph(path)
    packed = point_consumers(graph, ("M", 2, 1), cap=2)
    assert packed["n_point"] == 4
    assert packed["truncated"] == 2
    assert packed["n_point_same"] == 3
    assert packed["n_point_cross"] == 1
    assert len(packed["listed"]) == 2
    assert packed["listed"][0]["address"] == "A1"
    formulas = {item["formula"] for item in packed["all"]}
    assert "=SUM(B1:B20)" not in formulas


def test_all_conditions_are_independent_prompt_variants() -> None:
    assert CONDITIONS == ("C0", "C1", "C2")
    assert "formula" not in QUESTION_C0.lower() or "missing a formula" in QUESTION_C0
