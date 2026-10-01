"""Unit tests for the Phase-12 frozen verifier (synthetic workbooks)."""
import sys
import zipfile
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "phase12"))

from verification_block import verifier
from verification_block.verifier import (patch_full_calc, recalc, render,
                                         rewrite_filter, verify)


def _family_wb(path, break_cell=None, break_formula=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    for r in range(1, 11):
        ws.cell(row=r, column=1, value=r)
        ws.cell(row=r, column=2, value=f"=A{r}*2")
    if break_cell:
        ws[break_cell] = break_formula
    wb.save(path)


def test_rewrite_filter_keeps_localized_break(tmp_path):
    from phase11.mine import derive_workbook, structural_diff, adjacent_family_breaks
    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "phase11"))
    pre_p, post_p = tmp_path / "in.xlsx", tmp_path / "out.xlsx"
    _family_wb(pre_p)
    _family_wb(post_p, break_cell="B5", break_formula="=SUM(A1:A5)")
    pre, post = derive_workbook(str(pre_p)), derive_workbook(str(post_p))
    diff = structural_diff(pre, post)
    pre_cells = {b["cell"] for b in adjacent_family_breaks(pre)}
    fresh = [b for b in adjacent_family_breaks(post) if b["cell"] not in pre_cells]
    kept = rewrite_filter(fresh, pre, post, diff)
    assert [b["cell"] for b in kept] == ["S!B5"]


def test_rewrite_filter_suppresses_broad_rewrite(tmp_path):
    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "phase11"))
    from phase11.mine import derive_workbook, structural_diff, adjacent_family_breaks
    pre_p, post_p = tmp_path / "in.xlsx", tmp_path / "out.xlsx"
    _family_wb(pre_p)
    wb = openpyxl.load_workbook(post_p if False else pre_p)
    wb.save(post_p)
    wb = openpyxl.load_workbook(post_p)
    for r in range(1, 11):
        wb["S"].cell(row=r, column=2, value=f"=A{r}*3")
    wb.save(post_p)
    pre, post = derive_workbook(str(pre_p)), derive_workbook(str(post_p))
    diff = structural_diff(pre, post)
    pre_cells = {b["cell"] for b in adjacent_family_breaks(pre)}
    fresh = [b for b in adjacent_family_breaks(post) if b["cell"] not in pre_cells]
    # new coherent pattern installed everywhere: nothing survives
    assert rewrite_filter(fresh, pre, post, diff) == []


def test_render_caps_and_omits_empty_sections():
    report = {"signals": {
        "ERR": {"items": [{"cell": f"S!A{i}", "error": "#REF!"} for i in range(1, 15)],
                "count": 14, "sheets": 1, "types": {"#REF!": 14}, "changed_overlap": 1},
        "UNIF": {"items": [], "count": 0},
        "REF": {"items": [], "count": 0},
        "CHG": {"changed_cells": 3, "sheets": 1, "regions": []}}}
    text = render(report)
    assert "New formula errors (14)" in text
    assert "and 4 more" in text
    assert "pattern breaks" not in text
    assert "blank referenced" not in text
    for banned in ["fix", "wrong", "should", "must", "correct target"]:
        assert banned not in text.lower()


def test_verify_end_to_end_new_error(tmp_path):
    pre_p, post_p = tmp_path / "in.xlsx", tmp_path / "out.xlsx"
    _family_wb(pre_p)
    _family_wb(post_p, break_cell="B5", break_formula="=A5/0")
    report = verify(pre_p, post_p)
    assert report["status"] == "OK"
    assert report["positive"] is True
    assert "ERR" in report["positive_families"]
    assert any(i["cell"] == "S!B5" and i["error"] == "#DIV/0!"
               for i in report["signals"]["ERR"]["items"])
    assert report["model_block"] is not None


def test_verify_zero_signal_returns_none(tmp_path):
    pre_p, post_p = tmp_path / "in.xlsx", tmp_path / "out.xlsx"
    _family_wb(pre_p)
    _family_wb(post_p)
    report = verify(pre_p, post_p)
    assert report["status"] == "OK"
    assert report["positive"] is False
    assert report["model_block"] is None


def test_verify_unreadable_is_unavailable(tmp_path):
    missing = tmp_path / "nope.xlsx"
    pre_p = tmp_path / "in.xlsx"
    _family_wb(pre_p)
    report = verify(pre_p, missing)
    assert report["positive"] is None
    assert report["status"] == "VERIFIER_UNAVAILABLE"
