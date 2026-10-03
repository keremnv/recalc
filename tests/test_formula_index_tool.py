from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
LIB = ROOT / "benchmark/sweagent/formula_index/lib"
sys.path.insert(0, str(LIB))
sys.path.insert(0, str(ROOT / "benchmark"))

from fingerprint import relative_fingerprint  # noqa: E402
from ranges import compress_cells  # noqa: E402
from render import OBSERVATION_LIMIT, render_classes  # noqa: E402
from workbook import EquivalenceClass, build_index  # noqa: E402
from cli import main as cli_main  # noqa: E402


def _xlsx(tmp_path: Path, cells: dict[tuple[str, str], object]) -> Path:
    path = tmp_path / "book.xlsx"
    wb = openpyxl.Workbook()
    first = True
    sheets: dict[str, object] = {}
    for (sheet, addr), value in cells.items():
        if sheet not in sheets:
            if first:
                ws = wb.active
                ws.title = sheet
                first = False
            else:
                ws = wb.create_sheet(sheet)
            sheets[sheet] = ws
        sheets[sheet][addr] = value
    wb.save(path)
    return path


def test_range_returns_each_class_once(tmp_path: Path) -> None:
    path = _xlsx(
        tmp_path,
        {
            ("Model", "F10"): "=E10-E15",
            ("Model", "G10"): "=F10-F15",
            ("Model", "H10"): "=G10-G15",
            ("Model", "F11"): "=F10+1",
        },
    )
    index = build_index(path)
    records = index.classes_in_range("Model", "F10:H10")
    assert len(records) == 1
    assert records[0].n == 3
    assert "Model!F10:H10" in records[0].member_text()


def test_axis_works_on_a_blank_cell(tmp_path: Path) -> None:
    path = _xlsx(
        tmp_path,
        {
            ("Model", "F10"): "=E10-E15",
            ("Model", "G10"): "=F10-F15",
        },
    )
    index = build_index(path)
    records = index.classes_on_axis("Model", row=10, col=13)
    assert len(records) == 1
    assert records[0].eq_id == relative_fingerprint("=E10-E15", 6, 10, sheet="Model").eq_id


def test_axis_unions_row_and_column_without_ranking(tmp_path: Path) -> None:
    path = _xlsx(
        tmp_path,
        {
            ("Model", "A1"): "=B1",
            ("Model", "C3"): "=C2",
        },
    )
    index = build_index(path)
    records = index.classes_on_axis("Model", row=1, col=3)
    assert {record.eq_id for record in records} == {
        relative_fingerprint("=B1", 1, 1, sheet="Model").eq_id,
        relative_fingerprint("=C2", 3, 3, sheet="Model").eq_id,
    }


def test_lookup_returns_cross_sheet_members(tmp_path: Path) -> None:
    path = _xlsx(
        tmp_path,
        {
            ("Model", "F10"): "=E10",
            ("Schedule", "F10"): "=E10",
        },
    )
    index = build_index(path)
    eq_id = relative_fingerprint("=E10", 6, 10, sheet="Model").eq_id
    record = index.lookup(eq_id)
    assert record is not None
    members = record.member_text()
    assert "Model!F10" in members
    assert "Schedule!F10" in members


def test_opaque_cells_do_not_form_a_class(tmp_path: Path) -> None:
    path = _xlsx(
        tmp_path,
        {
            ("Model", "A1"): '=INDIRECT("B1")',
            ("Model", "A2"): '=INDIRECT("B1")',
            ("Model", "C1"): "=B1",
            ("Model", "C2"): "=B2",
        },
    )
    index = build_index(path)
    opaque = [record for record in index.classes.values() if record.opaque]
    canonical = [record for record in index.classes.values() if not record.opaque]
    assert len(opaque) == 2
    assert all(record.n == 1 for record in opaque)
    assert opaque[0].eq_id != opaque[1].eq_id
    assert len(canonical) == 1
    assert canonical[0].n == 2


def test_compress_cells_builds_rectangles() -> None:
    cells = {(c, r) for c in range(1, 4) for r in range(1, 3)}
    assert compress_cells(cells) == ["A1:C2"]


def test_cli_range_and_axis(tmp_path: Path) -> None:
    path = _xlsx(tmp_path, {("Model", "A1"): "=B1", ("Model", "A2"): "=B2"})
    assert cli_main(["range", str(path), "Model", "A1:A2"]) == 0
    assert cli_main(["axis", str(path), "Model", "--row", "1"]) == 0
    assert cli_main(["axis", str(path), "Model", "--row", "1", "--col", "1"]) == 0


def test_sweagent_accepts_formula_index_signature() -> None:
    """SWE-agent requires every argument name as <name> in the signature so invoke_format keys match."""
    import re
    import string

    import yaml
    pytest.importorskip("jinja2", reason="optional harness templating dependency")
    from jinja2 import Template

    spec = yaml.safe_load(
        (ROOT / "benchmark/sweagent/formula_index/config.yaml").read_text()
    )["tools"]["formula_index"]
    signature = spec["signature"]
    names = [arg["name"] for arg in spec["arguments"]]
    for name in names:
        assert (
            f"<{name}>" in signature
            or f"[<{name}>]" in signature
            or f"--{name}" in signature
        ), name
    invoke_format = re.sub(r"\[?<([a-zA-Z_][a-zA-Z0-9_-]*)>\]?", r"{\1}", signature)
    invoke_keys = {
        field_name
        for _, field_name, _, _ in string.Formatter().parse(invoke_format)
        if field_name is not None
    }
    assert invoke_keys == set(names)
    row_format = next(arg["argument_format"] for arg in spec["arguments"] if arg["name"] == "row")
    filled = {
        "command": "axis",
        "xlsx": "/tmp/book.xlsx",
        "sheet": "Model",
        "a1_range_or_eq_id": "",
        "row": Template(row_format).render(value=10),
        "col": "",
    }
    invoked = invoke_format.format(**filled).strip()
    assert invoked == "formula_index axis /tmp/book.xlsx Model  --row 10"


def _record(eq_id: str, formula: str, n: int = 1) -> EquivalenceClass:
    return EquivalenceClass(
        eq_id=eq_id,
        fingerprint=eq_id,
        opaque=False,
        reason=None,
        n=n,
        exemplar_sheet="Model",
        exemplar_col=1,
        exemplar_row=1,
        exemplar_formula=formula,
        members={"Model": {(1, 1)}},
    )


def test_truncation_keeps_every_class_identity() -> None:
    records = [_record(f"{index:010d}", "=A1+" + ("X" * 80), n=10) for index in range(40)]
    text = render_classes(
        kind="axis",
        header_fields=["sheet=Model", "row=1"],
        records=records,
        limit=2000,
    )
    assert len(text) <= 2000
    assert "truncated=true" in text.splitlines()[0]
    for record in records:
        assert f"eq {record.eq_id}" in text


def test_score_fingerprint_match_and_trajectory_retrieval(tmp_path: Path) -> None:
    from formula_index_score import occupancy_targets, parse_trajectory, score_targets
    from workbook import build_index

    path = _xlsx(
        tmp_path,
        {
            ("Model", "F10"): "=E10-E15",
            ("Model", "G10"): "=F10-F15",
        },
    )
    gpath = tmp_path / "golden.xlsx"
    wb = openpyxl.load_workbook(path)
    wb["Model"]["H10"] = "=G10-G15"
    wb.save(gpath)
    index = build_index(path)
    from formula_index_score import load_formulas, load_values

    inp, inv = load_formulas(path), load_values(path)
    gold, gval = load_formulas(gpath), load_values(gpath)
    targets = occupancy_targets(index, inp, inv, gold, gval)
    blank = [t for t in targets if t["blank_fill"]]
    assert blank
    eq_id = blank[0]["golden_eq_id"]
    traj = tmp_path / "t.traj"
    traj.write_text(
        json.dumps(
            {
                "info": {"model_stats": {"api_calls": 3, "tokens_sent": 1, "tokens_received": 1, "instance_cost": 0}},
                "trajectory": [
                    {
                        "action": "formula_index axis /tmp/book.xlsx Model --row 10",
                        "observation": f"eq {eq_id} n=2 sheets=1\n  source Model!F10 =E10-E15\n",
                    }
                ],
            }
        )
    )
    parsed = parse_trajectory(traj)
    assert parsed["invoked"] is True
    submitted = dict(inp)
    submitted[(blank[0]["sheet"], blank[0]["col"], blank[0]["row"])] = "=G10-G15"
    rows = score_targets(targets, submitted, parsed, task="Financial_Model:x")
    matched = [row for row in rows if row["blank_fill"]]
    assert matched[0]["fingerprint_match"] is True
    assert matched[0]["retrieved_class"] is True


def test_full_payload_stays_under_observation_cap_when_small() -> None:
    records = [_record("abcd123456", "=A1")]
    text = render_classes(kind="lookup", header_fields=["eq_id=abcd123456"], records=records)
    assert len(text) < OBSERVATION_LIMIT
    assert "truncated=false" in text
    assert "members " in text
