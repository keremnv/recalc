"""Unit tests for gold-blind block extent vs role extent predicates."""
from __future__ import annotations

import inspect
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "benchmark"))

import formula_hierarchy_probe  # noqa: E402
from formula_dependency_selection import build_graph  # noqa: E402
from formula_frontier import build_extent  # noqa: E402
from formula_hierarchy import (  # noqa: E402
    CONJUNCTIONS,
    DEFINITIONS,
    conjunctions,
    family_a,
    family_f,
    features_for_cell,
    qualifying_rows,
)
from formula_liveness import build_occupancy_peers  # noqa: E402
from formula_target_selection import occupancy_from_cells  # noqa: E402


def _save(tmp_path: Path, fill) -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "M"
    fill(sheet, workbook)
    path = tmp_path / "in.xlsx"
    workbook.save(path)
    workbook.close()
    return path


def test_definitions_are_gold_blind() -> None:
    assert DEFINITIONS["golden_in_generation"] is False
    source = inspect.getsource(sys.modules["formula_hierarchy"])
    assert "_golden_path" not in source
    assert "_label(" not in source
    freeze_src = inspect.getsource(formula_hierarchy_probe.cmd_freeze)
    assert "_golden_path" not in freeze_src
    assert "_label" not in freeze_src
    assert "golden" not in inspect.signature(features_for_cell).parameters


def test_aligned_block_frontier() -> None:
    cells = {}
    for row in range(1, 6):
        for col in range(1, 5):
            cells[(col, row)] = "F"
    grid = occupancy_from_cells("M", cells)
    extent = build_extent(grid)
    a = family_a(grid, extent, col=6, row=3, last_t=4)
    assert a["n"] >= 4
    assert a["M"] == 4
    assert a["A1"] is True
    assert a["A2"] is True
    assert a["Role1"] is True
    assert a["Role2"] is False
    f = family_f(grid, a["rows"], col=6, row=3, m=4)
    assert f["F1"] is True


def test_role_shorter_than_block() -> None:
    cells = {}
    for row in (1, 2, 4, 5):
        for col in range(1, 11):
            cells[(col, row)] = "F"
    for col in range(1, 4):
        cells[(col, 3)] = "F"
    grid = occupancy_from_cells("M", cells)
    extent = build_extent(grid)
    a = family_a(grid, extent, col=6, row=3, last_t=3)
    assert a["M"] == 10
    assert a["delta"] == -7
    assert a["Role2_d4"] is True
    assert a["A3"] is True
    assert a["Role1"] is False
    f = family_f(grid, a["rows"], col=6, row=3, m=a["M"])
    assert f["F2"] is True


def test_h_conjunctions() -> None:
    flags = {
        "A1": True,
        "A2": True,
        "B1": True,
        "D1": True,
        "F1": False,
        "G1": True,
        "Role2_d4": True,
        "block_continues": True,
        "E4": True,
        "E1": False,
        "E2": True,
    }
    out = conjunctions(flags)
    assert out["H1"] is True
    assert out["H2"] is True
    assert out["H3"] is True
    assert out["H6"] is True
    assert out["H7"] is True
    assert out["H8"] is True
    assert out["H10"] is True
    assert out["H9"] is False


def test_features_shape_and_g1(tmp_path: Path) -> None:
    def fill(sheet, _wb) -> None:
        for row in range(2, 7):
            sheet.cell(row=row, column=1, value=1)
            for col in range(2, 5):
                sheet.cell(row=row, column=col, value="=A" + str(row))
            # Downstream demand lives off the block rows so last ACTIVE stays at 4.
            sheet.cell(row=20, column=row, value="=F" + str(row))
        sheet["Z21"] = 1

    graph = build_graph(_save(tmp_path, fill))
    key = ("M", 6, 4)
    occ = build_occupancy_peers(graph.occupancy, {key})
    extent = build_extent(graph.occupancy["M"])
    feat = features_for_cell(
        graph,
        key,
        occ,
        extent,
        None,
        {},
        {"last_active_col": 4, "last_formula_col": 4},
    )
    assert set(CONJUNCTIONS) <= set(feat["flags"])
    assert feat["M"] == 4
    assert feat["flags"]["H1"] is True
    assert feat["flags"]["G1"] is True
    assert qualifying_rows(graph.occupancy["M"], 6, 4, 8, formula=False)
