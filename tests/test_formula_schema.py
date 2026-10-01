"""Unit tests for gold-blind S0/S1 schema extraction."""
from __future__ import annotations

import inspect
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "benchmark"))

import formula_schema_probe  # noqa: E402
from formula_schema import (  # noqa: E402
    DEFINITIONS,
    build_workbook_schema,
    canonical_row_label,
    candidate_row_labels,
    lexical_relation,
    parse_period,
    period_key,
    schema_for_cell,
)
from formula_schema_probe import align_refs, parse_refs  # noqa: E402


def _save(tmp_path: Path, fill) -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Fin"
    fill(sheet, workbook)
    path = tmp_path / "in.xlsx"
    workbook.save(path)
    workbook.close()
    return path


def test_definitions_are_gold_blind() -> None:
    assert DEFINITIONS["golden_in_generation"] is False
    source = inspect.getsource(sys.modules["formula_schema"])
    assert "_golden_path" not in source
    assert "eval_golden_formula" not in source
    freeze_src = inspect.getsource(formula_schema_probe.cmd_extract)
    assert "_golden_path" not in freeze_src
    assert "eval_golden_formula" not in freeze_src


def test_period_parse() -> None:
    assert parse_period(2023) == {"year": 2023}
    assert parse_period("FY2023") == {"year": 2023}
    assert parse_period("Q1 2024") == {"year": 2024, "quarter": 1}
    assert parse_period("Jan-24") == {"year": 2024, "month": 1}
    assert period_key({"year": 2024, "quarter": 1}) == "PERIOD(year=2024, quarter=1)"
    assert parse_period("Tax expense") is None


def test_lexical_jaccard() -> None:
    rel = lexical_relation("Tax expense (GAAP)", "Tax Expense")
    assert rel["ordered_equal_noparen"] is True
    assert rel["jaccard_noparen"] == 1.0
    rel2 = lexical_relation("Revenue growth", "Cost of sales")
    assert rel2["jaccard"] < 0.4


def test_row_label_and_stack(tmp_path: Path) -> None:
    def fill(sheet, wb) -> None:
        sheet["A1"] = "Income Statement"
        sheet.merge_cells("A1:F1")
        sheet["A10"] = "Tax expense"
        sheet["C9"] = "FY2023"
        sheet["D9"] = "FY2024"
        sheet["C10"] = 10
        sheet["D10"] = "=C10*1.1"
        other = wb.create_sheet("Assumptions")
        other["A27"] = "Tax expense"
        other["C27"] = 0.21

    path = _save(tmp_path, fill)
    book = build_workbook_schema(path)
    schema = schema_for_cell(book.sheets["Fin"], 4, 10)
    assert schema["has_immediate"] is True
    assert schema["label_raw"] == "Tax expense"
    assert schema["has_stack"] is True
    assert schema["stack"][0]["raw"] == "Income Statement"
    assert schema["period_key"] == "PERIOD(year=2024)"
    peers = book.label_rows[schema["label_norm"]]
    assert ("Fin", 10) in peers
    assert ("Assumptions", 27) in peers
    cands = candidate_row_labels(book.sheets["Fin"], 4, 10)
    assert canonical_row_label(cands)["raw"] == "Tax expense"


def test_slot_alignment_growth_driver() -> None:
    model = parse_refs("=J6*(1+K12)", "Assumptions")
    gold = parse_refs("=J6*(1+K27)", "Assumptions")
    aligned = align_refs(model, gold)
    assert aligned["same_span_set"] is False
    assert len(aligned["shared"]) == 1
    assert aligned["shared"][0]["ref"]["start"] == "J6"
    assert len(aligned["differing"]) == 1
    assert aligned["differing"][0]["model_ref"]["start"] == "K12"
    assert aligned["differing"][0]["gold_ref"]["start"] == "K27"
