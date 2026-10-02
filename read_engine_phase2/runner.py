"""Fresh-process child and offline Phase-2 orchestrator."""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import platform
import random
import resource
import shutil
import statistics
import subprocess
import sys
import tempfile
import time
import zipfile
from collections import Counter, defaultdict
from pathlib import Path
from xml.etree import ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / "read_engine_phase1"))
import openpyxl  # noqa: E402
from openpyxl.utils.cell import coordinate_to_tuple  # noqa: E402
from prototype import canonical  # noqa: E402
from persistence import acquire, no_openpyxl_load, paths, sha, tick  # noqa: E402

POP = HERE / "population.json"
TRACE_FILE = ROOT / "research/history/candidate_a_a1_checkpoint_rerun_01/contact_traces.jsonl"
SPEC_HASH = "227c7618a244ef9dc4960b6d1aed751165ec326bd2a6d9670895e773df7bddde"
P1_HASH = "a6e95f1503ec6090d6176065471372924bf447f907a2a428faceee12a2f72473"
POP_HASH = "6492d1cbc2163ea252c0d34cf346171d332a923bd5f05fc1fd54c3e851097afa"
TRACE_HASH = "665244f6150dfbb74ae647dfc8a7b8b034c02794f6de9cf7419f6d56da41696f"
VARIANTS = ["P0", "P1", "P2", "P3"]


def frozen(check_sources: bool = True) -> dict:
    for path, digest in ((HERE / "PREREGISTERED_SPEC_v2.md", SPEC_HASH),
                         (ROOT / "read_engine_phase1/prototype.py", P1_HASH),
                         (POP, POP_HASH), (TRACE_FILE, TRACE_HASH)):
        if sha(path) != digest:
            raise RuntimeError(f"Frozen identity changed: {path}")
    pop = json.loads(POP.read_text())
    if len(pop["workbooks"]) != 39 or len(pop["traces"]) != 51:
        raise RuntimeError("Population cardinality changed")
    if check_sources:
        for item in pop["workbooks"]:
            if sha(Path(item["path"])) != item["sha256"]:
                raise RuntimeError(f"Workbook changed: {item['path']}")
        for item in pop["traces"]:
            if sha(Path(item["snapshot_path"])) != item["snapshot_sha256"]:
                raise RuntimeError(f"Snapshot changed: {item['trace_id']}")
    return pop


def traces() -> dict:
    return {row["trace_id"]: row for row in map(json.loads, TRACE_FILE.read_text().splitlines())}


def trace_run(trace: dict, book) -> tuple[list[dict], dict, dict | None]:
    ws = cell = None
    values = []
    first_ns = 0
    op_count = 0
    t_all = tick()
    try:
        for event in trace["operations"]:
            op = event.get("operation")
            if op == "load_workbook":
                continue
            t = tick()
            if op == "Workbook.sheetnames":
                value = book.sheetnames
            elif op == "Workbook.__getitem__":
                ws = book[event["sheet"]]
                value = None
            elif op == "Worksheet.max_row":
                value = ws.max_row
            elif op == "Worksheet.max_column":
                value = ws.max_column
            elif op == "Worksheet.cell":
                row, col = coordinate_to_tuple(event["address"])
                cell = ws.cell(row=row, column=col)
                value = None
            elif op == "Cell.value":
                value = cell.value
            elif op == "Cell.data_type":
                value = cell.data_type
            else:
                continue
            if not op_count:
                first_ns = tick() - t
            op_count += 1
            if op not in ("Workbook.__getitem__", "Worksheet.cell"):
                values.append({"operation": op, "value": canonical(value)})
        return values, {"first_operation_ns": first_ns,
                        "trace_ns": tick() - t_all, "operation_count": op_count}, None
    except Exception as exc:
        return (values, {"first_operation_ns": first_ns,
                         "trace_ns": tick() - t_all, "operation_count": op_count},
                {"class": type(exc).__name__, "message": str(exc)})


def child_invoke(args) -> dict:
    source = Path(args.source)
    cache = Path(args.cache)
    trace = traces()[args.trace_id] if args.trace_id else None
    t0 = tick()
    if args.variant == "P0":
        try:
            book, witness, phases = acquire(source, cache, args.variant)
        except Exception as exc:
            if trace is not None:
                raise
            return {"variant": "P0", "trace_id": None, "witness": "REFERENCE_ORACLE_UNAVAILABLE",
                    "phases": {"build_ns": None}, "error": {"class": type(exc).__name__,
                    "message": str(exc)}, "engine_ns": tick() - t0,
                    "rss_kib": resource.getrusage(resource.RUSAGE_SELF).ru_maxrss}
        if trace:
            values, trace_profile, error = trace_run(trace, book)
        else:
            values, trace_profile, error = [], {}, None
    else:
        with no_openpyxl_load():
            book, witness, phases = acquire(source, cache, args.variant)
            if trace:
                values, trace_profile, error = trace_run(trace, book)
            else:
                values, trace_profile, error = [], {}, None
    t = tick()
    book.close()
    close_ns = tick() - t
    total_ns = tick() - t0
    encoded = json.dumps(values, sort_keys=True, separators=(",", ":"))
    row = {"variant": args.variant, "trace_id": args.trace_id,
           "source_sha256": phases.get("source_sha256"), "witness": witness,
           "phases": phases, "trace_profile": trace_profile, "error": error,
           "trace_digest": hashlib.sha256(encoded.encode()).hexdigest(),
           "value_count": len(values), "close_ns": close_ns,
           "engine_ns": total_ns, "rss_kib": resource.getrusage(resource.RUSAGE_SELF).ru_maxrss}
    if args.emit_values:
        row["values"] = values
    return row


def child_book(args) -> dict:
    source = Path(args.source)
    cache = Path(args.cache)
    results = {}
    books = {}
    for variant in ("P2", "P3"):
        with no_openpyxl_load():
            books[variant], witness, profile = acquire(source, cache, variant)
        results[variant] = {"witness": witness, "checked": 0, "mismatch_count": 0,
                            "first_mismatch": None, "artifact_bytes": profile.get("artifact_bytes")}
    try:
        ref = openpyxl.load_workbook(source, data_only=False)
    except Exception as exc:
        for book in books.values():
            book.close()
        return {"status": "REFERENCE_ORACLE_UNAVAILABLE", "error": type(exc).__name__,
                "message": str(exc), "variants": results}
    try:
        for variant, book in books.items():
            res = results[variant]
            if book.sheetnames != ref.sheetnames:
                res["mismatch_count"] += 1
                res["first_mismatch"] = {"kind": "sheetnames", "want": ref.sheetnames,
                                          "got": book.sheetnames}
            for ws in ref.worksheets:
                other = book[ws.title]
                if (ws.max_row, ws.max_column, ws.dimensions) != (other.max_row, other.max_column, other.dimensions):
                    res["mismatch_count"] += 1
                    res["first_mismatch"] = res["first_mismatch"] or {"kind": "bounds", "sheet": ws.title}
                with no_openpyxl_load():
                    for cell in ws._cells.values():
                        if cell.value is None:
                            continue
                        actual = other.cell(cell.row, cell.column)
                        got = (canonical(actual.value), actual.data_type)
                        want = (canonical(cell.value), cell.data_type)
                        res["checked"] += 1
                        if got != want:
                            res["mismatch_count"] += 1
                            res["first_mismatch"] = res["first_mismatch"] or {
                                "kind": "cell", "sheet": ws.title, "cell": cell.coordinate,
                                "want": want, "got": got}
        return {"status": "EXACT" if all(v["mismatch_count"] == 0 for v in results.values()) else "MISMATCH",
                "variants": results}
    finally:
        ref.close()
        for book in books.values():
            book.close()


def child(args) -> None:
    row = child_book(args) if args.mode == "book" else child_invoke(args)
    print(json.dumps(row, sort_keys=True, default=str))


def call_child(variant: str, source: Path, cache: Path, trace_id: str | None = None,
               emit_values: bool = False, mode: str = "invoke") -> tuple[dict, int]:
    cmd = [sys.executable, str(HERE / "runner.py"), "child", "--variant", variant,
           "--source", str(source), "--cache", str(cache), "--mode", mode]
    if trace_id:
        cmd += ["--trace-id", trace_id]
    if emit_values:
        cmd += ["--emit-values"]
    start = tick()
    completed = subprocess.run(cmd, capture_output=True, text=True)
    elapsed = tick() - start
    if completed.returncode:
        raise RuntimeError(f"Child {variant} {mode} failed: {completed.stderr[-3000:]}")
    return json.loads(completed.stdout), elapsed


def append(path: Path, row: dict) -> None:
    with path.open("a") as f:
        f.write(json.dumps(row, sort_keys=True, default=str) + "\n")


def initial_files() -> None:
    for name in ("raw_correctness.jsonl", "raw_timings.jsonl", "session_timings.jsonl", "invalidation.jsonl"):
        (HERE / name).write_text("")


def correctness(pop: dict) -> bool:
    out = HERE / "raw_correctness.jsonl"
    archive = traces()
    okay = True
    with tempfile.TemporaryDirectory(prefix="phase2-correct-") as d:
        root = Path(d)
        for i, item in enumerate(pop["traces"]):
            trace_id, source = item["trace_id"], Path(item["snapshot_path"])
            reference, _ = call_child("P0", source, root / f"trace-{i}" / "P0", trace_id, True)
            want = reference["values"]
            for variant in ("P1", "P2", "P3"):
                cache = root / f"trace-{i}" / variant
                built, _ = call_child(variant, source, cache, trace_id, True)
                reopened, _ = call_child(variant, source, cache, trace_id, True) if variant in ("P2", "P3") else (built, 0)
                exact = reopened["values"] == want and reopened["error"] == reference["error"]
                if variant in ("P2", "P3") and reopened["witness"] != "REUSED":
                    exact = False
                row = {"kind": "trace", "trace_id": trace_id, "snapshot_sha256": item["snapshot_sha256"],
                       "variant": variant, "exact": exact, "built_witness": built["witness"],
                       "reopen_witness": reopened["witness"], "value_count": len(want),
                       "first_mismatch": next((j for j,(a,b) in enumerate(zip(want,reopened["values"])) if a!=b), None),
                       "reference_error": reference["error"], "treatment_error": reopened["error"]}
                append(out, row)
                okay &= exact
            if (i + 1) % 10 == 0:
                print(f"correctness traces {i+1}/51", flush=True)
        for i, item in enumerate(pop["workbooks"]):
            book_cache = root / f"book-{i}"
            for variant in ("P2", "P3"):
                call_child(variant, Path(item["path"]), book_cache)
            row, _ = call_child("P2", Path(item["path"]), book_cache, mode="book")
            row.update({"kind": "book", "workbook_sha256": item["sha256"]})
            append(out, row)
            okay &= row["status"] in ("EXACT", "REFERENCE_ORACLE_UNAVAILABLE")
            okay &= all(v["witness"] == "REUSED" for v in row["variants"].values())
            if row["status"] == "REFERENCE_ORACLE_UNAVAILABLE" and item["sha256"] != "71233a0b02e0680361836392ca41c567aae657ce91dd24c4248f273f459fd53c":
                okay = False
            print(f"correctness workbooks {i+1}/39", flush=True)
    return okay


def mutate_scalar(source: Path, target: Path) -> dict:
    """Make an OOXML byte-different same-type scalar cell in staged copy."""
    target.parent.mkdir(parents=True, exist_ok=True)
    edited = None
    with zipfile.ZipFile(source) as zin, zipfile.ZipFile(target, "w") as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if edited is None and info.filename.startswith("xl/worksheets/sheet") and info.filename.endswith(".xml"):
                root = ET.fromstring(data)
                ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                for cell in root.findall(".//m:c", ns):
                    value = cell.find("m:v", ns)
                    if value is not None and value.text and cell.get("t") not in ("s", "e", "b"):
                        old = value.text
                        try:
                            float(old)
                        except ValueError:
                            continue
                        value.text = str(float(old) + 1)
                        data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                        edited = {"part": info.filename, "cell": cell.get("r"), "before": old, "after": value.text}
                        break
            zout.writestr(info, data)
    if edited is None:
        raise RuntimeError(f"No numeric scalar to mutate: {source}")
    return edited


def invalidation(pop: dict) -> bool:
    out = HERE / "invalidation.jsonl"
    ids = {"1768dcc784d3c9bd2319dabeaa3a7c27d4bc0ab048a63c7084450abab3f3a1a2",
           "1b64bd5ac59b13beceb019da28ee25a8e12d479e14d625ce3ead88ccc344c741"}
    selected = [x for x in pop["workbooks"] if x["sha256"] in ids]
    okay = len(selected) == 2
    with tempfile.TemporaryDirectory(prefix="phase2-invalidation-") as d:
        root = Path(d)
        for item in selected:
            source = Path(item["path"])
            stage = root / item["sha256"] / "input.xlsx"
            stage.parent.mkdir(parents=True)
            shutil.copyfile(source, stage)
            original = stage.read_bytes()
            h2 = stage.with_name("h2.xlsx")
            edit = mutate_scalar(source, h2)
            for variant in ("P2", "P3"):
                cache = root / item["sha256"] / variant
                sequence = []
                for label in ("H1_BUILD", "H1_REUSE", "H2_BUILD", "H1_RESTORE"):
                    if label == "H2_BUILD":
                        shutil.copyfile(h2, stage)
                    elif label == "H1_RESTORE":
                        stage.write_bytes(original)
                    row, _ = call_child(variant, stage, cache)
                    sequence.append({"step": label, "witness": row["witness"],
                                     "rebuild_reason": row["phases"].get("rebuild_reason"), "sha256": sha(stage)})
                artifact, sidecar, _ = paths(cache, item["sha256"], variant)
                artifact.unlink()
                missing, _ = call_child(variant, stage, cache)
                sequence.append({"step": "MISSING_REBUILD", "witness": missing["witness"],
                                 "rebuild_reason": missing["phases"].get("rebuild_reason")})
                artifact.write_bytes(artifact.read_bytes()[: max(1, artifact.stat().st_size // 3)])
                corrupt, _ = call_child(variant, stage, cache)
                sequence.append({"step": "CORRUPT_REBUILD", "witness": corrupt["witness"],
                                 "rebuild_reason": corrupt["phases"].get("rebuild_reason")})
                expected = ["BUILT", "REUSED", "BUILT", "REUSED", "BUILT", "BUILT"]
                exact = [x["witness"] for x in sequence] == expected
                okay &= exact
                append(out, {"workbook_sha256": item["sha256"], "variant": variant,
                             "h2_sha256": sha(h2), "edit": edit, "sequence": sequence, "exact": exact})
    return okay


def timing(pop: dict) -> None:
    raw, sessions = HERE / "raw_timings.jsonl", HERE / "session_timings.jsonl"
    with tempfile.TemporaryDirectory(prefix="phase2-timings-") as d:
        root = Path(d)
        for i, item in enumerate(pop["workbooks"]):
            source = Path(item["path"])
            order = VARIANTS[i % 4:] + VARIANTS[:i % 4]
            for variant in order:
                row, external = call_child(variant, source, root / "build" / str(i) / variant)
                append(raw, {"kind": "build", "workbook_sha256": item["sha256"],
                             "source_bytes": item["bytes"], "variant": variant,
                             "order_position": order.index(variant), "child": row,
                             "process_ns": external})
            print(f"timing workbooks {i+1}/39", flush=True)
        for i, item in enumerate(pop["traces"]):
            source, trace_id = Path(item["snapshot_path"]), item["trace_id"]
            order = VARIANTS[i % 4:] + VARIANTS[:i % 4]
            for variant in order:
                cache = root / "sessions" / str(i) / variant
                session_start = tick()
                for invocation in range(10):
                    row, external = call_child(variant, source, cache, trace_id)
                    cumulative = tick() - session_start
                    expected = "BUILT" if invocation == 0 else "REUSED"
                    if variant in ("P2", "P3") and row["witness"] != expected:
                        raise RuntimeError(f"Invalid session witness {trace_id} {variant} {invocation}: {row['witness']}")
                    append(raw, {"kind": "invocation", "trace_id": trace_id,
                                 "snapshot_sha256": item["snapshot_sha256"], "variant": variant,
                                 "invocation": invocation, "order_position": order.index(variant),
                                 "child": row, "process_ns": external})
                    if invocation + 1 in (1, 2, 3, 5, 10):
                        append(sessions, {"trace_id": trace_id, "snapshot_sha256": item["snapshot_sha256"],
                                          "variant": variant, "N": invocation + 1,
                                          "session_ns": cumulative, "order_position": order.index(variant)})
            print(f"timing traces {i+1}/51", flush=True)


def main() -> None:
    parser = argparse.ArgumentParser()
    sub = parser.add_subparsers(dest="command", required=True)
    child_parser = sub.add_parser("child")
    child_parser.add_argument("--variant", choices=VARIANTS, required=True)
    child_parser.add_argument("--source", required=True)
    child_parser.add_argument("--cache", required=True)
    child_parser.add_argument("--trace-id")
    child_parser.add_argument("--emit-values", action="store_true")
    child_parser.add_argument("--mode", choices=("invoke", "book"), default="invoke")
    sub.add_parser("all")
    args = parser.parse_args()
    if args.command == "child":
        child(args)
        return
    pop = frozen()
    initial_files()
    if not correctness(pop):
        raise RuntimeError("Correctness gate failed; no timing")
    if not invalidation(pop):
        raise RuntimeError("Invalidation gate failed; no timing")
    timing(pop)


if __name__ == "__main__":
    main()
