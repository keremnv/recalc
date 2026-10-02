"""Process-local interposition for the earned narrow openpyxl read surface."""
from __future__ import annotations

import atexit
import json
import os
import sys
import time
from pathlib import Path
from typing import Any

from . import artifact


def tick() -> int:
    return time.perf_counter_ns()


class Runtime:
    def __init__(self, context: dict, bootstrap_started_ns: int):
        self.context = context
        self.bootstrap_started_ns = bootstrap_started_ns
        self.events: list[dict] = []
        self.event_overflow = 0
        self.books: dict[str, Any] = {}
        self.proxies: list[ProxyWorkbook] = []
        self.counts = {"direct_served_loads": 0, "direct_served_reads": 0,
                       "reference_loads": 0, "fallback_loads": 0}
        self.times = {"artifact_load_ns": 0, "reference_parse_ns": 0,
                      "child_source_hash_ns": 0}
        self.original = None
        self.after_install_ns = 0

    def record(self, event: str, **fields: Any) -> None:
        if len(self.events) < 128:
            self.events.append({"event": event, **fields})
        else:
            self.event_overflow += 1

    def reference(self, reason: str, filename: Any, *args, **kwargs):
        self.counts["reference_loads"] += 1
        self.counts["fallback_loads"] += 1
        t = tick()
        try:
            return self.original(filename, *args, **kwargs)
        finally:
            elapsed = tick() - t
            self.times["reference_parse_ns"] += elapsed
            self.record("reference_parse", reason=reason, path=str(filename), duration_ns=elapsed)

    def load(self, filename: Any, *args, **kwargs):
        if not self.context["admitted"]:
            return self.reference("source_rejected_before_interposition", filename, *args, **kwargs)
        supported = (isinstance(filename, (str, os.PathLike)) and not args
                     and not kwargs.get("data_only", False) and not kwargs.get("read_only", False)
                     and not kwargs.get("write_only", False)
                     and Path(filename).suffix.lower() in {".xlsx", ".xlsm"}
                     and set(kwargs) <= {"data_only", "read_only", "write_only"})
        if not supported:
            return self.reference("unsupported_load_mode", filename, *args, **kwargs)
        path = str(Path(filename).resolve())
        entry = self.context["workbooks"].get(path)
        if not entry or entry.get("status") not in {"BUILT", "REUSED"}:
            return self.reference("missing_stale_corrupt_artifact", filename, *args, **kwargs)
        try:
            t = tick()
            digest = artifact.sha_file(Path(path))
            self.times["child_source_hash_ns"] += tick() - t
            if digest != entry["source_sha256"]:
                raise artifact.ArtifactError("source changed after bootstrap freshness verification")
            book = self.books.get(path)
            if book is None:
                t = tick()
                book = artifact.load(Path(entry["artifact_path"]), digest)
                self.times["artifact_load_ns"] += tick() - t
                self.books[path] = book
            proxy = ProxyWorkbook(self, book, path)
            self.proxies.append(proxy)
            self.counts["direct_served_loads"] += 1
            self.record("direct_served_load", path=path, artifact_status=entry["status"],
                        source_sha256=digest)
            return proxy
        except Exception as exc:
            self.record("artifact_runtime_failure", path=path, exception_class=type(exc).__name__,
                        reason=str(exc)[:300])
            return self.reference("runtime_failure", filename, *args, **kwargs)

    def install(self):
        import openpyxl
        self.original = openpyxl.load_workbook
        openpyxl.load_workbook = self.load
        self.after_install_ns = tick()
        self.record("bootstrap", status="INSTALLED", admitted=self.context["admitted"])
        atexit.register(self.finish)

    def finish(self):
        for proxy in self.proxies:
            try:
                proxy.close()
            except Exception:
                pass
        end = tick()
        profile = {"route": "DIRECT_WITH_FALLBACK" if self.counts["fallback_loads"]
                   else "DIRECT_RUNTIME", "times": self.times, "counts": self.counts,
                   "event_overflow": self.event_overflow,
                   "events": self.events,
                   "fallback_reasons": sorted({e["reason"] for e in self.events
                                               if e["event"] == "reference_parse"}),
                   "bootstrap_ns": self.after_install_ns - self.bootstrap_started_ns,
                   "post_bootstrap_to_exit_ns": end - self.after_install_ns,
                   "script_identity_verified": True}
        try:
            Path(self.context["runtime_state"]).write_text(
                json.dumps(profile, sort_keys=True, default=str) + "\n")
        except OSError:
            pass


def _serve_iter_rows(ws, args, kwargs):
    """Direct full-cell row iteration (probe) or fail-closed delegation.

    Mirrors openpyxl 3.1.5 Worksheet.iter_rows semantics exactly for the
    certified rectangle-of-cells shape, including merged ranges (children
    read as None/'n', matching MergedCell observables); anything else
    (values_only, uncertain emptiness, odd arguments) delegates to genuine
    openpyxl with a recorded reason.
    """
    runtime = ws._workbook._runtime
    real = lambda: ws._real_sheet().iter_rows(*args, **kwargs)  # noqa: E731
    if os.environ.get("RECALC_NO_ITERATION_PROBE") == "1":
        return real()
    # Bind through the exact openpyxl signature so arity/keyword errors
    # reproduce identically.
    def _bind(min_row=None, max_row=None, min_col=None, max_col=None,
              values_only=False):
        return min_row, max_row, min_col, max_col, values_only
    try:
        min_row, max_row, min_col, max_col, values_only = _bind(*args, **kwargs)
    except TypeError:
        # Arity/keyword errors reproduce byte-identically via genuine.
        runtime.record("iteration_reference", reason="signature_mismatch",
                       sheet=ws.title)
        return real()
    bounds = (min_row, max_row, min_col, max_col)
    if values_only or not all(b is None or isinstance(b, int)
                              for b in bounds):
        runtime.record("iteration_reference",
                       reason="values_only" if values_only else "dynamic_bounds",
                       sheet=ws.title)
        return real()
    info = ws._sheet.info
    explicit = any([min_col, min_row, max_col, max_row])
    if not explicit and not info.cells:
        # Indistinguishable without new state: truly empty, style-only
        # cells, or merges-only. Reference decides cheaply and correctly.
        runtime.record("iteration_reference", reason="sparse_state_empty",
                       sheet=ws.title)
        return real()
    # Mirror openpyxl's `or`-defaults exactly (falsy -> default).
    min_col = min_col or 1
    min_row = min_row or 1
    max_col = max_col or ws._sheet.max_column
    max_row = max_row or ws._sheet.max_row
    # Merged ranges are served, not escaped (see DEVIATIONS.md D1):
    # MemoryBook.cell maps merged children to (None, 'n'), which is exactly
    # MergedCell-observable state for the certified attribute contract.
    # Rich access on any served cell still escapes per-cell to genuine.
    runtime.record("iteration_direct", sheet=ws.title, min_row=min_row,
                   max_row=max_row, min_col=min_col, max_col=max_col)
    runtime.counts.setdefault("direct_iteration_rows", 0)
    runtime.counts.setdefault("direct_iteration_cells", 0)

    def _gen():
        # Lazily, like openpyxl: negative bounds fail on first advance
        # with the identical message openpyxl's cell() raises.
        if min_row < 1 or min_col < 1:
            raise ValueError("Row or column values must be at least 1")
        for row in range(min_row, max_row + 1):
            cells = tuple(ProxyCell(ws, row, column)
                          for column in range(min_col, max_col + 1))
            runtime.counts["direct_iteration_rows"] += 1
            runtime.counts["direct_iteration_cells"] += len(cells)
            yield cells

    return _gen()


class ProxyWorkbook:
    def __init__(self, runtime: Runtime, book, path: str):
        self._runtime, self._book, self._path, self._real = runtime, book, path, None

    @property
    def sheetnames(self):
        self._runtime.counts["direct_served_reads"] += 1
        return list(self._book.sheetnames)

    def __getitem__(self, name):
        if not isinstance(name, str) or name not in self._book.sheetnames:
            raise KeyError(name)
        self._runtime.counts["direct_served_reads"] += 1
        return ProxyWorksheet(self, name)

    def _real_workbook(self):
        if self._real is None:
            self._real = self._runtime.reference("proxy_operation_escape", self._path, data_only=False)
        return self._real

    def __iter__(self):
        return iter(self._real_workbook())

    def close(self):
        if self._real is not None:
            self._real.close()

    def __getattr__(self, name):
        return getattr(self._real_workbook(), name)


class ProxyWorksheet:
    def __init__(self, workbook: ProxyWorkbook, name: str):
        self._workbook, self.title = workbook, name
        self._sheet = workbook._book[name]

    @property
    def max_row(self):
        self._workbook._runtime.counts["direct_served_reads"] += 1
        return self._sheet.max_row

    @property
    def max_column(self):
        self._workbook._runtime.counts["direct_served_reads"] += 1
        return self._sheet.max_column

    @property
    def dimensions(self):
        self._workbook._runtime.counts["direct_served_reads"] += 1
        return self._sheet.dimensions

    def calculate_dimension(self):
        return self.dimensions

    def _real_sheet(self):
        return self._workbook._real_workbook()[self.title]

    def cell(self, row=1, column=1):
        if not isinstance(row, int) or not isinstance(column, int) or row < 1 or column < 1:
            raise ValueError("Row or column values must be at least 1")
        for r1, c1, r2, c2 in self._sheet.info.merged:
            if r1 <= row <= r2 and c1 <= column <= c2 and (row, column) != (r1, c1):
                runtime = self._workbook._runtime
                runtime.record("merged_child_contact", sheet=self.title, row=row, column=column)
                if runtime.context.get("merged_terminal_certified", False):
                    runtime.record("merged_child_direct", sheet=self.title, row=row, column=column)
                    runtime.counts["direct_served_reads"] += 1
                    return ProxyCell(self, row, column)
                runtime.record("merged_child_reference", sheet=self.title, row=row, column=column)
                return self._real_sheet().cell(row=row, column=column)
        self._workbook._runtime.counts["direct_served_reads"] += 1
        return ProxyCell(self, row, column)

    def __getitem__(self, key):
        from openpyxl.utils.cell import coordinate_to_tuple
        if isinstance(key, str) and ":" not in key:
            try:
                row, column = coordinate_to_tuple(key)
                return self.cell(row=row, column=column)
            except (ValueError, TypeError):
                pass
        return self._real_sheet()[key]

    def __iter__(self):
        return iter(self._real_sheet())

    def iter_rows(self, *args, **kwargs):
        return _serve_iter_rows(self, args, kwargs)

    def __getattr__(self, name):
        return getattr(self._real_sheet(), name)


class ProxyCell:
    def __init__(self, worksheet: ProxyWorksheet, row: int, column: int):
        from openpyxl.utils import get_column_letter
        self._worksheet = worksheet
        self.row, self.column = row, column
        self.coordinate = f"{get_column_letter(column)}{row}"

    @property
    def value(self):
        self._worksheet._workbook._runtime.counts["direct_served_reads"] += 1
        return self._worksheet._sheet.cell(self.row, self.column).value

    @property
    def data_type(self):
        self._worksheet._workbook._runtime.counts["direct_served_reads"] += 1
        return self._worksheet._sheet.cell(self.row, self.column).data_type

    def _real_cell(self):
        return self._worksheet._real_sheet().cell(self.row, self.column)

    def __getattr__(self, name):
        return getattr(self._real_cell(), name)


def install(context: dict, bootstrap_started_ns: int) -> None:
    script = str(Path(sys.argv[0]).resolve())
    if script != context["script"]:
        return
    Runtime(context, bootstrap_started_ns).install()
