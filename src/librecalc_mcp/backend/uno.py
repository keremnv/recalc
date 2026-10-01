from __future__ import annotations

import os
import re
import tempfile
import warnings
import xml.etree.ElementTree as ET
from contextlib import suppress
from copy import copy
from pathlib import Path
from typing import Any

from librecalc_mcp.backend.uno_charts import (
    chart_collection_name_for_id,
    export_charts_to_png,
    inspect_charts_from_document,
    upsert_chart_on_sheet,
)
from librecalc_mcp.backend.xlsx_metadata import repair_docprops
from librecalc_mcp.domain.charts import ChartSpec
from librecalc_mcp.domain.formulas import (
    is_escaped_text,
    is_formula_text,
    normalize_formula_argument_separators,
    translate_a1_formula,
    unescape_text,
)
from librecalc_mcp.domain.grid import column_label, validate_matrix_shape
from librecalc_mcp.domain.models import CalcOperation, CellFormat, Matrix, SheetInfo, WorkbookInfo
from librecalc_mcp.domain.write_policy import preserve_populated as _preserve_populated

_A1_RANGE = re.compile(r"^([A-Z]+)([1-9][0-9]*)(?::([A-Z]+)([1-9][0-9]*))?$", re.IGNORECASE)
_HEX_COLOR = re.compile(r"^#?([0-9A-F]{6})$", re.IGNORECASE)
_FORMAT_FIELDS = {
    "font_color",
    "background_color",
    "background_transparent",
    "font_weight",
}


def _column_index(label: str) -> int:
    value = 0
    for character in label.upper():
        value = value * 26 + ord(character) - ord("A") + 1
    return value - 1


def a1_range_address(cell_range: str) -> tuple[int, int, int, int]:
    """Return a 0-based (start_col, start_row, end_col, end_row) UNO range address."""
    match = _A1_RANGE.fullmatch(cell_range.strip())
    if match is None:
        raise ValueError(f"invalid A1 range: {cell_range}")
    start_column, start_row, end_column, end_row = match.groups()
    end_column = end_column or start_column
    end_row = end_row or start_row
    start_col = _column_index(start_column)
    start_row_number = int(start_row) - 1
    end_col = _column_index(end_column)
    end_row_number = int(end_row) - 1
    if end_col < start_col or end_row_number < start_row_number:
        raise ValueError(f"inverted A1 range: {cell_range}")
    return start_col, start_row_number, end_col, end_row_number


def _pythonize_uno_error(exc: BaseException) -> BaseException:
    """UNO exception objects are not safe to attach a Python traceback to."""
    module = type(exc).__module__ or ""
    if module == "builtins" or module.startswith("librecalc_mcp"):
        return exc
    return RuntimeError(f"{type(exc).__name__}: {exc}")


_EMPTY_CELL_TYPE_NAMES = frozenset({"EMPTY", "CELLTYPE_EMPTY"})


def _cell_is_populated(cell: Any) -> bool:
    """True when UNO CellType is not EMPTY.

    Container pyuno exposes CellType as ``enum.Enum``, so ``int(cell.Type)``
    raises TypeError and aborted preserve-populated fills (Template 06_24).
    """

    cell_type = getattr(cell, "Type", 0)
    if cell_type in (0, None):
        return False
    value = getattr(cell_type, "value", cell_type)
    if value in (0, None):
        return False
    name = str(getattr(cell_type, "name", "")).rsplit(".", 1)[-1].upper()
    if name in _EMPTY_CELL_TYPE_NAMES:
        return False
    if isinstance(value, int) and not isinstance(value, bool):
        return value != 0
    try:
        return int(value) != 0
    except (TypeError, ValueError):
        return True


def _set_cell_value(cell: Any, value: object) -> None:
    if value is None:
        cell.clearContents(31)
        return
    if isinstance(value, bool):
        cell.Value = 1.0 if value else 0.0
        return
    if isinstance(value, (int, float)):
        cell.Value = float(value)
        return
    if is_formula_text(value):
        cell.Formula = normalize_formula_argument_separators(str(value))
        return
    cell.String = unescape_text(str(value))


def _format_color(value: object, *, field: str) -> int:
    if isinstance(value, bool):
        raise ValueError(  # noqa: TRY004 - agent-recoverable tool observation
            f"format.{field} must be a #RRGGBB color or integer"
        )
    if isinstance(value, int) and 0 <= value <= 0xFFFFFF:
        return value
    if isinstance(value, str) and (match := _HEX_COLOR.fullmatch(value.strip())):
        return int(match.group(1), 16)
    raise ValueError(f"format.{field} must be a #RRGGBB color or integer")


def _argb(value: object, *, field: str) -> str:
    return f"FF{_format_color(value, field=field):06X}"


def _a1_addresses(cell_range: str) -> list[str]:
    start_col, start_row, end_col, end_row = a1_range_address(cell_range)
    return [
        f"{column_label(col + 1)}{row + 1}"
        for row in range(start_row, end_row + 1)
        for col in range(start_col, end_col + 1)
    ]


def _apply_format_fields(target: Any, cell_format: CellFormat) -> None:
    if "font_color" in cell_format:
        target.CharColor = _format_color(cell_format["font_color"], field="font_color")
    if "background_color" in cell_format:
        target.CellBackColor = _format_color(
            cell_format["background_color"], field="background_color"
        )
        if "background_transparent" not in cell_format:
            target.IsCellBackgroundTransparent = False
    if "background_transparent" in cell_format:
        transparent = cell_format["background_transparent"]
        if not isinstance(transparent, bool):
            raise ValueError("format.background_transparent must be boolean")
        target.IsCellBackgroundTransparent = transparent
    if "font_weight" in cell_format:
        weight = cell_format["font_weight"]
        if isinstance(weight, bool) or not isinstance(weight, (int, float)):
            raise ValueError("format.font_weight must be numeric")
        target.CharWeight = float(weight)


def _apply_format(target: Any, cell_format: CellFormat) -> int:
    unknown = sorted(set(cell_format) - _FORMAT_FIELDS)
    if unknown:
        raise ValueError(f"unsupported format fields: {', '.join(unknown)}")
    if not cell_format:
        raise ValueError("set_format requires at least one format field")

    # Per-cell, matching the container ScCellObj hole: a range-level CharColor set can
    # mutate a shared xf that a later set_format reverts on export.
    address = target.RangeAddress
    width = address.EndColumn - address.StartColumn + 1
    height = address.EndRow - address.StartRow + 1
    for row_offset in range(height):
        for column_offset in range(width):
            _apply_format_fields(target.getCellByPosition(column_offset, row_offset), cell_format)
    return width * height


def _theme_spec(color: Any) -> tuple[int, float] | None:
    if color is None or getattr(color, "type", None) != "theme":
        return None
    theme = color.theme
    if not isinstance(theme, int):
        return None
    tint = color.tint
    return theme, float(tint) if tint else 0.0


def _persist_xlsx_formats(
    path: str,
    patches: list[tuple[str, str, CellFormat]],
    *,
    source_path: str | None = None,
    restore_unpatched_themes: bool = True,
) -> None:
    """Write set_format RGB and put unpainted Excel theme fonts back.

    UNO CharColor can be live-true after save while openpyxl still sees the input RGB.
    Container LibreOffice also materializes workbook themes as RGB that is not the
    evaluator's Office theme map (03_04 LBO!C6 theme 8 → 2F5597 vs mapped 44749F).
    """
    suffix = Path(path).suffix.lower()
    if suffix not in {".xlsx", ".xlsm"}:
        return
    restore_from = (
        source_path
        if (
            restore_unpatched_themes
            and source_path
            and Path(source_path).suffix.lower() in {".xlsx", ".xlsm"}
            and os.path.abspath(source_path) != os.path.abspath(path)
        )
        else None
    )
    if not patches and restore_from is None:
        return
    try:
        import openpyxl
        from openpyxl.styles import Color, PatternFill
    except ImportError:
        return

    patched: set[tuple[str, str]] = {
        (sheet_name, address)
        for sheet_name, cell_range, _cell_format in patches
        for address in _a1_addresses(cell_range)
    }
    with tempfile.TemporaryDirectory(prefix="librecalc-docprops-") as temporary_directory:
        scratch = Path(temporary_directory)

        def load_tolerant(workbook_path: str, **kwargs: Any) -> Any:
            try:
                return openpyxl.load_workbook(workbook_path, **kwargs)
            except ET.ParseError:
                repaired = repair_docprops(Path(workbook_path), scratch)
                if repaired is None:
                    raise
                return openpyxl.load_workbook(repaired, **kwargs)

        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            workbook = load_tolerant(path, keep_vba=suffix == ".xlsm")
            source = load_tolerant(restore_from, read_only=False) if restore_from else None
        try:
            for sheet_name, cell_range, cell_format in patches:
                worksheet = workbook[sheet_name]
                for address in _a1_addresses(cell_range):
                    cell = worksheet[address]
                    if "font_color" in cell_format or "font_weight" in cell_format:
                        font = copy(cell.font)
                        if "font_color" in cell_format:
                            font.color = Color(
                                rgb=_argb(cell_format["font_color"], field="font_color")
                            )
                        if "font_weight" in cell_format:
                            weight = cell_format["font_weight"]
                            if isinstance(weight, bool) or not isinstance(weight, (int, float)):
                                raise ValueError("format.font_weight must be numeric")
                            font.bold = float(weight) >= 150
                        cell.font = font
                    if cell_format.get("background_transparent") is True:
                        cell.fill = PatternFill(fill_type=None)
                    elif "background_color" in cell_format:
                        fill_color = _argb(
                            cell_format["background_color"], field="background_color"
                        )
                        cell.fill = PatternFill(
                            fill_type="solid",
                            fgColor=fill_color,
                            bgColor=fill_color,
                        )
            if source is not None:
                for source_sheet in source.worksheets:
                    if source_sheet.title not in workbook.sheetnames:
                        continue
                    destination_sheet = workbook[source_sheet.title]
                    for row in source_sheet.iter_rows():
                        for source_cell in row:
                            if (
                                source_cell.value is None
                                and _theme_spec(source_cell.font.color) is None
                            ):
                                continue
                            address = source_cell.coordinate
                            if (source_sheet.title, address) in patched:
                                continue
                            spec = _theme_spec(source_cell.font.color)
                            if spec is None:
                                continue
                            theme, tint = spec
                            dest_cell = destination_sheet[address]
                            font = copy(dest_cell.font)
                            font.color = Color(theme=theme, tint=tint)
                            dest_cell.font = font
            workbook.save(path)
        finally:
            workbook.close()
            if source is not None:
                source.close()


def _apply_values(sheet: Any, cell_range: str, values: Matrix) -> int:
    start_col, start_row, end_col, end_row = a1_range_address(cell_range)
    validate_matrix_shape(cell_range, values)
    data = tuple(tuple(row) for row in values)
    skip_populated = _preserve_populated()
    if start_col == end_col and start_row == end_row:
        if len(data) != 1 or len(data[0]) != 1:
            raise ValueError("single-cell range requires a 1x1 values matrix")
        cell = sheet.getCellByPosition(start_col, start_row)
        if skip_populated and _cell_is_populated(cell):
            return 1
        # Container LibreOffice rejects setDataArray on ScCellObj (cellsuno.cxx:5014).
        _set_cell_value(cell, data[0][0])
        return 0
    cell_by_cell = skip_populated or any(
        is_formula_text(cell) or is_escaped_text(cell) for row in data for cell in row
    )
    if cell_by_cell:
        # Mixed/formula-bearing matrix, or the preserve-populated guard: setDataArray
        # would store formulas as text and would overwrite occupied cells.
        skipped = 0
        for row_offset, row in enumerate(data):
            for column_offset, cell_value in enumerate(row):
                cell = sheet.getCellByPosition(start_col + column_offset, start_row + row_offset)
                if skip_populated and _cell_is_populated(cell):
                    skipped += 1
                    continue
                _set_cell_value(cell, cell_value)
        return skipped
    sheet.getCellRangeByPosition(start_col, start_row, end_col, end_row).setDataArray(data)
    return 0


class _DocumentContext:
    def __init__(self, backend: UnoCalcBackend, path: str | None) -> None:
        self.backend = backend
        self.path = path
        self.doc: Any | None = None
        self.owned = False

    def __enter__(self) -> Any:
        expanded_path = os.path.abspath(os.path.expanduser(self.path)) if self.path else None
        if expanded_path is not None and not os.path.isfile(expanded_path):
            raise FileNotFoundError(f"Spreadsheet path does not exist: {expanded_path}")
        self.backend._connect()
        assert self.backend._desktop is not None
        self.owned = self.path is not None
        if expanded_path:
            uno = self.backend._uno()
            file_url = uno.systemPathToFileUrl(expanded_path)
            self.doc = self.backend._desktop.loadComponentFromURL(file_url, "_blank", 0, ())
        else:
            self.doc = self.backend._desktop.getCurrentComponent()
        if self.doc is None or not self.doc.supportsService(
            "com.sun.star.sheet.SpreadsheetDocument"
        ):
            if self.owned and self.doc is not None:
                with suppress(Exception):
                    self.doc.close(True)
            self.doc = None
            raise RuntimeError("No active Calc spreadsheet document. Open one or provide path=...")
        return self.doc

    def __exit__(
        self,
        _exc_type: type[BaseException] | None,
        exc: BaseException | None,
        _tb: object,
    ) -> bool:
        if self.owned and self.doc is not None:
            with suppress(Exception):
                self.doc.close(True)
        if exc is None:
            return False
        converted = _pythonize_uno_error(exc)
        if converted is exc:
            return False
        raise converted from None


class UnoUnavailable(RuntimeError):
    pass


class UnoCalcBackend:
    def __init__(self, host: str = "localhost", port: int = 2021) -> None:
        self.host = host
        self.port = port
        self._ctx: Any | None = None
        self._desktop: Any | None = None

    def _uno(self):
        try:
            import uno  # type: ignore
        except ImportError as exc:
            raise UnoUnavailable(
                "Python cannot import 'uno'. Use a Python environment with LibreOffice UNO bindings."
            ) from exc
        return uno

    def _connect(self) -> None:
        if self._desktop is not None:
            return
        uno = self._uno()
        local_ctx = uno.getComponentContext()
        resolver = local_ctx.ServiceManager.createInstanceWithContext(
            "com.sun.star.bridge.UnoUrlResolver", local_ctx
        )
        url = f"uno:socket,host={self.host},port={self.port};urp;StarOffice.ComponentContext"
        self._ctx = resolver.resolve(url)
        self._desktop = self._ctx.ServiceManager.createInstanceWithContext(
            "com.sun.star.frame.Desktop", self._ctx
        )

    def health(self) -> dict[str, object]:
        try:
            self._connect()
            return {"ok": True, "backend": "uno", "host": self.host, "port": self.port}
        except Exception as exc:
            return {
                "ok": False,
                "backend": "uno",
                "host": self.host,
                "port": self.port,
                "error": f"{type(exc).__name__}: {exc}",
            }

    def _document(self, path: str | None = None) -> _DocumentContext:
        # Class-based: @contextmanager throw() attaches a traceback to UNO exceptions
        # and pyuno then masks the real error as "Couldn't convert <traceback object>".
        return _DocumentContext(self, path)

    def _property(self, name: str, value: object) -> Any:
        prop = self._uno().createUnoStruct("com.sun.star.beans.PropertyValue")
        prop.Name = name
        prop.Value = value
        return prop

    @staticmethod
    def _filter_name(path: str) -> str:
        suffix = Path(path).suffix.lower()
        filters = {
            ".ods": "calc8",
            ".xlsx": "Calc MS Excel 2007 XML",
            ".xlsm": "Calc MS Excel 2007 VBA XML",
            ".xls": "MS Excel 97",
        }
        try:
            return filters[suffix]
        except KeyError as exc:
            raise ValueError(
                f"Unsupported spreadsheet output extension: {suffix or '<none>'}"
            ) from exc

    def _save_document(
        self,
        doc: Any,
        path: str | None,
        output_path: str | None,
    ) -> str | None:
        target_path = output_path or path
        if target_path is None:
            return None

        absolute_target = os.path.abspath(os.path.expanduser(target_path))
        source_path = os.path.abspath(os.path.expanduser(path)) if path else None
        os.makedirs(os.path.dirname(absolute_target), exist_ok=True)

        if source_path == absolute_target:
            doc.store()
        else:
            target_url = self._uno().systemPathToFileUrl(absolute_target)
            properties = (
                self._property("FilterName", self._filter_name(absolute_target)),
                self._property("Overwrite", True),
            )
            doc.storeToURL(target_url, properties)
        return absolute_target

    @staticmethod
    def _used_range(sheet: Any) -> str | None:
        cursor = sheet.createCursor()
        cursor.gotoEndOfUsedArea(True)
        address = cursor.RangeAddress

        # UNO columns/rows are zero-based; implement small A1 converter locally.
        def col_name(n: int) -> str:
            n += 1
            out = ""
            while n:
                n, rem = divmod(n - 1, 26)
                out = chr(65 + rem) + out
            return out

        return f"A1:{col_name(address.EndColumn)}{address.EndRow + 1}"

    @staticmethod
    def _sheet_range(sheet: Any, cell_range: str) -> Any:
        # getCellRangeByName("C9") returns ScCellObj. Container LibreOffice rejects
        # setDataArray on a cell object (cellsuno.cxx:5014); a 1x1 range from
        # getCellRangeByPosition does not.
        start_col, start_row, end_col, end_row = a1_range_address(cell_range)
        return sheet.getCellRangeByPosition(start_col, start_row, end_col, end_row)

    def inspect_workbook(self, path: str | None = None) -> WorkbookInfo:
        with self._document(path) as doc:
            sheets = doc.Sheets
            info: list[SheetInfo] = []
            for name in sheets.ElementNames:
                sheet = sheets.getByName(name)
                info.append(SheetInfo(name=name, used_range=self._used_range(sheet)))
            title = getattr(doc, "Title", None) or None
            url = getattr(doc, "URL", None) or None
            return WorkbookInfo(title=title, url=url, sheets=info)

    def inspect_charts(self, path: str | None = None) -> list[dict[str, object]]:
        with self._document(path) as doc:
            return inspect_charts_from_document(doc)

    def export_charts_png(
        self,
        path: str,
        output_dir: str,
        *,
        file_prefix: str = "chart",
    ) -> list[str]:
        with self._document(path) as doc:
            return export_charts_to_png(
                doc,
                output_dir=output_dir,
                file_prefix=file_prefix,
                uno_module=self._uno(),
                context=self._ctx,
            )

    @staticmethod
    def _read_target(target: Any, *, include_errors: bool) -> dict[str, Matrix]:
        values = [list(row) for row in target.getDataArray()]
        formulas = [list(row) for row in target.getFormulaArray()]
        errors: Matrix = [[None for _ in row] for row in formulas]
        if include_errors:
            # Cell-by-cell UNO access is expensive. Only formulas can produce calculated
            # spreadsheet errors, so constants and blanks never need an Error query.
            for row_offset, row in enumerate(formulas):
                for column_offset, formula in enumerate(row):
                    if not str(formula).startswith("="):
                        continue
                    cell = target.getCellByPosition(column_offset, row_offset)
                    if cell.Error:
                        errors[row_offset][column_offset] = cell.String
        return {"values": values, "formulas": formulas, "errors": errors}

    def read_range(self, sheet: str, cell_range: str, path: str | None = None) -> dict[str, Matrix]:
        with self._document(path) as doc:
            target = self._sheet_range(doc.Sheets.getByName(sheet), cell_range)
            return self._read_target(target, include_errors=True)

    def read_ranges(
        self,
        ranges: list[tuple[str, str]],
        path: str | None = None,
        *,
        include_errors: bool = True,
    ) -> list[dict[str, Matrix]]:
        """Read several semantic ranges while opening the workbook only once."""
        with self._document(path) as doc:
            return [
                self._read_target(
                    self._sheet_range(doc.Sheets.getByName(sheet), cell_range),
                    include_errors=include_errors,
                )
                for sheet, cell_range in ranges
            ]

    @staticmethod
    def _rgb_color(value: int) -> str | None:
        if value < 0:
            return None
        return f"#{value & 0xFFFFFF:06X}"

    def read_formats(
        self,
        cells: list[tuple[str, str]],
        path: str | None = None,
    ) -> list[CellFormat]:
        """Read compact formatting fingerprints for selected cells."""
        with self._document(path) as doc:
            fingerprints: list[CellFormat] = []
            for sheet, address in cells:
                cell = doc.Sheets.getByName(sheet).getCellRangeByName(address)
                fingerprints.append(
                    {
                        "font_color": self._rgb_color(int(cell.CharColor)),
                        "background_color": self._rgb_color(int(cell.CellBackColor)),
                        "background_transparent": bool(cell.IsCellBackgroundTransparent),
                        "font_weight": float(cell.CharWeight),
                    }
                )
            return fingerprints

    def write_range(
        self,
        sheet: str,
        cell_range: str,
        values: Matrix,
        path: str | None = None,
        output_path: str | None = None,
    ) -> dict[str, object]:
        with self._document(path) as doc:
            cells_skipped = _apply_values(doc.Sheets.getByName(sheet), cell_range, values)
            doc.calculateAll()
            saved_to = self._save_document(doc, path, output_path)
        if saved_to is not None:
            _persist_xlsx_formats(saved_to, [], source_path=path)
        return {
            "ok": True,
            "sheet": sheet,
            "range": cell_range,
            "rows": len(values),
            "cells_skipped": cells_skipped,
            "saved_to": saved_to,
        }

    def execute_program(
        self,
        operations: list[CalcOperation],
        path: str | None = None,
        output_path: str | None = None,
    ) -> dict[str, object]:
        with self._document(path) as doc:
            results: list[dict[str, object]] = []
            format_patches: list[tuple[str, str, CellFormat]] = []
            geometry_changed = False
            for operation in operations:
                if operation.op == "create_sheet":
                    if not operation.name:
                        raise ValueError("create_sheet requires name")
                    index = operation.index if operation.index is not None else doc.Sheets.Count
                    doc.Sheets.insertNewByName(operation.name, index)
                    results.append({"op": operation.op, "ok": True, "name": operation.name})
                    continue

                if operation.op in {"insert_row", "delete_row"}:
                    if not operation.sheet or operation.index is None:
                        raise ValueError(f"{operation.op} requires sheet and 1-based index")
                    count = 1 if operation.count is None else operation.count
                    if operation.index < 1 or count < 1:
                        raise ValueError(f"{operation.op} requires index >= 1 and count >= 1")
                    geometry_changed = True
                    rows = doc.Sheets.getByName(operation.sheet).Rows
                    uno_index = operation.index - 1
                    if operation.op == "insert_row":
                        rows.insertByIndex(uno_index, count)
                    else:
                        rows.removeByIndex(uno_index, count)
                    results.append(
                        {
                            "op": operation.op,
                            "ok": True,
                            "sheet": operation.sheet,
                            "index": operation.index,
                            "count": count,
                        }
                    )
                    continue

                if operation.op == "write_range":
                    if not operation.sheet or not operation.range:
                        raise ValueError("write_range requires sheet and range")
                    if operation.values is None:
                        raise ValueError("write_range requires values")
                    cells_skipped = _apply_values(
                        doc.Sheets.getByName(operation.sheet),
                        operation.range,
                        operation.values,
                    )
                    results.append({"op": operation.op, "ok": True, "cells_skipped": cells_skipped})
                    continue

                if operation.op == "upsert_chart":
                    if operation.chart is None:
                        raise ValueError("upsert_chart requires chart")
                    spec = ChartSpec.from_dict(operation.chart)
                    sheet = doc.Sheets.getByName(spec.sheet)
                    result = upsert_chart_on_sheet(
                        sheet,
                        spec,
                        doc=doc,
                        uno_module=self._uno(),
                        context=self._ctx,
                    )
                    results.append({"op": operation.op, **result})
                    continue

                if operation.op == "delete_chart":
                    if not operation.name:
                        raise ValueError("delete_chart requires name (chart id)")
                    deleted = False
                    for sheet_name in doc.Sheets.ElementNames:
                        sheet = doc.Sheets.getByName(sheet_name)
                        charts = sheet.getCharts()
                        stored_name = chart_collection_name_for_id(sheet, operation.name)
                        if stored_name is not None:
                            charts.removeByName(stored_name)
                            deleted = True
                    results.append({"op": operation.op, "ok": True, "deleted": deleted})
                    continue

                if not operation.sheet or not operation.range:
                    raise ValueError(f"{operation.op} requires sheet and range")
                target = self._sheet_range(doc.Sheets.getByName(operation.sheet), operation.range)

                cells_skipped = 0
                if operation.op == "set_formula":
                    if operation.formula is None:
                        raise ValueError("set_formula requires formula")
                    # v0 deliberately restricts this op to a single-cell range.
                    cell = target.getCellByPosition(0, 0)
                    if _preserve_populated() and _cell_is_populated(cell):
                        cells_skipped = 1
                    else:
                        cell.Formula = normalize_formula_argument_separators(operation.formula)
                elif operation.op == "fill_formula":
                    if operation.formula is None:
                        raise ValueError("fill_formula requires formula")
                    address = target.RangeAddress
                    width = address.EndColumn - address.StartColumn + 1
                    height = address.EndRow - address.StartRow + 1
                    for row_offset in range(height):
                        for column_offset in range(width):
                            cell = target.getCellByPosition(column_offset, row_offset)
                            if _preserve_populated() and _cell_is_populated(cell):
                                cells_skipped += 1
                                continue
                            translated = translate_a1_formula(
                                operation.formula,
                                column_offset=column_offset,
                                row_offset=row_offset,
                            )
                            cell.Formula = normalize_formula_argument_separators(translated)
                elif operation.op == "set_format":
                    if operation.cell_format is None:
                        raise ValueError("set_format requires format")
                    if not operation.sheet or not operation.range:
                        raise ValueError("set_format requires sheet and range")
                    cells_formatted = _apply_format(target, operation.cell_format)
                    format_patches.append(
                        (operation.sheet, operation.range, dict(operation.cell_format))
                    )
                    results.append(
                        {
                            "op": operation.op,
                            "ok": True,
                            "sheet": operation.sheet,
                            "range": operation.range,
                            "cells_formatted": cells_formatted,
                        }
                    )
                    continue
                elif operation.op == "clear_range":
                    # Values, dates, strings, annotations and formulas; formatting survives.
                    target.clearContents(31)
                else:
                    raise ValueError(f"Unsupported operation: {operation.op}")
                result = {"op": operation.op, "ok": True}
                if operation.op in {"set_formula", "fill_formula"}:
                    result["cells_skipped"] = cells_skipped
                results.append(result)

            doc.calculateAll()
            saved_to = self._save_document(doc, path, output_path)
        if saved_to is not None:
            # After close: a live LibreOffice handle can overwrite an in-place OOXML patch.
            _persist_xlsx_formats(
                saved_to,
                format_patches,
                source_path=path,
                restore_unpatched_themes=not geometry_changed,
            )
        return {"ok": True, "operations": results, "saved_to": saved_to}
