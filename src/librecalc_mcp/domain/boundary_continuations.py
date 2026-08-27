"""Gold-blind right-edge formula-continuation candidates.

A blank immediately after the rightmost exact horizontal translation run in a row,
still inside the used rectangle. The date+style arm is the only shipped inspect
payload: Financial Model census was 2/2, 100% precision (01_02 and 01_03 M3).
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from typing import Any
from zipfile import BadZipFile

from librecalc_mcp.domain.formulas import translate_a1_formula


def _is_formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=")


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


@dataclass(frozen=True)
class BoundaryContinuation:
    """One inferable formula immediately beyond a horizontal translation run."""

    sheet: str
    address: str
    inferred_formula: str
    run_start: str
    run_end: str
    run_length: int
    style_match: bool
    date_like: bool
    named_sheet: bool

    def rank_key(self) -> tuple[int, int, int, int, str, str]:
        return (
            -int(self.date_like),
            -int(self.style_match),
            -int(self.named_sheet),
            -self.run_length,
            self.sheet.casefold(),
            self.address,
        )

    def inspect_payload(self) -> dict[str, Any]:
        return {
            "sheet": self.sheet,
            "address": self.address,
            "inferred_formula": self.inferred_formula,
            "run": f"{self.run_start}:{self.run_end}",
            "run_length": self.run_length,
        }


def _date_like(raw_cell: Any, value_cell: Any) -> bool:
    from openpyxl.styles.numbers import is_date_format

    value = value_cell.value
    return isinstance(value, (date, datetime)) or is_date_format(raw_cell.number_format)


def enumerate_boundary_continuations(
    raw_workbook: Any,
    value_workbook: Any,
    *,
    named_sheets: set[str] | None = None,
    minimum_run: int = 3,
) -> list[BoundaryContinuation]:
    """Return gold-blind formula continuations inside each sheet's used rectangle."""

    from openpyxl.utils.cell import get_column_letter, range_boundaries

    if minimum_run < 2:
        raise ValueError("minimum_run must be at least 2")
    named = named_sheets or set()
    candidates: list[BoundaryContinuation] = []
    for raw_sheet in raw_workbook.worksheets:
        value_sheet = value_workbook[raw_sheet.title]
        min_column, min_row, max_column, max_row = range_boundaries(raw_sheet.calculate_dimension())
        if max_column <= min_column:
            continue
        for row in range(min_row, max_row + 1):
            formula_columns = [
                column
                for column in range(min_column, max_column + 1)
                if _is_formula(raw_sheet.cell(row, column).value)
            ]
            if not formula_columns:
                continue
            run_end_column = max(formula_columns)
            candidate_column = run_end_column + 1
            if candidate_column > max_column:
                continue
            candidate_cell = raw_sheet.cell(row, candidate_column)
            if not _is_blank(candidate_cell.value):
                continue

            run_start_column = run_end_column
            while run_start_column > min_column:
                previous = raw_sheet.cell(row, run_start_column - 1).value
                current = raw_sheet.cell(row, run_start_column).value
                if not (_is_formula(previous) and _is_formula(current)):
                    break
                try:
                    expected = translate_a1_formula(previous, column_offset=1, row_offset=0)
                except ValueError:
                    break
                if expected != current:
                    break
                run_start_column -= 1
            run_length = run_end_column - run_start_column + 1
            if run_length < minimum_run:
                continue
            last_cell = raw_sheet.cell(row, run_end_column)
            try:
                inferred = translate_a1_formula(str(last_cell.value), column_offset=1, row_offset=0)
            except ValueError:
                continue
            address = f"{get_column_letter(candidate_column)}{row}"
            candidates.append(
                BoundaryContinuation(
                    sheet=raw_sheet.title,
                    address=address,
                    inferred_formula=inferred,
                    run_start=f"{get_column_letter(run_start_column)}{row}",
                    run_end=f"{get_column_letter(run_end_column)}{row}",
                    run_length=run_length,
                    style_match=(
                        candidate_cell.style_id != 0
                        and candidate_cell.style_id == last_cell.style_id
                    ),
                    date_like=_date_like(last_cell, value_sheet.cell(row, run_end_column)),
                    named_sheet=raw_sheet.title in named,
                )
            )
    return sorted(candidates, key=BoundaryContinuation.rank_key)


def select_arm(candidates: list[BoundaryContinuation], arm: str) -> list[BoundaryContinuation]:
    if arm == "run3":
        return candidates
    if arm == "style":
        return [candidate for candidate in candidates if candidate.style_match]
    if arm == "date":
        return [candidate for candidate in candidates if candidate.date_like]
    if arm == "date+style":
        return [
            candidate for candidate in candidates if candidate.date_like and candidate.style_match
        ]
    if arm.startswith("top-"):
        return candidates[: int(arm.removeprefix("top-"))]
    raise ValueError(f"Unknown continuation arm: {arm}")


def payload_from_xlsx(path: str) -> dict[str, Any] | None:
    """Inspect payload for formula-patterns-v1. None if openpyxl cannot read the path."""

    try:
        import openpyxl
    except ImportError:
        return None
    try:
        raw = openpyxl.load_workbook(path, data_only=False)
        values = openpyxl.load_workbook(path, data_only=True)
    except (OSError, ValueError, BadZipFile, openpyxl.utils.exceptions.InvalidFileException):
        return None
    try:
        selected = select_arm(enumerate_boundary_continuations(raw, values), "date+style")
    finally:
        raw.close()
        values.close()
    if not selected:
        return None
    return {
        "candidates": [candidate.inspect_payload() for candidate in selected],
        "note": (
            "Heuristic one-cell formula continuations immediately right of a date run "
            "that shares number format with the blank. Not requirements."
        ),
    }


ARMS = ("run3", "style", "date", "date+style", "top-5", "top-10")

__all__ = [
    "ARMS",
    "BoundaryContinuation",
    "enumerate_boundary_continuations",
    "payload_from_xlsx",
    "select_arm",
]
